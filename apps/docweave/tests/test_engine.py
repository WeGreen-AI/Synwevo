"""Engine unit tests: value fidelity, formats, references, DOCX runs, XLSX rows, PDF redaction, naming."""
import datetime as dt
import io

import docx
import openpyxl
import pymupdf
import pytest
from docx.shared import RGBColor

from docweave.engine import docx_core as dc
from docweave.engine import pdf_engine as pe
from docweave.engine import tabular, transforms
from docweave.engine import values as V
from docweave.engine.detect import detect
from docweave.engine.numfmt import format_value
from docweave.engine.ooxml import Package
from docweave.engine.refshift import Shift, shift_formula, shift_sqref
from docweave.engine.xlsx_doc import XlsxDoc
from docweave.jobs import dedupe, render_name
from docweave.services import safe_output_name


# ---------------------------------------------------------------- values & number formats
@pytest.mark.parametrize("value,fmt,expect", [
    (1234.5, "#,##0.00", "1,234.50"), (123, "000000", "000123"), (0.256, "0.0%", "25.6%"),
    (-1234.5, "¥#,##0.00", "-¥1,234.50"), (dt.datetime(2024, 3, 5), 'yyyy"年"m"月"d"日"', "2024年3月5日"),
    (0, '#,##0.00;-#,##0.00;"-"', "-"), (-5, "0.00;(0.00)", "(5.00)"), (12345678901234567, "0", "12345678901234567"),
    (dt.time(15, 4), "h:mm AM/PM", "3:04 PM"), (1500000, '#,##0,"K"', "1,500K"),
])
def test_number_formats(value, fmt, expect):
    assert format_value(value, fmt)[0] == expect


def test_value_fidelity_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["id", "long", "zero", "empty", "bool", "money", "date", "padded"])
    ws.append(["000123", "123456789012345678", 0, None, True, 1234.5, dt.date(2024, 2, 29), 42])
    ws["F2"].number_format = "#,##0.00"
    ws["G2"].number_format = "yyyy-mm-dd"
    ws["H2"].number_format = "00000"
    b = io.BytesIO()
    wb.save(b)
    sh = tabular.read_xlsx(b.getvalue()).sheets[0]
    row = sh.rows[1]
    assert row[0].t == V.STR and row[0].text == "000123"          # leading zeros kept (text)
    assert row[1].text == "123456789012345678"                    # long number stored as text kept exactly
    assert row[2].t == V.INT and row[2].text == "0" and not row[2].is_empty   # 0 is not empty
    assert sh.cell(1, 3).is_empty                                   # empty is empty
    assert row[4].t == V.BOOL and row[4].v is True
    assert row[5].text == "1,234.50" and row[5].v == "1234.5"     # display vs exact value
    assert row[6].t == V.DATE and row[6].v == "2024-02-29"
    assert row[7].text == "00042" and row[7].v == "42"


def test_csv_gbk_quoted_newlines(fixtures):
    wb = tabular.read_csv((fixtures / "测试_GBK编码.csv").read_bytes())
    assert wb.info["encoding"] == "gb18030"
    rows = wb.sheets[0].rows
    assert len(rows) == 3                                       # quoted newline does NOT create a new row
    assert rows[1][0].text == "00017"                           # leading zero kept
    assert rows[1][2].text == "北京市测试区\n第二行地址"
    assert rows[2][2].text == "上海市, 测试路 1 号"
    u = tabular.read_csv((fixtures / "测试_UTF8.csv").read_bytes())
    assert u.info["encoding"] == "utf-8-sig" and u.sheets[0].rows[1][1].text == "测试张三"


def test_transforms_are_controlled():
    v = V.Value(V.DATE, "2024-03-05", "2024-03-05")
    assert transforms.apply(v, [{"op": "date_format", "pattern": "yyyy年m月d日"}]).text == "2024年3月5日"
    assert transforms.apply(V.text(" 12.345 "), [{"op": "trim"}, {"op": "decimals", "n": 2}]).text == "12.35"
    assert transforms.apply(V.empty(), [{"op": "default", "text": "—"}]).text == "—"
    assert transforms.validate([{"op": "exec", "code": "import os"}])  # unknown op rejected
    with pytest.raises(transforms.ConflictError):
        transforms.aggregate([V.text("A"), V.text("B")], "unique")
    assert transforms.aggregate([V.text("1.10"), V.text("2.2")], "sum").text == "3.3"
    assert transforms.combine([V.text("a"), V.text("b")], "newline").text == "a\nb"


# ---------------------------------------------------------------- detection
def test_detect_formats(fixtures):
    assert detect((fixtures / "测试_客户100.xlsx").read_bytes(), "x.xlsx")["format"] == "xlsx"
    info = detect((fixtures / "测试_客户100.xlsx").read_bytes(), "renamed.xls")
    assert info["format"] == "xlsx" and "实际内容是 XLSX" in info["notes"][0]
    with pytest.raises(ValueError):
        detect((fixtures / "测试_损坏文件.xlsx").read_bytes(), "测试_损坏文件.xlsx")
    with pytest.raises(ValueError):
        detect((fixtures / "测试_不支持.bin").read_bytes(), "x.bin")
    assert detect((fixtures / "测试_扫描单据.png").read_bytes(), "a.png")["family"] == "image"


# ---------------------------------------------------------------- references
def test_refshift_rules():
    sh = Shift(6, 3, expand_row=5)          # 3 rows inserted after template row 5
    assert shift_formula("SUM(E5:E5)", sh, "S", "S") == "SUM(E5:E8)"
    assert shift_formula("E6+$A$10", sh, "S", "S") == "E9+$A$13"
    assert shift_formula("'Other'!A10", sh, "S", "S") == "'Other'!A10"      # other sheet untouched
    assert shift_formula("S!A10", sh, "S", "Other") == "S!A13"
    assert shift_formula("A1:A4", sh, "S", "S") == "A1:A4"
    assert shift_sqref("B5:C5 D9", sh) == "B5:C8 D12"
    d = Shift(5, -1)
    assert shift_formula("E5", d, "S", "S") == "#REF!"
    assert shift_formula("SUM(E4:E6)", d, "S", "S") == "SUM(E4:E5)"


# ---------------------------------------------------------------- DOCX
def _doc_with_runs():
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("客户：").bold = True
    r = p.add_run("张")
    r.font.color.rgb = RGBColor(0xFF, 0, 0)
    p.add_run("三丰").italic = True
    p.add_run("，张三丰再次出现。")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_docx_cross_run_field_keeps_text_and_style():
    pkg = Package(_doc_with_runs())
    p = dc.get_block(dc.part_container(pkg, pkg.main_part()), "0")
    text = dc.para_text(p)
    s = text.index("张三丰")
    dc.wrap_range(p, s, s + 3, "dw:f1", "姓名")
    assert dc.para_text(p) == text                     # marking never changes text
    with pytest.raises(dc.DocxError):
        dc.wrap_range(p, s + 1, s + 4, "dw:f2", "x")   # overlap rejected
    for _, sdt in list(dc.iter_field_sdts(p)):
        dc.fill_sdt(sdt, "李四")
    out = docx.Document(io.BytesIO(pkg.save())).paragraphs[0]
    assert out.text == "客户：李四，张三丰再次出现。"   # second occurrence untouched
    assert out.runs[1].font.color.rgb == RGBColor(0xFF, 0, 0)  # style of the field's first char


# ---------------------------------------------------------------- XLSX
def test_xlsx_repeat_rows_shift_everything():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["A1"] = "标题"
    ws.merge_cells("A1:E1")
    ws["A4"], ws["B4"] = "品名", "数量"
    ws["A5"] = "样例"
    ws.merge_cells("B5:C5")
    ws["D5"] = 1
    ws["E5"] = "=D5*10"
    ws["E6"] = "=SUM(E5:E5)"
    ws2 = wb.create_sheet("汇总")
    ws2["A1"] = "=订单!E6"
    b = io.BytesIO()
    wb.save(b)
    x = XlsxDoc(b.getvalue())
    x.insert_copies_of_row("订单", 5, 2)
    for i, (n, q) in enumerate([("甲", "2"), ("乙", "3"), ("丙", "4")]):
        x.set_cell("订单", 4 + i, 0, V.text(n))
        x.set_cell("订单", 4 + i, 3, V.Value(V.DEC, q, q))
    out = openpyxl.load_workbook(io.BytesIO(x.finalize()))
    w = out["订单"]
    assert [w.cell(r, 1).value for r in (5, 6, 7)] == ["甲", "乙", "丙"]
    assert [w.cell(r, 5).value for r in (5, 6, 7)] == ["=D5*10", "=D6*10", "=D7*10"]
    assert w["E8"].value == "=SUM(E5:E7)"
    assert out["汇总"]["A1"].value == "=订单!E8"
    assert {str(m) for m in w.merged_cells.ranges} >= {"A1:E1", "B5:C5", "B6:C6", "B7:C7"}


# ---------------------------------------------------------------- PDF
@pytest.mark.parametrize("rot", [0, 90, 180, 270])
def test_pdf_coordinates_and_redaction(rot):
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    p.insert_font(fontname="f", fontbuffer=pe.cjk_font().buffer)
    p.insert_text((72, 100), "姓名：", fontname="f", fontsize=12)
    p.insert_text((120, 100), "旧值甲乙", fontname="f", fontsize=12)
    p.set_rotation(rot)
    data = doc.tobytes()
    d = pe.open_pdf(data)
    page = d[0]
    r = pymupdf.Rect(116, 86, 300, 104)
    nr = pe.unrot_to_visual_norm(page, r)
    back = pe.visual_norm_to_unrot(page, nr)
    assert all(abs(a - b) < 0.01 for a, b in zip(r, back))            # zoom/rotation independent round trip
    assert pe.extract_region_text(d, 0, back) == "旧值甲乙"
    slot = {"id": "s", "kind": "pdf_box", "name": "姓名", "locator": {"page": 0, "rect": list(back)}, "font_size": 12,
            "text_rotate": pe.text_rotation_in(page, back)}  # follow the direction of the replaced text
    work, warns = pe.make_working_copy(data, [slot])
    assert not warns
    out, issues = pe.fill_pdf(work, [slot], {"s": "新值"})
    txt = pe.open_pdf(out)[0].get_text()
    assert "旧值甲乙" not in txt and "新值" in txt                    # old value really removed


def test_pdf_overflow_is_an_error_not_truncation():
    doc = pymupdf.open()
    doc.new_page(width=300, height=200)
    slot = {"id": "s", "kind": "pdf_box", "name": "备注", "locator": {"page": 0, "rect": [10, 10, 60, 24]}, "font_size": 12, "min_font_size": 8}
    out, issues = pe.fill_pdf(doc.tobytes(), [slot], {"s": "这是一段非常非常长的文字，无论如何也放不进这个很小的框里" * 3})
    assert any(i["code"] == "overflow" and i["level"] == "error" for i in issues)


# ---------------------------------------------------------------- naming
def test_safe_names_and_dedupe():
    assert safe_output_name("../../etc/passwd") == "__etc_passwd" or "/" not in safe_output_name("../../etc/passwd")
    assert "/" not in safe_output_name("a/b\\c:d*e?f") and ".." not in safe_output_name("a..b")
    assert safe_output_name("CON") == "_CON"
    assert len(safe_output_name("长" * 200).encode()) <= 150
    assert dedupe(["a", "A", "a", "b"]) == ["a", "A (2)", "a (3)", "b"]
    assert render_name("{序号}_{不存在}", {"label": "x"}, None, [], 4, 3)[0] == "005_"


def test_combine_rejects_unknown_op():
    """DW-051: a mistyped combine op used to fall through to concatenation, so a mapping
    meant to add 100 + 200 silently delivered "100 200" with nothing reported.
    aggregate() already raised for the same mistake; combine() now matches it."""
    from docweave.engine import transforms as T
    from docweave.engine import values as V

    a, b = V.Value(V.DEC, "100", "100"), V.Value(V.DEC, "200", "200")
    assert T.combine([a, b], "sum").text == "300"
    assert T.combine([a, b], "concat", " ").text == "100 200"
    assert T.combine([a, b], "newline").text == "100\n200"

    for bad in ("sumn", "SUM", "add", "", "join"):
        with pytest.raises(T.TransformError) as e:
            T.combine([a, b], bad)
        assert "不支持的合并方式" in str(e.value) and "concat / newline / sum" in str(e.value)

    for bad in ("sumn", "SUM", "concat"):        # the aggregate side must stay consistent
        with pytest.raises(T.TransformError):
            T.aggregate([a, b], bad)
    assert T.aggregate([a, b], "join", "、").text == "100、200"      # join is an aggregate op, not a combine one
