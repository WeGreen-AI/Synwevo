"""Repeating content blocks (groups of paragraphs, not table rows) in DOCX templates."""
import io
import zipfile

import docx
import openpyxl
from lxml import etree

from conftest import _walk

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _paths(api, tid):
    v = api.view(tid)
    body = next(p for p in v["doc"]["parts"] if p["kind"] == "body")
    return {b["text"]: b for b in _walk(body["blocks"])}, v


def _orders(fixtures):
    ws = openpyxl.load_workbook(fixtures / "测试_订单明细100.xlsx")["明细"]
    out: dict[str, list] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        out.setdefault(r[0], []).append(r)
    for v in out.values():
        v.sort(key=lambda r: r[2])
    return out


def _setup(api, fixtures, **block):
    pid = api.project("段落重复块")
    api.upload(pid, fixtures / "测试_订单明细100.xlsx")
    t = api.template(pid, fixtures / "测试_订单段落模板.docx")
    tid = t["id"]
    blocks, _ = _paths(api, tid)
    first, last = blocks["行号：1"], blocks["单价（元）：0.00"]
    rep = api.slot(tid, kind="repeat_block", part=first["part"], first=first["path"], last=last["path"], name="明细项", **block)
    order = api.text_slot(tid, "DD-000", "订单号")
    return pid, tid, rep, order


def _map(api, pid, rep, order, cols=("行号", "品名", "数量", "单价")):
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "订单号"}], "sort": [{"header": "行号", "dir": "asc", "numeric": True}]}})
    by_name = {c["name"]: c["key"] for c in rep["columns"]}
    names = {"行号": "行号", "品名": "品名", "数量": "数量", "单价": "单价（元）"}
    api.maps(pid, [
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "订单号"}], "target": {"binding": order["binding"]}},
        {"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": rep["binding"]},
         "columns": [{"slot_col": by_name[names[c]], "source": {"header": c}} for c in cols]},
    ])


def test_block_columns_detected_and_filled(api, fixtures):
    pid, tid, rep, order = _setup(api, fixtures, separator="blank")
    assert rep["kind"] == "repeat" and rep["block"] and rep["paragraphs"] == 3
    assert [c["name"] for c in rep["columns"]] == ["行号", "品名", "数量", "单价（元）"]
    assert [c["original_text"] for c in rep["columns"]] == ["1", "示例商品", "1", "0.00"]
    # the view marks the block's paragraphs
    blocks, v = _paths(api, tid)
    assert blocks["行号：1"].get("blk") == rep["id"] and blocks["（明细结束）"].get("blk") is None
    _map(api, pid, rep, order)
    job = api.job(pid, naming="{订单号}")
    assert job["status"] == "completed", api.items(job["id"])[0]
    exp = _orders(fixtures)
    outs = api.outputs(job["id"])
    for key in ("DD-001", "DD-015"):
        d = docx.Document(io.BytesIO(outs[f"{key}.docx"]))
        paras = [p.text for p in d.paragraphs]
        i = paras.index(f"订单号：{key}")
        got = paras[i + 1: paras.index("（明细结束）")]
        want = []
        for k, r in enumerate(exp[key]):
            if k:
                want.append("")                                  # blank separator between copies
            want += [f"行号：{r[2]}", f"品名：{r[3]}　数量：{r[4]}", f"单价（元）：{r[5]}"]
        assert got == want
        assert d.paragraphs[i + 1].runs[0].bold                 # formatting of the copied paragraphs is kept
        xml = zipfile.ZipFile(io.BytesIO(outs[f"{key}.docx"])).read("word/document.xml")
        root = etree.fromstring(xml)
        assert not root.findall(f".//{{{W}}}sdt")               # no content controls left
        assert d.tables[0].cell(0, 0).text == "单元格内的段落"   # content after the block is untouched
    from docweave.engine.converters import convert, soffice_path
    if soffice_path():                                            # a real office suite opens and renders it
        import pymupdf
        pdf = pymupdf.open(stream=convert(outs["DD-001.docx"], "docx", "pdf"), filetype="pdf")
        text = "".join(p.get_text() for p in pdf).replace("\n", "")
        assert all(r[3] in text for r in exp["DD-001"])


def test_page_break_separator_and_zero_rows(api, fixtures):
    pid, tid, rep, order = _setup(api, fixtures, separator="page_break", empty_policy="placeholder")
    _map(api, pid, rep, order)
    job = api.job(pid, naming="{订单号}")
    out = api.outputs(job["id"])["DD-002.docx"]
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml"))
    n = len(_orders(fixtures)["DD-002"])
    assert len(root.findall(f".//{{{W}}}br[@{{{W}}}type='page']")) == n - 1
    # zero detail rows: placeholder policy keeps one copy with 「（无）」 in the first column
    from docweave.engine import templates as T
    tpl = api.ok(api.c.get(f"/api/templates/{tid}"))
    slot = next(s for s in tpl["version"]["slots"] if s["id"] == rep["id"])
    work = api.c.get(f"/api/templates/{tid}/download", params={"which": "working"})
    assert work.status_code == 200
    data = work.content
    filled, issues = T.docx_fill(data, [slot], {slot["id"]: {"kind": "rows", "rows": []}})
    d = docx.Document(io.BytesIO(filled))
    texts = [p.text for p in d.paragraphs]
    assert "行号：（无）" in texts and "品名：　数量：" in texts
    slot["empty_policy"] = "remove"
    filled, _ = T.docx_fill(data, [slot], {slot["id"]: {"kind": "rows", "rows": []}})
    texts = [p.text for p in docx.Document(io.BytesIO(filled)).paragraphs]
    assert not any(t.startswith("行号") for t in texts) and "（明细结束）" in texts


def test_columns_can_be_removed_and_added(api, fixtures):
    pid, tid, rep, order = _setup(api, fixtures, auto_columns=True)
    qty = next(c for c in rep["columns"] if c["name"] == "数量")
    upd = api.ok(api.c.patch(f"/api/templates/{tid}/slots/{rep['id']}", json={"remove_column": qty["key"]}))
    s = next(x for x in upd["version"]["slots"] if x["id"] == rep["id"])
    assert "数量" not in [c["name"] for c in s["columns"]]
    blocks, _ = _paths(api, tid)
    para = blocks["品名：示例商品　数量：1"]
    assert para.get("blk") == rep["id"]
    start = para["text"].index("1", para["text"].index("数量"))
    after = api.ok(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "block_column", "block_id": rep["id"], "part": para["part"],
                                                                     "path": para["path"], "start": start, "end": start + 1, "name": "件数"}))
    s = after["created_slot"]
    assert s["id"] == rep["id"] and s["columns"][-1]["name"] == "件数"
    assert len(after["version"]["slots"]) == 2                        # no new slot: the block got a column
    # selecting text outside the block is refused
    outside = blocks["（明细结束）"]
    r = api.c.post(f"/api/templates/{tid}/slots", json={"kind": "block_column", "block_id": rep["id"], "part": outside["part"],
                                                       "path": outside["path"], "start": 0, "end": 2, "name": "x"})
    assert r.status_code == 400 and "不在该重复内容块内" in r.text
    # removing the block restores the original paragraphs
    api.ok(api.c.delete(f"/api/templates/{tid}/slots/{rep['id']}"))
    blocks, _ = _paths(api, tid)
    assert "行号：1" in blocks and blocks["行号：1"].get("blk") is None and "单价（元）：0.00" in blocks


def test_block_selection_rules(api, fixtures):
    pid = api.project("段落重复块-规则")
    t = api.template(pid, fixtures / "测试_订单段落模板.docx")
    blocks, _ = _paths(api, t["id"])
    body_p, cell_p = blocks["行号：1"], blocks["单元格内的段落"]
    r = api.c.post(f"/api/templates/{t['id']}/slots", json={"kind": "repeat_block", "part": body_p["part"], "first": body_p["path"], "last": cell_p["path"], "name": "x"})
    assert r.status_code == 400 and "不在同一个位置" in r.text
    ok = api.slot(t["id"], kind="repeat_block", part=body_p["part"], first=body_p["path"], last=body_p["path"], name="单段")
    r = api.c.post(f"/api/templates/{t['id']}/slots", json={"kind": "repeat_block", "part": body_p["part"], "first": body_p["path"], "last": body_p["path"], "name": "y"})
    assert r.status_code == 400 and "不能嵌套" in r.text
    assert ok["columns"][0]["name"] == "行号"


def test_builtin_editor_repeat_block(api, fixtures):
    """The in-app document editor's 「重复内容块」 produces the same block structure."""
    pid = api.project("内置编辑器重复块")
    api.upload(pid, fixtures / "测试_订单明细100.xlsx")
    f = lambda i, n: {"type": "field", "attrs": {"id": i, "name": n}}  # noqa: E731
    editor = {"page": {"size": "A4", "orientation": "portrait", "margins": {"top": 25, "bottom": 25, "left": 30, "right": 30}},
              "body": {"type": "doc", "content": [
                  {"type": "paragraph", "content": [{"type": "text", "text": "订单："}, f("fo", "订单号")]},
                  {"type": "repeatBlock", "attrs": {"repeatId": "bk1", "repeatName": "商品", "separator": "blank", "emptyPolicy": "remove"}, "content": [
                      {"type": "paragraph", "content": [{"type": "text", "text": "品名："}, f("c_name", "品名")]},
                      {"type": "bulletList", "content": [{"type": "listItem", "content": [
                          {"type": "paragraph", "content": [{"type": "text", "text": "数量 "}, f("c_qty", "数量")]}]}]},
                  ]},
                  {"type": "paragraph", "content": [{"type": "text", "text": "完"}]},
              ]}}
    t = api.ok(api.c.post("/api/templates/builtin", json={"kind": "builtin_doc", "name": "块模板", "editor": editor, "project_id": pid}))
    blk = next(s for s in t["version"]["slots"] if s["id"] == "bk1")
    assert blk["kind"] == "repeat" and blk["block"] and blk["separator"] == "blank" and [c["name"] for c in blk["columns"]] == ["品名", "数量"]
    fo = next(s for s in t["version"]["slots"] if s["name"] == "订单号")
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "订单号"}], "sort": [{"header": "行号", "dir": "asc", "numeric": True}]}})
    api.maps(pid, [
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "订单号"}], "target": {"binding": fo["binding"]}},
        {"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": "bk1"},
         "columns": [{"slot_col": "c_name", "source": {"header": "品名"}}, {"slot_col": "c_qty", "source": {"header": "数量"}}]},
    ])
    job = api.job(pid, naming="{订单号}")
    assert job["status"] == "completed", api.items(job["id"])[0]
    d = docx.Document(io.BytesIO(api.outputs(job["id"])["DD-003.docx"]))
    texts = [p.text for p in d.paragraphs]
    rows = _orders(fixtures)["DD-003"]
    assert texts[0] == "订单：DD-003" and texts[-1] == "完"
    assert [t for t in texts if t.startswith("品名：")] == [f"品名：{r[3]}" for r in rows]
    assert [t for t in texts if t.startswith("数量 ")] == [f"数量 {r[4]}" for r in rows]
    assert texts.count("") == len(rows) - 1
