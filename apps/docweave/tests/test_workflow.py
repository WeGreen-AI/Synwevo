"""Editing workflow: per-file relocation, undo/redo, slot deletion cascades, preview invalidation."""
import io

import docx

from conftest import docx_text


def _form(label: str, value: str) -> bytes:
    d = docx.Document()
    d.add_paragraph("测试申请表")
    d.add_paragraph(f"{label}：{value}")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_relocate_mapping_for_one_file(api, fixtures):
    pid = api.project("relocate")
    api.upload_bytes(pid, "A.docx", _form("申请人", "测试甲"))
    api.upload_bytes(pid, "B.docx", _form("姓名（新版表单）", "测试乙"))
    srcs = {s["filename"]: s for s in api.ok(api.c.get(f"/api/projects/{pid}/sources"))}
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    s = api.text_slot(t["id"], "示例客户名称", "客户名称")
    api.rule(pid, {"mode": "per_file"})
    loc = api.ok(api.c.post(f"/api/sources/{srcs['A.docx']['id']}/locator", json={"project_id": pid, "selection": {
        "kind": "docx_text", "part": "word/document.xml", "path": "1", "start": 4, "end": 7}}))["proposals"][0]["locator"]
    maps = api.maps(pid, [{"type": "single", "sources": [loc], "target": {"binding": s["binding"]}}])["mappings"]
    val = {f["filename"]: list(f["results"].values())[0] for f in api.ok(api.c.get(f"/api/projects/{pid}/validate"))["files"]}
    assert val["A.docx"]["ok"] and val["B.docx"]["ok"] is False           # the other form has another label → not guessed
    # user re-selects the value in B.docx → stored as an override for that file only
    loc_b = api.ok(api.c.post(f"/api/sources/{srcs['B.docx']['id']}/locator", json={"project_id": pid, "selection": {
        "kind": "docx_text", "part": "word/document.xml", "path": "1", "start": 9, "end": 12}}))["proposals"][0]["locator"]
    maps[0]["overrides"] = {srcs["B.docx"]["id"]: {"source_index": 0, "locator": loc_b}}
    api.maps(pid, maps)
    loc_view = api.ok(api.c.get(f"/api/projects/{pid}/locate", params={"file": srcs["B.docx"]["id"]}))
    assert loc_view[maps[0]["id"]][0]["overridden"] and loc_view[maps[0]["id"]][0]["value"] == "测试乙"
    job = api.job(pid, naming="{来源文件名}")
    outs = api.outputs(job["id"])
    assert "测试甲" in docx_text(outs["A.docx"]) and "测试乙" in docx_text(outs["B.docx"])


def test_undo_redo_and_slot_deletion_cascade(api, fixtures):
    pid = api.project("undo")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    s = api.text_slot(t["id"], "示例客户名称", "客户名称")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 1}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
    pv1 = api.ok(api.c.get(f"/api/projects/{pid}/preview"))
    assert "测试钱磊敏" in docx_text(api.c.get(f"/api/preview-files/{pv1['output']['blob']}").content)
    # deleting the field removes its mapping and restores the original template text
    d = api.ok(api.c.delete(f"/api/templates/{t['id']}/slots/{s['id']}", params={"project_id": pid}))
    assert d["removed_mappings"] and not d["version"]["slots"]
    assert api.ok(api.c.get(f"/api/projects/{pid}/mappings"))["mappings"] == []
    pv2 = api.ok(api.c.get(f"/api/projects/{pid}/preview"))
    assert pv2["signature"] != pv1["signature"]                               # old preview invalidated
    assert "示例客户名称" in docx_text(api.c.get(f"/api/preview-files/{pv2['output']['blob']}").content)
    # undo both: template field back, mapping back
    api.ok(api.c.post(f"/api/templates/{t['id']}/undo"))
    m = api.ok(api.c.post(f"/api/projects/{pid}/mappings/undo"))
    assert len(m["mappings"]) == 1 and m["can_redo"]
    pv3 = api.ok(api.c.get(f"/api/projects/{pid}/preview"))
    assert "测试钱磊敏" in docx_text(api.c.get(f"/api/preview-files/{pv3['output']['blob']}").content)
    # redo mapping deletion again
    m = api.ok(api.c.post(f"/api/projects/{pid}/mappings/redo"))
    assert m["mappings"] == []
    # stale concurrent edit is rejected instead of overwriting
    r = api.c.put(f"/api/projects/{pid}/mappings", json={"mappings": [], "base_rev": 0})
    assert r.status_code == 409


def test_template_library_duplicate_isolated(api, fixtures):
    pid = api.project("lib")
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    api.text_slot(t["id"], "示例客户名称", "客户名称")
    dup = api.ok(api.c.post(f"/api/templates/{t['id']}/duplicate", json={}))
    assert [s["name"] for s in dup["version"]["slots"]] == ["客户名称"]
    api.text_slot(dup["id"], "2020-01-01", "日期")
    assert len(api.ok(api.c.get(f"/api/templates/{t['id']}"))["version"]["slots"]) == 1   # original untouched
    orig = api.c.get(f"/api/templates/{t['id']}/download", params={"which": "original"}).content
    assert orig == (fixtures / "测试_合同模板.docx").read_bytes()                        # original file byte-identical
