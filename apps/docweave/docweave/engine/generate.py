"""Produce one output document for one record.  Used by preview *and* batch jobs."""
from __future__ import annotations

import csv
import io

from . import templates as T
from .plan import Plan

NATIVE_EXT = {"docx": "docx", "xlsx": "xlsx", "pdf": "pdf", "pptx": "pptx", "builtin_doc": "docx", "builtin_sheet": "xlsx",
              "png": "png", "jpeg": "jpg"}
IMAGE_FORMATS = ("png", "jpg")
PDF_LIKE_TEMPLATES = ("pdf", "png", "jpeg")      # templates edited on page images with boxes


def native_format(template_format: str) -> str:
    return NATIVE_EXT.get(template_format, template_format)


def fill_native(template_format: str, work: bytes, slots: list[dict], plan: Plan, options: dict | None = None) -> tuple[bytes, list[dict]]:
    fmt = native_format(template_format)
    if fmt == "docx":
        return T.docx_fill(work, slots, plan.values, options)
    if fmt == "xlsx":
        return T.xlsx_fill(work, slots, plan.values, options)
    if fmt == "pdf":
        return T.pdf_fill(work, slots, plan.values, options)
    if fmt == "pptx":
        from .pptx_engine import pptx_fill
        return pptx_fill(work, slots, plan.values, options)
    if fmt in IMAGE_FORMATS:
        return T.image_fill(work, slots, plan.values, options)
    raise T.TemplateError("unsupported", f"不支持的模板格式：{template_format}")


LAYOUT_FORMATS = ("docx", "xlsx", "pptx")


def fill_checked(template_format: str, work: bytes, slots: list[dict], plan: Plan, options: dict | None,
                 layout_mode: str, tpl_pages: int | None) -> tuple[bytes, list[dict], dict | None]:
    """``fill_native`` plus the pagination guard (``layout_mode``: off / warn / fit).
    Returns (bytes, issues, layout report or None)."""
    options = dict(options or {})
    fmt = native_format(template_format)
    if layout_mode not in ("warn", "fit") or fmt not in LAYOUT_FORMATS or tpl_pages is None:
        out, issues = fill_native(template_format, work, slots, plan, options)
        return out, issues, None
    from . import layout as L
    try:
        if layout_mode == "fit" and fmt == "docx" and any(s.get("shrink") for s in slots):
            keep = bool(options.get("keep_fields"))
            out, issues = fill_native(template_format, work, slots, plan, {**options, "keep_fields": True})
            if any(i.get("level") == "error" for i in issues) or not out:
                return out, issues, None
            out, rep = L.fit_docx(out, slots, tpl_pages, keep)
            if rep["effective_factor"] < 1 and rep["pages"] <= tpl_pages:
                detail = f"，最小 {rep['min_pt']:.1f} 磅" if rep.get("min_pt") else ""
                pinned = f"；字段 {'、'.join(rep['clamped'])} 已到各自的最小字号" if rep.get("clamped") else ""
                issues.append({"level": "warning", "code": "shrunk_to_fit",
                               "message": f"为保持 {tpl_pages} 页，允许缩小的字段字号已缩小到原来的 "
                                          f"{rep['effective_factor']:.0%}{detail}{pinned}"})
            # Fitting the page budget is not success on its own: shrinking can push content out
            # of a fixed-height row or frame, leaving the right page count and the wrong document.
            if rep.get("lost_fields"):
                issues.append({"level": "error", "code": "content_lost",
                               "message": f"缩放后这些字段的内容在文档中消失了：{'、'.join(rep['lost_fields'])}。"
                                          f"页数满足了要求，但内容不完整，因此不产出这一份"})
            elif rep.get("truncated_fields"):
                issues.append({"level": "error", "code": "content_truncated",
                               "message": f"缩放后这些字段的内容被截断：{'、'.join(rep['truncated_fields'])}。"
                                          f"页数满足了要求，但内容不完整，因此不产出这一份"})
            iss = L.page_issue(rep["pages"], tpl_pages, strict=True, slot_hint=True)
            if iss:
                issues.append(iss)
            return out, issues, rep
        out, issues = fill_native(template_format, work, slots, plan, options)
        if any(i.get("level") == "error" for i in issues) or not out:
            return out, issues, None
        pages = L.page_count(out, fmt)
        iss = L.page_issue(pages, tpl_pages, strict=layout_mode == "fit", slot_hint=False)
        if iss:
            issues.append(iss)
        return out, issues, {"template_pages": tpl_pages, "pages": pages, "factor": 1.0, "renders": 1}
    except L.LayoutUnavailable as e:
        out, issues = fill_native(template_format, work, slots, plan, options)
        issues.append({"level": "warning", "code": "layout_unavailable", "message": str(e)})
        return out, issues, None


def template_pages(template_format: str, work: bytes) -> int | None:
    """Page count of the template itself (the reference for the pagination guard)."""
    fmt = native_format(template_format)
    if fmt not in LAYOUT_FORMATS:
        return None
    from . import layout as L
    try:
        return L.page_count(work, fmt)
    except L.LayoutUnavailable:
        return None


def export_formats(template_format: str, soffice_available: bool) -> list[dict]:
    """Formats offered for a template, each with an honest availability flag."""
    fmt = native_format(template_format)
    out = [{"format": fmt, "label": f"{fmt.upper()}（模板原生格式）", "available": True, "native": True}]
    if fmt in ("docx", "xlsx", "pptx"):
        out.append({"format": "pdf", "label": "PDF（经 LibreOffice 转换）", "available": soffice_available,
                    "reason": None if soffice_available else "未检测到 LibreOffice，无法转换 PDF"})
    if fmt == "xlsx":
        out.append({"format": "csv", "label": "CSV（仅第一个工作表的值，不含样式）", "available": True,
                    "note": "CSV 只保存单元格值，不保存字体、颜色、合并单元格和多工作表；公式写出计算结果"})
    if fmt == "docx":
        out.append({"format": "xlsx", "label": "XLSX", "available": False,
                    "reason": "文档型模板无法直接导出为 Excel。替代方案：改用表格型模板（内置表格模板或上传 XLSX 模板）生成 Excel，或导出 DOCX / PDF"})
        out.append({"format": "csv", "label": "CSV", "available": False,
                    "reason": "文档型模板没有行列结构，无法导出 CSV。替代方案：使用表格型模板，并选择 CSV 导出"})
    if fmt in IMAGE_FORMATS:
        out[0]["label"] = f"{fmt.upper()}（与模板图片同尺寸）"
        out.append({"format": "pdf", "label": "PDF（单页，嵌入生成的图片）", "available": True})
    if fmt == "pdf":
        out.append({"format": "docx", "label": "DOCX", "available": False,
                    "reason": "PDF 模板没有可编辑的文档结构，不能无损转换为 Word。请上传 DOCX 模板或用内置文档模板重建"})
    return out


def convert_output(native_bytes: bytes, native_fmt: str, target: str, options: dict | None = None) -> tuple[bytes, list[dict]]:
    issues: list[dict] = []
    if target == native_fmt:
        return native_bytes, issues
    if target == "pdf" and native_fmt in ("docx", "xlsx", "pptx"):
        from .converters import convert
        return convert(native_bytes, native_fmt, "pdf"), issues
    if target == "csv" and native_fmt == "xlsx":
        return xlsx_to_csv(native_bytes, issues), issues
    if target == "pdf" and native_fmt in IMAGE_FORMATS:
        from .image_template import to_pdf
        return to_pdf(native_bytes), issues
    raise T.TemplateError("unsupported_export", f"不支持从 {native_fmt.upper()} 导出 {target.upper()}")


def xlsx_to_csv(data: bytes, issues: list[dict]) -> bytes:
    """Values of the first sheet.  Formula results come from the cached values DocWeave's
    formula engine (or LibreOffice for unsupported functions) wrote into the output."""
    import openpyxl
    from .numfmt import format_value
    wbf = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    wbv = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws, wv = wbf.worksheets[0], wbv.worksheets[0]
    missing = [c.coordinate for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith("=") and wv[c.coordinate].value is None]
    if missing:
        raise T.TemplateError("formula_uncomputed", f"第一个工作表有 {len(missing)} 个公式无法在服务器计算（如 {'、'.join(missing[:3])}），CSV 无法写出它们的结果；请改为导出 XLSX 或安装 LibreOffice")
    out = io.StringIO()
    wr = csv.writer(out, lineterminator="\r\n")
    for row, vrow in zip(ws.iter_rows(), wv.iter_rows()):
        vals = []
        for c, vc in zip(row, vrow):
            v = vc.value
            text, _ = format_value(v, c.number_format)
            vals.append(text)
        while vals and vals[-1] == "":
            vals.pop()
        wr.writerow(vals)
    if len(wbf.worksheets) > 1:
        issues.append({"level": "warning", "code": "csv_single_sheet", "message": f"CSV 只包含第一个工作表「{ws.title}」，其余 {len(wbf.worksheets) - 1} 个工作表未导出"})
    return ("\ufeff" + out.getvalue()).encode("utf-8")
