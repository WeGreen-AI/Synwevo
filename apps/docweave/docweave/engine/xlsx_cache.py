"""Optional: compute formula results with LibreOffice and write them back as
cached values into *our* output file (the LibreOffice-converted file itself is
discarded, so its re-serialisation never replaces the template's structure)."""
from __future__ import annotations

import io

from lxml import etree

from .converters import ConversionError, recalc_xlsx
from .xlsx_doc import XlsxDoc, s


def fill_cached_values(data: bytes) -> tuple[bytes, list[dict]]:
    import openpyxl
    try:
        calc = recalc_xlsx(data)
    except ConversionError as e:
        return data, [{"level": "warning", "code": "recalc_failed", "message": f"公式未能用 LibreOffice 计算：{e}；Excel 打开时会重新计算"}]
    wb = openpyxl.load_workbook(io.BytesIO(calc), data_only=True, read_only=True)
    x = XlsxDoc(data)
    n = 0
    for sh in x.sheets:
        if sh["name"] not in wb.sheetnames:
            continue
        ws = wb[sh["name"]]
        values = {}
        for row in ws.iter_rows():
            for c in row:
                if getattr(c, "value", None) is not None:
                    values[c.coordinate] = c.value
        root = x.sheet_root(sh["name"])
        for c in root.iter(s("c")):
            f = c.find(s("f"))
            if f is None:
                continue
            val = values.get(c.get("r"))
            for old in c.findall(s("v")):
                c.remove(old)
            if val is None:
                continue
            v = etree.SubElement(c, s("v"))
            if isinstance(val, bool):
                c.set("t", "b")
                v.text = "1" if val else "0"
            elif isinstance(val, (int, float)):
                if "t" in c.attrib:
                    del c.attrib["t"]
                from .values import float_repr
                v.text = float_repr(float(val))
            else:
                c.set("t", "str")
                v.text = str(val)
            n += 1
    wb.close()
    out = x.pkg.save()
    return out, [{"level": "info", "code": "recalc", "message": f"已用 LibreOffice 计算 {n} 个公式的结果并写入缓存值"}]
