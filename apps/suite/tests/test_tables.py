"""Isolated D4 tests using real Store, capability transactions, HTTP and fixture identities."""
import csv
import io
import json
from concurrent.futures import ThreadPoolExecutor

import pytest
from synwevo_core.errors import ServiceError


def call(s, operation, data=None, who='alice', ws=None, key=None):
    return s.caps.invoke('tables.'+operation, data or {}, user_id=s.people.uid(who),
                         workspace_id=ws or s.people.a, idempotency_key=key)


def out(s, op, data=None, **kw):
    return call(s, op, data, **kw)['output']


def table(s, name='Fixture contacts'):
    return out(s, 'table.create', {'name': name})['id']


def get(s, tid, **kw):
    return out(s, 'table.get', {'table_id': tid}, **kw)


def field(s, tid, name='Name', type='text', **kw):
    args = {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'field': {'name': name, 'type': type, **kw}}
    return out(s, 'field.save', args)['field_id']


def record(s, tid, cells=None, **kw):
    return out(s, 'record.create', {'table_id': tid, 'cells': cells or {}, **kw})


def query(s, tid, **kw):
    who = kw.pop('who', 'alice')
    return out(s, 'record.query', {'table_id': tid, **kw}, who=who)


def denied(code, fn, status=None):
    with pytest.raises(ServiceError) as e:
        fn()
    assert e.value.code == code, (e.value.code, e.value.detail)
    if status:
        assert e.value.status == status


def test_auto_discovery_and_worktree_imports(suite):
    from pathlib import Path
    import synwevo_core.store as st
    import synwevo_suite.tables as tb
    root = Path(__file__).resolve().parents[3]
    assert Path(st.__file__).resolve().is_relative_to(root)
    assert Path(tb.__file__).resolve().is_relative_to(root)
    assert suite.caps.registry.app('tables')
    assert suite.people.store.one("SELECT name FROM sqlite_master WHERE name='tb_records'")
    assert len([op for op in suite.caps.registry.operations() if op.app == 'tables']) >= 18


def test_revision_uniqueness_noop_events_and_replay(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid, is_unique=True, required=True)
    args = {'table_id': tid, 'cells': {fid: 'ABC'}}
    first = call(s, 'record.create', args, key='record-intent')
    assert call(s, 'record.create', args, key='record-intent')['replayed']
    rid = first['output']['id']
    denied('unique_conflict', lambda: record(s, tid, {fid: ' abc '}))
    before = s.people.store.one("SELECT COUNT(*) AS n FROM app_events WHERE app='tables'")['n']
    assert not out(s, 'record.update', {'table_id': tid, 'record_id': rid, 'rev': 1, 'cells': {fid: 'ABC'}})['changed']
    assert s.people.store.one("SELECT COUNT(*) AS n FROM app_events WHERE app='tables'")['n'] == before
    out(s, 'record.update', {'table_id': tid, 'record_id': rid, 'rev': 1, 'cells': {fid: 'DEF'}})
    denied('revision_conflict', lambda: out(s, 'record.update', {'table_id': tid, 'record_id': rid, 'rev': 1, 'cells': {fid: 'lost'}}))
    assert query(s, tid)['records'][0]['cells'][fid] == 'DEF'
    denied('idempotency_conflict', lambda: call(s, 'record.create', {'table_id': tid, 'cells': {fid: 'other'}}, key='record-intent'))


def test_acl_query_search_sort_aggregate_and_history(suite):
    s = suite
    tid = table(s)
    name = field(s, tid)
    secret = field(s, tid, 'Restricted', 'number', read_role='admin', write_role='admin')
    calc = field(s, tid, 'Derived', 'formula', config={'expression': '{Restricted} * 2'})
    public = record(s, tid, {name: 'ordinary', secret: '99'})
    record(s, tid, {name: 'hidden-record', secret: '500'}, read_role='admin')
    result = query(s, tid, who='bob')
    assert result['total'] == 1
    assert set(result['records'][0]['cells']) == {name}
    assert not query(s, tid, who='bob', q='99')['records']
    assert not query(s, tid, who='bob', q='hidden-record')['records']
    for clause in ({'fields': [secret]}, {'sort': [{'field_id': secret}]}, {'filters': [{'field_id': calc, 'value': '198'}]}):
        denied('field_forbidden', lambda: query(s, tid, who='bob', **clause))
    denied('field_forbidden', lambda: out(s, 'record.aggregate', {'table_id': tid, 'function': 'sum', 'field_id': secret}, who='bob'))
    assert out(s, 'record.aggregate', {'table_id': tid}, who='bob')['value'] == 1
    history = out(s, 'history.list', {'table_id': tid}, who='bob')
    assert secret not in json.dumps(history) and calc not in json.dumps(history)
    denied('field_forbidden', lambda: out(s, 'record.update', {'table_id': tid, 'record_id': public['id'], 'rev': 1, 'cells': {secret: '0'}}, who='bob'))
    denied('read_only_role', lambda: out(s, 'record.create', {'table_id': tid}, who='vera'))
    denied('app_access_denied', lambda: get(s, tid, who='out'))
    denied('table_not_found', lambda: get(s, tid, who='oscar', ws=s.people.b))


def test_relations_redact_after_restriction_and_archive(suite):
    s = suite
    target = table(s, 'Related')
    primary = field(s, target)
    tr = record(s, target, {primary: 'Visible relation'})
    source = table(s)
    rel = field(s, source, 'Related', 'link', config={'table_id': target, 'multiple': True})
    record(s, source, {rel: [tr['id']]})
    assert query(s, source, who='bob')['records'][0]['relations'][rel][0]['title'] == 'Visible relation'
    out(s, 'record.update', {'table_id': target, 'record_id': tr['id'], 'rev': 1, 'read_role': 'admin'})
    assert query(s, source, who='bob')['records'][0]['cells'][rel] is None
    assert query(s, source, who='bob', q=tr['id'])['total'] == 0
    denied('record_not_found', lambda: out(s, 'record.create', {'table_id': source, 'cells': {rel: [tr['id']]}}, who='bob'))
    assert tr['id'] not in json.dumps(out(s, 'history.list', {'table_id': source}, who='bob'))
    out(s, 'record.state', {'table_id': target, 'record_id': tr['id'], 'rev': 2, 'deleted': True})
    assert query(s, source)['records'][0]['cells'][rel] is None
    out(s, 'record.state', {'table_id': target, 'record_id': tr['id'], 'rev': 3, 'deleted': False})
    assert query(s, source)['records'][0]['cells'][rel] == [tr['id']]


def test_pagination_typed_sort_filter_views_and_batch(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid, 'Amount', 'number', is_unique=True)
    data = {'table_id': tid, 'atomic': False, 'items': [{'action': 'create', 'cells': {fid: x}} for x in ['10', '2', '2', '30']]}
    res = call(s, 'record.batch', data)
    assert res['status'] == 'partial' and res['output']['failed'] == 1
    q = {'table_id': tid, 'sort': [{'field_id': fid}], 'filters': [{'field_id': fid, 'op': 'gte', 'value': '2'}], 'limit': 2}
    page = out(s, 'record.query', q)
    assert [r['cells'][fid] for r in page['records']] == ['2', '10']
    assert page['next_offset'] == 2
    assert out(s, 'record.query', dict(q, offset=2))['records'][0]['cells'][fid] == '30'
    vid = out(s, 'view.save', {'table_id': tid, 'name': 'Amounts', 'query': q})['id']
    assert query(s, tid, view_id=vid)['total'] == 3
    denied('unique_conflict', lambda: out(s, 'record.batch', dict(data, atomic=True, items=[{'action': 'create', 'cells': {fid: '70'}}, {'action': 'create', 'cells': {fid: '2'}}])))
    assert query(s, tid)['total'] == 3
    assert out(s, 'record.aggregate', {'table_id': tid, 'function': 'sum', 'field_id': fid})['value'] == '42'
    denied('field_dependencies', lambda: out(s, 'field.save', {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'field_id': fid, 'field': {'name': 'Amount', 'type': 'text'}}))


def test_formula_stable_ids_error_and_acl_propagation(suite):
    s = suite
    tid = table(s)
    num = field(s, tid, 'Price', 'number', config={'decimals': 2})
    den = field(s, tid, 'Quantity', 'number')
    calc = field(s, tid, 'Ratio', 'formula', config={'expression': '{Price} / {Quantity}'})
    rec = record(s, tid, {num: '10.20', den: '2'})
    assert query(s, tid)['records'][0]['cells'][calc] == '5.10'
    rev = get(s, tid)['schema_rev']
    out(s, 'field.save', {'table_id': tid, 'schema_rev': rev, 'field_id': num,
                         'field': {'name': 'Renamed price', 'type': 'number', 'config': {'decimals': 2}, 'read_role': 'admin'}})
    assert calc not in query(s, tid, who='bob')['records'][0]['cells']
    out(s, 'record.update', {'table_id': tid, 'record_id': rec['id'], 'rev': 1, 'cells': {den: '0'}})
    assert query(s, tid)['records'][0]['formula_errors'][calc] == 'division_by_zero'
    denied('invalid_field', lambda: field(s, tid, 'Bad', 'formula', config={'expression': '__import__("os")'}))


def test_schema_preview_loss_history_restore_and_deleted_field(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid)
    rec = record(s, tid, {fid: '001.20'})
    args = {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'field_id': fid,
            'field': {'name': 'Amount', 'type': 'number', 'config': {'decimals': 1}}}
    preview = out(s, 'field.save', dict(args, preview=True))
    assert preview['counts']['lossy'] == 1
    denied('preview_required', lambda: out(s, 'field.save', args))
    out(s, 'field.save', dict(args, preview_token=preview['preview_token']))
    row = query(s, tid)['records'][0]
    assert row['cells'][fid] == '1.2' and row['rev'] == 2
    assert s.people.store.one('SELECT * FROM tb_leftovers WHERE record_id=?', (rec['id'],))['original_json'] == '"001.20"'
    args = {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'field_id': fid, 'deleted': True}
    preview = out(s, 'field.state', dict(args, preview=True))
    out(s, 'field.state', dict(args, preview_token=preview['preview_token']))
    assert not query(s, tid)['records'][0]['cells']
    args.update(schema_rev=get(s, tid)['schema_rev'], deleted=False)
    preview = out(s, 'field.state', dict(args, preview=True))
    out(s, 'field.state', dict(args, preview_token=preview['preview_token']))
    assert query(s, tid)['records'][0]['cells'][fid] == '1.2'


def test_archive_restore_uniqueness_and_table_acl(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid, is_unique=True)
    rec = record(s, tid, {fid: 'same'})
    out(s, 'record.state', {'table_id': tid, 'record_id': rec['id'], 'rev': 1, 'deleted': True})
    assert query(s, tid)['total'] == 0 and query(s, tid, deleted=True)['total'] == 1
    record(s, tid, {fid: 'same'})
    denied('unique_conflict', lambda: out(s, 'record.state', {'table_id': tid, 'record_id': rec['id'], 'rev': 2, 'deleted': False}))
    out(s, 'table.update', {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'read_role': 'admin'})
    denied('table_not_found', lambda: get(s, tid, who='bob'))
    assert not out(s, 'table.list', who='bob')['tables']


def test_concurrent_updates_one_winner(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid)
    rec = record(s, tid, {fid: 'original'})
    def update(value):
        try:
            return out(s, 'record.update', {'table_id': tid, 'record_id': rec['id'], 'rev': 1, 'cells': {fid: value}})
        except ServiceError as e:
            return e.code
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(update, ['one', 'two']))
    assert sum(isinstance(r, dict) for r in results) == 1
    assert 'revision_conflict' in results
    assert query(s, tid)['records'][0]['rev'] == 2


def test_import_mapping_preview_partial_dedup_and_export_acl(suite):
    s = suite
    s.sign_in('alice')
    tid = table(s)
    fid = field(s, tid, is_unique=True)
    secret = field(s, tid, 'Private', read_role='admin')
    rec = record(s, tid, {fid: '=HYPERLINK("fixture")', secret: 'hidden-content'})
    response = s.post(f'/api/suite/tables/{tid}/imports', params={'workspace_id': s.people.a},
                      files={'file': ('fixture.csv', b'Name\n001\n002\n001\n', 'text/csv')}, headers={'X-Synwevo-CSRF': 'fixture-csrf'})
    assert response.status_code == 200, response.text
    args = {'table_id': tid, 'import_id': response.json()['id'], 'mapping': {'Name': fid}, 'atomic': False}
    preview = out(s, 'import.apply', args)
    assert preview['valid'] == 2 and len(preview['issues']) == 1
    assert query(s, tid)['total'] == 1
    first = call(s, 'import.apply', dict(args, preview=False))
    assert first['status'] == 'partial' and first['output']['succeeded'] == 2
    out(s, 'import.apply', dict(args, preview=False))
    assert query(s, tid)['total'] == 3
    s.sign_in('bob')
    export = out(s, 'export.create', {'table_id': tid}, who='bob')
    data = s.get(export['url'])
    assert data.status_code == 200 and 'hidden-content' not in data.text and 'Private' not in data.text
    assert "'=HYPERLINK" in data.text
    # Tighten a record AFTER ticket creation. A download rebuilds the current projection.
    out(s, 'record.update', {'table_id': tid, 'record_id': rec['id'], 'rev': 1, 'read_role': 'admin'})
    assert 'HYPERLINK' not in s.get(export['url']).text
    s.people.ident.remove_member(s.people.a, s.people.uid('bob'))
    assert s.get(export['url']).status_code == 403


def test_export_xlsx_string_cells_and_private_ticket(suite):
    from openpyxl import load_workbook
    s = suite
    tid = table(s)
    fid = field(s, tid)
    record(s, tid, {fid: '=1+1'})
    export = out(s, 'export.create', {'table_id': tid, 'format': 'xlsx'})
    s.sign_in('bob')
    assert s.get(export['url']).status_code == 404
    s.sign_in('alice')
    response = s.get(export['url'])
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.active['A2'].data_type == 's' and book.active['A2'].value == '=1+1'
    book.close()


def test_formula_and_numeric_limits_are_validation_errors(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid, 'Bounded', 'number', config={'decimals': 2, 'min': '0', 'max': '100'})
    for value in ['-1', '100.01', '9'*300, 'NaN', '1.234']:
        denied('invalid_cell', lambda: record(s, tid, {fid: value}))
    denied('invalid_field', lambda: field(s, tid, 'Deep', 'formula', config={'expression': '('*150+'1'+')'*150}))
    denied('invalid_field', lambda: field(s, tid, 'Deep unary', 'formula', config={'expression': '-'*150+'1'}))
    denied('invalid_field', lambda: field(s, tid, 'Unsupported option', config={'made_up': True}))
    assert record(s, tid, {fid: '100.00'})


def test_stale_migration_preview_required_defaults_and_history_undo(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid)
    rec = record(s, tid, {fid: '4'})
    args = {'table_id': tid, 'field_id': fid, 'schema_rev': get(s, tid)['schema_rev'],
            'field': {'name': 'Number', 'type': 'number'}}
    preview = out(s, 'field.save', dict(args, preview=True))
    out(s, 'record.update', {'table_id': tid, 'record_id': rec['id'], 'rev': 1, 'cells': {fid: '5'}})
    denied('preview_required', lambda: out(s, 'field.save', dict(args, preview_token=preview['preview_token'])))
    changes = out(s, 'history.list', {'table_id': tid, 'record_id': rec['id']})['changes']
    out(s, 'record.restore_change', {'table_id': tid, 'record_id': rec['id'], 'rev': 2, 'change_id': changes[0]['id']})
    assert query(s, tid)['records'][0]['cells'][fid] == '4'
    denied('migration_constraints', lambda: field(s, tid, 'Required', required=True))
    default = field(s, tid, 'Default', required=True, default='present')
    assert query(s, tid)['records'][0]['cells'][default] == 'present'


def test_table_archive_and_permission_limited_restore(suite):
    s = suite
    tid = table(s)
    fid = field(s, tid)
    rec = record(s, tid, {fid: 'fixture'}, write_role='admin')
    denied('record_read_only', lambda: out(s, 'record.update', {'table_id': tid, 'record_id': rec['id'], 'rev': 1, 'cells': {fid: 'changed'}}, who='bob'))
    out(s, 'table.update', {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'archived': True})
    denied('table_archived', lambda: record(s, tid, {fid: 'new'}))
    assert query(s, tid)['total'] == 1
    out(s, 'table.update', {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'archived': False})
    assert record(s, tid, {fid: 'new'})


def test_import_replay_rechecks_resource_permissions(suite):
    s = suite
    s.sign_in('bob')
    tid = table(s)
    fid = field(s, tid)
    response = s.post(f'/api/suite/tables/{tid}/imports', params={'workspace_id': s.people.a},
                      files={'file': ('fixture.csv', b'Name\n001\n', 'text/csv')}, headers={'X-Synwevo-CSRF': 'fixture-csrf'})
    assert response.status_code == 200
    args = {'table_id': tid, 'import_id': response.json()['id'], 'mapping': {'Name': fid}, 'preview': False}
    result = out(s, 'import.apply', args, who='bob')
    rid = result['items'][0]['id']
    out(s, 'record.update', {'table_id': tid, 'record_id': rid, 'rev': 1, 'read_role': 'admin'})
    denied('record_not_found', lambda: out(s, 'import.apply', args, who='bob'))
    assert query(s, tid)['total'] == 1


def guard_call(s, user='bob'):
    from types import SimpleNamespace
    return SimpleNamespace(user_id=s.people.uid(user), workspace_id=s.people.a)


def test_result_guard_current_record_field_relation_acl_and_no_execution(suite):
    from synwevo_suite.tables.guard import result_guard
    s = suite
    service = s.api.state.tables
    guard = result_guard(service)
    target, tid = table(s, 'Target'), table(s)
    tf, fid = field(s, target), field(s, tid)
    linked = record(s, target, {tf: 'Related'})
    rel = field(s, tid, 'Relation', 'link', config={'table_id': target})
    args = {'table_id': tid, 'cells': {fid: 'visible', rel: [linked['id']]}}
    result = call(s, 'record.create', args, who='bob', key='guard-intent')
    historical = s.people.store.one('SELECT * FROM op_runs WHERE id=?', (result['run_id'],))
    guard(guard_call(s), historical)
    # Exercise both JSON encodings allowed by the new contract.
    historical['input'] = json.dumps(args)
    historical['output'] = json.dumps(result['output'])
    guard(guard_call(s), historical)
    rid = result['output']['id']
    out(s, 'record.update', {'table_id': target, 'record_id': linked['id'], 'rev': 1, 'read_role': 'admin'})
    denied('record_not_found', lambda: guard(guard_call(s), historical))
    out(s, 'record.update', {'table_id': target, 'record_id': linked['id'], 'rev': 2, 'read_role': 'viewer'})
    out(s, 'field.save', {'table_id': tid, 'schema_rev': get(s, tid)['schema_rev'], 'field_id': fid, 'field': {'name': 'Name', 'read_role': 'admin'}})
    denied('historical_result_unavailable', lambda: guard(guard_call(s), historical))
    assert s.people.store.one('SELECT rev FROM tb_records WHERE id=?', (rid,))['rev'] == 1
    assert s.people.store.one('SELECT COUNT(*) AS n FROM tb_records WHERE table_id=?', (tid,))['n'] == 1


def test_registry_history_and_idempotency_with_new_guard_contract(suite):
    from synwevo_core.capabilities import Operation
    if 'result_guard' not in Operation.__dataclass_fields__:
        pytest.skip('Coordinator integration required: this worktree core has no Operation.result_guard yet')
    s = suite
    tid = table(s)
    fid = field(s, tid)
    args = {'table_id': tid, 'cells': {fid: 'original'}}
    result = call(s, 'record.create', args, who='bob', key='contract-intent')
    out(s, 'record.update', {'table_id': tid, 'record_id': result['output']['id'], 'rev': 1, 'read_role': 'admin'})
    with pytest.raises(ServiceError):
        call(s, 'record.create', args, who='bob', key='contract-intent')
    receipt = s.caps.run(result['run_id'], user_id=s.people.uid('bob'))
    assert receipt['snapshot_hidden'] is True
    assert receipt['input'] is None and receipt['output'] is None and receipt['object'] is None
    assert query(s, tid)['total'] == 1


def test_record_links_and_agent_confirmation_bind_composite_identity(suite):
    from synwevo_suite.agent.tools import resource_reads
    s = suite
    tid = table(s)
    fid = field(s, tid)
    created = out(s, 'record.create', {'table_id': tid, 'cells': {fid: 'fixture record'}})
    rid = created['id']
    assert out(s, 'record.locate', {'record_id': rid}, who='bob') == {'table_id': tid, 'record_id': rid}
    reads = resource_reads(s.caps, s.caps.registry.get('tables.record.update'),
                           {'table_id': tid, 'record_id': rid, 'rev': 1, 'cells': {fid: 'next'}})
    assert ('tables.table.get', {'table_id': tid}) in reads
    assert ('tables.record.get', {'table_id': tid, 'record_id': rid}) in reads
    out(s, 'record.update', {'table_id': tid, 'record_id': rid, 'rev': 1, 'read_role': 'admin'})
    with pytest.raises(ServiceError):
        out(s, 'record.locate', {'record_id': rid}, who='bob')


def test_representative_scale_query_and_batch_budget(suite):
    from time import perf_counter
    import platform
    s = suite
    tid = table(s)
    fid = field(s, tid, 'Numeric fixture', 'number', is_unique=True)
    field(s, tid, 'Computed fixture', 'formula', config={'expression': '{Numeric fixture} * 2'})
    started = perf_counter()
    for start in range(0, 1000, 100):
        out(s, 'record.batch', {'table_id': tid, 'items': [{'action': 'create', 'cells': {fid: str(n)}} for n in range(start, start+100)]})
    inserted = perf_counter() - started
    started = perf_counter()
    result = query(s, tid, filters=[{'field_id': fid, 'op': 'gte', 'value': '500'}], sort=[{'field_id': fid, 'direction': 'desc'}], limit=50)
    elapsed = perf_counter() - started
    assert result['total'] == 500 and len(result['records']) == 50
    assert result['records'][0]['cells'][fid] == '999'
    assert elapsed < 5, '1,000-row query exceeds the isolated five-second budget'
    print(f'TABLES_BUDGET platform={platform.system()}/{platform.machine()} rows=1000 fields=2 concurrency=1 insert_10x100_s={inserted:.3f} filter_sort_page_s={elapsed:.3f}')


def test_schema_migration_reopen_preserves_data(suite):
    from synwevo_core.store import Store
    s = suite
    tid = table(s)
    fid = field(s, tid)
    rec = record(s, tid, {fid: 'survives restart'})
    original = s.people.store
    reopened = Store(original.root)
    assert reopened.one('SELECT cells FROM tb_records WHERE id=?', (rec['id'],))['cells'] == json.dumps({fid: 'survives restart'}, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def test_related_table_restriction_hides_schema_target_and_export(suite):
    s = suite
    target, source = table(s, 'Private source'), table(s, 'Readable table')
    tf = field(s, target)
    linked = record(s, target, {tf: 'target content'})
    lf = field(s, source, 'Link', 'link', config={'table_id': target})
    record(s, source, {lf: [linked['id']]})
    out(s, 'table.update', {'table_id': target, 'schema_rev': get(s, target)['schema_rev'], 'read_role': 'admin'})
    definition = get(s, source, who='bob')
    assert definition['fields'][0]['config']['table_id'] is None
    assert not definition['fields'][0]['can_write']
    assert target not in json.dumps(definition)
    s.sign_in('bob')
    ticket = out(s, 'export.create', {'table_id': source}, who='bob')
    exported = s.get(ticket['url'])
    assert exported.status_code == 200
    assert linked['id'] not in exported.text and 'target content' not in exported.text
