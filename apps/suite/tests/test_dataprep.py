"""Isolated fixtures: real CSV/XLSX, persisted processing, permissions and crash recovery."""
import io
import json
import time
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from synwevo_core.errors import ServiceError
from synwevo_suite.dataprep.engine import Table, Row, execute
from synwevo_suite.dataprep.files import read_file, export_csv
from synwevo_suite.dataprep.models import Definition, ImportOptions


def op(c, name, payload=None, *, ws=None, key=None, status=200):
    response = c.post('/api/suite/ops/dataprep.' + name, headers={'X-Synwevo-CSRF': 'fixture-csrf'},
                      json={'workspace_id': ws or c.people.a, 'input': payload or {}, **({'idempotency_key': key} if key else {})})
    assert response.status_code == status, response.text
    return response.json() if status != 200 else response.json()['output']


def dataset(c, rows=None):
    return op(c, 'dataset.from_records', {'name': 'Fixture amounts', 'columns': ['id', 'amount'],
              'records': rows if rows is not None else [{'id': '001', 'amount': '0.10'}, {'id': '002', 'amount': '0.20'}]})


def recipe(c, steps=None):
    r = op(c, 'recipe.create', {'name': 'Fixture recipe', 'definition': {'steps': steps or []}})
    op(c, 'recipe.publish', {'recipe_id': r['id'], 'draft_rev': 1})
    return r


def start(c, r, d, **kwargs):
    return op(c, 'run.start', {'recipe_id': r['id'], 'version': 1, 'inputs': {'main': d['id']}}, **kwargs)


def finish(c, run):
    assert c.api.state.dataprep.tick()
    return op(c, 'run.get', {'run_id': run['id']})


def test_csv_range_null_missing_leading_zero_and_strict_shape():
    data = b'Ignored\nID,Amount,Note\n001,0.10,\n002,0.20,\\N\n003,invalid\n004,1,x,extra\n'
    table, _ = read_file(data, 'csv', ImportOptions(start_row=2), 'fixture')
    assert [r.values['ID'] for r in table.rows] == ['001', '002', '003']
    assert table.rows[0].values['Note'] == ''
    assert table.rows[1].values['Note'] is None
    assert 'Note' not in table.rows[2].values
    assert table.issues[0]['source_row'] == 6
    assert table.rows[0].origins == [['fixture', 3]]
    selected, _ = read_file(data, 'csv', ImportOptions(start_row=2, end_row=3, end_column=2), 'fixture')
    assert selected.columns == ['ID', 'Amount'] and len(selected.rows) == 1
    with pytest.raises(ServiceError, match='数据') as e:
        read_file(b'ID,ID\nx,y', 'csv', ImportOptions(), 'fixture')
    assert e.value.code == 'invalid_columns'


def workbook():
    book = Workbook()
    book.active.title = 'Skip'
    sheet = book.create_sheet('Fixture')
    sheet.append(['ID', 'Amount', 'Formula'])
    sheet.append(['0001', 123, '=1+2'])
    buffer = io.BytesIO()
    book.save(buffer)
    # Real OOXML numeric lexical content longer than a binary double's exact precision.
    source, target = ZipFile(io.BytesIO(buffer.getvalue())), io.BytesIO()
    with ZipFile(target, 'w') as dest:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet2.xml':
                data = data.replace(b'<v>123</v>', b'<v>12345678901234567890.12</v>').replace(b'<f>1+2</f><v></v>', b'<f>1+2</f><v>3</v>')
            dest.writestr(entry, data)
    return target.getvalue()


def test_xlsx_sheet_formula_and_lexical_precision():
    data = workbook()
    with pytest.raises(ServiceError) as e:
        read_file(data, 'xlsx', ImportOptions(sheet='Fixture'), 'fixture')
    assert e.value.code == 'formula_rejected'
    table, notes = read_file(data, 'xlsx', ImportOptions(sheet='Fixture', formulas='cached'), 'fixture')
    assert table.rows[0].values == {'ID': '0001', 'Amount': '12345678901234567890.12', 'Formula': '3'}
    assert 'formula_cache_not_recalculated' in notes
    table, _ = read_file(data, 'xlsx', ImportOptions(sheet='Fixture', end_column=2), 'fixture')
    assert len(table.rows) == 1


def test_engine_exact_aggregate_and_invalid_cast():
    source = Table(['id', 'amount'], [Row({'id': '001', 'amount': v}, [['fixture', i + 2]]) for i, v in enumerate(['0.10', '0.20', 'oops'])])
    definition = Definition.model_validate({'steps': [
        {'id': 'money', 'type': 'cast', 'column': 'amount', 'to': 'decimal'},
        {'id': 'total', 'type': 'aggregate', 'group_by': ['id'], 'measures': [{'op': 'sum', 'column': 'amount', 'target': 'total'}]}]})
    out, issues, reports = execute(definition, {'main': source}, {})
    assert out.rows[0].values == {'id': '001', 'total': '0.30'}
    assert issues[0]['source_row'] == 4 and issues[0]['code'] == 'invalid_decimal'
    assert reports[0]['output_rows'] == 2


def test_timezone_precision_and_null_rules():
    source = Table(['date'], [Row({'date': v}) for v in ['2026-09-30T09:00:00+08:00', '2026-09-30T09:00:00', None]])
    definition = Definition.model_validate({'steps': [{'id': 'time', 'type': 'cast', 'column': 'date', 'to': 'datetime'}]})
    out, issues, _ = execute(definition, {'main': source}, {})
    assert [r.values['date'] for r in out.rows] == ['2026-09-30T01:00:00Z', None]
    assert issues[0]['code'] == 'timezone_required'
    definition.steps[0].offset = '+08:00'
    out, issues, _ = execute(definition, {'main': source}, {})
    assert len(out.rows) == 3 and not issues


def test_many_to_many_join_nulls_and_stable_order():
    a = Table(['key', 'left'], [Row({'key': 'x', 'left': str(i)}) for i in range(2)] + [Row({'key': None})])
    b = Table(['key', 'right'], [Row({'key': 'x', 'right': str(i)}) for i in range(2)] + [Row({'key': None})])
    step = {'id': 'join', 'type': 'join', 'source': 'other', 'left': ['key'], 'right': ['key'], 'matches': 'one'}
    with pytest.raises(ServiceError) as e:
        execute(Definition.model_validate({'steps': [step]}), {'main': a, 'other': b}, {})
    assert e.value.code == 'duplicate_match'
    out, _, _ = execute(Definition.model_validate({'steps': [{**step, 'matches': 'all'}]}), {'main': a, 'other': b}, {})
    assert [(r.values['left'], r.values['right_right']) for r in out.rows[:4]] == [('0', '0'), ('0', '1'), ('1', '0'), ('1', '1')]
    assert 'right_key' not in out.rows[-1].values


def test_version_full_run_preview_issues_and_events(suite):
    suite.sign_in('bob')
    d = dataset(suite, [{'id': f'{i:03}', 'amount': 'oops' if i == 9 else '0.10'} for i in range(10)])
    steps = [{'id': 'money', 'type': 'cast', 'column': 'amount', 'to': 'decimal'}]
    r = recipe(suite, steps)
    preview = op(suite, 'recipe.preview', {'definition': {'steps': steps}, 'inputs': {'main': d['id']}, 'limit': 2})
    assert preview['sample_only'] and preview['ranges']['main'] == {'offset': 0, 'selected': 2, 'total': 10}
    assert preview['issues'] == []
    run = start(suite, r, d, key='fixture-full-run')
    assert start(suite, r, d, key='fixture-full-run')['id'] == run['id']
    op(suite, 'recipe.save', {'recipe_id': r['id'], 'draft_rev': 1, 'name': 'Changed', 'definition': {'steps': []}})
    op(suite, 'recipe.publish', {'recipe_id': r['id'], 'draft_rev': 2})
    result = finish(suite, run)
    assert result['recipe_version'] == 1 and result['status'] == 'partial'
    assert result['counts']['output_rows'] == 9
    issues = op(suite, 'run.issues', {'run_id': run['id']})
    assert issues['total'] == 1 and issues['items'][0]['source_row'] == 10
    events = suite.people.store.all("SELECT type FROM app_events WHERE object_id=?", (run['id'],))
    assert {e['type'] for e in events} == {'dataprep.run.queued', 'dataprep.run.partial'}
    assert suite.people.store.one('SELECT status FROM op_runs WHERE id=?', (run['op_run_id'],))['status'] == 'partial'


def test_authorization_isolation_conflicts_and_export(suite):
    suite.sign_in('alice')
    d = dataset(suite)
    r = recipe(suite)
    op(suite, 'recipe.save', {'recipe_id': r['id'], 'draft_rev': 7, 'name': 'stale'}, status=409)
    suite.sign_in('vera')
    op(suite, 'dataset.get', {'dataset_id': d['id']})
    op(suite, 'recipe.create', {'name': 'No'}, status=403)
    assert suite.get(f"/api/suite/dataprep/datasets/{d['id']}/export", params={'workspace_id': suite.people.a}).status_code == 403
    suite.sign_in('oscar')
    op(suite, 'dataset.get', {'dataset_id': d['id']}, ws=suite.people.b, status=404)
    suite.sign_in('out')
    op(suite, 'dataset.get', {'dataset_id': d['id']}, status=403)


def test_cancel_recovery_and_fenced_results(suite):
    suite.sign_in('bob')
    d, r = dataset(suite), recipe(suite)
    queued = start(suite, r, d)
    assert op(suite, 'run.cancel', {'run_id': queued['id']})['status'] == 'cancelled'
    run = start(suite, r, d)
    suite.people.store.update('dataprep_runs', {'id': run['id']}, {'status': 'running', 'heartbeat_at': time.time() - 200, 'worker': 'fixture-dead', 'attempts': 1})
    result = finish(suite, run)
    assert result['status'] == 'succeeded' and result['attempts'] == 2
    assert not suite.api.state.dataprep.tick()
    outputs = suite.people.store.all('SELECT id FROM dataprep_datasets WHERE run_id=?', (run['id'],))
    assert len(outputs) == 1
    retry = op(suite, 'run.retry', {'run_id': run['id']})
    result2 = finish(suite, retry)
    assert result2['output_hash'] == result['output_hash'] and retry['id'] != run['id']


def test_revoke_queued_authority_stops_execution(suite):
    suite.sign_in('bob')
    run = start(suite, recipe(suite), dataset(suite))
    suite.people.store.execute('DELETE FROM memberships WHERE user_id=? AND workspace_id=?', (suite.people.uid('bob'), suite.people.a))
    assert suite.api.state.dataprep.tick()
    suite.sign_in('alice')
    result = op(suite, 'run.get', {'run_id': run['id']})
    assert result['status'] == 'failed' and result['output_dataset_id'] is None


def test_actual_upload_import_and_safe_exports(suite):
    suite.sign_in('bob')
    args = {'params': {'workspace_id': suite.people.a}, 'headers': {'X-Synwevo-CSRF': 'fixture-csrf', 'Idempotency-Key': 'fixture-upload'}}
    file = suite.post('/api/suite/dataprep/files', files={'file': ('fixture.csv', b'id,amount\n001,=1+1\n002,+cmd\n')}, **args)
    assert file.status_code == 200, file.text
    fid = file.json()['id']
    again = suite.post('/api/suite/dataprep/files', files={'file': ('fixture.csv', b'id,amount\n001,=1+1\n002,+cmd\n')}, **args)
    assert again.json()['id'] == fid
    d = op(suite, 'dataset.import', {'file_id': fid, 'name': 'Imported'})
    csv = suite.get(f"/api/suite/dataprep/datasets/{d['id']}/export", params={'workspace_id': suite.people.a}).text
    assert "'=1+1" in csv and "'+cmd" in csv
    xlsx = suite.get(f"/api/suite/dataprep/datasets/{d['id']}/export", params={'workspace_id': suite.people.a, 'format': 'xlsx'})
    cells = list(load_workbook(io.BytesIO(xlsx.content)).active.rows)
    assert cells[1][0].value == '001' and cells[1][1].data_type == 's'
    out = op(suite, 'dataset.get', {'dataset_id': d['id']})
    assert out['row_count'] == 2


def test_append_split_merge_filter_dedupe_parameter_and_validation():
    a = Table(['name', 'id'], [Row({'name': ' A B ', 'id': '01'}), Row({'name': 'A B', 'id': '01'})])
    b = Table(['id', 'name'], [Row({'name': 'C D', 'id': '02'})])
    definition = Definition.model_validate({'parameters': {'exclude': '03'}, 'steps': [
        {'id': 'append', 'type': 'append', 'source': 'other'},
        {'id': 'trim', 'type': 'trim', 'columns': ['name']},
        {'id': 'dedupe', 'type': 'dedupe', 'columns': ['id']},
        {'id': 'split', 'type': 'split', 'column': 'name', 'into': ['first', 'last'], 'separator': ' '},
        {'id': 'merge', 'type': 'merge', 'columns': ['last', 'first'], 'target': 'label', 'separator': ', '},
        {'id': 'filter', 'type': 'filter', 'condition': {'column': 'id', 'op': 'ne', 'value': '$exclude'}},
        {'id': 'sort', 'type': 'sort', 'columns': ['id'], 'descending': True},
        {'id': 'validate', 'type': 'validate', 'condition': {'column': 'first', 'op': 'eq', 'value': 'A'}, 'on_error': 'keep'}]})
    out, issues, _ = execute(definition, {'main': a, 'other': b}, {})
    assert [r.values['label'] for r in out.rows] == ['D, C', 'B, A']
    assert len(issues) == 1


def test_limits_no_silent_truncation_and_no_executable_steps():
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        Definition.model_validate({'steps': [{'id': 'unsafe', 'type': 'python', 'code': 'print(1)'}]})
    with pytest.raises(ServiceError) as e:
        read_file(b'id\n' + b'x\n' * 50_001, 'csv', ImportOptions(end_row=2), 'fixture')
    assert e.value.code == 'row_limit'
    a = Table(['key'], [Row({'key': 'x'}) for _ in range(250)])
    with pytest.raises(ServiceError) as e:
        execute(Definition.model_validate({'steps': [{'id': 'join', 'type': 'join', 'source': 'other', 'left': ['key'], 'right': ['key'], 'matches': 'all'}]}), {'main': a, 'other': a}, {})
    assert e.value.code == 'join_limit'


def test_running_cancel_and_permission_recheck_before_commit(suite, monkeypatch):
    import synwevo_suite.dataprep.service as module
    suite.sign_in('bob')
    run = start(suite, recipe(suite), dataset(suite))
    real = module.execute

    def cancel_during_transform(*args, **kwargs):
        output = real(*args, **kwargs)
        suite.people.store.update('dataprep_runs', {'id': run['id']}, {'cancel_requested': 1})
        return output

    monkeypatch.setattr(module, 'execute', cancel_during_transform)  # isolated interruption injection
    result = finish(suite, run)
    assert result['status'] == 'cancelled'
    assert suite.people.store.all('SELECT id FROM dataprep_datasets WHERE run_id=?', (run['id'],)) == []


def test_worker_fencing_prevents_stale_publish(suite, monkeypatch):
    import synwevo_suite.dataprep.service as module
    suite.sign_in('bob')
    run = start(suite, recipe(suite), dataset(suite))
    real = module.execute

    def replace_lease(*args, **kwargs):
        result = real(*args, **kwargs)
        suite.people.store.update('dataprep_runs', {'id': run['id']}, {'worker': 'fixture-new-worker'})
        return result

    monkeypatch.setattr(module, 'execute', replace_lease)  # isolated competing worker fixture
    result = finish(suite, run)
    assert result['status'] == 'running' and result['output_dataset_id'] is None
    assert suite.people.store.all('SELECT id FROM dataprep_datasets WHERE run_id=?', (run['id'],)) == []


def test_large_scope_performance_and_all_bad_rows_retained(suite):
    suite.sign_in('bob')
    count = 20_000
    content = ('id,amount\n' + ''.join(f'{n:06},{"bad" if n % 100 == 0 else "0.01"}\n' for n in range(count))).encode()
    upload = suite.post('/api/suite/dataprep/files', params={'workspace_id': suite.people.a},
                        headers={'X-Synwevo-CSRF': 'fixture-csrf'}, data={'upload_key': 'fixture-large'},
                        files={'file': ('fixture-large.csv', content)})
    assert upload.status_code == 200
    d = op(suite, 'dataset.import', {'file_id': upload.json()['id'], 'name': 'Large fixture'})
    r = recipe(suite, [{'id': 'cast', 'type': 'cast', 'column': 'amount', 'to': 'decimal'}])
    began = time.perf_counter()
    run = finish(suite, start(suite, r, d))
    elapsed = time.perf_counter() - began
    assert run['status'] == 'partial' and run['counts']['output_rows'] == 19_800 and run['issue_total'] == 200
    page = op(suite, 'run.issues', {'run_id': run['id'], 'offset': 150, 'limit': 50})
    assert len(page['items']) == 50 and page['next_offset'] is None
    report = suite.get(f"/api/suite/dataprep/runs/{run['id']}/issues", params={'workspace_id': suite.people.a})
    assert len(report.json()) == 200 and report.json()[-1]['source_row'] == 19_902
    print(f'\nDataPrep 20,000 rows / 2 columns / 200 issues: {elapsed:.3f}s, CSV {len(content)} bytes')


def test_xlsx_missing_formula_cache_dates_and_untrusted_xml():
    book = Workbook()
    sheet = book.active
    from datetime import datetime
    sheet.append(['date', 'id', 'calculated'])
    sheet.append([datetime(2026, 9, 30, 12, 34, 56, 789000), '0007', '=2+2'])
    content = io.BytesIO(); book.save(content)
    table, _ = read_file(content.getvalue(), 'xlsx', ImportOptions(formulas='cached'), 'fixture')
    assert not table.rows and table.issues[0]['code'] == 'formula_cache_missing'
    table, _ = read_file(content.getvalue(), 'xlsx', ImportOptions(end_column=2), 'fixture')
    assert table.rows[0].values == {'date': '2026-09-30T12:34:56.789000', 'id': '0007'}
    src = ZipFile(io.BytesIO(content.getvalue())); target = io.BytesIO()
    with ZipFile(target, 'w') as dest:
        for entry in src.infolist():
            data = src.read(entry.filename)
            if entry.filename == 'xl/workbook.xml':
                data = ('<?xml version="1.0" encoding="utf-16"?><!DOCTYPE x [<!ENTITY bomb "bad">]><x>&bomb;</x>').encode('utf-16')
            dest.writestr(entry, data)
    with pytest.raises(ServiceError) as e:
        read_file(target.getvalue(), 'xlsx', ImportOptions(), 'fixture')
    assert e.value.code == 'invalid_workbook'


def test_every_write_has_guard_and_rechecks_revoked_workspace(suite):
    from types import SimpleNamespace
    suite.sign_in('bob')
    d = dataset(suite)
    operation = suite.caps.registry.get('dataprep.dataset.from_records')
    history = suite.people.store.one("SELECT * FROM op_runs WHERE op='dataprep.dataset.from_records' AND user_id=?", (suite.people.uid('bob'),))
    call = SimpleNamespace(workspace_id=suite.people.a, user_id=suite.people.uid('bob'), op=operation)
    before = suite.people.store.one('SELECT COUNT(*) AS n FROM op_runs')['n']
    for operation_ in suite.caps.registry.operations('dataprep'):
        if operation_.side_effect != 'none':
            assert callable(operation_.result_guard)
    operation.result_guard(call, history)
    assert suite.people.store.one('SELECT COUNT(*) AS n FROM op_runs')['n'] == before
    suite.people.store.execute("UPDATE memberships SET role='viewer' WHERE user_id=? AND workspace_id=?", (call.user_id, call.workspace_id))
    with pytest.raises(ServiceError):
        operation.result_guard(call, history)
    assert suite.people.store.one('SELECT COUNT(*) AS n FROM dataprep_datasets')['n'] == 1
    assert suite.people.store.one('SELECT COUNT(*) AS n FROM op_runs')['n'] == before


def test_guard_rechecks_source_and_target_and_fails_closed_without_inspect(suite):
    from types import SimpleNamespace
    suite.sign_in('bob')
    d = dataset(suite)
    operation = suite.caps.registry.get('dataprep.dataset.from_records')
    history = suite.people.store.one("SELECT * FROM op_runs WHERE op='dataprep.dataset.from_records' AND user_id=?", (suite.people.uid('bob'),))
    call = SimpleNamespace(workspace_id=suite.people.a, user_id=suite.people.uid('bob'), op=operation)
    # Isolated provenance fixture for a source dependency; no fake production reader.
    suite.people.store.update('dataprep_datasets', {'id': d['id']}, {'origin': 'form', 'source_json': json.dumps({'submission_ids': ['fsub_000000000000']})})
    with pytest.raises(ServiceError):
        operation.result_guard(call, history)
    with pytest.raises(ServiceError):
        operation.result_guard(SimpleNamespace(workspace_id=suite.people.b, user_id=call.user_id, op=operation), history)


def test_shared_history_and_replay_gate_after_revocation(suite):
    from synwevo_core.capabilities import Operation
    if 'result_guard' not in Operation.__dataclass_fields__:
        pytest.skip('Coordinator-owned Operation.result_guard/history integration is not in this worktree yet')
    suite.sign_in('bob')
    payload = {'name': 'Guard fixture', 'columns': ['id'], 'records': [{'id': '001'}]}
    d = op(suite, 'dataset.from_records', payload, key='guard-fixture')
    suite.people.store.update('dataprep_datasets', {'id': d['id']}, {'origin': 'form', 'source_json': json.dumps({'submission_ids': ['fsub_000000000000']})})
    response = suite.post('/api/suite/ops/dataprep.dataset.from_records', headers={'X-Synwevo-CSRF': 'fixture-csrf'},
                          json={'workspace_id': suite.people.a, 'input': payload, 'idempotency_key': 'guard-fixture'})
    assert response.status_code in (403, 404, 409), response.text
    historic = suite.people.store.one("SELECT id FROM op_runs WHERE op='dataprep.dataset.from_records' AND idempotency_key='guard-fixture'")
    history_response = suite.get('/api/suite/runs/' + historic['id'])
    # Core returns a redacted execution receipt so the caller can reconcile without
    # repeating the write. Assert the actual privacy contract, including its marker.
    assert history_response.status_code == 200
    receipt = history_response.json()
    assert receipt['snapshot_hidden'] is True
    assert receipt['input'] is None and receipt['output'] is None and receipt['object'] is None
    assert suite.people.store.one('SELECT COUNT(*) AS n FROM dataprep_datasets')['n'] == 1


def test_forms_import_uses_real_registered_submission(suite):
    suite.sign_in('bob')
    def forms(name, input):
        response = suite.post('/api/suite/ops/forms.' + name, headers={'X-Synwevo-CSRF': 'fixture-csrf'}, json={'workspace_id': suite.people.a, 'input': input})
        assert response.status_code == 200, response.text
        return response.json()['output']
    form = forms('form.create', {'title': 'Fixture source', 'definition': {'fields': [{'id': 'f_amount01', 'type': 'text', 'label': 'ID'}]}})
    forms('form.publish', {'form_id': form['id'], 'draft_rev': 1})
    request = forms('request.create', {'form_id': form['id']})
    token = request['url'].rsplit('/', 1)[1]
    response = suite.post(f'/api/suite/public/forms/r/{token}/submit', headers={'X-Synwevo-CSRF': 'fixture-csrf'},
                          json={'version': 1, 'answers': {'f_amount01': '0001'}, 'idempotency_key': 'fixture-submit-dataprep'})
    assert response.status_code == 200, response.text
    sub = suite.people.store.one('SELECT id FROM form_submissions WHERE workspace_id=?', (suite.people.a,))
    d = op(suite, 'dataset.from_forms', {'name': 'Form snapshot', 'submission_ids': [sub['id']]})
    table = suite.api.state.dataprep.load(suite.people.a, d['id'])
    assert table.rows[0].values == {'f_amount01': '0001'}
    assert table.rows[0].origins == [[sub['id'], 1]]
    assert d['source']['answer_fingerprints'][sub['id']]


def test_fractional_precision_rejection_rounding_and_large_sum():
    definition = Definition.model_validate({'steps': [{'id': 'money', 'type': 'cast', 'column': 'v', 'to': 'decimal', 'scale': 2}]})
    table = Table(['v'], [Row({'v': v}) for v in ['1.005', '2.675', '123456789012345678901234567890.12']])
    out, issues, _ = execute(definition, {'main': table}, {})
    assert len(issues) == 2 and all(i['code'] == 'precision_loss' for i in issues)
    assert out.rows[0].values['v'] == '123456789012345678901234567890.12'
    definition.steps[0].rounding = 'half_even'
    out, issues, _ = execute(definition, {'main': table}, {})
    assert [r.values['v'] for r in out.rows[:2]] == ['1.00', '2.68']
    large = Table(['v'], [Row({'v': '123456789012345678901234567890.12'}) for _ in range(2)])
    out, _, _ = execute(Definition.model_validate({'steps': [{'id': 'sum', 'type': 'aggregate', 'measures': [{'op': 'sum', 'column': 'v', 'target': 'total'}]}]}), {'main': large}, {})
    assert out.rows[0].values['total'] == '246913578024691357802469135780.24'


def test_count_is_an_exact_integer_while_sum_min_max_keep_their_decimal_text():
    """Release acceptance: 分组汇总 with 行数 read "4.000000" in the preview, the result grid and
    all three exports. A count counts rows; it is an integer, never a formatted float."""
    source = Table(['region', 'amount'], [Row({'region': r, 'amount': a}, [['fixture', n + 2]]) for n, (r, a) in
                                          enumerate([('N', '0.10'), ('N', '0.20'), ('S', None), ('N', '1.00')])])
    definition = Definition.model_validate({'steps': [{'id': 'totals', 'type': 'aggregate', 'group_by': ['region'], 'measures': [
        {'op': 'count', 'target': 'orders'}, {'op': 'sum', 'column': 'amount', 'target': 'total'},
        {'op': 'min', 'column': 'amount', 'target': 'low'}, {'op': 'max', 'column': 'amount', 'target': 'high'}]}]})
    out, issues, _ = execute(definition, {'main': source}, {})
    assert not issues
    assert [r.values for r in out.rows] == [
        {'region': 'N', 'orders': '3', 'total': '1.30', 'low': '0.10', 'high': '1.00'},
        {'region': 'S', 'orders': '1', 'total': None, 'low': None, 'high': None}]
    many = Table(['id'], [Row({'id': str(n)}) for n in range(1234)])
    out, _, _ = execute(Definition.model_validate({'steps': [{'id': 'n', 'type': 'aggregate', 'measures': [{'op': 'count', 'target': 'n'}]}]}), {'main': many}, {})
    assert out.rows[0].values == {'n': '1234'}


def test_count_reads_as_an_integer_in_preview_result_and_every_export(suite):
    suite.sign_in('bob')
    d = dataset(suite, [{'id': '001', 'amount': '0.10'}, {'id': '001', 'amount': '0.20'}, {'id': '002', 'amount': '5'}])
    steps = [{'id': 'totals', 'type': 'aggregate', 'group_by': ['id'], 'measures': [
        {'op': 'count', 'target': 'orders'}, {'op': 'sum', 'column': 'amount', 'target': 'total'}]}]
    preview = op(suite, 'recipe.preview', {'definition': {'steps': steps}, 'inputs': {'main': d['id']}})
    assert [r['values'] for r in preview['rows']] == [{'id': '001', 'orders': '2', 'total': '0.30'}, {'id': '002', 'orders': '1', 'total': '5'}]
    run = finish(suite, start(suite, recipe(suite, steps), d))
    assert run['status'] == 'succeeded'
    grid = op(suite, 'dataset.get', {'dataset_id': run['output_dataset_id']})
    assert [r['values']['orders'] for r in grid['rows']] == ['2', '1']
    url, params = f"/api/suite/dataprep/datasets/{run['output_dataset_id']}/export", {'workspace_id': suite.people.a}
    csv = suite.get(url, params=params).text
    assert csv.splitlines()[1:] == ['"001","2","0.30"', '"002","1","5"']
    cells = [[c.value for c in row] for row in load_workbook(io.BytesIO(suite.get(url, params={**params, 'format': 'xlsx'}).content)).active.rows]
    assert cells[1:] == [['001', '2', '0.30'], ['002', '1', '5']]
    exported = suite.get(url, params={**params, 'format': 'json'}).json()
    assert [r['values']['orders'] for r in exported['rows']] == ['2', '1']


def test_isolated_import_roots():
    from pathlib import Path
    import synwevo_core.store
    import synwevo_suite.dataprep
    import synwevo_auth_client
    tree = Path(__file__).resolve().parents[3]
    assert Path(synwevo_core.store.__file__).is_relative_to(tree / 'packages/core')
    assert Path(synwevo_suite.dataprep.__file__).is_relative_to(tree / 'apps/suite')
    assert Path(synwevo_auth_client.__file__).is_relative_to(tree / 'packages/auth-client/python')


def test_a_runs_correlation_and_operation_run_reach_only_its_starter(suite):
    """Review R5: dataprep.run.list and .get returned whole rows, so every member who can read
    DataPrep saw the starter's correlation — their own tracing text, or their own run id — and
    their operation run id, which the run receipts hide from anyone else."""
    c = suite
    c.sign_in('bob')
    d = dataset(c)
    r = recipe(c)
    start = {'recipe_id': r['id'], 'version': 1, 'inputs': {'main': d['id']}}
    response = c.post('/api/suite/ops/dataprep.run.start', headers={'X-Synwevo-CSRF': 'fixture-csrf'},
                      json={'workspace_id': c.people.a, 'input': start, 'correlation_id': 'fixture-bob-private-trace'})
    assert response.status_code == 200, response.text
    named = response.json()['output']
    # the starter still reads his own
    mine = op(c, 'run.get', {'run_id': named['id']})
    assert mine['correlation_id'] == 'fixture-bob-private-trace' and mine['op_run_id']
    c.sign_in('alice')
    listed = op(c, 'run.list', {'offset': 0, 'limit': 20})['items']
    assert listed and all('correlation_id' not in x and 'op_run_id' not in x for x in listed), listed
    got = op(c, 'run.get', {'run_id': named['id']})
    assert got['id'] == named['id'] and 'correlation_id' not in got and 'op_run_id' not in got
