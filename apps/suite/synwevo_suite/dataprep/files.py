"""CSV and OOXML readers with explicit ranges and lexical numeric preservation.

XLSX is parsed as bounded XML instead of passing financial values through binary floats. Only
Excel date serials use openpyxl's date conversion (millisecond resolution, workbook epoch).
Formula expressions and macros are never executed; cached values require an explicit opt-in.
"""
from __future__ import annotations

import csv
import io
import posixpath
from datetime import datetime
from decimal import Decimal
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET

from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils.datetime import from_excel, CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

from .engine import Row, Table, bounds, refuse
from .models import ImportOptions, MAX_EXPANDED, MAX_FILE, MAX_ROWS, MAX_COLUMNS, MAX_CELLS, MAX_TEXT

S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def xml(data):
    if b'\x00' in data or b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        refuse('invalid_workbook')
    try:
        return ET.fromstring(data)
    except ET.ParseError:
        refuse('invalid_workbook')


class Workbook:
    def __init__(self, data):
        try:
            self.zip = ZipFile(io.BytesIO(data))
            entries = self.zip.infolist()
            if (len(entries) > 500 or sum(i.file_size for i in entries) > MAX_EXPANDED
                    or any(i.flag_bits & 1 or (not i.filename.startswith('xl/worksheets/') and i.file_size > 8 * 1024 * 1024) for i in entries)):
                refuse('workbook_limit')
            if len({i.filename for i in entries}) != len(entries):
                refuse('invalid_workbook')
            if any('vbaproject' in i.filename.lower() or '/externallinks/' in i.filename.lower() for i in entries):
                refuse('active_workbook')
            book = xml(self.zip.read('xl/workbook.xml'))
            rels = xml(self.zip.read('xl/_rels/workbook.xml.rels'))
            paths = {}
            for rel in rels:
                if rel.get('TargetMode') == 'External':
                    continue
                target = rel.get('Target', '')
                path = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
                if not path.startswith('xl/'):
                    refuse('invalid_workbook')
                paths[rel.get('Id')] = path
            self.sheets = [(s.get('name'), paths.get(s.get(R + 'id'))) for s in book.findall(S + 'sheets/' + S + 'sheet')]
            if not self.sheets or len(self.sheets) > 50 or any(not n or not p for n, p in self.sheets):
                refuse('invalid_workbook')
            props = book.find(S + 'workbookPr')
            self.epoch = CALENDAR_MAC_1904 if props is not None and props.get('date1904') in ('1', 'true') else CALENDAR_WINDOWS_1900
            self.strings = []
            if 'xl/sharedStrings.xml' in self.zip.namelist():
                for si in xml(self.zip.read('xl/sharedStrings.xml')).findall(S + 'si'):
                    self.strings.append(''.join(t.text or '' for t in si.iter(S + 't')))
                    if len(self.strings) > MAX_CELLS or len(self.strings[-1]) > MAX_TEXT:
                        refuse('workbook_limit')
            self.date_styles = set()
            if 'xl/styles.xml' in self.zip.namelist():
                styles = xml(self.zip.read('xl/styles.xml'))
                formats = {**BUILTIN_FORMATS, **{int(e.get('numFmtId')): e.get('formatCode', '') for e in styles.findall(S + 'numFmts/' + S + 'numFmt')}}
                self.date_styles = {i for i, xf in enumerate(styles.findall(S + 'cellXfs/' + S + 'xf'))
                                    if is_date_format(formats.get(int(xf.get('numFmtId', '0')), ''))}
        except (BadZipFile, KeyError, ValueError, OverflowError, RuntimeError, TypeError):
            refuse('invalid_workbook')

    def rows(self, sheet, options):
        target = next((p for n, p in self.sheets if n == sheet), None)
        if target is None:
            refuse('sheet_missing')
        payload = self.zip.read(target)
        if b'\x00' in payload or b'<!DOCTYPE' in payload.upper() or b'<!ENTITY' in payload.upper():
            refuse('invalid_workbook')
        seen, cells = 0, 0
        try:
            for _, e in ET.iterparse(io.BytesIO(payload), events=('end',)):
                if e.tag != S + 'row':
                    continue
                rn = int(e.get('r', '0'))
                if rn <= seen or rn > MAX_ROWS + 1:
                    refuse('row_limit')
                seen = rn
                values, problems = {}, []
                for c in e.findall(S + 'c'):
                    cells += 1
                    if cells > MAX_CELLS + MAX_COLUMNS:
                        refuse('table_limit')
                    letters, cr = coordinate_from_string(c.get('r', ''))
                    col = column_index_from_string(letters)
                    if cr != rn or col in values:
                        refuse('invalid_workbook')
                    # Only the selected cells' formula policy matters.
                    if col < options.start_column or (options.end_column and col > options.end_column):
                        continue
                    if rn < options.start_row or (options.end_row and rn > options.end_row):
                        continue
                    formula, value, typ = c.find(S + 'f'), c.find(S + 'v'), c.get('t', 'n')
                    raw = value.text if value is not None else None
                    if formula is not None and options.formulas == 'reject':
                        refuse('formula_rejected', {'row': rn, 'column': col})
                    if formula is not None and raw is None:
                        problems.append((col, 'formula_cache_missing'))
                        values[col] = None
                        continue
                    if typ == 'e':
                        problems.append((col, 'spreadsheet_error'))
                        values[col] = raw
                    elif typ == 's':
                        idx = int(raw)
                        if not 0 <= idx < len(self.strings):
                            refuse('invalid_workbook')
                        values[col] = self.strings[idx]
                    elif typ == 'inlineStr':
                        values[col] = ''.join(t.text or '' for t in c.iter(S + 't'))
                    elif typ == 'b':
                        values[col] = 'true' if raw == '1' else 'false'
                    elif raw is not None and int(c.get('s', '0')) in self.date_styles and typ == 'n':
                        try:
                            serial = Decimal(raw)
                            # Excel's fictitious 1900-02-29 is not silently mapped onto February 28.
                            if self.epoch == CALENDAR_WINDOWS_1900 and 60 <= serial < 61:
                                raise ValueError()
                            converted = from_excel(float(serial), self.epoch)
                            values[col] = converted.isoformat()
                        except (ValueError, OverflowError):
                            values[col] = raw
                            problems.append((col, 'invalid_date'))
                    else:
                        values[col] = raw
                yield rn, values, problems
                e.clear()
        except (ET.ParseError, ValueError, KeyError, IndexError, TypeError, OverflowError):
            refuse('invalid_workbook')


def inspect(data: bytes, format: str):
    if not data or len(data) > MAX_FILE:
        refuse('file_limit')
    return {'sheets': [n for n, _ in Workbook(data).sheets] if format == 'xlsx' else [],
            'max_file_bytes': MAX_FILE, 'max_rows': MAX_ROWS, 'max_columns': MAX_COLUMNS,
            'max_cells': MAX_CELLS, 'encodings': ['utf-8-sig', 'utf-16', 'gb18030']}


def read_file(data: bytes, format: str, options: ImportOptions, source: str):
    inspect(data, format)
    notes = []
    if format == 'xlsx':
        book = Workbook(data)
        sheet = options.sheet or book.sheets[0][0]
        incoming = book.rows(sheet, options)
        notes = ['xlsx_dates_millisecond_local', 'xlsx_numbers_lexical']
        if options.formulas == 'cached':
            notes.append('formula_cache_not_recalculated')
    else:
        try:
            text = data.decode(options.encoding, errors='strict')
        except UnicodeError:
            refuse('invalid_encoding')
        if '\x00' in text:
            refuse('invalid_encoding')
        reader = csv.reader(io.StringIO(text, newline=''), delimiter=options.delimiter, strict=True)

        def csv_rows():
            try:
                for n, row in enumerate(reader, 1):
                    if n > MAX_ROWS + 1:
                        refuse('row_limit')
                    if len(row) > MAX_COLUMNS:
                        refuse('table_limit')
                    yield n, {i + 1: v for i, v in enumerate(row)}, []
            except csv.Error:
                refuse('invalid_csv')
        incoming = csv_rows()
    header, columns, out, issues = None, [], [], []
    for rn, values, problems in incoming:
        if rn < options.start_row or (options.end_row is not None and rn > options.end_row):
            continue
        end = options.end_column or (max(values, default=0) if header is None else header[-1])
        if end - options.start_column + 1 > MAX_COLUMNS:
            refuse('table_limit')
        selected = list(range(options.start_column, end + 1))
        if header is None:
            header = selected
            columns = [values.get(c) for c in header]
            if problems or not columns or any(not isinstance(c, str) or not c for c in columns) or len(set(columns)) != len(columns):
                refuse('invalid_columns')
            continue
        row_problems = list(problems)
        if options.end_column is None and any(c > end for c in values):
            row_problems.append((None, 'row_width'))
        vals = {name: None if values[c] in options.null_tokens else values[c] for c, name in zip(header, columns) if c in values}
        for col, code in row_problems:
            issues.append({'severity': 'error', 'code': code, 'input_name': source, 'source_row': rn,
                           'column_name': columns[header.index(col)] if col in header else None, 'message': code})
        if row_problems:
            continue
        out.append(Row(vals, [[source, rn]]))
        if len(out) * len(columns) > MAX_CELLS:
            refuse('table_limit')
    if not columns:
        refuse('invalid_columns')
    table = bounds(Table(columns, out, issues))
    return table, notes


def export_csv(table: Table):
    """All fields quoted. Prefix potential spreadsheet formulas, including whitespace/control
    prefixes. CSV is a display interchange: null/missing become empty; JSON export is lossless."""
    def safe(value):
        text = value or ''
        return "'" + text if text.lstrip(''.join(chr(c) for c in range(33)) + '\ufeff').startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')) else text
    output = io.StringIO(newline='')
    writer = csv.writer(output, quoting=csv.QUOTE_ALL)
    writer.writerow([safe(c) for c in table.columns])
    for row in table.rows:
        writer.writerow([safe(row.values.get(c)) for c in table.columns])
    return ('\ufeff' + output.getvalue()).encode('utf-8')


def export_xlsx(table: Table):
    from openpyxl import Workbook as Writer
    from openpyxl.cell import WriteOnlyCell
    book = Writer(write_only=True)
    sheet = book.create_sheet('Data')
    # Explicit string cells prevent formula execution and preserve identifiers and precision.
    from itertools import chain
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if any(ILLEGAL_CHARACTERS_RE.search(v) for v in table.columns + [v for r in table.rows for v in r.values.values() if v]):
        refuse('xlsx_invalid_char')
    for values in chain([table.columns], ([r.values.get(c) for c in table.columns] for r in table.rows)):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = 's'
            cells.append(cell)
        sheet.append(cells)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
