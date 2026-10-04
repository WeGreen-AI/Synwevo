"""Bounded CSV/XLSX staging and fresh-authorized export tickets. No stored blob references.

A ticket belongs to whoever made it: the person in the page, or one client acting for them. It is
stored under its maker (`stored_ticket`), so the same id asked for by the page, or by another
client the person also authorised, names nothing — the rule run receipts and idempotency keys
already follow.
"""
import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from types import SimpleNamespace
from urllib.parse import urlencode

from fastapi import Request, UploadFile, File
from fastapi.responses import Response
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import Workbook, load_workbook

from synwevo_core.capabilities import canonical
from synwevo_core.store import new_id, now
from synwevo_core.errors import ServiceError
from . import fields as ft
from .models import Batch, BatchItem, Query
from .service import refuse, unpack

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 500
_TICKET = re.compile(r'(imp|exp)_[a-f0-9]{12}')


def stored_ticket(c, ticket_id, missing):
    """A ticket's id as it is stored: as issued for a person, and under the client's name for a
    credential (`caller_key`, as its idempotency keys are). Anything that is not an id this module
    issued — a stored id passed back in, say — is the same 404 as a ticket that does not exist."""
    if not isinstance(ticket_id, str) or not _TICKET.fullmatch(ticket_id):
        refuse(missing, 404)
    if getattr(c, 'via', 'ui') != 'api':
        return ticket_id
    client = (getattr(c, 'delegation', None) or {}).get('client_id')
    if not client:
        refuse(missing, 404)
    from ..common.machine import caller_key
    return caller_key(client, ticket_id)


def _fold(name):
    return unicodedata.normalize('NFKC', name).strip().casefold()


def link_cell(field, raw, column, names):
    """A related-records cell of an import file, as record ids. Each related record is named the way
    export writes it (its display name; several are joined by commas) or by its id. A name matches
    exactly, else ignoring case and width. A name that matches no candidate or several, or an id
    that is no candidate, is refused for this row, naming the column and the value."""
    ids, exact, folded = names
    text = raw.strip()
    if not text:
        return None

    def one(name):
        if ft.RECORD_ID.match(name):
            hits = [name] if name in ids else []
        else:
            hits = exact.get(name) or folded.get(_fold(name), [])
        if len(hits) != 1:
            refuse('link_ambiguous' if hits else 'link_not_found', column=column, value=name[:200])
        return hits[0]
    # A name may itself contain a comma: the whole cell is one name when it names a candidate.
    whole = ft.RECORD_ID.match(text) or text in exact or _fold(text) in folded
    parts = [text] if whole or ',' not in text else [p.strip() for p in text.split(',') if p.strip()]
    return ft.normalize(field, [one(p) for p in parts])


class Transfers:
    def sweep(self):
        with self.store.tx():
            self.store.execute('DELETE FROM tb_imports WHERE expires_at<?', (now(),))
            self.store.execute('DELETE FROM tb_exports WHERE expires_at<?', (now(),))

    def stage_import(self, c, tid, filename, content):
        self.table(c, tid, write=True, active=True)
        self.sweep()
        if len(content) > MAX_FILE_BYTES:
            refuse('file_too_large', 413, limit=MAX_FILE_BYTES)
        fmt = filename.rsplit('.', 1)[-1].lower()
        if fmt not in ('csv', 'xlsx'):
            refuse('unsupported_format')
        rows = []
        try:
            if fmt == 'csv':
                reader = csv.reader(io.StringIO(content.decode('utf-8-sig')), strict=True)
                for row in reader:
                    rows.append(row)
                    if len(rows) > MAX_IMPORT_ROWS+1:
                        refuse('import_row_limit', 413, limit=MAX_IMPORT_ROWS)
            else:
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    if len(z.infolist()) > 500 or sum(x.file_size for x in z.infolist()) > 16*1024*1024:
                        refuse('file_too_large', 413)
                book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
                try:
                    if len(book.worksheets) != 1:
                        refuse('one_sheet_required')
                    sheet = book.worksheets[0]
                    if (sheet.max_column or 0) > 100 or (sheet.max_row or 0) > MAX_IMPORT_ROWS+1:
                        refuse('import_row_limit', 413, limit=MAX_IMPORT_ROWS)
                    for row in sheet.iter_rows():
                        if any(cell.data_type == 'f' for cell in row):
                            refuse('file_formula_not_allowed')
                        values = []
                        for cell in row:
                            value = cell.value
                            if isinstance(value, (date, datetime)):
                                value = value.isoformat()
                            values.append('' if value is None else str(value))
                        rows.append(values)
                        if len(rows) > MAX_IMPORT_ROWS+1:
                            refuse('import_row_limit', 413, limit=MAX_IMPORT_ROWS)
                finally:
                    book.close()
        except (UnicodeError, csv.Error, zipfile.BadZipFile, ValueError, KeyError, OSError, ParseError, InvalidFileException):
            refuse('invalid_file')
        if not rows or not rows[0] or len(rows[0]) > 100:
            refuse('invalid_file')
        columns = [str(v).strip() for v in rows[0]]
        if any(not x or len(x) > 200 for x in columns) or len(columns) != len(set(columns)):
            refuse('import_headers')
        data = []
        for n, row in enumerate(rows[1:], 2):
            if len(row) > len(columns) or any(len(str(v)) > ft.LONG_TEXT_MAX for v in row):
                refuse('invalid_file_row', row=n)
            if any(v != '' for v in row):
                data.append(row + [''] * (len(columns)-len(row)))
        if not data:
            refuse('import_empty')
        count = self.store.one('SELECT COUNT(*) AS n FROM tb_imports WHERE user_id=? AND workspace_id=?', (c.user_id, c.workspace_id))['n']
        if count >= 10:
            refuse('import_ticket_limit', 409)
        iid, ts = new_id('imp'), now()
        self.store.execute('INSERT INTO tb_imports (id,table_id,workspace_id,user_id,filename,format,columns_json,rows_json,row_count,status,created_at,expires_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                           (stored_ticket(c, iid, 'import_not_found'), tid, c.workspace_id, c.user_id, filename.replace('\\', '/').rsplit('/', 1)[-1][:200], fmt, canonical(columns), canonical(data), len(data), 'ready', ts, ts+3600))
        return {'id': iid, 'columns': columns, 'rows': data[:20], 'row_count': len(data), 'expires_at': ts+3600}

    def import_row(self, c, tid, iid):
        self.table(c, tid, write=True, active=True)
        row = self.store.one('SELECT * FROM tb_imports WHERE id=? AND table_id=? AND workspace_id=? AND user_id=? AND expires_at>?',
                             (stored_ticket(c, iid, 'import_not_found'), tid, c.workspace_id, c.user_id, now()))
        if not row:
            refuse('import_not_found', 404)
        return row

    def import_get(self, c, i):
        row = self.import_row(c, i.table_id, i.import_id)
        return {'id': i.import_id, 'columns': unpack(row['columns_json']), 'rows': unpack(row['rows_json'])[:20],
                'row_count': row['row_count'], 'status': row['status'], 'outcome': unpack(row['outcome_json']), 'expires_at': row['expires_at']}

    def link_names(self, c, field):
        """What an importer can name in `field`'s related table: the live records they can read,
        by id and by the name export writes for them (`record_title`). A record they cannot read is
        no candidate, so its name reads as unknown — never as ambiguous, never as existing."""
        table = self.table(c, field['config']['table_id'])
        if table['status'] != 'active':
            return set(), {}, {}
        rank = ft.ROLE_RANK[self.role(c)]
        allowed = [r for r in ft.READ_ROLES if ft.ROLE_RANK[r] <= rank]
        rows = self.store.all(f"SELECT id, seq, cells FROM tb_records WHERE table_id=? AND workspace_id=? AND deleted_at IS NULL AND read_role IN ({','.join('?' * len(allowed))})",
                              (table['id'], c.workspace_id, *allowed))
        title_field = self.title_field(c, table)
        ids, exact, folded = set(), {}, {}
        for r in rows:
            ids.add(r['id'])
            title = self.record_title(title_field, r).strip()
            if title:
                exact.setdefault(title, []).append(r['id'])
                folded.setdefault(_fold(title), []).append(r['id'])
        return ids, exact, folded

    def import_apply(self, c, i):
        row = self.import_row(c, i.table_id, i.import_id)
        settings = {'mapping': i.mapping, 'atomic': i.atomic}
        if row['status'] in ('done', 'partial'):
            if unpack(row['mapping_json']) != settings:
                refuse('import_already_applied', 409)
            from .guard import result_guard
            result_guard(self)(c, {'op': 'tables.import.apply', 'op_version': 1, 'user_id': c.user_id,
                                   'workspace_id': c.workspace_id, 'input': i.model_dump(), 'output': row['outcome_json']})
            return unpack(row['outcome_json'])
        columns, data = unpack(row['columns_json']), unpack(row['rows_json'])
        if not set(i.mapping) <= set(columns) or len(set(i.mapping.values())) != len(i.mapping):
            refuse('invalid_mapping')
        fields = self.fields(i.table_id)
        writable = {f['id']: f for f in self.visible_fields(c, fields) if f['can_write']}
        if not set(i.mapping.values()) <= writable.keys():
            refuse('field_forbidden', 403)
        column_of = {fid: col for col, fid in i.mapping.items()}

        def located(detail, column=None):
            """A refused row's detail, with the file column it is about when it is about one."""
            column = column or column_of.get((detail or {}).get('field_id'))
            return dict(detail or {}, **({'column': column} if column else {}))

        items, errors, names = [], [], {}
        for n, values in enumerate(data):
            cells = {}
            try:
                for col, fid in i.mapping.items():
                    f, raw = writable[fid], values[columns.index(col)]
                    if f['type'] == 'link':
                        if fid not in names:
                            names[fid] = self.link_names(c, f)
                        cells[fid] = link_cell(f, raw, col, names[fid])
                    else:
                        cells[fid] = ft.parse_text(f, raw)
                self.normalize_cells(c, fields, cells)
                items.append((n, BatchItem(action='create', cells=cells)))
            except ft.CellError as e:
                errors.append({'index': n, 'ok': False, 'error': 'invalid_cell', 'detail': located({'reason': e.code}, col)})
            except ServiceError as e:
                errors.append({'index': n, 'ok': False, 'error': e.code, 'detail': located(e.detail)})
        # Preview uses savepoints and the real writer, including uniqueness against existing rows
        # and previous input rows. It leaves no records, revisions or events behind.
        if i.preview:
            n_events = len(c.events)
            self.store.execute('SAVEPOINT tables_preview')
        try:
            if errors and i.atomic and not i.preview:
                refuse('import_validation', 409, issues=errors)
            outcomes = list(errors)
            for start in range(0, len(items), 100):
                chunk = items[start:start+100]
                result = self.batch(c, Batch(table_id=i.table_id, items=[v for _, v in chunk], atomic=i.atomic and not i.preview))
                outcomes.extend(dict(v, index=chunk[v['index']][0], **({} if v['ok'] else {'detail': located(v['detail'])}))
                                for v in result['items'])
            outcomes.sort(key=lambda v: v['index'])
            if i.preview:
                return {'preview': True, 'rows': len(data), 'valid': sum(x['ok'] for x in outcomes),
                        'issues': [v for v in outcomes if not v['ok']]}
            failed = sum(not x['ok'] for x in outcomes)
            output = {'atomic': i.atomic, 'items': outcomes, 'succeeded': len(data)-failed, 'failed': failed,
                      '__run_status__': 'partial' if failed else 'succeeded'}
            self.store.execute('UPDATE tb_imports SET status=?,mapping_json=?,outcome_json=?,run_id=?,finished_at=? WHERE id=?',
                               ('partial' if failed else 'done', canonical(settings), canonical(output), c.run_id, now(), row['id']))
            return output
        finally:
            if i.preview:
                self.store.execute('ROLLBACK TO tables_preview')
                self.store.execute('RELEASE tables_preview')
                del c.events[n_events:]

    def export_right(self, c):
        self.role(c)
        member = self.identity.membership(c.workspace_id, c.user_id)
        if not member or not (member['role'] in ('owner', 'admin', 'editor') or member.get('can_export')):
            refuse('export_not_allowed', 403)

    def export_create(self, c, i):
        self.export_right(c)
        resolved, _, _ = self.resolve_query(c, i)
        # Freeze the saved view definition, not a reference that someone can later widen.
        query = {k: v for k, v in resolved.model_dump().items() if k in Query.model_fields}
        self.sweep()
        count = self.store.one('SELECT COUNT(*) AS n FROM tb_exports WHERE user_id=? AND workspace_id=?', (c.user_id, c.workspace_id))['n']
        if count >= 50:
            refuse('export_ticket_limit', 409)
        eid, ts = new_id('exp'), now()
        self.store.execute('INSERT INTO tb_exports (id,table_id,workspace_id,user_id,format,query_json,run_id,created_at,expires_at) VALUES (?,?,?,?,?,?,?,?,?)',
                           (stored_ticket(c, eid, 'export_not_found'), i.table_id, c.workspace_id, c.user_id, i.format, canonical(query), c.run_id, ts, ts+600))
        c.touches('table', i.table_id)
        # Where its maker downloads it: the page's route for a person, the machine API's for a client.
        from synwevo_core.http import SUITE_V1
        base = SUITE_V1 if c.via == 'api' else '/api/suite'
        return {'id': eid, 'url': base+'/tables/exports/'+eid+'?'+urlencode({'workspace_id': c.workspace_id}),
                'expires_at': ts+600, 'policy': 'current_authorized_records_at_download'}

    def export_download(self, c, eid):
        self.export_right(c)
        row = self.store.one('SELECT * FROM tb_exports WHERE id=? AND workspace_id=? AND user_id=? AND expires_at>?',
                             (stored_ticket(c, eid, 'export_not_found'), c.workspace_id, c.user_id, now()))
        if not row:
            refuse('export_not_found', 404)
        rows, fields, _ = self.query_rows(c, Query.model_validate(unpack(row['query_json'])))
        values = [[f['name'] for f in fields]]
        for r in rows:
            values.append([ft.as_text(f, r['cells'].get(f['id']), titles={x['id']: x['title'] for x in r['relations'].get(f['id'], [])}) for f in fields])
        # Bound output memory and synchronous response cost explicitly.
        if sum(len(v.encode()) for r in values for v in r) > 16*1024*1024:
            refuse('export_too_large', 413)
        if row['format'] == 'csv':
            stream = io.StringIO(newline='')
            writer = csv.writer(stream)
            for r in values:
                writer.writerow(["'"+v if v.lstrip().startswith(('=', '+', '-', '@')) or v.startswith(('\t', '\r', '\n')) else v for v in r])
            content, mime = stream.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8'
        else:
            book = Workbook(write_only=True)
            sheet = book.create_sheet('Tables')
            from openpyxl.cell import WriteOnlyCell
            for r in values:
                cells = []
                for v in r:
                    cell = WriteOnlyCell(sheet, value=v)
                    cell.data_type = 's'
                    cells.append(cell)
                sheet.append(cells)
            stream = io.BytesIO()
            book.save(stream)
            content, mime = stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        self.store.execute('UPDATE tb_exports SET downloads=downloads+1 WHERE id=?', (row['id'],))
        return Response(content, media_type=mime, headers={'Content-Disposition': 'attachment; filename="tables.'+row['format']+'"',
                                                          'Cache-Control': 'no-store', 'X-Content-Type-Options': 'nosniff'})


def install_routes(ctx, service):
    from synwevo_core.http import SUITE_V1
    from ..common.machine import acting

    def context(who, workspace_id):
        return SimpleNamespace(user_id=who.user['id'], workspace_id=workspace_id, via=who.via,
                               delegation=who.invoke.get('delegation'))

    # Each route is served under the page's prefix (a session) and under the machine API's (a
    # credential holding the scope named here); `acting` decides which, the rest is one path.
    @ctx.api.post('/api/suite/tables/{table_id}/imports', tags=['tables'])
    @ctx.api.post(SUITE_V1 + '/tables/{table_id}/imports', tags=['tables'])
    async def upload(table_id: str, workspace_id: str, request: Request, file: UploadFile = File(...)):
        with acting(ctx, request, workspace_id, app='tables', scope='tables.write') as who:
            c = context(who, workspace_id)
            service.table(c, table_id, write=True, active=True)
            content = await file.read(MAX_FILE_BYTES+1)
            await file.close()
            with service.store.tx():
                return service.stage_import(c, table_id, file.filename or '', content)

    @ctx.api.get('/api/suite/tables/exports/{export_id}', tags=['tables'])
    @ctx.api.get(SUITE_V1 + '/tables/exports/{export_id}', tags=['tables'])
    def download(export_id: str, workspace_id: str, request: Request):
        with acting(ctx, request, workspace_id, app='tables', scope='tables.read') as who:
            with service.store.tx():
                return service.export_download(context(who, workspace_id), export_id)
