"""Opt-in, real Docs UI → engine → downloaded XLSX acceptance in all three locales.

    npm run build -w @synwevo/docs-web
    .venv/bin/python apps/docs/web/tests/browser_i18n_artifacts.py --run

Default: a temporary, test-marked, loopback Docs service, with UI_AUTH=none (no
claim of login/SSO coverage). No route mocking or business API fixture writes.
Every project/upload/field/rule/mapping/job is created through the actual UI.
For an existing test service add --base-url and optionally --storage-state from
a real login. Never saves cookies, auth headers, tokens or a browser trace.
Evidence and synthetic input files default to a new temporary directory.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import hashlib
import http.client
import io
import ipaddress
import json
from pathlib import Path
import re
import shutil
import secrets
import subprocess
import tempfile
import time
import traceback
from urllib.parse import urlsplit
from email import policy
from email.parser import Parser
from email.utils import parseaddr
import uuid
import zipfile

import openpyxl
from playwright.sync_api import expect, sync_playwright

from browser_review import load

REPO = Path(__file__).resolve().parents[4]
LOCALES = ("zh-Hans", "zh-Hant", "en")
ROWS = (("T001", "仅测试简体客户", 12.5), ("T002", "測試繁體客戶", 0), ("T003", "Test client 中文", -3.25))
STATIC = "仅测试：业务简体内容不可翻译"
LITERAL = "English 原文 <b>{count}</b>"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_inputs(folder):
    folder.mkdir(parents=True, exist_ok=True)
    src = folder / "仅测试来源.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "原始数据"
    ws.append(["编号", "姓名", "金额"])
    for row in ROWS:
        ws.append(row)
    wb.save(src)
    tpl = folder / "仅测试模板.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "原文工作表"
    ws.append([STATIC, None, None, LITERAL])
    ws.append(["编号", "旧编号"])
    ws.append(["姓名", "旧姓名"])
    ws.append(["金额", -999])
    ws.append(["双倍金额", "=B4*2"])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 28
    wb.save(tpl)
    return src, tpl


def labels():
    # Read the production catalog; assertions still target actual accessible UI labels.
    # Type-only imports disappear in Node; this does not load a browser fixture.
    js = """
      import {messages as docs} from './apps/docs/web/src/messages.ts';
      import {serverMessages} from './apps/docs/web/src/messages.server.ts';
      import {messages as common} from './packages/ui/src/messages.ts';
      console.log(JSON.stringify({...docs, ...serverMessages, ...common}));
    """
    return json.loads(subprocess.check_output(
        ["node", "--experimental-strip-types", "--disable-warning=ExperimentalWarning", "--input-type=module", "-e", js],
        cwd=REPO, text=True))


def translator(catalog, locale):
    def t(source, **params):
        assert source in catalog, f"Missing UI test resource: {source}"
        value = source if locale == "zh-Hans" else catalog[source][locale]
        if isinstance(value, dict):
            value = value["one" if params.get("count") == 1 else "other"]
        return re.sub(r"\{(\w+)\}", lambda m: str(params.get(m[1], m[0])), value)
    return t


def validate_base(base, env):
    parsed = urlsplit(base)
    try:
        loopback = ipaddress.ip_address(parsed.hostname or "").is_loopback
    except ValueError:
        loopback = parsed.hostname == "localhost" or (parsed.hostname or '').endswith('.localhost')
    if not loopback or parsed.scheme != "http" or parsed.username or parsed.path not in ("", "/") or parsed.query or parsed.fragment:
        raise SystemExit("Only an explicit local HTTP origin is permitted")
    # Do not honor ALLOW_LIVE_TESTS: the server-recorded marker is mandatory.
    if (parsed.hostname or '').endswith('.localhost'):
        # Health probe goes directly to loopback with the intended Host. No proxy/DNS changes.
        conn = http.client.HTTPConnection('127.0.0.1', parsed.port or 80, timeout=5)
        try:
            conn.request('GET', '/healthz', headers={'Host': parsed.netloc})
            response = conn.getresponse()
            health = json.loads(response.read()) if response.status == 200 else {}
        finally:
            conn.close()
    else:
        health = env.health(base)
    if health.get("instance") != "test":
        raise SystemExit("Refusing to create fixtures: /healthz must report instance=test")


@contextmanager
def target(args, env):
    if args.base_url:
        base = args.base_url.rstrip("/")
        validate_base(base, env)
        yield base, "existing test-marked service; real session state if supplied"
    else:
        with env.isolated() as instance:
            validate_base(instance.base, env)
            try:
                yield instance.base, "isolated real Docs service; UI_AUTH=none; no authentication coverage"
            finally:
                # Service log only belongs to our throwaway fixture, not an existing instance.
                log = instance.data_dir.parent / "server.log"
                if log.exists():
                    shutil.copyfile(log, args.evidence / "isolated-server.log")


def readback(raw, row):
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    assert openpyxl.load_workbook(io.BytesIO(raw), data_only=False).active['B5'].value == '=B4*2'
    assert wb.sheetnames == ["原文工作表"], wb.sheetnames
    ws = wb.active
    assert ws["A1"].value == STATIC
    assert ws["D1"].value == LITERAL
    got = [ws[f"B{i}"].value for i in (2, 3, 4, 5)]
    assert got == [*row, row[2] * 2], (row, got)
    assert isinstance(got[2], (int, float)), "Numeric business value was converted to UI text"
    return {"id": got[0], "name": got[1], "amount": got[2], "formula_result": got[3], "formula_preserved": True, "static_content_unchanged": True}


def register_fixture(browser, args, report):
    """One independent real account; all auth/verification/onboarding happens through UI.

    Passwords, file-mail tokens and the returned storage state stay only in memory.
    No screenshots or Playwright error strings from credential forms are persisted.
    """
    mail = args.mail_dir.resolve(strict=True)
    if not any(mail.is_relative_to(Path(p).resolve()) for p in ('/tmp', tempfile.gettempdir())) or mail.stat().st_mode & 0o077:
        raise SystemExit('File mail must be in a private temporary directory')
    email = args.resume_email or 'a-docs-' + uuid.uuid4().hex[:12] + '@example.test'
    password = 'Synthetic isolated lunar rivers ' + secrets.token_hex(12) + '!'
    replacement = 'Verified isolated lunar rivers ' + secrets.token_hex(12) + '!'
    context = browser.new_context(locale='en-US', viewport={'width': 1440, 'height': 1000})
    page = context.new_page()
    page.set_default_timeout(20000)
    report['registration'] = {'synthetic_email': email, 'stage': 'registration', 'status': 'running'}
    reg = report['registration']
    reg['requests'] = []
    page.on('response', lambda r: reg['requests'].append({'path': urlsplit(r.url).path, 'status': r.status}) if urlsplit(r.url).path.startswith('/api/') else None)
    try:
        if not args.resume_email:
            page.goto(args.portal_url + '/?lang=en')
            page.locator('a.pt-nav-register').click()
            page.wait_for_url(re.compile(re.escape(args.accounts_url) + r'/register'))
            expect(page.locator('html')).to_have_attribute('lang', 'en')
            page.locator('#register-name').fill('Synthetic Docs artifact acceptance')
            page.locator('#register-email').fill(email)
            page.locator('input[autocomplete=new-password]').fill(password)
            page.locator('input[name=intent][value=personal]').check()
            with page.expect_response(lambda r: urlsplit(r.url).path == '/api/auth/register' and r.request.method == 'POST') as registered:
                page.get_by_role('button', name='Create account and verify email', exact=True).click()
            assert registered.value.status == 202, 'Registration was not accepted'
            # The subsequent full-page navigation can evict Response.json() in Chrome.
            # Proof of file delivery is the actual recipient-bound mail read below.
        reg['stage'] = 'private file-mail verification'
        link = None
        deadline = time.monotonic() + 30
        while not link and time.monotonic() < deadline:
            for path in sorted(mail.glob('*.eml'), key=lambda p: p.stat().st_mtime, reverse=True):
                # Inspect only recipient headers until finding this run's own file mail.
                # Never read another agent's message body or verification token.
                with path.open() as fh:
                    header_lines = []
                    for line in fh:
                        header_lines.append(line)
                        if not line.strip():
                            break
                header = Parser(policy=policy.default).parsestr(''.join(header_lines), headersonly=True)
                if parseaddr(header.get('To', ''))[1] != email:
                    continue
                message = Parser(policy=policy.default).parsestr(path.read_text())
                body = message.get_body(preferencelist=('plain',))
                for match in re.finditer(r'https?://[^\s<>]+', body.get_content() if body else ''):
                    parsed = urlsplit(match[0])
                    if parsed.path == '/verify-email' and parsed.netloc == urlsplit(args.accounts_url).netloc and parsed.fragment:
                        link = match[0]
                        break
            if not link:
                page.wait_for_timeout(200)
        assert link, 'No private file-mail verification link for this run'
        page.goto(link)
        page.get_by_role('button', name='Verify this email', exact=True).click()
        page.locator('input[autocomplete=new-password]').fill(replacement)
        page.get_by_role('button', name='Set new password and verify email', exact=True).click()
        expect(page.get_by_role('heading', name='Email verified', exact=True)).to_be_visible()
        reg['stage'] = 'real login and personal onboarding'
        page.goto(args.accounts_url + '/login?next=%2Fonboarding&lang=en')
        page.locator('#login-email').fill(email)
        page.locator('input[autocomplete=current-password]').fill(replacement)
        page.get_by_role('button', name='Sign in', exact=True).click()
        page.wait_for_url('**/onboarding')
        page.get_by_role('button', name='Use personal workspace', exact=True).click()
        reg['stage'] = 'portal and Docs SSO launch'
        # Personal onboarding now completes the workspace-bound portal SSO handoff.
        # Require the actual landing; an extra goto/login would hide a broken handoff.
        page.wait_for_url(re.compile(re.escape(args.portal_url) + r'/workspace(?:\?.*)?$'))
        expect(page.locator('.pt-space.is-current')).to_have_count(1)
        page.locator('a.pt-app-card', has_text='Synwevo Docs').click()
        page.wait_for_url(re.compile(re.escape(args.base_url) + r'/'))
        expect(page.get_by_role('button', name='＋ New project', exact=True)).to_be_visible()
        me = context.request.get(args.base_url + '/api/auth/me')
        assert me.status == 200 and me.json()['email'] == email, 'Docs session is not our verified test account'
        reg.update(status='passed', stage='completed', workspace_id=me.json()['workspace_id'])
        return context.storage_state()  # Never write this to disk or logs.
    except Exception as error:
        reg['status'] = 'failed'
        reg['failure_type'] = type(error).__name__
        reg['page_path'] = urlsplit(page.url).path
        text = page.locator('body').inner_text()
        for secret in (password, replacement):
            text = text.replace(secret, '[REDACTED]')
        reg['visible_ui'] = text
        (args.evidence / 'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
        raise RuntimeError('Synthetic account preparation failed at ' + reg['stage'] + '; credential details suppressed') from None
    finally:
        context.close()


def flow(page, base, locale, catalog, inputs, folder, result):
    t = translator(catalog, locale)
    folder.mkdir(parents=True, exist_ok=True)
    step = lambda name: result.setdefault("steps", []).append(name)
    button = lambda source: page.get_by_role("button", name=t(source), exact=True)
    cell = lambda pane, value: page.locator(f".pane.{pane} .gcell", has_text=re.compile(f"^{re.escape(str(value))}$")).first
    snap = lambda name: page.screenshot(path=str(folder / f"{name}.png"), full_page=True)
    page.goto(base + "/?lang=" + locale)
    # A signed-in account hydrates its saved language after /me. Select through the
    # live UI after RequireSession has mounted Home, then verify the requested locale.
    page.locator('.language-menu select').first.select_option(locale)
    expect(page.locator("html")).to_have_attribute("lang", locale)
    expect(page).to_have_title(t("Synwevo Docs · 文档工作台"))
    button("＋ 新建项目").click()
    page.locator(".modal input").fill("仅测试三语言生成_" + locale)
    page.locator(".modal").get_by_role("button", name=t("确定"), exact=True).click()
    page.wait_for_url(re.compile(r"/p/p_"))
    result["project_id"] = page.url.rsplit("/", 1)[-1]
    step("project created through UI")
    page.locator(".side input[type=file]").set_input_files(str(inputs[0]))
    expect(page.locator(".fileitem", has_text=inputs[0].name).get_by_text(t("可用"), exact=True)).to_be_visible(timeout=30000)
    page.get_by_role("button", name=t("设置模板"), exact=True).first.click()
    page.locator(".modal input[type=file]").set_input_files(str(inputs[1]))
    expect(cell("right", STATIC)).to_be_visible(timeout=30000)
    snap("01-uploaded")
    step("synthetic XLSX source and template uploaded and parsed")
    for value, field in (("旧编号", "编号"), ("旧姓名", "姓名"), (-999, "金额")):
        cell("right", value).click()
        button("✎ 设为待填区域").click()
        dlg = page.locator(".modal")
        dlg.get_by_placeholder(t("例如：客户名称"), exact=True).fill(field)
        dlg.get_by_role("button", name=t("创建字段"), exact=True).click()
        expect(dlg).to_have_count(0)
    expect(page.locator(".pane.right .gmark .tag")).to_have_count(3)
    step("three template slots created through real cell selection")
    page.locator("button.step", has_text=t("定义样本")).click()
    dlg = page.locator(".modal")
    dlg.locator("select").filter(has_text=t("（不设置）")).select_option("编号")
    expect(dlg.locator(".statbox .v").first).to_have_text("3", timeout=15000)
    dlg.get_by_role("button", name=t("确认样本规则"), exact=True).click()
    expect(dlg).to_have_count(0)
    for index, value in enumerate(ROWS[0]):
        cell("left", value).click()
        page.locator(".pane.right .gmark .tag").nth(index).click()
        button("⇄ 建立映射").click()
        page.locator(".modal").get_by_role("button", name=t("保存映射"), exact=True).click()
        expect(page.locator(".modal")).to_have_count(0)
    page.locator(".tabs").get_by_role("button", name=t("映射 ") + "3", exact=True).click()
    expect(page.locator(".mapitem")).to_have_count(3)
    snap("02-mapped")
    step("three per-row mappings saved through UI")
    # Switch the real production tree and return while retaining selection/mappings.
    other = LOCALES[(LOCALES.index(locale) + 1) % len(LOCALES)]
    page.get_by_label(t("语言"), exact=True).select_option(other)
    expect(page.locator("html")).to_have_attribute("lang", other)
    expect(page.locator(".mapitem")).to_have_count(3)
    page.get_by_label(translator(catalog, other)("语言"), exact=True).select_option(locale)
    expect(page.locator("html")).to_have_attribute("lang", locale)
    button("👁 预览").click()
    expect(page.get_by_text(t("真实样本预览"), exact=True)).to_be_visible()
    expect(cell("right", ROWS[0][1])).to_be_visible(timeout=30000)
    button("下一条 ›").click()
    expect(cell("right", ROWS[1][1])).to_be_visible(timeout=30000)
    expect(page.locator(".notice.info", has_text=t("内置公式引擎已计算公式并写入结果。"))).to_be_visible()
    expect(page.get_by_placeholder(t("按样本标识查找"), exact=True)).to_be_visible()
    snap("03-preview")
    with page.expect_download() as download:
        button("下载此预览").click()
    preview = folder / "preview.xlsx"
    download.value.save_as(preview)
    result["preview_readback"] = readback(preview.read_bytes(), ROWS[1])
    step("real preview, next sample, production download helper and workbook readback")
    button("▶ 生成").click()
    dlg = page.locator(".modal")
    dlg.get_by_label(t("文件命名规则"), exact=True).fill("{编号}_{姓名}")
    start = dlg.get_by_role("button", name=t("开始生成 {v1} 个文件", v1=3), exact=True)
    expect(start).to_be_enabled(timeout=30000)
    snap("04-generate")
    start.click()
    page.wait_for_url(re.compile(r"/jobs/j_"))
    result["job_id"] = page.url.rsplit("/", 1)[-1]
    expect(page.get_by_text(t("已完成"), exact=True)).to_be_visible(timeout=120000)
    snap("05-completed")
    with page.expect_download() as download:
        button("⬇ 下载全部（ZIP）").click()
    archive = folder / "generated.zip"
    download.value.save_as(archive)
    result["download_filename"] = download.value.suggested_filename
    with zipfile.ZipFile(archive) as z:
        expected = {f"{row[0]}_{row[1]}.xlsx": row for row in ROWS}
        assert set(z.namelist()) == set(expected) | {"生成清单.csv"}, z.namelist()
        result["generated_readback"] = [readback(z.read(name), row) for name, row in expected.items()]
    step("real batch engine completed; ZIP downloaded; all three workbooks and formulas read back")
    page.reload()
    expect(page.locator("html")).to_have_attribute("lang", locale)
    expect(page.get_by_text(t("已完成"), exact=True)).to_be_visible()
    step("job and selected locale survived reload")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", action="store_true", help="explicitly create synthetic projects/jobs on a test-marked service")
    parser.add_argument("--base-url")
    parser.add_argument("--storage-state", type=Path, help="optional existing Playwright state from a real test-user login")
    parser.add_argument('--accounts-url', help='with --portal-url and --mail-dir: register a unique a-docs-* fixture through the real UI')
    parser.add_argument('--portal-url')
    parser.add_argument('--mail-dir', type=Path)
    parser.add_argument('--resume-email', help='resume mailbox activation of this runner\'s own pending a-docs-* fixture after a failed driver run')
    parser.add_argument("--evidence", type=Path)
    args = parser.parse_args()
    if not args.run:
        parser.error("Opt-in required: --run (creates real synthetic projects/jobs)")
    if args.storage_state and not args.base_url:
        parser.error("--storage-state requires --base-url")
    signup = any((args.accounts_url, args.portal_url, args.mail_dir))
    if signup and (not all((args.accounts_url, args.portal_url, args.mail_dir, args.base_url)) or args.storage_state):
        parser.error('Registration requires --base-url, --accounts-url, --portal-url, --mail-dir and no --storage-state')
    if args.resume_email and (not signup or not re.fullmatch(r'a-docs-[a-f0-9]{12}@example\.test', args.resume_email)):
        parser.error('--resume-email only accepts this runner\'s pending a-docs-* synthetic fixture with registration options')
    args.evidence = (args.evidence or Path(tempfile.mkdtemp(prefix="synwevo-docs-i18n-artifacts-"))).resolve()
    if args.evidence == REPO or REPO / "data" in [args.evidence, *args.evidence.parents]:
        parser.error("Evidence must not be the repository root or real data directory")
    args.evidence.mkdir(parents=True, exist_ok=True)
    inputs = make_inputs(args.evidence / "synthetic-inputs")
    catalog = labels()
    env = load("_testenv")
    if signup:
        for origin in (args.accounts_url, args.portal_url):
            validate_base(origin, env)
    report = {"script_sha256": digest(Path(__file__)), "mocked_responses": False, "results": []}
    with target(args, env) as (base, mode), sync_playwright() as pw:
        report.update(base=base, mode=mode, health_instance="test", started_at=time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()))
        browser = pw.chromium.launch(channel="chrome", headless=True)
        try:
            state = register_fixture(browser, args, report) if signup else str(args.storage_state) if args.storage_state else None
            for locale in LOCALES:
                context = browser.new_context(viewport={"width": 1600, "height": 1000}, accept_downloads=True,
                                              locale=locale, storage_state=state)
                page = context.new_page()
                page.set_default_timeout(15000)
                result = {"locale": locale, "page_errors": [], "requests": []}
                page.on("pageerror", lambda e: result["page_errors"].append(str(e)))
                # Metadata only; never serialize response bodies, request auth headers or query tokens.
                def response(r):
                    path = urlsplit(r.url).path
                    if path.startswith("/api/"):
                        result["requests"].append({"method": r.request.method, "path": path, "status": r.status,
                                                   "workspace_bound": "x-synwevo-workspace" in r.request.headers})
                page.on("response", response)
                start = time.monotonic()
                try:
                    flow(page, base, locale, catalog, inputs, args.evidence / locale, result)
                    assert not result["page_errors"], result["page_errors"]
                    assert not [r for r in result["requests"] if r["status"] >= 400], "Unexpected API failure; see request ledger"
                    if state:
                        resources = [r for r in result['requests'] if r['method'] not in ('GET', 'HEAD') or '/preview-files/' in r['path'] or r['path'].endswith('/zip')]
                        assert resources and all(r['workspace_bound'] for r in resources), 'A mutation or download omitted the tab workspace header'
                        result['workspace_binding_checked'] = True
                    result["status"] = "passed"
                except Exception:
                    result.update(status="failed", error=traceback.format_exc())
                    page.screenshot(path=str(args.evidence / f"{locale}-failure.png"), full_page=True)
                    (args.evidence / f"{locale}-failure.txt").write_text(page.locator("body").inner_text())
                finally:
                    result["seconds"] = round(time.monotonic() - start, 2)
                    context.close()
                report["results"].append(result)
                (args.evidence / "results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
                print(json.dumps({k: v for k, v in result.items() if k != "requests"}, ensure_ascii=False), flush=True)
        finally:
            browser.close()
    report["evidence"] = [{"path": str(p.relative_to(args.evidence)), "sha256": digest(p)} for p in sorted(args.evidence.rglob("*")) if p.is_file() and p.name != "results.json"]
    (args.evidence / "results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"Evidence: {args.evidence / 'results.json'}", flush=True)
    return 0 if all(r["status"] == "passed" for r in report["results"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
