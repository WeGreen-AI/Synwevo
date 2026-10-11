"""Complete Agent context, using real readers and entirely synthetic in-memory fixtures.

No scanner, OCR, converter, store or network is involved. Test-only PDF fixtures are built
with the same installed native library used by the Docs engine, not mocked parser output.
"""
from __future__ import annotations

import csv
import io
import json
import traceback

import pymupdf
import pytest

from synwevo_core.errors import ServiceError
from synwevo_docs.engine import pdf_engine as pe
from synwevo_docs.ingest import CONTEXT_LIMITS, FORMATS, DocumentIngest


@pytest.fixture(autouse=True)
def full_pdf_engine():
    previous = pe.use_engine("pymupdf")
    try:
        yield
    finally:
        pe.use_engine(previous)


def extract(data, name="fixture.txt", **kwargs):
    result = DocumentIngest().extract_context(data, name, **kwargs)
    assert result["status"] == "ready" and result["error"] is None
    assert result["lines_truncated"] is False
    assert all(table["truncated"] is False for table in result["tables"])
    assert result["ocr"]["applied"] is False
    assert len(json.dumps(result, ensure_ascii=False, separators=(",", ":"))) <= 60000
    return result


def refused(data, name, code, limit=None, **kwargs):
    with pytest.raises(ServiceError) as error:
        extract(data, name, **kwargs)
    assert error.value.code == code
    assert error.value.status == (413 if code == "context_limit_exceeded" else 400)
    if limit:
        assert error.value.detail == {"limit": limit, "maximum": CONTEXT_LIMITS[limit]}
    return error.value


def csv_bytes(rows, delimiter=","):
    buffer = io.StringIO(newline="")
    csv.writer(buffer, delimiter=delimiter).writerows(rows)
    return buffer.getvalue().encode("utf-8")


def pdf_bytes(pages):
    with pymupdf.open() as document:
        for lines in pages:
            page = document.new_page(width=3000, height=3000)
            for index, text in enumerate(lines):
                page.insert_text((40, 40 + 20 * index), text, fontsize=4)
        return document.tobytes()


@pytest.mark.parametrize("name", ["fixture.txt", "fixture.TXT", "fixture.text", "fixture.log", "fixture.md", "fixture.markdown"])
def test_utf8_text_keeps_long_lines_whitespace_blanks_and_source(name):
    lines = ["  合成夹具\t🙂  ", "", "x" * 1500 + "尾部凭据", "\t  ", "final"]
    result = extract(("\ufeff" + "\r\n".join(lines) + "\r\n").encode(), name)
    assert result["lines"] == [{"text": line, "source": {"line": i}} for i, line in enumerate(lines, 1)]
    assert result["tables"] == []


def test_markdown_remains_complete_source_even_with_over_fifty_tables():
    # TEST FIXTURE: links are literal text, not requests to external servers.
    lines = ["```md", "| A | B |", "| -- | -- |", "| literal | escaped \\| pipe |", "```", "[link](https://example.invalid)"]
    lines += ["| a | b |\n| -- | -- |\n| value | end |\n" for _ in range(51)]
    text = "\n".join(lines)
    result = extract(text.encode(), "fixture.md")
    assert [line["text"] for line in result["lines"]] == text.splitlines()
    assert result["tables"] == []


@pytest.mark.parametrize("name", ["empty.txt", "empty.md", "empty.csv", "empty.tsv"])
def test_empty_text_is_complete(name):
    result = extract(b"", name)
    assert result["lines"] == []
    assert all(table["rows"] == [] for table in result["tables"])


@pytest.mark.parametrize("name,delimiter", [("fixture.csv", ","), ("fixture.tsv", "\t")])
def test_delimited_text_keeps_all_cells_and_logical_row_numbers(name, delimiter):
    rows = [["header", "label", "trailing"],
            ["中文" * 800 + "END", "two\r\nphysical\nlines", ""],
            [], ["", "00012", '=SUM(A1:A2)'], ['a "quoted" value', "a,b;c|d\te", "  "]]
    result = extract(b"\xef\xbb\xbf" + csv_bytes(rows, delimiter), name)
    assert result["lines"] == []
    assert result["tables"] == [{"name": "CSV", "source": {"sheet": "CSV"}, "rows": rows,
                                 "row_numbers": [1, 2, 3, 4, 5], "truncated": False}]


def test_csv_does_not_guess_another_delimiter():
    assert extract(b"one;two|three\nvalue;tail|end", "fixture.csv")["tables"][0]["rows"] == [
        ["one;two|three"], ["value;tail|end"]]


def test_csv_reader_cannot_silently_normalize_a_second_bom():
    refused("\ufeff\ufefffield,value".encode(), "fixture.csv", "context_parse_error")


def test_json_escaping_also_counts_against_the_output_budget():
    # Payload itself fits; JSON escaping would exceed the transport character budget.
    refused((('"' * 9000 + "\n") * 4).encode(), "fixture.txt", "context_limit_exceeded", "output_chars")


@pytest.mark.parametrize("name", ["fixture.txt", "fixture.md", "fixture.csv", "fixture.tsv"])
@pytest.mark.parametrize("data", [b"secret\xff", b"secret\x00tail", b"secret\x1btail", b"secret\x7ftail"])
def test_non_utf8_or_binary_text_is_rejected_without_replacement(name, data):
    error = refused(data, name, "context_parse_error")
    assert "secret" not in str(error) + json.dumps(error.detail)


@pytest.mark.parametrize("name,content", [("fixture.csv", 'ok,value\n"secret'),
                                         ("fixture.tsv", 'ok\tvalue\n"secret'),
                                         ("fixture.csv", 'ok,value\n"secret"suffix,x')])
def test_malformed_csv_does_not_return_the_valid_prefix(name, content):
    refused(content.encode(), name, "context_parse_error")


@pytest.mark.parametrize("data,name,limit", [
    (b"x" * 10001, "long.txt", "line_chars"),
    (b"x" * 60001, "large.txt", "chars"),
    (b"\n" * 5001, "blank.txt", "lines"),
    (b"," * 100, "wide.csv", "columns"),
    (b"\n" * 2001, "empty_rows.csv", "rows"),
    (b"x" * 10001, "cell.csv", "cell_chars"),
    ((b"," * 99 + b"\n") * 101, "empty_cells.csv", "cells"),
    (b"x\n" * 2000, "metadata.txt", "output_chars"),
    (b"x" * (20 * 1024 * 1024 + 1), "huge.txt", "file_bytes"),
], ids=["line_chars", "chars", "lines", "columns", "rows", "cell_chars", "cells", "output_chars", "file_bytes"])
def test_real_oversized_inputs_fail_whole(data, name, limit):
    refused(data, name, "context_limit_exceeded", limit)


def test_limit_edges_and_large_complete_tail():
    assert extract(b"x" * 10000)["lines"][0]["text"] == "x" * 10000
    assert len(extract(b"," * 99, "columns.csv")["tables"][0]["rows"][0]) == 100
    assert len(extract(b"\n" * 2000, "rows.csv")["tables"][0]["rows"]) == 2000
    assert len(extract((b"," * 99 + b"\n") * 100, "cells.csv")["tables"][0]["rows"]) == 100
    text = "\n".join(["x" * 9800] * 5 + ["y" * 9700 + "TAIL"])
    result = extract(text.encode())
    assert "\n".join(line["text"] for line in result["lines"]) == text


@pytest.mark.parametrize("name", ["fixture.docm", "fixture.xlsm", "fixture.pptx", "fixture.json", "fixture.html", "fixture.png", "fixture"])
def test_unverified_formats_are_explicitly_unsupported(name):
    refused(b"TEST FIXTURE", name, "context_format_unsupported")


def test_ocr_flag_is_explicitly_unsupported():
    refused(b"TEST FIXTURE", "fixture.txt", "context_ocr_unsupported", ocr=True)


def test_text_pdf_keeps_short_long_blank_and_final_page_sources():
    long_line = "x" * 1200 + "END"
    result = extract(pdf_bytes([["A", long_line], [], ["final page", "last line"]]), "fixture.pdf")
    assert result["pages"] == 3
    assert result["lines"] == [{"text": "A", "source": {"page": 1, "line": 1}},
                               {"text": long_line, "source": {"page": 1, "line": 2}},
                               {"text": "final page", "source": {"page": 3, "line": 1}},
                               {"text": "last line", "source": {"page": 3, "line": 2}}]


def test_cjk_rotated_and_cropped_pdf_text_is_not_silently_lost():
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((40, 40), "中文上下文", fontname="china-s")
        page.insert_text((350, 450), "outside crop")
        page.set_cropbox(pymupdf.Rect(0, 0, 200, 200))
        page.set_rotation(90)
        data = document.tobytes()
    result = extract(data, "fixture.pdf")
    assert [line["text"] for line in result["lines"]] == ["中文上下文", "outside crop"]


def test_pdf_page_limit_checks_the_whole_document():
    assert extract(pdf_bytes([[]] * 200), "edge.pdf")["pages"] == 200
    refused(pdf_bytes([[]] * 200 + [["last page must not disappear"]]), "oversized.pdf",
            "context_limit_exceeded", "pages")


@pytest.mark.parametrize("pages,limit", [
    ([["x" * 10001]], "line_chars"),
    ([["x" * 9000]] * 7, "chars"),
    ([["x"] * 100] * 51, "lines"),
])
def test_pdf_real_text_limits(pages, limit):
    # Very wide pages keep the generated long lines physically on-page.
    with pymupdf.open() as document:
        for lines in pages:
            page = document.new_page(width=30000, height=3000)
            for index, text in enumerate(lines):
                page.insert_text((40, 40 + index * 20), text, fontsize=4)
        data = document.tobytes()
    refused(data, "fixture.pdf", "context_limit_exceeded", limit)


@pytest.mark.parametrize("variant", ["scan", "partial", "small_image", "searchable", "invisible", "vector"])
def test_any_scan_or_unverifiable_part_refuses_the_entire_pdf(variant):
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "complete typed page")
        page = document.new_page()
        if variant in ("scan", "partial", "small_image", "searchable"):
            # TEST FIXTURE: raster words stand in for a scan; includes a tiny partly scanned region.
            with pymupdf.open() as image_doc:
                image_page = image_doc.new_page(width=100, height=100)
                image_page.insert_text((5, 30), "scanned secret", fontsize=10)
                pixmap = image_page.get_pixmap()
            rect = pymupdf.Rect(200, 200, 300, 300) if variant == "small_image" else page.rect
            page.insert_image(rect, pixmap=pixmap)
        if variant in ("partial", "small_image", "searchable", "invisible"):
            page.insert_text((72, 72), "recognized text " * 5,
                             render_mode=3 if variant in ("searchable", "invisible") else 0)
        if variant == "vector":
            page.draw_rect(pymupdf.Rect(10, 10, 50, 50))
        data = document.tobytes()
    error = refused(data, "fixture.pdf", "context_needs_ocr")
    assert error.detail == {"needsOCR": True}


@pytest.mark.parametrize("variant", ["annotation", "form", "attachment", "layer", "link", "xfa"])
def test_unrepresented_pdf_content_is_not_ignored(variant):
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "readable body")
        if variant == "annotation":
            page.add_text_annot((100, 100), "annotation secret")
        elif variant == "form":
            widget = pymupdf.Widget()
            widget.field_name, widget.field_type = "field", pymupdf.PDF_WIDGET_TYPE_TEXT
            widget.field_value, widget.rect = "form secret", pymupdf.Rect(10, 100, 100, 120)
            page.add_widget(widget)
        elif variant == "attachment":
            document.embfile_add("secret.txt", b"attachment secret")
        elif variant == "layer":
            document.add_ocg("layer")
        elif variant == "link":
            page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(70, 60, 150, 90),
                              "uri": "https://example.invalid/secret"})
        else:
            document.xref_set_key(document.pdf_catalog(), "AcroForm", "<< /XFA (unrepresented fixture) >>")
        data = document.tobytes()
    refused(data, "fixture.pdf", "context_format_unsupported")


@pytest.mark.parametrize("user_password", ["", "fixture-password"])
def test_encrypted_pdf_including_owner_only_is_rejected(user_password):
    with pymupdf.open(stream=pdf_bytes([["private fixture"]]), filetype="pdf") as document:
        data = document.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner-fixture",
                               user_pw=user_password)
    refused(data, "fixture.pdf", "context_encrypted")


@pytest.mark.parametrize("variant", ["not_pdf", "truncated", "repairable", "missing_content", "broken_stream"])
def test_damaged_pdf_is_never_a_successful_prefix_and_errors_are_private(variant, capfd):
    data = pdf_bytes([["first valid page"], ["SENSITIVE_FIXTURE_CONTENT"]])
    if variant == "not_pdf":
        data = b"%PDF-1.7\nSENSITIVE_FIXTURE_CONTENT\n%%EOF"
    elif variant == "truncated":
        data = data[:-100]
    elif variant == "repairable":
        data = data[:data.rfind(b"xref\n")] + b"\n%%EOF"
    else:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            content = document[1].get_contents()[0]
            if variant == "missing_content":
                document.update_object(content, "<< /Length 0 >>")
            else:
                document.update_stream(content, b"(SENSITIVE_FIXTURE_CONTENT) UnknownOperator\n")
            data = document.tobytes()
    error = refused(data, "SENSITIVE_FIXTURE_FILENAME.pdf", "context_parse_error")
    formatted = "".join(traceback.format_exception(error))
    output = capfd.readouterr()
    assert "SENSITIVE_FIXTURE" not in str(error) + json.dumps(error.detail) + formatted + output.out + output.err


def test_basic_pdf_engine_fails_closed_and_capabilities_preserve_old_fields(monkeypatch):
    monkeypatch.setenv("SYNWEVO_MALWARE_SCAN", "off")
    full = DocumentIngest().describe()
    assert full["formats"] == list(FORMATS)
    assert full["limits"] == {"text_pages": 200, "ocr_pages": 20, "table_rows": 2000}
    assert full["context_formats"] == ["txt", "md", "csv", "tsv", "pdf", "docx", "xlsx"]
    assert full["context_limits"] == CONTEXT_LIMITS and full["context_ocr"] is False
    pe.use_engine("pdfium")
    refused(pdf_bytes([["body"]]), "fixture.pdf", "context_format_unsupported")
    basic = DocumentIngest().describe()
    assert basic["context_formats"] == ["txt", "md", "csv", "tsv", "docx", "xlsx"]
    assert basic["limits"] == full["limits"] and basic["formats"] == full["formats"]


def test_old_extract_remains_compatible():
    reader = DocumentIngest()
    long_line = "x" * 1000 + "TAIL"
    legacy = reader.extract(long_line.encode(), "fixture.txt")
    assert legacy["lines"] == [{"text": "x" * 500, "source": {"line": 1}}]
    assert extract(long_line.encode())["lines"][0]["text"] == long_line
    data = csv_bytes([["name", "value"], ["fixture", "001"]])
    old, new = reader.extract(data, "fixture.csv"), extract(data, "fixture.csv")
    assert old["tables"] == new["tables"]
    data = pdf_bytes([["first", "second"], ["third"]])
    old, new = reader.extract(data, "fixture.pdf"), extract(data, "fixture.pdf")
    assert old["lines"] == new["lines"] and old["pages"] == new["pages"]


# ---------------------------------------------------------------- Office text documents (synthetic fixtures)

def docx_bytes(build=None):
    """A python-docx document; `build(document)` adds the fixture's content."""
    import docx
    document = docx.Document()
    if build:
        build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(build=None):
    import openpyxl
    workbook = openpyxl.Workbook()
    if build:
        build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def png_bytes():
    return pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), 0).tobytes("png")


def test_docx_text_is_read_completely_with_parts_tables_links_and_comments():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def build(document):
        document.add_heading("合成标题 Fixture", 1)
        paragraph = document.add_paragraph("Body line\twith tab and ")
        rid = paragraph.part.relate_to("https://example.invalid/fixture", RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "a link"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        document.add_paragraph("")                                   # formatting only
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1 first\nB1 second"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = ""
        document.sections[0].header.paragraphs[0].text = "Header 合成"
        document.sections[0].footer.paragraphs[0].text = "Footer 合成"
        commented = document.add_paragraph()
        document.add_comment(commented.add_run("commented text"), text="A reviewer's note", author="Fixture")
    result = extract(docx_bytes(build), "fixture.docx")
    assert result["tables"] == []
    assert [(line["text"], line["source"]) for line in result["lines"]] == [
        ("合成标题 Fixture", {"part": "body", "paragraph": 1}),
        ("Body line\twith tab and a link <https://example.invalid/fixture>", {"part": "body", "paragraph": 2}),
        ("A1\tB1 first\nB1 second", {"part": "body", "table": 1, "row": 1}),
        ("A2\t", {"part": "body", "table": 1, "row": 2}),
        ("commented text", {"part": "body", "paragraph": 4}),
        ("Header 合成", {"part": "header:1", "paragraph": 1}),
        ("Footer 合成", {"part": "footer:1", "paragraph": 1}),
        ("A reviewer's note", {"part": "comments", "paragraph": 1}),
    ]


def test_docx_text_boxes_are_read_once_from_the_document_branch():
    from docx.oxml import parse_xml
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent><mc:Choice Requires="wps">'
        '<w:drawing><wp:inline><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>Box line one</w:t></w:r></w:p><w:p><w:r><w:t>Box line two</w:t></w:r></w:p>'
        '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Box line one</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>')

    def build(document):
        paragraph = document.add_paragraph("Before the box ")
        paragraph._p.append(parse_xml(xml))
        document.add_paragraph("After the box")
    result = extract(docx_bytes(build), "fixture.docx")
    assert [(line["text"], line["source"]) for line in result["lines"]] == [
        ("Before the box ", {"part": "body", "paragraph": 1}),
        ("Box line one", {"part": "body", "paragraph": 1, "textbox": 1}),
        ("Box line two", {"part": "body", "paragraph": 1, "textbox": 1}),
        ("After the box", {"part": "body", "paragraph": 2}),
    ]


def test_docx_with_a_picture_needs_ocr_and_objects_or_equations_are_unsupported():
    from docx.oxml import parse_xml
    from docx.shared import Pt

    def picture(document):
        document.add_paragraph("text")
        document.add_picture(io.BytesIO(png_bytes()), width=Pt(4))
    error = refused(docx_bytes(picture), "fixture.docx", "context_needs_ocr")
    assert error.detail == {"needsOCR": True}
    ole = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:o="urn:schemas-microsoft-com:office:office"><w:object><o:OLEObject ProgID="Fixture.Object"/></w:object></w:r>')
    math = ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><m:r><m:t>x</m:t></m:r></m:oMath>')
    for fragment in (ole, math):
        def build(document, fragment=fragment):
            document.add_paragraph("plain").\
                _p.append(parse_xml(fragment))
        refused(docx_bytes(build), "fixture.docx", "context_format_unsupported")


def test_office_ole_containers_are_encrypted_or_unsupported_never_parsed():
    from synwevo_docs.ingest import OLE_MAGIC
    encrypted = OLE_MAGIC + b"\0" * 24 + "EncryptedPackage".encode("utf-16-le") + b"\0" * 64
    for name in ("fixture.docx", "fixture.xlsx"):
        refused(encrypted, name, "context_encrypted")
        refused(OLE_MAGIC + b"\0" * 120, name, "context_format_unsupported")
    for name in ("fixture.docm", "fixture.xlsm", "fixture.xls", "fixture.doc", "fixture.pptx"):
        refused(b"PK\x03\x04fixture", name, "context_format_unsupported")
    refused(b"PK\x03\x04 not a package", "fixture.docx", "context_parse_error")


def test_xlsx_sheets_are_read_as_shown_including_hidden_sheets_and_dates():
    import datetime

    def build(workbook):
        sheet = workbook.active
        sheet.title = "Main"
        sheet.append(["name", "amount", "when"])
        sheet.append(["x", 12.5, datetime.date(2026, 10, 11)])
        sheet.append(["00012", True, None])
        sheet.merge_cells("A5:B5")
        sheet["A5"] = "merged"
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden.append(["hidden fixture row"])
    result = extract(xlsx_bytes(build), "fixture.xlsx")
    assert result["lines"] == []
    assert result["tables"] == [
        {"name": "Main", "source": {"sheet": "Main"}, "row_numbers": [1, 2, 3, 4, 5], "truncated": False,
         "rows": [["name", "amount", "when"], ["x", "12.5", "2026-10-11"], ["00012", "TRUE"], [], ["merged"]]},
        {"name": "Hidden", "source": {"sheet": "Hidden", "hidden": True}, "row_numbers": [1], "truncated": False,
         "rows": [["hidden fixture row"]]},
    ]


def test_xlsx_with_uncached_formulas_charts_or_pictures_is_refused():
    import openpyxl.chart
    import openpyxl.drawing.image

    def formula(workbook):
        workbook.active.append([1, 2])
        workbook.active["C1"] = "=A1+B1"            # saved without Excel's result
    refused(xlsx_bytes(formula), "fixture.xlsx", "context_format_unsupported")

    def chart(workbook):
        workbook.active.append(["a", 1])
        bars = openpyxl.chart.BarChart()
        bars.add_data(openpyxl.chart.Reference(workbook.active, min_col=2, min_row=1, max_row=1))
        workbook.active.add_chart(bars, "D2")
    refused(xlsx_bytes(chart), "fixture.xlsx", "context_format_unsupported")

    def picture(workbook):
        workbook.active.append(["a"])
        workbook.active.add_image(openpyxl.drawing.image.Image(io.BytesIO(png_bytes())), "C3")
    refused(xlsx_bytes(picture), "fixture.xlsx", "context_needs_ocr")


def test_office_limits_fail_whole():
    def long_document(document):
        for index in range(5001):
            document.add_paragraph(f"line {index}")
    refused(docx_bytes(long_document), "fixture.docx", "context_limit_exceeded", "lines")

    def wide_sheet(workbook):
        for _ in range(3):
            workbook.active.append(["c"] * 101)
    refused(xlsx_bytes(wide_sheet), "fixture.xlsx", "context_limit_exceeded", "columns")

    def many_sheets(workbook):
        for index in range(33):
            workbook.create_sheet(f"S{index}")
    refused(xlsx_bytes(many_sheets), "fixture.xlsx", "context_limit_exceeded", "tables")
