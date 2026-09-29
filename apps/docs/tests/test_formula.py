"""Built-in formula engine: semantics and integration."""
import datetime as dt
import io

import openpyxl
import pytest

from synwevo_docs.engine import values as V
from synwevo_docs.engine.formula import Book, Unsupported, XLError
from synwevo_docs.engine.tabular import read_xlsx
from synwevo_docs.engine.xlsx_doc import XlsxDoc


def _book(formulas: dict, values: dict | None = None):
    cells = {"S": {}}
    for (r, c), v in (values or {}).items():
        cells["S"][(r, c)] = ("v", v)
    for (r, c), f in formulas.items():
        cells["S"][(r, c)] = ("f", f)
    return Book(cells)


VALS = {(0, 0): 3.0, (1, 0): 5.0, (2, 0): "7", (3, 0): None, (4, 0): True, (0, 1): "苹果", (1, 1): "香蕉", (2, 1): "苹果派",
        (0, 2): 2.5, (1, 2): 1.5, (2, 2): 4.0}


@pytest.mark.parametrize("formula,expect", [
    ("1+2*3", 7.0), ("(1+2)*3", 9.0), ("-2^2", 4.0), ("2^3^2", 64.0), ("10%*50", 5.0), ("7/2", 3.5),
    ("SUM(A1:A5)", 8.0),                         # text "7" and TRUE inside a range are ignored
    ("A3+1", 8.0),                               # direct reference to numeric text is coerced
    ("COUNT(A1:A5)", 2.0), ("COUNTA(A1:A5)", 4.0), ("AVERAGE(A1:A2)", 4.0),
    ("SUMIF(B1:B3,\"苹果*\",C1:C3)", 6.5), ("COUNTIF(C1:C3,\">2\")", 2.0), ("SUMPRODUCT(A1:A2,C1:C2)", 15.0),
    ("IF(A1>2,\"多\",\"少\")", "多"), ("IFERROR(1/0,\"错\")", "错"), ("AND(A1>1,A2>1)", True), ("OR(A1>9,FALSE)", False),
    ("ROUND(2.345,2)", 2.35), ("ROUND(-2.5,0)", -3.0), ("ROUNDDOWN(-2.349,2)", -2.34), ("INT(-3.7)", -4.0), ("MOD(-7,3)", 2.0),
    ("VLOOKUP(\"香蕉\",B1:C3,2,FALSE)", 1.5), ("INDEX(A1:C3,2,3)", 1.5), ("MATCH(\"香蕉\",B1:B3,0)", 2.0),
    ("B1&\"-\"&A1", "苹果-3"), ("TEXT(1234.5,\"#,##0.00\")", "1,234.50"), ("LEFT(\"测试文本\",2)", "测试"), ("LEN(\"中文ab\")", 4.0),
    ("SUBSTITUTE(\"a-b-c\",\"-\",\"+\")", "a+b+c"), ("TRIM(\"  a  b \")", "a b"), ("UPPER(\"abc\")", "ABC"),
    ("YEAR(DATE(2024,14,1))", 2025.0), ("ISERROR(1/0)", True), ("ISERR(NA())", False), ("ISNA(VLOOKUP(\"无\",B1:C3,2,FALSE))", True), ("\"a\"=\"A\"", True), ("CHOOSE(2,\"x\",\"y\")", "y"), ("MAX(A:A)", 5.0),
])
def test_semantics(formula, expect):
    b = _book({(9, 9): formula}, VALS)
    v = b.value("S", 9, 9)
    assert v == pytest.approx(expect) if isinstance(expect, float) else v == expect


def test_errors_and_unsupported():
    b = _book({(9, 0): "1/0", (9, 1): "A10+1", (9, 2): "VLOOKUP(\"无\",B1:C3,2,FALSE)", (9, 3): "FOOBAR(1)", (9, 4): "A10"}, VALS)
    assert b.value("S", 9, 0) == XLError("#DIV/0!") and b.value("S", 9, 1) == XLError("#DIV/0!")
    assert b.value("S", 9, 2) == XLError("#N/A")
    with pytest.raises(Unsupported):
        b.value("S", 9, 3)
    cyc = _book({(0, 0): "B1", (0, 1): "A1"})
    with pytest.raises(Unsupported):
        cyc.value("S", 0, 0)


def test_outputs_carry_computed_values_without_libreoffice(monkeypatch):
    import synwevo_docs.engine.converters as C
    monkeypatch.setattr(C, "soffice_path", lambda: None)          # prove no LibreOffice is involved
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "订单"
    ws.append(["数量", "单价", "小计"])
    ws.append([1, 2.5, "=A2*B2"])
    ws["C4"] = "=SUM(C2:C2)"
    ws["D4"] = '=IF(C4>10,"大单","小单")'
    b = io.BytesIO()
    wb.save(b)
    x = XlsxDoc(b.getvalue())
    x.insert_copies_of_row("订单", 2, 1)
    x.set_cell("订单", 1, 0, V.Value(V.INT, "4", "4"))
    x.set_cell("订单", 2, 0, V.Value(V.INT, "2", "2"))
    x.set_cell("订单", 2, 1, V.Value(V.DEC, "3", "3"))
    out = x.finalize()
    assert x.calc_report["computed"] == 4 and not x.calc_report["unsupported"]
    wv = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["订单"]
    assert wv["C2"].value == 10 and wv["C3"].value == 6 and wv["C5"].value == 16 and wv["D5"].value == "大单"


def test_source_formulas_without_cache_are_computed():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", "b", "合计"])
    ws.append([2, 3, "=A2*B2"])
    ws.append([dt.date(2024, 1, 31), None, '=TEXT(EDATE(A3,1),"yyyy-mm-dd")'])
    b = io.BytesIO()
    wb.save(b)                                      # openpyxl writes formulas without cached results
    sh = read_xlsx(b.getvalue()).sheets[0]
    assert sh.cell(1, 2).text == "6" and sh.cell(2, 2).text == "2024-02-29"
    assert any("内置公式引擎计算" in n for n in sh.notes)


CROSS = [
    "SUM(A2:A6)", "AVERAGE(A2:A6)", "COUNT(A2:C6)", "COUNTA(A2:C6)", "COUNTBLANK(A2:C7)", "MAX(A2:A6)-MIN(A2:A6)",
    "MEDIAN(A2:A6)", "LARGE(A2:A6,2)", "SMALL(A2:A6,2)", "SUMIF(B2:B6,\"甲\",A2:A6)", "SUMIFS(A2:A6,B2:B6,\"甲\",A2:A6,\">1\")",
    "COUNTIF(B2:B6,\"甲\")", "COUNTIFS(B2:B6,\"<>甲\",A2:A6,\"<10\")", "AVERAGEIF(B2:B6,\"乙\",A2:A6)", "SUMPRODUCT(A2:A6,C2:C6)",
    "ROUND(A2/3,2)", "ROUNDUP(A3/7,1)", "ROUNDDOWN(-A4/7,1)", "INT(-A5/4)", "MOD(A6,4)", "ABS(-A2)", "POWER(A3,3)", "SQRT(A6)",
    "CEILING(A4/3,0.5)", "FLOOR(A4/3,0.5)", "TRUNC(-A5/3,1)", "SIGN(-A2)", "PRODUCT(A2:A4)",
    "IF(A2>A3,\"大\",\"小\")", "IFERROR(A2/0,-1)", "AND(A2>0,A3>0)", "OR(A2>100,A3>100)", "NOT(A2>1)", "CHOOSE(2,\"一\",\"二\",\"三\")",
    "VLOOKUP(\"丙\",B2:C6,2,FALSE)", "HLOOKUP(\"数量\",A1:C6,3,FALSE)", "INDEX(A2:C6,3,2)", "MATCH(4,A2:A6,0)", "MATCH(5,A2:A6,1)",
    "ROWS(A2:C6)*COLUMNS(A2:C6)", "_xlfn.XLOOKUP(\"丁\",B2:B6,A2:A6)",
    "B2&\"-\"&A2", "CONCATENATE(B2,B3,B4)", "_xlfn.TEXTJOIN(\",\",TRUE,B2:B6)", "LEFT(D2,3)", "RIGHT(D2,2)", "MID(D2,2,3)", "LEN(D2)",
    "FIND(\"b\",D2)", "SEARCH(\"B\",D2)", "SUBSTITUTE(D2,\"b\",\"X\")", "REPLACE(D2,1,2,\"##\")", "TRIM(\"  x  y  \")", "UPPER(D2)",
    "LOWER(\"ABC\")", "PROPER(\"hello world\")", "REPT(\"*\",3)", "EXACT(D2,\"abc\")", "VALUE(\"12.5\")+1",
    "TEXT(A2*1234.5,\"#,##0.00\")", "TEXT(E2,\"yyyy-mm-dd\")", "TEXT(0.256,\"0.0%\")",
    "YEAR(E2)", "MONTH(E2)", "DAY(E2)", "EDATE(E2,1)", "EOMONTH(E2,0)", "DATEDIF(E2,E3,\"d\")", "_xlfn.DAYS(E3,E2)", "WEEKDAY(E2,2)",
    "DATE(2024,2,30)", "ISNUMBER(A2)", "ISTEXT(B2)", "ISERROR(1/0)", "ISNA(NA())", "N(TRUE)", "2^10/4-3*2", "-A2^2", "15%*A3",
    "\"A\"<\"b\"", "A2=\"1\"", "SUM(Sheet2!A1:A3)", "Sheet2!B1*2",
]


@pytest.mark.skipif(not __import__("synwevo_docs.engine.converters", fromlist=["soffice_path"]).soffice_path(), reason="LibreOffice 未安装")
def test_cross_check_against_libreoffice():
    """The same formulas computed by the built-in engine and by LibreOffice Calc must agree."""
    from synwevo_docs.engine.converters import recalc_xlsx
    from synwevo_docs.engine.formula import compute_workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["数量", "名称", "单价", "文本", "日期"])
    rows = [(1, "甲", 2.5, "abcde", dt.datetime(2024, 1, 31)), (4, "乙", 3, None, dt.datetime(2024, 3, 15)),
            (7, "甲", 1.25, None, None), (9, "丙", 8, None, None), (13, "丁", 0.5, None, None)]
    for r in rows:
        ws.append(list(r))
    for i, f in enumerate(CROSS):
        ws.cell(row=10 + i, column=7, value="=" + f)
    s2 = wb.create_sheet("Sheet2")
    for r in ((10, 3), (20, 4), (30, 5)):
        s2.append(list(r))
    b = io.BytesIO()
    wb.save(b)
    x = XlsxDoc(b.getvalue())
    rep = compute_workbook(x)
    assert not rep["unsupported"], rep["unsupported"]
    ours = openpyxl.load_workbook(io.BytesIO(x.pkg.save()), data_only=True)["Sheet1"]
    lo = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(b.getvalue())), data_only=True)["Sheet1"]
    diffs = []
    for i, f in enumerate(CROSS):
        a, e = ours.cell(row=10 + i, column=7).value, lo.cell(row=10 + i, column=7).value
        if isinstance(a, dt.datetime):
            a = a.date()
        if isinstance(e, dt.datetime):
            e = e.date()
        if isinstance(e, dt.date) and isinstance(a, (int, float)):      # LibreOffice also applies a date format to the cell
            e = (e - dt.date(1899, 12, 30)).days
        same = (abs(a - e) < 1e-9) if isinstance(a, (int, float)) and isinstance(e, (int, float)) and not isinstance(a, bool) else a == e
        if not same:
            diffs.append((f, a, e))
    assert not diffs, diffs


def _counta_book() -> bytes:
    """TEST FIXTURE: A2 is a formula returning "" — not an empty cell for COUNTA."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"], ws["A2"], ws["A3"] = 1, '=""', "x"
    for ref, f in (("C1", "=COUNTA(A1:A4)"), ("C2", "=COUNTBLANK(A1:A4)"), ("C3", '=COUNTA("")')):
        ws[ref] = f
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_counta_counts_a_formula_that_returns_empty_text():
    """COUNTA skipped cells whose formula returns "" (engine 2, Excel and LibreOffice 3); COUNTBLANK
    counts them as blank, which it already did."""
    from synwevo_docs.engine.formula import compute_workbook
    x = XlsxDoc(_counta_book())
    compute_workbook(x)
    ws = openpyxl.load_workbook(io.BytesIO(x.pkg.save()), data_only=True)["S"]
    assert (ws["C1"].value, ws["C2"].value, ws["C3"].value) == (3, 2, 1)


@pytest.mark.skipif(not __import__("synwevo_docs.engine.converters", fromlist=["soffice_path"]).soffice_path(), reason="LibreOffice 未安装")
def test_counta_of_empty_text_agrees_with_libreoffice():
    from synwevo_docs.engine.converters import recalc_xlsx
    from synwevo_docs.engine.formula import compute_workbook
    data = _counta_book()
    x = XlsxDoc(data)
    compute_workbook(x)
    ours = openpyxl.load_workbook(io.BytesIO(x.pkg.save()), data_only=True)["S"]
    lo = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(data)), data_only=True)["S"]
    assert [ours[r].value for r in ("C1", "C2", "C3")] == [lo[r].value for r in ("C1", "C2", "C3")]
