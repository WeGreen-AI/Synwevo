"""Inserting detail rows moves everything below them — including what the row shift used to miss:
a pivot table on the same sheet, threaded comments, and the sheet's last row.

Before this, a pivot table below a detail row kept its old `location` while its cells moved, so
LibreOffice (and Excel) laid it out again over the inserted rows; a threaded comment stayed on the
old cell while its legacy note moved; and a detail long enough to pass row 1,048,576 produced rows
no spreadsheet program opens, while a validation or a sum reaching the last row was pushed past it.

TEST FIXTURES: every workbook here is synthetic, built in memory. The pivot table and the threaded
comment are written as XML parts by `_with_pivot_and_thread`, because openpyxl cannot create them.
"""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

from synwevo_docs.engine import templates as T
from synwevo_docs.engine import values as V
from synwevo_docs.engine.refshift import Shift

MAX_ROW = 1_048_576          # rows in an Excel worksheet

soffice = pytest.mark.skipif(not __import__("synwevo_docs.engine.converters", fromlist=["x"]).soffice_path(),
                             reason="需要 LibreOffice")

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
TC = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"


def _with_pivot_and_thread(pivot_at: int = 10, thread_at: str = "B8") -> bytes:
    """Sheet 报表: header row 1, detail row 2, 合计 at row 5, a threaded comment at `thread_at` and a
    pivot table over 数据!A1:B3 whose four rows start at `pivot_at`."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报表"
    ws["A1"], ws["B1"] = "品名", "金额"
    ws["A2"], ws["B2"] = "示例", 0
    ws["A5"], ws["B5"] = "合计", "=SUM(B2:B2)"
    ws[thread_at] = "备注格"
    d = wb.create_sheet("数据")
    for row in (["品名", "金额"], ["甲", 1], ["乙", 2]):
        d.append(row)
    for i, (label, value) in enumerate((("行标签", "求和项:金额"), ("甲", 1), ("乙", 2), ("总计", 3))):
        ws.cell(pivot_at + i, 1, label)
        ws.cell(pivot_at + i, 2, value)
    b = io.BytesIO()
    wb.save(b)
    files = {n: data for n, data in ((n, zipfile.ZipFile(b).read(n)) for n in zipfile.ZipFile(b).namelist())}
    loc = f"A{pivot_at}:B{pivot_at + 3}"
    files["xl/pivotCache/pivotCacheDefinition1.xml"] = (
        f'<pivotCacheDefinition xmlns="{MAIN}" xmlns:r="{REL}" r:id="rId1" refreshOnLoad="1" recordCount="2">'
        '<cacheSource type="worksheet"><worksheetSource ref="A1:B3" sheet="数据"/></cacheSource><cacheFields count="2">'
        '<cacheField name="品名" numFmtId="0"><sharedItems count="2"><s v="甲"/><s v="乙"/></sharedItems></cacheField>'
        '<cacheField name="金额" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1" '
        'containsInteger="1" minValue="1" maxValue="2"/></cacheField></cacheFields></pivotCacheDefinition>').encode()
    files["xl/pivotCache/pivotCacheRecords1.xml"] = (
        f'<pivotCacheRecords xmlns="{MAIN}" count="2"><r><x v="0"/><n v="1"/></r><r><x v="1"/><n v="2"/></r></pivotCacheRecords>').encode()
    files["xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels"] = (
        f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{REL}/pivotCacheRecords" Target="pivotCacheRecords1.xml"/></Relationships>').encode()
    files["xl/pivotTables/pivotTable1.xml"] = (
        f'<pivotTableDefinition xmlns="{MAIN}" name="透视表1" cacheId="7" applyNumberFormats="0" applyBorderFormats="0" '
        'applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" applyWidthHeightFormats="1" dataCaption="值" '
        'updatedVersion="6" minRefreshableVersion="3" createdVersion="6" indent="0" outline="1" outlineData="1">'
        f'<location ref="{loc}" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/><pivotFields count="2">'
        '<pivotField axis="axisRow" showAll="0"><items count="3"><item x="0"/><item x="1"/><item t="default"/></items></pivotField>'
        '<pivotField dataField="1" showAll="0"/></pivotFields><rowFields count="1"><field x="0"/></rowFields>'
        '<rowItems count="3"><i><x/></i><i><x v="1"/></i><i t="grand"><x/></i></rowItems><colItems count="1"><i/></colItems>'
        '<dataFields count="1"><dataField name="求和项:金额" fld="1" baseField="0" baseItem="0"/></dataFields></pivotTableDefinition>').encode()
    files["xl/pivotTables/_rels/pivotTable1.xml.rels"] = (
        f'<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{REL}/pivotCacheDefinition" '
        'Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>').encode()
    files["xl/threadedComments/threadedComment1.xml"] = (
        f'<ThreadedComments xmlns="{TC}"><threadedComment ref="{thread_at}" dT="2024-01-01T00:00:00.00" '
        'personId="{11111111-1111-1111-1111-111111111111}" id="{22222222-2222-2222-2222-222222222222}"><text>请核对</text>'
        '</threadedComment></ThreadedComments>').encode()
    files["xl/persons/person.xml"] = (
        f'<personList xmlns="{TC}"><person displayName="测试" id="{{11111111-1111-1111-1111-111111111111}}" userId="t" '
        'providerId="None"/></personList>').encode()
    files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(b"</Types>", (
        '<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>'
        '<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"/>'
        '<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
        '<Override PartName="/xl/threadedComments/threadedComment1.xml" ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
        '<Override PartName="/xl/persons/person.xml" ContentType="application/vnd.ms-excel.person+xml"/></Types>').encode())
    files["xl/workbook.xml"] = files["xl/workbook.xml"].replace(
        b"</workbook>", f'<pivotCaches><pivotCache xmlns:r="{REL}" cacheId="7" r:id="rIdPc1"/></pivotCaches></workbook>'.encode())
    files["xl/_rels/workbook.xml.rels"] = files["xl/_rels/workbook.xml.rels"].replace(b"</Relationships>", (
        f'<Relationship Id="rIdPc1" Type="{REL}/pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/>'
        '<Relationship Id="rIdPer" Type="http://schemas.microsoft.com/office/2017/10/relationships/person" Target="persons/person.xml"/>'
        '</Relationships>').encode())
    files["xl/worksheets/_rels/sheet1.xml.rels"] = (
        f'<Relationships xmlns="{PKG}"><Relationship Id="rIdPt1" Type="{REL}/pivotTable" Target="../pivotTables/pivotTable1.xml"/>'
        '<Relationship Id="rIdTc1" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
        'Target="../threadedComments/threadedComment1.xml"/></Relationships>').encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)
    return out.getvalue()


def _detail(work: bytes, row: int = 1):
    return T.xlsx_add_slot(work, {"kind": "repeat", "sheet": "报表", "row": row, "c1": 0, "c2": 1, "name": "明细"}, [])


def _rows(n: int) -> dict:
    return {"kind": "rows", "rows": [{"c0": V.text(f"品{i}"), "c1": V.Value(V.INT, i, str(i))} for i in range(n)]}


def _thread_refs(out: bytes) -> list[str]:
    from lxml import etree
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("xl/threadedComments/threadedComment1.xml"))
    return [t.get("ref") for t in root.iter(f"{{{TC}}}threadedComment")]


def test_a_pivot_table_below_the_detail_row_moves_with_its_cells():
    work, rep = _detail(_with_pivot_and_thread())
    out, issues = T.xlsx_fill(work, [rep], {rep["id"]: _rows(4)})
    assert not [i for i in issues if i["level"] == "error"], issues
    ws = openpyxl.load_workbook(io.BytesIO(out))["报表"]
    # three rows were inserted below row 2: the pivot's cells are at 13–16 and so is its definition
    assert [p.location.ref for p in ws._pivots] == ["A13:B16"]
    assert [ws.cell(r, 1).value for r in range(13, 17)] == ["行标签", "甲", "乙", "总计"]
    assert ws["A8"].value == "合计" and ws["B8"].value == "=SUM(B2:B5)"


@soffice
def test_libreoffice_lays_the_moved_pivot_table_out_where_its_cells_are():
    from synwevo_docs.engine.converters import recalc_xlsx
    work, rep = _detail(_with_pivot_and_thread())
    out, _ = T.xlsx_fill(work, [rep], {rep["id"]: _rows(4)})
    ws = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(out)))["报表"]
    # with the old location LibreOffice rebuilt the pivot over rows 10–13, on top of the total's neighbours
    assert [p.location.ref for p in ws._pivots] == ["A13:B16"]
    assert [ws.cell(r, 1).value for r in range(9, 13)] == [None, None, None, None]
    assert ws["A8"].value == "合计" and ws["A14"].value == "甲" and ws["A15"].value == "乙"


def test_a_detail_row_inside_a_pivot_table_is_refused():
    work = _with_pivot_and_thread(pivot_at=10)
    with pytest.raises(T.TemplateError) as e:
        _detail(work, row=10)                       # row 11: inside A10:B13
    assert e.value.code == "pivot_overlap"
    assert str(e.value) == "第 11 行与数据透视表「透视表1」（A10:B13）相交，不能作为重复行"
    # a slot stored before this check is stopped when the document is filled, not written broken
    slot = {"id": "fold", "kind": "repeat", "name": "明细", "locator": {"sheet": "报表", "row": 10, "c1": 0, "c2": 1},
            "columns": [{"key": "c0", "col": 0}, {"key": "c1", "col": 1}], "empty_policy": "remove_row"}
    out, issues = T.xlsx_fill(work, [slot], {"fold": _rows(3)})
    assert out == b""
    assert [(i["code"], i["message"]) for i in issues if i["level"] == "error"] == [
        ("pivot_overlap", "数据透视表「透视表1」（A10:B13）与插入或删除的行相交，无法随之移动：请把重复明细行移到透视表之外")]


def test_threaded_comments_move_with_their_cells():
    work, rep = _detail(_with_pivot_and_thread(thread_at="B8"))
    out, _ = T.xlsx_fill(work, [rep], {rep["id"]: _rows(4)})
    assert openpyxl.load_workbook(io.BytesIO(out))["报表"]["B11"].value == "备注格"
    assert _thread_refs(out) == ["B11"]
    # no detail rows and 「删除明细行」: the rows below move up, the comment with them
    out, _ = T.xlsx_fill(work, [rep], {rep["id"]: _rows(0)})
    assert openpyxl.load_workbook(io.BytesIO(out))["报表"]["B7"].value == "备注格"
    assert _thread_refs(out) == ["B7"]


def test_a_threaded_comment_on_a_deleted_detail_row_goes_with_it():
    work, rep = _detail(_with_pivot_and_thread(thread_at="B2"))
    out, _ = T.xlsx_fill(work, [rep], {rep["id"]: _rows(0)})
    assert _thread_refs(out) == []


def _xlsx(build) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报表"
    build(ws)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_a_detail_that_would_pass_the_last_row_is_refused():
    def build(ws):
        ws["A1"], ws["A2"] = "品名", "示例"
        ws.cell(MAX_ROW - 6, 1, "表尾说明")          # content 7 rows above the sheet's end
    work, rep = _detail(_xlsx(build))
    out, issues = T.xlsx_fill(work, [rep], {rep["id"]: _rows(8)})
    assert out == b""
    assert [(i["code"], i["message"]) for i in issues if i["level"] == "error"] == [
        ("row_limit", "重复明细「明细」有 8 行，插入后工作表「报表」共 1048577 行，超过 Excel 工作表的上限 1,048,576 行")]
    out, issues = T.xlsx_fill(work, [rep], {rep["id"]: _rows(7)})     # exactly to the last row
    assert out and not [i for i in issues if i["level"] == "error"]
    assert openpyxl.load_workbook(io.BytesIO(out))["报表"].cell(MAX_ROW, 1).value == "表尾说明"


def test_ranges_that_reach_the_last_row_stay_on_the_sheet():
    def build(ws):
        ws["A1"], ws["B1"], ws["C1"] = "品名", "类别", "金额"
        ws["A2"] = "示例"
        ws["E1"] = f"=SUM(C2:C{MAX_ROW})"
        dv = DataValidation(type="list", formula1='"甲,乙"')
        dv.add(f"B2:B{MAX_ROW}")
        ws.add_data_validation(dv)
        ws.conditional_formatting.add(f"C2:C{MAX_ROW}", CellIsRule(operator="greaterThan", formula=["100"]))
    work, rep = _detail(_xlsx(build))
    out, issues = T.xlsx_fill(work, [rep], {rep["id"]: _rows(3)})
    assert not [i for i in issues if i["level"] == "error"], issues
    ws = openpyxl.load_workbook(io.BytesIO(out))["报表"]
    # these used to become C2:C1048578 / B2:B1048578 — past the sheet, which Excel repairs away
    assert ws["E1"].value == f"=SUM(C2:C{MAX_ROW})"
    assert [str(d.sqref) for d in ws.data_validations.dataValidation] == [f"B2:B{MAX_ROW}"]
    assert [str(r.sqref) for r in ws.conditional_formatting] == [f"C2:C{MAX_ROW}"]


def test_the_shift_never_goes_past_the_last_row():
    sh = Shift(3, 5, expand_row=2)
    assert sh.span(2, MAX_ROW) == (2, MAX_ROW)          # a range to the end stays at the end
    assert sh.span(2, MAX_ROW - 2) == (2, MAX_ROW)      # a range pushed past it stops there
    assert sh.span(MAX_ROW - 1, MAX_ROW) is None        # a range pushed off entirely is gone
    assert sh.row(MAX_ROW - 5) == MAX_ROW and sh.row(MAX_ROW - 4) is None
    assert sh.span(2, 2) == (2, 7)                      # the template row still grows as before


def test_formulas_on_the_detail_row_are_filled_down_not_grown():
    """A running total on the detail row: line i sums up to its own line, as Excel's fill-down does.
    The template row's own `SUM($B$5:B5)` used to grow like a subtotal — `SUM($B$5:B7)` on the first
    line, `SUM($B$5:B9)` on the third, taking in the rows after it and the subtotal row."""
    def build(ws):
        ws["A4"], ws["B4"], ws["C4"], ws["D4"] = "品名", "金额", "余额", "累计"
        ws["A5"], ws["B5"], ws["C5"], ws["D5"] = "示例", 1, "=N(C4)+B5", "=SUM($B$5:B5)"
        ws["E5"] = "=B5/$B$6"                           # a line's share of the subtotal below
        ws["A6"], ws["B6"] = "小计", "=SUM(B5:B5)"
    work, rep = _detail(_xlsx(build), row=4)
    rows = {"kind": "rows", "rows": [{"c0": V.text(n), "c1": V.Value(V.INT, v, str(v))} for n, v in (("甲", 1), ("乙", 2), ("丙", 3))]}
    out, issues = T.xlsx_fill(work, [rep], {rep["id"]: rows})
    assert not [i for i in issues if i["level"] == "error"], issues
    ws = openpyxl.load_workbook(io.BytesIO(out))["报表"]
    assert [ws[f"D{r}"].value for r in (5, 6, 7)] == ["=SUM($B$5:B5)", "=SUM($B$5:B6)", "=SUM($B$5:B7)"]
    assert [ws[f"C{r}"].value for r in (5, 6, 7)] == ["=N(C4)+B5", "=N(C5)+B6", "=N(C6)+B7"]
    assert [ws[f"E{r}"].value for r in (5, 6, 7)] == ["=B5/$B$8", "=B6/$B$8", "=B7/$B$8"]   # the subtotal moved below
    assert ws["A8"].value == "小计" and ws["B8"].value == "=SUM(B5:B7)"
    calc = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["报表"]
    assert [calc[f"D{r}"].value for r in (5, 6, 7)] == [1, 3, 6]
    assert [calc[f"C{r}"].value for r in (5, 6, 7)] == [1, 3, 6] and calc["B8"].value == 6


# per-line formulas on the detail row (row 5), lines B = 3, 1, 4, 2 and C = 1, 2, 3, 4: the template's
# formula, the first and second lines' formulas, and each line's value — computed from those lines, as
# Excel fills down (relative end row) or reads the whole list (absolute end row)
LINE_FORMULAS = {
    "D": ("=B5*C5", "=B5*C5", "=B6*C6", [3, 2, 12, 8]),
    "E": ("=SUM($B$5:B5)", "=SUM($B$5:B5)", "=SUM($B$5:B6)", [3, 4, 8, 10]),
    "F": ("=N(F4)+B5", "=N(F4)+B5", "=N(F5)+B6", [3, 4, 8, 10]),
    "G": ("=B5*$K$1", "=B5*$K$1", "=B6*$K$1", [6, 2, 8, 4]),
    "H": ("=SUM(B5:C5)", "=SUM(B5:C5)", "=SUM(B6:C6)", [4, 3, 7, 6]),
    "I": ("=SUM(B$4:B5)", "=SUM(B$4:B5)", "=SUM(B$4:B6)", [3, 4, 8, 10]),
    "J": ("=B5/SUM($B$5:$B$5)", "=B5/SUM($B$5:$B$8)", "=B6/SUM($B$5:$B$8)", [0.3, 0.1, 0.4, 0.2]),
    "L": ("=RANK(B5,$B$5:$B$5)", "=RANK(B5,$B$5:$B$8)", "=RANK(B6,$B$5:$B$8)", [2, 4, 1, 3]),
    "M": ("=COUNTIF($A$5:$A$5,A5)", "=COUNTIF($A$5:$A$8,A5)", "=COUNTIF($A$5:$A$8,A6)", [1, 1, 1, 1]),
    "N": ("=MATCH(B5,$B$5:$B$5,0)", "=MATCH(B5,$B$5:$B$8,0)", "=MATCH(B6,$B$5:$B$8,0)", [1, 2, 3, 4]),
    "O": ("=MAX($B$1:$B$5)", "=MAX($B$1:$B$8)", "=MAX($B$1:$B$8)", [4, 4, 4, 4]),
}


def _line_template() -> bytes:
    def build(ws):
        ws["K1"] = 2
        ws["A4"], ws["B4"], ws["C4"] = "品名", "数量", "序号"
        ws["A5"], ws["B5"], ws["C5"] = "示例", 1, 1
        for col, (f, _, _, _) in LINE_FORMULAS.items():
            ws[f"{col}5"] = f
        ws["A7"], ws["B7"] = "合计", "=SUM(B5:B5)"
    return _xlsx(build)


def _four_lines(work, rep):
    rows = {"kind": "rows", "rows": [{"c0": V.text(a), "c1": V.Value(V.INT, b, str(b)), "c2": V.Value(V.INT, c, str(c))}
                                     for a, b, c in (("甲", 3, 1), ("乙", 1, 2), ("丙", 4, 3), ("丁", 2, 4))]}
    return T.xlsx_fill(work, [rep], {rep["id"]: rows})


def test_a_range_on_the_detail_row_with_an_absolute_end_is_the_whole_list():
    """Round 2, item 1: the fill-down above also held ranges whose end row is absolute on the first
    line — `B5/SUM($B$5:$B$5)` gave 1, 0.33, 1.33…, RANK and MATCH #N/A from the second line.
    Absolute end row: the whole list (grows, as before); relative end row: filled down."""
    work, rep = T.xlsx_add_slot(_line_template(), {"kind": "repeat", "sheet": "报表", "row": 4, "c1": 0, "c2": 2, "name": "明细"}, [])
    out, issues = _four_lines(work, rep)
    assert not [i for i in issues if i["level"] == "error"], issues
    ws = openpyxl.load_workbook(io.BytesIO(out))["报表"]
    calc = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["报表"]
    for col, (_, first, second, want) in LINE_FORMULAS.items():
        assert (ws[f"{col}5"].value, ws[f"{col}6"].value) == (first, second), col
        assert [calc[f"{col}{r}"].value for r in range(5, 9)] == pytest.approx(want), col


@soffice
def test_libreoffice_reads_the_detail_row_formulas_the_same_way():
    from synwevo_docs.engine.converters import recalc_xlsx
    work, rep = T.xlsx_add_slot(_line_template(), {"kind": "repeat", "sheet": "报表", "row": 4, "c1": 0, "c2": 2, "name": "明细"}, [])
    out, _ = _four_lines(work, rep)
    lo = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(out)), data_only=True)["报表"]
    for col, (_, _, _, want) in LINE_FORMULAS.items():
        assert [lo[f"{col}{r}"].value for r in range(5, 9)] == pytest.approx(want), col


def test_rules_and_array_formulas_on_the_detail_row_follow_their_line():
    """Round 2, item 7: a validation written on the detail row grew like a subtotal
    (`B5<=SUM($B$5:B5)` became `…:B7` on every line), and a one-row array formula kept a grown range
    (`D5:E7`) overlapping its copies — a file Excel repairs."""
    from openpyxl.worksheet.formula import ArrayFormula

    def build(ws):
        ws["A4"], ws["B4"], ws["C4"] = "品名", "数量", "单价"
        ws["A5"], ws["B5"], ws["C5"] = "示例", 1, 2
        ws["D5"] = ArrayFormula("D5:E5", "=B5:C5*2")
        ws["A6"], ws["B6"] = "小计", "=SUM(B5:B5)"
        dv = DataValidation(type="custom", formula1="B5<=SUM($B$5:B5)")
        dv.add("B5")
        ws.add_data_validation(dv)
        ws.conditional_formatting.add("B6", CellIsRule(operator="greaterThan", formula=["SUM(B5:B5)"]))   # on the subtotal row
    work, rep = _detail(_xlsx(build), row=4)
    out, issues = T.xlsx_fill(work, [rep], {rep["id"]: _rows(3)})
    assert not [i for i in issues if i["level"] == "error"], issues
    ws = openpyxl.load_workbook(io.BytesIO(out))["报表"]
    assert [(str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation] == [("B5:B7", "B5<=SUM($B$5:B5)")]
    assert [(str(r.sqref), [x.formula for x in r.rules]) for r in ws.conditional_formatting] == [("B8", [["SUM(B5:B7)"]])]
    sheet = zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml").decode()
    import re
    assert re.findall(r'<c r="(D\d+)"[^>]*><f t="array" ref="([^"]+)"', sheet) == [("D5", "D5:E5"), ("D6", "D6:E6"), ("D7", "D7:E7")]
