"""Synthetic templates authored entirely over HTTP; real engine and output readback."""
import io
import zipfile

import openpyxl
import pytest

from .test_api_v1 import _ok, _wait, keys  # noqa: F401
from .test_api_preview import _parts

DOC = {"body": {"type": "doc", "content": [{"type": "paragraph", "content": [
    {"type": "text", "text": "客户："}, {"type": "field", "attrs": {"id": "customer", "name": "客户", "required": True}}]}]}}
SHEET = {"sheets": [{"name": "结果", "cells": {"A1": {"v": "客户"}, "B1": {"v": "待填"}},
                      "fields": [{"id": "customer", "kind": "cell", "name": "客户", "ref": "B1", "required": True}]}]}


@pytest.mark.parametrize("kind,editor,fmt", [("builtin_doc", DOC, "docx"), ("builtin_sheet", SHEET, "xlsx")])
def test_author_edit_export_map_preview_generate_builtin(api, keys, kind, editor, fmt):
    c, h = api.c, keys["rw"]
    create = {"kind": kind, "editor": editor, "name": "合成验收模板"}
    t = _ok(c.post("/api/v1/templates/builtin", json=create, headers={**h, "Idempotency-Key": "builtin-create"}))
    tid = t["id"]
    replay = c.post("/api/v1/templates/builtin", json=create, headers={**h, "Idempotency-Key": "builtin-create"})
    assert replay.status_code == 201 and replay.json() == t and replay.headers["Idempotent-Replayed"] == "true"
    model = _ok(c.get(f"/api/v1/templates/{tid}/builtin", headers=h))
    assert model["editor"] == editor and model["version"] == 1
    body = {"editor": editor, "base_version": 1, "name": "改名后的合成模板"}
    edited = _ok(c.put(f"/api/v1/templates/{tid}/builtin", json=body, headers={**h, "Idempotency-Key": "builtin-save"}))
    assert edited["head"] == 2 and edited["name"] == body["name"]
    assert c.put(f"/api/v1/templates/{tid}/builtin", json=body, headers=h).status_code == 409
    again = c.put(f"/api/v1/templates/{tid}/builtin", json=body, headers={**h, "Idempotency-Key": "builtin-save"})
    assert again.json() == edited and again.headers["Idempotent-Replayed"] == "true"
    assert _ok(c.get(f"/api/v1/templates/{tid}/builtin?version=1", headers=h))["editor"] == editor
    # UI and API author exactly the same persisted model through the same service.
    ui = _ok(c.post("/api/templates/builtin", json=create))
    a = c.get(f"/api/v1/templates/{tid}/content?which=working", headers=h)
    b = c.get(f"/api/templates/{ui['id']}/download?which=working")
    assert a.status_code == b.status_code == 200 and _parts(a.content) == _parts(b.content)
    pid = _ok(c.post("/api/v1/projects", json={"name": "内置模板验收"}, headers=h))["id"]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "来源"
    ws.append(["姓名"]); ws.append(["甲客户"]); ws.append(["乙客户"])
    buf = io.BytesIO(); wb.save(buf)
    src = _ok(c.post(f"/api/v1/projects/{pid}/files", files={"file": ("合成.xlsx", buf.getvalue())}, headers=h))
    import time
    for _ in range(100):
        state = _ok(c.get(f"/api/v1/files/{src['id']}", headers=h))
        if state["status"] != "parsing": break
        time.sleep(.02)
    assert state["status"] == "ready"
    _ok(c.put(f"/api/v1/projects/{pid}/template", json={"template_id": tid, "version": 2}, headers=h))
    _ok(c.put(f"/api/v1/projects/{pid}/rule", json={"rule": {"mode": "per_row", "table": {"sheet": "来源", "header_row": 0}}}, headers=h))
    _ok(c.put(f"/api/v1/projects/{pid}/mappings", json={"base_rev": 0, "mappings": [
        {"type": "single", "field": "客户", "sources": [{"kind": "column", "sheet": "来源", "header": "姓名"}]}]}, headers=h))
    pre = _ok(c.post(f"/api/v1/projects/{pid}/preflight", json={}, headers=h))
    assert pre["valid"] == 2
    preview = c.post(f"/api/v1/projects/{pid}/preview", json={"index": 0}, headers=h)
    assert preview.status_code == 200 and preview.content.startswith(b"PK")
    j = _wait(api, keys, _ok(c.post(f"/api/v1/projects/{pid}/jobs", json={"formats": [fmt]}, headers=h))["id"])
    assert j["status"] == "completed" and j["counts"]["documents"] == 2
    items = _ok(c.get(f"/api/v1/jobs/{j['id']}/artifacts", headers=h))["items"]
    got = []
    for a in items:
        out = c.get(f"/api/v1/artifacts/{a['id']}/content", headers=h).content
        if fmt == "xlsx":
            got.append(openpyxl.load_workbook(io.BytesIO(out))["结果"]["B1"].value)
        else:
            with zipfile.ZipFile(io.BytesIO(out)) as z:
                xml = z.read("word/document.xml").decode()
            got.append(next(x for x in ("甲客户", "乙客户") if x in xml))
    assert sorted(got) == ["乙客户", "甲客户"]


def test_builtin_replay_and_models_obey_current_object_authorization(api, keys):
    c, h = api.c, keys["rw"]
    body = {"kind": "builtin_doc", "editor": DOC}
    kh = {**h, "Idempotency-Key": "deleted-builtin"}
    t = _ok(c.post("/api/v1/templates/builtin", json=body, headers=kh))
    for method, url, data in [("post", "/api/v1/templates/builtin", body),
                              ("put", f"/api/v1/templates/{t['id']}/builtin", {"editor": DOC, "base_version": 1})]:
        assert getattr(c, method)(url, json=data, headers=keys["ro"]).status_code == 403
    assert c.get(f"/api/v1/templates/{t['id']}/builtin", headers=keys["other"]).status_code == 404
    assert c.put(f"/api/v1/templates/{t['id']}/builtin", json={"editor": DOC, "base_version": 1}, headers=keys["other"]).status_code == 404
    _ok(c.delete(f"/api/templates/{t['id']}"))
    assert c.post("/api/v1/templates/builtin", json=body, headers=kh).status_code == 404


@pytest.mark.parametrize("editor", [[], {"unknown": 1}, {"page": {"size": "nonsense"}},
    {"body": {"type": "doc", "content": [{"type": "unknown", "text": "不能静默消失"}]}}])
def test_invalid_editor_never_creates_a_template(api, keys, editor):
    before = _ok(api.c.get("/api/v1/templates", headers=keys["rw"]))["total"]
    for path in ("/api/v1/templates/builtin", "/api/templates/builtin"):
        r = api.c.post(path, json={"kind": "builtin_doc", "editor": editor}, headers=keys["rw"])
        assert r.status_code == 400, r.text
    assert _ok(api.c.get("/api/v1/templates", headers=keys["rw"]))["total"] == before


def test_builtin_requires_revision_and_limits_streamed_json(api, keys, monkeypatch):
    from synwevo_docs.engine.limits import LIMITS
    c, h = api.c, keys["rw"]
    t = _ok(c.post("/api/v1/templates/builtin", json={"kind": "builtin_doc", "editor": DOC}, headers=h))
    for base in (None, True, 1.5, "1", -1):
        r = c.put(f"/api/v1/templates/{t['id']}/builtin", json={"editor": DOC, "base_version": base}, headers=h)
        assert r.status_code == 400
    monkeypatch.setattr(LIMITS, "max_editor_body_bytes", 100)
    r = c.post("/api/v1/templates/builtin", content=iter([b'{"editor":"', b'x'*101]), headers={**h, "Content-Type": "application/json"})
    assert r.status_code == 413


@pytest.mark.parametrize("kind,editor", [
    ("builtin_doc", {"page": {"size": "A5", "margins": {"left": 100, "right": 100}}}),
    ("builtin_doc", {"body": {"type": "doc", "content": [{"type": "image", "attrs": {"src": "data:image/svg+xml;base64,PHN2Zz4="}}]}}),
    ("builtin_doc", {"body": {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "原文", "marks": [{"type": []}]}]}]}}),
    ("builtin_doc", {"body": {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "field", "attrs": {"required": "false"}}]}]}}),
    ("builtin_sheet", {"sheets": [{"cells": {"A1": {"v": "金额待定", "t": "number"}}}]}),
    ("builtin_sheet", {"sheets": [{"cells": {"A1": {"v": "1e500", "t": "number"}}}]}),
    ("builtin_sheet", {"sheets": [{"cells": {"A1": {"v": "文字", "s": {"ha": "nonsense"}}}}]}),
    ("builtin_sheet", {"sheets": [{"fields": [{"id": "same", "ref": "A1"}, {"id": "same", "ref": "A2"}]}]}),
])
def test_editor_refuses_silent_data_or_layout_changes(api, keys, kind, editor):
    r = api.c.post("/api/v1/templates/builtin", json={"kind": kind, "editor": editor}, headers=keys["rw"])
    assert r.status_code == 400 and r.json()["error"]["code"] == "bad_editor", r.text
    assert _ok(api.c.get("/api/v1/templates", headers=keys["rw"]))["total"] == 0


def test_explicit_text_formula_and_numbers_keep_their_types_in_real_workbook(api, keys):
    editor = {"sheets": [{"name": "类型", "cells": {
        "A1": {"v": "=1+2", "t": "text"}, "B1": {"v": "=1+2", "t": "auto"},
        "A2": {"v": "00123", "t": "text"}, "B2": {"v": "1e3", "t": "number"}}}]}
    t = _ok(api.c.post("/api/v1/templates/builtin", json={"kind": "builtin_sheet", "editor": editor}, headers=keys["rw"]))
    data = api.c.get(f"/api/v1/templates/{t['id']}/content", headers=keys["rw"]).content
    ws = openpyxl.load_workbook(io.BytesIO(data))["类型"]
    assert (ws["A1"].value, ws["A1"].data_type) == ("=1+2", "s")
    assert (ws["B1"].value, ws["B1"].data_type) == ("=1+2", "f")
    assert ws["A2"].value == "00123" and ws["A2"].data_type == "s"
    assert ws["B2"].value == 1000 and ws["B2"].data_type == "n"
