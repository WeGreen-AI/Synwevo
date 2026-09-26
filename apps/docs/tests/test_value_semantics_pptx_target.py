"""Value semantics into a PPTX target: a slide gets the source's display text — leading zeros,
long identifiers, percent and date displays — exactly as the XLSX reader shows them."""
from __future__ import annotations

import datetime as dt
import io

import openpyxl
import pytest


def _xlsx() -> bytes:
    """The same source as test_value_semantics_e2e (text with leading zeros, long identifiers,
    percent and date formats, a zero-padding format), rebuilt here so the file imports nothing
    from a sibling test module."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["编号", "长号", "比例", "日期", "补零"])
    ws.append(["0007", "12345678901234567890", 0.125, dt.datetime(2026, 9, 26), 7])
    ws.append(["0008", "98765432109876543210", 1, dt.date(2026, 1, 5), 42])
    ws["C2"].number_format = "0.0%"; ws["C3"].number_format = "0%"
    ws["D2"].number_format = "yyyy-mm-dd"; ws["D3"].number_format = "yyyy/m/d"
    ws["E2"].number_format = "000000"; ws["E3"].number_format = "000000"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1.5))
    box.text_frame.text = "编号：X编号 长号：X长号 比例：X比例 日期：X日期 补零：X补零"
    buf = io.BytesIO(); prs.save(buf); return buf.getvalue()


def test_slide_text_gets_display_values(api, tmp_path):
    pytest.importorskip("pptx")
    pid = api.project("PPTX 目标取值")
    api.upload_bytes(pid, "数据.xlsx", _xlsx())
    tpl = tmp_path / "幻灯.pptx"; tpl.write_bytes(_pptx())
    t = api.template(pid, tpl)
    assert t["format"] == "pptx", t
    tid = t["id"]
    slides = api.view(tid)["slides"]["slides"]
    slots = {}
    for name in ("编号", "长号", "比例", "日期", "补零"):
        needle = f"X{name}"
        hit = next((sl["index"], sh["id"], pi, para["text"].index(needle)) for sl in slides for sh in sl["shapes"]
                   for pi, para in enumerate(sh.get("paras", [])) if needle in para["text"])
        r = api.c.post(f"/api/templates/{tid}/slots", json={"kind": "text", "slide": hit[0], "shape": hit[1], "para": hit[2], "start": hit[3], "end": hit[3] + len(needle), "name": name})
        assert r.status_code == 200, r.text
        slots[name] = r.json()["created_slot"]["binding"]
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "数据", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "数据", "header": h}], "target": {"binding": b}} for h, b in slots.items()])
    job = api.job(pid, naming="{编号}", formats=["pptx"])
    assert job["status"] == "completed", (job.get("error"), [(it["status"], it.get("issues")) for it in api.items(job["id"])])
    outs = api.outputs(job["id"])
    from pptx import Presentation
    texts = {}
    for name, data in outs.items():
        prs = Presentation(io.BytesIO(data))
        texts[name] = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert texts["0007.pptx"] == "编号：0007 长号：12345678901234567890 比例：12.5% 日期：2026-09-26 补零：000007", texts
    assert texts["0008.pptx"] == "编号：0008 长号：98765432109876543210 比例：100% 日期：2026/1/5 补零：000042", texts
