"""XLSX repeat content blocks (docs/platform/docs-xlsx-nested.md): a band of whole rows copied
once per order, the order's lines expanding inside each copy, formulas following their own copy.

TEST FIXTURES: every workbook is synthetic, built here with openpyxl; the slots are built by hand
(as the template service stores them) and the values as the plan hands them over. Every expected
number is computed from the rows written below, not read back from the engine.
"""
from __future__ import annotations

import io
import time
import zipfile
from decimal import Decimal

import openpyxl
import pytest
from lxml import etree
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from synwevo_docs.engine import templates as T
from synwevo_docs.engine import values as V
from synwevo_docs.engine.xlsx_doc import XlsxDoc, XlsxError

soffice = pytest.mark.skipif(not __import__("synwevo_docs.engine.converters", fromlist=["x"]).soffice_path(),
                             reason="需要 LibreOffice")
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

# customer A's orders: A-1 with two lines, A-2 with one (品名, 数量, 单价)
ORDERS = [("A-1", "2024-03-01", [("螺丝", 5, "2.00"), ("螺母", 3, "1.20")]),
          ("A-2", "2024-03-09", [("垫圈", 10, "0.50")])]


def _template(extra=None) -> bytes:
    """对账单: customer B1; block rows 3–6 (order header with B3/D3, line header, line row 5,
    subtotal row 6); grand totals below; a tax rate in H1; a summary sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "对账单"
    ws["A1"], ws["B1"], ws["G1"], ws["H1"] = "客户：", "C-0", "税率", 0.13
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "订单号：", "O-0", "日期：", "2024-01-01"
    for c, h in enumerate(("品名", "数量", "单价", "金额", "含税", "累计"), start=1):
        ws.cell(4, c, h)
    ws["A5"], ws["B5"], ws["C5"] = "示例", 1, 0
    ws["D5"], ws["E5"], ws["F5"] = "=B5*C5", "=D5*(1+H1)", "=SUM($D$5:D5)"
    ws["A6"], ws["D6"], ws["F6"] = "订单小计", "=SUM(D5:D5)", "=COUNTA(A5:A5)"
    ws["A8"], ws["D8"] = "合计", "=SUM(D6)"
    ws["A9"], ws["D9"] = "行数", "=COUNTA(A5:A5)"
    ws["A10"], ws["D10"] = "平均单价", "=AVERAGE(C5:C5)"
    ws["A11"], ws["D11"] = "明细合计", "=SUM(D5:D5)"
    ws["A12"], ws["D12"] = "块内行数", "=ROWS(A3:A6)"
    ws["A13"], ws["D13"] = "引用税率", "=H1"
    s2 = wb.create_sheet("汇总")
    s2["A1"], s2["B1"] = "总金额", "=SUM(对账单!D6)"
    s2["A2"], s2["B2"] = "最大单笔", "=MAX(对账单!D5:D5)"
    if extra:
        extra(wb)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _detail(work: bytes, row: int = 4, c2: int = 2, name: str = "明细", **kw):
    return T.xlsx_add_slot(work, {"kind": "repeat", "sheet": "对账单", "row": row, "c1": 0, "c2": c2, "name": name, **kw}, [])


def _band(first: int = 2, last: int = 5, **kw) -> dict:
    return {"id": "band", "binding": "band", "kind": "repeat", "block": True, "name": "订单段",
            "locator": {"sheet": "对账单", "first": first, "last": last},
            "columns": [{"key": "c0", "name": "订单号", "cell": "B3"}, {"key": "c1", "name": "日期", "cell": "D3"}],
            "empty_policy": "remove", "placeholder": "（无）", "separator": "none", **kw}


def _cust():
    return {"id": "cust", "binding": "cust", "kind": "cell", "name": "客户", "mode": "replace",
            "locator": {"sheet": "对账单", "cell": "B1"}}


def _values(band, rep, orders=ORDERS):
    lines = [[{"c0": V.text(n), "c1": V.from_python(q), "c2": V.Value(V.DEC, p, p)} for n, q, p in ls] for _, _, ls in orders]
    return {"cust": {"kind": "scalar", "value": V.text("客户 A")},
            band["id"]: {"kind": "rows", "rows": [{"c0": V.text(o), "c1": V.text(d)} for o, d, _ in orders],
                         "sections": [o for o, _, _ in orders]},
            rep["id"]: {"kind": "rows", "rows": [r for g in lines for r in g], "groups": lines, "sections": [o for o, _, _ in orders]}}


def _fill(work, slots, values):
    out, issues = T.xlsx_fill(work, slots, values)
    errors = [i for i in issues if i["level"] == "error"]
    return out, issues, errors


def _book(out, data_only=False):
    return openpyxl.load_workbook(io.BytesIO(out), data_only=data_only)


def _standard(extra=None, orders=ORDERS, band_kw=None, detail_kw=None):
    work, rep = _detail(_template(extra), **(detail_kw or {}))
    band = _band(**(band_kw or {}))
    slots = [_cust(), band, rep]
    return work, slots, band, rep, _values(band, rep, orders)


def test_each_order_copy_gets_its_own_lines_subtotal_and_the_total_lists_every_copy():
    work, slots, band, rep, values = _standard()
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    col = lambda c, a, b: [ws[f"{c}{r}"].value for r in range(a, b + 1)]  # noqa: E731
    assert ws["B1"].value == "客户 A"
    # copy 1: rows 3–7 (two lines), copy 2: rows 8–11 (one line); the rows below moved by 5
    assert col("A", 3, 11) == ["订单号：", "品名", "螺丝", "螺母", "订单小计", "订单号：", "品名", "垫圈", "订单小计"]
    assert [ws["B3"].value, ws["D3"].value, ws["B8"].value, ws["D8"].value] == ["A-1", "2024-03-01", "A-2", "2024-03-09"]
    assert col("B", 5, 6) == [5, 3] and ws["B10"].value == 10
    # formulas: each line its own row; the tax rate stays H1 in every copy (no $ needed)
    assert col("D", 5, 6) == ["=B5*C5", "=B6*C6"] and ws["D10"].value == "=B10*C10"
    assert col("E", 5, 6) == ["=D5*(1+H1)", "=D6*(1+H1)"] and ws["E10"].value == "=D10*(1+H1)"
    # running total inside the copy: $D$5 is the copy's first line
    assert col("F", 5, 6) == ["=SUM($D$5:D5)", "=SUM($D$5:D6)"] and ws["F10"].value == "=SUM($D$10:D10)"
    # the block subtotal covers only its own lines
    assert ws["D7"].value == "=SUM(D5:D6)" and ws["D11"].value == "=SUM(D10:D10)"
    assert ws["F7"].value == "=COUNTA(A5:A6)" and ws["F11"].value == "=COUNTA(A10:A10)"
    # below the block: aggregate arguments expand to every copy, a range over the block grows
    assert ws["A13"].value == "合计" and ws["D13"].value == "=SUM(D7,D11)"
    assert ws["D14"].value == "=COUNTA(A5:A6,A10:A10)"
    assert ws["D15"].value == "=AVERAGE(C5:C6,C10:C10)"
    assert ws["D16"].value == "=SUM(D5:D6,D10:D10)"
    assert ws["D17"].value == "=ROWS(A3:A11)"
    assert ws["D18"].value == "=H1"
    s2 = _book(out)["汇总"]
    assert s2["B1"].value == "=SUM(对账单!D7,对账单!D11)" and s2["B2"].value == "=MAX(对账单!D5:D6,对账单!D10:D10)"


def _expected():
    lines = [[q * Decimal(p) for _, q, p in ls] for _, _, ls in ORDERS]
    return lines, [sum(x) for x in lines]


def test_the_built_in_engine_computes_the_subtotals_and_the_total():
    work, slots, band, rep, values = _standard()
    out, issues, _ = _fill(work, slots, values)
    ws = _book(out, data_only=True)["对账单"]
    lines, subs = _expected()
    cent = lambda ref: Decimal(str(ws[ref].value)).quantize(Decimal("0.01"))  # noqa: E731 - 3×1.2 is 3.5999… in binary
    assert [cent(f"D{r}") for r in (5, 6, 10)] == [lines[0][0], lines[0][1], lines[1][0]]
    assert cent("D7") == subs[0] and cent("D11") == subs[1] and cent("D13") == sum(subs)
    assert ws["D14"].value == 3 and ws["D17"].value == 9
    assert Decimal(str(_book(out, data_only=True)["汇总"]["B1"].value)) == sum(subs)
    assert not [i for i in issues if i["code"] in ("formula_unsupported", "formula_error")], issues


@soffice
def test_libreoffice_recalculates_the_same_subtotals_and_total():
    from synwevo_docs.engine.converters import recalc_xlsx
    work, slots, band, rep, values = _standard()
    out, _, _ = _fill(work, slots, values)
    ours = _book(out, data_only=True)["对账单"]
    lo = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(out)), data_only=True)["对账单"]
    for ref in ("D5", "D6", "D7", "D10", "D11", "D13", "D14", "D15", "D16", "D17", "E5", "F6", "F7"):
        assert abs(float(ours[ref].value) - float(lo[ref].value)) < 1e-9, ref


def test_subtotal_arguments_expand_and_are_left_to_the_fallback_like_any_unsupported_function():
    def extra(wb):
        ws = wb["对账单"]
        ws["D6"] = "=SUBTOTAL(9,D5:D5)"
        ws["E8"] = "=SUBTOTAL(9,D6)"            # grand total over the subtotals
        ws["E9"] = "=SUBTOTAL(9,D3:D6)"         # over the whole block: SUBTOTAL skips the subtotals
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert ws["D7"].value == "=SUBTOTAL(9,D5:D6)" and ws["D11"].value == "=SUBTOTAL(9,D10:D10)"
    assert ws["E13"].value == "=SUBTOTAL(9,D7,D11)" and ws["E14"].value == "=SUBTOTAL(9,D3:D11)"
    # not computed on the server: reported the way every unsupported function is
    unsupported = [i for i in issues if i["code"] in ("formula_fallback", "formula_unsupported")]
    assert len(unsupported) == 1 and "SUBTOTAL" in unsupported[0]["message"]


@soffice
def test_libreoffice_computes_the_subtotal_grand_total_without_double_counting():
    from synwevo_docs.engine.converters import recalc_xlsx
    def extra(wb):
        wb["对账单"]["D6"] = "=SUBTOTAL(9,D5:D5)"
        wb["对账单"]["E8"] = "=SUBTOTAL(9,D3:D6)"
    work, slots, band, rep, values = _standard(extra)
    out, _, _ = _fill(work, slots, values)
    lo = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(out)), data_only=True)["对账单"]
    _, subs = _expected()
    assert Decimal(str(lo["E13"].value)).quantize(Decimal("0.01")) == sum(subs)


def test_structures_inside_the_block_are_copied_and_the_ones_below_move():
    def extra(wb):
        ws = wb["对账单"]
        ws.merge_cells("B3:C3")                                   # order header: merged per copy
        ws.merge_cells("E5:F5")                                   # a line's merge: one per line
        ws.merge_cells("A8:C8")                                   # below: moves
        ws.row_dimensions[4].height = 30
        ws.row_dimensions[6].hidden = True
        ws.conditional_formatting.add("C5", CellIsRule(operator="greaterThan", formula=["1"]))
        ws.conditional_formatting.add("A1:A20", FormulaRule(formula=["LEN(A1)>3"]))   # crosses the edge: split
        dv = DataValidation(type="whole", operator="between", formula1="0", formula2="999")
        dv.add("B5")
        ws.add_data_validation(dv)
        wb.defined_names["_xlnm.Print_Area"] = openpyxl.workbook.defined_name.DefinedName(
            "_xlnm.Print_Area", localSheetId=0, attr_text="对账单!$A$1:$H$13")
        wb.defined_names["税率"] = openpyxl.workbook.defined_name.DefinedName("税率", attr_text="对账单!$H$1")
        wb.defined_names["合计格"] = openpyxl.workbook.defined_name.DefinedName("合计格", attr_text="对账单!$D$8")
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    wb = _book(out)
    ws = wb["对账单"]
    merges = sorted(str(m) for m in ws.merged_cells.ranges)
    assert merges == sorted(["B3:C3", "B8:C8", "E5:F5", "E6:F6", "E10:F10", "A13:C13"])
    assert ws.row_dimensions[4].height == 30 and ws.row_dimensions[9].height == 30
    assert ws.row_dimensions[7].hidden and ws.row_dimensions[11].hidden
    # a format on the line's cell covers every line of every copy; one across the block's edge is
    # split: above it, each copy, and below it (moved) — still one rule, evaluated cell by cell
    assert sorted(str(r.sqref) for r in ws.conditional_formatting) == ["A1:A2 A3:A7 A8:A11 A12:A25", "C5:C6 C10"]
    assert [str(d.sqref) for d in ws.data_validations.dataValidation] == ["B5:B6 B10"]
    names = {n: wb.defined_names[n].attr_text for n in ("税率", "合计格")}
    assert names == {"税率": "对账单!$H$1", "合计格": "对账单!$D$13"}
    assert ws.print_area == "'对账单'!$A$1:$H$18"


def test_page_breaks_blank_separators_and_breaks_inside_the_block():
    def extra(wb):
        from openpyxl.worksheet.pagebreak import Break
        wb["对账单"].row_breaks.append(Break(id=4))          # after the line header, inside the block
    work, slots, band, rep, values = _standard(extra, band_kw={"separator": "page_break"})
    out, _, errors = _fill(work, slots, values)
    assert not errors
    ws = _book(out)["对账单"]
    assert sorted(b.id for b in ws.row_breaks.brk) == [4, 7, 9]   # 4 and 9: each copy's header; 7: after copy 1
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml"))
    rb = root.find(f"{{{MAIN}}}rowBreaks")
    assert rb.get("count") == "3" and rb.get("manualBreakCount") == "3"
    # one blank row between copies
    work, slots, band, rep, values = _standard(band_kw={"separator": "blank"})
    out, _, errors = _fill(work, slots, values)
    ws = _book(out)["对账单"]
    assert [ws[f"A{r}"].value for r in range(7, 10)] == ["订单小计", None, "订单号："]
    assert ws["D12"].value == "=SUM(D11:D11)" and ws["D14"].value == "=SUM(D7,D12)"


def test_empty_block_and_empty_orders_follow_their_policies():
    # no orders, 「删除整个内容块」: the rows go, the total's argument becomes #REF! with a warning
    work, slots, band, rep, values = _standard(orders=[])
    out, issues, errors = _fill(work, slots, values)
    assert not errors
    ws = _book(out)["对账单"]
    assert ws["A4"].value == "合计" and ws["D4"].value == "=SUM(#REF!)"
    assert ws["D8"].value == "=ROWS(#REF!)"               # the block it counted is gone
    warn = [i["message"] for i in issues if i["code"] == "ref_removed"]
    assert warn and warn[0].startswith("重复内容块「订单段」没有数据，已删除；引用它的公式 对账单!D4")
    # 「保留一份」: one copy, the first column shows the placeholder, the detail row follows its own policy
    work, slots, band, rep, values = _standard(orders=[], band_kw={"empty_policy": "placeholder"})
    out, issues, errors = _fill(work, slots, values)
    ws = _book(out)["对账单"]
    assert ws["B3"].value == "（无）" and ws["D3"].value is None
    assert ws["A5"].value == "订单小计" and ws["D5"].value == "=SUM(#REF!)"
    # an order without lines: 「删除明细行」 vs 「保留一行空白」
    orders = [ORDERS[0], ("A-3", "2024-04-01", [])]
    work, slots, band, rep, values = _standard(orders=orders)
    out, issues, errors = _fill(work, slots, values)
    ws = _book(out)["对账单"]
    assert [ws[f"A{r}"].value for r in range(8, 11)] == ["订单号：", "品名", "订单小计"]
    assert ws["D10"].value == "=SUM(#REF!)"
    assert [i["message"] for i in issues if i["code"] == "ref_removed"] == [
        "「A-3」这一段没有明细行，明细行已删除；该段的公式 D10、F10 变为 #REF!。如需保留公式，请把明细行的零条明细处理方式改为「保留一行空白」"]
    work, slots, band, rep, values = _standard(orders=orders, detail_kw={"empty_policy": "blank_row"})
    out, issues, errors = _fill(work, slots, values)
    ws = _book(out)["对账单"]
    assert [ws[f"A{r}"].value for r in range(8, 12)] == ["订单号：", "品名", None, "订单小计"]
    assert ws["D11"].value == "=SUM(D10:D10)" and not [i for i in issues if i["code"] == "ref_removed"]


def test_two_detail_rows_in_one_block_and_a_second_block_below():
    """Order block rows 1–5 with two detail rows (lines, payments); a notes block rows 7–8."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "对账单"
    ws["A1"], ws["B1"] = "订单号：", "O-0"
    ws["A2"], ws["B2"] = "品名", "金额"
    ws["A3"], ws["B3"] = "示例", 1
    ws["A4"], ws["B4"] = "付款", 0
    ws["A5"], ws["B5"] = "差额", "=SUM(B3:B3)-SUM(B4:B4)"
    ws["A7"], ws["B7"] = "备注：", "N-0"
    ws["A8"] = "——"
    b = io.BytesIO()
    wb.save(b)
    work = b.getvalue()
    work, lines = T.xlsx_add_slot(work, {"kind": "repeat", "sheet": "对账单", "row": 2, "c1": 0, "c2": 1, "name": "明细"}, [])
    work, pays = T.xlsx_add_slot(work, {"kind": "repeat", "sheet": "对账单", "row": 3, "c1": 0, "c2": 1, "name": "付款"}, [])
    band = {**_band(first=0, last=4), "columns": [{"key": "c0", "name": "订单号", "cell": "B1"}]}
    notes = {**_band(first=6, last=7), "id": "notes", "binding": "notes", "name": "备注段",
             "columns": [{"key": "c0", "name": "备注", "cell": "B7"}]}
    row = lambda *v: {f"c{i}": (V.text(x) if isinstance(x, str) else V.from_python(x)) for i, x in enumerate(v)}  # noqa: E731
    values = {"band": {"kind": "rows", "rows": [row("X-1"), row("X-2")], "sections": ["X-1", "X-2"]},
              lines["id"]: {"kind": "rows", "rows": [], "groups": [[row("甲", 5), row("乙", 7)], [row("丙", 1)]]},
              pays["id"]: {"kind": "rows", "rows": [], "groups": [[row("付款", 4)], [row("付款", 1), row("付款", 0)]]},
              "notes": {"kind": "rows", "rows": [row("第一条"), row("第二条"), row("第三条")]}}
    out, issues, errors = _fill(work, [band, lines, pays, notes], values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert [ws[f"A{r}"].value for r in range(1, 12)] == ["订单号：", "品名", "甲", "乙", "付款", "差额",
                                                        "订单号：", "品名", "丙", "付款", "付款"]
    assert ws["B6"].value == "=SUM(B3:B4)-SUM(B5:B5)" and ws["A12"].value == "差额"
    assert ws["B12"].value == "=SUM(B9:B9)-SUM(B10:B11)"
    assert [ws[f"B{r}"].value for r in (14, 16, 18)] == ["第一条", "第二条", "第三条"]


def test_a_standalone_detail_row_below_the_block_and_a_cell_field_inside_it():
    def extra(wb):
        ws = wb["对账单"]
        ws["A20"], ws["B20"] = "附注", "N"
        ws["G3"] = "客户"                                 # a cell field inside the block: same in every copy
    work, rep = _detail(_template(extra))
    work, notes = T.xlsx_add_slot(work, {"kind": "repeat", "sheet": "对账单", "row": 19, "c1": 0, "c2": 1, "name": "附注"}, [])
    band = _band()
    inner = {"id": "inner", "binding": "cust", "kind": "cell", "name": "块内客户", "mode": "replace",
             "locator": {"sheet": "对账单", "cell": "G3"}}
    values = _values(band, rep)
    values["inner"] = values["cust"]
    values[notes["id"]] = {"kind": "rows", "rows": [{"c0": V.text("n1"), "c1": V.text("x")}, {"c0": V.text("n2"), "c1": V.text("y")}]}
    out, issues, errors = _fill(work, [_cust(), band, rep, notes, inner], values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert ws["G3"].value == "客户 A" and ws["G8"].value == "客户 A"
    assert [ws["A25"].value, ws["A26"].value] == ["n1", "n2"]
    assert ws["D13"].value == "=SUM(D7,D11)"


# ------------------------------------------------------------------ refusals
def _refused(extra, code, message=None):
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert out == b"" and [e["code"] for e in errors] == [code], errors
    if message:
        assert errors[0]["message"] == message
    return errors[0]["message"]


def test_a_reference_to_part_of_the_block_from_outside_is_refused_unless_it_is_an_aggregate_argument():
    _refused(lambda wb: wb["对账单"].__setitem__("E8", "=D6*1.1"), "band_reference",
             "单元格 对账单!E8 的公式引用了重复内容块内的 D6：内容块会按段复制，这个引用只能指向其中一份。"
             "请改为引用整个内容块（如 3:6），或把它写成 SUM、COUNT、COUNTA、AVERAGE、MIN、MAX、PRODUCT、SUBTOTAL 的一个参数（生成时展开到每一份）")
    _refused(lambda wb: wb["对账单"].__setitem__("E8", "=SUM(D6)+VLOOKUP(1,A5:D5,4,0)"), "band_reference")
    _refused(lambda wb: wb["对账单"].__setitem__("E8", "=SUBTOTAL(D6,1)"), "band_reference")   # the function number


def test_references_and_merges_across_the_edge_are_refused():
    _refused(lambda wb: wb["对账单"].__setitem__("E8", "=SUM(D5:D8)"), "band_cross",
             "单元格 对账单!E8 的公式引用的 D5:D8 跨过了重复内容块的边界（第 3–6 行）：请让引用完全在块内，或包含整个内容块")
    _refused(lambda wb: wb["对账单"].__setitem__("E5", "=SUM(D1:D5)"), "band_cross")
    _refused(lambda wb: wb["对账单"].merge_cells("G6:G7"), "band_merge",
             "合并单元格 G6:G7 跨过了重复内容块的边界（第 3–6 行）：请让它完全在块内或块外")


def test_names_charts_formats_and_objects_touching_part_of_the_block_are_refused():
    from openpyxl.workbook.defined_name import DefinedName
    def name(wb):
        wb.defined_names["小计"] = DefinedName("小计", attr_text="对账单!$D$6")
    _refused(name, "band_name")
    def titles(wb):
        wb["对账单"].print_title_rows = "4:4"
    assert "名称「Print_Titles」" in _refused(titles, "band_name")
    def cf_split(wb):                  # inside and outside the block, looking at another row
        wb["对账单"].conditional_formatting.add("A3:A20", FormulaRule(formula=["A4>0"]))
    _refused(cf_split, "band_cf_split",
             "条件格式 A3:A20 既在重复内容块（第 3–6 行）内又在块外，而它的公式引用 A4 在每一份里指向的行不同："
             "请把块内和块外分成两条规则，或只引用本行")
    def cf_out(wb):
        wb["对账单"].conditional_formatting.add("D8", FormulaRule(formula=["D6>0"]))
    _refused(cf_out, "band_cf")
    def dv(wb):
        v = DataValidation(type="list", formula1="$A$5:$A$5")
        v.add("H9")
        wb["对账单"].add_data_validation(v)
    _refused(dv, "band_dv")
    def chart(wb):
        from openpyxl.chart import BarChart, Reference
        ch = BarChart()
        ch.add_data(Reference(wb["对账单"], min_col=4, min_row=5, max_row=5))
        wb["汇总"].add_chart(ch, "D2")
    _refused(chart, "band_chart")
    def picture_comment_link(kind):
        def build(wb):
            ws = wb["对账单"]
            if kind == "comment":
                from openpyxl.comments import Comment
                ws["C4"].comment = Comment("看这里", "测试")
            else:
                ws["C4"].hyperlink = "https://example.test/"
        return build
    assert _refused(picture_comment_link("comment"), "band_comment") == "单元格 C4 有批注：复制内容块时不会复制批注，请移到块外或删除"
    assert _refused(picture_comment_link("link"), "band_hyperlink") == "单元格 C4 有超链接：复制内容块时不会复制超链接，请移到块外或删除"
    def table(wb):
        from openpyxl.worksheet.table import Table
        wb["对账单"].add_table(Table(displayName="明细表", ref="A4:C5"))
    _refused(table, "band_table", "表格「明细表」（A4:C5）与重复内容块相交：Excel 表格不能随内容块复制，请转换为普通区域或移到块外")


def test_a_chart_or_name_over_the_whole_block_is_fine_and_grows():
    from openpyxl.chart import BarChart, Reference
    from openpyxl.workbook.defined_name import DefinedName
    def extra(wb):
        ch = BarChart()
        ch.add_data(Reference(wb["对账单"], min_col=4, min_row=3, max_row=6))
        wb["汇总"].add_chart(ch, "D2")
        wb.defined_names["全部订单"] = DefinedName("全部订单", attr_text="对账单!$A$3:$F$6")
    work, slots, band, rep, values = _standard(extra)
    out, _, errors = _fill(work, slots, values)
    assert not errors, errors
    wb = _book(out)
    assert wb.defined_names["全部订单"].attr_text == "对账单!$A$3:$F$11"
    chart = zipfile.ZipFile(io.BytesIO(out)).read("xl/charts/chart1.xml").decode()
    assert "$D$3:$D$11" in chart


def test_slot_shapes_around_a_block_are_checked_when_filling():
    work, slots, band, rep, values = _standard()
    region = {"id": "reg", "binding": "reg", "kind": "region", "name": "区域", "rows": 2, "cols": 2,
              "locator": {"sheet": "对账单", "range": "G5:H6"}}
    out, issues = T.xlsx_fill(work, slots + [region], values)
    assert out == b"" and [i["code"] for i in issues if i["level"] == "error"] == ["band_region"]
    other = {**_band(first=5, last=7), "id": "b2", "binding": "b2", "name": "另一段"}
    out, issues = T.xlsx_fill(work, slots + [other], values)
    assert [i["message"] for i in issues if i["level"] == "error"] == [
        "第 6–8 行与重复内容块「订单段」（第 3–6 行）重叠：内容块之间不能交叉或嵌套"]
    oncell = {**_cust(), "id": "c5", "binding": "c5", "locator": {"sheet": "对账单", "cell": "H5"}}
    out, issues = T.xlsx_fill(work, slots + [oncell], values)
    assert [i["code"] for i in issues if i["level"] == "error"] == ["band_detail_cell"]


def test_a_nested_detail_without_split_rows_is_refused_like_in_word():
    work, slots, band, rep, values = _standard()
    values[rep["id"]].pop("groups")
    out, issues = T.xlsx_fill(work, slots, values)
    assert out == b"" and [i["message"] for i in issues if i["level"] == "error"] == [
        "明细表「明细」在重复内容块「订单段」里：请在这个内容块的映射中设置分段字段，每一段才会只写入自己的明细行"]


def test_limits_on_rows_cells_and_expanded_formulas(monkeypatch):
    from synwevo_docs.engine import limits as L
    many = [(f"O-{i}", "2024-01-01", [("件", 1, "1.00")]) for i in range(300)]
    work, slots, band, rep, values = _standard(orders=many)
    out, issues, errors = _fill(work, slots, values)
    assert out == b"" and [e["code"] for e in errors] == ["band_formula_limit"]
    assert errors[0]["message"] == ("公式 对账单!D1204 展开到每一份后超过 Excel 的限制（每个函数最多 255 个参数、每个公式最多 8,192 个字符）："
                                    "请改为引用整个内容块")
    monkeypatch.setattr(L.LIMITS, "max_cells", 20)
    work, slots, band, rep, values = _standard()
    out, issues, errors = _fill(work, slots, values)
    assert [e["code"] for e in errors] == ["band_cell_limit"]
    def far(wb):
        wb["对账单"].cell(1_048_574, 1, "表尾")
    work, slots, band, rep, values = _standard(far)
    monkeypatch.setattr(L.LIMITS, "max_cells", 2_000_000)
    out, issues, errors = _fill(work, slots, values)
    assert [e["message"] for e in errors] == ["重复内容块「订单段」展开后工作表「对账单」共 1048579 行，超过 Excel 工作表的上限 1,048,576 行"]


def test_many_copies_fill_in_reasonable_time():
    many = [(f"O-{i}", "2024-01-01", [("件", 1, "1.00"), ("件", 2, "2.00"), ("件", 3, "3.00")]) for i in range(500)]
    def extra(wb):
        wb["对账单"]["D8"] = "=SUM(D3:D6)"                  # no per-copy argument list for 500 copies
        for r in (9, 10, 11):
            wb["对账单"][f"D{r}"] = None
        wb["汇总"]["B1"] = "=SUM(对账单!D3:D6)"
        wb["汇总"]["B2"] = None
    work, slots, band, rep, values = _standard(extra, orders=many)
    t = time.perf_counter()
    out, issues, errors = _fill(work, slots, values)
    took = time.perf_counter() - t
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert ws["D3002"].value == "=SUM(D2999:D3001)" and ws["B2997"].value == "O-499"
    assert took < 20, took


def test_the_scan_refuses_nothing_in_a_workbook_without_references_into_the_block():
    x = XlsxDoc(_template())
    x.expand_shared_formulas()
    from synwevo_docs.engine.xlsx_band import scan
    cand = scan(x, "对账单", 3, 6)
    assert sorted(c[2] for c in cand) == ["B1", "B2", "D10", "D11", "D8", "D9"]
    with pytest.raises(XlsxError) as e:
        scan(x, "对账单", 5, 8)          # as if the block were rows 5–8: ROWS(A3:A6) crosses its edge
    assert e.value.code == "band_cross" and "对账单!D12" in str(e.value)


# ------------------------------------------------------------------ independent review (2026-09-29)
def test_a_block_column_on_a_row_that_becomes_a_detail_row_is_refused_in_either_order():
    """Review item 1: column first, detail row second used to be accepted; the column's value was
    then written onto whatever row the detail row turned into (the subtotal row, or the first line)."""
    work = _template()
    work, band = T.xlsx_add_slot(work, {"kind": "repeat_block", "sheet": "对账单", "first": 2, "last": 5, "name": "订单段"}, [])
    work, band = T.xlsx_add_block_column(work, {"block_id": band["id"], "cell": "G5", "name": "备注"}, [band])
    with pytest.raises(T.TemplateError) as e:
        T.xlsx_add_slot(work, {"kind": "repeat", "sheet": "对账单", "row": 4, "c1": 0, "c2": 2, "name": "明细"}, [band])
    assert (e.value.code, str(e.value)) == ("band_detail_cell", "单元格 G5 在重复内容块「订单段」的明细行上：明细行上的单元格只能是明细列")
    # a version stored with both is stopped when it is filled, not written wrong
    work, rep = _detail(work)
    out, issues = T.xlsx_fill(work, [band, rep], _values(band, rep))
    assert out == b"" and [i["code"] for i in issues if i["level"] == "error"] == ["band_detail_cell"]


def test_a_pivot_source_given_as_whole_columns_is_fine_beside_a_block():
    """Review item 3: `worksheetSource ref="A:B"` (how Excel stores a source picked by column
    headers) made block creation fail with 「无效的单元格地址：A」 — a raw ValueError."""
    from .test_xlsx_row_shift import _with_pivot_and_thread
    zin = zipfile.ZipFile(io.BytesIO(_with_pivot_and_thread(pivot_at=10, thread_at="H30")))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n in zin.namelist():
            d = zin.read(n)
            if n == "xl/pivotCache/pivotCacheDefinition1.xml":
                d = d.replace('ref="A1:B3" sheet="数据"'.encode(), 'ref="A:B" sheet="报表"'.encode())
            z.writestr(n, d)
    work, band = T.xlsx_add_slot(buf.getvalue(), {"kind": "repeat_block", "sheet": "报表", "first": 1, "last": 4, "name": "段"}, [])
    band = {**band, "columns": [{"key": "c0", "name": "品名", "cell": "A2"}]}
    got, issues = T.xlsx_fill(work, [band], {band["id"]: {"kind": "rows", "rows": [{"c0": V.text("甲")}, {"c0": V.text("乙")}]}})
    assert got and not [i for i in issues if i["level"] == "error"], issues
    wb = openpyxl.load_workbook(io.BytesIO(got))
    assert [wb["报表"].cell(r, 1).value for r in (2, 6)] == ["甲", "乙"]
    assert [p.location.ref for p in wb["报表"]._pivots] == ["A14:B17"]     # moved down with its cells
    from synwevo_docs.engine.xlsx_band import _span
    assert [_span(r) for r in ("A:B", "$A:$B", "3:5", "B2:C9", "名称", "")] == \
        [(1, 1_048_576), (1, 1_048_576), (3, 5), (2, 9), None, None]


def test_formulas_inside_the_block_pointing_to_the_last_row_stay_on_the_sheet():
    """Review item 4: a range from inside the block down to the last row was pushed past it."""
    def extra(wb):
        ws = wb["对账单"]
        ws["G6"] = "=COUNTA(A10:A1048576)"
        ws["H6"] = "=B1048576"
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert [ws["G7"].value, ws["G11"].value] == ["=COUNTA(A15:A1048576)"] * 2
    assert [ws["H7"].value, ws["H11"].value] == ["=#REF!"] * 2              # pushed off the sheet, as a row shift does
    # not reported as a section without lines: every order here has lines
    assert not [i for i in issues if i["code"] == "ref_removed"]


def test_a_running_balance_on_a_detail_line_is_filled_down_within_its_copy():
    """Review item 2: `C5 = N(C4)+B5` gave 1, 2, 3 inside a block (every line looked at the header
    row) and 1, 3, 6 on a single detail row. A detail line's relative reference to a row above it
    in the block is filled down: line i looks at line i-1, the first line at the row above."""
    def extra(wb):
        ws = wb["对账单"]
        ws["G4"], ws["G5"] = "余额", "=N(G4)+B5"
        ws["H5"] = "=SUM(B4:B5)"                  # a two-row window, filled down
        ws["I5"] = "=N($G$4)+B5"                   # an absolute row stays where it is
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert [ws["G5"].value, ws["G6"].value, ws["G10"].value] == ["=N(G4)+B5", "=N(G5)+B6", "=N(G9)+B10"]
    assert [ws["H5"].value, ws["H6"].value, ws["H10"].value] == ["=SUM(B4:B5)", "=SUM(B5:B6)", "=SUM(B9:B10)"]
    assert [ws["I5"].value, ws["I6"].value, ws["I10"].value] == ["=N($G$4)+B5", "=N($G$4)+B6", "=N($G$9)+B10"]
    calc = _book(out, data_only=True)["对账单"]
    assert [calc["G5"].value, calc["G6"].value, calc["G10"].value] == [5, 8, 10]     # 5, 5+3; the next order starts again


@soffice
def test_libreoffice_agrees_with_the_filled_down_running_balance():
    from synwevo_docs.engine.converters import recalc_xlsx
    work, slots, band, rep, values = _standard(lambda wb: wb["对账单"].__setitem__("G5", "=N(G4)+B5"))
    out, _, _ = _fill(work, slots, values)
    lo = openpyxl.load_workbook(io.BytesIO(recalc_xlsx(out)), data_only=True)["对账单"]
    ours = _book(out, data_only=True)["对账单"]
    assert [lo[r].value for r in ("G5", "G6", "G10")] == [ours[r].value for r in ("G5", "G6", "G10")] == [5, 8, 10]


def _rules_of(out) -> list[tuple[str, list[str], str]]:
    """(sqref, formulas, priority) of each conditional format rule, in the file's order."""
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml"))
    return [(cf.get("sqref"), [f.text for f in rule.iter(f"{{{MAIN}}}formula")], rule.get("priority"))
            for cf in root.iter(f"{{{MAIN}}}conditionalFormatting") for rule in cf.iter(f"{{{MAIN}}}cfRule")]


def test_conditional_formats_follow_the_copies_like_cell_formulas():
    """Review item 5: a header-row rule looking at the order's subtotal (`$D6`, relative row) kept a
    fixed offset in every copy — pointing at a line whenever the order had other than one line; a
    threshold without `$` (`D5>H1`) drifted with each copy. Each copy now gets the reference a cell
    formula there would get; rules that read the same from their own first cell stay one rule."""
    def extra(wb):
        ws = wb["对账单"]
        ws.conditional_formatting.add("A3:D3", FormulaRule(formula=["$D6>100"]))      # header → subtotal
        ws.conditional_formatting.add("D5", FormulaRule(formula=["D5>H1"]))           # line → fixed threshold
        ws.conditional_formatting.add("D6", FormulaRule(formula=["D6>100"]))          # the subtotal row itself
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    rules = _rules_of(out)
    assert [(sq, f) for sq, f, _ in rules] == [
        ("A3:D3", ["$D$7>100"]), ("A8:D8", ["$D$11>100"]),       # each header looks at its own subtotal
        ("D5:D6 D10", ["D5>H$1"]),                                # every line, the same threshold: one rule
        ("D7 D11", ["D7>100"]),                                   # re-anchored where its first area now is
    ]
    assert sorted(int(p) for _, _, p in rules) == [1, 2, 3, 4]    # every rule its own priority, in order


def test_data_validations_follow_the_copies_like_cell_formulas():
    def extra(wb):
        dv = DataValidation(type="custom", formula1="B5<=B6")     # a line's quantity against the subtotal row
        dv.add("B5")
        wb["对账单"].add_data_validation(dv)
    work, slots, band, rep, values = _standard(extra)
    out, issues, errors = _fill(work, slots, values)
    assert not errors, errors
    ws = _book(out)["对账单"]
    assert [(str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation] == [("B5:B6", "B5<=B$7"), ("B10", "B10<=B$11")]
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml"))
    assert root.find(f"{{{MAIN}}}dataValidations").get("count") == "2"


def test_an_ordinary_rule_from_the_blocks_first_row_down_the_sheet_is_accepted():
    """Review item 6: `A3:H200 =$D3>100` (each row looks at its own row) was refused because the
    block starts at row 3; the same rule from row 1 was accepted."""
    for sqref, formula, want in (("A3:H200", "$D3>100", "A3:H7 A8:H11 A12:H205"),
                                 ("A1:H200", "$D1>100", "A1:H2 A3:H7 A8:H11 A12:H205"),
                                 ("A5:H5", "$D5>100", "A5:H6 A10:H10")):
        work, slots, band, rep, values = _standard(lambda wb: wb["对账单"].conditional_formatting.add(sqref, FormulaRule(formula=[formula])))
        out, issues, errors = _fill(work, slots, values)
        assert not errors, (sqref, errors)
        assert [(sq, f) for sq, f, _ in _rules_of(out)] == [(want, [formula])]


def test_page_break_counts_are_rewritten_when_the_block_goes():
    """Review item 7: breaks inside a block with no data were removed but the counts kept."""
    from openpyxl.worksheet.pagebreak import Break

    def extra(wb):
        wb["对账单"].row_breaks.append(Break(id=4))           # inside the block
        wb["对账单"].row_breaks.append(Break(id=12))          # below it
    work, slots, band, rep, values = _standard(extra, orders=[])
    out, _, errors = _fill(work, slots, values)
    assert not errors
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml"))
    rb = root.find(f"{{{MAIN}}}rowBreaks")
    assert [b.get("id") for b in rb] == ["8"] and (rb.get("count"), rb.get("manualBreakCount")) == ("1", "1")
    # only a break inside the block: the empty list goes (a count over no break is what Excel repairs)
    work, slots, band, rep, values = _standard(lambda wb: wb["对账单"].row_breaks.append(Break(id=4)), orders=[])
    out, _, _ = _fill(work, slots, values)
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("xl/worksheets/sheet1.xml"))
    assert root.find(f"{{{MAIN}}}rowBreaks") is None
