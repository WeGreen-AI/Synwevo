"""CI-103: a repeat row inside a repeat content block (or a block inside a repeat row) was
accepted and filled wrongly without a word — the first block copy took every detail row, the
other copies kept the template row still wrapped in field controls, and the fill reported no
issue. A block inside a repeat row is still refused. A repeat row inside a block is now a nested
detail (docs-nested-details.md, test_docx_nested.py): allowed, and filled per block copy only from
the sub-groups the plan makes — rows without them are refused at fill time, never guessed."""
from __future__ import annotations

import io

import docx
import pytest

from synwevo_docs.engine import docx_core as dc
from synwevo_docs.engine import templates as T
from synwevo_docs.engine import values as V
from synwevo_docs.engine.ooxml import Package


def _template() -> bytes:
    """TEST FIXTURE: 订单 paragraph, a two-row table (header, detail row), a closing paragraph."""
    d = docx.Document()
    d.add_paragraph("订单号：O-0")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "品名", "数量"
    t.cell(1, 0).text, t.cell(1, 1).text = "示例", "1"
    d.add_paragraph("（订单结束）")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _block_over_everything(work: bytes) -> tuple[bytes, list[dict]]:
    work, blk = T.docx_add_slot(work, {"kind": "repeat_block", "first": "0", "last": "2", "name": "订单段"}, [])
    return work, [blk]


def test_a_repeat_row_may_be_made_inside_a_repeat_block():
    work, slots = _block_over_everything(_template())
    work, rep = T.docx_add_slot(work, {"kind": "repeat", "table_path": "1", "row": 1, "name": "明细"}, slots)
    slots = T.assign_parents(work, slots + [rep])
    assert slots[1]["parent"] == slots[0]["id"] and "parent" not in slots[0]


def test_a_block_cannot_be_made_inside_a_repeat_row():
    # the table row repeats; a block inside one of its cells is refused
    d = docx.Document()
    t = d.add_table(rows=2, cols=1)
    t.cell(0, 0).text = "品名"
    t.cell(1, 0).text = "示例"
    t.cell(1, 0).add_paragraph("备注：无")
    buf = io.BytesIO()
    d.save(buf)
    work, rep = T.docx_add_slot(buf.getvalue(), {"kind": "repeat", "table_path": "0", "row": 1, "name": "明细"}, [])
    pkg = Package(work)
    body = dc.part_container(pkg, pkg.main_part())
    cell_paths = [p for p in _paths(body) if p.startswith("0:1:")]
    with pytest.raises(T.TemplateError) as e:
        T.docx_add_slot(work, {"kind": "repeat_block", "first": cell_paths[-1], "last": cell_paths[-1], "name": "备注段"}, [rep])
    assert e.value.code in ("inside_repeat_row", "nested_block")


def _paths(container) -> list[str]:
    out = []
    for i, el in enumerate(dc.iter_blocks(container)):
        out.append(str(i))
        if el.tag == dc.w("tbl"):
            for r, tr in enumerate(dc.iter_rows(el)):
                for c, tc in enumerate(dc.iter_cells(tr)):
                    for k, _ in enumerate(dc.iter_blocks(tc)):
                        out.append(f"{i}:{r}:{c}:{k}")
    return out


def test_nested_rows_without_sub_groups_are_refused_not_filled_wrongly():
    """TEST FIXTURE: the shape an earlier version stored (block made first, then the repeat row
    inside it, which nothing refused) — built here with the same wrapping calls — filled with
    detail rows that were not split per block copy (no `groups`)."""
    work, slots = _block_over_everything(_template())
    pkg = Package(work)
    body = dc.part_container(pkg, pkg.main_part())
    tbl = next(el for el in dc.iter_blocks(body) if el.tag == dc.w("tbl"))
    tr = list(dc.iter_rows(tbl))[1]
    rep = {"id": "rp_fixture", "name": "明细", "kind": "repeat", "binding": "rp_fixture", "locator": {"table_path": "1", "row": 1},
           "header_rows": 1, "empty_policy": "remove_row", "columns": []}
    for ci, tc in enumerate(dc.iter_cells(tr)):
        p = dc.cell_first_paragraph(tc)
        dc.wrap_range(p, 0, len(dc.para_text(p)), f"{dc.TAG_PREFIX}{rep['id']}#c{ci}", f"c{ci}")
        rep["columns"].append({"key": f"c{ci}", "name": f"c{ci}", "col": ci})
    work = pkg.save()
    slots = slots + [rep]
    row = lambda a, b: {"c0": V.from_python(a), "c1": V.from_python(b)}  # noqa: E731
    values = {slots[0]["id"]: {"kind": "rows", "rows": [{}, {}]},
              rep["id"]: {"kind": "rows", "rows": [row("螺丝", 2), row("螺母", 3)]}}
    out, issues = T.docx_fill(work, slots, values)
    assert [i["code"] for i in issues] == ["nested_unconfigured"] and issues[0]["level"] == "error"
    assert "明细" in issues[0]["message"] and "订单段" in issues[0]["message"]
    assert out == work, "nothing is written"
