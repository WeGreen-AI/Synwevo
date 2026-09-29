"""Nested details in XLSX templates through the service (docs/platform/docs-xlsx-nested.md): a
repeat content block over whole rows, its columns and the detail row inside it are created as
fields; the block's mapping sets 分段字段; each order copy gets only its own lines.

TEST FIXTURES: the source workbook comes from `test_docx_nested` (synthetic 客户 × 订单 × 明细
rows); the XLSX template is built here. Expected values are computed from those rows.
"""
from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest

from .test_api_v1 import _ok, keys  # noqa: F401  (keys is a fixture)
from .test_docx_nested import GROUP_BY, ROWS, TABLE, _workbook

SHEET = "对账单"


def _template_bytes(extra=None) -> bytes:
    """Customer B1; block rows 3–6: order header (B3 order no., D3 date), line header, the line
    row 5 (D5 = amount formula), the subtotal row 6; the grand total at D8."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"], ws["B1"] = "客户：", "C-0"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "订单号：", "O-0", "下单日期：", "2024-01-01"
    for c, h in enumerate(("品名", "数量", "单价", "金额"), start=1):
        ws.cell(4, c, h)
    ws["A5"], ws["B5"], ws["C5"], ws["D5"] = "示例", 1, 0, "=B5*C5"
    ws["A6"], ws["D6"] = "订单小计", "=SUM(D5:D5)"
    ws["A8"], ws["D8"] = "合计", "=SUM(D6)"
    if extra:
        extra(ws)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _template(tmp_path, extra=None):
    p = tmp_path / "对账单模板.xlsx"
    p.write_bytes(_template_bytes(extra))
    return p


def _refusal(r) -> tuple[str, str]:
    """(code, sentence) of a refused workbench request."""
    assert r.status_code == 400, r.text
    return r.json()["error"], r.json()["message"]


def _slots(api, tid):
    return {s["id"]: s for s in api.ok(api.c.get(f"/api/templates/{tid}"))["version"]["slots"]}


def _setup(api, tmp_path, rows=ROWS, row_first=False, block=None, row=None, extra=None):
    pid = api.project("表格嵌套明细")
    api.upload(pid, _workbook(tmp_path / "订单.xlsx", rows))
    tid = api.template(pid, _template(tmp_path, extra))["id"]
    make_block = lambda: api.slot(tid, kind="repeat_block", sheet=SHEET, first=2, last=5, name="订单段", **(block or {}))  # noqa: E731
    make_row = lambda: api.slot(tid, kind="repeat", sheet=SHEET, row=4, c1=0, c2=2, name="明细", **(row or {}))  # noqa: E731
    if row_first:
        rep = make_row()
        blk = make_block()
    else:
        blk = make_block()
        rep = make_row()
    api.slot(tid, kind="block_column", block_id=blk["id"], cell="B3", name="订单号")
    blk = api.slot(tid, kind="block_column", block_id=blk["id"], cell="D3", name="下单日期")
    cust = api.slot(tid, kind="cell", sheet=SHEET, cell="B1", name="客户")
    return pid, tid, blk, rep, cust


def _map(api, pid, blk, rep, cust, *, group_by=GROUP_BY, block_cols=None, child=True):
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "客户"}], "sort": [{"header": "订单号"}, {"header": "行号", "numeric": True}]}})
    bk = {c["name"]: c["key"] for c in blk["columns"]}
    rk = {c["name"]: c["key"] for c in rep["columns"]}
    blockmap = {"type": "repeat", "sources": [TABLE], "target": {"binding": blk["binding"]},
                "columns": block_cols or [{"slot_col": bk["订单号"], "source": {"header": "订单号"}},
                                          {"slot_col": bk["下单日期"], "source": {"header": "日期文本"}}]}
    if group_by is not None:
        blockmap["group_by"] = group_by
    maps = [{"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "客户"}], "target": {"binding": cust["binding"]}},
            blockmap]
    if child:
        maps.append({"type": "repeat", "sources": [TABLE], "target": {"binding": rep["binding"]},
                     "columns": [{"slot_col": rk["品名"], "source": {"header": "品名"}},
                                 {"slot_col": rk["数量"], "source": {"header": "数量"}},
                                 {"slot_col": rk["单价"], "source": {"header": "单价"}}]})
    return api.maps(pid, maps)


def _orders(customer: str, rows=ROWS) -> list:
    mine = sorted((r for r in rows if r[0] == customer), key=lambda r: (r[1], r[3] or 0))
    out: dict[str, list] = {}
    for r in mine:
        out.setdefault(r[1], []).append(r)
    return list(out.items())


def _read(data: bytes, customer: str, rows=ROWS) -> None:
    """The statement for one customer, cell by cell, against the fixture rows."""
    ws = openpyxl.load_workbook(io.BytesIO(data))[SHEET]
    calc = openpyxl.load_workbook(io.BytesIO(data), data_only=True)[SHEET]
    assert ws["B1"].value == customer
    r = 3
    subtotals = []
    for no, lines in _orders(customer, rows):
        assert [ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"D{r}"].value] == ["订单号：", no, lines[0][2].isoformat()]
        assert ws[f"A{r + 1}"].value == "品名"
        first = r + 2
        for i, line in enumerate(lines):
            assert [ws[f"A{first + i}"].value, ws[f"B{first + i}"].value, ws[f"D{first + i}"].value] == \
                [line[4], line[5], f"=B{first + i}*C{first + i}"]
            assert Decimal(str(ws[f"C{first + i}"].value)) == line[6]
        last = first + len(lines) - 1
        assert [ws[f"A{last + 1}"].value, ws[f"D{last + 1}"].value] == ["订单小计", f"=SUM(D{first}:D{last})"]
        amount = sum(line[5] * line[6] for line in lines)
        assert Decimal(str(calc[f"D{last + 1}"].value)).quantize(Decimal("0.01")) == amount
        subtotals.append(f"D{last + 1}")
        r = last + 2
    assert ws[f"A{r + 1}"].value == "合计" and ws[f"D{r + 1}"].value == f"=SUM({','.join(subtotals)})"
    total = sum(line[5] * line[6] for _, lines in _orders(customer, rows) for line in lines)
    assert Decimal(str(calc[f"D{r + 1}"].value)).quantize(Decimal("0.01")) == total


def test_each_order_copy_gets_only_its_own_lines(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    slots = _slots(api, tid)
    assert slots[rep["id"]].get("parent") == blk["id"] and "parent" not in slots[blk["id"]]
    assert [(c["key"], c["name"], c["cell"]) for c in slots[blk["id"]]["columns"]] == [("c0", "订单号", "B3"), ("c1", "下单日期", "D3")]
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    items = api.items(job["id"])
    assert job["status"] == "completed", items
    outs = api.outputs(job["id"])
    _read(outs["A.xlsx"], "A")
    _read(outs["B.xlsx"], "B")
    # the reopened output was checked value by value (texts; numbers are checked as numbers)
    ver = next(a for i in items for a in i["artifacts"] if a["filename"] == "A.xlsx")["verification"]
    assert ver["reopened"] and ver["checks"] >= 5 and ver["missing"] == []


def test_the_row_may_be_made_first_and_removing_the_block_clears_parent(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path, row_first=True)
    assert _slots(api, tid)[rep["id"]].get("parent") == blk["id"]
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed"
    _read(api.outputs(job["id"])["A.xlsx"], "A")
    api.ok(api.c.delete(f"/api/templates/{tid}/slots/{blk['id']}"))
    slots = _slots(api, tid)
    assert blk["id"] not in slots and "parent" not in slots[rep["id"]]


def _preview_issues(api, pid, key) -> list[dict]:
    return api.ok(api.c.get(f"/api/projects/{pid}/preview", params={"search": key}))["issues"]


def test_a_nested_detail_needs_the_block_to_be_split(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    _map(api, pid, blk, rep, cust, group_by=None)
    issues = [i for i in _preview_issues(api, pid, "A") if i["level"] == "error"]
    assert [i["code"] for i in issues] == ["nested_needs_group_by"]
    assert issues[0]["message"] == "明细表「明细」在重复内容块「订单段」里：请在这个内容块的映射中设置分段字段，每一段才会只写入自己的明细行"


def test_group_by_on_a_single_detail_row_is_still_refused(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    maps = _map(api, pid, blk, rep, cust)["mappings"]
    for m in maps:
        if (m.get("target") or {}).get("binding") == rep["binding"]:
            m["group_by"] = GROUP_BY
    api.maps(pid, maps)
    e = [i for i in _preview_issues(api, pid, "A") if i["code"] == "group_by_unsupported"]
    assert len(e) == 1 and e[0]["message"].endswith("：分段目前只支持 Word 和 Excel 模板的重复内容块；PPT / PDF 模板和表格行暂不支持嵌套明细")


EMPTY_ORDER = ("A", "A-3", ROWS[0][2], None, None, None, None)


@pytest.mark.parametrize("policy", ["remove_row", "blank_row", "placeholder"])
def test_an_order_without_lines_follows_the_detail_rows_empty_policy(api, tmp_path, policy):
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=ROWS + [EMPTY_ORDER], row={"empty_policy": policy})
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed", [(i["status"], i.get("issues")) for i in api.items(job["id"])]
    ws = openpyxl.load_workbook(io.BytesIO(api.outputs(job["id"])["A.xlsx"]))[SHEET]
    # A-1 (rows 3–7), A-2 (8–11), then A-3
    assert ws["B12"].value == "A-3" and ws["A13"].value == "品名"
    if policy == "remove_row":
        assert [ws["A14"].value, ws["D14"].value] == ["订单小计", "=SUM(#REF!)"]
        assert ws["D16"].value == "=SUM(D7,D11,D14)"
    else:
        assert ws["A14"].value == ("（无）" if policy == "placeholder" else None)
        assert [ws["A15"].value, ws["D15"].value] == ["订单小计", "=SUM(D14:D14)"]
        assert ws["D17"].value == "=SUM(D7,D11,D15)"


def test_a_line_without_an_order_number_is_an_error_unless_told_to_skip(api, tmp_path):
    stray = ("A", None, ROWS[0][2], 9, "散件", 1, Decimal("9.00"))
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=ROWS + [stray])
    _map(api, pid, blk, rep, cust)
    assert [i["code"] for i in _preview_issues(api, pid, "A") if i["level"] == "error"] == ["subgroup_key_empty"]
    _map(api, pid, blk, rep, cust, group_by={**GROUP_BY, "empty_key": "skip"})
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed"
    _read(api.outputs(job["id"])["A.xlsx"], "A")


def test_block_columns_are_checked_renamed_and_removed(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    def refused(code, **req):
        got, message = _refusal(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "block_column", "block_id": blk["id"], **req}))
        assert got == code, message
        return message
    assert refused("band_column_outside", cell="B8", name="合计") == "请在重复内容块「订单段」的行内（第 3–6 行）选择单元格"
    assert refused("band_detail_cell", cell="D5", name="金额") == "单元格 D5 在重复内容块「订单段」的明细行上：明细行上的单元格只能是明细列"
    assert refused("duplicate_column", cell="B3", name="又一次") == "B3 已是重复内容块「订单段」的列「订单号」"
    assert _refusal(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "cell", "sheet": SHEET, "cell": "B3", "name": "撞列"})) == \
        ("overlap", "与已有字段「订单段」重叠")
    api.ok(api.c.patch(f"/api/templates/{tid}/slots/{blk['id']}", json={"columns": [{"key": "c1", "name": "日期"}]}))
    assert [c["name"] for c in _slots(api, tid)[blk["id"]]["columns"]] == ["订单号", "日期"]
    api.ok(api.c.patch(f"/api/templates/{tid}/slots/{blk['id']}", json={"remove_column": "c1"}))
    assert [c["key"] for c in _slots(api, tid)[blk["id"]]["columns"]] == ["c0"]


def test_what_cannot_follow_the_copies_is_refused_when_the_block_is_made(api, tmp_path):
    pid = api.project("拒绝")
    def refused(extra, first=2, last=5, code=None):
        tid = api.template(pid, _template(tmp_path, extra))["id"]
        got, message = _refusal(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "repeat_block", "sheet": SHEET, "first": first, "last": last, "name": "段"}))
        assert got == code, message
        return tid, message
    _, msg = refused(lambda ws: ws.__setitem__("E8", "=D6*1.1"), code="band_reference")
    assert msg.startswith("单元格 对账单!E8 的公式引用了重复内容块内的 D6：")
    refused(lambda ws: ws.merge_cells("F6:F7"), code="band_merge")
    refused(None, first=2, last=2, code="band_rows")
    # a block over rows that already hold another block, a region inside a block
    tid = api.template(pid, _template(tmp_path))["id"]
    api.slot(tid, kind="repeat_block", sheet=SHEET, first=2, last=5, name="订单段")
    assert _refusal(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "repeat_block", "sheet": SHEET, "first": 4, "last": 7, "name": "另一段"})) == \
        ("band_overlap", "第 5–8 行与重复内容块「订单段」（第 3–6 行）重叠：内容块之间不能交叉或嵌套")
    assert _refusal(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "region", "sheet": SHEET, "range": "F4:G5", "name": "区域"}))[0] == "band_region"
    api.slot(tid, kind="repeat", sheet=SHEET, row=4, c1=0, c2=2, name="明细")
    assert _refusal(api.c.post(f"/api/templates/{tid}/slots", json={"kind": "cell", "sheet": SHEET, "cell": "F5", "name": "明细行上"}))[0] == "band_detail_cell"
    # a cell field elsewhere in the block is fine: every copy gets the same value
    api.slot(tid, kind="cell", sheet=SHEET, cell="F3", name="块内")


FIELDS = "docweave.template-fields"


def test_definitions_create_a_block_and_its_columns(api, keys, tmp_path):
    c, h = api.c, keys["rw"]
    p = _template(tmp_path)
    t = _ok(c.post("/api/v1/templates", files={"file": (p.name, p.read_bytes())}, headers=h))
    tid = t["id"]
    _ok(c.post(f"/api/v1/templates/{tid}/fields", headers=h, json={"schema": FIELDS, "version": 1, "fields": [
        {"kind": "repeat_block", "name": "订单段", "sheet": SHEET, "rows": "3:6", "separator": "blank"},
        {"kind": "repeat", "name": "明细", "sheet": SHEET, "row": 4, "c1": 0, "c2": 2},
        {"kind": "block_column", "name": "订单号", "block": "订单段", "cell": "B3"}]}))
    fields = {f["name"]: f for f in _ok(c.get(f"/api/v1/templates/{tid}/fields", headers=h))["fields"]}
    assert fields["订单段"]["locator"] == {"sheet": SHEET, "first": 2, "last": 5}
    assert [col["cell"] for col in fields["订单段"]["columns"]] == ["B3"]
    assert _slots(api, tid)[fields["明细"]["id"]]["parent"] == fields["订单段"]["id"]
    r = c.post(f"/api/v1/templates/{tid}/fields", headers=h, json={"schema": FIELDS, "version": 1, "fields": [
        {"kind": "repeat_block", "name": "坏", "sheet": SHEET, "rows": "三到六"}]})
    assert r.status_code == 400 and r.json()["error"]["field"] == "rows"
