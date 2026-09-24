"""Acceptance scenarios 1–4 through the HTTP API (real files, real generation)."""
import io
import re
import zipfile

import docx
import openpyxl

from .helpers import docx_text


def test_s1_per_row_100_contracts(api, fixtures):
    """XLSX 100 customers → 100 DOCX; first/last/sampled records correct; leading zeros kept."""
    pid = api.project("S1")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    tid = t["id"]
    s_name = api.text_slot(tid, "示例客户名称", "客户名称")            # spans 2 runs with different formatting
    s_id = api.text_slot(tid, "000000", "客户编号", ctx="客户编号 000000")
    s_date = api.text_slot(tid, "2020-01-01", "签约日期")
    s_amt = api.text_slot(tid, "0.00", "合同金额", ctx="合同金额")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}, "id_column": {"header": "客户编号"}})
    col = lambda h: {"kind": "column", "sheet": "客户", "header": h}
    api.maps(pid, [
        {"type": "single", "sources": [col("姓名")], "target": {"binding": s_name["binding"]}},
        {"type": "single", "sources": [col("客户编号")], "target": {"binding": s_id["binding"]}},
        {"type": "single", "sources": [col("签约日期")], "target": {"binding": s_date["binding"]}, "transforms": [{"op": "date_format", "pattern": "yyyy年m月d日"}]},
        {"type": "single", "sources": [col("金额")], "target": {"binding": s_amt["binding"]}},
    ])
    summ = api.ok(api.c.get(f"/api/projects/{pid}/generation-summary"))
    assert summ["records"] == 100 and summ["valid"] == 100 and summ["expected_outputs"] == 100
    job = api.job(pid, naming="{客户编号}_{客户名称}")
    assert job["status"] == "completed", job
    assert job["counts"]["success"] == 100 and job["counts"]["files"] == 100
    outs = api.outputs(job["id"])
    assert len(outs) == 100
    wb = openpyxl.load_workbook(fixtures / "测试_客户100.xlsx")
    ws = wb["客户"]
    for i in (1, 37, 100):
        cid, name, date, amt = ws.cell(i + 1, 1).value, ws.cell(i + 1, 2).value, ws.cell(i + 1, 3).value, ws.cell(i + 1, 4).value
        fname = f"{cid}_{name}.docx"
        assert fname in outs, fname
        text = docx_text(outs[fname])
        assert f"甲方（客户）：{name}，客户编号 {cid}。" in text
        assert f"签约日期：{date.year}年{date.month}月{date.day}日；合同金额：人民币 {amt:,.2f} 元。" in text
        assert "乙方：测试服务有限公司（固定文字，不应被修改）。" in text
    assert "000001_" in next(iter(sorted(outs)))                       # leading zeros preserved in names and content
    # ZIP contains exactly the files + manifest
    z = zipfile.ZipFile(io.BytesIO(api.c.get(f"/api/jobs/{job['id']}/zip").content))
    assert len([n for n in z.namelist() if n.endswith(".docx")]) == 100 and "生成清单.csv" in z.namelist()


def _report_template(api, pid, fixtures):
    t = api.template(pid, fixtures / "测试_检测报告模板.docx")
    tid = t["id"]
    s_no = api.text_slot(tid, "XXXX", "样品编号")
    s_org = api.text_slot(tid, "某某单位", "送检单位")
    s_date = api.text_slot(tid, "YYYY-MM-DD", "检测日期")
    s_con = api.text_slot(tid, "待定", "综合结论")
    v = api.view(tid)
    tbl = [b for b in v["doc"]["parts"][0]["blocks"] if b["t"] == "tbl"][0]
    rep = api.slot(tid, kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="检测明细")
    return tid, s_no, s_org, s_date, s_con, rep


def test_s2_cross_file_reuse_20(api, fixtures):
    """20 same-structure DOCX sources; mappings built once on the representative file."""
    pid = api.project("S2")
    files = sorted((fixtures / "检测表20").glob("*.docx"))
    api.upload(pid, *files)
    srcs = api.ok(api.c.get(f"/api/projects/{pid}/sources"))
    rep = next(s for s in srcs if s["representative"])
    tid, s_no, s_org, s_date, s_con, rep_slot = _report_template(api, pid, fixtures)
    api.rule(pid, {"mode": "per_file"})

    def loc(label_prefix):
        # build the locator from a selection on the representative file (server side), like the UI does
        view = api.ok(api.c.get(f"/api/sources/{rep['id']}/view"))
        for b in view["doc"]["parts"][0]["blocks"]:
            if b["t"] == "p" and b["text"].startswith(label_prefix):
                return api.ok(api.c.post(f"/api/sources/{rep['id']}/locator", json={
                    "project_id": pid, "selection": {"kind": "docx_text", "part": b["part"], "path": b["path"], "start": len(label_prefix), "end": len(b["text"])}}))["proposals"][0]["locator"]
        raise AssertionError(label_prefix)

    view = api.ok(api.c.get(f"/api/sources/{rep['id']}/view"))
    tbl = [b for b in view["doc"]["parts"][0]["blocks"] if b["t"] == "tbl"][0]
    detail = api.ok(api.c.post(f"/api/sources/{rep['id']}/locator", json={"project_id": pid, "selection": {"kind": "docx_table", "part": tbl["part"], "table_path": tbl["path"], "header_rows": 1}}))
    api.maps(pid, [
        {"type": "single", "sources": [loc("样品编号：")], "target": {"binding": s_no["binding"]}},
        {"type": "single", "sources": [loc("送检单位：")], "target": {"binding": s_org["binding"]}},
        {"type": "single", "sources": [loc("检测日期：")], "target": {"binding": s_date["binding"]}},
        {"type": "single", "sources": [loc("结论：")], "target": {"binding": s_con["binding"]}},
        {"type": "repeat", "sources": [detail["proposals"][0]["locator"]], "target": {"binding": rep_slot["binding"]},
         "columns": [{"slot_col": "c0", "source": {"header": "检测项目"}}, {"slot_col": "c1", "source": {"header": "结果"}}, {"slot_col": "c2", "source": {"header": "判定"}}]},
    ])
    val = api.ok(api.c.get(f"/api/projects/{pid}/validate"))
    assert all(all(r["ok"] for r in f["results"].values()) for f in val["files"]), val
    job = api.job(pid, naming="{样品编号}")
    assert job["status"] == "completed" and job["counts"]["success"] == 20
    outs = api.outputs(job["id"])
    assert len(outs) == 20
    for i in (1, 7, 20):
        src = docx.Document(files[i - 1])
        src_rows = [[c.text for c in r.cells] for r in src.tables[0].rows[1:]]
        text = docx_text(outs[f"YP-{i:04d}.docx"])
        assert f"报告对象：样品 YP-{i:04d}，送检单位 {src.paragraphs[2].text.split('：')[1]}。" in text
        assert src.paragraphs[3].text.split("：")[1] in text
        out = docx.Document(io.BytesIO(outs[f"YP-{i:04d}.docx"]))
        got = [[c.text for c in r.cells] for r in out.tables[0].rows[1:]]
        assert got == src_rows, (got, src_rows)       # every file its own detail rows, in order


def test_s3_group_orders_15(api, fixtures):
    """100 detail rows grouped by 订单号 → 15 documents with only their own rows, sorted by 行号."""
    pid = api.project("S3")
    api.upload(pid, fixtures / "测试_订单明细100.xlsx")
    t = api.template(pid, fixtures / "测试_订单模板.docx")
    tid = t["id"]
    s_order = api.text_slot(tid, "DD-000", "订单号")
    s_cust = api.text_slot(tid, "某客户", "客户")
    v = api.view(tid)
    tbl = [b for b in v["doc"]["parts"][0]["blocks"] if b["t"] == "tbl"][0]
    rep = api.slot(tid, kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="明细")
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "订单号"}], "sort": [{"header": "行号", "dir": "asc", "numeric": True}], "scope": "file"}})
    recs = api.ok(api.c.get(f"/api/projects/{pid}/records"))
    assert recs["total"] == 15
    api.maps(pid, [
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "订单号"}], "target": {"binding": s_order["binding"]}},
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "客户"}], "target": {"binding": s_cust["binding"]}},
        {"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": rep["binding"]},
         "columns": [{"slot_col": "c0", "source": {"header": "行号"}}, {"slot_col": "c1", "source": {"header": "品名"}},
                     {"slot_col": "c2", "source": {"header": "数量"}}, {"slot_col": "c3", "source": {"header": "单价"}}]},
    ])
    job = api.job(pid, naming="{订单号}")
    assert job["status"] == "completed" and job["counts"]["success"] == 15
    outs = api.outputs(job["id"])
    ws = openpyxl.load_workbook(fixtures / "测试_订单明细100.xlsx")["明细"]
    expected: dict[str, list] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        expected.setdefault(r[0], []).append(r)
    total_rows = 0
    for order, rows in expected.items():
        rows.sort(key=lambda r: r[2])
        d = docx.Document(io.BytesIO(outs[f"{order}.docx"]))
        assert f"订单号：{order}" in d.paragraphs[1].text and rows[0][1] in d.paragraphs[1].text
        got = [[c.text for c in r.cells] for r in d.tables[0].rows[1:]]
        assert got == [[str(r[2]), r[3], str(r[4]), f"{r[5]:g}" if isinstance(r[5], float) else str(r[5])] for r in rows], (order, got)
        total_rows += len(got)
    assert total_rows == 100


def test_s3b_group_into_xlsx_template_with_formulas(api, fixtures):
    pid = api.project("S3b")
    api.upload(pid, fixtures / "测试_订单明细100.xlsx")
    t = api.template(pid, fixtures / "测试_订单模板.xlsx")
    tid = t["id"]
    s_order = api.slot(tid, kind="cell", sheet="订单", cell="B2", name="订单号")
    rep = api.slot(tid, kind="repeat", sheet="订单", row=4, c1=0, c2=3, name="明细")
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "订单号"}], "sort": [{"header": "行号", "dir": "asc", "numeric": True}]}})
    api.maps(pid, [
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "订单号"}], "target": {"binding": s_order["binding"]}},
        {"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": rep["binding"]},
         "columns": [{"slot_col": "c0", "source": {"header": "行号"}}, {"slot_col": "c1", "source": {"header": "品名"}},
                     {"slot_col": "c2", "source": {"header": "数量"}}, {"slot_col": "c3", "source": {"header": "单价"}}]},
    ])
    job = api.job(pid, naming="{订单号}")
    assert job["status"] == "completed", api.items(job["id"])[0]
    outs = api.outputs(job["id"])
    wbf = openpyxl.load_workbook(io.BytesIO(outs["DD-001.xlsx"]))
    wbv = openpyxl.load_workbook(io.BytesIO(outs["DD-001.xlsx"]), data_only=True)
    ws, wv = wbf["订单"], wbv["订单"]
    n = sum(1 for r in openpyxl.load_workbook(fixtures / "测试_订单明细100.xlsx")["明细"].iter_rows(min_row=2, values_only=True) if r[0] == "DD-001")
    assert ws["B2"].value == "DD-001"
    last = 4 + n
    assert ws[f"E{last}"].value == f"=C{last}*D{last}"                   # row formula translated for every new row
    assert ws[f"E{last + 1}"].value == f"=SUM(E5:E{last})"               # total expanded over all detail rows
    total = sum(wv[f"C{r}"].value * wv[f"D{r}"].value for r in range(5, last + 1))
    assert abs(wv[f"E{last + 1}"].value - total) < 1e-6                   # cached values computed by the built-in engine
    assert ws[f"A{last + 3}"].value == "备注：测试模板"                    # content below moved, not overwritten
    assert "A1:E1" in {str(m) for m in ws.merged_cells.ranges}
    assert ws["A5"].border.left.style == "thin" and ws[f"A{last}"].border.left.style == "thin"  # template row style copied


def test_s4_region_semantics(api, fixtures):
    """The same 10 rows: explicitly '10 samples' (10 files) vs 'one detail table' (1 file) – counts shown before generating."""
    pid = api.project("S4")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    rows_rule = {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_start": 1, "data_end": 10}}
    assert api.ok(api.c.post(f"/api/projects/{pid}/rule-preview", json=rows_rule))["total"] == 10
    assert api.ok(api.c.post(f"/api/projects/{pid}/rule-preview", json={"mode": "per_file"}))["total"] == 1
    # (a) each row is a sample → 10 documents
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    s_name = api.text_slot(t["id"], "示例客户名称", "客户名称")
    api.rule(pid, rows_rule)
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s_name["binding"]}}])
    job = api.job(pid)
    assert job["counts"]["success"] == 10
    # (b) the same rows are the detail table of ONE document
    t2 = api.template(pid, fixtures / "测试_订单模板.docx")
    v = api.view(t2["id"])
    tbl = [b for b in v["doc"]["parts"][0]["blocks"] if b["t"] == "tbl"][0]
    rep = api.slot(t2["id"], kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="明细")
    api.rule(pid, {"mode": "per_file"})
    api.maps(pid, [{"type": "repeat", "sources": [{"kind": "table", "sheet": "客户", "header_row": 0, "data_start": 1, "data_end": 10}],
                    "target": {"binding": rep["binding"]},
                    "columns": [{"slot_col": "c0", "source": {"header": "客户编号"}}, {"slot_col": "c1", "source": {"header": "姓名"}}]}])
    job = api.job(pid)
    assert job["counts"]["success"] == 1
    d = docx.Document(io.BytesIO(next(iter(api.outputs(job["id"]).values()))))
    assert len(d.tables[0].rows) == 11 and d.tables[0].rows[10].cells[0].text == "000010"
