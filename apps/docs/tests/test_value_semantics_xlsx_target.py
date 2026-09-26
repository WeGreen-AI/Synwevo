"""Value semantics into an XLSX target (docs/FORMATS.md §写入值): numbers stay numbers, dates
stay dates (a serial under the target's date format), text stays text with its leading zeros,
and a source cell whose *display* is padded (000123) written into a General cell is warned
about rather than silently changed."""
from __future__ import annotations

import datetime as dt
import io

import openpyxl


def _source() -> bytes:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["编号", "数量", "日期", "比例", "补零", "金额"])
    ws.append(["0007", 12, dt.datetime(2026, 9, 26), 0.125, 7, 1234.5])
    ws.append(["0008", 3, dt.date(2026, 1, 5), 1, 42, 0.1 + 0.2])
    ws["E2"].number_format = "000000"; ws["E3"].number_format = "000000"
    ws["D2"].number_format = "0.0%"; ws["D3"].number_format = "0%"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _template() -> bytes:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "单据"
    ws["A1"] = "编号"; ws["B1"] = "X"
    ws["A2"] = "数量"; ws["B2"] = 0
    ws["A3"] = "日期"; ws["B3"] = dt.datetime(2000, 1, 1); ws["B3"].number_format = "yyyy-mm-dd"
    ws["A4"] = "比例"; ws["B4"] = 0; ws["B4"].number_format = "0.0%"
    ws["A5"] = "补零（常规）"; ws["B5"] = 0
    ws["A6"] = "金额"; ws["B6"] = 0; ws["B6"].number_format = "#,##0.00"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def test_values_written_into_xlsx_cells_keep_their_types(api, tmp_path):
    pid = api.project("XLSX 目标取值")
    api.upload_bytes(pid, "数据.xlsx", _source())
    tpl = tmp_path / "单据.xlsx"; tpl.write_bytes(_template())
    tid = api.template(pid, tpl)["id"]
    cells = {"编号": "B1", "数量": "B2", "日期": "B3", "比例": "B4", "补零": "B5", "金额": "B6"}
    slots = {name: api.slot(tid, kind="cell", sheet="单据", cell=ref, name=name)["binding"] for name, ref in cells.items()}
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "数据", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "数据", "header": h}], "target": {"binding": b}} for h, b in slots.items()])
    job = api.job(pid, naming="{编号}", formats=["xlsx"])
    assert job["status"] == "completed", (job.get("error"), [(it["status"], it.get("error"), it.get("issues")) for it in api.items(job["id"])])
    outs = api.outputs(job["id"])
    assert set(outs) == {"0007.xlsx", "0008.xlsx"}, sorted(outs)
    got = {}
    for name, data in outs.items():
        ws = openpyxl.load_workbook(io.BytesIO(data))["单据"]
        got[name] = {ref: (ws[ref].value, ws[ref].number_format) for ref in cells.values()}
    a, b = got["0007.xlsx"], got["0008.xlsx"]
    assert a["B1"][0] == "0007" and b["B1"][0] == "0008", "text keeps its leading zeros"
    assert a["B2"][0] == 12 and b["B2"][0] == 3 and isinstance(a["B2"][0], int), "a number is a number"
    assert isinstance(a["B3"][0], dt.datetime) and a["B3"][0].date() == dt.date(2026, 9, 26) and a["B3"][1] == "yyyy-mm-dd", "a date under a date format is a real date"
    assert isinstance(b["B3"][0], dt.datetime) and b["B3"][0].date() == dt.date(2026, 1, 5)
    assert a["B4"][0] == 0.125 and a["B4"][1] == "0.0%", "a ratio stays numeric under the target's percent format"
    assert b["B4"][0] == 1
    assert a["B6"][0] == 1234.5 and a["B6"][1] == "#,##0.00"
    assert b["B6"][0] == 0.3, "no binary noise is written"
    # the padded source display into a General target: the number, plus a warning that names it
    assert a["B5"][0] == 7 and b["B5"][0] == 42
    items = api.items(job["id"])
    notes = [i["message"] for it in items for i in (it.get("issues") or []) if i.get("level") == "warning"]
    assert any("B5" in m and "000007" in m for m in notes), notes
    assert all(it["status"] == "success" for it in items), [(it["status"], it.get("issues")) for it in items]
