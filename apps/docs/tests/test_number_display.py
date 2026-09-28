"""Q14: a number is shown the way Excel's General format shows it — at most 15 significant
digits. Excel stores 3*1.2 as the double 3.5999999999999996 and writes exactly that into the
sheet XML; it shows 3.6. Synwevo filled documents with 3.5999999999999996, from such cached
values and from formulas its own engine evaluated, and a sum of them carried the noise on
(13.5999999999999996). The value itself keeps the whole double: it is what goes back into a
spreadsheet target.

TEST FIXTURES: the workbooks below are built here; one cell is patched to hold the digits Excel
writes (openpyxl shortens 3.5999999999999996 to 3.6 when it saves, which hid the defect)."""
from __future__ import annotations

import io
import zipfile
from decimal import Decimal

import docx
import openpyxl
import pytest

from synwevo_docs.engine import transforms as T
from synwevo_docs.engine import values as V
from synwevo_docs.engine.numfmt import format_value


@pytest.mark.parametrize("x, shown", [
    (3 * 1.2, "3.6"), (0.1 + 0.2, "0.3"), (1 / 3, "0.333333333333333"), (2.675, "2.675"), (-0.5, "-0.5"),
    (1e-7, "0.0000001"), (123456789.12345679, "123456789.123457"),
    (12345678901234567.0, "12345678901234568"), (9007199254740993.0, "9007199254740992"),   # integral: every digit kept
])
def test_general_shows_what_excel_shows(x, shown):
    assert format_value(x, None)[0] == shown and format_value(x, "General")[0] == shown
    v = V.from_python(x)
    assert v.text == shown
    assert v.v == V.float_repr(x), "the value keeps the whole double"


def test_arithmetic_on_a_shown_number_is_on_what_excel_sees():
    noisy = V.from_python(3 * 1.2)
    assert noisy.as_decimal() == Decimal("3.6")
    assert T.aggregate([V.from_python(10.0), noisy], "sum").text == "13.6"
    assert T.combine([V.from_python(0.1 + 0.2), V.from_python(0.7)], "sum").text == "1"
    assert T.apply(noisy, [{"op": "decimals", "n": 2}], []).text == "3.60"
    assert V.from_python(12345678901234567.0).as_decimal() == Decimal("12345678901234568")


def _excel_like(tmp_path):
    """B2 holds the double Excel writes for =3*1.2; C2 is a formula with no cached value (the
    engine evaluates it); D2 a plain 10."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["编号", "单价", "金额", "运费"])
    ws.append(["A", 3.6, "=B2*1", 10])
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode()
                assert "<v>3.6</v>" in text, "fixture drifted"
                data = text.replace("<v>3.6</v>", "<v>3.5999999999999996</v>").encode()
            dst.writestr(item, data)
    path = tmp_path / "数据.xlsx"
    path.write_bytes(out.getvalue())
    return path


def test_a_document_reads_like_the_spreadsheet(api, tmp_path):
    pid = api.project("数字显示")
    api.upload(pid, _excel_like(tmp_path))
    d = docx.Document()
    d.add_paragraph("单价：X单价 金额：X金额 合计：X合计 两位：X两位")
    d.save(tmp_path / "t.docx")
    tid = api.template(pid, tmp_path / "t.docx")["id"]
    slots = {n: api.text_slot(tid, f"X{n}", n)["binding"] for n in ("单价", "金额", "合计", "两位")}
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "数据", "header_row": 0}})
    col = lambda h: {"kind": "column", "sheet": "数据", "header": h}  # noqa: E731
    api.maps(pid, [
        {"type": "single", "sources": [col("单价")], "target": {"binding": slots["单价"]}},
        {"type": "single", "sources": [col("金额")], "target": {"binding": slots["金额"]}},
        {"type": "single", "sources": [col("单价"), col("运费")], "combine": {"op": "sum"}, "target": {"binding": slots["合计"]}},
        {"type": "single", "sources": [col("单价")], "transforms": [{"op": "decimals", "n": 2}], "target": {"binding": slots["两位"]}},
    ])
    job = api.job(pid, naming="{编号}")
    assert job["status"] == "completed", api.items(job["id"])
    text = docx.Document(io.BytesIO(api.outputs(job["id"])["A.docx"])).paragraphs[0].text
    assert text == "单价：3.6 金额：3.6 合计：13.6 两位：3.60", text


# ------------------------------------------------------------------ found by the independent review (Q14)

def test_a_sum_of_doubles_adds_them_whole_and_rounds_the_total_once():
    third = V.from_python(100 / 3)
    assert T.aggregate([third, third, third], "sum").text == "100"          # not 99.9999999999999
    assert T.aggregate([third, V.from_python(0.05 / 3)], "sum").text == "33.35"
    assert T.combine([third, third, third], "sum").text == "100"
    assert T.aggregate([V.from_python(1), V.from_python(2)], "sum").text == "3"


@pytest.mark.parametrize("fmt, shown", [("0.00", "0.12"), ("¥#,##0.00", "¥0.12"), ("0.0%", "11.5%")])
def test_a_number_format_formats_what_excel_sees(fmt, shown):
    # 0.1*1.15 is the double 0.11499999999999999; Excel formats 0.115, as ROUND does
    assert format_value(0.1 * 1.15, fmt)[0] == shown
    assert T.apply(V.from_python(0.1 * 1.15), [{"op": "decimals", "n": 2}], []).text == "0.12"


def test_a_sheet_parsed_before_the_fix_is_shown_the_new_way():
    """TEST FIXTURE: a sheet as an earlier build stored it (no text-rules marker, the double's
    full expansion as its text)."""
    from synwevo_docs.engine import tabular
    old = {"name": "数据", "index": 0, "rows": [[["float", "3.5999999999999996", "3.5999999999999996"],
                                               ["float", "0.11499999999999999", "0.11", "0.00"], ["str", "x", "x"]]]}
    sh = tabular.Sheet.from_json(old)
    assert [c.text for c in sh.rows[0]] == ["3.6", "0.12", "x"]
    assert sh.rows[0][0].v == "3.5999999999999996"                          # the value is kept whole
    assert tabular.Sheet.from_json(sh.to_json()).rows[0][0].text == "3.6"
    kept = {**sh.to_json()}
    kept["rows"][0][0][2] = "left as stored"                                # new sheets are not rewritten
    assert tabular.Sheet.from_json(kept).rows[0][0].text == "left as stored"


def test_a_formula_without_a_saved_result_is_shown_in_its_cell_format():
    """TEST FIXTURE: formulas openpyxl saves without cached values, in formatted cells."""
    import datetime as dt
    from synwevo_docs.engine import tabular
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["单价", "金额", "比例", "到期"])
    ws.append([0.1, "=A2*1.15", "=A2/4", "=D3+30"])
    ws.append([None, None, None, dt.date(2026, 1, 1)])
    ws["B2"].number_format = "0.00"
    ws["C2"].number_format = "0.0%"
    ws["D2"].number_format = "yyyy-mm-dd"
    buf = io.BytesIO()
    wb.save(buf)
    row = tabular.read_xlsx(buf.getvalue()).sheets[0].rows[1]
    assert [c.text for c in row[1:4]] == ["0.12", "2.5%", "2026-01-31"]
    assert row[3].t == V.DATE
