"""CI-95: TODAY() and NOW() mean the date and time of the person a document is calculated for.

A job fixes one instant at submission — the submission time in the submitter's time zone (the
browser's, else the profile's, else UTC) — and computes every formula with it; a preview uses the
viewer's time zone at the current moment. Before this, both read the server's UTC clock: someone
in UTC+8 generating before 08:00 got yesterday's date, and a rerun after midnight another result."""
from __future__ import annotations

import datetime as dt
import io
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from synwevo_docs import clock
from synwevo_docs.engine.formula import Book, as_of
from synwevo_docs.engine.numfmt import serial_to_datetime

# 2026-09-27 20:30 UTC = 2026-09-28 04:30 in Shanghai = 2026-09-27 13:30 in Los Angeles
FIXED = dt.datetime(2026, 9, 27, 20, 30, tzinfo=dt.timezone.utc).timestamp()


def _value(formula: str):
    return Book({"S": {(0, 0): ("f", formula)}}).value("S", 0, 0)


def test_today_and_now_follow_the_moment_they_are_given():
    shanghai = clock.moment(FIXED, "Asia/Shanghai")
    with as_of(shanghai):
        assert serial_to_datetime(_value("TODAY()")).date() == dt.date(2026, 9, 28)
        assert serial_to_datetime(_value("NOW()")).replace(second=0, microsecond=0) == dt.datetime(2026, 9, 28, 4, 30)
        with as_of(clock.moment(FIXED, "America/Los_Angeles")):          # nested, then restored
            assert serial_to_datetime(_value("TODAY()")).date() == dt.date(2026, 9, 27)
        assert serial_to_datetime(_value("TODAY()")).date() == dt.date(2026, 9, 28)
    # outside any moment: the process clock, as before
    assert serial_to_datetime(_value("TODAY()")).date() == dt.date.today()


def test_timezones_are_resolved_strictly():
    assert clock.resolve_timezone("Asia/Shanghai", None, None) == "Asia/Shanghai"
    assert clock.resolve_timezone(None, None, None) == "UTC"
    for bad in ("Mars/Olympus", "../etc/passwd", "/UTC", "x" * 80, 5):
        with pytest.raises(Exception) as e:
            clock.resolve_timezone(bad, None, None)
        assert getattr(e.value, "code", None) == "invalid_field", bad
    assert clock.moment_of({"as_of": {"at": FIXED, "timezone": "Asia/Shanghai"}}) == dt.datetime(2026, 9, 28, 4, 30, tzinfo=ZoneInfo("Asia/Shanghai"))
    assert clock.moment_of({}) is None, "a job from before CI-95 keeps the process clock"


def _dated_template(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "名称", "占位", "=TODAY()", '=TEXT(TODAY(),"yyyy-mm-dd")'
    p = tmp_path / "日期模板.xlsx"
    wb.save(p)
    return p


def _project(api, tmp_path):
    pid = api.project("CI-95 日期")
    src = openpyxl.Workbook()
    src.active.title = "名单"
    src.active.append(["名称"])
    src.active.append(["测试甲"])
    buf = io.BytesIO()
    src.save(buf)
    api.upload_bytes(pid, "名单.xlsx", buf.getvalue())
    t = api.template(pid, _dated_template(tmp_path))
    s = api.slot(t["id"], kind="cell", sheet="表", cell="B1", name="名称")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "名单", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "名单", "header": "名称"}], "target": {"binding": s["binding"]}}])
    return pid


def _dates(api, jid):
    out = next(iter(api.outputs(jid).values()))
    wv = openpyxl.load_workbook(io.BytesIO(out), data_only=True)["表"]
    return wv["C1"].value, wv["D1"].value


def test_a_job_computes_today_in_the_submitters_zone_at_submission(api, tmp_path, monkeypatch):
    monkeypatch.setattr(clock, "current_instant", lambda: FIXED)
    pid = _project(api, tmp_path)
    shanghai = api.job(pid, timezone="Asia/Shanghai")
    assert shanghai["status"] == "completed", shanghai
    day, text = _dates(api, shanghai["id"])
    as_date = day.date() if hasattr(day, "date") else serial_to_datetime(day).date()   # an unformatted cell keeps the serial
    assert as_date == dt.date(2026, 9, 28) and text == "2026-09-28"
    la = api.job(pid, timezone="America/Los_Angeles")
    assert _dates(api, la["id"])[1] == "2026-09-27"
    # no zone and no profile: UTC, stated rather than guessed
    plain = api.job(pid)
    assert _dates(api, plain["id"])[1] == "2026-09-27"
    # the instant is part of the job: later runs of the same job compute with it, whatever the clock says
    snap = api.c.app_state.jobs.store.one("SELECT snapshot FROM jobs WHERE id=?", (shanghai["id"],))["snapshot"]
    assert snap["as_of"] == {"at": FIXED, "timezone": "Asia/Shanghai"}
    assert api.ok(api.c.get(f"/api/jobs/{shanghai['id']}"))["as_of"] == {"at": FIXED, "timezone": "Asia/Shanghai"}, "the job says whose today it used"
    r = api.c.post(f"/api/projects/{pid}/jobs", json={"timezone": "Mars/Olympus"})
    assert r.status_code == 400 and r.json()["error"] == "invalid_field" and r.json()["detail"]["field"] == "timezone"


def test_a_preview_uses_the_viewers_zone(api, tmp_path, monkeypatch):
    monkeypatch.setattr(clock, "current_instant", lambda: FIXED)
    pid = _project(api, tmp_path)
    out = api.ok(api.c.get(f"/api/projects/{pid}/preview", params={"tz": "Asia/Shanghai"}))["output"]
    data = api.ok_bytes(api.c.get(f"/api/preview-files/{out['blob']}", params={"fmt": "xlsx"}))
    assert openpyxl.load_workbook(io.BytesIO(data), data_only=True)["表"]["D1"].value == "2026-09-28"
    out_la = api.ok(api.c.get(f"/api/projects/{pid}/preview", params={"tz": "America/Los_Angeles"}))["output"]
    assert out_la["blob"] != out["blob"], "two zones on either side of midnight do not share a cached preview"
    data = api.ok_bytes(api.c.get(f"/api/preview-files/{out_la['blob']}", params={"fmt": "xlsx"}))
    assert openpyxl.load_workbook(io.BytesIO(data), data_only=True)["表"]["D1"].value == "2026-09-27"
