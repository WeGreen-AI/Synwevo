"""Acceptance scenarios 9–14."""
import io
import threading
import time
import zipfile

import docx
import openpyxl
import pymupdf
import pytest
from docx.shared import Pt

from .helpers import Api, docx_text
from synwevo_docs.engine.converters import convert


# ------------------------------------------------------------------ 9. template preservation
def test_s9_template_preservation(api, fixtures, tmp_path):
    d = docx.Document(fixtures / "测试_合同模板.docx")
    p = d.add_paragraph()
    r = p.add_run("特殊字体段落：")
    r.font.name = "不存在的字体ABC"
    p.add_run("占位")
    b = io.BytesIO()
    d.save(b)
    tpl_path = tmp_path / "模板_含特殊字体.docx"
    tpl_path.write_bytes(b.getvalue())
    pid = api.project("S9")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, tpl_path)
    assert any("不存在的字体ABC" in n for n in t["info"]["notes"])               # missing font reported
    s_name = api.text_slot(t["id"], "示例客户名称", "客户名称")
    s_long = api.text_slot(t["id"], "占位", "长文本")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 1}})
    long_text = "中文长文本测试，" * 300 + "结尾ABC"
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s_name["binding"]}},
                   {"type": "single", "sources": [{"kind": "constant", "value": long_text}], "target": {"binding": s_long["binding"]}}])
    job = api.job(pid, formats=["docx", "pdf"])
    assert job["status"] == "completed", api.items(job["id"])
    outs = api.outputs(job["id"])
    out = next(v for k, v in outs.items() if k.endswith(".docx"))
    orig = zipfile.ZipFile(io.BytesIO(tpl_path.read_bytes()))
    res = zipfile.ZipFile(io.BytesIO(out))
    changed = [n for n in orig.namelist() if n in res.namelist() and orig.read(n) != res.read(n)]
    assert changed == ["word/document.xml"], changed          # every other part (styles, header, footer, media...) byte-identical
    od, rd = docx.Document(io.BytesIO(tpl_path.read_bytes())), docx.Document(io.BytesIO(out))
    untouched = [i for i, para in enumerate(od.paragraphs) if "示例客户名称" not in para.text and "占位" not in para.text]
    assert all(od.paragraphs[i].text == rd.paragraphs[i].text for i in untouched)
    assert [[c.text for c in r.cells] for r in od.tables[0].rows] == [[c.text for c in r.cells] for r in rd.tables[0].rows]
    assert len(rd.inline_shapes) == 1
    assert long_text in rd.paragraphs[-1].text and rd.paragraphs[-1].runs[0].font.name == "不存在的字体ABC"   # style of field start kept
    pdf = pymupdf.open(stream=next(v for k, v in outs.items() if k.endswith(".pdf")), filetype="pdf")
    txt = "".join(p.get_text() for p in pdf).replace("\n", "")
    assert "结尾ABC" in txt and pdf.page_count >= 2                              # long text flows onto following pages, not cut


# ------------------------------------------------------------------ 10. PDF & OCR
def _pdf_template_slots(api, t):
    slots = []
    doc = pymupdf.open(stream=api.c.get(f"/api/templates/{t['id']}/download?which=original").content, filetype="pdf")
    for page in (0, 1):
        pg = doc[page]
        wr = [x for x in pg.get_text("words") if x[4] == "旧姓名样例"][0]
        r = pymupdf.Rect(wr[:4]) + (-2, -2, 60, 2)
        vr = r * pg.rotation_matrix                   # what the user sees and draws on the rotated page
        vr.normalize()
        nr = [vr.x0 / pg.rect.width, vr.y0 / pg.rect.height, vr.x1 / pg.rect.width, vr.y1 / pg.rect.height]
        slots.append(api.slot(t["id"], kind="pdf_box", page=page, nrect=nr, name=f"收件人{page + 1}"))
    return slots


def test_s10_text_pdf_source_into_pdf_template(api, fixtures):
    pid = api.project("S10")
    api.upload(pid, fixtures / "测试_申请表.pdf")
    src = api.ok(api.c.get(f"/api/projects/{pid}/sources"))[0]
    assert src["status"] == "ready"
    # select the value on page 2 from the text layer (normalised visual coords → stored in PDF space)
    words = api.ok(api.c.get(f"/api/sources/{src['id']}/words/1"))["words"]
    w = next(x for x in words if x["t"].startswith("TEST"))
    nrect = [w["r"][0] - 0.01, w["r"][1] - 0.005, w["r"][2] + 0.05, w["r"][3] + 0.005]
    prop = api.ok(api.c.post(f"/api/sources/{src['id']}/locator", json={"project_id": pid, "selection": {"kind": "pdf_region", "page": 1, "nrect": nrect}}))["proposals"][0]
    assert prop["sample"] == "TEST00000002" and prop["locator"]["anchor"]["text"] == "证件号："
    t = api.template(pid, fixtures / "测试_通知书模板.pdf")
    view = api.ok(api.c.get(f"/api/templates/{t['id']}/view"))
    assert [p["rotation"] for p in view["pdf"]["pages"]] == [0, 90]
    slots = _pdf_template_slots(api, t)
    assert all(s["original_text"] == "旧姓名样例" for s in slots)
    api.rule(pid, {"mode": "per_file"})
    api.maps(pid, [{"type": "single", "sources": [prop["locator"]], "target": {"binding": slots[0]["binding"]}},
                   {"type": "single", "sources": [prop["locator"]], "target": {"binding": slots[1]["binding"]}}])
    job = api.job(pid)
    assert job["status"] == "completed"
    pdf = pymupdf.open(stream=next(iter(api.outputs(job["id"]).values())), filetype="pdf")
    all_text = "".join(p.get_text() for p in pdf)
    assert "旧姓名样例" not in all_text                                      # old values really removed on both pages
    assert "TEST00000002" in pdf[0].get_text() and "TEST00000002" in pdf[1].get_text().replace("\n", "")
    # the value on the rotated page is written in the direction of the text it replaces
    for pg in pdf:
        hits = [x for x in pg.get_text("words") if "TEST" in x[4]]
        assert hits and all(len(x[4]) > 4 for x in hits)                     # not broken into single characters


def test_s10_scanned_image_ocr_and_correction(api, fixtures):
    pid = api.project("S10-OCR")
    api.upload(pid, fixtures / "测试_扫描单据.png")
    img = api.ok(api.c.get(f"/api/projects/{pid}/sources"))[0]
    assert img["status"] == "needs_ocr"
    ocr = api.ok(api.c.post(f"/api/sources/{img['id']}/ocr", json={}))       # local Tesseract (chi_sim + eng)
    assert ocr["words"] > 5
    iw = api.ok(api.c.get(f"/api/sources/{img['id']}/words/0"))
    assert iw["source"] == "ocr" and all("conf" in x for x in iw["words"])
    code = next(x for x in iw["words"] if "SC-" in x["t"] or "2026" in x["t"])
    loc = api.ok(api.c.post(f"/api/sources/{img['id']}/locator", json={"project_id": pid, "selection": {"kind": "pdf_region", "page": 0,
                "nrect": [code["r"][0] - 0.01, code["r"][1] - 0.01, code["r"][2] + 0.01, code["r"][3] + 0.01]}}))["proposals"][0]["locator"]
    assert loc["kind"] == "ocr_region"
    t = api.template(pid, fixtures / "测试_通知书模板.pdf")
    slots = _pdf_template_slots(api, t)
    api.rule(pid, {"mode": "per_file"})
    api.maps(pid, [{"type": "single", "sources": [loc], "target": {"binding": slots[1]["binding"]}}])
    mid = api.ok(api.c.get(f"/api/projects/{pid}/mappings"))["mappings"][0]["id"]
    # force the confirmation path: raise the threshold so the recognised value needs confirmation
    m = api.ok(api.c.get(f"/api/projects/{pid}/mappings"))["mappings"]
    m[0]["sources"][0]["min_conf"] = 101
    api.maps(pid, m)
    pv = api.ok(api.c.get(f"/api/projects/{pid}/preview"))
    assert pv.get("blocked") and any(i["code"] == "ocr_low_confidence" for i in pv["issues"])   # low confidence → pending confirmation
    api.ok(api.c.put(f"/api/sources/{img['id']}/corrections/{mid}", json={"value": "SC-20260922", "confirmed": True}))
    job = api.job(pid)
    assert job["status"] == "completed"
    pdf2 = pymupdf.open(stream=next(iter(api.outputs(job["id"]).values())), filetype="pdf")
    assert "SC-20260922" in pdf2[1].get_text().replace("\n", "") and "旧姓名样例" not in pdf2[1].get_text()


def test_s10b_pdf_form_template(api, fixtures):
    pid = api.project("S10b")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, fixtures / "测试_表单模板.pdf")
    assert t["info"]["pdf"]["fields"][0]["name"] == "applicant"
    s = api.slot(t["id"], kind="pdf_field", field="applicant", name="申请人")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 2}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
    job = api.job(pid)
    pdf = pymupdf.open(stream=next(iter(api.outputs(job["id"]).values())), filetype="pdf")
    assert "测试钱磊敏" in pdf[0].get_text() and not list(pdf[0].widgets() or [])      # filled and flattened


# ------------------------------------------------------------------ 11. save & reuse
def test_s11_persistence_reuse_and_versions(tmp_path, fixtures):
    from fastapi.testclient import TestClient

    from synwevo_docs.api.app import create_app
    data_dir = str(tmp_path / "data")
    app1 = create_app(data_dir)
    with TestClient(app1) as c1:
        a = Api(c1)
        pid = a.project("S11")
        a.upload(pid, fixtures / "测试_客户100.xlsx")
        t = a.template(pid, fixtures / "测试_合同模板.docx")
        s = a.text_slot(t["id"], "示例客户名称", "客户名称")
        a.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 2}})
        a.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
        tid, head = t["id"], a.ok(c1.get(f"/api/templates/{t['id']}"))["head"]
    app1.state.jobs.stop()
    # "restart": a new server process on the same data directory
    app2 = create_app(data_dir)
    with TestClient(app2) as c2:
        a = Api(c2)
        p = a.ok(c2.get(f"/api/projects/{pid}"))
        assert p["template"]["id"] == tid and p["template"]["version"]["no"] == head
        assert p["record_rule"]["mode"] == "per_row" and len(p["mappings"]["mappings"]) == 1
        # replace the source with a new batch of the same structure
        old = p["sources"][0]["id"]
        a.ok(c2.delete(f"/api/sources/{old}"))
        wb = openpyxl.load_workbook(fixtures / "测试_客户100.xlsx")
        wb["客户"]["B2"] = "新批次客户"
        b = io.BytesIO()
        wb.save(b)
        a.upload_bytes(pid, "下一批.xlsx", b.getvalue())
        job = a.job(pid)
        assert "新批次客户" in docx_text(a.outputs(job["id"])["001.docx"])
    app2.state.jobs.stop()


def test_s11b_job_snapshot_isolated_from_later_edits(tmp_path, fixtures, monkeypatch):
    """Template / mappings edited after the job was created must not leak into that job."""
    from fastapi.testclient import TestClient

    from synwevo_docs import jobs as J
    from synwevo_docs.api.app import create_app
    gate = threading.Event()
    orig_run = J.JobService._run

    def gated(self, jid):
        gate.wait(20)
        return orig_run(self, jid)
    monkeypatch.setattr(J.JobService, "_run", gated)
    app = create_app(str(tmp_path / "d"))
    with TestClient(app) as c:
        a = Api(c)
        pid = a.project("S11b")
        a.upload(pid, fixtures / "测试_客户100.xlsx")
        t = a.template(pid, fixtures / "测试_合同模板.docx")
        s = a.text_slot(t["id"], "示例客户名称", "客户名称")
        a.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 2}})
        a.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
        j = a.ok(c.post(f"/api/projects/{pid}/jobs", json={}))
        v_at_start = a.ok(c.get(f"/api/templates/{t['id']}"))["head"]
        # edits after job creation
        a.ok(c.delete(f"/api/templates/{t['id']}/slots/{s['id']}", params={"project_id": pid}))
        a.text_slot(t["id"], "0.00", "金额", ctx="合同金额")
        gate.set()
        for _ in range(100):
            j = a.ok(c.get(f"/api/jobs/{j['id']}"))
            if j["status"] not in ("queued", "running"):
                break
            time.sleep(0.1)
        assert j["status"] == "completed" and j["template"]["version_no"] == v_at_start
        text = docx_text(a.outputs(j["id"])["001.docx"])
        assert "测试钱磊敏" in text and "人民币 0.00 元" in text                  # old template version + old mappings used
    app.state.jobs.stop()


# ------------------------------------------------------------------ 12. cancel / partial failure / retry / names
def _xlsx_template_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"] = "备注"
    ws["B1"] = "（待填）"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_s12_partial_failure_retry_and_names(api, fixtures, tmp_path, monkeypatch):
    pid = api.project("S12")
    # CSV has no cell-length limit, so the 40 000-character value really exceeds Excel's 32 767 limit
    csv_text = "编号,备注\nA,正常\nA,正常2\nB," + "超长" * 20000 + "\nC,正常3\n"
    api.upload_bytes(pid, "数据.csv", csv_text.encode("utf-8"))
    tp = tmp_path / "表格模板.xlsx"
    tp.write_bytes(_xlsx_template_bytes())
    t = api.template(pid, tp)
    s = api.slot(t["id"], kind="cell", sheet="表", cell="B1", name="备注")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "CSV", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "CSV", "header": "备注"}], "target": {"binding": s["binding"]}}])
    # transient failure on first attempt of record 4 → retry fixes it
    from synwevo_docs.engine import generate as G
    real = G.fill_native
    calls = {"n": 0}

    def flaky(fmt, work, slots, plan, options=None):
        if plan.record["index"] == 3 and calls["n"] == 0:
            calls["n"] += 1
            raise RuntimeError("模拟的临时故障")
        return real(fmt, work, slots, plan, options)
    monkeypatch.setattr(G, "fill_native", flaky)
    job = api.job(pid, naming="{编号}")
    items = api.items(job["id"])
    st = [i["status"] for i in items]
    assert st == ["success", "success", "failed", "failed"], st           # 40k-char cell > Excel limit, and the transient error
    assert job["status"] == "completed_with_errors"
    names = sorted(a["filename"] for i in items for a in i["artifacts"])
    assert names == ["A (2).xlsx", "A.xlsx"]                             # same name never overwrites
    report = api.c.get(f"/api/jobs/{job['id']}/report.csv").content.decode("utf-8-sig")
    assert "32767" in report and "模拟的临时故障" in report
    first_blobs = {a["id"] for i in items for a in i["artifacts"]}
    api.ok(api.c.post(f"/api/jobs/{job['id']}/retry", json={}))
    for _ in range(100):
        j = api.ok(api.c.get(f"/api/jobs/{job['id']}"))
        if j["status"] not in ("queued", "running"):
            break
        time.sleep(0.1)
    items = api.items(job["id"])
    assert [i["status"] for i in items] == ["success", "success", "failed", "success"]
    assert items[3]["attempts"] == 2 and items[2]["attempts"] == 2
    assert {a["id"] for i in items[:2] for a in i["artifacts"]} == first_blobs   # earlier results untouched
    c = j["counts"]
    assert c["success"] == 3 and c["failed"] == 1 and c["files"] == 3 == sum(len(i["artifacts"]) for i in items)


def test_s12b_cancel(api, fixtures):
    pid = api.project("S12b")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    s = api.text_slot(t["id"], "示例客户名称", "客户名称")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
    j = api.ok(api.c.post(f"/api/projects/{pid}/jobs", json={"formats": ["docx", "pdf"]}))   # PDF conversion makes it slow enough
    time.sleep(1.5)
    api.ok(api.c.post(f"/api/jobs/{j['id']}/cancel"))
    for _ in range(300):
        j = api.ok(api.c.get(f"/api/jobs/{j['id']}"))
        if j["status"] not in ("queued", "running", "cancelling"):    # 「取消中」= 已接受请求，尚未到终态
            break
        time.sleep(0.1)
    assert j["status"] == "cancelled"
    items = api.items(j["id"])
    c = j["counts"]
    assert 0 < c.get("success", 0) < 100 and c.get("success", 0) + c.get("cancelled", 0) == 100
    assert c["files"] == 2 * c["success"] == sum(len(i["artifacts"]) for i in items)
    assert all(not i["artifacts"] for i in items if i["status"] == "cancelled")


def test_s12c_restart_recovery(tmp_path, fixtures):
    """A job interrupted by a server restart resumes from the persisted state."""
    from fastapi.testclient import TestClient

    from synwevo_docs.api.app import create_app
    from synwevo_core.store import Store
    data_dir = str(tmp_path / "data")
    app = create_app(data_dir)
    with TestClient(app) as c:
        a = Api(c)
        pid = a.project("S12c")
        a.upload(pid, fixtures / "测试_客户100.xlsx")
        t = a.template(pid, fixtures / "测试_合同模板.docx")
        s = a.text_slot(t["id"], "示例客户名称", "客户名称")
        a.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 10}})
        a.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
        app.state.jobs.stop()
        time.sleep(1.2)
        j = a.ok(c.post(f"/api/projects/{pid}/jobs", json={}))
    # simulate a crash in the middle: first item marked running, job running
    st = Store(data_dir)
    st.execute("UPDATE jobs SET status='running' WHERE id=?", (j["id"],))
    first = st.one("SELECT id FROM job_items WHERE job_id=? ORDER BY seq LIMIT 1", (j["id"],))
    st.execute("UPDATE job_items SET status='running' WHERE id=?", (first["id"],))
    app2 = create_app(data_dir)
    with TestClient(app2) as c2:
        a = Api(c2)
        for _ in range(200):
            j2 = a.ok(c2.get(f"/api/jobs/{j['id']}"))
            if j2["status"] not in ("queued", "running"):
                break
            time.sleep(0.1)
        assert j2["status"] == "completed" and j2["counts"]["success"] == 10 and j2["counts"]["files"] == 10
    app2.state.jobs.stop()


# ------------------------------------------------------------------ 13. native validation with a real office suite
@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pptx"])
def test_s13_outputs_open_in_libreoffice(api, fixtures, fmt, tmp_path):
    pid = api.project("S13")
    api.upload(pid, fixtures / "测试_订单明细100.xlsx")
    tpl = {"docx": "测试_订单模板.docx", "xlsx": "测试_订单模板.xlsx", "pptx": "测试_汇报模板.pptx"}[fmt]
    t = api.template(pid, fixtures / tpl)
    if fmt == "docx":
        v = api.view(t["id"])
        tbl = [b for b in v["doc"]["parts"][0]["blocks"] if b["t"] == "tbl"][0]
        rep = api.slot(t["id"], kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="明细")
        cols = ["行号", "品名", "数量", "单价"]
    elif fmt == "xlsx":
        rep = api.slot(t["id"], kind="repeat", sheet="订单", row=4, c1=0, c2=3, name="明细")
        cols = ["行号", "品名", "数量", "单价"]
    else:
        view = api.ok(api.c.get(f"/api/templates/{t['id']}/view"))
        table = next(sh for sh in view["slides"]["slides"][0]["shapes"] if sh["kind"] == "table")
        rep = api.slot(t["id"], kind="repeat", slide=0, shape=table["id"], row=1, name="明细")
        cols = ["品名", "数量", "单价"]
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0}, "group": {"keys": [{"header": "订单号"}]}})
    api.maps(pid, [{"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": rep["binding"]},
                    "columns": [{"slot_col": f"c{i}", "source": {"header": h}} for i, h in enumerate(cols)]}])
    job = api.job(pid, naming="{样本标识}")
    assert job["status"] == "completed"
    data = api.outputs(job["id"])["DD-001." + fmt]
    pdf = convert(data, fmt, "pdf")                     # LibreOffice opens and renders the file
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc.page_count >= 1 and "测试商品" in "".join(p.get_text() for p in doc)
    zipfile.ZipFile(io.BytesIO(data)).testzip()


# ------------------------------------------------------------------ 14. capability boundary
def test_s14_capability_boundary(api, fixtures):
    pid = api.project("S14")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    exports = {e["format"]: e for e in t["exports"]}
    assert exports["docx"]["available"] and not exports["xlsx"]["available"] and not exports["csv"]["available"]
    assert "表格型模板" in exports["xlsx"]["reason"]                              # explicit alternative path
    s = api.text_slot(t["id"], "示例客户名称", "客户名称")
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": s["binding"]}}])
    r = api.c.post(f"/api/projects/{pid}/jobs", json={"formats": ["xlsx"]})
    assert r.status_code == 400 and r.json()["error"] == "format_unavailable"
    assert not api.ok(api.c.get(f"/api/projects/{pid}/jobs"))                   # nothing shown as generated
    pdf_t = api.template(pid, fixtures / "测试_通知书模板.pdf")
    assert not {e["format"]: e for e in pdf_t["exports"]}["docx"]["available"]
    x_t = api.template(pid, fixtures / "测试_区域模板.xlsx")
    assert {e["format"]: e for e in x_t["exports"]}["csv"]["available"]
    cap = api.ok(api.c.get("/api/capabilities"))
    csv_row = next(r for r in cap["rows"] if r["format"].startswith("CSV"))
    assert csv_row["template"][0] == "no"
