"""Synthetic real-engine sources; machine gate, browser parity and live authorization.

These tests use the real Docs factory and its integrated machine router. All
files and identities below are test fixtures, built in the private pytest data dir.
No identity-service routes or mocked principals/parsers are used.
"""
import hashlib
import io
import json
import time

import docx
import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient

from synwevo_core.machine import docs_resource
from synwevo_core.oidc import Clients, Grants
from synwevo_docs.api.app import create_app
from synwevo_docs.api.v1 import openapi_subset
from synwevo_docs.api import v1_sources as sources_api


def ok(response):
    assert response.status_code < 400, (response.status_code, response.text[:500])
    return response.json()


def error(response, status, code=None):
    assert response.status_code == status, response.text[:500]
    body = response.json()["error"]
    assert body["code"] and body["message"] and body["request_id"]
    if code:
        assert body["code"] == code


def synthetic_files():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in (["Name", "Amount"], ["Alpha One", 10], ["Alpha Two", 20], ["Alpha Three", 30]):
        ws.append(row)
    x = io.BytesIO()
    wb.save(x)
    wb.close()
    d = docx.Document()
    d.add_paragraph("Record: Alpha One")
    d.add_paragraph("Amount: 10")
    d.add_paragraph("Record: Alpha Two")
    d.add_paragraph("Amount: 20")
    word = io.BytesIO()
    d.save(word)
    with pymupdf.open() as pdf:
        for text in ("Alpha One Beta Gamma", "Alpha Two Delta Epsilon"):
            page = pdf.new_page(width=300, height=200)
            page.insert_text((30, 50), text)
        pdf_bytes = pdf.tobytes()
    return {"xlsx": x.getvalue(), "docx": word.getvalue(), "pdf": pdf_bytes,
            "csv": b'Name;Amount\n"Alpha One";10\n"Alpha Two";20\n'}


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("SYNWEVO_DOCS_UI_DIR", str(tmp_path / "no-ui"))
    monkeypatch.setenv("SYNWEVO_API_RATE_PER_MIN", "10000")
    api = create_app(str(tmp_path / "synthetic-source-data"), start_jobs=False)
    app = api.state.app

    headers = {}
    for name, tenant, scopes in (("rw", "default", ("read", "write")),
                                 ("ro", "default", ("read",)),
                                 ("other", "source-other", ("read", "write"))):
        token, _ = api.state.auth.create_key(name, tenant, scopes)
        headers[name] = {"Authorization": f"Bearer {token}"}
    with TestClient(api) as client:
        pid = ok(client.post("/api/v1/projects", headers=headers["rw"], json={"name": "Synthetic source operations"}))["id"]
        files = synthetic_files()
        src = {fmt: upload(client, headers["rw"], pid, fmt, data) for fmt, data in files.items()}
        mappings = ok(client.put(f"/api/v1/projects/{pid}/mappings", headers=headers["rw"], json={
            "mappings": [{"sources": [{"kind": "cell", "sheet": "Data", "cell": "A2"}],
                          "target": {"binding": "synthetic-target"}}]}))
        yield {"client": client, "api": api, "app": app, "h": headers, "pid": pid,
               "src": src, "files": files, "mid": mappings["mappings"][0]["id"]}
    app.parse_pool.shutdown(wait=True)
    api.state.webhooks.stop()


def upload(client, headers, pid, fmt, data):
    src = ok(client.post(f"/api/v1/projects/{pid}/files", files={"file": (f"synthetic.{fmt}", data)}, headers=headers))
    for _ in range(250):
        src = ok(client.get(f"/api/v1/files/{src['id']}", headers=headers))
        if src["status"] != "parsing":
            break
        time.sleep(.02)
    assert src["status"] == "ready", src
    return src


@pytest.mark.parametrize("fmt,view", [("xlsx", "grid"), ("docx", "docx"), ("pdf", "pdf")])
def test_real_structure_view_and_original_content_match_browser(env, fmt, view):
    c, h = env["client"], env["h"]["rw"]
    sid = env["src"][fmt]["id"]
    base = f"/api/v1/files/{sid}"
    structure = ok(c.get(base + "/structure", headers=h))
    assert structure["structure"] == env["src"][fmt]["structure"]
    machine = ok(c.get(base + "/view", headers=h))
    browser = ok(c.get(f"/api/sources/{sid}/view"))
    assert machine == browser and machine["view"] == view
    original = c.get(base + "/content", headers=h)
    assert original.status_code == 200 and original.content == env["files"][fmt]
    assert original.content == c.get(f"/api/sources/{sid}/download").content
    assert original.headers["X-Content-SHA256"] == hashlib.sha256(original.content).hexdigest()
    assert original.headers["Cache-Control"] == "no-store"
    assert "attachment" in original.headers["Content-Disposition"]


def test_grid_and_search_pages_have_no_skipped_or_repeated_items(env):
    c, h = env["client"], env["h"]["rw"]
    sid = env["src"]["xlsx"]["id"]
    base = f"/api/v1/files/{sid}"
    grid = ok(c.get(base + "/view?sheet=Data&offset=2&limit=1", headers=h))["grid"]
    assert grid["offset"] == 2 and grid["rows"] == [[{"v": "Alpha Two"}, {"v": "20", "t": "int"}]]
    got = []
    for offset in range(3):
        response = c.get(base + f"/search?q=Alpha&limit=1&offset={offset}", headers=h)
        out = ok(response)
        got.extend(out["items"])
        assert out["total"] == 3 and out["limit"] == 1 and out["offset"] == offset
        assert out["next_offset"] == (offset + 1 if offset < 2 else None)
        if offset < 2:
            assert "rel=\"next\"" in response.headers["Link"]
    assert got == ok(c.get(f"/api/sources/{sid}/search?q=Alpha"))["hits"]
    assert ok(c.get(base + "/search?q=Alpha&offset=10", headers=h))["items"] == []


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_document_and_pdf_search_use_existing_positions(env, fmt):
    c, h = env["client"], env["h"]["rw"]
    sid = env["src"][fmt]["id"]
    expected = ok(c.get(f"/api/sources/{sid}/search?q=Alpha"))
    out = ok(c.get(f"/api/v1/files/{sid}/search?q=Alpha&offset=1&limit=1", headers=h))
    assert out["items"] == expected["hits"][1:2] and out["total"] == expected["total"]


def test_pdf_words_are_paginated_with_real_geometry(env):
    c, h = env["client"], env["h"]["rw"]
    sid = env["src"]["pdf"]["id"]
    browser = ok(c.get(f"/api/sources/{sid}/words/1"))
    words = []
    for offset in range(len(browser["words"])):
        out = ok(c.get(f"/api/v1/files/{sid}/words/1?offset={offset}&limit=1", headers=h))
        assert out["source"] == "text_layer" and out["index"] == 1 and out["total"] == len(browser["words"])
        words.extend(out["items"])
    assert words == browser["words"]
    assert [w["t"] for w in words] == ["Alpha", "Two", "Delta", "Epsilon"]
    assert all(len(w["r"]) == 4 and 0 <= w["r"][0] < w["r"][2] <= 1 for w in words)
    error(c.get(f"/api/v1/files/{sid}/words/2", headers=h), 404, "page_missing")


def test_page_image_matches_browser_and_has_verified_content_hash(env):
    c, h, sid = env["client"], env["h"]["rw"], env["src"]["pdf"]["id"]
    response = c.get(f"/api/v1/files/{sid}/page/1.png?zoom=2", headers=h)
    assert response.status_code == 200 and response.content.startswith(b"\x89PNG\r\n\x1a\n")
    assert response.content == c.get(f"/api/sources/{sid}/page/1.png?zoom=2").content
    assert response.headers["X-Content-SHA256"] == hashlib.sha256(response.content).hexdigest()
    assert response.headers["Cache-Control"] == "no-store"
    error(c.get(f"/api/v1/files/{sid}/page/2.png", headers=h), 404, "page_missing")
    for query in ("zoom=0", "zoom=5", "zoom=nan", "zoom=inf", "zoom=1&zoom=2", "other=1"):
        error(c.get(f"/api/v1/files/{sid}/page/0.png?{query}", headers=h), 400)


@pytest.mark.parametrize("change,status,code", [
    ("revoke", 401, "key_revoked"), ("delete", 404, "not_found"),
    ("delete_project", 404, "not_found"), ("version", 409, "source_changed"),
])
def test_page_export_checks_authority_again_after_rendering(env, monkeypatch, change, status, code):
    c, api, h, sid = env["client"], env["api"], env["h"]["rw"], env["src"]["pdf"]["id"]
    render = api.state.views.source_page_png
    key_id = ok(c.get("/api/v1/whoami", headers=h))["key_id"]

    def changed(*args):
        data = render(*args)  # Real render, then an isolated concurrent-lifecycle injection.
        if change == "revoke":
            api.state.auth.revoke_key(key_id)
        elif change == "delete":
            env["app"].remove_source(sid)
        elif change == "delete_project":
            env["app"].delete_project(env["pid"])
        else:
            env["app"].store.execute("UPDATE source_files SET version=version+1 WHERE id=?", (sid,))
        return data

    monkeypatch.setattr(api.state.views, "source_page_png", changed)
    error(c.get(f"/api/v1/files/{sid}/page/0.png", headers=h), status, code)


@pytest.mark.parametrize("fmt,selection", [
    ("xlsx", {"kind": "cells", "sheet": "Data", "r1": 1, "c1": 0, "r2": 1, "c2": 0}),
    ("docx", {"kind": "docx_text", "path": "0", "start": 8, "end": 17}),
    ("pdf", {"kind": "pdf_region", "page": 0, "nrect": [0.0, 0.0, 1.0, 1.0]}),
])
def test_machine_locator_uses_same_engine_as_browser(env, fmt, selection):
    c, h, pid, sid = env["client"], env["h"]["rw"], env["pid"], env["src"][fmt]["id"]
    body = {"selection": selection, "project_id": pid}
    out = ok(c.post(f"/api/v1/files/{sid}/locator", headers=h, json=body))
    assert out == ok(c.post(f"/api/sources/{sid}/locator", json=body))
    assert out["proposals"] and "Alpha One" in str(out["proposals"])
    assert ok(c.post(f"/api/v1/files/{sid}/locator", headers=env["h"]["ro"], json=body)) == out
    if fmt == "xlsx":
        proposed = out["proposals"][0]["locator"]
        saved = ok(c.put(f"/api/v1/projects/{pid}/mappings", headers=h, json={"mappings": [
            {"sources": [proposed], "target": {"binding": "synthetic-target"}}]}))
        located = ok(c.get(f"/api/v1/projects/{pid}/locate?file={sid}", headers=h))
        assert located[saved["mappings"][0]["id"]][0]["value"] == "Alpha One"


@pytest.mark.parametrize("fmt,selection", [
    ("xlsx", {"kind": "cells", "sheet": "Data", "r1": 1, "c1": 0, "r2": 999, "c2": 0}),
    ("xlsx", {"kind": "cells", "sheet": "Data", "r1": 1, "c1": 0, "r2": 1, "c2": 99}),
    ("docx", {"kind": "docx_text", "path": "0", "start": 8, "end": 999}),
    ("docx", {"kind": "docx_text", "path": "999", "start": 0, "end": 1}),
    ("pdf", {"kind": "pdf_region", "page": 2, "nrect": [0.0, 0.0, 1.0, 1.0]}),
    ("xlsx", {"kind": "pdf_region", "page": 0, "nrect": [0.0, 0.0, 1.0, 1.0]}),
])
def test_out_of_bounds_selections_are_not_silently_clipped(env, fmt, selection):
    c, h, sid = env["client"], env["h"]["rw"], env["src"][fmt]["id"]
    body = {"selection": selection}
    error(c.post(f"/api/v1/files/{sid}/locator", headers=h, json=body), 400, "bad_selection")
    browser = c.post(f"/api/sources/{sid}/locator", json=body)
    assert browser.status_code == 400 and browser.json()["error"] == "bad_selection"
    assert browser.json()["message"] and browser.json()["request_id"]


def test_locator_rejects_wrong_project_and_strict_inputs(env):
    c, h, sid = env["client"], env["h"]["rw"], env["src"]["pdf"]["id"]
    selection = {"kind": "pdf_region", "page": 0, "nrect": [0, 0, 1, 1]}
    pid = ok(c.post("/api/v1/projects", headers=h, json={"name": "Another synthetic project"}))["id"]
    body = {"project_id": pid, "selection": selection}
    error(c.post(f"/api/v1/files/{sid}/locator", headers=h, json=body), 404, "not_found")
    browser = c.post(f"/api/sources/{sid}/locator", headers=h, json=body)
    assert browser.status_code == 404 and browser.json()["error"] == "not_found"
    for body in ({}, {"selection": {**selection, "page": True}}, {"selection": {**selection, "nrect": [0, 0, 0, 1]}},
                 {"selection": {**selection, "unknown": 1}}, {"selection": selection, "block": "0"}):
        error(c.post(f"/api/v1/files/{sid}/locator", headers=h, json=body), 400)


def test_docx_blocks_and_sample_records_share_the_current_rule(env):
    c, h, pid = env["client"], env["h"]["rw"], env["pid"]
    sid = env["src"]["docx"]["id"]
    ok(c.put(f"/api/v1/projects/{pid}/rule", headers=h, json={
        "rule": {"mode": "per_block", "block": {"anchor": "Record:", "match": "starts"}}}))
    expected = ok(c.get(f"/api/sources/{sid}/blocks"))["blocks"]
    assert len(expected) == 2
    for offset in range(2):
        out = ok(c.get(f"/api/v1/files/{sid}/blocks?offset={offset}&limit=1", headers=h))
        assert out["items"] == expected[offset:offset + 1] and out["total"] == 2
    machine = ok(c.get(f"/api/v1/files/{sid}/view?block=1", headers=h))
    assert machine == ok(c.get(f"/api/sources/{sid}/view?block=1"))
    records = ok(c.get(f"/api/v1/projects/{pid}/records?offset=1&limit=1", headers=h))
    browser = ok(c.get(f"/api/projects/{pid}/records?offset=1&limit=1"))
    assert records["items"] == browser["records"] and records["total"] == browser["total"]
    assert records["items"][0]["key"].startswith(sid)
    assert env["app"].store.one("SELECT COUNT(*) AS n FROM jobs")["n"] == 0


def test_options_reparse_only_real_csv_controls(env):
    c, h = env["client"], env["h"]["rw"]
    sid = env["src"]["csv"]["id"]
    url = f"/api/v1/files/{sid}/options"
    out = ok(c.put(url, headers=h, json={"csv": {"delimiter": ",", "encoding": "utf-8"}}))
    assert out["version"] == env["src"]["csv"]["version"] + 1 and out["structure"]["sheets"][0]["cols"] == 1
    body = {"csv": {"delimiter": ";", "encoding": "utf-8", "quotechar": '"'}}
    out = ok(c.put(url, headers=h, json=body))
    assert out["structure"]["sheets"][0]["cols"] == 2
    browser = ok(c.put(f"/api/sources/{sid}/options", json=body))
    assert out["structure"] == browser["structure"] and out["options"] == browser["options"]
    error(c.put(f"/api/v1/files/{env['src']['xlsx']['id']}/options", headers=h, json=body), 400, "unsupported_options")
    assert c.get(f"/api/v1/files/{sid}/content", headers=h).content == env["files"]["csv"]


def test_corrections_current_mapping_and_shared_store(env):
    c, h, mid = env["client"], env["h"]["rw"], env["mid"]
    sid = env["src"]["xlsx"]["id"]
    url = f"/api/v1/files/{sid}/corrections/{mid}"
    correction = {"value": "Corrected synthetic value", "confirmed": False}
    assert ok(c.put(url, headers=h, json=correction)) == {mid: correction}
    assert ok(c.get(f"/api/v1/files/{sid}", headers=h))["corrections"] == {mid: correction}
    assert ok(c.put(f"/api/sources/{sid}/corrections/{mid}", json={"value": None})) == {}
    assert ok(c.put(url, headers=h, json={"value": "Again"}))[mid]["confirmed"] is True
    assert ok(c.put(url, headers=h, json={"value": None})) == {}
    # Once removed from the current revision, a mapping id is no longer a writable target.
    ok(c.put(f"/api/v1/projects/{env['pid']}/mappings", headers=h, json={"mappings": []}))
    error(c.put(url, headers=h, json=correction), 404)
    assert ok(c.get(f"/api/v1/files/{sid}", headers=h))["corrections"] == {}


def test_representative_order_and_locate_are_shared_state(env):
    c, h, pid = env["client"], env["h"]["rw"], env["pid"]
    sid = env["src"]["xlsx"]["id"]
    selected = env["src"]["docx"]["id"]
    assert ok(c.post(f"/api/v1/files/{selected}/representative", headers=h))["representative"]
    ids = [s["id"] for s in reversed(list(env["src"].values()))]
    out = ok(c.put(f"/api/v1/projects/{pid}/files/order", headers=h, json={"ids": ids}))
    assert [s["id"] for s in out] == ids and sum(s["representative"] for s in out) == 1
    assert out == ok(c.get(f"/api/projects/{pid}/sources"))
    machine = ok(c.get(f"/api/v1/projects/{pid}/locate?file={sid}", headers=h))
    assert machine == ok(c.get(f"/api/projects/{pid}/locate?file={sid}"))
    assert machine[env["mid"]][0]["value"] == "Alpha One"


def source_calls(env):
    sid, mid = env["src"]["pdf"]["id"], env["mid"]
    base = f"/api/v1/files/{sid}"
    return [("GET", base + suffix, None) for suffix in (
        "/structure", "/view", "/blocks", "/search?q=Alpha", "/words/0", "/page/0.png", "/content")] + [
        ("POST", base + "/locator", {"selection": {"kind": "pdf_region", "page": 0, "nrect": [0, 0, 1, 1]}}),
        ("PUT", base + "/options", {"csv": {}}),
        ("PUT", base + f"/corrections/{mid}", {"value": "x"}),
        ("POST", base + "/representative", None)]


def test_missing_credential_wrong_tenant_and_read_only_scope(env):
    c = env["client"]
    for method, url, body in source_calls(env):
        kw = {"json": body} if body is not None else {}
        error(c.request(method, url, **kw), 401)
        error(c.request(method, url, headers=env["h"]["other"], **kw), 404)
        if url.endswith("/locator"):
            assert c.request(method, url, headers=env["h"]["ro"], **kw).status_code == 200
        elif method != "GET":
            error(c.request(method, url, headers=env["h"]["ro"], **kw), 403)


@pytest.mark.parametrize("remove_project", [False, True])
def test_removed_source_or_deleted_project_refuses_all_operations(env, remove_project):
    c, h = env["client"], env["h"]["rw"]
    sid = env["src"]["pdf"]["id"]
    # Warm every read path first, including content-addressed and parsed caches.
    for method, url, body in source_calls(env):
        if method == "GET":
            assert c.get(url, headers=h).status_code == 200
    if remove_project:
        ok(c.delete(f"/api/projects/{env['pid']}"))
    else:
        ok(c.delete(f"/api/v1/files/{sid}", headers=h))
    for method, url, body in source_calls(env):
        error(c.request(method, url, headers=h, **({"json": body} if body is not None else {})), 404)


def test_wrong_project_source_order_and_locate_never_partially_write(env):
    c, h, pid = env["client"], env["h"]["rw"], env["pid"]
    pid2 = ok(c.post("/api/v1/projects", headers=h, json={"name": "Other synthetic project"}))["id"]
    other = upload(c, h, pid2, "csv", b"Name\nOther\n")
    sid = env["src"]["xlsx"]["id"]
    before = ok(c.get(f"/api/projects/{pid}/sources"))
    for ids, status in (([sid, sid], 400), ([sid], 400), ([sid, other["id"]], 404), ([sid, "missing"], 404)):
        error(c.put(f"/api/v1/projects/{pid}/files/order", headers=h, json={"ids": ids}), status)
        assert ok(c.get(f"/api/projects/{pid}/sources")) == before
    error(c.get(f"/api/v1/projects/{pid}/locate?file={other['id']}", headers=h), 404)
    for path in (f"/api/v1/projects/{pid}/files/order", f"/api/v1/projects/{pid}/records",
                 f"/api/v1/projects/{pid}/locate?file={sid}"):
        response = c.put(path, headers=env["h"]["other"], json={"ids": []}) if path.endswith("order") else c.get(path, headers=env["h"]["other"])
        error(response, 404)
    error(c.put(f"/api/v1/projects/{pid}/files/order", headers=env["h"]["ro"], json={"ids": []}), 403)


@pytest.mark.parametrize("suffix", ["/view?offset=-1", "/view?limit=201", "/view?sheet=missing",
    "/view?block=-1", "/view?block=5000", "/view?block=0", "/blocks?limit=0",
    "/search?q=%20", "/search?q=x&offset=10001", "/search?q=" + "x" * 257,
    "/words/-1", "/words/0?limit=0", "/words/0?offset=-1", "/view?limit=1.0",
    "/view?unknown=1", "/blocks?offset=0&offset=1", "/content?which=working"])
def test_bounded_bad_query_inputs(env, suffix):
    error(env["client"].get(f"/api/v1/files/{env['src']['xlsx']['id']}" + suffix, headers=env["h"]["rw"]), 400)


@pytest.mark.parametrize("body", [{}, [], {"csv": []}, {"csv": {"encoding": "rot_13"}},
    {"csv": {"delimiter": "two"}}, {"csv": {"quotechar": False}}, {"csv": {"extra": 1}},
    {"ocr": True}, {"csv": {}, "extra": 1}])
def test_invalid_options_never_mutate(env, body):
    c, h, sid = env["client"], env["h"]["rw"], env["src"]["csv"]["id"]
    before = ok(c.get(f"/api/v1/files/{sid}", headers=h))
    error(c.put(f"/api/v1/files/{sid}/options", headers=h, json=body), 400)
    assert ok(c.get(f"/api/v1/files/{sid}", headers=h)) == before


@pytest.mark.parametrize("body", [{}, {"value": 1}, {"value": "x", "confirmed": "false"},
    {"value": "x", "confirmed": 0}, {"value": "x", "unknown": 1}, {"value": "x" * 20001}])
def test_invalid_corrections_never_mutate(env, body):
    c, h, sid = env["client"], env["h"]["rw"], env["src"]["xlsx"]["id"]
    before = ok(c.get(f"/api/v1/files/{sid}", headers=h))
    error(c.put(f"/api/v1/files/{sid}/corrections/{env['mid']}", headers=h, json=body), 400)
    assert ok(c.get(f"/api/v1/files/{sid}", headers=h)) == before


def test_streamed_body_limit_and_malformed_json(env):
    c, sid, h = env["client"], env["src"]["csv"]["id"], env["h"]["rw"]
    url = f"/api/v1/files/{sid}/options"
    json_headers = {**h, "Content-Type": "application/json"}
    error(c.put(url, headers=json_headers, content=iter([b'{"csv":', b' ' * (sources_api.MAX_BODY_BYTES + 1)])), 413)
    error(c.put(url, headers=json_headers, content=b'{'), 400)
    error(c.put(url, headers=json_headers, content=b'\xff'), 400)
    error(c.put(url, headers=json_headers, content=b'{"csv":{"delimiter":",","delimiter":";"}}'), 400)
    error(c.put(url, headers=h, content=b'{}'), 400)
    error(c.put(f"/api/v1/projects/{env['pid']}/files/order", headers=h, json={"ids": [True]}), 400)
    error(c.post(f"/api/v1/files/{sid}/representative", headers=h, json={"unknown": True}), 400)


def test_search_reports_its_bounded_continuation_window(env):
    c, h = env["client"], env["h"]["rw"]
    # Real synthetic CSV large enough to exercise the documented search window.
    src = upload(c, h, env["pid"], "csv", b"Name\n" + b"Match\n" * 10003)
    url = f"/api/v1/files/{src['id']}/search?q=Match&limit=1"
    previous = ok(c.get(url + "&offset=9999", headers=h))
    assert previous["next_offset"] == 10000 and not previous["truncated"]
    response = c.get(url + "&offset=10000", headers=h)
    last = ok(response)
    assert last["total"] == 10003 and last["truncated"] and last["next_offset"] is None
    assert last["items"][0]["r"] == previous["items"][0]["r"] + 1
    assert "Link" not in response.headers


def test_real_delegated_viewer_export_and_live_membership(env, monkeypatch):
    c, api = env["client"], env["api"]
    ident, store = api.state.identity, env["app"].store
    user = ident.create_user("source-viewer@example.test", "synthetic-test-passphrase-long", status="active")
    wid = ident.workspaces_for(user["id"])[0]["id"]
    owner = ident.create_user("source-owner@example.test", "synthetic-test-passphrase-long", status="active")
    ident.add_member(wid, owner["id"], "owner")
    # FIXTURE: real grants-store issuance, without invoking another application's routes.
    clients = Clients(store)
    clients.register(client_id="synthetic-source-reader", app_id="docweave", name="Synthetic source reader",
                     redirect_uris=["http://localhost/test-callback"], scopes="openid docweave.read docweave.write")
    token = Grants(store).issue_tokens(client=clients.get("synthetic-source-reader"),
        code_row={"user_id": user["id"], "scope": "docweave.read docweave.write", "central_session_id": None,
                  "workspace_id": wid}, with_refresh=True, audience=docs_resource())["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    pid = ok(c.post("/api/v1/projects", headers=h, json={"name": "Delegation fixture"}))["id"]
    src = upload(c, h, pid, "pdf", env["files"]["pdf"])
    sid = src["id"]
    # FIXTURE: a real stored session is deliberately insufficient for /api/v1.
    from synwevo_core.sessions import COOKIE_NAME
    cookie, _ = api.state.sessions.create(user["id"], workspace_id=wid, amr="pwd")
    c.cookies.set(COOKIE_NAME, cookie)
    error(c.get(f"/api/v1/files/{sid}/view"), 401)
    error(c.get(f"/api/v1/files/{sid}/content"), 401)
    c.cookies.clear()
    ident.add_member(wid, user["id"], "viewer", can_export=False)
    assert c.get(f"/api/v1/files/{sid}/view", headers=h).status_code == 200
    error(c.post(f"/api/v1/files/{sid}/representative", headers=h), 403)
    error(c.get(f"/api/v1/files/{sid}/content", headers=h), 403, "export_not_allowed")
    error(c.get(f"/api/v1/files/{sid}/page/0.png", headers=h), 403, "export_not_allowed")
    ident.add_member(wid, user["id"], "viewer", can_export=True)
    assert c.get(f"/api/v1/files/{sid}/content", headers=h).content == env["files"]["pdf"]
    assert c.get(f"/api/v1/files/{sid}/page/0.png", headers=h).content.startswith(b"\x89PNG\r\n\x1a\n")
    render = api.state.views.source_page_png

    def withdraw_export(*args):
        data = render(*args)  # Real render, then a concurrent permission change in the test store.
        ident.add_member(wid, user["id"], "viewer", can_export=False)
        return data

    monkeypatch.setattr(api.state.views, "source_page_png", withdraw_export)
    error(c.get(f"/api/v1/files/{sid}/page/0.png", headers=h), 403, "export_not_allowed")
    ident.remove_member(wid, user["id"])
    error(c.get(f"/api/v1/files/{sid}/view", headers=h), 401)
    error(c.get(f"/api/v1/files/{sid}/content", headers=h), 401)
    error(c.get(f"/api/v1/files/{sid}/page/0.png", headers=h), 401)


def test_openapi_explicit_operations_scopes_and_truthful_schemas(env):
    spec = openapi_subset(env["api"])
    operations = {
        ("/files/{fid}/structure", "get"): ("get_source_structure", "read"),
        ("/files/{fid}/view", "get"): ("get_source_view", "read"),
        ("/files/{fid}/blocks", "get"): ("list_source_blocks", "read"),
        ("/files/{fid}/search", "get"): ("search_source", "read"),
        ("/files/{fid}/words/{page}", "get"): ("list_source_page_words", "read"),
        ("/files/{fid}/page/{page}.png", "get"): ("render_source_page", "read"),
        ("/files/{fid}/locator", "post"): ("build_source_locator", "read"),
        ("/files/{fid}/options", "put"): ("update_source_options", "write"),
        ("/files/{fid}/corrections/{mid}", "put"): ("update_source_correction", "write"),
        ("/files/{fid}/representative", "post"): ("select_representative_source", "write"),
        ("/projects/{pid}/files/order", "put"): ("reorder_project_sources", "write"),
        ("/files/{fid}/content", "get"): ("download_source_content", "read"),
        ("/projects/{pid}/locate", "get"): ("locate_source_mappings", "read"),
        ("/projects/{pid}/records", "get"): ("list_source_records", "read"),
    }
    for (path, method), (opid, scope) in operations.items():
        op = spec["paths"]["/api/v1" + path][method]
        assert op["operationId"] == opid and {"apiKey": [scope]} in op["security"]
        assert "422" not in op["responses"]
        content = op["responses"]["200"]["content"]
        assert any(v.get("schema") for v in content.values())
        if method == "put":
            body = op["requestBody"]["content"]["application/json"]["schema"]
            assert body["additionalProperties"] is False and body["required"]
    options = spec["paths"]["/api/v1/files/{fid}/options"]["put"]["requestBody"]["content"]["application/json"]["schema"]
    assert "gb18030" in options["properties"]["csv"]["properties"]["encoding"]["enum"]
    correction = spec["paths"]["/api/v1/files/{fid}/corrections/{mid}"]["put"]["requestBody"]["content"]["application/json"]["schema"]
    assert {"type": "null"} in correction["properties"]["value"]["anyOf"]
    locator = spec["paths"]["/api/v1/files/{fid}/locator"]["post"]["requestBody"]["content"]["application/json"]["schema"]
    selection = locator["properties"]["selection"]
    assert "discriminator" not in selection  # Inline alternatives have no named schema targets.
    assert {option["properties"]["kind"]["const"] for option in selection["oneOf"]} == {
        "cells", "docx_text", "docx_cell", "docx_table", "pdf_region", "text_span", "pptx_text", "pptx_cell"}
    assert all("kind" in option["required"] for option in selection["oneOf"])
    # Every exported reference must resolve, including the recursive JSON payload type.
    def check_refs(node):
        if isinstance(node, dict):
            if "$ref" in node:
                target = spec
                for part in node["$ref"].split("/")[1:]:
                    target = target[part.replace("~1", "/").replace("~0", "~")]
            for value in node.values():
                check_refs(value)
        elif isinstance(node, list):
            for value in node:
                check_refs(value)
    check_refs(spec)
    assert len(set(opid for opid, _ in operations.values())) == len(operations)
