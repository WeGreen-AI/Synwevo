"""Value semantics end to end (docs/FORMATS.md §XLSX, §写入值): what a cell *shows* is what a
document gets — leading zeros kept by a number format, long identifiers stored as text, large
integers never in scientific notation, percentages and dates as their display text, floats
without binary noise — and an empty row in the middle of the data is not a sample."""
from __future__ import annotations

import datetime as dt
import io

import docx
import openpyxl


def _xlsx() -> bytes:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["编号", "长号", "大数", "金额", "比例", "日期", "备注", "补零"])
    ws.append(["0007", "12345678901234567890", 1234567890123456, 0.1 + 0.2, 0.125, dt.datetime(2026, 9, 26), None, 7])
    ws.append([None, None, None, None, None, None, None, None])       # an empty row: not a sample
    ws.append(["0008", "98765432109876543210", 9007199254740993, 1234.5, 1, dt.date(2026, 1, 5), "有备注", 42])
    ws["E2"].number_format = "0.0%"; ws["E4"].number_format = "0%"
    ws["F2"].number_format = "yyyy-mm-dd"; ws["F4"].number_format = "yyyy/m/d"
    ws["H2"].number_format = "000000"; ws["H4"].number_format = "000000"
    ws["C2"].number_format = "0"; ws["C4"].number_format = "0"
    buf = io.BytesIO(); wb.save(buf)
    # openpyxl itself serialises the integer through a double (the sheet XML says …992), so the
    # file is patched to hold the digits a producer that keeps int64 identifiers would write.
    return _replace_in_sheet(buf.getvalue(), "<v>9007199254740992</v>", "<v>9007199254740993</v>")


def _replace_in_sheet(data: bytes, old: str, new: str) -> bytes:
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(data)); out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            payload = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = payload.decode(); assert old in text, "fixture drifted"; payload = text.replace(old, new).encode()
            dst.writestr(item, payload)
    return out.getvalue()


def _docx() -> bytes:
    d = docx.Document()
    d.add_paragraph("编号：X编号 长号：X长号 大数：X大数 金额：X金额 比例：X比例 日期：X日期 备注：X备注 补零：X补零")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def test_display_values_reach_the_document_and_empty_rows_are_not_samples(api, tmp_path):
    pid = api.project("取值语义")
    api.upload_bytes(pid, "数据.xlsx", _xlsx())
    tpl = tmp_path / "模板.docx"; tpl.write_bytes(_docx())
    t = api.template(pid, tpl); tid = t["id"]
    slots = {name: api.text_slot(tid, f"X{name}", name)["binding"] for name in ("编号", "长号", "大数", "金额", "比例", "日期", "备注", "补零")}
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "数据", "header_row": 0}})
    recs = api.ok(api.c.get(f"/api/projects/{pid}/records"))
    assert recs["total"] == 2, recs
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "数据", "header": h}], "target": {"binding": b}} for h, b in slots.items()])
    job = api.job(pid, naming="{编号}")
    assert job["status"] == "completed", job
    outs = api.outputs(job["id"])
    assert set(outs) == {"0007.docx", "0008.docx"}, sorted(outs)
    texts = {k: docx.Document(io.BytesIO(v)).paragraphs[0].text for k, v in outs.items()}
    print(texts)
    first, second = texts["0007.docx"], texts["0008.docx"]
    assert "编号：0007 " in first and "编号：0008 " in second, "text cells keep their leading zeros"
    assert "长号：12345678901234567890 " in first and "长号：98765432109876543210 " in second, "long identifiers stored as text are not numbers"
    assert "大数：1234567890123456 " in first, "a 16-digit integer is not written in scientific notation"
    assert "大数：9007199254740993 " in second, "an integer beyond 2^53 keeps every digit"
    assert "金额：0.3 " in first and "金额：1234.5 " in second, "floats are shown without binary noise"
    assert "比例：12.5% " in first and "比例：100% " in second, "percent number formats give their display text"
    assert "日期：2026-09-26 " in first and "日期：2026/1/5 " in second, "date number formats give their display text"
    assert "备注： " in first and "备注：有备注 " in second, "an empty cell is empty, not 'None'"
    assert "补零：000007" in first and "补零：000042" in second, "a zero-padding number format is honoured"


def test_csv_keeps_every_cell_as_text(api, tmp_path):
    """CSV has no types (docs/FORMATS.md): leading zeros and long numbers are text and stay so."""
    csv = "编号,长号,大数,金额\n0007,12345678901234567890,9007199254740993,1.10\n0008,00000000000000000001,1e5,2.50\n".encode("utf-8-sig")
    pid = api.project("CSV 取值")
    api.upload_bytes(pid, "数据.csv", csv)
    tpl = tmp_path / "模板.docx"
    d = docx.Document(); d.add_paragraph("编号：X编号 长号：X长号 大数：X大数 金额：X金额"); d.save(tpl)
    tid = api.template(pid, tpl)["id"]
    slots = {name: api.text_slot(tid, f"X{name}", name)["binding"] for name in ("编号", "长号", "大数", "金额")}
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "数据", "header_row": 0}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "数据", "header": h}], "target": {"binding": b}} for h, b in slots.items()])
    job = api.job(pid, naming="{编号}")
    assert job["status"] == "completed", job
    texts = {k: docx.Document(io.BytesIO(v)).paragraphs[0].text for k, v in api.outputs(job["id"]).items()}
    assert texts["0007.docx"] == "编号：0007 长号：12345678901234567890 大数：9007199254740993 金额：1.10", texts
    assert texts["0008.docx"] == "编号：0008 长号：00000000000000000001 大数：1e5 金额：2.50", texts


def test_xls_reader_gives_the_same_values_and_display_text():
    """The XLS (97–2003) path reads through xlrd when there is no LibreOffice at run time; the
    values it yields must mean the same as the XLSX reader's. XLS stores every number as a
    double, so the integer beyond 2^53 is …992 in the file itself — that is the format, not
    the reader — and the assertion says so."""
    import pytest
    from synwevo_docs.engine.converters import convert, soffice_path
    from synwevo_docs.engine.tabular import read_xls, read_xlsx
    if not soffice_path():
        pytest.skip("LibreOffice is needed to produce the XLS fixture")
    data = _xlsx()
    xls = convert(data, "xlsx", "xls")
    a, b = read_xlsx(data).sheets[0], read_xls(xls).sheets[0]
    assert len(a.rows) == len(b.rows) == 4
    for ra, rb in zip(a.rows, b.rows):
        for va, vb in zip(ra, rb):
            if va.v == "9007199254740993":
                assert vb.v == "9007199254740992", "XLS holds doubles: the 16th digit is lost by the format"
                continue
            assert (va.t, va.v, va.text) == (vb.t, vb.v, vb.text), (va, vb)
