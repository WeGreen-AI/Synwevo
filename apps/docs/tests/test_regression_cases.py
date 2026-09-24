"""Regression cases: capture a project with its expected results, replay it, detect changes.
Also runs every committed case in regression/ (real, desensitized files once available)."""
import io
import json

import docx
import openpyxl
import pytest

from docweave import regression as RG


def test_committed_cases_still_pass():
    cases = RG.list_cases()
    if not cases:
        pytest.skip("regression/ 中还没有回归用例（拿到脱敏的真实文件后，在界面中「保存为回归用例」即可加入）")
    bad = [(c["id"], r) for c in cases for r in [RG.run(c["id"])] if not r["ok"]]
    assert not bad, json.dumps(bad, ensure_ascii=False, indent=1)[:4000]


@pytest.fixture()
def case_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCWEAVE_REGRESSION_DIR", str(tmp_path / "regression"))
    return tmp_path / "regression"


def _project(api, fixtures):
    pid = api.project("回归用例")
    api.upload(pid, fixtures / "测试_订单明细100.xlsx")
    t = api.template(pid, fixtures / "测试_订单模板.xlsx")
    s_order = api.slot(t["id"], kind="cell", sheet="订单", cell="B2", name="订单号")
    rep = api.slot(t["id"], kind="repeat", sheet="订单", row=4, c1=0, c2=3, name="明细")
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "订单号"}], "sort": [{"header": "行号", "dir": "asc", "numeric": True}]}})
    api.maps(pid, [
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "订单号"}], "target": {"binding": s_order["binding"]}},
        {"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": rep["binding"]},
         "columns": [{"slot_col": "c0", "source": {"header": "行号"}}, {"slot_col": "c1", "source": {"header": "品名"}},
                     {"slot_col": "c2", "source": {"header": "数量"}}, {"slot_col": "c3", "source": {"header": "单价"}}]},
    ])
    return pid


def test_capture_replay_and_detect_changes(api, fixtures, case_dir):
    pid = _project(api, fixtures)
    c = api.ok(api.c.post(f"/api/projects/{pid}/regression-cases", json={"name": "订单（测试数据）", "note": "15 个订单"}))
    assert c["samples"] == 15 and c["records_total"] == 15 and (case_dir / c["id"] / "case.json").exists()
    assert [x["id"] for x in api.ok(api.c.get("/api/regression-cases"))] == [c["id"]]
    r = api.ok(api.c.post(f"/api/regression-cases/{c['id']}/run"))
    assert r["ok"] and r["passed"] == 15, r
    # simulate a behaviour change: the stored expectation no longer matches what the code produces
    case = json.loads((case_dir / c["id"] / "case.json").read_text(encoding="utf-8"))
    case["samples"][0]["values"]["订单号"] = "DD-999"
    case["samples"][1]["text_sha1"] = "0" * 40
    (case_dir / c["id"] / "case.json").write_text(json.dumps(case, ensure_ascii=False), encoding="utf-8")
    r = api.ok(api.c.post(f"/api/regression-cases/{c['id']}/run"))
    assert not r["ok"] and r["passed"] == 13
    diffs = [d for s in r["samples"] for d in s["diffs"]]
    assert any("字段「订单号」" in d and "DD-999" in d for d in diffs) and any("输出文字变化" in d for d in diffs)
    z = api.c.get(f"/api/regression-cases/{c['id']}/download")
    assert z.status_code == 200 and z.content[:2] == b"PK"
    assert api.ok(api.c.delete(f"/api/regression-cases/{c['id']}"))["ok"] and not api.ok(api.c.get("/api/regression-cases"))
    assert api.c.post("/api/regression-cases/..%2Fetc/run").status_code in (400, 404, 405)
    with pytest.raises(RG.RegressionError):
        RG.load("../etc")


def test_cli(api, fixtures, case_dir):
    import subprocess
    import sys
    from pathlib import Path
    pid = _project(api, fixtures)
    api.ok(api.c.post(f"/api/projects/{pid}/regression-cases", json={"name": "CLI"}))
    root = Path(__file__).resolve().parents[1]
    p = subprocess.run([sys.executable, str(root / "scripts" / "regression.py")], capture_output=True, text=True,
                       env={**__import__("os").environ, "DOCWEAVE_REGRESSION_DIR": str(case_dir)})
    assert p.returncode == 0 and "[通过]" in p.stdout, p.stdout + p.stderr


def test_desensitize_copies(api, fixtures):
    pid = api.project("脱敏")
    api.upload(pid, fixtures / "测试_客户100.xlsx", fixtures / "测试_GBK编码.csv", fixtures / "测试_申请表.pdf")
    d = docx.Document()
    p = d.add_paragraph("客户：")
    p.add_run("测试赵一").bold = True
    p.add_run("，手机 13912345678，身份证 110101199003074518")
    b = io.BytesIO()
    d.save(b)
    api.upload_bytes(pid, "测试_联系单.docx", b.getvalue())
    res = api.ok(api.c.post(f"/api/projects/{pid}/desensitize", json={
        "words": ["测试赵一"], "columns": {"姓名": "pseudonym", "金额": "scale", "客户编号": "digits"}, "add_to_project": True}))
    names = {f["source"]: f for f in res["files"]}
    assert set(names) == {"测试_客户100.xlsx", "测试_GBK编码.csv", "测试_联系单.docx"}
    assert res["skipped"][0]["file"] == "测试_申请表.pdf" and "PDF" in res["skipped"][0]["reason"]
    assert "seed" not in res
    x = openpyxl.load_workbook(io.BytesIO(api.c.get(f"/api/desensitized/{names['测试_客户100.xlsx']['blob']}").content)).active
    orig = openpyxl.load_workbook(fixtures / "测试_客户100.xlsx").active
    assert x["B2"].value == "姓名0001" and x["B3"].value == "姓名0002" and x["A2"].value != orig["A2"].value and len(x["A2"].value) == 6
    assert x["D2"].value != orig["D2"].value and x["C2"].value == orig["C2"].value      # dates untouched unless chosen
    dd = docx.Document(io.BytesIO(api.c.get(f"/api/desensitized/{names['测试_联系单.docx']['blob']}").content)).paragraphs[0]
    assert "测试赵一" not in dd.text and "13912345678" not in dd.text and "19900307" in dd.text
    assert any(r.bold and r.text.startswith("脱敏词") for r in dd.runs)                  # formatting kept
    srcs = [s["filename"] for s in api.ok(api.c.get(f"/api/projects/{pid}/sources"))]
    assert "脱敏_测试_客户100.xlsx" in srcs and "脱敏_测试_联系单.docx" in srcs
