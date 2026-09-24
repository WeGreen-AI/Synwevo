"""Parsing untrusted files: every limit holds before the expensive work starts, and no single
upload can hold the process (and its GIL) for long.  All inputs are TEST fixtures built in
memory here; nothing reads the data directory."""
from __future__ import annotations

import io
import time
import zipfile

import pytest

from synwevo_docs.engine import detect as D
from synwevo_docs.engine import sources as S
from synwevo_docs.engine import tabular
from synwevo_docs.engine.limits import LIMITS, LimitError


# ---------------------------------------------------------------- zip sniffing (F174)
def _zip(entries: dict[str, bytes]) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n, d in entries.items():
            z.writestr(n, d)
    return b.getvalue()


def test_detect_refuses_an_oversized_content_types_part_before_inflating_it():
    # 8 MB of spaces compress to a few KB; the part is read only if its declared size is sane
    data = _zip({"[Content_Types].xml": b" " * (8 * 1024 * 1024), "word/document.xml": b"<x/>"})
    assert len(data) < 100_000
    with pytest.raises(LimitError):
        D.detect(data, "a.docx")                      # before: read all 8 MB and answered "docx"


def test_detect_refuses_a_zip_whose_declared_total_is_over_the_limit(monkeypatch):
    monkeypatch.setattr(LIMITS, "max_unzipped_bytes", 1024 * 1024)
    data = _zip({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<x/>", "word/media/big.bin": b"\0" * (2 * 1024 * 1024)})
    with pytest.raises(LimitError):
        D.detect(data, "a.docx")
    ps = S.parse_source(data, "a.docx")
    assert ps.status == "error" and "上限" in ps.error


def test_detect_refuses_a_falsified_member_size():
    """A member whose header claims a small size but inflates to more is cut at the cap."""
    data = bytearray(_zip({"mimetype": b"application/vnd.oasis.opendocument.text" + b" " * 4000, "content.xml": b"<x/>"}))
    # patch the central-directory size of "mimetype" down to 40 bytes
    cd = data.rfind(b"PK\x01\x02", 0, data.rfind(b"content.xml"))
    data[cd + 24:cd + 28] = (40).to_bytes(4, "little")
    lh = data.find(b"PK\x03\x04")
    data[lh + 22:lh + 26] = (40).to_bytes(4, "little")
    try:
        info = D.detect(bytes(data), "a.odt")
    except (LimitError, ValueError, zipfile.BadZipFile):
        return                                         # refused: fine
    assert info["format"] in ("odt", None)             # never read past the declared size


# ---------------------------------------------------------------- content sniffing (F184)
def test_text_starting_with_bm_is_not_a_bitmap():
    assert D.detect(b"BMI,name\n22,a\n", "x.csv")["format"] == "csv"      # was "bmp"
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (4, 4)).save(b, "BMP")
    assert D.detect(b.getvalue(), "x.bmp")["format"] == "bmp"


def test_pdf_marker_inside_text_is_not_a_pdf():
    assert D.detect(b"name,note\na,%PDF-1.4 is a version\n", "x.csv")["format"] == "csv"   # was "pdf"
    assert D.detect(b"%PDF-1.7\n%...", "x.bin")["format"] == "pdf"


# ---------------------------------------------------------------- HTML (F175)
def test_unclosed_html_tables_parse_in_linear_time():
    data = b"<table>" * 40000                       # 280 KB: minutes with the old lazy regexes
    t = time.perf_counter()
    ps = S.parse_source(data, "x.html")
    assert time.perf_counter() - t < 5
    assert ps.status == "ready"


def test_html_tables_and_text_are_read_as_before():
    html = ("<html><head><style>p{}</style><script>var a='<table>';</script></head><body><h1>标题</h1>"
            "<p>第一段 &amp; 其他<br>换行</p><table><tr><th>姓名</th><th colspan=2>电话</th></tr>"
            "<tr><td>张三</td><td>1</td><td>2</td></tr></table></body></html>")
    assert S.html_to_text(html) == "标题\n第一段 & 其他\n换行\n姓名\t电话\n张三\t1\t2"
    rows = [[c.text for c in r] for r in S.html_tables(html)[0]]
    assert rows == [["姓名", "电话", ""], ["张三", "1", "2"]]


def test_html_colspan_is_capped():
    rows = S.html_tables('<table><tr><td colspan="1000000000">x</td></tr></table>')[0]
    assert len(rows[0]) == 1000                        # was a billion cells


# ---------------------------------------------------------------- spreadsheet rows (F176)
def _xlsx_with_cell_at_row(r: int) -> bytes:
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
          '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
          '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    wb = ('<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
          '<sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>')
    wbrels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
    sh = ('<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
          f'<row r="{r}"><c r="A{r}" t="inlineStr"><is><t>x</t></is></c></row></sheetData></worksheet>')
    return _zip({"[Content_Types].xml": ct.encode(), "_rels/.rels": rels.encode(), "xl/workbook.xml": wb.encode(),
                 "xl/_rels/workbook.xml.rels": wbrels.encode(), "xl/worksheets/sheet1.xml": sh.encode()})


def test_a_row_gap_cannot_expand_a_tiny_xlsx_into_millions_of_rows(monkeypatch):
    monkeypatch.setattr(LIMITS, "max_rows", 10_000)
    data = _xlsx_with_cell_at_row(2_000_000)
    assert len(data) < 5_000
    ps = S.parse_source(data, "a.xlsx")
    assert ps.status == "error" and "行数超过上限" in ps.error      # was status=ready with 2,000,000 rows


def test_blank_csv_lines_count_as_rows(monkeypatch):
    monkeypatch.setattr(LIMITS, "max_rows", 1000)
    with pytest.raises(LimitError):
        tabular.read_csv(b"a\n" + b"\n" * 5000 + b"b\n")
    assert tabular.read_csv(b"a\n\nb\n").sheets[0].n_rows == 3


# ---------------------------------------------------------------- XML parts (F189)
def test_a_single_oversized_xml_part_is_refused(monkeypatch):
    import docx
    from synwevo_docs.engine.ooxml import Package
    d = docx.Document()
    for _ in range(200):
        d.add_paragraph("填充" * 50)
    b = io.BytesIO()
    d.save(b)
    monkeypatch.setattr(LIMITS, "max_xml_part_bytes", 10_000)
    pkg = Package(b.getvalue())
    with pytest.raises(LimitError):
        pkg.xml(pkg.main_part())


# ---------------------------------------------------------------- page rendering (F201)
def test_page_rendering_stays_within_the_pixel_budget(monkeypatch):
    pymupdf = pytest.importorskip("pymupdf")
    from PIL import Image
    from synwevo_docs.engine import pdf_engine as pe
    d = pymupdf.open()
    p = d.new_page(width=200, height=200)
    p.insert_text((20, 50), "x")
    d.xref_set_key(p.xref, "MediaBox", "[0 0 3000 3000]")     # page size is the uploader's choice
    data = d.tobytes()
    monkeypatch.setattr(LIMITS, "max_render_pixels", 2_000_000)
    png = pe.render_page_png(data, 0, 4.0)                     # asked for 12000 x 12000
    im = Image.open(io.BytesIO(png))
    assert im.width * im.height <= 2_000_000 * 1.01


# ---------------------------------------------------------------- images (F216)
def test_image_between_one_and_two_times_the_limit_is_refused(monkeypatch):
    from PIL import Image
    from synwevo_docs.engine import images
    b = io.BytesIO()
    Image.new("RGB", (1500, 1000), "white").save(b, "PNG")
    monkeypatch.setattr(LIMITS, "max_image_pixels", 1_000_000)      # 1.5 MP: Pillow alone only warns
    with pytest.raises(LimitError):
        images.inspect_image(b.getvalue())


def test_animated_gif_is_not_decoded_frame_by_frame():
    from PIL import Image
    from synwevo_docs.engine import images
    frames = [Image.new("RGB", (8, 8), (i * 5, 255 - i * 5, 0)) for i in range(50)]
    b = io.BytesIO()
    frames[0].save(b, "GIF", save_all=True, append_images=frames[1:])
    meta = images.inspect_image(b.getvalue())
    assert meta["animated"] is True and "frames" not in meta


# ---------------------------------------------------------------- built-in editor input (F214)
def test_builtin_table_spans_are_bounded():
    from synwevo_docs.engine import builtin
    cell = lambda a: {"type": "tableCell", "attrs": a, "content": [{"type": "paragraph", "content": [{"type": "text", "text": "x"}]}]}  # noqa: E731
    for attrs in ({"colspan": 10 ** 7}, {"rowspan": 50}, {"colspan": -3}):
        ed = {"body": {"type": "doc", "content": [{"type": "table", "content": [{"type": "tableRow", "content": [cell(attrs)]}]}]}}
        t = time.perf_counter()
        with pytest.raises(builtin.BuiltinError):
            builtin.build_docx(ed)
        assert time.perf_counter() - t < 2
    with pytest.raises(builtin.BuiltinError):
        builtin.build_xlsx({"sheets": [{"name": "S", "merges": ["A1:XFD1048576"]}]})


# ---------------------------------------------------------------- marker regex (F190 / F112)
def test_catastrophic_marker_regex_times_out_without_stalling_other_threads(monkeypatch):
    import threading
    from synwevo_docs.engine import blocks as B
    from synwevo_docs.engine.resolve import ResolveError, SourceDoc
    monkeypatch.setattr(B, "REGEX_TIMEOUT_S", 2)
    src = SourceDoc("T", "t.txt", "txt", "text")
    src._text = "a" * 40 + "!"
    ticks, stop = [], threading.Event()

    def tick():
        while not stop.is_set():
            ticks.append(time.perf_counter())
            time.sleep(0.01)
    th = threading.Thread(target=tick)
    th.start()
    try:
        with pytest.raises(ResolveError) as e:
            B.split(src, {"anchor": r"(a+)+$", "match": "regex"})
    finally:
        stop.set()
        th.join()
    assert e.value.code == "regex_timeout"
    assert max(b - a for a, b in zip(ticks, ticks[1:])) < 1.0      # before: one gap as long as the whole match


def test_marker_regex_still_splits_and_is_length_limited():
    from synwevo_docs.engine import blocks as B
    from synwevo_docs.engine.resolve import ResolveError, SourceDoc
    src = SourceDoc("T", "t.txt", "txt", "text")
    src._text = "前言\n人员登记表 1\n姓名 甲\n人员登记表 2\n姓名 乙"
    assert [b.label for b in B.split(src, {"anchor": r"^人员登记表 \d", "match": "regex"})] == ["人员登记表 1", "人员登记表 2"]
    with pytest.raises(ResolveError):
        B.matcher({"anchor": "a" * (LIMITS.max_regex_chars + 1), "match": "regex"})


# ---------------------------------------------------------------- configuration (F363)
def test_limits_honour_the_pre_rebrand_variable_names(monkeypatch):
    from synwevo_docs.engine import limits
    monkeypatch.delenv("SYNWEVO_DOCS_MAX_FILE_MB", raising=False)
    monkeypatch.delenv("SYNWEVO_MAX_FILE_MB", raising=False)
    monkeypatch.setenv("DOCWEAVE_MAX_FILE_MB", "7")
    assert limits._env_int("MAX_FILE_MB", 100) == 7                  # was silently 100
    monkeypatch.setenv("SYNWEVO_DOCS_MAX_FILE_MB", "9")
    assert limits._env_int("MAX_FILE_MB", 100) == 9
    monkeypatch.setenv("SYNWEVO_DOCS_MAX_FILE_MB", "lots")
    assert limits._env_int("MAX_FILE_MB", 100) == 100


# ---------------------------------------------------------------- LibreOffice child (F181 / F194)
def test_libreoffice_gets_no_server_secrets(monkeypatch):
    from synwevo_docs.engine import converters as C
    monkeypatch.setenv("SYNWEVO_SECRET", "TEST-FIXTURE-SECRET")
    monkeypatch.setenv("SYNWEVO_DOCS_WEBHOOK_SECRET", "TEST-FIXTURE-SECRET")
    env = C._child_env()
    assert "TEST-FIXTURE-SECRET" not in env.values() and not any(k.startswith(("SYNWEVO", "DOCWEAVE")) for k in env)
    assert env["HOME"] == str(C._PROFILE_DIR)


def test_libreoffice_profile_parent_must_belong_to_this_user(monkeypatch, tmp_path):
    import os
    from synwevo_docs.engine import converters as C
    target = tmp_path / "elsewhere"
    target.mkdir()
    link = tmp_path / "lo"
    link.symlink_to(target)                       # a pre-created path this process does not own
    monkeypatch.setattr(C, "_PROFILE_BASE", link)
    monkeypatch.setattr(C, "_BASE_CHECKED", False)
    with pytest.raises(C.ConversionError):
        C._private_base()
    good = tmp_path / "mine"
    monkeypatch.setattr(C, "_PROFILE_BASE", good)
    C._private_base()
    assert good.is_dir() and (os.stat(good).st_mode & 0o077) == 0


def test_libreoffice_profile_is_seeded_with_macros_and_links_disabled(monkeypatch, tmp_path):
    from synwevo_docs.engine import converters as C
    monkeypatch.setattr(C, "_PROFILE_DIR", tmp_path / "profile-1")
    C._seed_registry()
    C._seed_registry()                            # idempotent
    xcu = (tmp_path / "profile-1" / "user" / "registrymodifications.xcu").read_text(encoding="utf-8")
    for name in ("DisableMacrosExecution", "DisableActiveContent", "BlockUntrustedRefererLinks"):
        assert xcu.count(f'oor:name="{name}"') == 1


@pytest.mark.skipif(not __import__("synwevo_docs.engine.converters", fromlist=["x"]).soffice_path(), reason="需要 LibreOffice")
def test_workbook_with_network_formula_is_not_recalculated_by_libreoffice(monkeypatch):
    import openpyxl
    from synwevo_docs.engine import xlsx_cache
    from synwevo_docs.engine import templates as T
    calls = []
    monkeypatch.setattr(xlsx_cache, "recalc_xlsx", lambda data: calls.append(1) or data)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["B1"] = '=WEBSERVICE("http://127.0.0.1:9/probe")'
    b = io.BytesIO()
    wb.save(b)
    work, slot = T.xlsx_add_slot(b.getvalue(), {"sheet": ws.title, "kind": "cell", "cell": "A1", "name": "a"}, [])
    from synwevo_docs.engine import values as V
    out, issues = T.xlsx_fill(work, [slot], {slot["id"]: {"kind": "scalar", "value": V.text("y")}})
    assert calls == []                                            # LibreOffice was never asked
    assert any(i["code"] == "formula_unsupported" and "联网" in i["message"] for i in issues)


# ---------------------------------------------------------------- OCR (F217)
def test_tesseract_language_list_is_asked_once(monkeypatch):
    from synwevo_docs.engine import ocr
    calls = []

    class R:
        stdout = "List of available languages (2):\nchi_sim\neng\n"
    monkeypatch.setattr(ocr, "tesseract_path", lambda: "/fixture/tesseract")
    monkeypatch.setattr(ocr, "tessdata_dir", lambda: None)
    monkeypatch.setattr(ocr, "_LANGS", {})
    monkeypatch.setattr(ocr.subprocess, "run", lambda *a, **k: calls.append(k.get("env", {}).get("OMP_THREAD_LIMIT")) or R())
    assert ocr.available_langs() == ["chi_sim", "eng"] and ocr.available_langs() == ["chi_sim", "eng"]
    assert calls == ["1"]
