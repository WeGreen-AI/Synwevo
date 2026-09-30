"""The committed OpenAPI file and the example definition must match the running code."""
import json
import subprocess
import sys

import yaml

ROOT = __import__("synwevo_core.http", fromlist=["repo_root"]).repo_root()
SPEC_FILE = ROOT / "docs" / "integration" / "openapi.yaml"
METHODS = {"get", "put", "post", "delete", "patch", "head", "options", "trace"}


def _spec() -> dict:
    return yaml.safe_load(SPEC_FILE.read_text(encoding="utf-8"))


def _codes(op: dict) -> dict:
    # Response codes are compared as strings: the CLI's YAML writer does not quote them, so a
    # YAML parser reads `202:` as an integer key.
    return {str(k): v for k, v in op["responses"].items()}


def test_openapi_file_is_in_sync(tmp_path):
    out = tmp_path / "openapi.yaml"
    p = subprocess.run([sys.executable, "-m", "synwevo_docs.cli", "--data", str(tmp_path / "data"),
                        "openapi", "--out", str(out)], capture_output=True, text=True)
    assert p.returncode == 0, p.stderr[-800:]
    committed = SPEC_FILE.read_text(encoding="utf-8")
    assert out.read_text(encoding="utf-8") == committed, "docs/integration/openapi.yaml 与代码不一致：运行 python -m synwevo_docs.cli openapi"


def test_every_machine_route_is_documented_and_nothing_else(client):
    """Drift in the one direction a text diff explains badly: which operations exist.

    Compared against the *running application's* routes, not against a second export of the
    spec, so a route hidden from the schema (include_in_schema=False) or added by another module
    (governance, webhooks, uploads) without being in the committed file fails here by name.
    """
    from fastapi import routing
    # FastAPI 0.14x keeps included routers as lazy wrappers in `app.routes`; its own iterator
    # yields the effective routes with their full paths (older versions list them flat).
    walk = getattr(routing, "iter_route_contexts", None)
    items = walk(client.app.routes) if walk else client.app.routes
    live = {(m.lower(), r.path) for r in items if (getattr(r, "path", "") or "").startswith("/api/v1")
            for m in (getattr(r, "methods", None) or ()) if m != "HEAD"}
    assert len(live) > 40, "没有取到运行中的机器接口路由"
    documented = {(m, p) for p, ops in _spec()["paths"].items() if p.startswith("/api/v1") for m in ops if m in METHODS}
    missing, extra = sorted(live - documented), sorted(documented - live)
    assert not missing and not extra, (f"未写入 openapi.yaml 的路由：{missing}；openapi.yaml 中已不存在的路由：{extra}。"
                                       "运行 python -m synwevo_docs.cli openapi 重新生成")


def test_openapi_describes_only_machine_routes_and_needs_auth():
    spec = _spec()
    assert str(spec["openapi"]).startswith("3.1")
    paths = set(spec["paths"])
    assert {"/healthz", "/readyz"} <= paths
    assert all(p.startswith("/api/v1") for p in paths - {"/healthz", "/readyz"})
    assert not any(p.startswith("/api/") and "/v1" not in p for p in paths)      # no UI-internal routes
    # both machine credentials are described, and either one satisfies the global requirement
    schemes = spec["components"]["securitySchemes"]
    assert schemes["apiKey"]["scheme"] == "bearer"
    assert schemes["platformToken"]["type"] == "oauth2"
    assert set(schemes["platformToken"]["flows"]["clientCredentials"]["scopes"]) == \
        {"docweave.read", "docweave.write", "docweave.manage"}
    assert spec["security"] == [{"apiKey": []}, {"platformToken": []}]
    for probe in ("/healthz", "/readyz"):
        assert spec["paths"][probe]["get"]["security"] == []
    # the signed download link is its own credential: documented as needing no key
    assert spec["paths"]["/api/v1/download/{token}"]["get"]["security"] == []
    # per-operation scopes in both vocabularies
    assert spec["paths"]["/api/v1/projects"]["post"]["security"] == [{"apiKey": ["write"]}, {"platformToken": ["docweave.write"]}]
    assert spec["paths"]["/api/v1/projects/{pid}"]["delete"]["security"] == \
        [{"apiKey": ["admin"]}, {"platformToken": ["docweave.manage"]}]
    # Component schemas are pruned to what the kept paths reach. FastAPI emits one per route in
    # the whole app, so without pruning this file describes UI and OIDC request bodies it does
    # not document — including, once the account service was added, the OIDC form bodies.
    names = " ".join(spec["components"].get("schemas", {}))
    assert "oidc" not in names.lower()
    for name in spec["components"].get("schemas", {}):
        if name.startswith("Body_"):
            assert "_api_v1_" in name, f"{name} 属于非 v1 路由，不应出现在机器接口文档中"


def test_openapi_describes_the_error_and_header_contract():
    spec = _spec()
    comps = spec["components"]
    err = comps["schemas"]["Error"]["properties"]["error"]
    assert {"code", "message", "request_id"} <= set(err["required"]) and "field" in err["properties"]
    assert {"type", "title", "status", "detail", "instance"} <= set(comps["schemas"]["Problem"]["properties"])
    assert {"X-Request-ID", "Idempotent-Replayed", "RateLimit", "RateLimit-Policy", "Deprecation", "Retry-After"} <= set(comps["headers"])
    creating = [("post", "/api/v1/projects"), ("post", "/api/v1/projects/{pid}/files"), ("post", "/api/v1/templates"),
                ("post", "/api/v1/templates/{tid}/fields"), ("post", "/api/v1/projects/{pid}/jobs"),
                ("post", "/api/v1/jobs/{jid}/retry"), ("post", "/api/v1/artifacts/{aid}/link"),
                ("post", "/api/v1/jobs/{jid}/archive-link"), ("post", "/api/v1/uploads")]
    for m, p in creating:
        op = spec["paths"][p][m]
        assert {"$ref": "#/components/parameters/IdempotencyKey"} in op["parameters"], (m, p)
        ok = next(v for k, v in _codes(op).items() if k.startswith("2"))
        assert "Idempotent-Replayed" in ok["headers"], (m, p)
    job = spec["paths"]["/api/v1/projects/{pid}/jobs"]["post"]
    assert job["requestBody"]["content"]["application/json"]["schema"]["properties"]["idempotency_key"]["deprecated"] is True
    assert "Deprecation" in _codes(job)["202"]["headers"]
    # every authenticated operation documents the error envelope, and never FastAPI's 422
    for p, ops in spec["paths"].items():
        for m, op in ops.items():
            if m not in METHODS or p in ("/healthz", "/readyz"):
                continue
            assert "422" not in _codes(op), (m, p)
            assert "400" in _codes(op), (m, p)
    lst = spec["paths"]["/api/v1/projects"]["get"]
    assert {p["name"] for p in lst["parameters"] if "name" in p} >= {"offset", "limit", "cursor"}
    assert "Link" in _codes(lst)["200"]["headers"]
    # binary downloads are typed as binary, not as JSON
    dl = _codes(spec["paths"]["/api/v1/artifacts/{aid}/content"]["get"])["200"]["content"]
    assert "application/json" not in dl and dl["application/octet-stream"]["schema"]["format"] == "binary"
    # A self-authenticating download URL still refuses an invalid/revoked issuer
    # or an expired signature. Rendering also refuses changed/blocked inputs.
    signed = _codes(spec["paths"]["/api/v1/download/{token}"]["get"])
    assert signed["403"]["$ref"] == "#/components/responses/Forbidden"
    assert signed["410"]["$ref"] == "#/components/responses/Gone"
    preview = _codes(spec["paths"]["/api/v1/projects/{pid}/preview"]["post"])
    assert preview["409"]["$ref"] == "#/components/responses/Conflict"


def test_upload_contract_declares_resume_failures_and_all_public_states():
    spec = _spec()
    operations = {
        ("post", "/api/v1/uploads"): ("201", {"409", "410", "413"}),
        ("get", "/api/v1/uploads/{uid}"): ("200", set()),
        ("put", "/api/v1/uploads/{uid}"): ("200", {"409", "410", "416"}),
        ("post", "/api/v1/uploads/{uid}/complete"): ("200", {"409", "410"}),
        ("post", "/api/v1/uploads/{uid}/abort"): ("200", {"409"}),
    }
    for (method, path), (success, failures) in operations.items():
        responses = _codes(spec["paths"][path][method])
        assert failures <= set(responses), (method, path)
        assert responses[success]["content"]["application/json"]["schema"] == {"$ref": "#/components/schemas/UploadView"}
        for failure in failures:
            component = spec["components"]["responses"][responses[failure]["$ref"].split("/")[-1]]
            assert set(component["content"]) == {"application/json", "application/problem+json"}
            assert "X-Request-ID" in component["headers"]
    shape = spec["components"]["schemas"]["UploadView"]
    assert set(shape["properties"]["status"]["enum"]) == {"open", "complete", "consumed", "expired", "aborted"}
    assert {"sha256", "purpose", "complete", "received", "upload_url"} <= set(shape["required"])
    assert not {"tenant", "principal", "blob", "consumed_by"} & set(shape["properties"])


def test_copy_contract_requires_user_delegation_and_types_public_results():
    from synwevo_docs.api.v1 import SCOPE_NAMES

    spec = _spec()
    copies = {
        ("get", "/api/v1/copy-targets"): ("read", "CopyTargetsResponse"),
        ("get", "/api/v1/projects/{pid}/copy-manifest"): ("read", "CopyManifestResponse"),
        ("get", "/api/v1/projects/{pid}/copy-requests"): ("read", "CopyRequestList"),
        ("post", "/api/v1/projects/{pid}/copy-requests"): ("write", "CopyRequestView"),
        ("get", "/api/v1/copy-requests/{cid}"): ("read", "CopyRequestView"),
        ("post", "/api/v1/copy-requests/{cid}/withdraw"): ("write", "CopyRequestView"),
    }
    delegated = set()
    for path, ops in spec["paths"].items():
        for method, op in ops.items():
            if method not in METHODS:
                continue
            if (method, path) not in copies:
                # Adding a restricted flow must not silently remove the credentials accepted
                # by any existing operation, including probes and signed download links.
                assert not any("delegatedToken" in scheme for scheme in op.get("security", []))
                continue
            scope, response = copies[(method, path)]
            assert op["x-synwevo-identity"] == "delegated"
            assert op["security"] == [{"delegatedToken": [SCOPE_NAMES[scope]]}]
            ok = next(v for k, v in _codes(op).items() if k.startswith("2"))
            assert ok["content"]["application/json"]["schema"] == {"$ref": f"#/components/schemas/{response}"}
            assert {"401", "403"} <= set(_codes(op))
            delegated.add((method, path))
    assert delegated == set(copies)
    create = spec["paths"]["/api/v1/projects/{pid}/copy-requests"]["post"]
    assert create["requestBody"]["content"]["application/json"]["schema"] == {"$ref": "#/components/schemas/CopyRequestCreate"}
    for operation in (create, spec["paths"]["/api/v1/copy-requests/{cid}/withdraw"]["post"]):
        assert _codes(operation)["409"]["$ref"] == "#/components/responses/Conflict"
    for name, schema in spec["components"]["schemas"].items():
        if name.startswith("Copy"):
            assert "requested_authority" not in json.dumps(schema)


def test_delegated_oauth_flow_uses_actual_account_endpoints(client):
    from synwevo_core import oidc
    from synwevo_docs.api.v1 import SCOPE_NAMES, openapi_subset

    discovery = oidc.discovery()
    spec = openapi_subset(client.app, issuer=discovery["issuer"] + "/")
    schemes = spec["components"]["securitySchemes"]
    delegated = schemes["delegatedToken"]
    assert delegated["type"] == "oauth2" and delegated["x-synwevo-identity"] == "delegated"
    assert set(delegated["flows"]) == {"authorizationCode"}  # client_credentials cannot act for an owner
    flow = delegated["flows"]["authorizationCode"]
    assert flow["authorizationUrl"] == discovery["authorization_endpoint"]
    assert flow["tokenUrl"] == discovery["token_endpoint"]
    assert "authorization_code" in discovery["grant_types_supported"]
    assert set(flow["scopes"]) == set(SCOPE_NAMES.values())
    assert set(flow["scopes"]) <= set(discovery["scopes_supported"])
    assert schemes["platformToken"]["flows"]["authorizationCode"] == flow
    assert schemes["platformToken"]["flows"]["clientCredentials"]["tokenUrl"] == flow["tokenUrl"]
    custom = openapi_subset(client.app, issuer="https://accounts.example.test/custom/")
    custom_flow = custom["components"]["securitySchemes"]["delegatedToken"]["flows"]["authorizationCode"]
    assert custom_flow["authorizationUrl"] == "https://accounts.example.test/custom/oidc/authorize"
    assert custom_flow["tokenUrl"] == "https://accounts.example.test/custom/oidc/token"


def test_openapi_describes_webhooks():
    spec = _spec()
    from synwevo_docs import webhooks as W
    hooks = spec["webhooks"]
    assert set(hooks) == set(W.EVENT_TYPES) | {W.PING_EVENT}
    op = hooks["job.completed"]["post"]
    headers = {p["name"]: p for p in op["parameters"]}
    assert headers["webhook-signature"]["required"] is True and headers["webhook-id"]["required"] is True
    assert "X-DocWeave-Signature" in headers
    body = op["requestBody"]["content"]["application/json"]["schema"]
    assert body == {"$ref": "#/components/schemas/WebhookEvent"}
    assert {"id", "type", "version", "occurred_at"} <= set(spec["components"]["schemas"]["WebhookEvent"]["required"])


def test_example_definition_produces_the_right_text(api, fixtures):
    """The definition shipped with the docs really drives a template + rule + mappings — and the
    documents it produces say what they should.

    The old assertion (`"年" in text and "," in text`) passed while every output read
    「甲方（客户）：示例测试钱磊敏」: the example's `find` covered only 「客户名称」 and left the
    template's 「示例」 in every contract.
    """
    doc = json.loads((ROOT / "docs" / "integration" / "example-definition.json").read_text(encoding="utf-8"))
    pid = api.project("示例定义")
    api.upload(pid, fixtures / "测试_客户100.xlsx")
    t = api.template(pid, fixtures / "测试_合同模板.docx")
    from synwevo_docs import definitions as D
    created = D.import_template_fields(api.c.app_state.app, t["id"], {"schema": "docweave.template-fields", "version": 1,
                                                                      "fields": doc["fields"]})
    names = [c["name"] for c in created["created"]]
    assert names == [f["name"] for f in doc["fields"]]
    api.ok(api.c.patch(f"/api/projects/{pid}", json={"template_id": t["id"], "record_rule": doc["rule"]}))
    fields = {s["name"]: s for s in api.ok(api.c.get(f"/api/templates/{t['id']}"))["version"]["slots"]}
    maps = []
    for m in doc["mappings"]:
        m = dict(m)
        maps.append({**m, "target": {"binding": fields[m.pop("field")]["binding"]}})
    api.maps(pid, maps)
    job = api.job(pid, naming="{客户编号}_{客户名称}")
    assert job["status"] == "completed" and job["counts"]["success"] == 10
    import io

    import docx
    import openpyxl
    out = api.outputs(job["id"])
    ws = openpyxl.load_workbook(fixtures / "测试_客户100.xlsx")["客户"]
    rows = {str(r[0]): r for r in ws.iter_rows(min_row=2, max_row=11, values_only=True)}
    for name in sorted(out)[:3]:
        text = [p.text for p in docx.Document(io.BytesIO(out[name])).paragraphs]
        cid = name.split("_", 1)[0]
        _, person, signed, amount, _ = rows[cid]
        assert text[1] == f"甲方（客户）：{person}，客户编号 {cid}。", text[1]
        assert text[2] == (f"签约日期：{signed.year}年{signed.month}月{signed.day}日；"
                           f"合同金额：人民币 {amount:,.2f} 元。"), text[2]
        assert "示例" not in "\n".join(text)
