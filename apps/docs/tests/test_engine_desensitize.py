"""Desensitized copies must not contain what the rules matched — anywhere in the package.

Every document here is a TEST fixture built in memory; the sensitive-looking values are
synthetic.  Each test reads the produced file back (all of its XML, not only the body text).
"""
from __future__ import annotations

import io
import zipfile

import docx
import openpyxl
import pytest
from docx.oxml.ns import qn
from lxml import etree

from synwevo_docs.engine import desensitize as D
from synwevo_docs.engine import docx_core as dc

SECRETS = ("测试赵一", "13912345678", "13800001111", "13700002222", "13600003333", "bob@corp.com", "corp.com")


def _all_xml(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    return "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith((".xml", ".rels")))


def _docx_fixture() -> bytes:
    """TEST fixture: the same data in every place a Word document can hold text."""
    d = docx.Document()
    d.core_properties.author = "测试赵一"
    d.core_properties.last_modified_by = "测试赵一"
    d.core_properties.title = "测试赵一的联系单"
    d.add_paragraph("联系人 测试赵一 手机 13912345678")
    # a value already inside a field control (keep_fields output / a Synwevo template)
    p2 = d.add_paragraph("手机：")
    sdt = dc.make_sdt("dw:phone", "手机")
    r = etree.SubElement(sdt.find(qn("w:sdtContent")), qn("w:r"))
    etree.SubElement(r, qn("w:t")).text = "13800001111"
    p2._p.append(sdt)
    # a match that starts in plain text and continues into a hyperlink run
    p3 = d.add_paragraph("邮箱 bob")
    hl = etree.SubElement(p3._p, qn("w:hyperlink"))
    r3 = etree.SubElement(hl, qn("w:r"))
    etree.SubElement(r3, qn("w:t")).text = "@corp.com 电话"
    # a tracked deletion by a named reviewer, and a field code
    p4 = d.add_paragraph("旧号码 ")
    de = etree.SubElement(p4._p, qn("w:del"), {qn("w:id"): "1", qn("w:author"): "测试赵一", qn("w:date"): "2026-01-01T00:00:00Z"})
    r4 = etree.SubElement(de, qn("w:r"))
    etree.SubElement(r4, qn("w:delText")).text = "13700002222"
    p5 = d.add_paragraph()
    r5 = etree.SubElement(p5._p, qn("w:r"))
    etree.SubElement(r5, qn("w:instrText")).text = ' HYPERLINK "mailto:bob@corp.com" '
    # header and a relationship to mailto:
    d.sections[0].header.paragraphs[0].text = "页眉 13600003333"
    d.part.relate_to("mailto:bob@corp.com", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                     is_external=True)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_docx_copy_contains_none_of_the_matched_values():
    sess = D.Session(words=["测试赵一"])
    out, ext = D.desensitize(_docx_fixture(), "x.docx", sess)
    assert ext == "docx"
    xml = _all_xml(out)
    leaked = [s for s in SECRETS if s in xml]
    # before: the SDT value, the hyperlink-split e-mail, the deleted text, the field code, the
    # relationship target and every document property were left in place
    assert leaked == []
    body = [dc.para_text(p._p) for p in docx.Document(io.BytesIO(out)).paragraphs]
    assert body[0].startswith("联系人 脱敏词0001 手机 ") and body[2] == "邮箱 user0001@example.com 电话"
    props = docx.Document(io.BytesIO(out)).core_properties
    assert props.author == "" and props.last_modified_by == "" and "脱敏词0001" in props.title
    # consistent: the same e-mail got the same pseudonym in the text, the field code and the link
    assert xml.count("user0001@example.com") >= 3


def test_docx_formatting_of_the_replaced_text_is_kept():
    d = docx.Document()
    p = d.add_paragraph("客户：")
    p.add_run("测试赵一").bold = True
    b = io.BytesIO()
    d.save(b)
    out, _ = D.desensitize(b.getvalue(), "x.docx", D.Session(words=["测试赵一"]))
    runs = docx.Document(io.BytesIO(out)).paragraphs[0].runs
    assert any(r.bold and r.text == "脱敏词0001" for r in runs)


def test_docm_copy_is_a_plain_docx_without_the_vba_project():
    d = docx.Document()
    d.add_paragraph("测试赵一")
    b = io.BytesIO()
    d.save(b)
    zin = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                                    b"application/vnd.ms-word.document.macroEnabled.main+xml")
                data = data.replace(b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
            z.writestr(n, data)
        z.writestr("word/vbaProject.bin", "Sub A: MsgBox \"测试赵一\": End Sub".encode("utf-8"))   # TEST fixture VBA
    out, ext = D.desensitize(buf.getvalue(), "x.docm", D.Session(words=["测试赵一"]))
    z = zipfile.ZipFile(io.BytesIO(out))
    assert ext == "docx" and "word/vbaProject.bin" not in z.namelist()
    assert b"macroEnabled" not in z.read("[Content_Types].xml")
    docx.Document(io.BytesIO(out))            # opens as a normal document


def test_docx_with_an_unknown_embedded_object_is_refused():
    d = docx.Document()
    d.add_paragraph("x")
    b = io.BytesIO()
    d.save(b)
    zin = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n in zin.namelist():
            z.writestr(n, zin.read(n))
        z.writestr("word/embeddings/oleObject1.bin", b"\0" * 32)
    with pytest.raises(D.DesensitizeError) as e:
        D.desensitize(buf.getvalue(), "x.docx", D.Session())
    assert e.value.code == "embedded_object"


def _xlsx_fixture() -> bytes:
    from openpyxl.comments import Comment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试赵一"
    ws.append(["姓名", "手机", "金额"])
    ws.append(["测试赵一", 13800138000, 12.5])                   # the phone is a number, as Excel stores it
    ws.append(["李四", 13912345678.0, 7.25])
    ws["A2"].comment = Comment("联系 13700002222", "测试赵一")
    ws["D2"] = "官网"
    ws["D2"].hyperlink = "mailto:bob@corp.com"
    ws["E2"] = '=IF(A2="测试赵一","VIP","")'
    ws.oddHeader.center.text = "13600003333"
    wb.properties.creator = "测试赵一"
    wb.properties.title = "测试赵一的名单"
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_xlsx_numbers_comments_headers_links_and_properties_are_scrubbed():
    # Synthetic seed that previously gave the phone a leading zero, losing a digit in Excel.
    sess = D.Session(seed="leading-zero-14", words=["测试赵一"])
    out, ext = D.desensitize(_xlsx_fixture(), "x.xlsx", sess)
    xml = _all_xml(out)
    assert [s for s in SECRETS if s in xml] == []
    wb = openpyxl.load_workbook(io.BytesIO(out))
    ws = wb.worksheets[0]
    assert ws.title == "脱敏词0001"
    assert isinstance(ws["B2"].value, int) and ws["B2"].value != 13800138000 and len(str(ws["B2"].value)) == 11
    assert isinstance(ws["B3"].value, (int, float)) and int(ws["B3"].value) != 13912345678
    assert ws["C2"].value == 12.5                                   # an amount matches no rule: unchanged
    assert ws["E2"].value == '=IF(A2="脱敏词0001","VIP","")'        # the literal follows the cell it compares with
    assert ws["A2"].value == "脱敏词0001" and wb.properties.creator in ("", None)


def test_xlsx_formula_referencing_a_renamed_sheet_follows_it():
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "测试赵一"
    a["A1"] = 5
    b = wb.create_sheet("汇总")
    b["A1"] = "='测试赵一'!A1*2"
    b["A2"] = "=测试赵一!A1+1"
    buf = io.BytesIO()
    wb.save(buf)
    out, _ = D.desensitize(buf.getvalue(), "x.xlsx", D.Session(words=["测试赵一"]))
    wb2 = openpyxl.load_workbook(io.BytesIO(out))
    assert wb2.sheetnames == ["脱敏词0001", "汇总"]
    assert wb2["汇总"]["A1"].value == "='脱敏词0001'!A1*2" and wb2["汇总"]["A2"].value == "='脱敏词0001'!A1+1"


def test_digits_strategy_changes_floats_and_keeps_their_type():
    sess = D.Session(seed="fixture")
    v = sess.column_value("digits", "编号", 123456.75)
    assert isinstance(v, float) and v != 123456.75 and len(D._num_text(v).replace(".", "")) <= 8


@pytest.mark.parametrize("value", [13800138000, -13800138000, 13800138000.0, -13800138000.0])
@pytest.mark.parametrize("strategy", ["patterns", "digits"])
def test_numeric_replacements_keep_width_and_match_the_same_text(value, strategy):
    sess = D.Session(seed="leading-zero-14")
    text = D._num_text(value)
    if strategy == "patterns":
        number, written = sess.replace_value(value), sess.replace_text(text)
    else:
        number = sess.column_value("digits", "synthetic number", value)
        written = sess.column_value("digits", "synthetic number", text)
    assert type(number) is type(value) and number != value
    assert D._num_text(number) == written  # Numeric and text copies still join on the same key.
    assert len(written.lstrip("-")) == len(text.lstrip("-"))


def test_pseudonyms_stay_consistent_and_numbered_per_kind():
    sess = D.Session()
    names = [sess.pseudonym("姓名", f"人{i}") for i in range(3)] + [sess.pseudonym("姓名", "人1")]
    assert names == ["姓名0001", "姓名0002", "姓名0003", "姓名0002"]
    assert sess.pseudonym("公司", "某公司") == "公司0001"


def test_word_list_prefers_the_longest_word_at_a_position():
    sess = D.Session(words=["测试", "测试赵一"])
    assert sess.replace_text("测试赵一和测试") == "脱敏词0001和脱敏词0002"


def test_scaling_never_copies_a_non_zero_amount_unchanged():
    """CI-84: the factor used to be any of 0.6…1.4 rounded to 0.001, so about one session in 800
    drew exactly 1.000 and every "scaled" amount came out as the real one."""
    factors = [D.Session(seed=f"scale-{i}").factor for i in range(3000)]
    assert all(abs(f - 1) >= 0.05 and 0.6 <= f <= 1.4 for f in factors)
    assert any(f > 1 for f in factors) and any(f < 1 for f in factors)
    sess = D.Session(seed="fixture")
    for factor in (1.0, 1.06, 0.94):
        sess.factor = factor                     # the old failure (1.000) and values that round back
        for v in (1123.45, 7, -3, 0.01, 1, 9, -2.5):
            out = sess.column_value("scale", "金额", v)
            assert out != v and type(out) is type(v), (factor, v, out)
            if factor != 1.0:
                assert (abs(out) > abs(v)) == (factor > 1), (factor, v, out)   # the direction follows the factor
    assert sess.column_value("scale", "金额", 0) == 0 and sess.column_value("scale", "金额", 0.0) == 0.0
