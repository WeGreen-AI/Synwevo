"""Server-computed values must be what Excel shows (the values CSV export, previews and output
verification read), and edge cases must become Excel error values, not failed documents.
Also: value transforms and non-finite numbers.  All inputs are TEST fixtures."""
from __future__ import annotations

import io

import openpyxl
import pytest

from synwevo_docs.engine import transforms
from synwevo_docs.engine import values as V
from synwevo_docs.engine.formula import Book, XLError


def ev(formula: str, cells: dict | None = None, date1904: bool = False):
    b = Book({"S": dict(cells or {})}, {}, date1904)
    b.cells["S"][(99, 0)] = ("f", formula)
    b.dims = {"S": (99, 5)}
    v = b.value("S", 99, 0)
    return v.code if isinstance(v, XLError) else v


# ---------------------------------------------------------------- text functions (F163, F168)
@pytest.mark.parametrize("formula,want", [
    ('=RIGHT("abc",5)', "abc"),              # was "bc"
    ('=RIGHT("abc",2)', "bc"),
    ('=LEFT("abc",-1)', "#VALUE!"),          # was ""
    ('=RIGHT("abc",-1)', "#VALUE!"),
    ('=REPLACE("abcdef",0,2,"X")', "#VALUE!"),   # was "abcdeXbcdef"
    ('=REPLACE("abcdef",2,2,"X")', "aXdef"),
    ('=SUBSTITUTE("aaa","a","b",0)', "#VALUE!"),  # was "aabaaa"
    ('=SUBSTITUTE("aaa","a","b",2)', "aba"),
])
def test_text_functions(formula, want):
    assert ev(formula) == want


# ---------------------------------------------------------------- lookups (F163)
LOOK = {(0, 0): ("v", 5.0), (1, 0): ("v", "5"), (2, 0): ("v", ">1"),
        (0, 1): ("v", "number"), (1, 1): ("v", "text"), (2, 1): ("v", "operator")}


@pytest.mark.parametrize("formula,want", [
    ('=VLOOKUP("5",A1:B3,2,FALSE)', "text"),          # was "number": text "5" is not the number 5
    ('=VLOOKUP(5,A1:B3,2,FALSE)', "number"),
    ('=VLOOKUP(">1",A1:B3,2,FALSE)', "operator"),     # was a comparison, not a lookup of the text
    ('=MATCH("5",A1:A3,0)', 2.0),
    ('=MATCH("op*",B1:B3,0)', 3.0),                   # wildcards still work in exact mode
    ('=VLOOKUP("nope",A1:B3,2,FALSE)', "#N/A"),
])
def test_exact_match_lookups(formula, want):
    assert ev(formula, LOOK) == want


XL = {(0, 0): ("v", 1.0), (1, 0): ("v", 2.0), (2, 0): ("v", 2.0), (3, 0): ("v", 5.0),
      (0, 1): ("v", "a"), (1, 1): ("v", "b"), (2, 1): ("v", "c"), (3, 1): ("v", "d")}


@pytest.mark.parametrize("formula,want", [
    ('=XLOOKUP(2,A1:A4,B1:B4)', "b"),
    ('=XLOOKUP(2,A1:A4,B1:B4,"nf",0,-1)', "c"),       # last-to-first: was the first match
    ('=XLOOKUP(3,A1:A4,B1:B4,"nf",-1)', "b"),         # next smaller: was "nf"
    ('=XLOOKUP(3,A1:A4,B1:B4,"nf",1)', "d"),          # next larger
    ('=XLOOKUP(3,A1:A4,B1:B4,"nf")', "nf"),
    ('=XLOOKUP(2,A1:A4,B1:B4,"nf",7)', "#VALUE!"),
])
def test_xlookup_modes(formula, want):
    assert ev(formula, XL) == want


def test_xlookup_binary_search_is_not_guessed():
    from synwevo_docs.engine.formula import Unsupported
    with pytest.raises(Unsupported):
        ev('=XLOOKUP(2,A1:A4,B1:B4,"nf",0,2)', XL)


# ---------------------------------------------------------------- numeric edge cases (F164, F168)
@pytest.mark.parametrize("formula,want", [
    ("=POWER(0,-1)", "#DIV/0!"),             # was ZeroDivisionError: the whole document failed
    ("=POWER(-8,1/3)", "#NUM!"),             # was a complex number written as text
    ("=1E308*10", "#NUM!"),                  # was <v>inf</v>
    ("=DATE(10000,1,1)", "#NUM!"),           # was ValueError
    ("=TRUNC(1.15,2)", 1.15),                # was 1.14
    ("=ROUND(0.4999999999,0)", 0.0),         # was 1.0
    ("=ROUND(2.675,2)", 2.68),
    ("=ROUND(1234.5,-2)", 1200.0),
    ("=ROUNDUP(1.001,2)", 1.01),
    ("=ROUNDDOWN(-1.999,2)", -1.99),
    ("=ROUND(123456789,20)", 123456789.0),
])
def test_numeric_edge_cases(formula, want):
    got = ev(formula)
    assert got == (pytest.approx(want) if isinstance(want, float) else want)


def test_error_values_propagate_to_dependent_cells():
    assert ev("=A1+1", {(0, 0): ("f", "=POWER(0,-1)")}) == "#DIV/0!"


def test_ifs_ranges_of_different_size_are_value_errors():
    cells = {(0, 0): ("v", 1.0), (1, 0): ("v", 2.0), (2, 0): ("v", 3.0), (0, 1): ("v", "a"), (1, 1): ("v", "a")}
    assert ev('=AVERAGEIFS(A1:A3,B1:B2,"a")', cells) == "#VALUE!"     # was IndexError
    assert ev('=SUMIFS(A1:A3,B1:B2,"a")', cells) == "#VALUE!"
    assert ev('=COUNTIFS(A1:A3,">0",B1:B2,"a")', cells) == "#VALUE!"


def test_index_rejects_negative_and_out_of_range_positions():
    cells = {(0, 0): ("v", 1.0), (1, 0): ("v", 2.0), (2, 0): ("v", 3.0)}
    assert ev("=INDEX(A1:A3,-1)", cells) == "#VALUE!"                 # was the wrong row
    assert ev("=INDEX(A1:A3,4)", cells) == "#REF!"
    assert ev("=INDEX(A1:A3,2)", cells) == 2.0


def test_text_uses_the_workbook_date_system():
    assert ev('=TEXT(DATE(2024,1,5),"yyyy-mm-dd")', date1904=True) == "2024-01-05"   # was 2020-01-04
    assert ev('=TEXT(DATE(2024,1,5),"yyyy-mm-dd")') == "2024-01-05"


def test_filled_workbook_carries_error_values_not_a_failure():
    from synwevo_docs.engine import templates as T
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 0
    ws["B1"] = "=POWER(A1,-1)"
    ws["C1"] = "=B1*2"
    b = io.BytesIO()
    wb.save(b)
    work, slot = T.xlsx_add_slot(b.getvalue(), {"sheet": "Sheet", "kind": "cell", "cell": "A1", "name": "a"}, [])
    out, issues = T.xlsx_fill(work, [slot], {slot["id"]: {"kind": "scalar", "value": V.Value(V.INT, "0", "0")}})
    assert not [i for i in issues if i["level"] == "error"]
    ws2 = openpyxl.load_workbook(io.BytesIO(out), data_only=True).active
    assert ws2["B1"].value == "#DIV/0!" and ws2["C1"].value == "#DIV/0!"


# ---------------------------------------------------------------- transforms (F161, F162)
def test_prefix_and_suffix_survive_later_text_operations():
    v = transforms.apply(V.text("abc"), [{"op": "prefix", "text": "No."}, {"op": "upper"}])
    assert v.text == "NO.ABC"                                        # was "ABC"
    v = transforms.apply(V.text(" 12 "), [{"op": "suffix", "text": "元"}, {"op": "trim"}])
    assert v.text == "12 元"                                          # was "12"
    d = transforms.apply(V.Value(V.DEC, "12.50", "12.50"), [{"op": "prefix", "text": "¥"}])
    assert d.t == V.DEC and d.v == "12.50" and d.text == "¥12.50"


@pytest.mark.parametrize("ops", [
    [{"op": "pad_left", "width": 10 ** 9}],               # gigabytes per record before
    [{"op": "pad_left", "width": 6, "char": "00"}],
    [{"op": "map", "table": [1]}],                        # TypeError before
    [{"op": "decimals", "n": 2, "thousands": "yes"}],
    [{"op": "prefix", "text": "x" * 1000}],
    [{"op": "trim", "unknown": 1}],                       # unknown options were accepted
    [{"op": "map", "table": {"a": {"nested": 1}}}],
    ["trim"],
])
def test_bad_transform_parameters_are_refused(ops):
    assert transforms.validate(ops)
    with pytest.raises(transforms.TransformError):
        transforms.apply(V.text("1"), ops)


def test_valid_transforms_still_pass_validation():
    ops = [{"op": "trim"}, {"op": "decimals", "n": 2, "thousands": True}, {"op": "pad_left", "width": 6, "char": "0"},
           {"op": "map", "table": {"TRUE": "是"}}, {"op": "date_format", "pattern": "yyyy-mm-dd"}, {"op": "prefix", "text": "No."}]
    assert transforms.validate(ops) == []


def test_decimals_on_long_and_non_finite_numbers():
    v = transforms.apply(V.text("123456789012345678901234567"), [{"op": "decimals", "n": 2}])
    assert v.text == "123456789012345678901234567.00"                # was decimal.InvalidOperation
    warns: list[str] = []
    v = transforms.apply(V.text("Infinity"), [{"op": "decimals", "n": 2}], warns)
    assert v.text == "Infinity" and warns                             # a warning, not InvalidOperation


# ---------------------------------------------------------------- non-finite numbers (F169)
def test_nan_and_infinity_text_is_not_a_number():
    assert V.text("NaN").as_decimal() is None and V.text("Infinity").as_decimal() is None
    assert V.text("sNaN").as_decimal() is None and V.text("1,234.5").as_decimal() is not None
    with pytest.raises(transforms.TransformError):
        transforms.aggregate([V.text("1"), V.text("NaN")], "sum")     # was the text "NaN" in a money field
