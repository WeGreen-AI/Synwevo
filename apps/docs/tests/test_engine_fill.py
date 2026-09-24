"""Generated documents are what the template and the data say — read back from the output.

Every template and value here is a TEST fixture built in memory.
"""
from __future__ import annotations

import io
import zipfile

import docx
import openpyxl
import pytest
from lxml import etree

from synwevo_docs.engine import docx_blocks as db
from synwevo_docs.engine import docx_core as dc
from synwevo_docs.engine import generate as G
from synwevo_docs.engine import pdf_engine as pe
from synwevo_docs.engine import templates as T
from synwevo_docs.engine import values as V
from synwevo_docs.engine.ooxml import Package
from synwevo_docs.engine.plan import Plan

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
soffice = pytest.mark.skipif(not __import__("synwevo_docs.engine.converters", fromlist=["x"]).soffice_path(),
                             reason="需要 LibreOffice")


def _rows(*vals):
    return {"kind": "rows", "rows": [{"c0": V.text(v)} for v in vals]}


# ---------------------------------------------------------------- DOCX repeat blocks (F177)
def _block_template(header: bool, footer: bool = False) -> tuple[bytes, dict]:
    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = "信头 TEST"
    if footer:
        d.sections[0].footer.paragraphs[0].text = "页脚 TEST"
    for t in ("前", "品名：样例", "后"):
        d.add_paragraph(t)
    b = io.BytesIO()
    d.save(b)
    pkg = Package(b.getvalue())
    main = pkg.main_part()
    sdt = db.wrap_blocks(dc.part_container(pkg, main), "1", "1", dc.TAG_PREFIX + "blk1", "明细")
    cols = db.auto_columns(sdt, "blk1")
    slot = {"id": "blk1", "binding": "blk1", "name": "明细", "kind": "repeat", "block": True, "columns": cols,
            "separator": "none", "empty_policy": "remove", "locator": {"part": main, "block": True}}
    return pkg.save(), slot


@pytest.mark.parametrize("header,footer", [(False, False), (True, False), (True, True)])
def test_repeat_block_with_headers_and_footers_fills_without_false_errors(header, footer):
    work, slot = _block_template(header, footer)
    key = slot["columns"][0]["key"]
    out, issues = T.docx_fill(work, [slot], {"blk1": {"kind": "rows", "rows": [{key: V.text("甲")}, {key: V.text("乙")}]}})
    assert [i for i in issues if i["level"] == "error"] == []          # was block_lost once per header / footer
    texts = [p.text for p in docx.Document(io.BytesIO(out)).paragraphs]
    assert texts == ["前", "品名：甲", "品名：乙", "后"]


def test_a_repeat_missing_from_every_part_is_still_reported():
    work, slot = _block_template(True)
    gone = {**slot, "id": "nope"}
    _, issues = T.docx_fill(work, [slot, gone], {"blk1": {"kind": "rows", "rows": []}})
    assert [(i["slot"], i["code"]) for i in issues if i["level"] == "error"] == [("nope", "block_lost")]
    row = {"id": "row_gone", "name": "明细行", "kind": "repeat", "columns": [], "locator": {}}
    _, issues = T.docx_fill(work, [row], {})
    assert [i["code"] for i in issues] == ["repeat_lost"]               # was silently skipped


def test_detail_rows_keep_the_template_row_properties():
    d = docx.Document()
    t = d.add_table(rows=2, cols=1)
    t.cell(0, 0).text = "品名"
    t.cell(1, 0).text = "样例"
    b = io.BytesIO()
    d.save(b)
    work, slot = T.docx_add_slot(b.getvalue(), {"kind": "repeat", "table_path": "0", "row": 1, "name": "明细"}, [])
    out, _ = T.docx_fill(work, [slot], {slot["id"]: _rows("甲", "乙")})
    rows = list(Package(out).xml("word/document.xml").iter(dc.w("tr")))
    assert len(rows) == 3 and not any(r.find(f"{dc.w('trPr')}/{dc.w('cantSplit')}") is not None for r in rows)


# ---------------------------------------------------------------- fit-mode output (F179, F178)
def test_strip_marks_keeps_every_prefix_markup_compatibility_names():
    from synwevo_docs.engine import layout as L
    d = docx.Document()
    d.add_paragraph("x")
    b = io.BytesIO()
    d.save(b)
    for unwrap in (True, False):
        out = L._strip_marks(b.getvalue(), unwrap)
        root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml"))
        ign = root.get("{%s}Ignorable" % MC)
        assert ign and [p for p in ign.split() if p not in root.nsmap] == []    # was ['w14', 'wp14']


def _clipping_docx() -> bytes:
    """TEST fixture: a field inside an exact-height table row, filled with far too much text."""
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Pt
    d = docx.Document()
    d.add_paragraph("标题")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].height = Pt(14)
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    t.cell(0, 0).paragraphs[0].text = "占位"
    d.add_paragraph("结尾")
    b = io.BytesIO()
    d.save(b)
    pkg = Package(b.getvalue())
    p = list(dc.part_container(pkg, pkg.main_part()).iter(dc.w("tc")))[0].find(dc.w("p"))
    sdt = dc.wrap_range(p, 0, 2, dc.TAG_PREFIX + "s1", "备注")
    dc.fill_sdt(sdt, "很长的内容第" + "".join(f"{i}段" for i in range(1, 80)) + "完结", True)
    return pkg.save()


@soffice
def test_fit_mode_detects_text_clipped_by_an_exact_height_row():
    from synwevo_docs.engine import layout as L
    out, rep = L.fit_docx(_clipping_docx(), [{"id": "s1", "name": "备注", "shrink": True, "shrink_min": 7}], 1, False)
    assert rep["pages"] == 1
    assert rep["truncated_fields"] + rep["lost_fields"] == ["备注"]     # was always [] (tautological check)


# ---------------------------------------------------------------- pagination policy (F157 / F185)
def test_fit_without_a_template_page_count_is_not_a_silent_success():
    d = docx.Document()
    d.add_paragraph("x")
    b = io.BytesIO()
    d.save(b)
    _, issues, lay = G.fill_checked("docx", b.getvalue(), [], Plan(record={}), {}, "fit", None)
    assert lay is None and [(i["level"], i["code"]) for i in issues] == [("error", "layout_unavailable")]
    _, issues, _ = G.fill_checked("docx", b.getvalue(), [], Plan(record={}), {}, "warn", None)
    assert [(i["level"], i["code"]) for i in issues] == [("warning", "layout_unavailable")]
    _, issues, _ = G.fill_checked("docx", b.getvalue(), [], Plan(record={}), {}, "off", None)
    assert issues == []


# ---------------------------------------------------------------- CSV export (F159 / F186 / F276, F160)
def _xlsx(cells: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for ref, v in cells.items():
        ws[ref] = v
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_csv_export_neutralises_formula_text_but_not_numbers():
    work = _xlsx({"A1": "名称", "B1": "金额"})
    work, a2 = T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "cell", "cell": "A2", "name": "名称"}, [])
    work, b2 = T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "cell", "cell": "B2", "name": "金额"}, [a2])
    evil = '=HYPERLINK("http://example.invalid","click")'
    out, _ = T.xlsx_fill(work, [a2, b2], {a2["id"]: {"kind": "scalar", "value": V.text(evil)},
                                           b2["id"]: {"kind": "scalar", "value": V.Value(V.INT, "-5", "-5")}})
    text = G.xlsx_to_csv(out, []).decode("utf-8-sig")
    line = text.splitlines()[1]
    assert line.startswith("\"'=HYPERLINK(") and line.endswith(",-5")
    for bad in ("+1", "-x", "@SUM(A1)", "\tx"):
        assert G.neutralize_csv_cell(bad).startswith("'")
    assert G.neutralize_csv_cell("正常") == "正常"


def test_formula_whose_result_is_an_empty_string_exports_to_csv():
    work = _xlsx({"A1": "x", "B1": '=IF(A1="x","",A1)'})
    work, a1 = T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "cell", "cell": "A1", "name": "a"}, [])
    out, _ = T.xlsx_fill(work, [a1], {a1["id"]: {"kind": "scalar", "value": V.text("x")}})
    assert G.xlsx_to_csv(out, []).decode("utf-8-sig").splitlines()[0] == "x"   # was TemplateError formula_uncomputed


# ---------------------------------------------------------------- XLSX structure (F172, F191, F192)
def test_region_on_a_row_removed_by_an_empty_detail_row_is_reported():
    work = _xlsx({"A1": "品名", "A2": "样例", "B2": "区域", "A3": "尾"})
    work, rep = T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "repeat", "row": 1, "c1": 0, "c2": 0, "name": "明细"}, [])
    work, reg = T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "region", "range": "B2:B2", "name": "区域"}, [rep])
    out, issues = T.xlsx_fill(work, [rep, reg], {rep["id"]: {"kind": "rows", "rows": []},
                                                 reg["id"]: {"kind": "matrix", "rows": [[V.text("v")]]}})
    assert [(i["slot"], i["code"]) for i in issues if i["level"] == "error"] == [(reg["id"], "slot_removed")]   # was a TypeError


def test_calcpr_is_inserted_in_schema_order():
    from synwevo_docs.engine.xlsx_doc import XlsxDoc, s
    work = _xlsx({"A1": 1, "B1": "=A1+1"})
    pkg = Package(work)
    wbx = pkg.xml("xl/workbook.xml")
    for tag in ("calcPr", "definedNames"):
        el = wbx.find(s(tag))
        if el is not None:
            wbx.remove(el)
    etree.SubElement(wbx, s("pivotCaches"))           # after calcPr in CT_Workbook
    x = XlsxDoc(pkg.save())
    x.touched = True
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(x.finalize())).read("xl/workbook.xml"))
    order = [etree.QName(c).localname for c in root]
    assert order.index("calcPr") < order.index("pivotCaches")          # was placed after pivotCaches


def test_array_formula_copies_point_at_their_own_row():
    from synwevo_docs.engine.xlsx_doc import XlsxDoc, s
    work = _xlsx({"A1": "h", "A2": 1, "B2": 2})
    pkg = Package(work)
    sh = pkg.xml("xl/worksheets/sheet1.xml")
    for c in sh.iter(s("c")):
        if c.get("r") == "B2":
            for v in list(c):
                c.remove(v)
            f = etree.SubElement(c, s("f"), {"t": "array", "ref": "C2"})
            f.text = "A2*2"
    x = XlsxDoc(pkg.save())
    for c in x.sheet_root("Sheet").iter(s("c")):
        if c.get("r") == "B2":
            c.set("r", "C2")
    x.insert_copies_of_row("Sheet", 2, 2)
    refs = {c.get("r"): c.find(s("f")).get("ref") for c in x.sheet_root("Sheet").iter(s("c")) if c.find(s("f")) is not None}
    assert refs == {"C2": "C2", "C3": "C3", "C4": "C4"}                 # copies kept ref="C2"


def test_many_detail_rows_fill_in_linear_time():
    import time
    work = _xlsx({"A1": "a", "B1": "b", "C1": "c", "A2": "x"})
    work, rep = T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "repeat", "row": 1, "c1": 0, "c2": 2, "name": "明细"}, [])
    rows = [{f"c{i}": V.text(f"{r}-{i}") for i in range(3)} for r in range(8000)]
    t = time.perf_counter()
    out, _ = T.xlsx_fill(work, [rep], {rep["id"]: {"kind": "rows", "rows": rows}})
    # measured: 0.3 s now; 20.9 s when every written cell rescanned all rows
    assert time.perf_counter() - t < 5
    ws = openpyxl.load_workbook(io.BytesIO(out)).active
    assert ws["A2"].value == "0-0" and ws["C8001"].value == "7999-2" and ws.max_row == 8001


# ---------------------------------------------------------------- slot options (F171)
def test_unknown_slot_options_are_refused_when_the_field_is_made():
    work = _xlsx({"A1": "x"})
    with pytest.raises(T.TemplateError) as e:
        T.xlsx_add_slot(work, {"sheet": "Sheet", "kind": "cell", "cell": "A1", "name": "a", "mode": "insert-after"}, [])
    assert e.value.code == "bad_value"                                  # was silently "replace"
    d = docx.Document()
    d.add_paragraph("示例文字")
    b = io.BytesIO()
    d.save(b)
    with pytest.raises(T.TemplateError):
        T.docx_add_slot(b.getvalue(), {"kind": "text", "path": "0", "start": 0, "end": 2, "name": "a", "mode": "append"}, [])


# ---------------------------------------------------------------- control characters (F182)
def test_docx_value_with_xml_illegal_characters_is_written():
    d = docx.Document()
    d.add_paragraph("备注：样例")
    b = io.BytesIO()
    d.save(b)
    work, slot = T.docx_add_slot(b.getvalue(), {"kind": "text", "path": "0", "start": 3, "end": 5, "name": "备注"}, [])
    out, issues = T.docx_fill(work, [slot], {slot["id"]: {"kind": "scalar", "value": V.text("第一行\x0b第二行\x00尾")}})
    p = docx.Document(io.BytesIO(out)).paragraphs[0]
    assert p.text == "备注：第一行\n第二行尾"                                # was ValueError from lxml
    assert [i["code"] for i in issues] == ["control_chars"]


# ---------------------------------------------------------------- PDF boxes and form fields (F202, F206, F207)
def test_box_whose_size_is_below_the_default_minimum_still_fills():
    pymupdf = pytest.importorskip("pymupdf")
    d = pymupdf.open()
    d.new_page(width=300, height=200)
    slot = {"id": "s", "kind": "pdf_box", "name": "小字", "locator": {"page": 0, "rect": [10, 10, 200, 30]},
            "font_size": 5.5, "min_font_size": 6}
    out, issues = pe.fill_pdf(d.tobytes(), [slot], {"s": "ABC"})
    assert not [i for i in issues if i["level"] == "error"]            # was overflow for every record
    _, issues = pe.fill_pdf(d.tobytes(), [slot], {"s": ""})
    assert not [i for i in issues if i["level"] == "error"]


def _form_pdf() -> bytes:
    pymupdf = pytest.importorskip("pymupdf")
    d = pymupdf.open()
    for pno in range(2):
        p = d.new_page(width=300, height=200)
        w = pymupdf.Widget()
        w.field_name, w.field_type = "name", pymupdf.PDF_WIDGET_TYPE_TEXT
        w.rect = pymupdf.Rect(20, 20, 200, 40)
        w.field_value = "样例"
        p.add_widget(w)
        if pno == 0:
            cb = pymupdf.Widget()
            cb.field_name, cb.field_type = "agree", pymupdf.PDF_WIDGET_TYPE_CHECKBOX
            cb.rect = pymupdf.Rect(20, 60, 35, 75)
            p.add_widget(cb)
    return d.tobytes()


def test_pdf_field_fills_every_widget_and_refuses_non_text_fields():
    pymupdf = pytest.importorskip("pymupdf")
    if pe.ENGINE != "pymupdf":
        pytest.skip("PDF 模板需要 PyMuPDF")
    data = _form_pdf()
    slot = {"id": "n", "kind": "pdf_field", "name": "姓名", "locator": {"page": 0, "field": "name"}, "font_size": 10}
    out, issues = pe.fill_pdf(data, [slot], {"n": "TEST-VALUE"})
    doc = pymupdf.open(stream=out, filetype="pdf")
    assert all("TEST-VALUE" in doc[i].get_text() for i in range(2))     # page 2 kept the sample before
    assert not any(w.field_name == "name" for pg in doc for w in (pg.widgets() or []))
    cb = {"id": "c", "kind": "pdf_field", "name": "同意", "locator": {"page": 0, "field": "agree"}}
    _, issues = pe.fill_pdf(data, [cb], {"c": "是"})
    assert [i["code"] for i in issues] == ["field_type_unsupported"]   # was drawn as text over a check box
    with pytest.raises(T.TemplateError):
        T.pdf_add_slot(data, {"kind": "pdf_field", "field": "agree"}, [])


def test_box_over_an_ocr_text_layer_defaults_to_clearing_pixels():
    pymupdf = pytest.importorskip("pymupdf")
    if pe.ENGINE != "pymupdf":
        pytest.skip("PDF 模板需要 PyMuPDF")
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (400, 200), "white")
    ImageDraw.Draw(im).rectangle([40, 40, 200, 60], fill="black")        # the visible "text" is pixels
    b = io.BytesIO()
    im.save(b, "PNG")
    d = pymupdf.open()
    p = d.new_page(width=400, height=200)
    p.insert_image(p.rect, stream=b.getvalue())
    p.insert_text((40, 58), "INVISIBLE OCR LAYER TEXT HERE, LONG ENOUGH", render_mode=3)
    assert pe.page_kind(p) == "text"
    slot = T.pdf_add_slot(d.tobytes(), {"kind": "pdf_box", "page": 0, "rect": [35, 35, 300, 65], "name": "x"}, [])
    assert slot["clear"] == "text_pixels"                              # was "text": only the hidden layer went


# ---------------------------------------------------------------- raster image templates (F197)
def test_raster_fill_clears_the_template_sample_under_an_empty_value():
    from PIL import Image, ImageDraw
    from synwevo_docs.engine import fonts
    from synwevo_docs.engine.image_raster import fill
    if not fonts.cjk_font_file():
        pytest.skip("本机没有可用的中文字体文件")
    meta = {"native": "png", "width": 200, "height": 60, "mode": "RGB", "dpi": [72, 72], "alpha": False}
    im = Image.new("RGB", (200, 60), "white")
    ImageDraw.Draw(im).rectangle([20, 20, 120, 35], fill="black")      # the template's sample content
    buf = io.BytesIO()
    im.save(buf, "PNG")
    slot = {"id": "s1", "name": "姓名", "locator": {"rect": [10, 10, 150, 45]}, "clear": "text_pixels",
            "font_size": 12, "color": "#000000"}
    out, issues = fill(buf.getvalue(), [slot], {"s1": {"kind": "scalar", "value": V.empty()}}, meta)
    assert issues == []
    dark = sum(1 for px in Image.open(io.BytesIO(out)).convert("L").crop((10, 10, 150, 45)).getdata() if px < 128)
    assert dark == 0                                                  # was the untouched black sample


# ---------------------------------------------------------------- images without PyMuPDF (F203)
def test_output_image_to_pdf_without_pymupdf(monkeypatch):
    import sys
    from PIL import Image
    from synwevo_docs.engine import image_template
    monkeypatch.setitem(sys.modules, "pymupdf", None)                  # an AGPL-free deployment
    b = io.BytesIO()
    Image.new("RGB", (120, 80), "white").save(b, "PNG", dpi=(150, 150))
    pdf = image_template.to_pdf(b.getvalue())
    assert pdf.startswith(b"%PDF")                                    # was ModuleNotFoundError


# ---------------------------------------------------------------- PPTX render (F210)
def test_render_pptx_can_skip_picture_payloads():
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from synwevo_docs.engine.pptx_engine import render_pptx
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    img = io.BytesIO()
    Image.new("RGB", (40, 40), "red").save(img, "PNG")
    img.seek(0)
    s.shapes.add_picture(img, Inches(1), Inches(1))
    b = io.BytesIO()
    prs.save(b)
    pics = lambda r: [sh for sl in r["slides"] for sh in sl["shapes"] if sh["kind"] == "picture"]  # noqa: E731
    assert pics(render_pptx(b.getvalue()))[0]["img"].startswith("data:image/png")
    assert pics(render_pptx(b.getvalue(), images=False))[0]["img"] is None


# ---------------------------------------------------------------- recalculation fallback (F183)
@soffice
def test_libreoffice_fallback_only_touches_cells_the_engine_could_not_compute():
    import datetime as dt
    from synwevo_docs.engine.xlsx_cache import fill_cached_values
    from synwevo_docs.engine.xlsx_doc import XlsxDoc, s
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = dt.date(2026, 1, 5)
    ws["A1"].number_format = "yyyy-mm-dd"
    ws["B1"] = "=A1+1"
    ws["B1"].number_format = "yyyy-mm-dd"
    ws["C1"] = "=1/0"
    ws["D1"] = "=ROMAN(4)"                                           # not in the built-in engine
    b = io.BytesIO()
    wb.save(b)
    x = XlsxDoc(b.getvalue())
    from synwevo_docs.engine.formula import compute_workbook
    rep = compute_workbook(x)
    assert [r for _, r, _ in rep["unsupported"]] == ["D1"]
    data, _ = fill_cached_values(x.pkg.save(), {(sh, r) for sh, r, _ in rep["unsupported"]})
    cells = {c.get("r"): c for c in XlsxDoc(data).sheet_root("Sheet").iter(s("c"))}
    assert cells["B1"].findtext(s("v")) == "46028" and cells["B1"].get("t") is None   # was a "2026-01-06 00:00:00" string
    assert cells["C1"].get("t") == "e" and cells["C1"].findtext(s("v")) == "#DIV/0!"
    assert cells["D1"].findtext(s("v")) == "IV"
