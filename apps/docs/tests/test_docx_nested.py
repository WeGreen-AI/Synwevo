"""Nested details (docs-nested-details.md): a repeat content block whose mapping sets 分段字段
(`group_by`) gets one copy per sub-group of the detail rows, and a detail table inside the block
gets, in each copy, only the rows of that copy's sub-group.

TEST FIXTURES: the workbook and the Word template below are synthetic, built here; every expected
value is computed from the rows written into the workbook, not read back from the engine."""
from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import docx
import openpyxl
import pytest

from .helpers import _walk

# 客户 × 订单 × 明细, deliberately out of order (the rule's sort puts them in order) and with an
# order's lines not adjacent (sub-groups are not required to be contiguous)
ROWS = [
    # 客户, 订单号, 下单日期, 行号, 品名, 数量, 单价
    ("A", "A-2", dt.date(2024, 3, 9), 1, "垫圈", 10, Decimal("0.50")),
    ("B", "B-1", dt.date(2024, 3, 2), 1, "螺栓", 4, Decimal("3.25")),
    ("A", "A-1", dt.date(2024, 3, 1), 2, "螺母", 3, Decimal("1.20")),
    ("A", "A-1", dt.date(2024, 3, 1), 1, "螺丝", 5, Decimal("2.00")),
]


def _workbook(path, rows=ROWS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["客户", "订单号", "下单日期", "行号", "品名", "数量", "单价", "金额", "日期文本"])
    for i, r in enumerate(rows, start=2):
        lined = r[5] is not None
        ws.append([*r[:5], r[5], float(r[6]) if lined else None, f"=F{i}*G{i}" if lined else None,
                   f'=TEXT(C{i},"yyyy-mm-dd")' if r[2] else None])
    wb.save(path)
    return path


def _template(path):
    d = docx.Document()
    d.add_paragraph("客户：C-0")
    d.add_paragraph("订单号：O-0　下单日期：2024-01-01")
    t = d.add_table(rows=2, cols=4)
    for c, h in enumerate(("品名", "数量", "单价", "金额")):
        t.cell(0, c).text = h
    for c, v in enumerate(("示例", "1", "0.00", "0.00")):
        t.cell(1, c).text = v
    d.add_paragraph("订单小计：0.00")
    d.add_paragraph("（订单结束）")
    d.save(path)
    return path


def _body(api, tid):
    v = api.view(tid)
    return next(p for p in v["doc"]["parts"] if p["kind"] == "body")


def _setup(api, tmp_path, rows=ROWS, row_first=False, block=None, row=None):
    """The template's customer field, the 订单段 block over the order paragraphs and the table,
    and the 明细 repeat row inside it (made in either order)."""
    pid = api.project("嵌套明细")
    api.upload(pid, _workbook(tmp_path / "订单.xlsx", rows))
    tid = api.template(pid, _template(tmp_path / "订单模板.docx"))["id"]
    body = _body(api, tid)
    ps = {b["text"]: b for b in _walk(body["blocks"])}
    first, last = ps["订单号：O-0　下单日期：2024-01-01"], ps["订单小计：0.00"]
    tbl = next(b for b in body["blocks"] if b["t"] != "p")
    make_block = lambda: api.slot(tid, kind="repeat_block", part=first["part"], first=first["path"], last=last["path"],  # noqa: E731
                                  name="订单段", **(block or {}))
    make_row = lambda: api.slot(tid, kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="明细",  # noqa: E731
                                header_rows=1, **(row or {}))
    if row_first:
        rep = make_row()
        blk = make_block()
    else:
        blk = make_block()
        rep = make_row()
    cust = api.text_slot(tid, "C-0", "客户")
    return pid, tid, blk, rep, cust


TABLE = {"kind": "table", "sheet": "明细", "header_row": 0}
GROUP_BY = {"keys": [{"header": "订单号"}]}


def _map(api, pid, blk, rep, cust, *, group_by=GROUP_BY, block_cols=None, child_cols=None, child=True, sort=None):
    api.rule(pid, {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                   "group": {"keys": [{"header": "客户"}],
                             "sort": sort or [{"header": "订单号"}, {"header": "行号", "numeric": True}]}})
    bk = {c["name"]: c["key"] for c in blk["columns"]}
    rk = {c["name"]: c["key"] for c in rep["columns"]}
    two = [{"op": "decimals", "n": 2}]
    block_cols = block_cols or [
        {"slot_col": bk["订单号"], "source": {"header": "订单号"}},
        {"slot_col": bk["下单日期"], "source": {"header": "日期文本"}},
        {"slot_col": bk["订单小计"], "source": {"header": "金额"}, "agg": "sum", "transforms": two},
    ]
    child_cols = child_cols or [
        {"slot_col": rk["品名"], "source": {"header": "品名"}},
        {"slot_col": rk["数量"], "source": {"header": "数量"}},
        {"slot_col": rk["单价"], "source": {"header": "单价"}, "transforms": two},
        {"slot_col": rk["金额"], "source": {"header": "金额"}, "transforms": two},
    ]
    blockmap = {"type": "repeat", "sources": [TABLE], "target": {"binding": blk["binding"]}, "columns": block_cols}
    if group_by is not None:
        blockmap["group_by"] = group_by
    maps = [{"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "客户"}], "target": {"binding": cust["binding"]}},
            blockmap]
    if child:
        maps.append({"type": "repeat", "sources": [TABLE], "target": {"binding": rep["binding"]}, "columns": child_cols})
    return api.maps(pid, maps)


def _flat(data: bytes) -> list:
    """The document body in order: a paragraph's text, or a table as its rows of cell texts."""
    d = docx.Document(io.BytesIO(data))
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append("".join(t.text or "" for t in el.iter(f"{{{W}}}t")))
        elif tag == "tbl":
            out.append([["".join(t.text or "" for t in tc.iter(f"{{{W}}}t")) for tc in tr.iter(f"{{{W}}}tc")]
                        for tr in el.iter(f"{{{W}}}tr")])
    return out


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HEAD = ["品名", "数量", "单价", "金额"]


def _money(x) -> str:
    return f"{Decimal(str(x)):.2f}"


def _expected(customer: str, rows=ROWS) -> list:
    """What the document for one customer must read, computed from the fixture rows."""
    mine = sorted((r for r in rows if r[0] == customer), key=lambda r: (r[1], r[3]))
    out: list = [f"客户：{customer}"]
    orders: dict[str, list] = {}
    for r in mine:
        orders.setdefault(r[1], []).append(r)
    for no, lines in orders.items():
        out.append(f"订单号：{no}　下单日期：{lines[0][2].isoformat()}")
        out.append([HEAD] + [[r[4], str(r[5]), _money(r[6]), _money(r[5] * r[6])] for r in lines])
        out.append(f"订单小计：{_money(sum(r[5] * r[6] for r in lines))}")
    out.append("（订单结束）")
    return out


def test_each_order_copy_gets_only_its_own_lines(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    tpl = api.c.get(f"/api/templates/{tid}").json()
    assert {s["id"]: s.get("parent") for s in tpl["version"]["slots"]}[rep["id"]] == blk["id"]
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed", api.items(job["id"])
    outs = api.outputs(job["id"])
    assert _flat(outs["A.docx"]) == _expected("A")
    assert _flat(outs["B.docx"]) == _expected("B")
    # every copy's table keeps its header row marked to repeat on each page
    d = docx.Document(io.BytesIO(outs["A.docx"]))
    assert len(d.tables) == 2
    for t in d.tables:
        assert t.rows[0]._tr.find(f"{{{W}}}trPr/{{{W}}}tblHeader") is not None
    assert not d.element.body.findall(f".//{{{W}}}sdt")      # no field controls left


def _preview_issues(api, pid, key) -> list[dict]:
    return api.ok(api.c.get(f"/api/projects/{pid}/preview", params={"search": key}))["issues"]


def test_the_row_may_be_made_first_and_removing_the_block_clears_parent(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path, row_first=True)
    slots = {s["id"]: s for s in api.c.get(f"/api/templates/{tid}").json()["version"]["slots"]}
    assert slots[rep["id"]].get("parent") == blk["id"] and "parent" not in slots[blk["id"]]
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed"
    assert _flat(api.outputs(job["id"])["A.docx"]) == _expected("A")
    api.ok(api.c.delete(f"/api/templates/{tid}/slots/{blk['id']}"))
    slots = {s["id"]: s for s in api.c.get(f"/api/templates/{tid}").json()["version"]["slots"]}
    assert blk["id"] not in slots and "parent" not in slots[rep["id"]]


# an order row with no lines: only 订单号 and 下单日期 are filled in
EMPTY_ORDER = ("A", "A-3", dt.date(2024, 3, 12), None, None, None, None)


@pytest.mark.parametrize("policy", ["remove_row", "placeholder", "remove_table"])
def test_an_order_without_lines_follows_the_tables_empty_policy(api, tmp_path, policy):
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=ROWS + [EMPTY_ORDER], row={"empty_policy": policy})
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed", [(i["status"], i.get("issues")) for i in api.items(job["id"])]
    got = _flat(api.outputs(job["id"])["A.docx"])
    want = _expected("A")[:-1]            # the two orders with lines, then A-3
    tail = got[len(want):]
    assert got[:len(want)] == want
    assert tail[0] == "订单号：A-3　下单日期：2024-03-12"
    if policy == "remove_row":
        assert tail[1] == [HEAD]                                  # the header stays, no line
    elif policy == "placeholder":
        assert tail[1] == [HEAD, ["（无）", "", "", ""]]
    else:
        assert not isinstance(tail[1], list)                      # no table in this copy
    assert tail[-2:] == ["订单小计：", "（订单结束）"]              # nothing to sum: left empty, not 0.00


def test_an_order_without_lines_is_an_error_when_the_table_requires_rows(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=ROWS + [EMPTY_ORDER], row={"empty_policy": "error"})
    _map(api, pid, blk, rep, cust)
    issues = _preview_issues(api, pid, "A")
    nd = [i for i in issues if i["code"] == "no_detail"]
    assert len(nd) == 1 and nd[0]["level"] == "error" and "A-3" in nd[0]["message"] and nd[0]["slot"] == rep["id"]
    assert not [i for i in _preview_issues(api, pid, "B") if i["level"] == "error"]


NO_KEY = ("C", None, dt.date(2024, 3, 5), 1, "扳手", 1, Decimal("9.90"))


def test_a_line_without_an_order_number_is_an_error_unless_told_to_skip(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=ROWS + [NO_KEY])
    _map(api, pid, blk, rep, cust)
    e = [i for i in _preview_issues(api, pid, "C") if i["code"] == "subgroup_key_empty"]
    assert len(e) == 1 and e[0]["level"] == "error" and "1 行明细的分段字段为空：明细!B6。" in e[0]["message"]   # the cell is named
    r = api.job(pid, naming="{客户}")
    assert r.status_code == 409 and [x["record"] for x in r.json()["detail"]["invalid"]] == ["C"]
    # skip: the line is left out with a warning; customer C has no order and the block is removed
    _map(api, pid, blk, rep, cust, group_by={**GROUP_BY, "empty_key": "skip"})
    issues = _preview_issues(api, pid, "C")
    assert not [i for i in issues if i["level"] == "error"]
    assert any(i["level"] == "warning" and "分段字段为空" in i["message"] for i in issues)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed"
    assert _flat(api.outputs(job["id"])["C.docx"]) == ["客户：C", "（订单结束）"]


def test_a_missing_sub_group_column_is_named(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    _map(api, pid, blk, rep, cust, group_by={"keys": [{"header": "订单编号"}]})
    e = [i for i in _preview_issues(api, pid, "A") if i["level"] == "error"]
    assert [i["code"] for i in e] == ["missing_column"] and "订单编号" in e[0]["message"]


def test_two_values_in_one_order_are_a_conflict_until_a_rule_is_chosen(api, tmp_path):
    rows = [r if r[3] != 2 else (*r[:2], dt.date(2024, 3, 3), *r[3:]) for r in ROWS]   # A-1's second line: another date
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=rows)
    _map(api, pid, blk, rep, cust)
    e = [i for i in _preview_issues(api, pid, "A") if i["level"] == "error"]
    assert [i["code"] for i in e] == ["group_conflict"]
    assert "A-1" in e[0]["message"] and "日期文本" in e[0]["message"] and set(e[0]["values"]) == {"2024-03-01", "2024-03-03"}
    bk = {c["name"]: c["key"] for c in blk["columns"]}
    _map(api, pid, blk, rep, cust, block_cols=[
        {"slot_col": bk["订单号"], "source": {"header": "订单号"}},
        {"slot_col": bk["下单日期"], "source": {"header": "日期文本"}, "agg": "first"},   # after the sort: line 1
        {"slot_col": bk["订单小计"], "source": {"header": "金额"}, "agg": "sum", "transforms": [{"op": "decimals", "n": 2}]},
    ])
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed"
    assert _flat(api.outputs(job["id"])["A.docx"]) == _expected("A")


def test_a_nested_table_needs_the_block_to_be_split(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    _map(api, pid, blk, rep, cust, group_by=None)
    e = [i for i in _preview_issues(api, pid, "A") if i["level"] == "error"]
    assert [i["code"] for i in e] == ["nested_needs_group_by"] and e[0]["slot"] == rep["id"]
    r = api.job(pid, naming="{客户}")                     # refused up front, not a wrong document
    assert r.status_code == 409 and r.json()["error"] == "has_errors"
    assert {i["code"] for rec in r.json()["detail"]["invalid"] for i in rec["issues"]} == {"nested_needs_group_by"}


def test_a_split_block_whose_table_is_not_mapped_leaves_each_table_empty(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    _map(api, pid, blk, rep, cust, child=False)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed", api.items(job["id"])
    got = _flat(api.outputs(job["id"])["A.docx"])
    assert [x for x in got if isinstance(x, list)] == [[HEAD], [HEAD]]


def test_group_by_is_refused_off_a_word_block(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    _map(api, pid, blk, rep, cust)
    maps = api.ok(api.c.get(f"/api/projects/{pid}/mappings"))["mappings"]
    for m in maps:
        if (m.get("target") or {}).get("binding") == rep["binding"]:
            m["group_by"] = GROUP_BY
    api.maps(pid, maps)
    e = [i for i in _preview_issues(api, pid, "A") if i["code"] == "group_by_unsupported"]
    assert len(e) == 1 and e[0]["level"] == "error"
    # the same check for a spreadsheet / slide / PDF repeat (no `block`)
    from synwevo_docs.engine.plan import static_check
    xl = {"id": "s1", "binding": "s1", "kind": "repeat", "name": "明细", "columns": []}
    issues = static_check([{"id": "m1", "no": 1, "type": "repeat", "target": {"binding": "s1"}, "group_by": GROUP_BY, "sources": [TABLE]}],
                          [xl], {"mode": "per_group"})
    assert [i["code"] for i in issues] == ["group_by_unsupported"]


def test_kept_field_controls_have_unique_ids(api, tmp_path):
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}", keep_fields=True)
    assert job["status"] == "completed"
    data = api.outputs(job["id"])["A.docx"]
    d = docx.Document(io.BytesIO(data))
    ids = [e.get(f"{{{W}}}val") for e in d.element.body.iter(f"{{{W}}}id") if e.getparent().tag == f"{{{W}}}sdtPr"]
    assert ids and len(ids) == len(set(ids))
    assert [x for x in _flat(data) if isinstance(x, list)] == [x for x in _expected("A") if isinstance(x, list)]


def test_a_long_order_keeps_its_lines_and_header_across_pages(api, tmp_path):
    long = [("A", "A-1", dt.date(2024, 3, 1), k, f"零件{k:03d}", k % 7 + 1, Decimal("1.10")) for k in range(1, 121)]
    rows = long + [("A", "A-2", dt.date(2024, 3, 9), 1, "垫圈", 10, Decimal("0.50"))]
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=rows)
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed"
    data = api.outputs(job["id"])["A.docx"]
    assert _flat(data) == _expected("A", rows)
    from synwevo_docs.engine.converters import convert, soffice_path
    if soffice_path():                                        # a real office suite paginates it
        import pymupdf
        pdf = pymupdf.open(stream=convert(data, "docx", "pdf"), filetype="pdf")
        assert pdf.page_count > 1
        for page in pdf:
            text = page.get_text()
            if "零件" in text:
                assert "品名" in text, f"page {page.number + 1}: the header row does not repeat"


def test_a_version_stored_without_parent_is_nested_from_the_file(api, tmp_path):
    """TEST FIXTURE: the version's slots are rewritten without `parent`, as a version stored,
    restored or copied before nested details has them; planning reads the nesting from the file."""
    import json
    pid, tid, blk, rep, cust = _setup(api, tmp_path)
    store = api.c.app_state.app.store
    head = api.c.get(f"/api/templates/{tid}").json()["head"]
    row = store.one("SELECT slots FROM template_versions WHERE template_id=? AND version_no=?", (tid, head))
    stored = row["slots"] if isinstance(row["slots"], list) else json.loads(row["slots"])
    assert any(s.get("parent") for s in stored)
    bare = [{k: v for k, v in s.items() if k != "parent"} for s in stored]
    store.execute("UPDATE template_versions SET slots=? WHERE template_id=? AND version_no=?", (json.dumps(bare, ensure_ascii=False), tid, head))
    detail = {s["id"]: s for s in api.c.get(f"/api/templates/{tid}").json()["version"]["slots"]}
    assert detail[rep["id"]].get("parent") == blk["id"]          # what the mapping dialog reads
    _map(api, pid, blk, rep, cust)
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed", api.items(job["id"])
    assert _flat(api.outputs(job["id"])["A.docx"]) == _expected("A")


# ------------------------------------------------------------------ found by the independent review (CI-104)

def _book(head, rows, fmt=None):
    """TEST FIXTURE: an in-memory workbook read the way an uploaded source is."""
    from synwevo_docs.engine import tabular
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(head)
    for r in rows:
        ws.append(list(r))
    if fmt:
        for (col, f) in fmt.items():
            for cell in ws[col][1:]:
                cell.number_format = f
    buf = io.BytesIO()
    wb.save(buf)
    return tabular.read_xlsx(buf.getvalue()).to_json()


def _engine(books, maps, slots, scope="file"):
    from synwevo_docs.engine.plan import build_plan
    from synwevo_docs.engine.records import build_records
    from synwevo_docs.engine.resolve import SourceDoc
    srcs = {f"F{i}": SourceDoc(f"F{i}", f"f{i}.xlsx", "xlsx", "table", workbook_json=b) for i, b in enumerate(books)}
    rule = {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0}, "group": {"keys": [{"header": "客户"}], "scope": scope}}
    rec = build_records(list(srcs.values()), rule)["records"][0]
    return build_plan(rec, rule, srcs, maps, slots)


def _cols(*pairs, **agg):
    return [{"slot_col": k, "source": {"header": h}, **({"agg": agg[k]} if k in agg else {})} for k, h in pairs]


BLOCK = {"id": "B1", "binding": "B1", "name": "按订单", "kind": "repeat", "block": True,
         "columns": [{"key": "c0", "name": "订单号"}, {"key": "c1", "name": "件数"}]}
ROW = {"id": "R1", "binding": "RB", "name": "明细1", "kind": "repeat", "parent": "B1", "columns": [{"key": "c0", "name": "品名"}]}
SPLIT = {"id": "m1", "no": 1, "type": "repeat", "sources": [TABLE], "target": {"binding": "B1"}, "group_by": GROUP_BY,
         "columns": _cols(("c0", "订单号"))}
LINES = {"id": "m2", "no": 2, "type": "repeat", "sources": [TABLE], "target": {"binding": "RB"}, "columns": _cols(("c0", "品名"))}


def test_rows_are_matched_by_the_source_row_they_came_from(api, tmp_path):
    """每文件一份: the block reads columns A–B and the table C–F. Both skip a different blank row,
    so their row counts agree — the lines used to land in the next order's copy with no issue."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["订单号", "日期文本", "品名", "数量", "单价", "金额"])
    ws.append(["A-1", "2024-03-01", "螺丝", 5, 2, 10])
    ws.append(["A-2", "2024-03-09", None, None, None, None])     # an order without lines
    ws.append(["A-3", "2024-03-12", "扳手", 1, 9, 9])
    ws.append([None, None, "垫圈", 10, 1, 10])                   # a line without an order
    wb.save(tmp_path / "订单.xlsx")
    pid = api.project("嵌套明细·列范围")
    api.upload(pid, tmp_path / "订单.xlsx")
    tid = api.template(pid, _template(tmp_path / "订单模板.docx"))["id"]
    body = _body(api, tid)
    ps = {b["text"]: b for b in _walk(body["blocks"])}
    first, last = ps["订单号：O-0　下单日期：2024-01-01"], ps["订单小计：0.00"]
    tbl = next(b for b in body["blocks"] if b["t"] != "p")
    blk = api.slot(tid, kind="repeat_block", part=first["part"], first=first["path"], last=last["path"], name="订单段")
    rep = api.slot(tid, kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="明细", header_rows=1)
    api.rule(pid, {"mode": "per_file"})
    bk = {c["name"]: c["key"] for c in blk["columns"]}
    rk = {c["name"]: c["key"] for c in rep["columns"]}
    table = {"kind": "table", "sheet": "明细", "header_row": 0, "data_start": 1, "data_end": None}
    api.maps(pid, [
        {"type": "repeat", "sources": [{**table, "first_col": 0, "last_col": 1}], "target": {"binding": blk["binding"]}, "group_by": GROUP_BY,
         "columns": [{"slot_col": bk["订单号"], "source": {"header": "订单号"}}, {"slot_col": bk["下单日期"], "source": {"header": "日期文本"}}]},
        {"type": "repeat", "sources": [{**table, "first_col": 2, "last_col": 5}], "target": {"binding": rep["binding"]},
         "columns": [{"slot_col": rk["品名"], "source": {"header": "品名"}}, {"slot_col": rk["数量"], "source": {"header": "数量"}}]},
    ])
    issues = api.ok(api.c.get(f"/api/projects/{pid}/preview", params={"index": 0}))["issues"]
    e = [i for i in issues if i["level"] == "error"]
    assert [i["code"] for i in e] == ["nested_source_mismatch"] and "同一范围" in e[0]["message"], issues
    # the same range on both sides: every order gets exactly its own lines, in the order they first appear
    # (the line without an order is an error by default; skipped here, with a warning)
    maps = api.ok(api.c.get(f"/api/projects/{pid}/mappings"))["mappings"]
    for m in maps:
        m["sources"] = [{**table, "first_col": 0, "last_col": 5}]
        if m.get("group_by"):
            m["group_by"] = {**GROUP_BY, "empty_key": "skip"}
    api.maps(pid, maps)
    job = api.job(pid)
    assert job["status"] == "completed", api.items(job["id"])
    got = _flat(next(iter(api.outputs(job["id"]).values())))
    tables = [x for x in got if isinstance(x, list)]
    assert tables == [[HEAD, ["螺丝", "5", "", ""]], [HEAD], [HEAD, ["扳手", "1", "", ""]]], got   # the line without an order is not placed


def test_a_block_field_inside_the_table_row_keeps_the_blocks_value(api, tmp_path):
    """The detail row's cell has a second paragraph 「规格：M6」, which the block took as its own
    column. The table filler used to fill it with the table's column of the same key."""
    d = docx.Document()
    d.add_paragraph("客户：C-0")
    d.add_paragraph("订单号：O-0　下单日期：2024-01-01")
    t = d.add_table(rows=2, cols=4)
    for c, h in enumerate(("品名", "数量", "单价", "金额")):
        t.cell(0, c).text = h
    for c, v in enumerate(("示例", "1", "0.00", "0.00")):
        t.cell(1, c).text = v
    t.cell(1, 0).add_paragraph("规格：M6")
    d.add_paragraph("订单小计：0.00")
    d.add_paragraph("（订单结束）")
    d.save(tmp_path / "t.docx")
    pid = api.project("嵌套明细·块字段在行内")
    api.upload(pid, _workbook(tmp_path / "订单.xlsx"))
    tid = api.template(pid, tmp_path / "t.docx")["id"]
    body = _body(api, tid)
    ps = {b["text"]: b for b in _walk(body["blocks"])}
    first, last = ps["订单号：O-0　下单日期：2024-01-01"], ps["订单小计：0.00"]
    tbl = next(b for b in body["blocks"] if b["t"] != "p")
    blk = api.slot(tid, kind="repeat_block", part=first["part"], first=first["path"], last=last["path"], name="订单段")
    rep = api.slot(tid, kind="repeat", part=tbl["part"], table_path=tbl["path"], row=1, name="明细", header_rows=1)
    cust = api.text_slot(tid, "C-0", "客户")
    bk = {c["name"]: c["key"] for c in blk["columns"]}
    assert "规格" in bk
    _map(api, pid, blk, rep, cust, block_cols=[
        {"slot_col": bk["订单号"], "source": {"header": "订单号"}},
        {"slot_col": bk["下单日期"], "source": {"header": "日期文本"}},
        {"slot_col": bk["订单小计"], "source": {"header": "金额"}, "agg": "sum", "transforms": [{"op": "decimals", "n": 2}]},
        {"slot_col": bk["规格"], "source": {"header": "客户"}},
    ])
    job = api.job(pid, naming="{客户}")
    assert job["status"] == "completed", api.items(job["id"])
    lines = [r for x in _flat(api.outputs(job["id"])["A.docx"]) if isinstance(x, list) for r in x[1:]]
    assert [r[0] for r in lines] == ["螺丝规格：A", "螺母规格：A", "垫圈规格：A"], lines


def test_a_placeholder_block_still_enforces_a_table_that_needs_rows(api, tmp_path):
    rows = [("C", None, dt.date(2024, 3, 5), 1, "扳手", 1, Decimal("9.90"))]
    pid, tid, blk, rep, cust = _setup(api, tmp_path, rows=rows, row={"empty_policy": "error"}, block={"empty_policy": "placeholder"})
    _map(api, pid, blk, rep, cust, group_by={**GROUP_BY, "empty_key": "skip"})
    e = [i for i in _preview_issues(api, pid, "C") if i["level"] == "error"]
    assert [(i["code"], i["slot"]) for i in e] == [("no_detail", rep["id"])]


def test_a_nested_table_cannot_share_its_binding():
    book = _book(["客户", "订单号", "仓库", "品名"], [("A", "A-1", "W1", "螺丝"), ("A", "A-1", "W2", "螺母"), ("A", "A-2", "W1", "垫圈")])
    # two nested tables in two blocks on one binding: each would get the other's split
    b2 = {**BLOCK, "id": "B2", "binding": "B2", "name": "按仓库"}
    r2 = {**ROW, "id": "R2", "parent": "B2", "name": "明细2"}
    split2 = {**SPLIT, "id": "m9", "no": 9, "target": {"binding": "B2"}, "group_by": {"keys": [{"header": "仓库"}]}, "columns": _cols(("c0", "仓库"))}
    p = _engine([book], [SPLIT, split2, LINES], [BLOCK, b2, ROW, r2])
    assert {i["code"] for i in p.issues if i["level"] == "error"} == {"nested_shared_binding"}
    assert "RB" not in p.values and "R1" not in p.values and "R2" not in p.values
    # sharing with a table outside any block, and with its own block (which used to raise IndexError)
    outside = {"id": "R9", "binding": "RB", "name": "全部明细", "kind": "repeat", "empty_policy": "error", "columns": [{"key": "c0", "name": "品名"}]}
    p = _engine([book], [SPLIT, LINES], [BLOCK, ROW, outside])
    assert [i["code"] for i in p.issues if i["level"] == "error"] == ["nested_shared_binding"]
    own = {**ROW, "binding": "B1"}
    p = _engine([book], [SPLIT], [BLOCK, own])
    assert [i["code"] for i in p.issues if i["level"] == "error"] == ["nested_shared_binding"]


def test_count_in_a_section_counts_the_rows_that_have_a_value():
    book = _book(["客户", "订单号", "品名"], [("A", "A-1", "螺丝"), ("A", "A-1", "螺母"), ("A", "A-3", None)])
    split = {**SPLIT, "columns": _cols(("c0", "订单号"), ("c1", "品名"), c1="count")}
    p = _engine([book], [split, LINES], [BLOCK, ROW])
    assert [(r["c0"].text, r["c1"].text) for r in p.values["B1"]["rows"]] == [("A-1", "2"), ("A-3", "0")]
    assert [len(g) for g in p.values["R1"]["groups"]] == [2, 0]


def test_a_merged_file_without_the_split_column_is_named_not_called_empty():
    a = _book(["客户", "订单号", "品名"], [("A", "A-1", "螺丝")])
    b = _book(["客户", "单号", "品名"], [("A", "A-2", "螺母")])
    for policy in ("error", "skip"):
        p = _engine([a, b], [{**SPLIT, "group_by": {**GROUP_BY, "empty_key": policy}}, LINES], [BLOCK, ROW], scope="cross_file")
        e = [i for i in p.issues if i["level"] == "error"]
        assert [i["code"] for i in e] == ["missing_column"] and "订单号" in e[0]["message"] and "f1.xlsx" in e[0]["message"], e


def test_an_unsplit_block_is_reported_before_any_preview():
    from synwevo_docs.engine.plan import static_check
    unsplit = {k: v for k, v in SPLIT.items() if k != "group_by"}
    issues = static_check([unsplit, LINES], [BLOCK, ROW], {"mode": "per_group"})
    assert [(i["code"], i["mapping"], i["slot"]) for i in issues] == [("nested_needs_group_by", "m2", "R1")]
    assert static_check([SPLIT, LINES], [BLOCK, ROW], {"mode": "per_group"}) == []


def test_a_subtotal_reads_like_the_numbers_it_adds_up():
    book = _book(["客户", "订单号", "金额"], [("A", "A-1", 10), ("A", "A-1", 3.6)], fmt={"C": "#,##0.00"})
    split = {**SPLIT, "columns": _cols(("c0", "订单号"), ("c1", "金额"), c1="sum")}
    p = _engine([book], [split], [BLOCK])
    assert p.values["B1"]["rows"][0]["c1"].text == "13.60"
