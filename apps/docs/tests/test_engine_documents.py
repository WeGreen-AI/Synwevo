"""Wrong-document regressions in the engine: each test builds its own small TEST fixture in
memory and reads the generated document back.

* OCR confirmations stay with the sample (and the box) they were confirmed on;
* the built-in editor's 「重复内容块」 wraps exactly its own paragraphs;
* PPTX single cells / areas below an expanded detail row keep their target;
* optional single-value fields that end up empty are reported (not silent).
"""
from __future__ import annotations

import io

import docx
import pytest
from pptx import Presentation
from pptx.util import Inches

from synwevo_docs.engine import builtin, pptx_engine as P, values as V
from synwevo_docs.engine import pdf_engine as pe
from synwevo_docs.engine.plan import build_plan
from synwevo_docs.engine.records import build_records
from synwevo_docs.engine import resolve as R
from synwevo_docs.engine.resolve import SourceDoc


# ====================================================================== OCR corrections
def _two_form_pdf() -> bytes:
    """TEST fixture: two forms, one per page, each starting with the marker FORM-START."""
    pymupdf = pytest.importorskip("pymupdf")
    d = pymupdf.open()
    for _ in range(2):
        p = d.new_page(width=400, height=400)
        p.insert_text((40, 50), "FORM-START", fontsize=12)
    data = d.tobytes()
    d.close()
    return data


def _ocr(value_by_page: dict[int, str]) -> dict:
    # TEST fixture: what Tesseract would have stored for the value box of each page
    return {str(pno): {"w": 400, "h": 400, "engine": "fixture",
                       "words": [{"t": v, "r": [0.1, 0.30, 0.4, 0.34], "conf": 90.0, "line": "1.1.1", "n": 1}]}
            for pno, v in value_by_page.items()}


SLOT = {"id": "s1", "binding": "s1", "name": "编号", "kind": "pdf_box", "required": True}
LOC = {"kind": "ocr_region", "page": 0, "nrect": [0.0, 0.0, 1.0, 1.0], "min_conf": 101}   # 101: always needs confirming
RULE = {"mode": "per_block", "block": {"anchor": "FORM-START"}}


def _plans(corrections: dict, loc: dict = LOC):
    if pe.ENGINE != "pymupdf":
        pytest.skip("page-internal block split needs PyMuPDF")
    src = SourceDoc("F1", "forms.pdf", "pdf", "pdf", data=_two_form_pdf(), ocr=_ocr({0: "A-001", 1: "B-002"}),
                    corrections=corrections)
    recs = build_records([src], RULE)["records"]
    assert [r["file_id"] for r in recs] == ["F1~0", "F1~1"]
    m = {"id": "m1", "no": 1, "type": "single", "sources": [loc], "target": {"binding": "s1"}}
    return [build_plan(r, RULE, {"F1": src}, [m], [SLOT]) for r in recs]


def _value(plan):
    fv = plan.values.get("s1")
    return fv["value"].text if fv else None


def _codes(plan):
    return [i["code"] for i in plan.issues if i["level"] == "error"]


def test_file_level_confirmation_does_not_leak_into_every_block():
    # before the fix every block of the file returned "A-001" (and skipped the confidence gate)
    plans = _plans({"m1": {"value": "A-001", "confirmed": True}})
    assert [_value(p) for p in plans] == [None, None]
    assert [_codes(p) for p in plans] == [["ocr_low_confidence"], ["ocr_low_confidence"]]
    # the issue says which sample and which box is being confirmed
    details = [next(i for i in p.issues if i["code"] == "ocr_low_confidence")["detail"] for p in plans]
    assert [d["doc"] for d in details] == ["F1~0", "F1~1"] and [d["value"] for d in details] == ["A-001", "B-002"]
    assert details[0]["fp"] == details[1]["fp"]


def test_block_confirmation_applies_to_that_block_only():
    fp = _plans({})[1].issues[-1]["detail"]["fp"]
    parent, key = R.correction_key("F1~1", "m1")
    assert (parent, key) == ("F1", "m1~1")
    plans = _plans({key: {"value": "B-002", "confirmed": True, "fp": fp}})
    assert _codes(plans[0]) == ["ocr_low_confidence"] and _value(plans[0]) is None
    assert _codes(plans[1]) == [] and _value(plans[1]) == "B-002"
    assert plans[1].values["s1"]["value"].src["corrected"] is True


def test_block_confirmation_does_not_reach_block_eleven():
    from synwevo_docs.engine.blocks import _virtual
    parent = SourceDoc("F1", "f.pdf", "pdf", "pdf", corrections={"m1~1": {"value": "B"}, "m1~11": {"value": "L"}, "m1": {"value": "X"}})
    assert _virtual(parent, 1).corrections == {"m1": {"value": "B"}}
    assert _virtual(parent, 11).corrections == {"m1": {"value": "L"}}
    assert _virtual(parent, 2).corrections == {}


def test_confirmation_stops_applying_when_the_mapping_is_redrawn():
    fp = _plans({})[1].issues[-1]["detail"]["fp"]
    moved = {**LOC, "nrect": [0.0, 0.2, 1.0, 0.5]}
    plans = _plans({"m1~1": {"value": "B-002", "confirmed": True, "fp": fp}}, moved)
    assert _codes(plans[1]) == ["ocr_low_confidence"]


def test_legacy_confirmation_is_not_applied_to_a_multi_source_mapping():
    if pe.ENGINE != "pymupdf":
        pytest.skip("needs PyMuPDF")
    src = SourceDoc("F1", "form.pdf", "pdf", "pdf", data=_two_form_pdf(), ocr=_ocr({0: "A-001", 1: "B-002"}),
                    corrections={"m1": {"value": "X", "confirmed": True}})
    rule = {"mode": "per_file"}
    rec = build_records([src], rule)["records"][0]
    one = {"id": "m1", "no": 1, "type": "single", "sources": [LOC], "target": {"binding": "s1"}}
    two = {**one, "sources": [LOC, {**LOC, "page": 1}], "combine": {"op": "concat", "sep": "/"}}
    assert _value(build_plan(rec, rule, {"F1": src}, [one], [SLOT])) == "X"          # whole-file, one box: unchanged
    p2 = build_plan(rec, rule, {"F1": src}, [two], [SLOT])
    assert _codes(p2) == ["ocr_low_confidence"]                                     # was "X/X"


# ====================================================================== built-in repeat block
def _field(i, n):
    return {"type": "field", "attrs": {"id": i, "name": n}}


def _editor(n_before: int, header: bool = False) -> dict:
    block = {"type": "repeatBlock", "attrs": {"repeatId": "bk1", "repeatName": "商品"}, "content": [
        {"type": "paragraph", "content": [{"type": "text", "text": "品名："}, _field("c_name", "品名")]},
        {"type": "paragraph", "content": [{"type": "text", "text": "数量："}, _field("c_qty", "数量")]}]}
    body = [{"type": "paragraph", "content": [{"type": "text", "text": f"前{i}"}]} for i in range(n_before)]
    body += [{"type": "paragraph", "content": [{"type": "text", "text": "订单："}, _field("fo", "订单号")]}, block,
             {"type": "paragraph", "content": [{"type": "text", "text": "完"}]}]
    ed = {"body": {"type": "doc", "content": body}}
    if header:
        ed["header"] = {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "页眉甲"}]},
                                                   {"type": "paragraph", "content": [{"type": "text", "text": "页眉乙"}]}]}
    return ed


def _blocks(container):
    """[(is block sdt, [paragraph texts])] of the top-level children, in document order."""
    from synwevo_docs.engine import docx_core as dc
    out = []
    for el in container:
        if el.tag == dc.w("p"):
            out.append((False, [dc.para_text(el)]))
        elif el.tag == dc.w("sdt"):
            out.append((True, [dc.para_text(p) for p in el.iter(dc.w("p"))]))
    return out


@pytest.mark.parametrize("n_before", [0, 1, 3, 10])
def test_builtin_repeat_block_wraps_exactly_its_paragraphs(n_before):
    data, slots = builtin.build_docx(_editor(n_before, header=True))
    d = docx.Document(io.BytesIO(data))
    got = _blocks(d.element.body)
    want = [(False, [f"前{i}"]) for i in range(n_before)]
    want += [(False, ["订单：【订单号】"]), (True, ["品名：【品名】", "数量：【数量】"]), (False, ["完"])]
    assert got == want
    blk = next(s for s in slots if s["id"] == "bk1")
    assert blk["paragraphs"] == 2 and [c["name"] for c in blk["columns"]] == ["品名", "数量"]


def test_builtin_repeat_block_in_header_wraps_its_paragraphs():
    ed = _editor(0)
    ed["header"] = {"type": "doc", "content": [
        {"type": "paragraph", "content": [{"type": "text", "text": "页眉0"}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "页眉1"}]},
        {"type": "repeatBlock", "attrs": {"repeatId": "hb", "repeatName": "页眉块"}, "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": "项："}, _field("h1", "项")]}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "页眉尾"}]}]}
    data, _ = builtin.build_docx(ed)
    d = docx.Document(io.BytesIO(data))
    assert _blocks(d.sections[0].header._element) == [(False, ["页眉0"]), (False, ["页眉1"]), (True, ["项：【项】"]),
                                                     (False, ["页眉尾"])]


# ====================================================================== PPTX tables
def _deck(rows: list[list[str]]) -> tuple[bytes, str]:
    """TEST fixture: one slide with one table."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    gf = s.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(1), Inches(4), Inches(2))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            gf.table.cell(r, c).text = v
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), str(gf.shape_id)


def _table_texts(data: bytes) -> list[list[str]]:
    prs = Presentation(io.BytesIO(data))
    tbls = [sh for sh in prs.slides[0].shapes if sh.has_table]
    if not tbls:
        return []
    t = tbls[0].table
    return [[t.cell(r, c).text for c in range(len(t.columns))] for r in range(len(t.rows))]


def _rows(*pairs):
    return {"kind": "rows", "rows": [{"c0": V.text(a), "c1": V.text(b)} for a, b in pairs]}


def test_pptx_cell_and_region_below_detail_rows_keep_their_target():
    work, sid = _deck([["品名", "数量"], ["样例", "0"], ["合计", "0"], ["签字", "某人"]])
    slots = []
    for req in ({"kind": "repeat", "row": 1, "name": "明细"}, {"kind": "cell", "row": 2, "col": 1, "name": "合计"},
                {"kind": "region", "r1": 3, "c1": 0, "r2": 3, "c2": 1, "name": "签字区"}):
        work, sl = P.pptx_add_slot(work, {"slide": 0, "shape": sid, **req}, slots)
        slots.append(sl)
    rep, cell, reg = slots
    vals = {rep["id"]: _rows(("A", "1"), ("B", "2"), ("C", "3")), cell["id"]: {"kind": "scalar", "value": V.text("6")},
            reg["id"]: {"kind": "matrix", "rows": [[V.text("审核"), V.text("张三")]]}}
    out, issues = P.pptx_fill(work, slots, vals)
    assert issues == []
    # before the fix: [['品名','数量'],['A','1'],['B','6'],['C','3'],['合计','0'], ...]
    assert _table_texts(out) == [["品名", "数量"], ["A", "1"], ["B", "2"], ["C", "3"], ["合计", "6"], ["审核", "张三"]]


def test_pptx_two_detail_rows_in_one_table():
    work, sid = _deck([["甲", "乙"], ["a", "0"], ["中间", "-"], ["b", "0"]])
    slots = []
    for row in (1, 3):
        work, sl = P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "repeat", "row": row, "name": f"明细{row}"}, slots)
        slots.append(sl)
    out, issues = P.pptx_fill(work, slots, {slots[0]["id"]: _rows(("A1", "1"), ("A2", "2")),
                                            slots[1]["id"]: _rows(("B1", "3"), ("B2", "4"), ("B3", "5"))})
    assert issues == []
    assert _table_texts(out) == [["甲", "乙"], ["A1", "1"], ["A2", "2"], ["中间", "-"], ["B1", "3"], ["B2", "4"], ["B3", "5"]]


def test_pptx_single_cell_on_a_detail_row_is_refused():
    from synwevo_docs.engine.templates import TemplateError
    work, sid = _deck([["品名", "数量"], ["样例", "0"], ["合计", "0"]])
    slots = []
    work, rep = P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "repeat", "row": 1, "name": "明细"}, slots)
    slots.append(rep)
    with pytest.raises(TemplateError) as e:
        P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "cell", "row": 1, "col": 1, "name": "x"}, slots)
    assert e.value.code == "overlap"
    # a template saved before the check still gets a located error, never a silently lost value
    cell = {"id": "old", "name": "旧字段", "kind": "pptx_text", "locator": {"slide": 0, "shape": sid, "row": 1, "col": 1}}
    out, issues = P.pptx_fill(work, slots + [cell], {rep["id"]: _rows(("A", "1")), "old": {"kind": "scalar", "value": V.text("9")}})
    assert [(i["slot"], i["code"]) for i in issues] == [("old", "overlap")]


def test_pptx_remove_table_policy_removes_the_table():
    work, sid = _deck([["品名", "数量"], ["样例", "0"]])
    work, rep = P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "repeat", "row": 1, "name": "明细",
                                       "empty_policy": "remove_table"}, [])
    out, issues = P.pptx_fill(work, [rep], {rep["id"]: {"kind": "rows", "rows": []}})
    assert issues == [] and _table_texts(out) == []                  # was a blank row


def test_pptx_multiline_value_in_empty_cell_is_schema_valid():
    from lxml import etree
    from synwevo_docs.engine.ooxml import Package
    work, sid = _deck([["名称", ""]])
    # an empty cell as PowerPoint saves it: the paragraph carries only a:endParaRPr
    pkg = Package(work)
    part = P.slide_parts(pkg)[0]
    tc = list(pkg.xml(part).iter(P.a("tc")))[1]
    para = tc.find(f"{P.a('txBody')}/{P.a('p')}")
    for child in list(para):
        para.remove(child)
    etree.SubElement(para, P.a("endParaRPr"), {"lang": "zh-CN", "sz": "1400"})
    work = pkg.save()
    work, cell = P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "cell", "row": 0, "col": 1, "name": "x"}, [])
    out, _ = P.pptx_fill(work, [cell], {cell["id"]: {"kind": "scalar", "value": V.text("一\n二")}})
    pkg = Package(out)
    xml = pkg.xml(P.slide_parts(pkg)[0])
    brs = list(xml.iter(P.a("br")))
    assert brs and all(child.tag == P.a("rPr") for br in brs for child in br)      # never a:endParaRPr under a:br
    assert _table_texts(out) == [["名称", "一\x0b二"]]                             # python-pptx reads a:br as \v


def test_pptx_table_past_the_slide_bottom_fails_the_copy_only_when_keeping_the_layout():
    """layout=fit promises "keep the template's pages, or that copy fails". A presentation keeps
    its slide count however long a table grows, so the table running off the slide is the copy not
    fitting: under fit it is an error (the copy is not produced); otherwise a warning, as before."""
    from synwevo_docs.engine import generate as G
    from synwevo_docs.engine.plan import Plan
    work, sid = _deck([["名称", "数量"], ["a", "0"]])
    work, rep = P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "repeat", "row": 1, "name": "明细"}, [])
    vals = {rep["id"]: _rows(*[(f"第{i}项", str(i)) for i in range(40)])}
    _, issues = P.pptx_fill(work, [rep], vals)
    assert [(i["level"], i["code"]) for i in issues] == [("warning", "slide_overflow")]
    _, issues = P.pptx_fill(work, [rep], vals, {"layout_mode": "fit"})
    assert [(i["level"], i["code"]) for i in issues] == [("error", "slide_overflow")]
    assert issues[0]["message"] == "重复明细「明细」共 40 行，表格超出幻灯片底部；已选择保持模板页数，因此不产出这一份。请减少行数或调整模板"
    # the guard hands the choice to the engine; a short table fits and is no issue in any mode
    plan = Plan(record={}, values=vals)
    _, issues, _ = G.fill_checked("pptx", work, [rep], plan, None, "fit", 1)
    assert [(i["level"], i["code"]) for i in issues] == [("error", "slide_overflow")]
    _, issues, _ = G.fill_checked("pptx", work, [rep], plan, None, "off", None)
    assert [(i["level"], i["code"]) for i in issues] == [("warning", "slide_overflow")]
    _, issues = P.pptx_fill(work, [rep], {rep["id"]: _rows(("a", "1"))}, {"layout_mode": "fit"})
    assert issues == []


def _long_deck(chart: bool = False, notes: bool = False) -> tuple[bytes, str, str]:
    """TEST fixture: slide 1 has a title 「客户：甲」 and a table (header, detail, total rows of 0.5 in)
    whose top edge is 1.5 in down a 7.5 in slide; slide 2 says 「结束」. Returns (bytes, title id, table id)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    title = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(0.6))
    title.text_frame.text = "客户：甲"
    gf = s.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(6), Inches(1.5))
    for r, row in enumerate([["名称", "数量"], ["a", "0"], ["合计", "0"]]):
        for c, v in enumerate(row):
            gf.table.cell(r, c).text = v
    if chart:
        data = CategoryChartData()
        data.categories = ["x"]
        data.add_series("s", (1,))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(0.5), Inches(2), Inches(1), data)
    if notes:
        s.notes_slide.notes_text_frame.text = "演讲者备注"
    end = prs.slides.add_slide(prs.slide_layouts[6])
    end.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "结束"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), str(title.shape_id), str(gf.shape_id)


def _continuing(chart=False, notes=False, overflow="continue"):
    work, title_id, sid = _long_deck(chart, notes)
    work, name = P.pptx_add_slot(work, {"slide": 0, "shape": title_id, "kind": "text", "para": 0, "start": 3, "end": 4, "name": "客户"}, [])
    work, rep = P.pptx_add_slot(work, {"slide": 0, "shape": sid, "kind": "repeat", "row": 1, "name": "明细", "overflow": overflow}, [name])
    return work, [name, rep]


def test_pptx_detail_past_the_slide_continues_on_copies_of_the_slide():
    """overflow=continue: the table carries on over copies of its slide placed right after it; the
    header row is on every copy, the rows below the detail only on the last, and every other field
    of the slide reads the same on each copy. Rows are counted at their template height (0.5 in):
    11 fit under the header, 10 when the total row must fit too."""
    work, slots = _continuing(notes=True)
    name, rep = slots
    vals = {name["id"]: {"kind": "scalar", "value": V.text("乙公司")},
            rep["id"]: _rows(*[(f"第{i}项", str(i)) for i in range(40)])}
    out, issues = P.pptx_fill(work, slots, vals)
    assert issues == []
    prs = Presentation(io.BytesIO(out))
    assert len(prs.slides) == 5
    tables, titles = [], []
    for slide in list(prs.slides)[:4]:
        tables.append([[c.text for c in row.cells] for row in next(sh for sh in slide.shapes if sh.has_table).table.rows])
        titles.append(next(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame))
    assert titles == ["客户：乙公司"] * 4
    assert [len(t) - 1 for t in tables] == [11, 11, 11, 8]                    # header + detail rows (+ the total on the last)
    assert all(t[0] == ["名称", "数量"] for t in tables)
    detail = [row for t in tables for row in t[1:] if row[0] != "合计"]
    assert detail == [[f"第{i}项", str(i)] for i in range(40)]              # every row once, in order
    assert [row for t in tables for row in t if row[0] == "合计"] == [["合计", "0"]] and tables[-1][-1] == ["合计", "0"]
    assert [sh.text_frame.text for sh in list(prs.slides)[4].shapes if sh.has_text_frame] == ["结束"]
    # the copies fit the slide; the notes stay with the original only
    for slide in list(prs.slides)[:4]:
        frame = next(sh for sh in slide.shapes if sh.has_table)
        assert frame.top + frame.height <= prs.slide_height
    assert [sl.has_notes_slide for sl in list(prs.slides)[:4]] == [True, False, False, False]
    assert len({sl.slide_id for sl in prs.slides}) == 5
    # the renderer the job verifies with reads every copy
    rend = P.render_pptx(out)
    assert len(rend["slides"]) == 5


def test_pptx_continuation_that_fits_adds_no_slide_and_the_other_choices():
    work, slots = _continuing()
    name, rep = slots
    few = {name["id"]: {"kind": "scalar", "value": V.text("乙")}, rep["id"]: _rows(("a", "1"), ("b", "2"))}
    out, issues = P.pptx_fill(work, slots, few)
    assert issues == [] and len(Presentation(io.BytesIO(out)).slides) == 2
    def many(slots):
        return {slots[0]["id"]: {"kind": "scalar", "value": V.text("乙")}, slots[1]["id"]: _rows(*[(f"第{i}项", str(i)) for i in range(40)])}
    work_e, slots_e = _continuing(overflow="error")
    _, issues = P.pptx_fill(work_e, slots_e, many(slots_e))
    assert [(i["level"], i["code"]) for i in issues] == [("error", "slide_overflow")]
    assert "设置为超出时视为错误" in issues[0]["message"]
    work_w, slots_w = _continuing(overflow="warn")
    out, issues = P.pptx_fill(work_w, slots_w, many(slots_w))
    assert [(i["level"], i["code"]) for i in issues] == [("warning", "slide_overflow")]
    assert "复制本张幻灯片续写" in issues[0]["message"] and len(Presentation(io.BytesIO(out)).slides) == 2


def test_pptx_slide_with_a_chart_is_not_copied():
    work, slots = _continuing(chart=True)
    name, rep = slots
    _, issues = P.pptx_fill(work, slots, {name["id"]: {"kind": "scalar", "value": V.text("乙")},
                                          rep["id"]: _rows(*[(f"第{i}项", str(i)) for i in range(40)])})
    assert [(i["level"], i["code"]) for i in issues] == [("error", "continue_uncopyable")]


def test_pptx_continuation_chunks():
    h = 100
    assert P._continuation_chunks("x", 40, 11 * h, 0, h) == [11, 11, 11, 7]
    assert P._continuation_chunks("x", 40, 11 * h + 50, h, h) == [11, 11, 11, 7]     # the last holds 10
    assert P._continuation_chunks("x", 12, 11 * h, h, h) == [11, 1]
    assert P._continuation_chunks("x", 10, 8 * h, 7 * h, h) == [8, 1, 1]              # the last holds 1
    assert P._continuation_chunks("x", 9, 8 * h, 8 * h, h) == [8, 1, 0]               # only the totals on the last
    with pytest.raises(P.PptxError):
        P._continuation_chunks("x", 3, h - 1, 0, h)
    with pytest.raises(P.PptxError):
        P._continuation_chunks("x", 3, 5 * h, 0, 0)


def test_pptx_source_with_fewer_paragraphs_names_the_slide_shape_and_paragraph():
    """A source deck whose text box has fewer paragraphs than the representative file says which
    slide, which shape and which paragraph the value was read from — not just "differs"."""
    from synwevo_docs.engine.ooxml import Package
    prs = Presentation()                          # TEST fixture: one text box with one paragraph
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.name = "客户信息"
    box.text_frame.text = "客户：甲"
    buf = io.BytesIO()
    prs.save(buf)
    pkg = Package(buf.getvalue())
    loc = {"slide": 0, "shape": str(box.shape_id), "shape_name": "客户信息", "para": 2}
    with pytest.raises(R.ResolveError) as e:
        P.extract_text(pkg, loc)
    assert e.value.code == "path_missing"
    assert str(e.value) == "第 1 张幻灯片的形状「客户信息」有 1 段，代表文件中取值的是第 3 段"
    assert P.extract_text(pkg, {**loc, "para": 0}) == "客户：甲"


# ====================================================================== R02: optional fields left empty
def test_optional_empty_single_field_is_a_warning():
    src = SourceDoc("T", "t.csv", "csv", "table", workbook_json=None)
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in (["姓名", "备注", "电话"], ["张三", "", "1"], ["李四", "有", "2"]):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    from synwevo_docs.engine import tabular
    src = SourceDoc("T", "t.xlsx", "xlsx", "table", workbook_json=tabular.read_xlsx(buf.getvalue()).to_json())
    rule = {"mode": "per_row", "table": {"sheet": "Sheet1", "header_row": 0}}
    recs = build_records([src], rule)["records"]
    slots = [{"id": "a", "binding": "a", "name": "姓名", "kind": "text", "required": True},
             {"id": "b", "binding": "b", "name": "备注", "kind": "text"},
             {"id": "c", "binding": "c", "name": "未映射", "kind": "text"},
             {"id": "d", "binding": "d", "name": "有默认", "kind": "text", "default": 0}]
    maps = [{"id": "m1", "no": 1, "type": "single", "sources": [{"kind": "column", "sheet": "Sheet1", "header": "姓名"}], "target": {"binding": "a"}},
            {"id": "m2", "no": 2, "type": "single", "sources": [{"kind": "column", "sheet": "Sheet1", "header": "备注"}], "target": {"binding": "b"}}]
    p0, p1 = (build_plan(r, rule, {"T": src}, maps, slots) for r in recs)
    w0 = sorted((i["slot"], i.get("mapping")) for i in p0.issues if i["code"] == "optional_empty")
    assert w0 == [("b", "m2"), ("c", None)] and not p0.blocking
    assert all(i["level"] == "warning" for i in p0.issues if i["code"] == "optional_empty")
    assert sorted(i["slot"] for i in p1.issues if i["code"] == "optional_empty") == ["c"]
    # a default of 0 is a real value, never an empty cell
    assert p0.values["d"]["value"].text == "0"


def test_slot_default_fills_an_empty_mapped_value():
    from synwevo_docs.engine import tabular
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["备注"])
    ws.append([""])
    ws.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    src = SourceDoc("T", "t.xlsx", "xlsx", "table", workbook_json=tabular.read_xlsx(buf.getvalue()).to_json())
    rule = {"mode": "per_row", "table": {"sheet": "S", "header_row": 0, "skip_blank": False}}
    recs = build_records([src], rule)["records"]
    slots = [{"id": "b", "binding": "b", "name": "备注", "kind": "text", "required": True, "default": "无"}]
    maps = [{"id": "m", "no": 1, "type": "single", "sources": [{"kind": "column", "sheet": "S", "header": "备注"}], "target": {"binding": "b"}}]
    got = [build_plan(r, rule, {"T": src}, maps, slots) for r in recs]
    assert [p.values["b"]["value"].text for p in got] == ["无", "x"] and not any(p.blocking for p in got)


def test_bad_slot_format_is_a_located_issue_not_an_exception():
    src = SourceDoc("F", "f.txt", "txt", "text")
    src._text = "x"
    rule = {"mode": "per_file"}
    rec = build_records([src], rule)["records"][0]
    slots = [{"id": "a", "binding": "a", "name": "金额", "kind": "text", "format": [{"op": "no_such_op"}]}]
    p = build_plan(rec, rule, {"F": src}, [], slots)
    assert [(i["code"], i["slot"]) for i in p.blocking] == [("format_error", "a")]


def test_binding_shared_across_slot_kinds_is_reported():
    src = SourceDoc("F", "f.txt", "txt", "text")
    src._text = "x"
    rule = {"mode": "per_file"}
    rec = build_records([src], rule)["records"][0]
    slots = [{"id": "a", "binding": "a", "name": "单值", "kind": "cell"},
             {"id": "b", "binding": "a", "name": "区域", "kind": "region", "rows": 1, "cols": 1}]
    maps = [{"id": "m", "no": 1, "type": "single", "sources": [{"kind": "constant", "value": "v"}], "target": {"binding": "a"}}]
    p = build_plan(rec, rule, {"F": src}, maps, slots)       # raised KeyError('rows') before
    assert "binding_mixed" in [i["code"] for i in p.blocking]


# ====================================================================== per-block locator overrides (F173)
def test_file_override_applies_to_the_blocks_of_that_file():
    src = SourceDoc("F", "t.txt", "txt", "text")
    src._text = "登记表\n姓名 甲\n电话 111\n登记表\n姓名 乙\n电话 222"
    rule = {"mode": "per_block", "block": {"anchor": "登记表", "match": "starts"}}
    recs = build_records([src], rule)["records"]
    loc = {"kind": "text_span", "line": 1, "prefix": "姓名 "}
    override = {"kind": "text_span", "line": 2, "prefix": "电话 "}
    m = {"id": "m1", "no": 1, "type": "single", "sources": [loc], "target": {"binding": "s1"},
         "overrides": {"F": {"source_index": 0, "locator": override}}}          # set on the real file id
    slot = {"id": "s1", "binding": "s1", "name": "值", "kind": "text"}
    got = [build_plan(r, rule, {"F": src}, [m], [slot]).values["s1"]["value"].text for r in recs]
    assert got == ["111", "222"]                                                  # was ["甲", "乙"]: override ignored


# ====================================================================== cross-file groups (F167)
def test_detail_column_missing_in_one_file_of_a_cross_file_group_is_reported():
    from synwevo_docs.engine import tabular
    a = tabular.read_csv("订单号,品名,数量\nD1,甲,1\n".encode("utf-8"))
    b = tabular.read_csv("订单号,品名\nD1,乙\n".encode("utf-8"))            # no 数量 column
    srcs = {"A": SourceDoc("A", "a.csv", "csv", "table", workbook_json=a.to_json()),
            "B": SourceDoc("B", "b.csv", "csv", "table", workbook_json=b.to_json())}
    rule = {"mode": "per_group", "table": {"sheet": "CSV", "header_row": 0},
            "group": {"keys": [{"header": "订单号"}], "scope": "cross_file"}}
    rec = build_records(list(srcs.values()), rule)["records"][0]
    slot = {"id": "r", "binding": "r", "name": "明细", "kind": "repeat", "columns": [{"key": "c0"}, {"key": "c1"}]}
    m = {"id": "m", "no": 1, "type": "repeat", "sources": [{"kind": "table", "sheet": "CSV", "header_row": 0}],
         "target": {"binding": "r"}, "columns": [{"slot_col": "c0", "source": {"header": "品名"}},
                                                 {"slot_col": "c1", "source": {"header": "数量"}}]}
    p = build_plan(rec, rule, srcs, [m], [slot])
    err = [i for i in p.blocking if i["code"] == "missing_column"]
    assert err and "b.csv" in err[0]["message"]                  # was: b.csv's row got a blank 数量
    m_ok = {**m, "columns": [{"slot_col": "c0", "source": {"header": "品名"}}]}
    rows = build_plan(rec, rule, srcs, [m_ok], [slot]).values["r"]["rows"]
    assert [r["c0"].text for r in rows] == ["甲", "乙"]
