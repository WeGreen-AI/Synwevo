"""Optional: compute formula results with LibreOffice and write them back as
cached values into *our* output file (the LibreOffice-converted file itself is
discarded, so its re-serialisation never replaces the template's structure)."""
from __future__ import annotations

import datetime as _dt
import io

from lxml import etree

from .converters import ConversionError, recalc_xlsx
from .xlsx_doc import XlsxDoc, s

_ERROR_CODES = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA", "#SPILL!", "#CALC!"}


def fill_cached_values(data: bytes, only: set[tuple[str, str]] | None = None) -> tuple[bytes, list[dict]]:
    """Write LibreOffice's results into the formula cells of ``data``.

    ``only``: the (sheet, cell) pairs the built-in engine could not compute.  Every other
    formula keeps the cache the built-in engine wrote: overwriting all of them replaced
    correct numbers with openpyxl's reading of LibreOffice's file — dates as
    "2026-10-24 00:00:00" strings, error values as plain text."""
    import openpyxl
    from openpyxl.utils.datetime import to_excel
    try:
        calc = recalc_xlsx(data)
    except ConversionError as e:
        return data, [{"level": "warning", "code": "recalc_failed", "message": f"公式未能由排版引擎计算：{e}；Excel 打开时会重新计算"}]
    wb = openpyxl.load_workbook(io.BytesIO(calc), data_only=True, read_only=True)
    x = XlsxDoc(data)
    epoch = getattr(wb, "epoch", None)
    n = 0
    for sh in x.sheets:
        if sh["name"] not in wb.sheetnames:
            continue
        wanted = None if only is None else {ref for sname, ref in only if sname == sh["name"]}
        if wanted is not None and not wanted:
            continue
        ws = wb[sh["name"]]
        values = {}
        for row in ws.iter_rows():
            for c in row:
                if getattr(c, "value", None) is not None and (wanted is None or c.coordinate in wanted):
                    values[c.coordinate] = c.value
        root = x.sheet_root(sh["name"])
        for c in root.iter(s("c")):
            f = c.find(s("f"))
            if f is None or (wanted is not None and c.get("r") not in wanted):
                continue
            val = values.get(c.get("r"))
            for old in c.findall(s("v")):
                c.remove(old)
            if "t" in c.attrib:
                del c.attrib["t"]
            if val is None:
                continue
            v = etree.SubElement(c, s("v"))
            if isinstance(val, bool):
                c.set("t", "b")
                v.text = "1" if val else "0"
            elif isinstance(val, int):
                # exact: an identifier-sized integer (beyond 2**53) must not pass through a double
                v.text = str(val)
            elif isinstance(val, float):
                from .values import float_repr
                v.text = float_repr(val)
            elif isinstance(val, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
                # a date-formatted number: back to the serial it is in the file
                from .values import float_repr
                v.text = float_repr(float(to_excel(val, epoch) if epoch is not None else to_excel(val)))
            elif isinstance(val, str) and val in _ERROR_CODES:
                c.set("t", "e")
                v.text = val
            else:
                c.set("t", "str")
                v.text = str(val)
            n += 1
    wb.close()
    out = x.pkg.save()
    return out, [{"level": "info", "code": "recalc", "message": f"已由排版引擎计算 {n} 个公式的结果并写入缓存值"}]
