"""Tabular sources: CSV, XLSX/XLSM, XLS (and ODS/others via conversion).

A parsed workbook is a list of :class:`Sheet` objects holding typed values for
*every* used cell – nothing is truncated; the UI pages through it.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

from . import values as V
from .limits import LIMITS, LimitError
from .numfmt import format_value


@dataclass
class Sheet:
    name: str
    index: int
    rows: list[list[V.Value]] = field(default_factory=list)
    merged: list[list[int]] = field(default_factory=list)  # [r1, c1, r2, c2] 0-based inclusive
    hidden: bool = False
    notes: list[str] = field(default_factory=list)
    formula_cells: int = 0
    formula_without_cache: list[str] = field(default_factory=list)

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def cell(self, r: int, c: int) -> V.Value:
        if 0 <= r < len(self.rows) and 0 <= c < len(self.rows[r]):
            return self.rows[r][c]
        return V.empty()

    def row_is_blank(self, r: int) -> bool:
        if r >= len(self.rows):
            return True
        return all(v.is_empty or (v.t == V.STR and str(v.v).strip() == "") for v in self.rows[r])

    def to_json(self) -> dict:
        return {
            "name": self.name, "index": self.index, "hidden": self.hidden, "merged": self.merged,
            "notes": self.notes, "formula_cells": self.formula_cells,
            "formula_without_cache": self.formula_without_cache[:50],
            "rows": [[[c.t, c.v, c.text] + ([c.fmt] if c.fmt else []) for c in row] for row in self.rows],
        }

    @staticmethod
    def from_json(d: dict) -> "Sheet":
        rows = [[V.Value(c[0], c[1], c[2], c[3] if len(c) > 3 else None) for c in row] for row in d["rows"]]
        return Sheet(d["name"], d["index"], rows, d.get("merged", []), d.get("hidden", False), d.get("notes", []),
                     d.get("formula_cells", 0), d.get("formula_without_cache", []))


@dataclass
class Workbook:
    sheets: list[Sheet]
    info: dict = field(default_factory=dict)

    def sheet(self, name: str | None = None, index: int | None = None) -> Sheet | None:
        if name is not None:
            exact = [s for s in self.sheets if s.name == name]
            if exact:
                return exact[0]
            return None
        if index is not None and 0 <= index < len(self.sheets):
            return self.sheets[index]
        return None

    def to_json(self) -> dict:
        return {"sheets": [s.to_json() for s in self.sheets], "info": self.info}

    @staticmethod
    def from_json(d: dict) -> "Workbook":
        return Workbook([Sheet.from_json(s) for s in d["sheets"]], d.get("info", {}))


# ---------------------------------------------------------------- helpers
def col_letter(c: int) -> str:
    s = ""
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        s = chr(65 + rem) + s
    return s


def col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")


def parse_cell(ref: str) -> tuple[int, int]:
    m = _CELL_RE.match(ref.strip())
    if not m:
        raise ValueError(f"无效的单元格地址：{ref}")
    return int(m.group(2)) - 1, col_index(m.group(1))


def cell_ref(r: int, c: int) -> str:
    return f"{col_letter(c)}{r + 1}"


def range_ref(r1: int, c1: int, r2: int, c2: int) -> str:
    a, b = cell_ref(r1, c1), cell_ref(r2, c2)
    return a if a == b else f"{a}:{b}"


def parse_range(ref: str) -> tuple[int, int, int, int]:
    if ":" in ref:
        a, b = ref.split(":", 1)
        r1, c1 = parse_cell(a)
        r2, c2 = parse_cell(b)
    else:
        r1, c1 = parse_cell(ref)
        r2, c2 = r1, c1
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


def normalize_header(s: str) -> str:
    s = unicodedata.normalize("NFKC", s or "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _check_cells(n: int):
    if n > LIMITS.max_cells:
        raise LimitError(f"单元格数量超过上限 {LIMITS.max_cells:,}（可通过 SYNWEVO_DOCS_MAX_CELLS 调整）")


# ---------------------------------------------------------------- CSV
def detect_encoding(data: bytes) -> tuple[str, bool]:
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig", True
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return "utf-16", True
    try:
        data.decode("utf-8")
        return "utf-8", False
    except UnicodeDecodeError:
        pass
    try:
        data.decode("gb18030")
        return "gb18030", False
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data).best()
        if best is not None:
            return best.encoding, False
    except Exception:  # pragma: no cover - library optional
        pass
    return "latin-1", False


def detect_delimiter(text: str) -> str:
    sample = text[:64 * 1024]
    try:
        d = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        return d
    except csv.Error:
        pass
    first = sample.split("\n", 1)[0]
    counts = {d: first.count(d) for d in [",", "\t", ";", "|"]}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def read_csv(data: bytes, options: dict | None = None) -> Workbook:
    options = options or {}
    enc = options.get("encoding")
    had_bom = False
    if not enc or enc == "auto":
        enc, had_bom = detect_encoding(data)
    try:
        text_ = data.decode(enc)
    except (UnicodeDecodeError, LookupError) as e:
        raise ValueError(f"无法按编码 {enc} 解码 CSV：{e}") from e
    if text_.startswith("﻿"):
        text_ = text_[1:]
    delim = options.get("delimiter") or "auto"
    if delim == "auto":
        delim = detect_delimiter(text_)
    elif delim == "\\t":
        delim = "\t"
    quote = options.get("quotechar") or '"'
    reader = csv.reader(io.StringIO(text_, newline=""), delimiter=delim, quotechar=quote, strict=False)
    rows: list[list[V.Value]] = []
    total = 0
    for rec in reader:
        total += len(rec)
        _check_cells(total)
        rows.append([V.Value(V.STR, s, s) if s != "" else V.empty() for s in rec])
    replaced = text_.count("�")
    info = {"encoding": enc, "had_bom": had_bom, "delimiter": delim, "quotechar": quote,
            "replacement_chars": replaced}
    sh = Sheet("CSV", 0, rows)
    if replaced:
        sh.notes.append(f"解码后出现 {replaced} 个无法识别的字符，请确认文件编码")
    return Workbook([sh], info)


# ---------------------------------------------------------------- XLSX
def read_xlsx(data: bytes) -> Workbook:
    import openpyxl
    from lxml import etree
    from .ooxml import Package, q

    pkg = Package(data)
    if pkg.has("EncryptedPackage") or not pkg.has("xl/workbook.xml"):
        raise ValueError("不是有效的 XLSX 工作簿（可能已加密或损坏）")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except zipfile.BadZipFile as e:
        raise ValueError("XLSX 压缩结构损坏") from e
    # sheet part paths for merge / formula scan
    wbx = pkg.xml("xl/workbook.xml")
    sheet_parts = {}
    for sh in wbx.iter(q("s", "sheet")):
        rid = sh.get("{%s}id" % "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
        sheet_parts[sh.get("name")] = pkg.rel_target("xl/workbook.xml", rid)
    sheets: list[Sheet] = []
    total = 0
    for idx, ws in enumerate(wb.worksheets):
        rows: list[list[V.Value]] = []
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        for r_i, row in enumerate(ws.iter_rows()):
            out: list[V.Value] = []
            for c_i, cell in enumerate(row):
                val = getattr(cell, "value", None)
                fmt = getattr(cell, "number_format", None)
                if val is None:
                    out.append(V.empty())
                    continue
                disp, _exact = format_value(val, fmt)
                out.append(V.from_python(val, fmt=None if fmt == "General" else fmt, display=disp))
            while out and out[-1].is_empty and out[-1].t == V.EMPTY:
                out.pop()
            rows.append(out)
            total += len(out)
            _check_cells(total)
        while rows and not rows[-1]:
            rows.pop()
        sh = Sheet(ws.title, idx, rows, hidden=(ws.sheet_state != "visible"))
        part = sheet_parts.get(ws.title)
        if part and pkg.has(part):
            _scan_sheet_xml(pkg.read(part), sh)
        sheets.append(sh)
    wb.close()
    missing = [(sh.name, ref) for sh in sheets for ref in sh.formula_without_cache]
    if missing:
        _compute_missing(data, sheets, missing)
    info = {"date1904": bool(getattr(wb, "epoch", None) and wb.epoch.year == 1904), "inventory": pkg.inventory()}
    return Workbook(sheets, info)


def _scan_sheet_xml(xml: bytes, sh: Sheet) -> None:
    from lxml import etree
    S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    for ev, el in etree.iterparse(io.BytesIO(xml), events=("end",), huge_tree=True):
        tag = el.tag
        if tag == S + "c":
            f = el.find(S + "f")
            if f is not None:
                sh.formula_cells += 1
                v = el.find(S + "v")
                if v is None or (v.text or "") == "":
                    if len(sh.formula_without_cache) < 200000:
                        sh.formula_without_cache.append(el.get("r"))
            el.clear()
        elif tag == S + "mergeCell":
            r1, c1, r2, c2 = parse_range(el.get("ref"))
            sh.merged.append([r1, c1, r2, c2])
        elif tag == S + "row":
            el.clear()
    if sh.formula_cells:
        sh.notes.append(f"含 {sh.formula_cells} 个公式单元格：读取文件中保存的计算结果（与 Excel 最后一次计算一致）")



def _compute_missing(data: bytes, sheets: list[Sheet], missing: list[tuple[str, str]]) -> None:
    """Formulas saved without a cached result (e.g. files written by scripts): compute them
    with the built-in engine so the source value is real, and say so."""
    from .formula import XLError, evaluate_cells
    try:
        vals = evaluate_cells(data, missing)
    except Exception as e:  # noqa: BLE001
        for sh in sheets:
            if sh.formula_without_cache:
                sh.notes.append(f"{len(sh.formula_without_cache)} 个公式没有缓存结果且无法计算（{e}），这些单元格读取为空")
        return
    by_sheet = {sh.name: sh for sh in sheets}
    done: dict[str, int] = {}
    failed: dict[str, list[str]] = {}
    for (sname, ref), v in vals.items():
        sh = by_sheet[sname]
        r, c = parse_cell(ref)
        if isinstance(v, Exception):
            failed.setdefault(sname, []).append(f"{ref}（{v}）")
            continue
        while len(sh.rows) <= r:
            sh.rows.append([])
        row = sh.rows[r]
        while len(row) <= c:
            row.append(V.empty())
        if isinstance(v, XLError):
            row[c] = V.Value(V.ERROR, v.code, v.code)
        elif isinstance(v, bool):
            row[c] = V.from_python(v)
        elif isinstance(v, float):
            row[c] = V.from_python(v, display=format_value(v, row[c].fmt)[0] if row[c].fmt else None)
        elif v is None:
            row[c] = V.empty()
        else:
            row[c] = V.text(str(v))
        done[sname] = done.get(sname, 0) + 1
    for sh in sheets:
        if done.get(sh.name):
            sh.notes.append(f"{done[sh.name]} 个公式在文件中没有缓存结果，已由内置公式引擎计算")
        if failed.get(sh.name):
            sh.notes.append(f"{len(failed[sh.name])} 个公式没有缓存结果且内置引擎不支持，读取为空：{'、'.join(failed[sh.name][:3])}")


# ---------------------------------------------------------------- XLS (BIFF)
def read_xls(data: bytes) -> Workbook:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
    except xlrd.biffh.XLRDError as e:
        msg = str(e)
        if "encrypted" in msg.lower():
            raise ValueError("XLS 文件已加密，无法读取") from e
        raise ValueError(f"无法读取 XLS：{msg}") from e
    sheets = []
    total = 0
    for idx in range(book.nsheets):
        sh = book.sheet_by_index(idx)
        rows = []
        for r in range(sh.nrows):
            out = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                fmt = None
                try:
                    xf = book.xf_list[cell.xf_index]
                    fmt = book.format_map[xf.format_key].format_str
                except Exception:
                    pass
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    out.append(V.empty())
                elif cell.ctype == xlrd.XL_CELL_TEXT:
                    out.append(V.Value(V.STR, cell.value, cell.value))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    disp, _ = format_value(cell.value, fmt)
                    out.append(V.from_python(cell.value, fmt=None if fmt == "General" else fmt, display=disp))
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    dt = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    disp, _ = format_value(dt, fmt)
                    out.append(V.from_python(dt, fmt=fmt, display=disp))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    out.append(V.from_python(bool(cell.value)))
                else:
                    out.append(V.Value(V.ERROR, str(cell.value), xlrd.error_text_from_code.get(cell.value, "#ERR")))
            while out and out[-1].t == V.EMPTY:
                out.pop()
            rows.append(out)
            total += len(out)
            _check_cells(total)
        while rows and not rows[-1]:
            rows.pop()
        merged = [[r1, c1, r2 - 1, c2 - 1] for (r1, r2, c1, c2) in sh.merged_cells]
        s = Sheet(sh.name, idx, rows, merged, hidden=bool(getattr(sh, "visibility", 0)))
        s.notes.append("XLS 为旧格式：直接读取单元格值与数字格式；公式读取的是文件中保存的结果")
        sheets.append(s)
    return Workbook(sheets, {"date1904": bool(book.datemode)})


# ---------------------------------------------------------------- table views
@dataclass
class TableView:
    sheet: Sheet
    header_row: int | None          # 0-based; None = no header row
    headers: list[str]
    data_rows: list[int]            # 0-based sheet row indices of data rows (after filters)
    excluded_blank: int = 0
    excluded_filtered: int = 0
    first_col: int = 0
    last_col: int | None = None

    def header_index(self) -> dict[str, list[int]]:
        out: dict[str, list[int]] = {}
        for i, h in enumerate(self.headers):
            out.setdefault(normalize_header(h), []).append(i + self.first_col)
        return out


def build_table_view(sheet: Sheet, header_row: int | None = 0, data_start: int | None = None,
                     data_end: int | None = None, skip_blank: bool = True, filters: list[dict] | None = None,
                     first_col: int = 0, last_col: int | None = None) -> TableView:
    ncols = sheet.n_cols
    lc = ncols - 1 if last_col is None else min(last_col, max(ncols - 1, last_col))
    headers: list[str] = []
    if header_row is not None:
        for c in range(first_col, lc + 1):
            headers.append(sheet.cell(header_row, c).text or "")
    else:
        headers = [col_letter(c) for c in range(first_col, lc + 1)]
    start = data_start if data_start is not None else ((header_row + 1) if header_row is not None else 0)
    end = data_end if data_end is not None else sheet.n_rows - 1
    rows, blank, filtered = [], 0, 0
    tv = TableView(sheet, header_row, headers, rows, first_col=first_col, last_col=lc)
    hidx = tv.header_index()
    for r in range(start, end + 1):
        if skip_blank and _row_blank(sheet, r, first_col, lc):
            blank += 1
            continue
        if filters and not _passes(sheet, r, filters, hidx):
            filtered += 1
            continue
        rows.append(r)
    tv.excluded_blank, tv.excluded_filtered = blank, filtered
    return tv


def _row_blank(sheet: Sheet, r: int, c1: int, c2: int) -> bool:
    for c in range(c1, c2 + 1):
        v = sheet.cell(r, c)
        if not (v.is_empty or (v.t == V.STR and str(v.v).strip() == "")):
            return False
    return True


def _passes(sheet: Sheet, r: int, filters: list[dict], hidx: dict[str, list[int]]) -> bool:
    for f in filters:
        cols = hidx.get(normalize_header(f.get("column", "")), [])
        if not cols:
            continue
        v = sheet.cell(r, cols[0]).text
        op, val = f.get("op", "nonempty"), str(f.get("value", ""))
        if op == "nonempty" and not v.strip():
            return False
        if op == "empty" and v.strip():
            return False
        if op == "eq" and v != val:
            return False
        if op == "neq" and v == val:
            return False
        if op == "contains" and val not in v:
            return False
        if op == "not_contains" and val in v:
            return False
    return True


class ColumnError(ValueError):
    def __init__(self, code: str, message: str, candidates: list[int] | None = None):
        super().__init__(message)
        self.code = code
        self.candidates = candidates or []


def resolve_column(tv: TableView, binding: dict) -> int:
    """Resolve a column binding {'header': str, 'occurrence': int|None, 'by': 'header'|'position', 'col': int}."""
    by = binding.get("by", "header")
    if by == "position":
        return int(binding["col"])
    name = normalize_header(binding.get("header", ""))
    idx = tv.header_index().get(name, [])
    if not idx:
        raise ColumnError("missing_column", f"找不到列「{binding.get('header')}」")
    occ = binding.get("occurrence")
    if len(idx) > 1:
        if occ is None:
            raise ColumnError("ambiguous_column", f"存在 {len(idx)} 个同名列「{binding.get('header')}」，需要指定是第几个", idx)
        if occ < 1 or occ > len(idx):
            raise ColumnError("missing_column", f"同名列「{binding.get('header')}」没有第 {occ} 个")
        return idx[occ - 1]
    if occ not in (None, 1):
        raise ColumnError("missing_column", f"同名列「{binding.get('header')}」只有 1 个，找不到第 {occ} 个")
    return idx[0]
