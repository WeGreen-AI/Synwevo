"""Built-in template editors → real DOCX / XLSX files with structured fields.

Document editor JSON is a ProseMirror/TipTap document; spreadsheet editor JSON is
a compact grid description.  Users never write template syntax.
"""
from __future__ import annotations

import base64
import io
import re

from . import docx_core as dc
from . import tabular
from .ooxml import Package

PAPER_MM = {"A4": (210, 297), "A3": (297, 420), "A5": (148, 210), "Letter": (215.9, 279.4), "Legal": (215.9, 355.6), "B5": (176, 250)}


class BuiltinError(ValueError):
    pass


def _mm(v):
    from docx.shared import Mm
    return Mm(float(v))


# ====================================================================== document
def build_docx(editor: dict) -> tuple[bytes, list[dict]]:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = docx.Document()
    # East Asian default font so Chinese text renders with a CJK face
    normal = d.styles["Normal"]
    normal.font.size = Pt(float(editor.get("base_size", 11)))
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), editor.get("font_ea", "宋体"))
    page = editor.get("page") or {}
    sec = d.sections[0]
    w_mm, h_mm = PAPER_MM.get(page.get("size", "A4"), PAPER_MM["A4"])
    if page.get("orientation") == "landscape":
        sec.orientation = WD_ORIENT.LANDSCAPE
        w_mm, h_mm = h_mm, w_mm
    sec.page_width, sec.page_height = _mm(w_mm), _mm(h_mm)
    m = page.get("margins") or {}
    sec.top_margin = _mm(m.get("top", 25.4))
    sec.bottom_margin = _mm(m.get("bottom", 25.4))
    sec.left_margin = _mm(m.get("left", 31.8))
    sec.right_margin = _mm(m.get("right", 31.8))
    ctx = {"fields": {}, "repeats": {}, "doc": d, "content_width_mm": w_mm - float(m.get("left", 31.8)) - float(m.get("right", 31.8))}
    body = editor.get("body") or {"type": "doc", "content": []}
    # remove the default empty paragraph
    for p in list(d.paragraphs):
        p._element.getparent().remove(p._element)
    _emit_blocks(d, body.get("content", []), ctx)
    if not d.paragraphs and not d.tables:
        d.add_paragraph()
    for key, part in (("header", sec.header), ("footer", sec.footer)):
        content = (editor.get(key) or {}).get("content") or []
        if content:
            part.is_linked_to_previous = False
            for p in list(part.paragraphs):
                p._element.getparent().remove(p._element)
            _emit_blocks(part, content, ctx)
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    slots = _collect_docx_slots(data, ctx)
    return data, slots


def _emit_blocks(container, nodes: list[dict], ctx: dict, repeat_id: str | None = None):
    from docx.enum.text import WD_BREAK
    for node in nodes:
        t = node.get("type")
        if t in ("paragraph", "heading"):
            style = None
            if t == "heading":
                style = f"Heading {min(max(int((node.get('attrs') or {}).get('level', 1)), 1), 6)}"
            p = container.add_paragraph(style=style) if style else container.add_paragraph()
            _align(p, (node.get("attrs") or {}).get("textAlign"))
            _emit_inline(p, node.get("content") or [], ctx, repeat_id)
        elif t == "repeatBlock" and repeat_id is None:
            _emit_repeat_block(container, node, ctx)
        elif t in ("bulletList", "orderedList"):
            style = "List Bullet" if t == "bulletList" else "List Number"
            for li in node.get("content") or []:
                for sub in li.get("content") or []:
                    if sub.get("type") == "paragraph":
                        p = container.add_paragraph(style=style)
                        _emit_inline(p, sub.get("content") or [], ctx, repeat_id)
                    else:
                        _emit_blocks(container, [sub], ctx, repeat_id)
        elif t == "table":
            _emit_table(container, node, ctx)
        elif t == "pageBreak":
            p = container.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        elif t == "horizontalRule":
            p = container.add_paragraph()
            _hr(p)
        elif t == "image":
            p = container.add_paragraph()
            _align(p, (node.get("attrs") or {}).get("align", "center"))
            _image(p, node.get("attrs") or {}, ctx)
        elif t == "blockquote":
            _emit_blocks(container, node.get("content") or [], ctx)


def _emit_repeat_block(container, node: dict, ctx: dict):
    """Emit the block's children, then wrap exactly those new elements in a block-level
    content control ``dw:<id>`` (the same structure an uploaded DOCX block gets)."""
    attrs = node.get("attrs") or {}
    bid = attrs.get("repeatId") or ("b" + base64.b16encode(bytes(str(len(ctx["repeats"]) + 1), "ascii")).decode().lower())
    parent = container.element.body if hasattr(container, "element") and hasattr(container.element, "body") else container._element
    before = {id(c) for c in parent}
    info = ctx["repeats"].setdefault(bid, {"columns": []})
    info.update(name=attrs.get("repeatName") or "明细", block=True,
                separator=attrs.get("separator") if attrs.get("separator") in ("none", "blank", "page_break") else "none",
                empty_policy=attrs.get("emptyPolicy") if attrs.get("emptyPolicy") in ("remove", "placeholder", "error") else "remove")
    _emit_blocks(container, node.get("content") or [], ctx, bid)
    new = [c for c in parent if id(c) not in before and c.tag in (dc.w("p"), dc.w("tbl"))]
    if not new:
        return
    sdt = dc.make_sdt(dc.TAG_PREFIX + bid, info["name"])
    new[0].addprevious(sdt)
    content = sdt.find(dc.w("sdtContent"))
    for el in new:
        content.append(el)
    info["paragraphs"] = sum(1 for el in new if el.tag == dc.w("p"))


def _align(p, align):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    m = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
    if align in m:
        p.alignment = m[align]


def _hr(p):
    from docx.oxml.ns import qn
    ppr = p._element.get_or_add_pPr()
    bdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = bdr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "auto"})
    bdr.append(bottom)
    ppr.append(bdr)


def _image(p, attrs: dict, ctx: dict):
    from docx.shared import Mm
    src = attrs.get("src") or ""
    m = re.match(r"data:image/(png|jpeg|jpg|gif|bmp);base64,(.*)$", src, re.S)
    if not m:
        return
    data = base64.b64decode(m.group(2))
    width = attrs.get("width")
    run = p.add_run()
    maxw = ctx["content_width_mm"]
    if width:
        wmm = min(float(width) * 25.4 / 96, maxw)
        run.add_picture(io.BytesIO(data), width=Mm(wmm))
    else:
        run.add_picture(io.BytesIO(data))
        shp = run._element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
        if shp:
            cx = int(shp[0].get("cx"))
            maxemu = int(maxw * 36000)
            if cx > maxemu:
                ratio = maxemu / cx
                for e in run._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent",
                                           "{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
                    e.set("cx", str(int(int(e.get("cx")) * ratio)))
                    e.set("cy", str(int(int(e.get("cy")) * ratio)))


def _emit_inline(p, nodes: list[dict], ctx: dict, repeat_id: str | None = None):
    from docx.shared import Pt, RGBColor
    for node in nodes:
        t = node.get("type")
        if t == "text":
            run = p.add_run(node.get("text", ""))
            _marks(run, node.get("marks") or [])
        elif t == "hardBreak":
            p.add_run().add_break()
        elif t == "field":
            attrs = node.get("attrs") or {}
            fid = attrs.get("id") or ("f" + base64.b16encode(bytes(str(len(ctx["fields"]) + 1), "ascii")).decode().lower())
            name = attrs.get("name") or "字段"
            run = p.add_run(f"【{name}】")
            _marks(run, node.get("marks") or [])
            tag = dc.TAG_PREFIX + (f"{repeat_id}#{fid}" if repeat_id else fid)
            sdt = dc.make_sdt(tag, name, run._element.find(dc.w("rPr")))
            run._element.addprevious(sdt)
            sdt.find(dc.w("sdtContent")).append(run._element)
            if repeat_id:
                ctx["repeats"].setdefault(repeat_id, {"columns": []})["columns"].append({"key": fid, "name": name})
            else:
                ctx["fields"][fid] = {"name": name, "required": bool(attrs.get("required")), "data_type": attrs.get("dataType") or "auto"}


def _marks(run, marks):
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor
    for mk in marks:
        t = mk.get("type")
        a = mk.get("attrs") or {}
        if t == "bold":
            run.bold = True
        elif t == "italic":
            run.italic = True
        elif t == "underline":
            run.underline = True
        elif t == "strike":
            run.font.strike = True
        elif t == "textStyle":
            col = a.get("color")
            if col and re.fullmatch(r"#?[0-9a-fA-F]{6}", col):
                run.font.color.rgb = RGBColor.from_string(col.lstrip("#").upper())
            fs = a.get("fontSize")
            if fs:
                m = re.match(r"([\d.]+)", str(fs))
                if m:
                    run.font.size = Pt(float(m.group(1)))
            ff = a.get("fontFamily")
            if ff:
                run.font.name = ff
                rpr = run._element.get_or_add_rPr()
                from docx.oxml.ns import qn
                rf = rpr.find(qn("w:rFonts"))
                if rf is not None:
                    rf.set(qn("w:eastAsia"), ff)
        elif t == "highlight":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _emit_table(container, node, ctx):
    from docx.oxml.ns import qn
    rows = [r for r in node.get("content") or [] if r.get("type") == "tableRow"]
    if not rows:
        return
    # compute grid size honoring colspan/rowspan
    ncols = 0
    occupancy: dict[tuple, bool] = {}
    layout = []
    for ri, row in enumerate(rows):
        ci = 0
        cells = []
        for cell in row.get("content") or []:
            while occupancy.get((ri, ci)):
                ci += 1
            a = cell.get("attrs") or {}
            cs, rs = int(a.get("colspan", 1) or 1), int(a.get("rowspan", 1) or 1)
            for dr in range(rs):
                for dcol in range(cs):
                    occupancy[(ri + dr, ci + dcol)] = True
            cells.append((ci, cs, rs, cell))
            ci += cs
        ncols = max(ncols, ci)
        layout.append(cells)
    if hasattr(container, "add_table"):
        try:
            tbl = container.add_table(rows=len(rows), cols=ncols)
        except TypeError:
            from docx.shared import Mm
            tbl = container.add_table(rows=len(rows), cols=ncols, width=Mm(ctx["content_width_mm"]))
    else:
        return
    try:
        tbl.style = ctx["doc"].styles["Table Grid"]
    except KeyError:
        pass
    for ri, row in enumerate(rows):
        attrs = row.get("attrs") or {}
        repeat_id = attrs.get("repeatId") if attrs.get("repeat") else None
        is_header = all(c[3].get("type") == "tableHeader" for c in layout[ri]) and layout[ri]
        for ci, cs, rs, cell in layout[ri]:
            target = tbl.cell(ri, ci)
            if cs > 1 or rs > 1:
                target = target.merge(tbl.cell(ri + rs - 1, ci + cs - 1))
            p = target.paragraphs[0]
            first = True
            for sub in cell.get("content") or []:
                if sub.get("type") in ("paragraph", "heading"):
                    if not first:
                        p = target.add_paragraph()
                    first = False
                    _align(p, (sub.get("attrs") or {}).get("textAlign"))
                    _emit_inline(p, sub.get("content") or [], ctx, repeat_id)
            bg = (cell.get("attrs") or {}).get("background")
            if bg and re.fullmatch(r"#[0-9a-fA-F]{6}", bg):
                tcpr = target._tc.get_or_add_tcPr()
                shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): bg[1:]})
                tcpr.append(shd)
        if is_header:
            dc.set_row_header(tbl.rows[ri]._tr, True)
        if repeat_id:
            info = ctx["repeats"].setdefault(repeat_id, {"columns": []})
            info["name"] = attrs.get("repeatName") or "明细"
            info["header_rows"] = sum(1 for k in range(ri) if all(c[3].get("type") == "tableHeader" for c in layout[k])) or (1 if ri > 0 else 0)
            info["empty_policy"] = attrs.get("emptyPolicy") or "remove_row"
            # cells without a field become empty columns (still repeated)
            for ci, cs, rs, cell in layout[ri]:
                has_field = any(n.get("type") == "field" for sub in cell.get("content") or [] for n in sub.get("content") or [])
                if not has_field:
                    tc = tbl.cell(ri, ci)._tc
                    p = dc.cell_first_paragraph(tc)
                    key = f"c{ci}"
                    sdt = dc.make_sdt(f"{dc.TAG_PREFIX}{repeat_id}#{key}", f"第{ci + 1}列", None, [dc.etree.Element(dc.w("r"))])
                    p.append(sdt)
                    info["columns"].append({"key": key, "name": f"第{ci + 1}列"})


def _collect_docx_slots(data: bytes, ctx: dict) -> list[dict]:
    pkg = Package(data)
    slots = []
    seen = set()
    for info in dc.list_parts(pkg):
        for tag, sdt in dc.iter_field_sdts(pkg.xml(info["part"])):
            base = tag.split("#")[0]
            if base in seen:
                continue
            seen.add(base)
            if base in ctx["repeats"] and ctx["repeats"][base].get("block"):
                rp = ctx["repeats"][base]
                slots.append({"id": base, "binding": base, "name": rp.get("name", "明细"), "kind": "repeat", "block": True,
                              "columns": [{"key": c["key"], "name": c["name"]} for c in rp["columns"]],
                              "separator": rp.get("separator", "none"), "empty_policy": rp.get("empty_policy", "remove"),
                              "placeholder": "（无）", "paragraphs": rp.get("paragraphs"), "locator": {"part": info["part"], "block": True},
                              "data_type": "auto", "required": False, "format": [], "default": None})
            elif base in ctx["repeats"]:
                rp = ctx["repeats"][base]
                slots.append({"id": base, "binding": base, "name": rp.get("name", "明细"), "kind": "repeat",
                              "columns": [{"key": c["key"], "name": c["name"], "col": i} for i, c in enumerate(rp["columns"])],
                              "header_rows": rp.get("header_rows", 1), "empty_policy": rp.get("empty_policy", "remove_row"),
                              "placeholder": "（无）", "repeat_header": True, "locator": {"part": info["part"]},
                              "data_type": "auto", "required": False, "format": [], "default": None})
            elif base in ctx["fields"]:
                f = ctx["fields"][base]
                slots.append({"id": base, "binding": base, "name": f["name"], "kind": "text", "mode": "replace",
                              "original_text": f"【{f['name']}】", "locator": {"part": info["part"]},
                              "data_type": f.get("data_type", "auto"), "required": f.get("required", False), "format": [], "default": None})
    # fields with the same name share one binding unless the editor said otherwise
    by_name: dict[str, str] = {}
    for s in slots:
        if s["kind"] == "text":
            if s["name"] in by_name:
                s["binding"] = by_name[s["name"]]
            else:
                by_name[s["name"]] = s["id"]
    return slots


# ====================================================================== spreadsheet
def build_xlsx(editor: dict) -> tuple[bytes, list[dict]]:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from .templates import xlsx_add_slot

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = editor.get("sheets") or [{"name": "Sheet1"}]
    names = set()
    for sh in sheets:
        name = re.sub(r"[\\/*?:\[\]]", "_", (sh.get("name") or "Sheet").strip())[:31] or "Sheet"
        if name in names:
            raise BuiltinError(f"工作表名称重复：{name}")
        names.add(name)
        ws = wb.create_sheet(name)
        for ci, col in enumerate(sh.get("cols") or []):
            if col and col.get("w"):
                ws.column_dimensions[get_column_letter(ci + 1)].width = max(1, (float(col["w"]) - 5) / 7)
        for ri, row in enumerate(sh.get("rows") or []):
            if row and row.get("h"):
                ws.row_dimensions[ri + 1].height = float(row["h"]) * 0.75
        thin = Side(style="thin", color="000000")
        for ref, cell in (sh.get("cells") or {}).items():
            r, c = tabular.parse_cell(ref)
            target = ws.cell(r + 1, c + 1)
            v = cell.get("v")
            if v not in (None, ""):
                target.value = _typed(v, cell.get("t"))
            st = cell.get("s") or {}
            if st:
                target.font = Font(bold=bool(st.get("b")), italic=bool(st.get("i")), underline="single" if st.get("u") else None,
                                   size=float(st.get("sz", 11)), color=(st.get("color") or "#000000").lstrip("#").upper()[:6])
                if st.get("bg"):
                    target.fill = PatternFill("solid", fgColor=st["bg"].lstrip("#").upper()[:6])
                target.alignment = Alignment(horizontal=st.get("ha"), vertical=st.get("va") or "center", wrap_text=bool(st.get("wrap")))
                if st.get("bd"):
                    target.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if st.get("nf"):
                    target.number_format = st["nf"]
        for mref in sh.get("merges") or []:
            ws.merge_cells(mref)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    slots: list[dict] = []
    for sh, name in zip(sheets, [s.title for s in wb.worksheets]):
        for f in sh.get("fields") or []:
            req = {"sheet": name, "name": f.get("name"), "kind": f.get("kind", "cell")}
            if req["kind"] == "cell":
                req["cell"] = f["ref"]
            elif req["kind"] == "region":
                req["range"] = f["ref"]
            elif req["kind"] == "repeat":
                req.update(row=int(f["row"]), c1=int(f.get("c1", 0)), c2=int(f.get("c2", 0)), empty_policy=f.get("empty_policy", "remove_row"))
            data, slot = xlsx_add_slot(data, req, slots)
            if f.get("id"):
                slot["id"] = slot["binding"] = f["id"]
            slot["required"] = bool(f.get("required"))
            slots.append(slot)
    # defined names were created with generated ids; rebuild with final ids
    data = _rename_defined_names(data, slots)
    return data, slots


def _rename_defined_names(data: bytes, slots: list[dict]) -> bytes:
    from .ooxml import q
    from .xlsx_doc import XlsxDoc
    x = XlsxDoc(data)
    dn = x.wb.find(q("s", "definedNames"))
    if dn is not None:
        for el in list(dn):
            if (el.get("name") or "").startswith("_dw_"):
                dn.remove(el)
    for s in slots:
        loc = s["locator"]
        ref = loc.get("cell") or loc.get("range") or tabular.range_ref(loc["row"], loc.get("c1", 0), loc["row"], loc.get("c2", 0))
        x.set_defined_name("_dw_" + s["id"], loc["sheet"], ref)
    return x.pkg.save()


def _typed(v, t):
    if t == "number":
        try:
            return float(v) if "." in str(v) else int(v)
        except ValueError:
            return v
    if isinstance(v, str) and v.startswith("="):
        return v  # formula typed in the editor
    return v
