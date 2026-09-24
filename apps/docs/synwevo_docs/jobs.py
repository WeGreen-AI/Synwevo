"""Generation jobs: snapshot inputs, plan and run items in a background thread, verify
outputs, keep state in SQLite so a page refresh or server restart can resume."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import secrets
import socket
import sqlite3
import tempfile
import threading
import time
import traceback
import zipfile
from collections import OrderedDict
from contextlib import contextmanager
from pathlib import Path

from synwevo_core import authority, env
from synwevo_core.context import tenant as current_tenant
from synwevo_core.errors import ServiceError, not_found
from synwevo_core.store import new_id, now

from .engine import generate as G
from .engine import records as R
from .engine.limits import LIMITS
from .engine.plan import build_plan, static_check
from .engine.resolve import SourceDoc
from .services import App, safe_output_name

TOKEN_RE = re.compile(r"\{([^{}]{1,40})\}")


# ------------------------------------------------------------------ naming
def naming_tokens(slots: list[dict], columns: list[str] | None = None) -> list[str]:
    names = ["序号", "样本标识", "来源文件名", "工作表", "块号"]
    for s in slots:
        if s["kind"] not in ("repeat", "region") and s["name"] not in names:
            names.append(s["name"])
    for c in columns or []:
        if c and c not in names:
            names.append(c)
    return names


def render_name(pattern: str, record: dict, plan, slots: list[dict], index: int, width: int,
                column_value=None) -> tuple[str, list[str]]:
    """Build an output file name.  Tokens: built-ins, template field names, then source column
    headers (value of the current sample's row; first row of a group)."""
    warns = []
    by_name = {}
    for s in slots:
        fv = plan.values.get(s["id"]) if plan else None
        if fv and fv.get("kind") == "scalar":
            by_name.setdefault(s["name"], fv["value"].text)

    def sub(m):
        key = m.group(1).strip()
        if key == "序号":
            return str(index + 1).zfill(width)
        if key == "样本标识":
            return record.get("ident") or record.get("label", "")
        if key == "来源文件名":
            return (record.get("file_name") or "").rsplit(".", 1)[0]
        if key == "工作表":
            return record.get("sheet", "")
        if key == "块号":
            return str(record["block"] + 1) if record.get("block") is not None else ""
        if key in by_name:
            v = by_name[key]
            if not v:
                warns.append(f"文件名字段「{key}」为空")
            return v
        if column_value is not None:
            v = column_value(record, key)
            if v is not None:
                return v
        warns.append(f"文件名中的「{{{key}}}」不是可用字段")
        return ""
    name = TOKEN_RE.sub(sub, pattern or "{序号}")
    return safe_output_name(name.replace("\n", " ")), warns


def dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        k = n.lower()
        if k not in seen:
            seen[k] = 1
            out.append(n)
            continue
        i = seen[k] + 1
        while f"{n} ({i})".lower() in seen:
            i += 1
        seen[k] = i
        cand = f"{n} ({i})"
        seen[cand.lower()] = 1
        out.append(cand)
    return out


# ------------------------------------------------------------------ project context
class ProjectContext:
    """Everything needed to build plans for a project at a given moment."""

    def __init__(self, app: App, pid: str):
        self.app = app
        self.project = app.get_project(pid)
        self.rule = self.project.get("record_rule") or {}
        tid = self.project.get("template_id")
        self.template = None
        self.template_missing = False
        if tid:
            try:
                self.template = app.get_template(tid)
            except ServiceError as e:
                if e.status != 404:
                    raise
                # A template deleted while this project still pointed at it (possible before
                # deletion refused templates in use). The project must still open; it behaves
                # as a project with no template chosen, and says so.
                self.template_missing = True
        # A pinned version (PUT /projects/{pid}/template with "version", or an imported
        # definition) is what the caller asked to generate from. Using head instead meant any
        # later edit of the template silently changed every subsequent job.
        self.pinned_version = self.project.get("template_version") if self.template else None
        self.version = (app.get_version(tid, self.pinned_version or self.template["head"])
                        if self.template else None)
        self.slots = (self.version or {}).get("slots") or []
        self.mappings = app.get_mappings(pid)["mappings"]
        self.source_rows = app.project_sources(pid)
        self.all_source_rows = app.project_sources(pid, ready_only=False)
        self.srcs: dict[str, SourceDoc] = {r["id"]: app._doc_from_row(r) for r in self.source_rows}
        self._records = None

    def signature(self) -> str:
        parts = [self.project["rev"], self.project["mapping_rev"], (self.template or {}).get("id"), (self.version or {}).get("version_no"),
                 json.dumps(self.rule, sort_keys=True), [(r["id"], r["version"], r["blob"]) for r in self.source_rows]]
        return hashlib.sha1(json.dumps(parts, default=str).encode()).hexdigest()[:16]

    @property
    def records(self) -> dict:
        if self._records is None:
            if not self.rule.get("mode"):
                raise ServiceError(400, "no_rule", "请先设置样本模式")
            try:
                self._records = R.build_records([self.srcs[r["id"]] for r in self.source_rows], self.rule)
            except ValueError as e:
                raise ServiceError(400, "bad_rule", str(e)) from e
            _name_records(self._records["records"], {r["id"]: r["filename"] for r in self.source_rows})
        return self._records

    def plan(self, record: dict):
        return build_plan(record, self.rule, self.srcs, self.mappings, self.slots)

    def column_headers(self) -> list[str]:
        return column_headers(self.rule, self.srcs, self.source_rows)

    def column_value(self, record: dict, header: str):
        return record_column_value(self.rule, self.srcs, record, header)

    def working(self) -> bytes:
        return self.app.store.get_blob(self.version["working_blob"])


def _name_records(records: list[dict], names: dict[str, str]) -> None:
    for rec in records:
        rec["file_name"] = names.get(rec["file_ids"][0] if rec.get("file_ids") else rec["file_id"], "")


def column_headers(rule: dict, srcs: dict, rows: list[dict]) -> list[str]:
    if rule.get("mode") not in ("per_row", "per_group") or not rule.get("table"):
        return []
    from .engine.resolve import ResolveError
    for r in rows:
        src = srcs.get(r["id"])
        if src is None or src.workbook_json is None:
            continue
        try:
            return [h for h in R.cached_table_view(src, rule["table"]).headers if h]
        except ResolveError:
            continue
    return []


def record_column_value(rule: dict, srcs: dict, record: dict, header: str):
    if rule.get("mode") not in ("per_row", "per_group"):
        return None
    from .engine import tabular
    from .engine.resolve import ResolveError
    fid = record.get("file_id")
    src = srcs.get(fid)
    if src is None or src.workbook_json is None:
        return None
    try:
        # The view is cached on the document itself, so it can never be a view of another
        # file version, and the cache goes away with the document.
        tv = R.cached_table_view(src, rule["table"])
    except ResolveError:
        return None
    idx = tv.header_index().get(tabular.normalize_header(header))
    if not idx or len(idx) > 1:
        return None
    row = record.get("row")
    if row is None and record.get("rows"):
        row = record["rows"][0][2]
    if row is None:
        return None
    return tv.sheet.cell(row, idx[0]).text



# ------------------------------------------------------------------ verification
def verify_output(fmt: str, data: bytes, slots: list[dict], plan) -> dict:
    """Reopen the produced file and check that filled values are present."""
    res = {"reopened": False, "checks": 0, "missing": [], "engine": None}
    try:
        if fmt == "docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            from .engine.docx_render import render_docx
            rend = render_docx(data)
            text = "\n".join(_blocks_text(p["blocks"]) for p in rend["parts"])
            res.update(reopened=True, engine="python-docx " + getattr(docx, "__version__", ""), paragraphs=len(d.paragraphs))
        elif fmt == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data))
            text = "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
            res.update(reopened=True, engine="openpyxl", sheets=len(wb.worksheets))
        elif fmt == "pdf":
            from .engine import pdf_engine as pe
            text = pe.pdf_text(data)
            res.update(reopened=True, engine={"pymupdf": "PyMuPDF", "pdfium": "PDFium"}[pe.ENGINE], pages=pe.page_count(data))
        elif fmt == "pptx":
            import pptx
            prs = pptx.Presentation(io.BytesIO(data))
            from .engine.pptx_engine import render_pptx
            rend = render_pptx(data)
            text = "\n".join(p["text"] for s in rend["slides"] for sh in s["shapes"] for p in sh.get("paras", [])) + "\n" + \
                "\n".join(p["text"] for s in rend["slides"] for sh in s["shapes"] for row in sh.get("rows", []) for c in row for p in c["paras"])
            res.update(reopened=True, engine="python-pptx", slides=len(prs.slides))
        elif fmt == "csv":
            text = data.decode("utf-8-sig")
            res.update(reopened=True, engine="csv")
        elif fmt in ("png", "jpg"):
            from PIL import Image
            im = Image.open(io.BytesIO(data))
            im.verify()
            im = Image.open(io.BytesIO(data))
            res.update(reopened=True, engine="Pillow", size=list(im.size),
                       note="图片中的文字无法直接读取：已在栅格化前的页面中逐值核对（未写入则该样本失败）")
            return res
        else:
            return res
    except Exception as e:  # noqa: BLE001
        res["error"] = f"输出文件无法重新打开：{e}"
        return res
    norm = _squash(text)
    for s in slots:
        fv = plan.values.get(s["id"]) if plan else None
        if not fv:
            continue
        vals = []
        if fv["kind"] == "scalar":
            vals = [fv["value"].text]
        elif fv["kind"] == "rows":
            vals = [v.text for row in fv["rows"][:50] for v in row.values()]
        elif fv["kind"] == "matrix":
            vals = [c.text for row in fv["rows"] for c in row]
        for v in vals:
            if not v or (fmt == "xlsx" and _numeric(v)):
                continue
            res["checks"] += 1
            if _squash(v) not in norm:
                res["missing"].append({"slot": s["name"], "value": v[:80]})
    return res


def _numeric(s: str) -> bool:
    return bool(re.fullmatch(r"-?[\d,]+(\.\d+)?", s.strip()))


def _squash(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def _blocks_text(blocks) -> str:
    out = []
    for b in blocks:
        if b["t"] == "p":
            out.append(b["text"])
        else:
            for row in b["rows"]:
                for c in row["cells"]:
                    out.append(_blocks_text(c["blocks"]))
    return "\n".join(out)



# ------------------------------------------------------------------ job service
LEASE_S = float(env.docs("JOB_LEASE_S", "30"))
IDEMPOTENCY_TTL_S = float(env.docs("IDEMPOTENCY_TTL_HOURS", "24")) * 3600
TERMINAL = ("completed", "completed_with_errors", "failed", "cancelled")
ACTIVE = ("queued", "running", "cancelling")
EVENT_VERSION = 1
QUEUE_RETRY_AFTER_S = 30
# Housekeeping the supervisor does every ten minutes (the first time right after start):
# expired idempotency records of both families, expired resumable uploads, old temp files.
MAINTENANCE_S = 600.0

# A preflight simulation runs the real renderer inside the request, so it is bounded: at most
# ``simulate_max()`` samples per project state (``simulate=true`` means the first
# SIMULATE_DEFAULT), at most SIMULATE_BUDGET_S seconds of rendering per request, one request per
# project state at a time and SIMULATE_CONCURRENCY in the whole process. Whatever it has rendered
# is kept per project state, so asking again (or for more) continues instead of starting over.
SIMULATE_DEFAULT = 50
SIMULATE_HARD_MAX = 200
SIMULATE_BUDGET_S = 60.0
SIMULATE_WAIT_S = 5.0
SIMULATE_CONCURRENCY = 2
SIMULATE_RETRY_AFTER_S = 5

# Why a job stopped because the party that asked for it lost the right to. Shown on the items
# that were not generated, so the person reading the report knows it was not a failure of theirs.
HALT_MESSAGES = {
    "credential_revoked": "提交任务所用的密钥或客户端已被吊销或停用",
    "scope_withdrawn": "提交任务所用凭据的写权限已被收回",
    "workspace_frozen": "工作空间已被冻结",
    "workspace_deleted": "工作空间已被删除",
    "workspace_unavailable": "工作空间当前不可用",
    "workspace_mismatch": "提交者已不属于这个工作空间",
    "not_a_member": "提交者已不是这个工作空间的成员",
    "role_insufficient": "提交者的角色已没有写权限",
    "account_unavailable": "提交者的账户已停用或删除",
    "entitlement_revoked": "提交者的应用使用资格已被撤销",
    "no_entitlement": "提交者没有这个应用的使用资格",
    "app_disabled": "应用已被停用",
    "app_not_enabled_for_workspace": "这个工作空间已不再启用该应用",
    "grant_revoked": "提交者对该客户端的授权已被撤销",
    "wrong_audience": "提交所用的授权属于另一个应用",
    "unknown_principal": "无法识别提交者的身份",
    "project_deleted": "项目已被删除",
    "restored_from_backup": "服务已从备份恢复，恢复前未完成的任务不会自动续跑；需要时请重试",
}


def halt_message(halt_reason: str | None) -> str | None:
    """What a stored ``halt_reason`` means, for the person reading the job or its report."""
    if not halt_reason:
        return None
    why = halt_reason.split(":", 1)[1] if halt_reason.startswith("authorization_revoked:") else halt_reason
    return HALT_MESSAGES.get(why, why)


def _env_int(name: str, default: int, lo: int = 1) -> int:
    """Read at call time: an operator changing the value (or a test) takes effect without a restart."""
    try:
        return max(lo, int(env.docs(name, str(default))))
    except (TypeError, ValueError):
        return default


def sync_plan_max() -> int:
    """Up to this many samples are planned inside the submitting request; larger batches are
    planned by the runner in slices, so a 20,000-sample submission does not hold a request open."""
    return _env_int("SYNC_PLAN_MAX", 1000, lo=0)


def max_active_jobs() -> int:
    return _env_int("MAX_ACTIVE_JOBS", 10)


def plan_batch() -> int:
    return _env_int("PLAN_BATCH", 500)


def runner_slots() -> int:
    return _env_int("RUNNER_SLOTS", 4)


def simulate_max() -> int:
    """The most samples one preflight simulation renders (``SYNWEVO_DOCS_SIMULATE_MAX``, never
    above SIMULATE_HARD_MAX). Anything larger is what submitting the job is for: every item of a
    job is rendered and verified by the runner, outside any request."""
    return min(SIMULATE_HARD_MAX, _env_int("SIMULATE_MAX", SIMULATE_HARD_MAX))


def slice_seconds() -> float:
    try:
        return max(0.05, float(env.docs("JOB_SLICE_S", "1.0")))
    except (TypeError, ValueError):
        return 1.0


def request_hash(obj) -> str:
    """Stable hash of a normalised request (key order and whitespace do not matter)."""
    return hashlib.sha256(json.dumps(obj, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def _pid_alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    except OSError:
        return False
    return True


def _pid_namespace() -> str:
    """Which PID namespace this process lives in ("-" where there is no such thing).

    Two containers can share a hostname (network_mode: host) and still have separate PID
    namespaces; there a live worker's PID looks dead from next door, and the lease used to be
    taken from under it at once. A PID is only evidence inside the namespace it was issued in.
    """
    try:
        return format(os.stat("/proc/self/ns/pid").st_ino, "x")
    except OSError:
        return "-"


def _idem_owner(desc: dict | None) -> str:
    """Whose idempotency namespace a key belongs to.

    Keys are chosen by clients, and two integrations in one workspace can easily pick the same
    string; they must not receive each other's job. An API key is identified by its lineage, so a
    rotated key still replays what its predecessor submitted. Operator requests (the legacy
    browser modes) keep the workspace-wide namespace they always had.
    """
    d = desc or {}
    k = d.get("kind")
    if k == "api_key":
        return f"api:{d.get('lineage') or d.get('key_id')}"
    if k == "service_account":
        return f"svc:{d.get('client_id')}"
    if k == "delegated":
        return f"dlg:{d.get('client_id')}:{d.get('user_id')}"
    if k == "user":
        return f"user:{d.get('user_id')}"
    return ""


def _csv_cell(v):
    """Neutralise spreadsheet formulas in text taken from sources and file names.

    The report is opened in Excel by the people who run the batch, and a sample label is a cell
    value from someone else's file: "=HYPERLINK(…)" or a DDE payload there would execute. A
    leading quote makes Excel show the text as text (OWASP CSV injection guidance)."""
    if isinstance(v, str) and v and v[0] in "=+-@\t\r":
        return "'" + v
    return v


def _loads(raw):
    if raw is None or isinstance(raw, (dict, list)):
        return raw
    try:
        return json.loads(raw)
    except (TypeError, ValueError):
        return None


class _Lease:
    """A job this runner holds, and what it needs to keep working on it between slices."""

    def __init__(self, jid: str, tenant: str, created_at: float):
        self.jid = jid
        self.tenant = tenant or "default"
        self.created_at = created_at or 0.0
        self.served_at = 0.0
        self.cursor = 0                      # next item seq to look at: finding work stays O(1)
        self.loaded = False
        self.records: list[dict] | None = None


class JobService:
    """Runs generation jobs in a background thread.

    Reliability model (see docs/integration/api.md「任务语义」):

    * a job is *claimed* with an atomic ``UPDATE … WHERE status='queued'`` and held with a
      lease that a heartbeat renews; a job whose lease expired (worker gone) is recovered:
      unfinished items go back to ``pending`` and the job to ``queued`` (or ``cancelled``);
    * an item's output files and its ``success`` status are published in one transaction –
      an interruption leaves either nothing or a complete result, never a half result;
    * the terminal status is written with a conditional update inside a transaction that also
      reads the cancel flag, so a cancel and the completion of the last item cannot both win;
    * every state change appends an event (``job_events``) that pollers and webhooks read.

    Scheduling:

    * large batches are planned by the runner (``phase='planning'``) in batches whose progress is
      visible as ``counts.planned`` / ``counts.total``; a crash resumes from what was stored;
    * the runner holds several jobs and works on them in short slices, round-robin by workspace,
      so one workspace's 20,000-sample batch does not stop everyone else's jobs;
    * a workspace can have at most ``SYNWEVO_DOCS_MAX_ACTIVE_JOBS`` jobs queued or running,
      checked inside the transaction that inserts the job;
    * the authority of whoever asked for the work is re-checked when a run starts, between
      samples and immediately before anything is published (synwevo_core.authority).
    """

    def __init__(self, app: App, start: bool = True):
        self.app = app
        self.store = app.store
        self._ns = _pid_namespace()
        self.worker_id = f"{socket.gethostname()}:{os.getpid()}:{secrets.token_hex(3)}:{self._ns}"
        self._wake = threading.Event()
        self._stop = False
        self._thread: threading.Thread | None = None
        self._supervisor: threading.Thread | None = None
        self._held: dict[str, _Lease] = {}
        self._held_lock = threading.Lock()
        self._last_tenant: str | None = None
        self._last_fill = 0.0
        self._last_maintenance = 0.0
        self._analysis_cache: OrderedDict = OrderedDict()
        self._analysis_lock = threading.Lock()
        self._sim_flights: dict = {}                    # project state -> [lock, users]
        self._sim_flights_lock = threading.Lock()
        self._sim_slots = threading.BoundedSemaphore(SIMULATE_CONCURRENCY)
        self.last_heartbeat = 0.0
        self.tmp_dir = Path(self.store.root) / "tmp"
        self.tmp_dir.mkdir(exist_ok=True)
        self.recover()
        if start:
            self._start_runner()
            self._supervisor = threading.Thread(target=self._supervise, name="job-supervisor", daemon=True)
            self._supervisor.start()

    # -------------------------------------------------------------- lifecycle
    def _start_runner(self) -> None:
        self._thread = threading.Thread(target=self._loop, name="job-runner", daemon=True)
        self._thread.start()

    def runner_alive(self) -> bool:
        return bool(self._thread and self._thread.is_alive())

    def stop(self):
        self._stop = True
        self._wake.set()

    def held(self) -> list[str]:
        with self._held_lock:
            return list(self._held)

    def _supervise(self) -> None:
        """Heartbeat for held leases, recovery of jobs whose worker is gone, restart of a dead
        runner thread and, every MAINTENANCE_S, :meth:`maintain`."""
        tick = max(0.5, min(5.0, LEASE_S / 3))
        while not self._stop:
            try:
                t = now()
                # Renew only what the runner actually holds. Renewing every lease this worker id
                # owns kept a job orphaned by a crashed runner thread "running" for as long as
                # the process lived.
                ids = self.held()
                if ids:
                    self.store.execute(f"UPDATE jobs SET lease_until=? WHERE lease_owner=? AND status IN ('running','cancelling') "
                                       f"AND id IN ({','.join('?' * len(ids))})", (t + LEASE_S, self.worker_id, *ids))
                self.last_heartbeat = t
                self.recover()
                if not self.runner_alive() and not self._stop:
                    self._release_held()
                    self._start_runner()
                if t - self._last_maintenance >= MAINTENANCE_S:
                    self._last_maintenance = t
                    self.maintain(t)
            except Exception:  # noqa: BLE001 - the supervisor must keep running; errors surface through /readyz
                traceback.print_exc()
            time.sleep(tick)

    def maintain(self, t: float | None = None) -> dict:
        """Housekeeping that keeps tables and the temp directory from growing without bound.

        * job and retry idempotency keys past their TTL (only the two namespaces owned here);
        * the machine API's stored responses past theirs (``api.idempotency.purge_expired``);
        * resumable uploads past their deadline, and their parts (``api.v1_uploads.cleanup_expired``);
        * temporary files (ZIP builds) older than an hour.

        Each part runs on its own, so one failing does not stop the others; the errors go to
        the server log. Returns what was removed, for the caller and for tests."""
        t = now() if t is None else t
        out: dict = {}
        steps = (("idempotency_jobs", lambda: self.store.execute(
                     "DELETE FROM idempotency WHERE created_at < ? AND (scope LIKE 'job:%' OR scope LIKE 'retry:%')",
                     (t - IDEMPOTENCY_TTL_S,))),
                 ("idempotency_api", lambda: self._api_hook("idempotency", "purge_expired", t)),
                 ("uploads_expired", lambda: self._api_hook("v1_uploads", "cleanup_expired", t)),
                 ("tmp_files", lambda: self._sweep_tmp(t)))
        for name, fn in steps:
            try:
                out[name] = fn()
            except Exception:  # noqa: BLE001 - one broken step must not stop the rest
                traceback.print_exc()
                out[name] = None
        return out

    def _api_hook(self, module: str, fn_name: str, t: float):
        # These belong to the machine API; imported late because the api package imports this
        # module, and optional because a deployment may not carry that module.
        import importlib
        try:
            mod = importlib.import_module(f"{__package__}.api.{module}")
        except ImportError:
            return None
        fn = getattr(mod, fn_name, None)
        return fn(self.store, t) if fn else None

    def _sweep_tmp(self, t: float) -> int:
        n = 0
        for f in self.tmp_dir.glob("*"):
            if f.is_file() and f.stat().st_mtime < t - 3600:
                f.unlink(missing_ok=True)
                n += 1
        return n

    def _lease_expired(self, j: dict, t: float) -> bool:
        owner = j.get("lease_owner")
        if not owner:
            return True                                     # pre-lease job or released lease
        if owner == self.worker_id:
            # Ours, but no longer held by the runner (its thread died between claim and
            # release): once the lease runs out nobody renews it, so it is recoverable.
            with self._held_lock:
                held = j["id"] in self._held
            return not held and (j.get("lease_until") or 0) < t
        parts = owner.split(":")
        host, pid = parts[0], (parts[1] if len(parts) > 1 else "")
        ns = parts[3] if len(parts) > 3 else None
        if host == socket.gethostname() and ns == self._ns and pid.isdigit() and not _pid_alive(int(pid)):
            return True                                     # its process is gone: no need to wait for the lease
        return (j.get("lease_until") or 0) < t

    def recover(self) -> list[str]:
        """Resume jobs whose worker stopped: running items back to pending, the job back to
        the queue (or to ``cancelled`` when a cancel had been accepted)."""
        t = now()
        recovered = []
        rows = self.store.all("SELECT id, status, lease_owner, lease_until FROM jobs WHERE status IN ('running','cancelling')")
        for j in rows:
            if not self._lease_expired(j, t):
                continue
            with self.store.tx() as c:
                cur = c.execute("SELECT status, lease_owner, tenant, run_no FROM jobs WHERE id=?", (j["id"],)).fetchone()
                if cur is None or cur["status"] not in ("running", "cancelling") or cur["lease_owner"] != j["lease_owner"]:
                    continue                                # changed meanwhile
                self._requeue_tx(c, j["id"], cur)
            recovered.append(j["id"])
        if recovered:
            self._wake.set()
        return recovered

    def _requeue_tx(self, c: sqlite3.Connection, jid: str, cur) -> None:
        c.execute("UPDATE job_items SET status='pending' WHERE job_id=? AND status='running'", (jid,))
        if cur["status"] == "cancelling":
            self._finalize_tx(c, jid, expect_owner=cur["lease_owner"])
        else:
            c.execute("UPDATE jobs SET status='queued', lease_owner=NULL, lease_until=NULL WHERE id=?", (jid,))
            self._emit(c, jid, cur["tenant"], "job.recovered", {"status": "queued", "run_no": cur["run_no"],
                                                                "previous_worker": bool(cur["lease_owner"])})

    def _release_held(self) -> None:
        """The runner thread died: give back what it held, so a fresh runner (here or elsewhere)
        picks it up at once instead of after the lease runs out — or never, as it used to."""
        with self._held_lock:
            ids = list(self._held)
            self._held.clear()
        for jid in ids:
            with self.store.tx() as c:
                cur = c.execute("SELECT status, lease_owner, tenant, run_no FROM jobs WHERE id=?", (jid,)).fetchone()
                if cur is None or cur["status"] not in ("running", "cancelling") or cur["lease_owner"] != self.worker_id:
                    continue
                self._requeue_tx(c, jid, cur)
        if ids:
            self._wake.set()

    # -------------------------------------------------------------- events
    def _emit(self, c: sqlite3.Connection, jid: str, tenant: str | None, etype: str, data: dict) -> None:
        c.execute("INSERT INTO job_events (id, job_id, tenant, type, version, occurred_at, data) VALUES (?,?,?,?,?,?,?)",
                  (new_id("ev"), jid, tenant or "default", etype, EVENT_VERSION, now(), json.dumps(data, ensure_ascii=False)))
        hook = getattr(self.app, "on_event", None)
        if hook:
            hook()

    def events(self, jid: str, after: int = 0, limit: int = 200) -> list[dict]:
        rows = self.store.all("SELECT seq, id, job_id, type, version, occurred_at, data FROM job_events WHERE job_id=? AND seq>? ORDER BY seq LIMIT ?",
                              (jid, after, limit))
        return [{"seq": r["seq"], "id": r["id"], "type": r["type"], "version": r["version"], "job_id": r["job_id"],
                 "occurred_at": r["occurred_at"], "data": r["data"]} for r in rows]

    # -------------------------------------------------------------- summary before generation
    @staticmethod
    def _no_template(ctx: ProjectContext) -> ServiceError:
        if ctx.template_missing:
            return ServiceError(400, "no_template", "项目使用的模板已被删除，请重新选择或创建模板")
        return ServiceError(400, "no_template", "请先选择或创建模板")

    def _cache_get(self, key):
        with self._analysis_lock:
            hit = self._analysis_cache.get(key)
            if hit is None:
                return None
            self._analysis_cache.move_to_end(key)
            return hit[0]

    ANALYSIS_BUDGET = 60_000          # samples held across all cached analyses (about 1 KB each)

    def _cache_put(self, key, value, weight: int = 1) -> None:
        # Bounded by the samples held, not by the number of entries: sixteen cached 20,000-sample
        # projects would be hundreds of megabytes on a container limited to 3 GB.
        with self._analysis_lock:
            self._analysis_cache[key] = (value, max(1, weight))
            self._analysis_cache.move_to_end(key)
            while len(self._analysis_cache) > 1 and sum(w for _, w in self._analysis_cache.values()) > self.ANALYSIS_BUDGET:
                self._analysis_cache.popitem(last=False)

    def _analysis(self, ctx: ProjectContext) -> dict:
        """Plan every sample once per project signature.

        The signature covers the project revision, mappings, template version, rule and every
        source version, so any edit produces a new key; the workspace is part of the key because
        two workspaces working from the same files can produce the same signature. Opening the
        submit dialog, changing the naming pattern and paging through the issue list no longer
        re-plan 20,000 samples each time.
        """
        key = ("plan", current_tenant(), ctx.signature())
        hit = self._cache_get(key)
        if hit is not None:
            return hit
        recs = ctx.records
        static = static_check(ctx.mappings, ctx.slots, ctx.rule)
        static_keys = {json.dumps(i, sort_keys=True, default=str) for i in static}
        slot_names = {s["id"]: s.get("name") for s in ctx.slots}
        by_binding = {}
        for s in ctx.slots:
            by_binding.setdefault(s.get("binding") or s["id"], s.get("name"))
        mapping_target = {m.get("id"): by_binding.get((m.get("target") or {}).get("binding") or (m.get("target") or {}).get("slot_id"))
                          for m in ctx.mappings}

        def slot_of(i: dict):
            if i.get("slot"):
                return slot_names.get(i["slot"], i["slot"])
            if i.get("mapping"):
                return mapping_target.get(i["mapping"])
            return None

        def row(rec, i):
            return {"seq": rec["index"] if rec else None, "label": rec["label"] if rec else None, "key": rec["key"] if rec else None,
                    "file_name": rec.get("file_name") if rec else None, "where": rec.get("where") if rec else None,
                    "slot": slot_of(i), "level": i.get("level"), "code": i.get("code"), "message": i.get("message")}

        valid, invalid, warnings = [], [], 0
        issues_by_code: dict[str, int] = {}
        issue_rows = [row(None, i) for i in static]
        for rec in recs["records"]:
            pl = ctx.plan(rec)
            warns = [i for i in pl.issues if i["level"] == "warning"]
            warnings += len(warns)
            blocking = pl.blocking
            if blocking:
                invalid.append((rec, blocking))
                for i in blocking:
                    issues_by_code[i["code"]] = issues_by_code.get(i["code"], 0) + 1
            else:
                valid.append(rec)
            for i in blocking + warns:
                if static_keys and json.dumps(i, sort_keys=True, default=str) in static_keys:
                    continue                          # reported once above, not once per sample
                issue_rows.append(row(rec, i))
        out = {"records": recs, "valid": valid, "invalid": invalid, "warnings": warnings, "issues_by_code": issues_by_code,
               "static": static, "issue_rows": issue_rows}
        self._cache_put(key, out, len(recs["records"]) + len(issue_rows))
        return out

    @contextmanager
    def _sim_flight(self, key):
        """One simulation per project state at a time, and SIMULATE_CONCURRENCY in the process.

        A second request for the same state waits (briefly) for the first and then reads what it
        rendered; beyond that, or when every render slot is busy, the answer is 429 with
        ``retry_after`` rather than a request thread parked for minutes."""
        with self._sim_flights_lock:
            entry = self._sim_flights.setdefault(key, [threading.Lock(), 0])
            entry[1] += 1
        try:
            if not entry[0].acquire(timeout=SIMULATE_WAIT_S):
                raise ServiceError(429, "simulation_busy", "这个项目正在进行另一次模拟生成；请稍后再试",
                                   {"retry_after": SIMULATE_RETRY_AFTER_S})
            try:
                if not self._sim_slots.acquire(timeout=SIMULATE_WAIT_S):
                    raise ServiceError(429, "simulation_busy", "服务器正在进行其他模拟生成；请稍后再试",
                                       {"retry_after": SIMULATE_RETRY_AFTER_S})
                try:
                    yield
                finally:
                    self._sim_slots.release()
            finally:
                entry[0].release()
        finally:
            with self._sim_flights_lock:
                entry[1] -= 1
                if entry[1] <= 0 and self._sim_flights.get(key) is entry:
                    del self._sim_flights[key]

    def _simulate(self, ctx: ProjectContext, ana: dict, limit: int) -> dict:
        """Run the real renderer over the first ``limit`` valid samples, in memory (R01).

        Nothing is stored: no blob, no artifact row, no event, no webhook. What it finds are the
        problems that only appear when the document is actually filled — lost or truncated
        content, overflow, pagination — which used to surface only as failed items after the
        batch had been submitted.

        Bounded, because it runs inside a request that any member may make: the caller caps
        ``limit`` (:func:`simulate_max`), rendering stops after SIMULATE_BUDGET_S, and what was
        rendered is kept per project state — the samples are the same ones in the same order for
        a given signature — so asking for 1, 2, 3 … 200 renders each sample once, not 20,000
        times, and a request cut short by the budget is continued by the next one.
        """
        out_cfg = ctx.project.get("output") or {}
        layout = out_cfg.get("layout") if out_cfg.get("layout") in ("off", "warn", "fit") else "off"
        keep = bool(out_cfg.get("keep_fields"))
        valid = ana["valid"]
        want = min(max(0, int(limit)), len(valid))
        key = ("sim", current_tenant(), ctx.signature(), layout, keep)
        st = self._cache_get(key)
        stopped = None
        if st is None or st["done"] < want:
            with self._sim_flight(key):
                st = self._cache_get(key) or {"done": 0, "failed": [], "warns": []}
                if st["done"] < want:
                    st = self._simulate_more(ctx, valid, st, want, layout, keep)
                    self._cache_put(key, st, 1 + len(st["failed"]))
                    if st["done"] < want:
                        stopped = "time_budget"
        n = min(want, st["done"])
        failed = [(rec, issues) for pos, rec, issues in st["failed"] if pos < n]
        return {"checked": n, "failed": failed, "warnings": sum(st["warns"][:n]), "layout": layout,
                "limited": n < len(valid), "stopped": stopped or ("limit" if n < len(valid) else None)}

    def _simulate_more(self, ctx: ProjectContext, valid: list, st: dict, want: int, layout: str, keep: bool) -> dict:
        """Render valid[st.done:want] (at least one sample, then until the time budget runs out)."""
        fmt = ctx.template["format"]
        native = G.native_format(fmt)
        work = ctx.working()
        opts = {**self.app.with_image_bytes(self.app.fill_options(ctx.template)), "keep_fields": keep}
        tpl_pages = G.template_pages(fmt, work) if layout != "off" else None
        failed, warns = list(st["failed"]), list(st["warns"])
        pos = st["done"]
        deadline = time.monotonic() + SIMULATE_BUDGET_S
        while pos < want and not self._stop:
            rec = valid[pos]
            plan = ctx.plan(rec)
            try:
                data, issues, _lay = G.fill_checked(fmt, work, ctx.slots, plan, opts, layout, tpl_pages)
            except Exception as e:  # noqa: BLE001 - reported on the sample, like the job would
                data, issues = None, [{"level": "error", "code": "exception", "message": f"{type(e).__name__}: {e}"}]
            blocking = [i for i in issues if i.get("level") == "error"]
            if not blocking and not data:
                blocking = [{"level": "error", "code": "render_failed", "message": "填充没有产生文件"}]
            if not blocking:
                ver = verify_output(native, data, ctx.slots, plan)
                if not ver.get("reopened"):
                    blocking = [{"level": "error", "code": "export_failed", "message": ver.get("error") or "输出校验失败"}]
                elif ver["missing"]:
                    blocking = [{"level": "error", "code": "export_failed",
                                 "message": "输出文件中找不到已填写的值：" + "、".join(f"{m['slot']}={m['value']}" for m in ver["missing"][:3])}]
            warns.append(sum(1 for i in issues if i.get("level") == "warning"))
            if blocking:
                failed.append((pos, rec, [{**i, "stage": "render"} for i in blocking]))
            pos += 1
            if time.monotonic() >= deadline:
                break
        return {"done": pos, "failed": failed, "warns": warns}

    def summary(self, pid: str, formats: list[str] | None = None, naming: str | None = None, simulate=False) -> dict:
        """What a job would do, before it is submitted.

        ``simulate`` runs the renderer over the valid samples without storing anything; ``True``
        covers the first SIMULATE_DEFAULT valid samples, an integer the first N — both capped at
        :func:`simulate_max`. ``simulation.limited`` says that not every valid sample was
        rendered, and ``simulation.stopped`` why: ``limit`` (the cap) or ``time_budget`` (asking
        again continues where this one stopped). The job itself renders and verifies every item.
        """
        ctx = ProjectContext(self.app, pid)
        if not ctx.template:
            raise self._no_template(ctx)
        ana = self._analysis(ctx)
        recs = ana["records"]
        valid = list(ana["valid"])
        invalid = list(ana["invalid"])
        issues_by_code = dict(ana["issues_by_code"])
        warnings = ana["warnings"]
        simulation = None
        if simulate:
            cap = simulate_max()
            limit = min(SIMULATE_DEFAULT, cap)
            if not isinstance(simulate, bool):
                try:
                    limit = max(1, min(int(simulate), cap))
                except (TypeError, ValueError) as e:
                    raise ServiceError(400, "bad_value", "simulate 只能是 true/false 或正整数（模拟的样本数）", {"field": "simulate"}) from e
            sim = self._simulate(ctx, ana, limit)
            bad = {r["key"] for r, _ in sim["failed"]}
            valid = [r for r in valid if r["key"] not in bad]
            invalid = invalid + sim["failed"]
            for _, issues in sim["failed"]:
                for i in issues:
                    issues_by_code[i["code"]] = issues_by_code.get(i["code"], 0) + 1
            warnings += sim["warnings"]
            simulation = {"checked": sim["checked"], "failed": len(sim["failed"]), "limited": sim["limited"], "layout": sim["layout"],
                          "limit": limit, "max": cap, "stopped": sim["stopped"]}
        invalid.sort(key=lambda x: x[0]["index"])
        sample_issues = [{"record": rec["label"], "index": rec["index"], "key": rec["key"], "file_name": rec.get("file_name"),
                          "where": rec.get("where"), "issues": issues[:3]} for rec, issues in invalid[:50]]
        exports = self.app._exports_for(ctx.template["format"])
        fmts = formats or [G.native_format(ctx.template["format"])]
        fmt_info = []
        for f in fmts:
            ex = next((e for e in exports if e["format"] == f), None)
            if ex is None:
                fmt_info.append({"format": f, "available": False, "reason": f"不支持导出 {f.upper()}"})
            else:
                fmt_info.append(ex)
        pattern = naming or (ctx.project.get("output") or {}).get("naming") or "{序号}"
        names = []
        width = max(3, len(str(len(recs["records"]))))
        for rec in valid[:5]:
            nm, _ = render_name(pattern, rec, ctx.plan(rec), ctx.slots, rec["index"], width, ctx.column_value)
            names.append(nm)
        n_ok_formats = len([f for f in fmt_info if f.get("available")])
        errors = sum(1 for r in ana["issue_rows"] if r["level"] == "error")
        out = {"input_files": len(ctx.all_source_rows), "ready_files": len(ctx.source_rows),
               "file_errors": recs["file_errors"], "records": len(recs["records"]), "valid": len(valid), "invalid": len(invalid),
               "excluded": recs["excluded"], "warnings": warnings, "record_warnings": recs["warnings"],
               "issues_by_code": issues_by_code, "invalid_samples": sample_issues, "formats": fmt_info,
               "expected_outputs": len(valid) * n_ok_formats, "expected_if_block": len(recs["records"]) * n_ok_formats if not invalid else 0,
               "naming_examples": names, "static_issues": ana["static"], "tokens": naming_tokens(ctx.slots, ctx.column_headers()),
               "signature": ctx.signature(), "template_version": ctx.version["version_no"],
               "template_pinned": bool(ctx.pinned_version),
               "issues_total": {"error": errors, "warning": len(ana["issue_rows"]) - errors},
               "sources_parsing": sum(1 for r in ctx.all_source_rows if r["status"] == "parsing")}
        if simulation is not None:
            out["simulation"] = simulation
        return out

    def preflight_issues(self, pid: str, level: str | None = None, offset: int = 0, limit: int = 100,
                         formats: list[str] | None = None, naming: str | None = None) -> dict:
        """Every located issue of every sample, paged (R02) — not the first 50 with 3 issues each.

        Mapping-level (static) issues are listed once, with no sample, before the per-sample
        rows. ``formats`` and ``naming`` are accepted for symmetry with the summary; they do not
        change which data issues exist.
        """
        if level not in (None, "", "error", "warning"):
            raise ServiceError(400, "bad_value", "level 只能是 error 或 warning", {"field": "level"})
        try:
            offset, limit = int(offset), int(limit)
        except (TypeError, ValueError) as e:
            raise ServiceError(400, "bad_value", "offset / limit 必须是整数", {"field": "offset"}) from e
        if offset < 0:
            raise ServiceError(400, "bad_value", "offset 不能为负数", {"field": "offset"})
        if not 1 <= limit <= 1000:
            raise ServiceError(400, "bad_value", "limit 取值范围是 1–1000", {"field": "limit"})
        ctx = ProjectContext(self.app, pid)
        if not ctx.template:
            raise self._no_template(ctx)
        rows = self._analysis(ctx)["issue_rows"]
        if level:
            rows = [r for r in rows if r["level"] == level]
        page = rows[offset:offset + limit]
        nxt = offset + len(page)
        return {"items": page, "total": len(rows), "offset": offset, "limit": limit, "next_offset": nxt if nxt < len(rows) else None}

    # -------------------------------------------------------------- idempotency
    def _idem_find(self, tenant: str, scopes: list[str], key: str, hashes: tuple, c: sqlite3.Connection | None = None) -> str | None:
        """The job an earlier request with this key created, or None. Raises on a different body.

        ``scopes`` lists the caller's own namespace first and, for callers that have one, the
        workspace-wide namespace used before keys were scoped to the caller — a replay that
        straddles the upgrade must not create a second job."""
        for i, scope in enumerate(scopes):
            sql = "SELECT request_hash, job_id, created_at, principal FROM idempotency WHERE tenant=? AND scope=? AND key=?"
            row = (c.execute(sql, (tenant, scope, key)).fetchone() if c is not None else self.store.one(sql, (tenant, scope, key)))
            if not row:
                continue
            if i > 0 and row["principal"] is not None:
                continue                                  # someone else's workspace-wide key, not a legacy row
            if row["created_at"] < now() - IDEMPOTENCY_TTL_S:
                if c is None:
                    self.store.execute("DELETE FROM idempotency WHERE tenant=? AND scope=? AND key=?", (tenant, scope, key))
                continue
            if row["request_hash"] not in hashes:
                raise ServiceError(409, "idempotency_conflict", "该幂等键已用于另一个不同内容的请求；请换一个新的幂等键，或发送与原请求完全相同的内容",
                                   {"job_id": row["job_id"]})
            return row["job_id"]
        return None

    @staticmethod
    def _idem_scopes(base: str, owner: str) -> list[str]:
        return [f"{base}|{owner}", base] if owner else [base]

    def _admit(self, c: sqlite3.Connection, tenant: str) -> None:
        """Per-workspace admission, inside the transaction that inserts the job.

        Checked there because a count taken before the transaction can be raced: two requests
        both see nine active jobs and both insert the tenth. ``BEGIN IMMEDIATE`` serialises the
        writers, so the count and the insert are one decision."""
        limit = max_active_jobs()
        n = c.execute("SELECT COUNT(*) AS n FROM jobs WHERE status IN ('queued','running','cancelling') AND tenant=?",
                      (tenant,)).fetchone()["n"]
        if n >= limit:
            raise ServiceError(429, "queue_full",
                               f"本工作空间已有 {n} 个任务在排队或进行中（上限 {limit}）；请等部分任务结束后再提交",
                               {"retry_after": QUEUE_RETRY_AFTER_S, "active_jobs": n, "limit": limit})

    # -------------------------------------------------------------- create
    def create(self, pid: str, options: dict, tenant: str | None = None, idem_key: str | None = None, origin: str = "ui",
               principal: dict | None = None) -> dict:
        options = options if isinstance(options, dict) else {}
        tenant = tenant or current_tenant()
        idem_key = idem_key or options.get("request_id")
        desc = authority.loads(principal)
        ctx = ProjectContext(self.app, pid)
        if not ctx.template:
            raise self._no_template(ctx)
        native = G.native_format(ctx.template["format"])
        formats = options.get("formats") or [native]
        if not isinstance(formats, list) or not all(isinstance(f, str) for f in formats):
            raise ServiceError(400, "bad_formats", "formats 必须是格式名称的数组", {"field": "formats"})
        policy = options.get("error_policy", "block")
        if policy not in ("block", "skip"):
            raise ServiceError(400, "bad_error_policy", "error_policy 只能是 block 或 skip", {"field": "error_policy"})
        project_out = ctx.project.get("output") or {}
        # An unknown value must not silently become "off": a caller asking for the strictest
        # pagination check would then get no check at all and a 202 saying it worked.
        if "layout" in options and options["layout"] is not None:
            layout = options["layout"]
            if layout not in ("off", "warn", "fit"):
                raise ServiceError(400, "bad_layout", "layout 只能是 off / warn / fit", {"field": "layout"})
        else:
            layout = project_out.get("layout") or "off"
            if layout not in ("off", "warn", "fit"):
                # The same rule for the project's stored default: it was downgraded to "off".
                raise ServiceError(400, "bad_layout", f"项目默认设置中的 layout「{layout}」无效；请改为 off / warn / fit",
                                   {"field": "output.layout"})
        pattern = options.get("naming") or project_out.get("naming") or "{序号}"
        if not isinstance(pattern, str):
            raise ServiceError(400, "bad_naming", "naming 必须是文本", {"field": "naming"})
        keep = bool(options.get("keep_fields"))
        # The fingerprint is what the client sent, not what the server derived from it: hashing
        # the resolved defaults meant an identical replay was refused as a "different request"
        # after someone changed the project's naming or layout.
        h = request_hash({"project": pid, "formats": options.get("formats") or None, "error_policy": options.get("error_policy"),
                          "naming": options.get("naming") or None, "keep_fields": keep, "layout": options.get("layout")})
        h_legacy = request_hash({"project": pid, "formats": list(formats), "error_policy": policy, "naming": pattern,
                                 "keep_fields": keep, "layout": layout})
        hashes = (h, h_legacy)
        owner = _idem_owner(desc)
        scopes = self._idem_scopes(f"job:{pid}", owner)
        if idem_key:
            existing = self._idem_find(tenant, scopes, idem_key, hashes)
            if existing:
                return {**self.get(existing, tenant), "idempotent_replay": True}
        exports = {e["format"]: e for e in self.app._exports_for(ctx.template["format"])}
        for f in formats:
            if f not in exports or not exports[f].get("available"):
                reason = (exports.get(f) or {}).get("reason") or f"不支持导出 {f.upper()}"
                raise ServiceError(400, "format_unavailable", f"无法导出 {f.upper()}：{reason}", {"field": "formats"})
        parsing = [r["filename"] for r in ctx.all_source_rows if r["status"] == "parsing"]
        if parsing:
            # A job over the files that happen to be ready would complete "successfully" with
            # the others' documents missing and nothing saying so.
            raise ServiceError(409, "sources_not_ready", f"还有 {len(parsing)} 个来源文件正在解析（如「{parsing[0]}」）；请等解析完成后再提交",
                               {"files": parsing[:20], "retry_after": 2})
        recs = ctx.records
        records = recs["records"]
        n = len(records)
        if not n:
            raise ServiceError(400, "no_records", "没有识别到任何样本，无法生成")
        if n > LIMITS.max_records_per_job:
            raise ServiceError(400, "too_many", f"样本数 {n} 超过单次任务上限 {LIMITS.max_records_per_job}")
        width = max(3, len(str(n)))
        sync = n <= sync_plan_max()
        rows: list[tuple] = []
        invalid_n = 0
        if sync:
            plans = [(rec, ctx.plan(rec)) for rec in records]
            invalid = [(rec, pl) for rec, pl in plans if pl.blocking]
            if invalid and policy == "block":
                raise ServiceError(409, "has_errors", f"{len(invalid)} 个样本存在错误，已按「有错误则不开始」停止。请修正或选择「跳过错误样本」",
                                   {"invalid": [{"record": r["label"], "key": r["key"], "index": r["index"], "file_name": r.get("file_name"),
                                                 "where": r.get("where"), "issues": p.blocking[:3]} for r, p in invalid[:50]],
                                    "invalid_total": len(invalid)})
            names = dedupe([render_name(pattern, rec, pl, ctx.slots, rec["index"], width, ctx.column_value)[0] for rec, pl in plans])
            rows = [(rec, pl.blocking, name) for (rec, pl), name in zip(plans, names)]
            invalid_n = len(invalid)
        jid = new_id("j")
        snapshot = {
            "project_id": pid, "signature": ctx.signature(), "rule": ctx.rule, "mappings": ctx.mappings,
            "mapping_rev": ctx.project.get("mapping_rev"),
            "template": {"id": ctx.template["id"], "name": ctx.template["name"], "format": ctx.template["format"],
                         "version_no": ctx.version["version_no"], "pinned": bool(ctx.pinned_version),
                         "working_blob": ctx.version["working_blob"], "slots": ctx.slots,
                         "fill_options": self.app.fill_options(ctx.template)},
            "sources": [{k: r.get(k) for k in ("id", "filename", "version", "blob", "working_blob", "parsed_blob", "format", "family",
                                              "ocr", "corrections")} for r in ctx.source_rows],
            # Files that were uploaded but cannot be read are not part of the job; say which.
            "excluded_files": [{"id": r["id"], "filename": r["filename"], "status": r["status"], "error": r.get("error")}
                               for r in ctx.all_source_rows if r["status"] not in ("ready", "needs_ocr")],
            "records_total": n, "naming_width": width,
        }
        opts = {"formats": formats, "error_policy": policy, "naming": pattern, "keep_fields": keep, "layout": layout}
        phase = "generating" if sync else "planning"
        summary = {"records": n, "planned": n if sync else 0, "invalid": invalid_n if sync else None}
        t = now()
        replay = None
        try:
            with self.store.tx() as c:
                if idem_key:
                    replay = self._idem_find(tenant, scopes, idem_key, hashes, c)
                if replay is None:
                    self._admit(c, tenant)
                    if idem_key:
                        c.execute("INSERT INTO idempotency (tenant, scope, key, request_hash, job_id, created_at, resource_id, principal) "
                                  "VALUES (?,?,?,?,?,?,?,?)", (tenant, scopes[0], idem_key, h, jid, t, jid, owner))
                    c.execute("INSERT INTO jobs (id, project_id, status, created_at, snapshot, options, counts, request_id, tenant, run_no, origin, "
                              "principal, phase, summary) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                              (jid, pid, "queued", t, json.dumps(snapshot, ensure_ascii=False), json.dumps(opts, ensure_ascii=False),
                               "{}", None, tenant, 1, origin, authority.dumps(desc), phase, json.dumps(summary)))
                    c.executemany("INSERT INTO job_items (id, job_id, seq, record_key, record, status, issues, file_base, label, run_no) "
                                  "VALUES (?,?,?,?,?,?,?,?,?,1)",
                                  [(new_id("i"), jid, rec["index"], rec["key"], json.dumps(rec, ensure_ascii=False, default=str),
                                    "skipped" if blocking else "pending", json.dumps(blocking or [], ensure_ascii=False), name, rec["label"])
                                   for rec, blocking, name in rows])
                    c.execute("INSERT INTO job_runs (job_id, run_no, trigger, requested_by, items, created_at, status, principal) VALUES (?,?,?,?,?,?,?,?)",
                              (jid, 1, "create", origin, (n - invalid_n) if sync else None, t, "queued", authority.dumps(desc)))
                    self._update_counts(jid, c)
                    self._emit(c, jid, tenant, "job.created", {"status": "queued", "phase": phase, "run_no": 1, "samples": n,
                                                                "skipped": invalid_n if sync else None, "formats": formats})
        except sqlite3.IntegrityError:
            # another process inserted the same idempotency key between our check and our insert
            replay = self._idem_find(tenant, scopes, idem_key or "", hashes) if idem_key else None
            if not replay:
                raise
        if replay:
            return {**self.get(replay, tenant), "idempotent_replay": True}
        self._wake.set()
        return self.get(jid, tenant)

    # -------------------------------------------------------------- queries
    def get(self, jid: str, tenant: str | None = None) -> dict:
        # The template and source count come out of the snapshot with json_extract: loading and
        # decoding the whole snapshot (mappings, slots, per-source OCR words) on every poll was
        # most of the cost of a status request.
        j = self.store.one(
            "SELECT id, project_id, status, created_at, started_at, finished_at, options, counts, error, cancel_requested, "
            "cancel_requested_at, request_id, tenant, run_no, origin, phase, halt_reason, principal, summary, "
            "json_extract(snapshot, '$.template.id') AS tpl_id, json_extract(snapshot, '$.template.name') AS tpl_name, "
            "json_extract(snapshot, '$.template.format') AS tpl_format, json_extract(snapshot, '$.template.version_no') AS tpl_version, "
            "json_array_length(snapshot, '$.sources') AS source_count FROM jobs WHERE id=?", (jid,))
        tenant = tenant or current_tenant()
        if not j or (j["tenant"] or "default") != tenant:
            not_found("任务")
        j["template"] = {"id": j.pop("tpl_id"), "name": j.pop("tpl_name"), "format": j.pop("tpl_format"), "version_no": j.pop("tpl_version")}
        j["cancel_requested"] = bool(j["cancel_requested"])
        j["phase"] = j.get("phase") or "generating"           # jobs before v5 were planned synchronously
        summary = j.pop("summary", None) or {}
        j["error_code"] = summary.get("error_code") if isinstance(summary, dict) else None
        desc = authority.loads(j.pop("principal", None))
        j["submitted_by"] = authority.label(desc) if desc else None
        j["halt_message"] = halt_message(j.get("halt_reason"))
        return j

    @staticmethod
    def _public_issues(issues):
        # Tracebacks stay in the server log: they carry file paths, library versions and
        # sometimes document text, and were returned to every reader of the job.
        if not isinstance(issues, list):
            return issues
        return [{k: v for k, v in i.items() if k != "trace"} if isinstance(i, dict) else i for i in issues]

    def items(self, jid: str, status: str | None = None, offset: int = 0, limit: int = 500) -> list[dict]:
        sql = "SELECT id, seq, record_key, label, status, issues, attempts, file_base, started_at, finished_at, run_no FROM job_items WHERE job_id=?"
        args: tuple = (jid,)
        if status:
            sql += " AND status=?"
            args += (status,)
        sql += " ORDER BY seq LIMIT ? OFFSET ?"
        args += (limit, offset)
        rows = self.store.all(sql, args)
        ids = [r["id"] for r in rows]
        by_item: dict[str, list] = {}
        if ids:
            arts = self.store.all(f"SELECT id, item_id, format, filename, blob, size, verification, run_no FROM artifacts "
                                  f"WHERE job_id=? AND deleted_at IS NULL AND item_id IN ({','.join('?' * len(ids))})", (jid, *ids))
            for a in arts:
                a["sha256"] = a.pop("blob")
                by_item.setdefault(a["item_id"], []).append(a)
        halted = self._halt_issue(jid)
        for r in rows:
            r["artifacts"] = by_item.get(r["id"], [])
            r["issues"] = self._public_issues(r["issues"])
            if halted and r["status"] == "cancelled" and not r["issues"]:
                # Items stopped by a restore (or anything else that stopped the job without
                # writing a reason on each item) say why, instead of an unexplained "cancelled".
                r["issues"] = [halted]
        return rows

    def _halt_issue(self, jid: str) -> dict | None:
        row = self.store.one("SELECT halt_reason FROM jobs WHERE id=?", (jid,))
        why = (row or {}).get("halt_reason")
        if not why:
            return None
        code = why.split(":", 1)[0]
        return {"level": "warning", "code": code, "stage": "halt", "message": halt_message(why)}

    def count_items(self, jid: str, status: str | None = None) -> int:
        if status:
            return self.store.one("SELECT COUNT(*) AS n FROM job_items WHERE job_id=? AND status=?", (jid, status))["n"]
        return self.store.one("SELECT COUNT(*) AS n FROM job_items WHERE job_id=?", (jid,))["n"]

    def runs(self, jid: str) -> list[dict]:
        rows = self.store.all("SELECT run_no, trigger, requested_by, items, created_at, started_at, finished_at, status, principal "
                              "FROM job_runs WHERE job_id=? ORDER BY run_no", (jid,))
        for r in rows:
            desc = authority.loads(r.pop("principal", None))
            r["requested_by_principal"] = authority.label(desc) if desc else None
        return rows

    def list_for_project(self, pid: str) -> list[dict]:
        return self.store.all("SELECT id, status, phase, halt_reason, created_at, started_at, finished_at, counts, options, run_no FROM jobs "
                              "WHERE project_id=? AND tenant=? ORDER BY created_at DESC LIMIT 50", (pid, current_tenant()))

    def usage(self, tenant: str | None = None) -> dict:
        """What this workspace currently has in the queue, against its limits."""
        tenant = tenant or current_tenant()
        by = {r["status"]: r["n"] for r in self.store.all(
            "SELECT status, COUNT(*) AS n FROM jobs WHERE status IN ('queued','running','cancelling') AND tenant=? GROUP BY status", (tenant,))}
        pending = self.store.one(
            "SELECT COUNT(*) AS n FROM job_items i JOIN jobs j ON j.id=i.job_id "
            "WHERE j.status IN ('queued','running','cancelling') AND j.tenant=? AND i.status IN ('pending','running')", (tenant,))["n"]
        unplanned = 0
        planning = 0
        for r in self.store.all("SELECT summary, counts FROM jobs WHERE status IN ('queued','running','cancelling') AND tenant=? "
                                "AND phase='planning'", (tenant,)):
            planning += 1
            unplanned += max(0, int((r["summary"] or {}).get("records") or 0) - int((r["counts"] or {}).get("planned") or 0))
        return {"active_jobs": sum(by.values()), "queued_jobs": by.get("queued", 0), "running_jobs": by.get("running", 0),
                "cancelling_jobs": by.get("cancelling", 0), "planning_jobs": planning, "pending_items": pending + unplanned,
                "limits": {"max_active_jobs": max_active_jobs(), "max_records_per_job": LIMITS.max_records_per_job,
                           "sync_plan_max": sync_plan_max()}}

    # -------------------------------------------------------------- control
    def cancel(self, jid: str, tenant: str | None = None) -> dict:
        """Cancel.  queued → cancelled at once; running → cancelling (no new item starts; the item
        in progress finishes and keeps its result); terminal → unchanged, reported as such."""
        self.get(jid, tenant)
        accepted, reason = True, None
        with self.store.tx() as c:
            j = c.execute("SELECT status, tenant, run_no FROM jobs WHERE id=?", (jid,)).fetchone()
            st = j["status"]
            if st in TERMINAL:
                accepted, reason = False, "already_terminal"
            elif st == "queued":
                c.execute("UPDATE job_items SET status='cancelled' WHERE job_id=? AND status IN ('pending','running')", (jid,))
                c.execute("UPDATE jobs SET status='cancelled', cancel_requested=1, cancel_requested_at=?, finished_at=? WHERE id=?", (now(), now(), jid))
                c.execute("UPDATE job_runs SET status='cancelled', finished_at=? WHERE job_id=? AND run_no=?", (now(), jid, j["run_no"]))
                counts = self._update_counts(jid, c)
                self._emit(c, jid, j["tenant"], "job.cancelled", {"status": "cancelled", "run_no": j["run_no"], "counts": counts})
            elif st == "running":
                c.execute("UPDATE jobs SET status='cancelling', cancel_requested=1, cancel_requested_at=? WHERE id=?", (now(), jid))
                self._emit(c, jid, j["tenant"], "job.cancel_requested", {"status": "cancelling", "run_no": j["run_no"]})
            else:
                reason = "already_requested"                  # cancelling
        out = self.get(jid, tenant)
        out["cancel"] = {"accepted": accepted, "reason": reason}
        return out

    def retry(self, jid: str, include_skipped: bool = False, tenant: str | None = None, idem_key: str | None = None,
              requested_by: str = "ui", principal: dict | None = None, *, failed_only: bool = False) -> dict:
        for name, value in (("include_skipped", include_skipped), ("failed_only", failed_only)):
            if type(value) is not bool:
                raise ServiceError(400, "invalid", f"{name} 必须是布尔值", {"field": name})
        if failed_only and include_skipped:
            raise ServiceError(400, "invalid", "failed_only 不能与 include_skipped 同时启用", {"field": "include_skipped"})
        j0 = self.get(jid, tenant)
        tenant = j0["tenant"]
        if j0.get("error_code") == "has_errors":
            raise ServiceError(409, "has_errors", "这个任务因样本错误没有开始生成（有错误则不开始）；请修正数据后重新提交，或以「跳过错误样本」提交新任务")
        desc = authority.loads(principal)
        fingerprint = {"job": jid, "include_skipped": include_skipped}
        if failed_only:
            fingerprint["failed_only"] = True     # default keeps pre-existing fingerprints valid
        h = request_hash(fingerprint)
        owner = _idem_owner(desc)
        scopes = self._idem_scopes(f"retry:{jid}", owner)
        if idem_key:
            existing = self._idem_find(tenant, scopes, idem_key, (h,))
            if existing:
                return {**self.get(jid, tenant), "idempotent_replay": True}
        statuses = ("failed",) if failed_only else ("failed", "cancelled") + (("skipped",) if include_skipped else ())
        replay = None
        try:
            with self.store.tx() as c:
                # The key is looked at again inside the transaction and *before* the active
                # check: the loser of two concurrent retries with one key used to get
                # 409 job_active, because the winner had already re-queued the job.
                if idem_key:
                    replay = self._idem_find(tenant, scopes, idem_key, (h,), c)
                if replay is None:
                    j = c.execute("SELECT status, tenant, run_no, phase, summary FROM jobs WHERE id=?", (jid,)).fetchone()
                    if j["status"] in ACTIVE:
                        raise ServiceError(409, "job_active", "任务仍在进行中；请等它结束（或取消）后再重试失败项")
                    self._admit(c, tenant)
                    run_no = (j["run_no"] or 1) + 1
                    n = c.execute(f"UPDATE job_items SET status='pending', run_no=? WHERE job_id=? AND status IN ({','.join('?' * len(statuses))})",
                                  (run_no, jid) + statuses).rowcount
                    summary = _loads(j["summary"]) or {}
                    planned = c.execute("SELECT COUNT(*) AS n FROM job_items WHERE job_id=?", (jid,)).fetchone()["n"]
                    unplanned = max(0, int(summary.get("records") or 0) - planned) if j["phase"] == "planning" and not failed_only else 0
                    if not n and not unplanned:
                        raise ServiceError(409, "nothing_to_retry", "没有失败的样本可重试" if failed_only else "没有失败或已取消的样本可重试")
                    if idem_key:
                        c.execute("INSERT INTO idempotency (tenant, scope, key, request_hash, job_id, created_at, resource_id, principal) "
                                  "VALUES (?,?,?,?,?,?,?,?)", (tenant, scopes[0], idem_key, h, jid, now(), jid, owner))
                    # halt_reason is cleared: the new run is authorised by whoever asked for it.
                    c.execute("UPDATE jobs SET status='queued', run_no=?, cancel_requested=0, cancel_requested_at=NULL, finished_at=NULL, "
                              "error=NULL, halt_reason=NULL WHERE id=?", (run_no, jid))
                    if failed_only:
                        c.execute("UPDATE jobs SET phase='generating' WHERE id=?", (jid,))
                    c.execute("INSERT INTO job_runs (job_id, run_no, trigger, requested_by, items, created_at, status, principal) VALUES (?,?,?,?,?,?,?,?)",
                              (jid, run_no, "retry", requested_by, n + unplanned, now(), "queued", authority.dumps(desc)))
                    counts = self._update_counts(jid, c)
                    self._emit(c, jid, tenant, "job.retry_requested", {"status": "queued", "run_no": run_no, "items": n + unplanned, "counts": counts})
        except sqlite3.IntegrityError:
            replay = self._idem_find(tenant, scopes, idem_key or "", (h,)) if idem_key else None
            if not replay:
                raise
        if replay:
            return {**self.get(jid, tenant), "idempotent_replay": True}
        self._wake.set()
        return self.get(jid, tenant)

    def _update_counts(self, jid: str, c: sqlite3.Connection | None = None) -> dict:
        def run(c):
            rows = c.execute("SELECT status, COUNT(*) AS n FROM job_items WHERE job_id=? GROUP BY status", (jid,)).fetchall()
            cnt = {r["status"]: r["n"] for r in rows}
            planned = sum(cnt.values())
            srow = c.execute("SELECT summary FROM jobs WHERE id=?", (jid,)).fetchone()
            records = int((_loads(srow["summary"] if srow else None) or {}).get("records") or 0)
            cnt["planned"] = planned
            cnt["total"] = max(planned, records)
            if records > planned:
                cnt["unplanned"] = records - planned
            cnt["files"] = c.execute("SELECT COUNT(*) AS n FROM artifacts WHERE job_id=? AND deleted_at IS NULL", (jid,)).fetchone()["n"]
            cnt["documents"] = cnt.get("success", 0)
            c.execute("UPDATE jobs SET counts=? WHERE id=?", (json.dumps(cnt), jid))
            return cnt
        if c is not None:
            return run(c)
        with self.store.tx() as c2:
            return run(c2)

    # -------------------------------------------------------------- authority
    def _authorized(self, jid: str) -> tuple[bool, str | None]:
        """Is whoever asked for the current run of this job still allowed to have it done?

        The run's own principal wins (a retry is authorised by whoever retried); jobs created
        before descriptors existed have none and count as operator work."""
        j = self.store.one("SELECT tenant, principal, run_no, project_id FROM jobs WHERE id=?", (jid,))
        if not j:
            return False, "job_missing"
        p = self.store.one("SELECT deleted FROM projects WHERE id=?", (j["project_id"],))
        if not p or p["deleted"]:
            return False, "project_deleted"
        run = self.store.one("SELECT principal FROM job_runs WHERE job_id=? AND run_no=?", (jid, j["run_no"] or 1))
        raw = (run or {}).get("principal")
        desc = authority.loads(j["principal"] if raw is None or raw == "" else raw)
        return authority.check(self.store, desc, need="write", tenant=j["tenant"] or "default")

    def _halt(self, jid: str, why: str | None) -> bool:
        """Stop a job whose submitter lost the right to it: nothing more is generated or published.

        The unfinished samples are cancelled with a reason the reader can act on, and the event
        carries ``halt_reason`` so a webhook consumer can tell this apart from a user's cancel."""
        why = why or "unknown_principal"
        deleted = why == "project_deleted"
        halt_reason = "project_deleted" if deleted else f"authorization_revoked:{why}"
        code = "project_deleted" if deleted else "authorization_revoked"
        detail = HALT_MESSAGES.get(why, why)
        issue = {"level": "error", "code": code, "stage": "authorize",
                 "message": f"{detail}，其余样本不再生成" if not deleted else "项目已被删除，其余样本不再生成"}
        t = now()
        with self.store.tx() as c:
            j = c.execute("SELECT status, tenant, run_no, lease_owner FROM jobs WHERE id=?", (jid,)).fetchone()
            if not j or j["status"] in TERMINAL:
                return False
            if j["status"] in ("running", "cancelling") and j["lease_owner"] != self.worker_id:
                return False                                   # another worker holds it and will decide
            c.execute("UPDATE job_items SET status='cancelled', issues=?, finished_at=? WHERE job_id=? AND status IN ('pending','running')",
                      (json.dumps([issue], ensure_ascii=False), t, jid))
            counts = self._update_counts(jid, c)
            c.execute("UPDATE jobs SET status='cancelled', halt_reason=?, finished_at=?, lease_owner=NULL, lease_until=NULL WHERE id=?",
                      (halt_reason, t, jid))
            c.execute("UPDATE job_runs SET status='cancelled', finished_at=? WHERE job_id=? AND run_no=?", (t, jid, j["run_no"]))
            self._emit(c, jid, j["tenant"], "job.cancelled", {"status": "cancelled", "run_no": j["run_no"], "counts": counts,
                                                               "halt_reason": halt_reason})
        with self._held_lock:
            self._held.pop(jid, None)
        return True

    # -------------------------------------------------------------- runner
    def _claim(self, jid: str | None = None) -> str | None:
        """Claim one queued job (the given one, or the oldest). Returns its id or None."""
        if jid is None:
            row = self.store.one("SELECT id FROM jobs WHERE status='queued' ORDER BY created_at LIMIT 1")
            if not row:
                return None
            jid = row["id"]
        row = self.store.one("SELECT id, tenant, created_at FROM jobs WHERE id=? AND status='queued'", (jid,))
        if not row:
            return None
        # Before the run starts: a job queued by a credential revoked since must not start at all.
        ok, why = self._authorized(jid)
        if not ok:
            self._halt(jid, why)
            return None
        t = now()
        with self.store.tx() as c:
            n = c.execute("UPDATE jobs SET status='running', lease_owner=?, lease_until=?, started_at=COALESCE(started_at, ?) "
                          "WHERE id=? AND status='queued'", (self.worker_id, t + LEASE_S, t, jid)).rowcount
            if not n:
                return None                                  # another worker claimed it first
            j = c.execute("SELECT tenant, run_no, phase FROM jobs WHERE id=?", (jid,)).fetchone()
            run = c.execute("SELECT started_at FROM job_runs WHERE job_id=? AND run_no=?", (jid, j["run_no"])).fetchone()
            c.execute("UPDATE job_runs SET status='running', started_at=COALESCE(started_at, ?) WHERE job_id=? AND run_no=?",
                      (t, jid, j["run_no"]))
            if run is None or run["started_at"] is None:
                # Once per run. A job resumed after recovery already said it started; saying it
                # again made consumers count one run twice.
                self._emit(c, jid, j["tenant"], "job.started", {"status": "running", "run_no": j["run_no"],
                                                                 "phase": j["phase"] or "generating"})
        with self._held_lock:
            self._held[jid] = _Lease(jid, row["tenant"], row["created_at"])
        return jid

    def held_max(self) -> int:
        """How many jobs this runner holds at most: the slots, and room for one job of each of
        a few more workspaces so a long batch does not make everyone else wait for it. Every
        held job keeps its snapshot's sources and template in memory and shows as running."""
        return runner_slots() * 4

    def _fill(self) -> None:
        """Take on queued work, fairly across workspaces.

        A workspace with nothing running here gets its oldest job claimed even when the slots
        are full — that is what stops one workspace's long batch from making everyone else wait
        for it to finish — but never past :meth:`held_max`: the bound is checked before every
        claim, not once before the loop (sixty queued workspaces used to become sixty held,
        "running" jobs). Workspaces beyond it are taken on, oldest first, as held jobs finish.
        The remaining free slots then go to the oldest queued jobs."""
        self._last_fill = time.monotonic()
        cap = runner_slots()
        hard = self.held_max()
        with self._held_lock:
            held_tenants = {l.tenant for l in self._held.values()}
            n_held = len(self._held)
        if n_held >= hard:
            return
        # The oldest queued job of each workspace, oldest workspace first.
        firsts = self.store.all(
            "SELECT id, tenant FROM (SELECT id, COALESCE(tenant, 'default') AS tenant, created_at, "
            "ROW_NUMBER() OVER (PARTITION BY COALESCE(tenant, 'default') ORDER BY created_at, id) AS rn "
            "FROM jobs WHERE status='queued') WHERE rn=1 ORDER BY created_at, id LIMIT 200")
        for r in firsts:
            if len(self.held()) >= hard:
                break
            if r["tenant"] in held_tenants:
                continue
            if self._claim(r["id"]):
                held_tenants.add(r["tenant"])
        if len(self.held()) >= cap:
            return
        for r in self.store.all("SELECT id FROM jobs WHERE status='queued' ORDER BY created_at, id LIMIT 200"):
            if len(self.held()) >= cap:
                break
            self._claim(r["id"])

    def _next_lease(self) -> _Lease | None:
        """Round-robin by workspace; within a workspace, the job served least recently."""
        with self._held_lock:
            held = list(self._held.values())
        if not held:
            return None
        tenants = sorted({l.tenant for l in held})
        nxt = next((t for t in tenants if self._last_tenant is not None and t > self._last_tenant), tenants[0])
        lease = min((l for l in held if l.tenant == nxt), key=lambda l: (l.served_at, l.created_at))
        self._last_tenant = nxt
        lease.served_at = time.monotonic()
        return lease

    def _loop(self) -> None:
        while not self._stop:
            if self._wake.is_set() or time.monotonic() - self._last_fill > 1.0:
                self._wake.clear()
                self._fill()
            lease = self._next_lease()
            if lease is None:
                self._wake.wait(1.0)
                continue
            try:
                done = self._slice(lease)
            except Exception as e:  # noqa: BLE001 - job-level failure (inputs unreadable…): recorded, runner continues
                traceback.print_exc()
                self._fail_job(lease.jid, e)
                done = True
            if done:
                with self._held_lock:
                    self._held.pop(lease.jid, None)

    def _run(self, jid: str) -> None:
        """Run a claimed job to its end in this thread (used by tools and tests; the runner
        interleaves jobs in slices instead)."""
        with self._held_lock:
            lease = self._held.get(jid)
        if lease is None:
            j = self.store.one("SELECT tenant, created_at FROM jobs WHERE id=?", (jid,))
            lease = _Lease(jid, j["tenant"], j["created_at"])
            with self._held_lock:
                self._held[jid] = lease
        try:
            while not self._stop and not self._slice(lease):
                pass
        finally:
            with self._held_lock:
                self._held.pop(jid, None)

    def _slice(self, lease: _Lease) -> bool:
        """Work on one held job for a short while. Returns True when the runner should let go of it."""
        ctl = self.store.one("SELECT status, cancel_requested, lease_owner, phase FROM jobs WHERE id=?", (lease.jid,))
        if not ctl or ctl["lease_owner"] != self.worker_id or ctl["status"] not in ("running", "cancelling"):
            return True                                       # lease lost or the job ended elsewhere
        if ctl["cancel_requested"]:
            with self.store.tx() as c:
                self._finalize_tx(c, lease.jid, self.worker_id)
            return True
        self._load(lease)
        if (ctl["phase"] or "generating") == "planning":
            done = self._plan_slice(lease)
        else:
            done = self._generate_slice(lease)
        if not done:
            self.store.execute("UPDATE jobs SET last_slice_at=? WHERE id=? AND lease_owner=?", (now(), lease.jid, self.worker_id))
        return done

    def _load(self, L: _Lease) -> None:
        L.run_no = (self.store.one("SELECT run_no FROM jobs WHERE id=?", (L.jid,)) or {}).get("run_no") or 1
        if L.loaded:
            return                                        # the snapshot is decoded once per lease, not per slice
        j = self.store.one("SELECT snapshot, options FROM jobs WHERE id=?", (L.jid,))
        snap, opts = j["snapshot"], j["options"]
        srcs = {s["id"]: self.app._doc_from_row({**s, "family": s["family"]}) for s in snap["sources"]}
        L.snap, L.opts, L.srcs = snap, opts, srcs
        L.tpl = snap["template"]
        L.work = self.store.get_blob(L.tpl["working_blob"])
        L.slots = L.tpl["slots"]
        L.native = G.native_format(L.tpl["format"])
        L.layout_mode = opts.get("layout") or "off"
        L.tpl_pages = None
        L.tpl_pages_done = False
        L.width = snap.get("naming_width") or max(3, len(str(snap.get("records_total") or 0)))
        L.loaded = True

    def _plan_slice(self, L: _Lease) -> bool:
        """Plan the next batches of a large job; persist each batch with its progress.

        Batches are inserted with a check that the number of stored items is still what this
        worker last saw, so a crash between batches resumes exactly where the last committed
        batch ended and never stores a sample twice."""
        jid = L.jid
        if L.records is None:
            recs = R.build_records([L.srcs[s["id"]] for s in L.snap["sources"]], L.snap["rule"])
            _name_records(recs["records"], {s["id"]: s["filename"] for s in L.snap["sources"]})
            if len(recs["records"]) != (L.snap.get("records_total") or len(recs["records"])):
                raise RuntimeError(f"来源数据与提交时不一致：提交时 {L.snap.get('records_total')} 个样本，现在 {len(recs['records'])} 个")
            L.records = recs["records"]
        records = L.records
        rule, mappings = L.snap["rule"], L.snap["mappings"]
        column_value = (lambda rec, h: record_column_value(rule, L.srcs, rec, h))
        deadline = time.monotonic() + slice_seconds()
        while not self._stop:
            ok, why = self._authorized(jid)
            if not ok:
                self._halt(jid, why)
                return True
            planned = self.store.one("SELECT COUNT(*) AS n FROM job_items WHERE job_id=?", (jid,))["n"]
            batch = records[planned:planned + plan_batch()]
            rows = []
            for rec in batch:
                pl = build_plan(rec, rule, L.srcs, mappings, L.slots)
                name = render_name(L.opts["naming"], rec, pl, L.slots, rec["index"], L.width, column_value)[0]
                rows.append((new_id("i"), jid, rec["index"], rec["key"], json.dumps(rec, ensure_ascii=False, default=str),
                             "skipped" if pl.blocking else "pending", json.dumps(pl.blocking or [], ensure_ascii=False), name,
                             rec["label"], L.run_no))
            with self.store.tx() as c:
                cur = c.execute("SELECT status, lease_owner, cancel_requested FROM jobs WHERE id=?", (jid,)).fetchone()
                if not cur or cur["lease_owner"] != self.worker_id or cur["status"] not in ("running", "cancelling"):
                    return True
                if cur["cancel_requested"]:
                    self._finalize_tx(c, jid, self.worker_id)
                    return True
                if c.execute("SELECT COUNT(*) AS n FROM job_items WHERE job_id=?", (jid,)).fetchone()["n"] != planned:
                    continue                                  # changed under us; re-read on the next round
                c.executemany("INSERT INTO job_items (id, job_id, seq, record_key, record, status, issues, file_base, label, run_no) "
                              "VALUES (?,?,?,?,?,?,?,?,?,?)", rows)
                summary = _loads(c.execute("SELECT summary FROM jobs WHERE id=?", (jid,)).fetchone()["summary"]) or {}
                summary["planned"] = planned + len(rows)
                c.execute("UPDATE jobs SET summary=? WHERE id=?", (json.dumps(summary), jid))
                self._update_counts(jid, c)
            if planned + len(rows) >= len(records):
                return self._finish_planning(L)
            if time.monotonic() >= deadline:
                return False
        return False

    def _finish_planning(self, L: _Lease) -> bool:
        jid = L.jid
        t = now()
        with self.store.tx() as c:
            j = c.execute("SELECT status, tenant, run_no, lease_owner, summary FROM jobs WHERE id=?", (jid,)).fetchone()
            if not j or j["lease_owner"] != self.worker_id or j["status"] not in ("running", "cancelling"):
                return True
            items = c.execute("SELECT id, file_base FROM job_items WHERE job_id=? ORDER BY seq", (jid,)).fetchall()
            names = dedupe([r["file_base"] for r in items])
            c.executemany("UPDATE job_items SET file_base=? WHERE id=?",
                          [(nm, r["id"]) for r, nm in zip(items, names) if nm != r["file_base"]])
            invalid = c.execute("SELECT COUNT(*) AS n FROM job_items WHERE job_id=? AND status='skipped'", (jid,)).fetchone()["n"]
            summary = _loads(j["summary"]) or {}
            summary.update(planned=len(items), invalid=invalid)
            if invalid and L.opts.get("error_policy", "block") == "block":
                # The same rule as a small job, which is refused before it is created; here the
                # job already exists, so it ends failed and keeps every located issue on its items.
                msg = f"{invalid} 个样本存在错误，已按「有错误则不开始」停止。请修正或选择「跳过错误样本」"
                c.execute("UPDATE job_items SET status='cancelled', issues=?, finished_at=? WHERE job_id=? AND status='pending'",
                          (json.dumps([{"level": "error", "code": "not_started", "stage": "plan",
                                        "message": "其他样本存在错误，按「有错误则不开始」未生成本样本"}], ensure_ascii=False), t, jid))
                summary["error_code"] = "has_errors"
                c.execute("UPDATE jobs SET summary=? WHERE id=?", (json.dumps(summary), jid))
                counts = self._update_counts(jid, c)
                c.execute("UPDATE jobs SET status='failed', error=?, finished_at=?, lease_owner=NULL, lease_until=NULL WHERE id=?",
                          (msg, t, jid))
                c.execute("UPDATE job_runs SET status='failed', finished_at=? WHERE job_id=? AND run_no=?", (t, jid, j["run_no"]))
                self._emit(c, jid, j["tenant"], "job.failed", {"status": "failed", "run_no": j["run_no"], "counts": counts,
                                                                "error": msg, "error_code": "has_errors", "invalid": invalid})
                L.records = None
                return True
            pending = c.execute("SELECT COUNT(*) AS n FROM job_items WHERE job_id=? AND status='pending'", (jid,)).fetchone()["n"]
            c.execute("UPDATE jobs SET phase='generating', summary=? WHERE id=?", (json.dumps(summary), jid))
            c.execute("UPDATE job_runs SET items=? WHERE job_id=? AND run_no=?", (pending, jid, j["run_no"]))
            counts = self._update_counts(jid, c)
            self._emit(c, jid, j["tenant"], "job.planned", {"status": j["status"], "phase": "generating", "run_no": j["run_no"],
                                                             "counts": counts})
        # Every sample is stored as an item now; the records (up to 20,000 dicts with their
        # rows) were only needed to plan them and were held until the job ended.
        L.records = None
        L.cursor = 0
        return False

    def _generate_slice(self, L: _Lease) -> bool:
        jid = L.jid
        if not L.tpl_pages_done:
            L.tpl_pages = G.template_pages(L.tpl["format"], L.work) if L.layout_mode != "off" else None
            L.tpl_pages_done = True
        deadline = time.monotonic() + slice_seconds()
        last_count = time.monotonic()
        while not self._stop:
            ctl = self.store.one("SELECT cancel_requested, lease_owner, status FROM jobs WHERE id=?", (jid,))
            if not ctl or ctl["lease_owner"] != self.worker_id or ctl["status"] not in ("running", "cancelling"):
                return True                                   # lease lost (e.g. paused too long): another worker resumes
            if ctl["cancel_requested"]:
                break
            ok, why = self._authorized(jid)                   # between samples
            if not ok:
                self._halt(jid, why)
                return True
            item = self.store.one("SELECT * FROM job_items WHERE job_id=? AND status='pending' AND seq>=? ORDER BY seq LIMIT 1",
                                  (jid, L.cursor))
            if item is None and L.cursor:
                # Nothing after the cursor; make sure nothing before it was put back meanwhile.
                item = self.store.one("SELECT * FROM job_items WHERE job_id=? AND status='pending' ORDER BY seq LIMIT 1", (jid,))
            if item is None:
                break
            L.cursor = item["seq"]
            with self.store.tx() as c:
                n = c.execute("UPDATE job_items SET status='running', started_at=?, attempts=attempts+1, run_no=? "
                              "WHERE id=? AND status='pending' AND EXISTS (SELECT 1 FROM jobs j WHERE j.id=job_items.job_id AND j.lease_owner=?)",
                              (now(), L.run_no, item["id"], self.worker_id)).rowcount
            if not n:
                L.cursor = item["seq"] + 1
                continue
            attempt = item["attempts"] + 1
            try:
                self._process_item(jid, item, attempt, L.run_no, L.snap, L.opts, L.srcs, L.tpl, L.work, L.slots, L.native,
                                   L.layout_mode, L.tpl_pages)
            except Exception as e:  # noqa: BLE001 - item-level failure: recorded on the item, trace in the server log
                print(f"[synwevo] job {jid} item {item['id']} failed\n{traceback.format_exc()}", flush=True)
                self._item_done(item["id"], "failed", [{"level": "error", "code": "exception", "stage": "generate",
                                                         "message": f"{type(e).__name__}: {e}"}])
            L.cursor = item["seq"] + 1
            if time.monotonic() - last_count > 0.5:
                self._update_counts(jid)
                last_count = time.monotonic()
            if time.monotonic() >= deadline:
                self._update_counts(jid)
                return False
        if self._stop:
            return False
        with self.store.tx() as c:
            self._finalize_tx(c, jid, self.worker_id)
        return True

    def _fail_job(self, jid: str, e: Exception) -> None:
        with self.store.tx() as c:
            j = c.execute("SELECT status, tenant, run_no, lease_owner FROM jobs WHERE id=?", (jid,)).fetchone()
            if not j or j["lease_owner"] != self.worker_id or j["status"] not in ("running", "cancelling"):
                return
            c.execute("UPDATE job_items SET status='failed', finished_at=?, issues=? WHERE job_id=? AND status IN ('pending','running')",
                      (now(), json.dumps([{"level": "error", "code": "job_error", "message": f"任务无法继续：{type(e).__name__}: {e}"}], ensure_ascii=False), jid))
            counts = self._update_counts(jid, c)
            c.execute("UPDATE jobs SET status='failed', error=?, finished_at=?, lease_owner=NULL, lease_until=NULL WHERE id=?",
                      (f"{type(e).__name__}: {e}", now(), jid))
            c.execute("UPDATE job_runs SET status='failed', finished_at=? WHERE job_id=? AND run_no=?", (now(), jid, j["run_no"]))
            self._emit(c, jid, j["tenant"], "job.failed", {"status": "failed", "run_no": j["run_no"], "counts": counts,
                                                            "error": f"{type(e).__name__}: {e}"})

    def _holds_lease(self, jid: str) -> bool:
        row = self.store.one("SELECT lease_owner FROM jobs WHERE id=?", (jid,))
        return bool(row) and row["lease_owner"] == self.worker_id

    def _finalize_tx(self, c: sqlite3.Connection, jid: str, expect_owner: str | None) -> str | None:
        """Decide the terminal status inside a transaction.  A cancel that was accepted cancels the
        items that have not started; if none were left, the job completes normally."""
        j = c.execute("SELECT status, tenant, run_no, cancel_requested, lease_owner, phase FROM jobs WHERE id=?", (jid,)).fetchone()
        if not j or j["status"] not in ("running", "cancelling") or j["lease_owner"] != expect_owner:
            return None
        cancelled_now = 0
        if j["cancel_requested"]:
            cancelled_now = c.execute("UPDATE job_items SET status='cancelled' WHERE job_id=? AND status IN ('pending','running')", (jid,)).rowcount
        counts = self._update_counts(jid, c)
        if cancelled_now or (j["cancel_requested"] and (counts.get("cancelled") or j["phase"] == "planning")):
            status = "cancelled"
        elif counts.get("success", 0) == 0:
            status = "failed"
        elif counts.get("failed") or counts.get("skipped") or counts.get("cancelled") or counts.get("unplanned"):
            status = "completed_with_errors"
        else:
            status = "completed"
        t = now()
        c.execute("UPDATE jobs SET status=?, finished_at=?, lease_owner=NULL, lease_until=NULL WHERE id=?", (status, t, jid))
        c.execute("UPDATE job_runs SET status=?, finished_at=? WHERE job_id=? AND run_no=?", (status, t, jid, j["run_no"]))
        self._emit(c, jid, j["tenant"], f"job.{status}", {"status": status, "run_no": j["run_no"], "counts": counts})
        return status

    def _item_done(self, item_id: str, status: str, issues: list[dict]) -> None:
        # Guarded by the lease: a worker that lost its lease must not fail an item another
        # worker is now processing (its good result would then be refused at publication).
        self.store.execute("UPDATE job_items SET status=?, issues=?, finished_at=? WHERE id=? AND status='running' "
                           "AND EXISTS (SELECT 1 FROM jobs j WHERE j.id=job_items.job_id AND j.lease_owner=?)",
                           (status, json.dumps(issues, ensure_ascii=False), now(), item_id, self.worker_id))

    def _process_item(self, jid, item, attempt, run_no, snap, opts, srcs, tpl, work, slots, native, layout_mode, tpl_pages) -> None:
        plan = build_plan(item["record"], snap["rule"], srcs, snap["mappings"], slots)
        if plan.blocking:
            self._item_done(item["id"], "failed", [{**i, "stage": i.get("stage", "resolve")} for i in plan.blocking])
            return
        out, issues, lay = G.fill_checked(tpl["format"], work, slots, plan,
                                          {**self.app.with_image_bytes(tpl.get("fill_options") or {}), "keep_fields": opts.get("keep_fields")},
                                          layout_mode, tpl_pages)
        blocking = [i for i in issues if i.get("level") == "error"]
        if blocking or not out:
            failed = [{**i, "stage": i.get("stage", "fill")} for i in (blocking or issues)]
            if lay:
                # Keep the layout facts precisely when the sample failed — that is when the
                # user needs to know the achieved page count, the factor tried and the floor
                # that stopped it.  They used to be thrown away here.
                failed.append({"level": "info", "code": "layout_report", "stage": "layout",
                               "message": f"版式：模板 {lay.get('template_pages')} 页，本次排版 {lay.get('pages')} 页；"
                                          f"实际缩放 {lay.get('effective_factor', 1.0):.0%}"
                                          + (f"，最小 {lay['min_pt']:.1f} 磅" if lay.get("min_pt") else "")
                                          + (f"；已到最小字号的字段：{'、'.join(lay['clamped'])}" if lay.get("clamped") else "")
                                          + f"；共排版 {lay.get('renders')} 次",
                               "layout": lay})
            self._item_done(item["id"], "failed", failed)
            return
        produced = []
        for f in opts["formats"]:
            try:
                data, conv_issues = G.convert_output(out, native, f)
                issues.extend(conv_issues)
            except Exception as e:  # noqa: BLE001
                self._item_done(item["id"], "failed", issues + [{"level": "error", "code": "export_failed", "stage": "export",
                                                                  "message": f"{f.upper()}：{e}"}])
                return
            ver = verify_output(f, data, slots if f == native else [], plan)
            if lay and f == native:
                ver["layout"] = lay
            problem = None if ver.get("reopened") else (ver.get("error") or "输出校验失败")
            if not problem and ver["missing"]:
                problem = "输出文件中找不到已填写的值：" + "、".join(f"{m['slot']}={m['value']}" for m in ver["missing"][:3])
            if problem:
                self._item_done(item["id"], "failed", issues + [{"level": "error", "code": "export_failed", "stage": "verify",
                                                                  "message": f"{f.upper()}：{problem}"}])
                return
            produced.append((f, data, ver))
        # Immediately before anything becomes visible: rendering one sample can take minutes, and
        # a key revoked meanwhile must not get this sample's files either.
        ok, why = self._authorized(jid)
        if not ok:
            self._halt(jid, why)
            return
        stored = [(f, self.store.put_blob(data), len(data), ver) for f, data, ver in produced]   # content-addressed, atomic
        # publish: output records and the success status become visible together, or not at all
        with self.store.tx() as c:
            # Recheck under the writer transaction: a revocation can commit while blob bytes
            # are being stored after the earlier check. It must win over artifact publication.
            ok, why = self._authorized(jid)
            if not ok:
                self._halt(jid, why)
                return
            live = c.execute("SELECT i.status, j.lease_owner, j.status AS job_status FROM job_items i JOIN jobs j ON j.id=i.job_id WHERE i.id=?",
                             (item["id"],)).fetchone()
            if not live or live["status"] != "running" or live["lease_owner"] != self.worker_id \
                    or live["job_status"] not in ("running", "cancelling"):
                return                                        # recovered or halted meanwhile: do not publish
            for f, blob, size, ver in stored:
                c.execute("INSERT INTO artifacts (id, job_id, item_id, format, filename, blob, size, trace, verification, created_at, attempt, run_no) "
                          "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                          (new_id("a"), jid, item["id"], f, f"{item['file_base']}.{f}", blob, size,
                           json.dumps({"record": {k: item["record"].get(k) for k in ("key", "label", "file_id", "file_name", "where", "rows")},
                                       "slots": plan.trace, "template_version": tpl["version_no"]}, ensure_ascii=False, default=str),
                           json.dumps(ver, ensure_ascii=False, default=str), now(), attempt, run_no))
            c.execute("UPDATE job_items SET status='success', issues=?, finished_at=? WHERE id=?",
                      (json.dumps([i for i in issues if i.get("level") != "info"][:20], ensure_ascii=False), now(), item["id"]))

    # -------------------------------------------------------------- downloads
    def artifact(self, aid: str, tenant: str | None = None) -> tuple[dict, bytes]:
        a = self.store.one("SELECT a.*, j.tenant FROM artifacts a JOIN jobs j ON j.id=a.job_id WHERE a.id=?", (aid,))
        tenant = tenant or current_tenant()
        if not a or (a["tenant"] or "default") != tenant:
            not_found("输出文件")
        if a.get("deleted_at"):
            raise ServiceError(410, "artifact_deleted", "该输出文件已被删除，无法下载；可对原任务重新生成")
        if not self.store.has_blob(a["blob"]):
            raise ServiceError(410, "artifact_missing", "该输出文件的内容已不在存储中（可能已按保留期清理）")
        return a, self.store.get_blob(a["blob"])

    def zip_file(self, jid: str, fmt: str | None = None) -> Path:
        """Write the job's published outputs (+ manifest) to a temporary ZIP file and return its path;
        streaming from disk keeps memory flat for large jobs and supports range requests."""
        j = self.get(jid)
        arts = self.store.all("SELECT a.*, i.seq FROM artifacts a JOIN job_items i ON a.item_id=i.id "
                              "WHERE a.job_id=? AND a.deleted_at IS NULL AND i.status='success' ORDER BY i.seq, a.format", (jid,))
        fd, path = tempfile.mkstemp(prefix=f"{jid}-", suffix=".zip", dir=self.tmp_dir)
        os.close(fd)
        used = set()
        multi = len({x["format"] for x in arts}) > 1
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for a in arts:
                if fmt and a["format"] != fmt:
                    continue
                name = a["filename"]
                folder = f"{a['format']}/" if multi else ""
                full = folder + name
                k = 2
                while full.lower() in used:
                    base, _, ext = name.rpartition(".")
                    full = f"{folder}{base} ({k}).{ext}"
                    k += 1
                used.add(full.lower())
                z.write(self.store.blob_path(a["blob"]), full)
            # A fixed timestamp (the job's own) instead of "now": two builds of an unchanged job
            # are then byte-identical, so a resumed range download cannot splice two archives.
            stamp = time.localtime(j.get("finished_at") or j.get("created_at") or 0)[:6]
            zi = zipfile.ZipInfo("生成清单.csv", date_time=(max(1980, stamp[0]),) + tuple(stamp[1:]))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, self.report_csv(jid))
        return Path(path)

    # `zip()` used to live here: it called zip_file() and then read the whole archive back into
    # memory, undoing exactly what zip_file's docstring promises. Removed rather than deprecated,
    # so it cannot be reached for again — both callers now stream the file. Measured difference
    # on a large job: ~+222 MB RSS versus ~+13 MB.

    def report_csv(self, jid: str) -> bytes:
        # Two queries for the whole report: it used to run one SELECT per item after building
        # an IN (…) list with one parameter per item, which fails past SQLite's variable limit.
        items = self.store.all("SELECT id, seq, label, status, issues, json_extract(record, '$.file_name') AS file_name, "
                               "json_extract(record, '$.where') AS where_ FROM job_items WHERE job_id=? ORDER BY seq", (jid,))
        halted = self._halt_issue(jid)
        files: dict[str, list[str]] = {}
        for a in self.store.all("SELECT item_id, filename FROM artifacts WHERE job_id=? AND deleted_at IS NULL ORDER BY rowid", (jid,)):
            files.setdefault(a["item_id"], []).append(a["filename"])
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(["序号", "样本", "状态", "输出文件", "来源文件", "来源位置", "问题说明"])
        status_cn = {"success": "成功", "failed": "失败", "skipped": "已跳过（生成前校验错误）", "cancelled": "已取消", "pending": "等待中", "running": "进行中"}
        for it in items:
            its = it["issues"] or ([halted] if halted and it["status"] == "cancelled" else [])
            issues = "；".join(i.get("message", "") for i in its if isinstance(i, dict) and i.get("level") in ("error", "warning"))
            w.writerow([it["seq"] + 1, _csv_cell(it["label"]), status_cn.get(it["status"], it["status"]),
                        _csv_cell("；".join(files.get(it["id"], []))), _csv_cell(it["file_name"] or ""), _csv_cell(it["where_"] or ""),
                        _csv_cell(issues)])
        return ("﻿" + out.getvalue()).encode("utf-8")
