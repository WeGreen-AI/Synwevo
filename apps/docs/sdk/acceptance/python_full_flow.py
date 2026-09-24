"""Run the installed Python SDK and CLI against a real isolated Synwevo Docs service.

Only synthetic make_kit inputs. Core identity fixtures bootstrap a scoped service client;
all token and business operations use real HTTP APIs. No supplied/live URL or real data.
Server dependencies use the repository interpreter; --sdk-python may be a clean wheel install.
"""
import argparse
import hashlib
import importlib.util
import io
import json
import os
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

HERE = Path(__file__).resolve().parent


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@contextmanager
def services(testenv, out):
    """Two actual roles, one newly initialized temporary store, no browser session."""
    with tempfile.TemporaryDirectory(prefix="synwevo-python-service-") as tmp:
        root = Path(tmp)
        account_port, docs_port = testenv._free_port(), testenv._free_port()
        while docs_port == account_port:
            docs_port = testenv._free_port()
        account, docs = f"http://127.0.0.1:{account_port}", f"http://127.0.0.1:{docs_port}"
        env = testenv._child_env(root, docs_port)
        env.update(SYNWEVO_UI_AUTH="session", SYNWEVO_ISSUER=account, SYNWEVO_COOKIE_SECURE="0")
        launcher = HERE.parents[3] / "tests" / "integration" / "sdk_acceptance" / "python_service.py"
        launch = [sys.executable, str(launcher), "--root", str(root)]
        subprocess.run(launch + ["--bootstrap"], env=env, cwd=root, check=True, timeout=60)
        identity = json.loads((root / "identity.json").read_text())
        (out / "identity.json").write_text(json.dumps(identity, indent=2))
        processes = []
        try:
            for role, port, base in (("accounts", account_port, account), ("app", docs_port, docs)):
                log = out / (role + "-server.log")
                with log.open("wb") as f:
                    proc = subprocess.Popen(launch + ["--audit", str(out / (role + "-http.jsonl"))],
                                            cwd=root, env={**env, "SYNWEVO_ROLE": role, "SYNWEVO_PORT": str(port)},
                                            stdout=f, stderr=f)
                processes.append(proc)
                testenv._wait(base, proc, log)
                assert testenv.health(base)["instance"] == "test"
            yield {"SYNWEVO_BASE_URL": docs, "SYNWEVO_TOKEN_URL": account + "/oidc/token",
                   "SYNWEVO_CLIENT_ID": identity["client_id"], "SYNWEVO_CLIENT_SECRET_FILE": str(root / "client-secret"),
                   "SYNWEVO_TOKEN_SCOPE": "docweave.read docweave.write", "SYNWEVO_TOKEN_RESOURCE": "docweave-api"}
        finally:
            for proc in reversed(processes):
                proc.terminate()
                try:
                    proc.wait(timeout=15)
                except subprocess.TimeoutExpired:
                    proc.kill()
                    proc.wait(timeout=5)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sdk-python", default=sys.executable)
    parser.add_argument("--evidence", type=Path, required=True)
    args = parser.parse_args()
    # macOS system proxy settings also apply without HTTP_PROXY in the environment.
    # This harness communicates exclusively over loopback, including the health probe.
    os.environ["NO_PROXY"] = "127.0.0.1,localhost"
    out = args.evidence.resolve()
    out.mkdir(parents=True, exist_ok=True)
    if any(out.iterdir()):
        parser.error("--evidence must be empty; refusing to combine a new run with earlier evidence")
    kit = load("synthetic_kit", HERE / "make_kit.py")
    testenv = load("isolated_testenv", HERE.parents[1] / "scripts" / "_testenv.py")
    expected = kit.make(out / "kit")
    env = {k: v for k, v in os.environ.items() if not k.startswith(("SYNWEVO_", "DOCWEAVE_"))}
    # A larger, separately labelled cancellation fixture reuses only the same synthetic rows.
    # The primary SDK and CLI flows still use the unchanged 12-row shared kit.
    import openpyxl
    workbook = openpyxl.load_workbook(out / "kit" / "customers.xlsx")
    sheet = workbook.active
    for _ in range(24):
        for row in kit.CUSTOMERS:
            sheet.append(list(row))
    workbook.save(out / "kit" / "cancel-synthetic.xlsx")
    with services(testenv, out) as credentials:
        child = {**env, **credentials}
        cmd = [args.sdk_python, str(HERE / "python_worker.py"), "--kit", str(out / "kit"), "--out", str(out / "flow")]
        with (out / "worker.log").open("w") as log:
            result = subprocess.run(cmd, env=child, stdout=log, stderr=subprocess.STDOUT, timeout=600)
        assert result.returncode == 0, f"Worker exited {result.returncode}; see {out / 'worker.log'}"
        results = json.loads((out / "flow" / "results.json").read_text())
        checks = []
        for channel, flow in results.items():
            files = flow["artifacts"]
            assert len(files) == expected["samples"]
            by_name = {f["filename"]: f for f in files}
            for item in expected["items"]:
                artifact = by_name[item["file_base"] + ".docx"]
                data = Path(artifact["path"]).read_bytes()
                assert hashlib.sha256(data).hexdigest() == artifact["sha256"]
                problems = kit.check_docx(data, item, expected["forbidden_text"])
                assert not problems, (channel, artifact["filename"], problems)
                checks.append({"client": channel, "filename": artifact["filename"], "sha256": artifact["sha256"], "content": "passed"})
        preview_checks = []
        for filename in ("sdk-preview.docx", "sdk-preview-key.docx", "cli-preview.docx"):
            data = (out / "flow" / filename).read_bytes()
            assert not kit.check_docx(data, expected["items"][0], expected["forbidden_text"])
            preview_checks.append({"filename": filename, "sha256": hashlib.sha256(data).hexdigest(), "content": "passed"})
        builtin_checks = []
        builtin_flows = json.loads((out / "flow" / "builtins.json").read_text())
        assert len(builtin_flows) == 4
        for flow in builtin_flows:
            by_name = {a["filename"]: a for a in flow["artifacts"]}
            pairs = [(by_name[item["file_base"] + "." + flow["format"]], item) for item in expected["items"]]
            pairs.append((flow["preview"], expected["items"][0]))
            for artifact, item in pairs:
                data = Path(artifact["path"]).read_bytes()
                assert hashlib.sha256(data).hexdigest() == artifact["sha256"]
                if flow["format"] == "docx":
                    import docx
                    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
                    assert not kit.check_docx(data, item, expected["forbidden_text"])
                else:
                    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
                    text = "\n".join(str(cell.value) for ws in wb for row in ws for cell in row if cell.value is not None)
                    wb.close()
                    assert all(value in text for value in item["must_contain"]), (artifact["path"], text)
                    assert not any(value in text for value in expected["forbidden_text"])
                assert flow["marker"] in text, (artifact["path"], "saved revision marker absent")
                builtin_checks.append({"client": flow["client"], "kind": flow["kind"], "path": artifact["path"],
                                       "sha256": artifact["sha256"], "content": "passed"})
    # Services have stopped and their temporary database, blobs and secret are removed.
    audits = [json.loads(line) for role in ("accounts", "app")
              for line in (out / (role + "-http.jsonl")).read_text().splitlines()]
    assert not any(r["cookie"] or r["set_cookie"] or r["authorization"] == "legacy_api_key" for r in audits)
    tokens = [r for r in audits if r["path"] == "/oidc/token"]
    assert tokens and all(r["status"] == 200 and r["grant"]["grant_type"] == ["client_credentials"] for r in tokens)
    assert all(set(r["token_response"]["scope"].split()) == {"docweave.read", "docweave.write"} for r in tokens)
    business = [r for r in audits if r["path"].startswith("/api/v1/")]
    assert business and all(r["authorization"] == "bearer" and r["principal"]["kind"] == "service_account" and
                            set(r["principal"]["scopes"]) == {"docweave.read", "docweave.write"} for r in business)
    summary = {"instance": "test", "authorization": "Account client_credentials over HTTP; Docs read/write only",
               "clients": list(results), "documents_verified": len(checks), "checks": checks,
               "token_requests": len(tokens), "machine_requests": len(business), "cookies": 0, "admin_scope": False,
               "preview": preview_checks, "cancellation_retry": "two real 300-item batches; see flow/calls.json",
               "builtin_documents_verified": 48, "builtin_previews_verified": 4, "builtin_checks": builtin_checks,
               "failed_only": "cancelled-item exclusion and fingerprint conflict; no injected failed-item fixture",
               "result": "passed"}
    (out / "acceptance.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    print(json.dumps({"result": "passed", "clients": list(results), "documents_verified": len(checks),
                      "builtin_outputs_verified": len(builtin_checks), "evidence": str(out)}))


if __name__ == "__main__":
    main()
