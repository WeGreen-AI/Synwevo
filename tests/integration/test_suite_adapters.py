"""WP7: real factories/services, isolated fixture people, files and controlled conversation.

No model or customer service is called. Fault injection is explicitly limited to simulated
lost acknowledgements; all document rendering and message writes are the real implementations.
"""
from __future__ import annotations

import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import docx
import openpyxl
import pytest
from fastapi.testclient import TestClient

from synwevo_core.apps import Apps, DOCS_APP_ID, MESSAGES_APP_ID
from synwevo_core.context import as_tenant
from synwevo_core.errors import ServiceError
from synwevo_core.organizations import Organizations
from synwevo_core.sessions import COOKIE_NAME, CSRF_COOKIE, Sessions
from synwevo_docs.definitions import TEMPLATE_FIELDS_SCHEMA, import_template_fields
from synwevo_server.app import create_app
from synwevo_server.suite_adapters import prepare, install

ORIGIN = "http://suite.testserver"
PW = "harbor lights guide the evening ferry home"


def wait_for(fn):
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        result = fn()
        if result:
            return result
        time.sleep(0.03)
    raise AssertionError("real fixture parse did not finish")


@pytest.fixture
def world(tmp_path, monkeypatch, request):
    for key, value in {"UI_AUTH": "session", "COOKIE_SECURE": "0", "ROLE": "all", "APP_URL": ORIGIN,
                       "DATA_DIR": str(tmp_path / "data"), "MAIL_MODE": "file", "MAIL_DIR": str(tmp_path / "mail"),
                       "SUITE_APPS": getattr(request, "param", "forms,agent"), "SCANNER_MODE": "off"}.items():
        monkeypatch.setenv("SYNWEVO_" + key, value)
    api = create_app(str(tmp_path / "data"), role="all", start_jobs=False)
    ctx = api.state.suite
    adapters = api.state.suite_adapters
    assert adapters._ctx is ctx  # deployed factory registers the actual services
    ident, store = ctx.identity, ctx.store
    users = {}
    for name in ("owner", "editor", "viewer", "other", "admin"):
        u = ident.create_user(f"{name}@example.test", PW, status="active", display_name=name,
                              platform_role="admin" if name == "admin" else "none")
        store.execute("UPDATE users SET email_verified_at=1 WHERE id=?", (u["id"],))
        users[name] = u["id"]
    orgs = Organizations(store, ident)
    made = orgs.create(users["owner"], "WP7 Fixture Organization", "Fixture A", "wp7-fixture-organization")
    oid, a = made["organization"]["id"], made["workspace"]["id"]
    b = orgs.create_workspace(oid, users["owner"], "Fixture B")["id"]
    for who, role in (("editor", "editor"), ("viewer", "viewer"), ("other", "editor")):
        store.execute("INSERT INTO organization_memberships (organization_id,user_id,role,created_at,updated_at,created_by) "
                      "VALUES (?,?,'member',1,1,?)", (oid, users[who], users["owner"]))
        ident.add_member(a if who != "other" else b, users[who], role)
    # Owner belongs to both; editor gets read-only access to B.
    ident.add_member(b, users["editor"], "viewer")
    apps = Apps(store, ident)
    for wid in (a, b):
        apps.enable_for_workspace(wid, DOCS_APP_ID)
    messaging = api.state.messaging
    group = messaging.create_group(oid, users["owner"], name="WP7 authorized test destination",
                                   member_ids=[users["editor"], users["viewer"]], workspace_id=a)
    with TestClient(api, base_url=ORIGIN) as client:
        w = SimpleNamespace(api=api, ctx=ctx, caps=ctx.capabilities, adapters=adapters, ident=ident,
                            store=store, users=users, a=a, b=b, oid=oid, cid=group["id"], client=client, apps=apps)
        def invoke(suffix, payload=None, *, app=DOCS_APP_ID, user="editor", ws=None, key=None, **kw):
            return w.caps.invoke(f"{app}.{suffix}", payload or {}, user_id=users[user], workspace_id=ws or a,
                                 idempotency_key=key, **kw)
        w.invoke = invoke
        def read(suffix, payload=None, **kw):
            return invoke(suffix, payload, **kw)["output"]
        w.read = read
        yield w


def configure(w, *, wid=None, name="WP7 actual document fixture", value="FORM-0007", required=False):
    """Real business-service setup: input workbook, uploaded template, published field mapping."""
    app = w.adapters.docs
    wid = wid or w.a
    with as_tenant(wid):
        pid = app.create_project(name, origin="test:wp7")["id"]
        wb = openpyxl.Workbook()
        wb.active.title = "Input"
        wb.active.append(["Reference", "Marker"])
        for entry in value if isinstance(value, list) else [value]:
            wb.active.append([entry, "fixture-row"])
        data = io.BytesIO(); wb.save(data)
        source = app.add_source(pid, "fixture.xlsx", data.getvalue())
        wait_for(lambda: app.get_source(source["id"])["status"] == "ready")
        template = docx.Document(); template.add_paragraph("Receipt: PLACEHOLDER")
        data = io.BytesIO(); template.save(data)
        tpl = app.upload_template("fixture.docx", data.getvalue())
        import_template_fields(app, tpl["id"], {"schema": TEMPLATE_FIELDS_SCHEMA, "version": 1,
                                         "fields": [{"kind": "text", "name": "Reference", "find": "PLACEHOLDER", "required": required}]})
        head = app.get_template(tpl["id"])["head"]
        field = app.get_version(tpl["id"], head)["slots"][0]
        app.update_project(pid, {"template_id": tpl["id"], "template_version": head,
                                "record_rule": {"mode": "per_row", "table": {"sheet": "Input", "header_row": 0}}})
        app.save_mappings(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "Input", "header": "Reference"}],
                                "target": {"binding": field["binding"]}}])
    return pid, tpl["id"], head


def job_input(w, pid, **kw):
    p = w.read("project.get", {"project_id": pid}, **kw)
    return {"project_id": pid, "configuration": p["configuration"], "formats": ["docx"]}


def finish(w, jid):
    jobs = w.adapters.jobs
    assert jobs._claim(jid) == jid
    jobs._run(jid)
    out = w.read("job.get", {"job_id": jid})
    assert out["status"] == "completed", out
    w.adapters.sync_jobs()
    return out


def send_input(w, **kw):
    conv = w.read("conversation.get", {"conversation_id": w.cid}, app=MESSAGES_APP_ID)
    return {"conversation_id": w.cid, "audience": conv["audience"], "body": "Authorized fixture delivery", **kw}


def sign_in(w, who):
    raw, _ = Sessions(w.store).create(w.users[who], workspace_id=w.b, amr="password")
    w.client.cookies.set(COOKIE_NAME, raw)
    w.client.cookies.set(CSRF_COOKIE, "fixture-csrf")


def test_real_docx_job_protected_download_and_test_message(world):
    w = world
    # Real Forms submission -> explicit fixture handoff to the existing Docs input service.
    # This tests the adapter boundary, not the still-to-be-composed middle applications.
    w.apps.enable_for_workspace(w.a, "forms")
    form = w.read("form.create", {"title": "WP7 receipt source", "definition": {"fields": [
        {"id": "f_ref00001", "type": "text", "label": "Reference", "required": True}]}}, app="forms")
    w.read("form.publish", {"form_id": form["id"], "draft_rev": form["draft_rev"]}, app="forms")
    request = w.read("request.create", {"form_id": form["id"]}, app="forms")
    token = request["url"].rsplit("/", 1)[1]
    sign_in(w, "editor")
    submitted = w.client.post(f"/api/suite/public/forms/r/{token}/submit",
                              headers={"X-Synwevo-CSRF": "fixture-csrf"},
                              json={"version": 1, "answers": {"f_ref00001": "FORM-0007"}, "idempotency_key": "wp7-form-submit"})
    assert submitted.status_code == 200, submitted.text
    submission = w.read("submission.get", {"submission_id": submitted.json()["submission_id"]}, app="forms")
    pid, tid, version = configure(w, value=submission["answers"]["f_ref00001"])
    inp = job_input(w, pid)
    result = w.invoke("job.create", inp, key="business-request-1", via="flow", delegation={"flow_run": "fixture-chain"},
                      principal={"kind": "operator"})
    assert result["status"] == "queued"
    jid = result["output"]["job_id"]
    stored = w.store.one("SELECT principal,snapshot FROM jobs WHERE id=?", (jid,))
    principal = json.loads(stored["principal"])
    assert principal == {"kind": "user", "user_id": w.users["editor"], "tenant": w.a}
    assert stored["snapshot"]["template"]["version_no"] == version
    # The job pins an immutable version, even if the template head moves before execution.
    with as_tenant(w.a):
        import_template_fields(w.adapters.docs, tid, {"schema": TEMPLATE_FIELDS_SCHEMA, "version": 1,
                               "fields": [{"kind": "text", "name": "Label", "find": "Receipt"}]})
    finish(w, jid)
    assert w.caps.run(result["run_id"], user_id=w.users["editor"])["status"] == "succeeded"
    aid = w.read("job.artifacts", {"job_id": jid})["items"][0]["artifact_id"]
    link = w.read("artifact.get", {"artifact_id": aid})["download_url"]
    sign_in(w, "editor")
    response = w.client.get(link)  # session is in B; the link explicitly authorizes A
    assert response.status_code == 200, response.text
    document = docx.Document(io.BytesIO(response.content))
    assert "FORM-0007" in "\n".join(p.text for p in document.paragraphs)
    payload = send_input(w, artifact_ids=[aid])
    sent = w.invoke("message.send", payload, app=MESSAGES_APP_ID, key="send-business-request-1", via="flow")
    assert sent["status"] == "succeeded" and sent["output"]["delivery"] == "recorded"
    stored_msg = w.store.one("SELECT * FROM messages WHERE id=?", (sent["output"]["message_id"],))
    assert stored_msg["sender_id"] == w.users["editor"] and link in stored_msg["body"]
    assert w.invoke("message.send", payload, app=MESSAGES_APP_ID, key="send-business-request-1")["replayed"]
    assert w.store.one("SELECT COUNT(*) AS n FROM messages WHERE client_id=?", (sent["run_id"],))["n"] == 1
    assert w.store.one("SELECT COUNT(*) AS n FROM mail_outbox")["n"] == 0
    sign_in(w, "other")
    assert w.client.get(link).status_code == 403


def test_concurrent_idempotency_and_restart_keep_one_job(world):
    w = world
    pid, _, _ = configure(w)
    inp = job_input(w, pid)
    with ThreadPoolExecutor(max_workers=4) as pool:
        runs = list(pool.map(lambda _: w.invoke("job.create", inp, key="concurrent"), range(4)))
    assert len({r["output"]["job_id"] for r in runs}) == 1
    assert w.store.one("SELECT COUNT(*) AS n FROM jobs WHERE project_id=?", (pid,))["n"] == 1
    with pytest.raises(ServiceError) as err:
        w.invoke("job.create", {**inp, "naming": "changed"}, key="concurrent")
    assert err.value.code == "idempotency_conflict"
    # Fresh service and registry, same isolated persistent data (no in-memory deduplication).
    from synwevo_core.capabilities import Registry, Capabilities
    from synwevo_docs.jobs import JobService
    from synwevo_server.suite_adapters import SuiteAdapters
    w.adapters.jobs.stop()
    jobs = JobService(w.adapters.docs, start=False)
    docs_api = SimpleNamespace(state=SimpleNamespace(app=w.adapters.docs, jobs=jobs))
    fresh = SuiteAdapters(docs_api, w.api, public_url=ORIGIN)
    registry = Registry(); caps = Capabilities(w.store, w.ident, registry)
    ctx = SimpleNamespace(store=w.store, identity=w.ident, registry=registry, capabilities=caps,
                          start_jobs=False, api=w.api)
    fresh.register(ctx)
    replay = caps.invoke(f"{DOCS_APP_ID}.job.create", inp, user_id=w.users["editor"], workspace_id=w.a, idempotency_key="concurrent")
    assert replay["replayed"] and replay["run_id"] == runs[0]["run_id"]
    jobs.stop()


def test_configuration_and_acl_fail_closed(world):
    w = world
    pid, _, _ = configure(w)
    inp = job_input(w, pid)
    with as_tenant(w.a):
        w.adapters.docs.update_project(pid, {"output": {"naming": "changed"}})
    with pytest.raises(ServiceError) as err:
        w.invoke("job.create", inp, key="changed")
    assert err.value.code == "configuration_changed"
    for who, ws in (("other", w.a), ("admin", w.a), ("editor", w.b)):
        with pytest.raises(ServiceError):
            w.read("project.get", {"project_id": pid}, user=who, ws=ws)
    with pytest.raises(ServiceError) as err:
        w.invoke("job.create", job_input(w, pid), key="viewer-write", user="viewer")
    assert err.value.code == "read_only_role"
    w.apps.revoke(w.a, DOCS_APP_ID, w.users["editor"])
    with pytest.raises(ServiceError):
        w.read("project.list")


def test_job_worker_checks_current_person_and_old_result_acl(world):
    w = world
    pid, _, _ = configure(w)
    inp = job_input(w, pid)
    run = w.invoke("job.create", inp, key="queued-revoke")
    jid = run["output"]["job_id"]
    w.apps.revoke(w.a, DOCS_APP_ID, w.users["editor"])
    assert w.adapters.jobs._claim(jid) is None
    row = w.store.one("SELECT status,halt_reason FROM jobs WHERE id=?", (jid,))
    assert row["status"] == "cancelled" and row["halt_reason"].startswith("authorization_revoked:")
    assert w.store.one("SELECT COUNT(*) AS n FROM artifacts WHERE job_id=?", (jid,))["n"] == 0
    assert w.caps.run(run["run_id"], user_id=w.users["editor"])["snapshot_hidden"]


def test_deleted_project_and_export_withdrawal_hide_old_results(world):
    w = world
    pid, _, _ = configure(w)
    inp = job_input(w, pid)
    run = w.invoke("job.create", inp, key="delete-after-result")
    jid = run["output"]["job_id"]
    finish(w, jid)
    aid = w.read("job.artifacts", {"job_id": jid})["items"][0]["artifact_id"]
    with pytest.raises(ServiceError) as err:
        w.read("artifact.get", {"artifact_id": aid}, user="viewer")
    assert err.value.code == "export_not_allowed"
    w.store.execute("UPDATE projects SET deleted=1 WHERE id=?", (pid,))
    with pytest.raises(ServiceError):
        w.read("artifact.get", {"artifact_id": aid})
    with pytest.raises(ServiceError) as err:
        w.invoke("job.create", inp, key="delete-after-result")
    assert err.value.code == "run_result_unavailable"
    assert w.caps.run(run["run_id"], user_id=w.users["editor"])["snapshot_hidden"]


def test_message_audience_scope_and_revocation(world):
    w = world
    inp = send_input(w)
    with pytest.raises(ServiceError):
        w.read("conversation.get", {"conversation_id": w.cid}, app=MESSAGES_APP_ID, ws=w.b)
    assert w.read("conversation.list", app=MESSAGES_APP_ID)["items"][0]["conversation_id"] == w.cid
    w.api.state.messaging.remove_member(w.cid, w.users["owner"], w.users["viewer"])
    with pytest.raises(ServiceError) as err:
        w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="stale-audience")
    assert err.value.code == "audience_changed"
    inp = send_input(w)
    sent = w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="current-audience")
    w.api.state.messaging.remove_member(w.cid, w.users["owner"], w.users["editor"])
    with pytest.raises(ServiceError) as err:
        w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="current-audience")
    assert err.value.code == "run_result_unavailable"
    assert w.caps.run(sent["run_id"], user_id=w.users["editor"])["snapshot_hidden"]


def test_lost_ack_reconciles_read_only_and_never_resends(world, monkeypatch):
    w = world
    inp = send_input(w)
    real_finish = w.caps._finish
    def crash_after_business_commit(*args, **kwargs):
        raise RuntimeError("TEST FIXTURE: process lost after committed message, before operation receipt")
    monkeypatch.setattr(w.caps, "_finish", crash_after_business_commit)
    with pytest.raises(RuntimeError):
        w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="lost-ack")
    monkeypatch.setattr(w.caps, "_finish", real_finish)
    row = w.store.one("SELECT * FROM op_runs WHERE idempotency_key='lost-ack'")
    w.caps.reconcile_stale(older_than_s=-1)
    before = w.store.one("SELECT COUNT(*) AS n FROM messages")["n"]
    receipt = w.caps.inspect(f"{MESSAGES_APP_ID}.message.inspect", {"run_id": row["id"]},
                             user_id=w.users["editor"], workspace_id=w.a)
    assert receipt["delivery"] == "recorded"
    replay = w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="lost-ack")
    assert replay["status"] == "unknown" and replay["replayed"]
    assert w.store.one("SELECT COUNT(*) AS n FROM messages")["n"] == before
    with pytest.raises(ServiceError):
        w.caps.inspect(f"{MESSAGES_APP_ID}.message.inspect", {"run_id": row["id"]},
                       user_id=w.users["owner"], workspace_id=w.a)


def test_inspect_does_not_write_or_send_and_policies_are_explicit(world):
    w = world
    pid, _, _ = configure(w)
    # sqlite total_changes is connection-local; all reads here use this same test thread.
    before = w.store._c().total_changes
    for suffix, inp, app in (("project.list", {}, DOCS_APP_ID), ("project.get", {"project_id": pid}, DOCS_APP_ID),
                              ("conversation.list", {}, MESSAGES_APP_ID), ("conversation.get", {"conversation_id": w.cid}, MESSAGES_APP_ID)):
        w.caps.inspect(f"{app}.{suffix}", inp, user_id=w.users["editor"], workspace_id=w.a)
    assert w.store._c().total_changes == before
    for app, suffix, inp in ((DOCS_APP_ID, "job.create", job_input(w, pid)),
                             (MESSAGES_APP_ID, "message.send", send_input(w))):
        op = w.ctx.registry.get(f"{app}.{suffix}")
        assert op.agent == "confirm" and op.result_guard is not None
        with pytest.raises(ServiceError) as err:
            w.caps.inspect(op.id, inp, user_id=w.users["editor"], workspace_id=w.a)
        assert err.value.code == "source_unverifiable"
        with pytest.raises(ServiceError) as err:
            w.invoke(suffix, inp, app=app)
        assert err.value.code == "idempotency_key_required"


def test_real_process_rebuild_reconciles_message_without_resend(world, monkeypatch):
    w = world
    payload = send_input(w)
    real_finish = w.caps._finish
    def lose_receipt(*args, **kwargs):
        raise RuntimeError("TEST FIXTURE: server dies after the business transaction")
    monkeypatch.setattr(w.caps, "_finish", lose_receipt)
    with pytest.raises(RuntimeError):
        w.invoke("message.send", payload, app=MESSAGES_APP_ID, key="restart-send")
    monkeypatch.setattr(w.caps, "_finish", real_finish)
    row = w.store.one("SELECT id FROM op_runs WHERE idempotency_key='restart-send'")
    fresh_api = create_app(str(w.store.root), role="all", start_jobs=False)
    fresh_ctx = fresh_api.state.suite
    fresh = fresh_api.state.suite_adapters
    assert fresh._ctx is fresh_ctx
    with TestClient(fresh_api, base_url=ORIGIN):
        fresh_ctx.capabilities.reconcile_stale(older_than_s=-1)
        got = fresh_ctx.capabilities.inspect(f"{MESSAGES_APP_ID}.message.inspect", {"run_id": row["id"]},
                                             user_id=w.users["editor"], workspace_id=w.a)
        assert got["delivery"] == "recorded"
        replay = fresh_ctx.capabilities.invoke(f"{MESSAGES_APP_ID}.message.send", payload,
                                               user_id=w.users["editor"], workspace_id=w.a,
                                               idempotency_key="restart-send")
        assert replay["status"] == "unknown" and replay["replayed"]
        assert fresh.store.one("SELECT COUNT(*) AS n FROM messages WHERE client_id=?", (row["id"],))["n"] == 1


def test_artifact_acl_is_rechecked_for_sent_message_receipts(world):
    w = world
    pid, _, _ = configure(w)
    job = w.invoke("job.create", job_input(w, pid), key="source-job")["output"]
    finish(w, job["job_id"])
    aid = w.read("job.artifacts", {"job_id": job["job_id"]})["items"][0]["artifact_id"]
    inp = send_input(w, artifact_ids=[aid])
    sent = w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="source-message")
    w.apps.revoke(w.a, DOCS_APP_ID, w.users["editor"])
    # Message membership persists; the source reference has a separate live authorization.
    assert w.api.state.messaging.get(w.cid, w.users["editor"])
    assert w.caps.run(sent["run_id"], user_id=w.users["editor"])["snapshot_hidden"]
    with pytest.raises(ServiceError):
        w.read("message.inspect", {"run_id": sent["run_id"]}, app=MESSAGES_APP_ID)


def test_unknown_send_without_record_remains_unknown(world, monkeypatch):
    w = world
    inp = send_input(w)
    def timeout(*args, **kwargs):
        raise TimeoutError("TEST FIXTURE: timeout before durable acknowledgement")
    monkeypatch.setattr(w.adapters.sender, "send", timeout)
    count = w.store.one("SELECT COUNT(*) AS n FROM messages")["n"]
    sent = w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="timeout")
    assert sent["status"] == "unknown"
    assert w.read("message.inspect", {"run_id": sent["run_id"]}, app=MESSAGES_APP_ID)["delivery"] == "unknown"
    assert w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="timeout")["replayed"]
    assert w.store.one("SELECT COUNT(*) AS n FROM messages")["n"] == count


def test_message_concurrent_retries_and_archived_refusal(world):
    w = world
    inp = send_input(w)
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda _: w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="send-concurrent"), range(4)))
    assert len({r["run_id"] for r in results}) == 1
    assert w.store.one("SELECT COUNT(*) AS n FROM messages WHERE client_id=?", (results[0]["run_id"],))["n"] == 1
    with pytest.raises(ServiceError) as err:
        w.invoke("message.send", {**inp, "body": "different"}, app=MESSAGES_APP_ID, key="send-concurrent")
    assert err.value.code == "idempotency_conflict"
    w.api.state.messaging.archive(w.cid, w.users["owner"])
    inp = send_input(w)
    with pytest.raises(ServiceError) as err:
        w.invoke("message.send", inp, app=MESSAGES_APP_ID, key="archived")
    assert err.value.code == "archived"


def test_prepare_before_install_private_docs_lifecycle(tmp_path, monkeypatch):
    """Real private Docs installer + real Suite install; no server app.py modifications."""
    from synwevo_core.http import build_platform, new_api, install as install_platform
    from synwevo_core.apps import bootstrap_docs_app
    from synwevo_suite.app import install as install_suite
    monkeypatch.setenv("SYNWEVO_UI_AUTH", "session")
    monkeypatch.setenv("SYNWEVO_COOKIE_SECURE", "0")
    monkeypatch.setenv("SYNWEVO_APP_URL", ORIGIN)
    plat = build_platform(str(tmp_path / "prepared"))
    api = new_api("WP7 real composition fixture")
    api.state.identity, api.state.sessions = plat.identity, plat.sessions
    install_platform(api, plat, role="portal")
    # Test fixture deployment registration. prepare itself must not grant app access.
    bootstrap_docs_app(plat.store, plat.identity)
    grants = plat.store.one("SELECT COUNT(*) AS n FROM workspace_apps")["n"]
    prepared = prepare(api, plat, start_jobs=True, public_url=ORIGIN)
    assert prepared.jobs.runner_alive()
    ctx = install_suite(api, plat, docs=prepared, start_jobs=True, released=frozenset())
    assert ctx.registry.get(f"{DOCS_APP_ID}.job.create") is not None
    assert ctx.registry.get(f"{MESSAGES_APP_ID}.message.send") is None  # no detached Message writer
    assert plat.store.one("SELECT COUNT(*) AS n FROM workspace_apps")["n"] == grants
    with TestClient(api, base_url=ORIGIN):
        assert prepared.ticker.thread.is_alive()
        user = plat.identity.create_user("private-worker@example.test", PW, status="active")
        ws = plat.identity.create_workspace("Private worker fixture", owner_id=user["id"])["id"]
        Apps(plat.store, plat.identity).enable_for_workspace(ws, DOCS_APP_ID)
        w = SimpleNamespace(adapters=prepared, a=ws)
        pid, _, _ = configure(w, value="ASYNC-001")
        config = ctx.capabilities.inspect(f"{DOCS_APP_ID}.project.get", {"project_id": pid},
                                          user_id=user["id"], workspace_id=ws)
        run = ctx.capabilities.invoke(f"{DOCS_APP_ID}.job.create",
                                      {"project_id": pid, "configuration": config["configuration"], "formats": ["docx"]},
                                      user_id=user["id"], workspace_id=ws, idempotency_key="real-background-job")
        done = wait_for(lambda: (r := ctx.capabilities.run(run["run_id"], user_id=user["id"]))["status"] == "succeeded" and r)
        assert done["output"]["counts"]["files"] == 1
        artifact = plat.store.one("SELECT id FROM artifacts WHERE job_id=?", (run["output"]["job_id"],))
        _, data = prepared.jobs.artifact(artifact["id"], ws)
        assert "ASYNC-001" in docx.Document(io.BytesIO(data)).paragraphs[0].text
    wait_for(lambda: not prepared.jobs.runner_alive())
    wait_for(lambda: not prepared.ticker.thread.is_alive())


def test_same_user_cannot_send_another_workspaces_artifact(world):
    w = world
    pid, _, _ = configure(w, wid=w.b)
    inp = job_input(w, pid, user="owner", ws=w.b)
    job = w.invoke("job.create", inp, user="owner", ws=w.b, key="workspace-b")["output"]
    assert w.adapters.jobs._claim(job["job_id"]) == job["job_id"]
    w.adapters.jobs._run(job["job_id"])
    aid = w.store.one("SELECT id FROM artifacts WHERE job_id=?", (job["job_id"],))["id"]
    count = w.store.one("SELECT COUNT(*) AS n FROM messages")["n"]
    with pytest.raises(ServiceError) as err:
        w.invoke("message.send", send_input(w, artifact_ids=[aid]), app=MESSAGES_APP_ID,
                 user="owner", key="cross-space-link")
    assert err.value.status == 404
    assert w.store.one("SELECT COUNT(*) AS n FROM messages")["n"] == count


def test_object_conversation_rechecks_docs_app_grant(world):
    from synwevo_core.messaging_bridge import ObjectConversations
    w = world
    pid, _, _ = configure(w)
    objects = ObjectConversations(w.store, w.ident, w.api.state.messaging)
    opened = objects.open(DOCS_APP_ID, "project", pid, w.users["editor"])
    cid = opened["conversation"]["id"]
    assert w.read("conversation.get", {"conversation_id": cid}, app=MESSAGES_APP_ID)
    w.apps.revoke(w.a, DOCS_APP_ID, w.users["editor"])
    with pytest.raises(ServiceError):
        w.read("conversation.get", {"conversation_id": cid}, app=MESSAGES_APP_ID)
    assert cid not in {c["conversation_id"] for c in w.read("conversation.list", app=MESSAGES_APP_ID)["items"]}


@pytest.mark.parametrize("values,expected,files", [(["GOOD", ""], "partial", 1), ([""], "failed", 0)])
def test_real_partial_and_failed_jobs_are_not_reported_as_success(world, values, expected, files):
    w = world
    pid, _, _ = configure(w, value=values, required=True)
    inp = {**job_input(w, pid), "error_policy": "skip"}
    run = w.invoke("job.create", inp, key="partial-fixture")
    jid = run["output"]["job_id"]
    assert w.adapters.jobs._claim(jid) == jid
    w.adapters.jobs._run(jid)
    result = w.read("job.get", {"job_id": jid})
    assert result["state"] == expected, result
    assert result["counts"]["files"] == files and result["counts"]["skipped"] == 1
    w.adapters.sync_jobs()
    assert w.caps.run(run["run_id"], user_id=w.users["editor"])["status"] == expected
