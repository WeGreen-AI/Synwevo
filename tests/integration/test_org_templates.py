"""Organization-shared templates on the composed server: who may publish (editors of the source
workspace, in an organization workspace), what a release freezes, who can see and use it (members
of the organization, from any of its workspaces, never from a personal workspace or another
organization), that a copy made from a release is an ordinary template of the reader's workspace
which later releases and deprecation do not touch, and that maintenance (new release, deprecate,
restore) needs the maintainer's authority."""
from __future__ import annotations

import io
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

PW = "harbor lights guide the evening ferry home"
CSRF, WS = "X-Synwevo-CSRF", "X-Synwevo-Workspace"


def _docx() -> bytes:
    import docx
    d = docx.Document(); d.add_paragraph("尊敬的示例客户名称：您好。")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def _walk(blocks):
    for b in blocks:
        if b["t"] == "p":
            yield b
        else:
            for row in b["rows"]:
                for cell in row["cells"]:
                    yield from _walk(cell["blocks"])


@pytest.fixture()
def world(tmp_path, monkeypatch):
    for k, v in {"UI_AUTH": "session", "COOKIE_SECURE": "0", "DATA_DIR": str(tmp_path / "d"), "MAIL_MODE": "file", "MAIL_DIR": str(tmp_path / "mail"),
                 "MAIL_FROM": "synwevo@example.test", "PUBLIC_URL": "http://testserver", "ROLE": "all"}.items():
        monkeypatch.setenv("SYNWEVO_" + k, v)
    from synwevo_server.app import create_app
    api = create_app(str(tmp_path / "d"), role="all", start_jobs=True)
    ident = api.state.identity
    users = {}
    for key, name in (("ann", "Ann Editor"), ("bob", "Bob Viewer"), ("cid", "Cid Editor"), ("dan", "Dan Owner"), ("eve", "Eve Elsewhere")):
        u = ident.create_user(f"{key}@example.test", PW, status="active", display_name=name, with_personal_workspace=True)
        ident.store.execute("UPDATE users SET email_verified_at=1 WHERE id=?", (u["id"],))
        users[key] = u
    from synwevo_core.organizations import Organizations
    orgs = Organizations(ident.store, ident)
    created = orgs.create(users["dan"]["id"], "Template Org", "Team A", "fixture-idem-tpl-org-0001")
    oid, wa = created["organization"]["id"], created["workspace"]["id"]
    wb = orgs.create_workspace(oid, users["dan"]["id"], "Team B")["id"]
    other = orgs.create(users["eve"]["id"], "Other Org", "Other Team", "fixture-idem-tpl-org-0002")
    for key, wid, role in (("ann", wa, "editor"), ("bob", wa, "viewer"), ("cid", wb, "editor")):
        ident.store.execute("INSERT OR IGNORE INTO organization_memberships (organization_id,user_id,role,created_at,updated_at,created_by) VALUES (?,?,'member',1,1,?)",
                            (oid, users[key]["id"], users["dan"]["id"]))
        ident.store.execute("INSERT INTO memberships (workspace_id,user_id,role,can_export,created_at,updated_at,created_by) VALUES (?,?,?,1,1,1,?)",
                            (wid, users[key]["id"], role, users["dan"]["id"]))
    clients = {}
    for key, wid in (("ann", wa), ("bob", wa), ("cid", wb), ("dan", wa), ("eve", other["workspace"]["id"])):
        c = TestClient(api)
        assert c.post("/api/auth/login", json={"email": f"{key}@example.test", "password": PW}).status_code == 200
        c.headers.update({CSRF: c.cookies.get("dw_csrf"), WS: wid})
        clients[key] = c
    try:
        yield {"api": api, "ident": ident, "c": clients, "u": users, "oid": oid, "wa": wa, "wb": wb,
               "personal": {k: ident.ensure_personal_workspace(u["id"])["id"] for k, u in users.items()}}
    finally:
        for c in clients.values():
            c.close()
        if api.state.jobs:
            api.state.jobs.stop()


def build(c) -> tuple[str, str]:
    pid = c.post("/api/projects", json={"name": "共享模板项目"}).json()["id"]
    t = c.post("/api/templates/upload", files={"file": ("模板.docx", _docx())}, data={"project_id": pid}).json()
    view = c.get(f"/api/templates/{t['id']}/view").json()
    block = next(b for part in view["doc"]["parts"] for b in _walk(part["blocks"]) if "示例客户名称" in b["text"])
    s = block["text"].index("示例客户名称")
    assert c.post(f"/api/templates/{t['id']}/slots", json={"kind": "text", "part": block["part"], "path": block["path"], "start": s, "end": s + 6, "name": "客户名称"}).status_code == 200
    return pid, t["id"]


def test_publish_use_release_and_deprecate(world):
    C, U, W = world["c"], world["u"], world
    ann, bob, cid, dan, eve = (C[k] for k in ("ann", "bob", "cid", "dan", "eve"))
    pid, tid = build(ann)
    # who may publish: a viewer is refused at the route gate; a personal workspace has no organization
    assert bob.post(f"/api/templates/{tid}/publish", json={}).status_code == 403
    ann.headers.update({WS: W["personal"]["ann"]})
    ppid = ann.post("/api/projects", json={"name": "个人"}).json()["id"]
    pt = ann.post("/api/templates/upload", files={"file": ("个人.docx", _docx())}, data={"project_id": ppid}).json()
    assert ann.post(f"/api/templates/{pt['id']}/publish", json={}).json()["error"] == "no_organization"
    assert ann.get("/api/organization-templates").json() == {"organization": None, "items": []}
    ann.headers.update({WS: W["wa"]})
    # publish freezes the head version (v2: upload + one field)
    r = ann.post(f"/api/templates/{tid}/publish", json={"description": "  客户函件标准模板 "})
    assert r.status_code == 201, r.text
    pub = r.json()
    assert pub["version_no"] == 2 and pub["release_no"] == 1 and pub["slot_count"] == 1 and pub["description"] == "客户函件标准模板"
    assert pub["published_by"]["name"] == "Ann Editor" and pub["source_workspace"]["name"] == "Team A" and pub["can_manage"] is True and pub["source_ahead"] is False
    assert ann.post(f"/api/templates/{tid}/publish", json={}).status_code == 409
    assert ann.get("/api/templates").json()[0]["publication"]["release_no"] == 1
    # visible to the organization from another workspace; not from another organization
    seen = cid.get("/api/organization-templates").json()
    assert seen["organization"]["name"] == "Template Org" and [x["id"] for x in seen["items"]] == [pub["id"]]
    assert seen["items"][0]["can_manage"] is False and seen["items"][0]["source_ahead"] is None and seen["items"][0]["source_template_id"] is None
    assert eve.get("/api/organization-templates").json()["items"] == []
    assert eve.get(f"/api/organization-templates/{pub['id']}").status_code == 404
    assert eve.post(f"/api/organization-templates/{pub['id']}/use", json={}).status_code == 404
    # a viewer may browse but not copy; the owner (org owner, admin of Team A) may maintain
    assert bob.get(f"/api/organization-templates/{pub['id']}").status_code == 200
    assert bob.post(f"/api/organization-templates/{pub['id']}/use", json={}).status_code == 403
    assert cid.post(f"/api/organization-templates/{pub['id']}/update", json={}).status_code == 403
    # cid uses it in Team B: an ordinary template of Team B with the field and the origin recorded
    bpid = cid.post("/api/projects", json={"name": "B 项目"}).json()["id"]
    r = cid.post(f"/api/organization-templates/{pub['id']}/use", json={"project_id": bpid})
    assert r.status_code == 201, r.text
    copy = r.json()
    assert copy["id"] != tid and copy["head"] == 1 and [s["name"] for s in copy["version"]["slots"]] == ["客户名称"]
    assert copy["info"]["origin"]["publication_id"] == pub["id"] and copy["info"]["origin"]["release_no"] == 1
    assert cid.get(f"/api/projects/{bpid}").json()["template_id"] == copy["id"]
    assert cid.get(f"/api/organization-templates/{pub['id']}").json()["use_count"] == 1
    # the source moves on (a rename → new version? no: a new field) — the release and the copy do not
    view = ann.get(f"/api/templates/{tid}/view").json()
    block = next(b for part in view["doc"]["parts"] for b in _walk(part["blocks"]) if "您好" in b["text"])
    s = block["text"].index("您好")
    assert ann.post(f"/api/templates/{tid}/slots", json={"kind": "text", "part": block["part"], "path": block["path"], "start": s, "end": s + 2, "name": "问候"}).status_code == 200
    assert ann.get(f"/api/organization-templates/{pub['id']}").json()["version_no"] == 2
    assert ann.get(f"/api/organization-templates/{pub['id']}").json()["source_ahead"] is True
    assert len(cid.get(f"/api/templates/{copy['id']}").json()["version"]["slots"]) == 1
    # a new release: explicit; then the organization sees release 2 with two fields
    rel = ann.post(f"/api/organization-templates/{pub['id']}/update", json={}).json()
    assert rel["release_no"] == 2 and rel["version_no"] == 3 and rel["slot_count"] == 2 and rel["released"] is True
    assert ann.post(f"/api/organization-templates/{pub['id']}/update", json={}).json()["error"] == "unchanged"
    assert ann.post(f"/api/organization-templates/{pub['id']}/update", json={"description": "第二版"}).json()["description"] == "第二版"
    assert [x["release_no"] for x in cid.get(f"/api/organization-templates/{pub['id']}").json()["releases"]] == [2, 1]
    assert len(cid.get(f"/api/templates/{copy['id']}").json()["version"]["slots"]) == 1, "the copy keeps the release it was made from"
    # deprecate: hidden from the picker, refused for new uses, existing copy untouched; the org owner may restore
    assert ann.post(f"/api/organization-templates/{pub['id']}/deprecate", json={"note": "改用新版"}).json()["deprecated"] is True
    assert cid.get("/api/organization-templates").json()["items"] == []
    assert cid.get("/api/organization-templates?include_deprecated=1").json()["items"][0]["deprecation_note"] == "改用新版"
    assert cid.post(f"/api/organization-templates/{pub['id']}/use", json={}).json()["error"] == "deprecated"
    assert cid.get(f"/api/templates/{copy['id']}").status_code == 200
    assert dan.post(f"/api/organization-templates/{pub['id']}/deprecate", json={"deprecated": False}).json()["deprecated"] is False
    assert dan.post(f"/api/organization-templates/{pub['id']}/deprecate", json={"deprecated": False}).status_code == 409
    # the source template is deleted: the release stays usable, a new release is refused
    assert ann.patch(f"/api/projects/{pid}", json={"template_id": None}).status_code == 200
    assert ann.delete(f"/api/templates/{tid}").status_code == 200
    got = cid.get(f"/api/organization-templates/{pub['id']}").json()
    assert got["source_gone"] is True and got["slot_count"] == 2
    assert cid.post(f"/api/organization-templates/{pub['id']}/use", json={}).status_code == 201
    assert ann.post(f"/api/organization-templates/{pub['id']}/update", json={}).json()["error"] == "source_gone"


TESTDATA = Path(__file__).resolve().parents[2] / "testdata"
W, H = 1190, 900
NO_BOX = [226 / W, 214 / H, 520 / W, 254 / H]          # the order number box of the synthetic picture template
AREA = [146 / W, 346 / H, 1044 / W, 606 / H]           # its 6 body rows


def _xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "客户"
    ws.append(["姓名", "编号"]); ws.append(["张三", "0007"]); ws.append(["李四", "0008"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _wait_parsed(c, pid, timeout=60):
    t0 = time.time()
    while time.time() - t0 < timeout:
        srcs = c.get(f"/api/projects/{pid}/sources").json()
        if all(s["status"] not in ("parsing", "scanning") for s in srcs):
            assert [s["status"] for s in srcs] == ["ready"], srcs
            return srcs
        time.sleep(0.2)
    raise AssertionError("sources did not settle")


def _run(c, pid, formats, timeout=180):
    j = c.post(f"/api/projects/{pid}/jobs", json={"formats": formats}).json()
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/jobs/{j['id']}").json()
        if j["status"] not in ("queued", "running", "cancelling"):
            return j
        time.sleep(0.2)
    raise AssertionError("job timeout")


def test_copies_generate_like_ordinary_templates_in_every_format(world):
    """The claim behind "使用即复制" is that nothing downstream learns a new kind of template:
    a copy made from a release binds, maps and generates exactly like an uploaded one — for a
    DOCX template and for a picture (PNG) template with a repeat area."""
    C = world["c"]
    ann, cid = C["ann"], C["cid"]
    # ---- DOCX: published from Team A, generated in Team B from Team B's own rows
    pid_a, tid = build(ann)
    pub = ann.post(f"/api/templates/{tid}/publish", json={}).json()
    pid_b = cid.post("/api/projects", json={"name": "B 函件"}).json()["id"]
    copy = cid.post(f"/api/organization-templates/{pub['id']}/use", json={"project_id": pid_b}).json()
    assert cid.post(f"/api/projects/{pid_b}/sources", files={"files": ("客户.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).status_code in (200, 201)
    _wait_parsed(cid, pid_b)
    binding = copy["version"]["slots"][0]["binding"]
    assert cid.patch(f"/api/projects/{pid_b}", json={"record_rule": {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}}}).status_code == 200
    r = cid.put(f"/api/projects/{pid_b}/mappings", json={"mappings": [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": binding}}]})
    assert r.status_code == 200, r.text
    job = _run(cid, pid_b, ["docx"])
    assert job["status"] == "completed", job
    items = cid.get(f"/api/jobs/{job['id']}/items", params={"limit": 100}).json()
    import docx
    texts = ["\n".join(p.text for p in docx.Document(io.BytesIO(cid.get(f"/api/artifacts/{a['id']}").content)).paragraphs) for it in items for a in it["artifacts"]]
    assert len(texts) == 2 and any("张三" in t for t in texts) and any("李四" in t for t in texts), texts
    assert job["snapshot"]["template"]["id"] == copy["id"] if "snapshot" in job else True
    # ---- PNG: a picture template with a number box and a repeat area, published and used the same way
    png = (TESTDATA / "测试_发货单图片模板.png").read_bytes()
    orders = (TESTDATA / "测试_订单明细100.xlsx").read_bytes()
    pid_img = ann.post("/api/projects", json={"name": "图片模板"}).json()["id"]
    t = ann.post("/api/templates/upload", files={"file": ("发货单.png", png, "image/png")}, data={"project_id": pid_img}).json()
    assert t["format"] == "png", t
    assert ann.post(f"/api/templates/{t['id']}/slots", json={"kind": "pdf_box", "page": 0, "nrect": NO_BOX, "name": "单号"}).status_code == 200
    assert ann.post(f"/api/templates/{t['id']}/slots", json={"kind": "pdf_repeat", "page": 0, "nrect": AREA, "name": "明细", "overflow": "shrink"}).status_code == 200
    pub2 = ann.post(f"/api/templates/{t['id']}/publish", json={"description": "发货单图片"}).json()
    assert pub2["format"] == "png" and pub2["slot_count"] == 2
    pid_b2 = cid.post("/api/projects", json={"name": "B 发货单"}).json()["id"]
    copy2 = cid.post(f"/api/organization-templates/{pub2['id']}/use", json={"project_id": pid_b2}).json()
    assert copy2["format"] == "png" and copy2["info"]["image"]["width"] == W and copy2["info"]["origin"]["publication_id"] == pub2["id"]
    assert cid.post(f"/api/projects/{pid_b2}/sources", files={"files": ("订单.xlsx", orders, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).status_code in (200, 201)
    _wait_parsed(cid, pid_b2)
    slots = {s_["name"]: s_["binding"] for s_ in copy2["version"]["slots"]}
    assert cid.patch(f"/api/projects/{pid_b2}", json={"record_rule": {"mode": "per_group", "table": {"sheet": "明细", "header_row": 0},
                                                                       "group": {"keys": [{"header": "订单号"}], "sort": [{"header": "行号", "dir": "asc", "numeric": True}]}}}).status_code == 200
    r = cid.put(f"/api/projects/{pid_b2}/mappings", json={"mappings": [
        {"type": "single", "sources": [{"kind": "column", "sheet": "明细", "header": "订单号"}], "target": {"binding": slots["单号"]}},
        {"type": "repeat", "sources": [{"kind": "table", "sheet": "明细", "header_row": 0}], "target": {"binding": slots["明细"]},
         "columns": [{"slot_col": "c0", "source": {"header": "行号"}}, {"slot_col": "c1", "source": {"header": "品名"}},
                     {"slot_col": "c2", "source": {"header": "数量"}}, {"slot_col": "c3", "source": {"header": "单价"}}]}]})
    assert r.status_code == 200, r.text
    job = _run(cid, pid_b2, ["png"])
    assert job["status"] == "completed", job
    items = cid.get(f"/api/jobs/{job['id']}/items", params={"limit": 200}).json()
    assert items and all(it["status"] == "success" for it in items), [it["status"] for it in items]
    first = cid.get(f"/api/artifacts/{items[0]['artifacts'][0]['id']}").content
    assert first[:8] == b"\x89PNG\r\n\x1a\n"
    from PIL import Image
    im = Image.open(io.BytesIO(first))
    assert im.size == (W, H), "the output keeps the picture template's pixel size"
