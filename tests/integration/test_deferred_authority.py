"""Synthetic isolated regression fixtures for authority after the original request ends."""
from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from synwevo_core import authority
from synwevo_core.machine import MachineAuth
from synwevo_core.oidc import Clients, Grants
from synwevo_docs.jobs import JobService
from synwevo_server.app import create_app


@pytest.fixture()
def deferred(tmp_path, monkeypatch):
    monkeypatch.setenv("SYNWEVO_UI_AUTH", "session")
    monkeypatch.setenv("SYNWEVO_COOKIE_SECURE", "0")
    api = create_app(str(tmp_path / "data"), role="app", start_jobs=True)
    api.state.jobs.stop()  # Real queue/renderer, driven explicitly below to control revocation.
    store, ident = api.state.app.store, api.state.identity
    user = ident.create_user("deferred@example.test", "fixture-only-long-passphrase", status="active")
    wid = ident.workspaces_for(user["id"])[0]["id"]
    clients, grants = Clients(store), Grants(store)
    clients.register(client_id="deferred-user", app_id="docweave", name="Synthetic delegation",
                     redirect_uris=["http://localhost/callback"],
                     scopes="openid docweave.read docweave.write")
    client = clients.get("deferred-user")
    token = grants.issue_tokens(client=client, code_row={"user_id": user["id"],
        "workspace_id": wid, "scope": "docweave.read docweave.write", "central_session_id": None},
        with_refresh=True)
    machine = MachineAuth(store, ident, api.state.auth)
    principal = machine.verify("Bearer " + token["access_token"])
    with TestClient(api) as http:
        yield {"api": api, "store": store, "user": user, "wid": wid, "clients": clients,
               "grants": grants, "token": token, "machine": machine, "principal": principal,
               "desc": authority.describe(principal), "http": http}


def test_delegation_rechecks_registered_scopes(deferred):
    d = deferred
    assert authority.check(d["store"], d["desc"], need="write", tenant=d["wid"])[0]
    d["store"].execute("UPDATE oauth_clients SET scopes='openid docweave.read' WHERE client_id=?",
                       (d["principal"]["client_id"],))
    assert authority.check(d["store"], d["desc"], need="read", tenant=d["wid"])[0]
    assert authority.check(d["store"], d["desc"], need="write", tenant=d["wid"]) == (False, "scope_withdrawn")


def test_delegation_requires_its_live_grant(deferred):
    d = deferred
    d["store"].execute("DELETE FROM oauth_tokens WHERE family_id=?", (d["desc"]["family_id"],))
    assert authority.check(d["store"], d["desc"], need="read", tenant=d["wid"]) == (False, "grant_revoked")


@pytest.mark.parametrize("raw", ["{broken", "null", "[]", "{}", '{"tenant":"x"}'])
def test_corrupt_descriptor_never_becomes_legacy_operator(deferred, raw):
    desc = authority.loads(raw)
    assert not authority.check(deferred["store"], desc, tenant=deferred["wid"])[0]
    assert not authority.check(deferred["store"], authority.loads(authority.dumps(desc)),
                               tenant=deferred["wid"])[0]


def test_delegation_survives_normal_refresh_but_not_family_revocation(deferred):
    d = deferred
    d["grants"].rotate_refresh(d["token"]["refresh_token"], client=d["clients"].get("deferred-user"))
    assert authority.check(d["store"], d["desc"], need="write", tenant=d["wid"]) == (True, None)
    d["grants"].revoke_family(d["desc"]["family_id"], "security_revoked")
    assert authority.check(d["store"], d["desc"], need="write", tenant=d["wid"]) == (False, "grant_revoked")


def test_delegation_cannot_borrow_another_users_grant(deferred):
    d = deferred
    d["store"].execute("UPDATE oauth_tokens SET user_id='another-synthetic-user' WHERE family_id=?",
                       (d["desc"]["family_id"],))
    assert authority.check(d["store"], d["desc"], need="read", tenant=d["wid"]) == (False, "grant_revoked")


def test_service_authority_rechecks_application_and_read_scope(deferred):
    d = deferred
    d["clients"].register(client_id="deferred-service", app_id="docweave", name="Synthetic service",
        redirect_uris=[], kind="service", confidential=True, workspace_id=d["wid"],
        api_scopes="docweave.read docweave.write")
    desc = {"kind": "service_account", "client_id": "deferred-service", "tenant": d["wid"]}
    assert authority.check(d["store"], desc, need="export", tenant=d["wid"]) == (True, None)
    d["store"].execute("UPDATE oauth_clients SET api_scopes='docweave.write' WHERE client_id='deferred-service'")
    assert authority.check(d["store"], desc, need="export", tenant=d["wid"]) == (False, "scope_withdrawn")
    assert authority.check(d["store"], desc, need="write", tenant=d["wid"]) == (True, None)
    d["store"].execute("UPDATE oauth_clients SET app_id='another-app' WHERE client_id='deferred-service'")
    assert authority.check(d["store"], desc, need="write", tenant=d["wid"]) == (False, "wrong_audience")


def _ok(response):
    assert response.status_code < 300, response.text
    return response.json()


def _generated_output(d):
    """Create actual XLSX/DOCX inputs, fields, mapping and output through business APIs."""
    c = d["http"]
    headers = {"Authorization": "Bearer " + d["token"]["access_token"]}
    def request(method, path, **kw):
        return _ok(c.request(method, "/api/v1" + path, headers=headers, **kw))
    project = request("POST", "/projects", json={"name": "Synthetic download authority"})["id"]
    book = Workbook()
    book.active.title = "Rows"
    book.active.append(["Name"])
    book.active.append(["Synthetic Customer"])
    source = io.BytesIO()
    book.save(source)
    request("POST", f"/projects/{project}/files", files={"file": ("source.xlsx", source.getvalue())})
    doc = Document()
    doc.add_paragraph("Customer: PLACEHOLDER")
    template = io.BytesIO()
    doc.save(template)
    tid = request("POST", "/templates", files={"file": ("template.docx", template.getvalue())})["id"]
    request("POST", f"/templates/{tid}/fields", json={"schema": "docweave.template-fields", "version": 1,
        "fields": [{"kind": "text", "name": "Customer", "find": "PLACEHOLDER"}]})
    request("PUT", f"/projects/{project}/template", json={"template_id": tid})
    request("PUT", f"/projects/{project}/rule", json={"rule": {
        "mode": "per_row", "table": {"sheet": "Rows", "header_row": 0}}})
    binding = request("GET", f"/templates/{tid}")["fields"][0]["binding"]
    request("PUT", f"/projects/{project}/mappings", json={"mappings": [{"type": "single",
        "sources": [{"kind": "column", "sheet": "Rows", "header": "Name"}], "target": {"binding": binding}}]})
    jid = request("POST", f"/projects/{project}/jobs", json={"formats": ["docx"]})["id"]
    jobs = JobService(d["api"].state.app, start=False)
    assert jobs._claim(jid) == jid
    jobs._run(jid)
    assert request("GET", f"/jobs/{jid}")["status"] == "completed"
    artifact = request("GET", f"/jobs/{jid}/artifacts")["items"][0]
    content = c.get(f"/api/v1/artifacts/{artifact['id']}/content", headers=headers)
    assert content.status_code == 200
    assert "Synthetic Customer" in "\n".join(p.text for p in Document(io.BytesIO(content.content)).paragraphs)
    return jid, artifact["id"], headers


def test_export_withdrawal_invalidates_previously_minted_links(deferred):
    d = deferred
    jid, aid, headers = _generated_output(d)
    c = d["http"]
    links = [_ok(c.post(path, json={}, headers=headers))["url"] for path in
             (f"/api/v1/artifacts/{aid}/link", f"/api/v1/jobs/{jid}/archive-link")]
    for url in links:
        assert c.get(url).status_code == 200
    # FIXTURE mutation models an administrator withdrawing this member's export permission.
    d["store"].execute("UPDATE memberships SET role='viewer', can_export=0 WHERE user_id=? AND workspace_id=?",
                       (d["user"]["id"], d["wid"]))
    assert c.get(f"/api/v1/jobs/{jid}", headers=headers).status_code == 200
    assert c.get(f"/api/v1/artifacts/{aid}/content", headers=headers).status_code == 403
    for url in links:
        result = c.get(url)
        assert result.status_code == 403
        assert result.json()["error"]["code"] == "link_revoked"
    # A later explicit grant restores access; the link's original expiry still bounds it.
    d["store"].execute("UPDATE memberships SET can_export=1 WHERE user_id=? AND workspace_id=?",
                       (d["user"]["id"], d["wid"]))
    assert all(c.get(url).status_code == 200 for url in links)
