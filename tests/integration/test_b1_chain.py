"""Acceptance B1 (with B2, B3, B5): one business chain driven by the real flow engine.

Forms submission with a real PDF -> Extract import, parse and a person's confirmation -> DataPrep
dataset and published recipe run (partial on a row issue) -> Tables pending-review record ->
Approvals request decided by a different person -> Docs project from the dataset, job, artifact ->
Message send to a workspace-bound group.

Real composed factory (``synwevo_server.app.create_app(role="all")``), real services and the real
``synwevo_core.flows.Flows`` engine. Fixture people, organization, workspaces and the Message
group come from the WP7 ``world`` fixture; every business object below is made through its
application's registered operation or public route (the Docs template project through the Docs
service calls the WP7 ``configure`` helper uses). Nothing here writes a business table.

Fault injection is limited to: rebuilding the factory on the same data directory (B3), rewinding
the engine's own wait-event cursor (B3), stale approval event hints while the real request is
still waiting, losing a receipt update after Message has stored the real message (B4), and
ticking the engine from two threads (B5). No business operation or reader is replaced.
Workers that a deployment runs on timers are stepped explicitly: Extract's and DataPrep's
``tick``, Docs' ``_claim``/``_run``, Message's outbox worker and Suite's receipt reconcilers.
"""
from __future__ import annotations

import hashlib
import io
import json
import threading
from datetime import datetime, timezone
from types import SimpleNamespace

import docx
import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient

from synwevo_core.apps import DOCS_APP_ID, MESSAGES_APP_ID
from synwevo_core.capabilities import record_event
from synwevo_core.context import as_tenant
from synwevo_core.errors import ServiceError
from synwevo_core.flows import Flows
from synwevo_docs.definitions import TEMPLATE_FIELDS_SCHEMA, import_template_fields
from synwevo_server.app import create_app
from .test_suite_adapters import ORIGIN, deliver, sign_in, wait_for, world  # noqa: F401

pytestmark = pytest.mark.parametrize("world", ["forms,extract,dataprep,tables,approvals,agent"], indirect=True)

CSRF = {"X-Synwevo-CSRF": "fixture-csrf"}
REFERENCE, AMOUNT, AMOUNT_2DP = "INV-B1-0042", "12500.5", "12500.50"
# Every application event the chain emits, each exactly once per run.
CHAIN_EVENTS = (
    "forms.request.created", "forms.submission.created", "extract.file.uploaded", "extract.file.imported",
    "extract.task.queued", "extract.task.started", "extract.task.ready", "extract.task.confirmation",
    "dataprep.dataset.created", "dataprep.run.queued", "dataprep.run.partial", "tables.record.changed",
    "approvals.request.created", "approvals.stage.opened", "approvals.task.decided", "approvals.stage.approved",
    "approvals.request.approved", "docs.project.materialized", "docs.job.submitted", "messages.message.recorded")
# Every operation the flow itself calls (call steps), each exactly once per run.
FLOW_CALLS = (
    "forms.request.create", "forms.submission.get", "extract.file.from_forms", "extract.task.create",
    "dataprep.dataset.from_extract", "dataprep.run.start", "tables.record.create", "approvals.request.create",
    f"{DOCS_APP_ID}.project.from_dataset", f"{DOCS_APP_ID}.project.get", f"{DOCS_APP_ID}.job.create",
    f"{DOCS_APP_ID}.job.artifacts", f"{MESSAGES_APP_ID}.conversation.get", f"{MESSAGES_APP_ID}.message.send")


class Clock:
    """The engine's clock (fixture). Business services keep real time."""

    def __init__(self):
        self.at = datetime(2026, 9, 30, 0, 0, tzinfo=timezone.utc).timestamp()

    def __call__(self):
        return self.at

    def advance(self, seconds=31):
        self.at += seconds


def invoice_pdf(*, conflicting_reference=False):
    """A real one-page PDF with a text layer, as tests/integration/test_extract_reader.py makes it."""
    doc = fitz.open()
    page = doc.new_page()
    text = f"Invoice: {REFERENCE}\nAmount: {AMOUNT}\nIssued: 2026-09-30"
    if conflicting_reference:
        text += "\nInvoice: INV-B1-CONFLICT"
    page.insert_text((60, 80), text)
    data = doc.tobytes()
    doc.close()
    return data


def docs_template_project(w):
    """Docs template project (same service calls as test_suite_adapters.configure), mapped by the
    DataPrep column names so the materialized project reads the dataset's own columns."""
    app = w.adapters.docs
    with as_tenant(w.a):
        pid = app.create_project("B1 release document", origin="test:b1")["id"]
        wb = openpyxl.Workbook()
        wb.active.title = "Input"
        wb.active.append(["reference", "amount"])
        wb.active.append(["TEMPLATE-ROW", "0.00"])
        data = io.BytesIO(); wb.save(data)
        source = app.add_source(pid, "template-input.xlsx", data.getvalue())
        wait_for(lambda: app.get_source(source["id"])["status"] == "ready")
        template = docx.Document()
        template.add_paragraph("Invoice REFSLOT")
        template.add_paragraph("Amount AMTSLOT")
        data = io.BytesIO(); template.save(data)
        tpl = app.upload_template("b1-release.docx", data.getvalue())
        import_template_fields(app, tpl["id"], {"schema": TEMPLATE_FIELDS_SCHEMA, "version": 1, "fields": [
            {"kind": "text", "name": "Reference", "find": "REFSLOT"},
            {"kind": "text", "name": "Amount", "find": "AMTSLOT"}]})
        head = app.get_template(tpl["id"])["head"]
        slots = app.get_version(tpl["id"], head)["slots"]
        assert len(slots) == 2
        app.update_project(pid, {"template_id": tpl["id"], "template_version": head,
                                 "record_rule": {"mode": "per_row", "table": {"sheet": "Input", "header_row": 0}}})
        app.save_mappings(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "Input", "header": column}],
                                 "target": {"binding": slot["binding"]}}
                                for column, slot in zip(("reference", "amount"), slots)])
    return pid


def setup_apps(w):
    """Published Forms form, Extract template, DataPrep recipe, Tables table and Approvals template."""
    for app in ("forms", "extract", "dataprep", "tables", "approvals"):
        w.apps.enable_for_workspace(w.a, app)
    form = w.read("form.create", {"title": "B1 invoice intake", "definition": {"fields": [
        {"id": "f_invoice01", "type": "file", "label": "Invoice", "accept": [".pdf"], "required": True}]}}, app="forms")
    w.read("form.publish", {"form_id": form["id"], "draft_rev": form["draft_rev"]}, app="forms")
    extract = w.read("template.create", {"title": "B1 invoice", "definition": {"fields": [
        {"id": "reference", "label": "Reference", "rule": {"prefix": "Invoice:"}},
        {"id": "amount", "label": "Amount", "type": "decimal", "rule": {"prefix": "Amount:"}}]}}, app="extract")
    w.read("template.publish", {"template_id": extract["id"], "draft_rev": 1}, app="extract")
    recipe = w.read("recipe.create", {"name": "B1 invoice checks", "definition": {"steps": [
        {"id": "trim_ref", "type": "trim", "columns": ["reference"]},
        {"id": "amount_2dp", "type": "cast", "column": "amount", "to": "decimal", "scale": 2},
        {"id": "review_limit", "type": "validate", "message": "Amount above the review threshold", "on_error": "keep",
         "condition": {"column": "amount", "op": "le", "value": "10000", "numeric": True}}]}}, app="dataprep")
    w.read("recipe.publish", {"recipe_id": recipe["id"], "draft_rev": recipe["draft_rev"]}, app="dataprep")
    table = w.read("table.create", {"name": "B1 pending review"}, app="tables", user="owner")
    rev, fields = table["schema_rev"], {}
    for name, kind, config in (("Dataset", "text", {}), ("Issues", "number", {"decimals": 0}), ("Status", "text", {})):
        saved = w.read("field.save", {"table_id": table["id"], "schema_rev": rev,
                                      "field": {"name": name, "type": kind, "config": config}}, app="tables", user="owner")
        rev, fields[name] = saved["schema_rev"], saved["field_id"]
    approval = w.read("template.create", {"title": "B1 document release", "definition": {
        "fields": [{"id": "f_dataset", "type": "text", "label": "Dataset", "required": True},
                   {"id": "f_issues", "type": "number", "label": "Row issues"}],
        "stages": [{"id": "s_release", "name": "Release review", "mode": "all",
                    "approvers": {"users": [w.users["owner"]]}}]}}, app="approvals", user="owner")
    w.read("template.publish", {"template_id": approval["id"], "draft_rev": 1}, app="approvals", user="owner")
    return SimpleNamespace(form=form["id"], extract=extract["id"], recipe=recipe["id"], table=table["id"],
                           fields=fields, approval=approval["id"], template_project=docs_template_project(w))


def definition(w, s):
    ws = w.a

    def ref(path):
        return {"$ref": path}

    def call(id, operation, payload, **kw):
        return {"id": id, "kind": "call", "operation": operation, "workspace_id": ws, "input": payload, **kw}

    return {"steps": [
        call("request", "forms.request.create", {"form_id": s.form, "version": 1},
             result_read={"operation": "forms.request.get", "workspace_id": ws, "input": {"request_id": ref("receipt.request_id")}}),
        {"id": "collect", "kind": "wait", "read": {"operation": "forms.request.get", "workspace_id": ws,
                                                    "input": {"request_id": ref("steps.request.id")}},
         "until": {"left": ref("record.status"), "right": "submitted"}, "event_types": ["forms.submission.created"]},
        call("submission", "forms.submission.get", {"submission_id": ref("steps.collect.submissions.0.id")}),
        call("import", "extract.file.from_forms", {"submission_id": ref("steps.collect.submissions.0.id"),
                                                   "file_id": ref("steps.submission.files.0.id")}),
        call("extract", "extract.task.create", {"template_id": s.extract, "version": 1, "file_id": ref("steps.import.id")}),
        call("dataset", "dataprep.dataset.from_extract", {"name": "B1 confirmed invoice", "field_ids": ["reference", "amount"],
             "tasks": [{"task_id": ref("steps.extract.id"), "revision": ref("steps.extract.confirmed_revision")}]}),
        call("prep", "dataprep.run.start", {"recipe_id": s.recipe, "version": 1, "inputs": {"main": ref("steps.dataset.id")},
                                            "output_name": "B1 checked invoice"}),
        {"id": "issues", "kind": "branch", "condition": {"left": ref("outcome.prep"), "right": "partial"},
         "then": "record", "otherwise": "approval"},
        call("record", "tables.record.create", {"table_id": s.table, "cells": {
            s.fields["Dataset"]: ref("steps.prep.output_dataset_id"), s.fields["Issues"]: ref("steps.prep.issue_total"),
            s.fields["Status"]: "pending_review"}}),
        call("approval", "approvals.request.create", {"template_id": s.approval, "version": 1, "title": "Release the B1 invoice document",
             "values": {"f_dataset": ref("steps.prep.output_dataset_id"), "f_issues": ref("steps.prep.issue_total")},
             "basis": ref("steps.prep.output_dataset_id")},
             result_read={"operation": "approvals.request.get", "workspace_id": ws, "input": {"request_id": ref("receipt.id")}}),
        {"id": "review", "kind": "wait", "state": "waiting_approval",
         "read": {"operation": "approvals.request.get", "workspace_id": ws, "input": {"request_id": ref("steps.approval.id")}},
         "until": {"left": ref("record.status"), "op": "in", "right": ["approved", "rejected"]},
         "reject": {"left": ref("record.status"), "op": "in", "right": ["returned", "withdrawn"]},
         "event_types": ["approvals.request.approved", "approvals.request.rejected"]},
        {"id": "route", "kind": "branch", "condition": {"left": ref("steps.review.status"), "right": "approved"},
         "then": "materialize", "otherwise": "rejected"},
        call("materialize", f"{DOCS_APP_ID}.project.from_dataset",
             {"template_project_id": s.template_project, "dataset_id": ref("steps.prep.output_dataset_id")}),
        call("project", f"{DOCS_APP_ID}.project.get", {"project_id": ref("steps.materialize.project_id")}),
        call("job", f"{DOCS_APP_ID}.job.create", {"project_id": ref("steps.materialize.project_id"),
                                                  "configuration": ref("steps.project.configuration"), "formats": ["docx"]}),
        call("artifacts", f"{DOCS_APP_ID}.job.artifacts", {"job_id": ref("steps.job.job_id")}),
        call("conversation", f"{MESSAGES_APP_ID}.conversation.get", {"conversation_id": w.cid}),
        call("send", f"{MESSAGES_APP_ID}.message.send", {"conversation_id": w.cid, "audience": ref("steps.conversation.audience"),
             "body": "B1 release document approved", "artifact_ids": [ref("steps.artifacts.items.0.artifact_id")]}, next="end"),
        {"id": "rejected", "kind": "fail", "code": "approval_rejected"},
    ]}


class Chain:
    """Drives one flow run owned by the editor; the owner is the independent approver."""

    def __init__(self, w):
        self.w, self.clock, self.s = w, Clock(), setup_apps(w)
        self.owner = w.users["editor"]
        self.api, self.store, self.caps, self.adapters = w.api, w.store, w.caps, w.adapters
        self.engine = Flows(w.store, w.caps, clock=self.clock)
        self.fid = self.engine.create(user_id=self.owner, name="B1 invoice to released document",
                                      definition=definition(w, self.s))["flow_id"]
        self.rid = None

    def rebuild(self, api):
        """Use a fresh factory's services from here on (same data directory)."""
        self.api, ctx = api, api.state.suite
        self.store, self.caps, self.adapters = ctx.store, ctx.capabilities, api.state.suite_adapters
        self.engine = Flows(ctx.store, ctx.capabilities, clock=self.clock)

    def tick(self):
        self.clock.advance(31)  # async receipts are re-polled 30 s after release
        self.engine.tick()

    def run(self):
        return self.store.one("SELECT * FROM flow_runs WHERE id=?", (self.rid,))

    def steps(self):
        return {r["step_id"]: r for r in self.store.all("SELECT * FROM flow_steps WHERE run_id=?", (self.rid,))}

    def op(self, operation):
        rows = self.store.all("SELECT * FROM op_runs WHERE parent_run_id=? AND op=?", (self.rid, operation))
        assert len(rows) == 1, (operation, len(rows))
        return rows[0], decode(rows[0]["output"])

    def read(self, operation, payload, *, user="editor"):
        return self.caps.invoke(operation, payload, user_id=self.w.users[user], workspace_id=self.w.a)["output"]

    def start(self):
        w = self.w
        self.rid = self.engine.start(self.fid, user_id=self.owner, version=1, parameters={},
                                     idempotency_key="b1-chain")["id"]
        self.engine.tick()
        assert self.run()["status"] == "waiting_input" and self.run()["step_id"] == "collect"
        request = decode(self.steps()["request"]["value_json"])
        self.form_url = "/api/suite/public/forms/r/" + request["url"].rsplit("/", 1)[1]
        sign_in(w, "editor")

    def submit_file(self, pdf):
        """Submit actual bytes and let the registered reader parse them, without confirming."""
        w = self.w
        self.pdf_digest = hashlib.sha256(pdf).hexdigest()
        up = w.client.post(self.form_url + "/files", data={"field_id": "f_invoice01"}, headers=CSRF,
                           files={"file": ("invoice-b1.pdf", pdf, "application/pdf")})
        assert up.status_code == 200, up.text
        self.form_file_id = up.json()["id"]
        sub = w.client.post(self.form_url + "/submit", headers=CSRF, json={
            "version": 1, "answers": {"f_invoice01": [self.form_file_id]}, "idempotency_key": "b1-form-submit"})
        assert sub.status_code == 200, sub.text
        self.submission_id = sub.json()["submission_id"]
        self.tick()
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_input", "extract"), self.run()
        _, task = self.op("extract.task.create")
        self.task_id = task["id"]
        self.api.state.extract.tick()  # the deployed reader parses the actual PDF
        parsed = self.read("extract.task.get", {"task_id": task["id"]})
        assert parsed["status"] == "waiting_input", parsed
        return parsed

    def to_approval(self):
        """Start, a real submission, Extract parse + a person's confirmation, DataPrep run."""
        self.start()
        parsed = self.submit_file(invoice_pdf())
        assert parsed["values"]["reference"]["value"] == REFERENCE and parsed["values"]["amount"]["value"] == AMOUNT
        self.tick()  # a machine result is not a confirmed one: the flow keeps waiting for a person
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_input", "extract")
        # Person-only operation, called by the person while the flow waits.
        self.caps.invoke("extract.task.confirm", {"task_id": self.task_id, "revision": 1, "note": "Checked against the PDF"},
                         user_id=self.owner, workspace_id=self.w.a, via="ui")
        self.to_review()

    def to_review(self):
        """Continue the confirmed extraction through real DataPrep and the approval request."""
        self.tick()
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_input", "prep"), self.run()
        assert self.api.state.dataprep.tick()
        self.tick()
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_approval", "review"), self.run()
        self.approval_id = decode(self.steps()["approval"]["value_json"])["id"]

    def decide(self, decision):
        current = self.read("approvals.request.get", {"request_id": self.approval_id}, user="owner")
        task = next(t for t in current["tasks"] if t["assignee_id"] == self.w.users["owner"] and t["status"] == "pending")
        # Person-only operation (via="ui"), by someone other than the flow owner who requested it.
        self.caps.invoke("approvals.request.decide", {"request_id": self.approval_id, "task_id": task["id"],
                         "state_version": current["state_version"], "decision": decision, "note": "B1 independent review"},
                         user_id=self.w.users["owner"], workspace_id=self.w.a, via="ui")

    def to_send(self, tick=None):
        """Docs input and job settle; Message has accepted its durable queued send."""
        tick = tick or self.tick
        tick()
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_input", "materialize"), self.run()
        _, made = self.op(f"{DOCS_APP_ID}.project.from_dataset")
        with as_tenant(self.w.a):
            wait_for(lambda: self.adapters.docs.get_source(made["source_id"])["status"] == "ready")
        self.adapters.sync_inputs()
        tick()
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_input", "job"), self.run()
        _, job = self.op(f"{DOCS_APP_ID}.job.create")
        jobs = self.adapters.jobs
        assert jobs._claim(job["job_id"]) == job["job_id"]
        jobs._run(job["job_id"])
        self.adapters.sync_jobs()
        tick()
        assert (self.run()["status"], self.run()["step_id"]) == ("waiting_input", "send"), self.run()
        return made, job

    def to_end(self, tick=None):
        """Docs input, Docs job and Message delivery, each settled by its own worker."""
        tick = tick or self.tick
        made, job = self.to_send(tick)
        deliver(self.api, self.adapters)
        tick()
        return made, job


def decode(value):
    return json.loads(value) if isinstance(value, (str, bytes)) else value


def count(store, sql, args=()):
    return store.one("SELECT COUNT(*) AS n FROM " + sql, args)["n"]


def dispatches(c):
    """Call steps the engine started for this run: one per call step means nothing was sent twice."""
    return count(c.store, "flow_history WHERE run_id=? AND type='step_started'", (c.rid,))


def group_messages(store, cid):
    return store.all("SELECT * FROM messages WHERE conversation_id=? AND kind!='system'", (cid,))


def assert_no_release(c):
    assert not {"materialize", "project", "job", "artifacts", "conversation", "send"} & set(c.steps())
    assert count(c.store, "projects WHERE origin LIKE 'suite:%'") == 0
    assert count(c.store, "jobs") == 0 and count(c.store, "artifacts") == 0
    assert group_messages(c.store, c.w.cid) == []


def assert_one_document_and_message(c, w, made, job):
    store = c.store
    assert count(store, "projects WHERE origin LIKE 'suite:%'") == 1
    assert count(store, "jobs WHERE project_id=?", (made["project_id"],)) == 1
    assert count(store, "jobs") == 1
    artifacts = store.all("SELECT id FROM artifacts WHERE job_id=?", (job["job_id"],))
    assert len(artifacts) == 1
    messages = group_messages(store, w.cid)
    assert len(messages) == 1, messages
    link = c.read(f"{DOCS_APP_ID}.artifact.get", {"artifact_id": artifacts[0]["id"]})["download_url"]
    assert messages[0]["sender_id"] == w.users["editor"] and link in messages[0]["body"]
    return artifacts[0]["id"], link


def test_b1_form_pdf_to_approved_document_in_the_group(world):
    w = world
    c = Chain(w)
    c.to_approval()
    # DataPrep found a row issue: the step completed partial and the flow took the review branch.
    steps = c.steps()
    assert steps["prep"]["status"] == "partial"
    assert decode(steps["issues"]["value_json"]) == {"target": "record"}
    _, prep = c.op("dataprep.run.start")
    assert prep["status"] == "partial" and prep["issue_total"] == 1
    c.decide("approve")
    made, job = c.to_end()

    run = c.run()
    assert run["status"] == "partial" and run["error_code"] is None and run["finished_at"]  # never plain success
    steps = c.steps()
    assert {k: v["status"] for k, v in steps.items()} == {
        **{k: "succeeded" for k in ("request", "collect", "submission", "import", "extract", "dataset", "issues",
                                    "record", "approval", "review", "route", "materialize", "project", "job",
                                    "artifacts", "conversation", "send")}, "prep": "partial"}
    assert "rejected" not in steps

    # The document carries the PDF's values as Extract read them and DataPrep normalized them.
    aid, link = assert_one_document_and_message(c, w, made, job)
    _, data = c.adapters.jobs.artifact(aid, w.a)
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    assert f"Invoice {REFERENCE}" in text and f"Amount {AMOUNT_2DP}" in text and "TEMPLATE-ROW" not in text
    sign_in(w, "editor")
    download = w.client.get(link)
    assert download.status_code == 200 and download.content == data
    # Provenance: the dataset came from the confirmed revision of the task that read the form's file.
    dataset = c.read("dataprep.dataset.get", {"dataset_id": prep["output_dataset_id"]})
    assert dataset["rows"][0]["values"] == {"reference": REFERENCE, "amount": AMOUNT_2DP}
    _, imported = c.op("extract.file.from_forms")
    assert imported["source"] == {
        "app": "forms", "workspace_id": w.a, "submission_id": c.submission_id,
        "file_id": c.form_file_id, "form_id": c.s.form, "form_version": 1,
        "answers_hash": c.read("forms.submission.get", {"submission_id": c.submission_id})["answers_hash"]}
    assert imported["content_digest"] == c.pdf_digest
    confirmed = c.read("extract.task.result", {"task_id": c.task_id, "revision": 2})
    assert confirmed["file"]["id"] == imported["id"] and confirmed["file"]["content_digest"] == c.pdf_digest
    assert confirmed["review"]["actor_id"] == w.users["editor"]
    segments = c.read("extract.task.segments", {"task_id": c.task_id})["items"]
    for field, expected in (("reference", REFERENCE), ("amount", AMOUNT)):
        candidate, = confirmed["machine"][field]["candidates"]
        evidence = candidate["evidence"]
        assert candidate["verified"] and candidate["value"] == expected
        assert evidence["source"]["page"] == 1
        segment = next(s for s in segments if s["id"] == evidence["segment_id"])
        assert segment["text"][evidence["start"]:evidence["end"]] == expected
    _, source_dataset = c.op("dataprep.dataset.from_extract")
    source, = source_dataset["source"]["tasks"]
    assert {k: source[k] for k in ("task_id", "revision", "file_id", "content_digest", "template_id", "template_version")} == {
        "task_id": c.task_id, "revision": 2, "file_id": imported["id"], "content_digest": c.pdf_digest,
        "template_id": c.s.extract, "template_version": 1}
    assert dataset["source"]["inputs"]["main"]["dataset_id"] == source_dataset["id"]
    assert dataset["rows"][0]["origins"] == [[c.task_id, 1, 2]]
    issues = c.read("dataprep.run.issues", {"run_id": prep["run_id"]})
    assert issues["total"] == 1 and issues["next_offset"] is None
    issue, = issues["items"]
    assert {k: issue[k] for k in ("code", "step_id", "column_name", "value", "input_name", "source_row", "message")} == {
        "code": "rule_failed", "step_id": "review_limit", "column_name": "amount", "value": AMOUNT_2DP,
        "input_name": c.task_id, "source_row": 1, "message": "Amount above the review threshold"}
    assert issue["params"]["origins"] == [[c.task_id, 1, 2]]

    # Tables holds the pending-review record; Approvals the decided request.
    record_run, record = c.op("tables.record.create")
    stored = c.read("tables.record.get", {"table_id": c.s.table, "record_id": record["id"]})
    assert stored["cells"] == {c.s.fields["Dataset"]: prep["output_dataset_id"], c.s.fields["Issues"]: "1",
                               c.s.fields["Status"]: "pending_review"}
    assert stored["created_by"] == w.users["editor"]
    request = c.read("approvals.request.get", {"request_id": c.approval_id}, user="owner")
    assert request["status"] == "approved" and request["requester_id"] == w.users["editor"]
    assert request["snapshot"]["fields"] == {"f_dataset": prep["output_dataset_id"], "f_issues": "1"}
    assert [(t["assignee_id"], t["status"]) for t in request["tasks"]] == [(w.users["owner"], "approved")]

    # Every application event along the chain exists exactly once.
    counts = {r["type"]: r["n"] for r in c.store.all("SELECT type, COUNT(*) AS n FROM app_events GROUP BY type")}
    assert {t: counts.get(t, 0) for t in CHAIN_EVENTS} == {t: 1 for t in CHAIN_EVENTS}
    # And each flow operation ran once, with the flow's delegation.
    calls = c.store.all("SELECT op, via, user_id FROM op_runs WHERE parent_run_id=?", (c.rid,))
    assert sorted(r["op"] for r in calls) == sorted(FLOW_CALLS) and {r["via"] for r in calls} == {"flow"}
    assert {r["user_id"] for r in calls} == {w.users["editor"]}
    assert dispatches(c) == len(FLOW_CALLS) and counts["agent.flow.started"] == 1


def test_b2_missing_file_and_unresolved_extraction_do_not_start_downstream(world):
    c = Chain(world)
    c.start()
    missing = world.client.post(c.form_url + "/submit", headers=CSRF, json={
        "version": 1, "answers": {}, "idempotency_key": "b1-missing-file"})
    assert missing.status_code == 422
    assert ("f_invoice01", "required") in {(e["field"], e["code"]) for e in missing.json()["detail"]["errors"]}
    c.tick()
    assert (c.run()["status"], c.run()["step_id"]) == ("waiting_input", "collect")
    assert count(c.store, "form_submissions") == 0 and count(c.store, "extract_tasks") == 0
    assert set(c.steps()) == {"request", "collect"}
    assert_no_release(c)

    parsed = c.submit_file(invoice_pdf(conflicting_reference=True))
    assert parsed["values"]["reference"]["status"] == "conflict"
    assert {x["value"] for x in parsed["values"]["reference"]["candidates"]} == {REFERENCE, "INV-B1-CONFLICT"}
    with pytest.raises(ServiceError) as refused:
        c.caps.invoke("extract.task.confirm", {"task_id": c.task_id, "revision": 1,
                      "note": "Fixture: acknowledgement cannot resolve conflicting evidence", "acknowledge_unverified": True},
                      user_id=c.owner, workspace_id=world.a, via="ui")
    assert (refused.value.status, refused.value.code) == (422, "review_unresolved")
    for _ in range(2):
        c.tick()
    current = c.read("extract.task.get", {"task_id": c.task_id})
    assert current["confirmed_revision"] is None and current["revision"] == 1
    assert (c.run()["status"], c.run()["step_id"]) == ("waiting_input", "extract")
    assert not {"dataset", "prep", "record", "approval", "review"} & set(c.steps())
    for table in ("dataprep_datasets", "dataprep_runs", "tb_records", "approval_requests"):
        assert count(c.store, table) == 0, table
    assert_no_release(c)

    # A real person's explicit correction and confirmation unlock the same run, not a replacement.
    rid = c.rid
    c.caps.invoke("extract.task.correct", {"task_id": c.task_id, "revision": 1, "values": {"reference": REFERENCE},
                  "note": "Fixture: use the first invoice reference after reviewing both source lines"},
                  user_id=c.owner, workspace_id=world.a, via="ui")
    c.caps.invoke("extract.task.confirm", {"task_id": c.task_id, "revision": 2, "note": "Checked the correction",
                  "acknowledge_unverified": True}, user_id=c.owner, workspace_id=world.a, via="ui")
    c.to_review()
    assert c.rid == rid
    assert count(c.store, "form_submissions") == 1 and count(c.store, "extract_tasks") == 1
    assert count(c.store, "tb_records") == 1 and count(c.store, "approval_requests") == 1
    _, dataset = c.op("dataprep.dataset.from_extract")
    assert dataset["source"]["tasks"][0]["revision"] == 3
    assert_no_release(c)


def test_b2_rejected_approval_ends_in_the_fail_branch_without_document_or_message(world):
    w = world
    c = Chain(w)
    c.to_approval()
    c.decide("reject")
    c.tick()
    run = c.run()
    assert (run["status"], run["error_code"], run["step_id"]) == ("failed", "approval_rejected", "rejected")
    steps = c.steps()
    assert decode(steps["route"]["value_json"]) == {"target": "rejected"}
    assert not {"materialize", "project", "job", "artifacts", "conversation", "send"} & set(steps)
    assert c.read("approvals.request.get", {"request_id": c.approval_id}, user="owner")["status"] == "rejected"
    deliver(c.api, c.adapters)
    assert count(c.store, "projects WHERE origin LIKE 'suite:%'") == 0
    assert count(c.store, "jobs") == 0 and count(c.store, "artifacts") == 0
    assert group_messages(c.store, w.cid) == []
    assert count(c.store, "op_runs WHERE parent_run_id=? AND (op LIKE ? OR op LIKE ?)",
                 (c.rid, DOCS_APP_ID + ".%", MESSAGES_APP_ID + ".%")) == 0
    c.tick()
    assert c.run()["status"] == "failed"


def test_b3_rebuilt_factory_continues_after_the_decision_once(world):
    w = world
    c = Chain(w)
    c.to_approval()
    fresh = create_app(str(w.store.root), role="all", start_jobs=False)
    with TestClient(fresh, base_url=ORIGIN):
        c.rebuild(fresh)
        assert c.engine.store is not w.store and c.adapters is not w.adapters
        c.decide("approve")

        def replaying():
            # Fault injection: rewind the engine's own wait-event cursor so the decision event is
            # read again, then tick twice more, after every phase of the continuation.
            c.tick()
            c.store.execute("UPDATE app_event_cursors SET seq=0 WHERE consumer='flows.wait.v1'")
            c.tick()
            c.tick()

        made, job = c.to_end(tick=replaying)
        assert c.run()["status"] == "partial"
        replaying()
        deliver(c.api, c.adapters)
        assert c.run()["status"] == "partial"
        assert_one_document_and_message(c, w, made, job)
        assert count(c.store, "app_events WHERE type='approvals.request.approved'") == 1
        for operation in (f"{DOCS_APP_ID}.project.from_dataset", f"{DOCS_APP_ID}.job.create", f"{MESSAGES_APP_ID}.message.send"):
            c.op(operation)  # exactly one run of each
        assert count(c.store, "flow_history WHERE run_id=? AND type='advanced' AND status='partial'", (c.rid,)) == 1
        assert dispatches(c) == len(FLOW_CALLS)


def test_b4_delivered_message_survives_lost_receipt_update_without_replaying_effects(world, monkeypatch):
    c = Chain(world)
    c.to_approval()
    c.decide("approve")
    made, job = c.to_send()
    before = {sid: (s["status"], s["op_run_id"], s["key"]) for sid, s in c.steps().items() if sid != "send"}
    accepted, _ = c.op(f"{MESSAGES_APP_ID}.message.send")
    assert accepted["status"] == "queued"
    assert group_messages(c.store, world.cid) == []
    lost = []
    update_run = c.caps.update_run

    def lose_receipt(run_id, status, **kw):
        if run_id == accepted["id"] and status == "succeeded":
            lost.append(run_id)
            raise RuntimeError("B1 fault injection: delivery committed but receipt update lost")
        return update_run(run_id, status, **kw)

    # The real worker has committed its message; only the projection into the receipt is lost.
    with monkeypatch.context() as fault:
        fault.setattr(c.caps, "update_run", lose_receipt)
        with pytest.raises(RuntimeError, match="delivery committed but receipt update lost"):
            deliver(c.api, c.adapters)
    assert lost == [accepted["id"]]
    message, = group_messages(c.store, world.cid)
    assert message["client_id"] == accepted["id"]
    assert c.op(f"{MESSAGES_APP_ID}.message.send")[0]["status"] == "queued"
    assert count(c.store, "app_events WHERE type='messages.message.recorded'") == 0
    c.tick()
    assert (c.run()["status"], c.run()["step_id"]) == ("waiting_input", "send")
    assert c.steps()["send"]["status"] == "operation_wait"  # accepted queue is known, not an uncertain dispatch

    fresh = create_app(str(world.store.root), role="all", start_jobs=False)
    with TestClient(fresh, base_url=ORIGIN):
        c.rebuild(fresh)
        c.tick()  # the receipt has not been reconciled yet; never start another send
        assert (c.run()["status"], c.run()["step_id"]) == ("waiting_input", "send")
        inspected = c.read(f"{MESSAGES_APP_ID}.message.inspect", {"run_id": accepted["id"]})
        assert inspected["delivery"] == "recorded" and inspected["message_id"] == message["id"]
        assert c.op(f"{MESSAGES_APP_ID}.message.send")[0]["status"] == "queued"  # inspect is read-only
        c.adapters.sync_messages()  # the production reconciler, no invented success receipt
        c.tick()
        assert c.run()["status"] == "partial"
        settled, receipt = c.op(f"{MESSAGES_APP_ID}.message.send")
        assert settled["id"] == accepted["id"] and settled["status"] == "succeeded"
        assert receipt["message_id"] == message["id"] and receipt["seq"] == message["seq"]
        assert c.steps()["send"]["op_run_id"] == accepted["id"] and c.steps()["send"]["attempt"] == 1
        assert {sid: (s["status"], s["op_run_id"], s["key"]) for sid, s in c.steps().items() if sid != "send"} == before
        deliver(c.api, c.adapters)
        c.tick()
        assert_one_document_and_message(c, world, made, job)
        assert group_messages(c.store, world.cid)[0]["id"] == message["id"]
        assert count(c.store, "messaging_outbox WHERE event_id=?", (accepted["id"],)) == 1
        assert count(c.store, "app_events WHERE type='messages.message.recorded'") == 1
        assert dispatches(c) == len(FLOW_CALLS)


def test_b5_active_approval_wait_rereads_the_request_despite_stale_event_hints(world):
    c = Chain(world)
    c.to_approval()
    pending = c.read("approvals.request.get", {"request_id": c.approval_id}, user="owner")
    assert pending["status"] == "in_review"

    def hint(status):
        # Explicit event-delivery fault fixture, not a decision or a write to approval tables.
        # Retain the pending revision/state_version even when the actual request advances.
        with c.store.tx() as t:
            event = record_event(t, type_="approvals.request." + status, app="approvals", workspace_id=world.a,
                                 object_kind="approval_request", object_id=c.approval_id,
                                 data={"request_id": c.approval_id, "revision": pending["revision"],
                                       "state_version": pending["state_version"], "status": status})
        return c.store.one("SELECT seq FROM app_events WHERE id=?", (event["id"],))["seq"]

    for status in ("approved", "rejected", "approved"):
        checks = c.steps()["review"]["checks"]
        assert c.run()["due_at"] > c.clock()
        seq = hint(status)
        c.engine.tick()  # clock does not advance: only consuming the event can wake this wait
        assert c.caps.cursor("flows.wait.v1") >= seq
        assert c.steps()["review"]["checks"] == checks + 1
        assert (c.run()["status"], c.run()["step_id"]) == ("waiting_approval", "review")
        current = c.read("approvals.request.get", {"request_id": c.approval_id}, user="owner")
        assert current["status"] == "in_review" and current["state_version"] == pending["state_version"]
        assert_no_release(c)

    c.decide("approve")
    hint("rejected")  # an older contradictory hint arrives after the real approval, before continuation
    made, job = c.to_end()
    assert c.run()["status"] == "partial"
    assert decode(c.steps()["route"]["value_json"]) == {"target": "materialize"}
    assert c.read("approvals.request.get", {"request_id": c.approval_id}, user="owner")["status"] == "approved"
    assert_one_document_and_message(c, world, made, job)
    assert dispatches(c) == len(FLOW_CALLS)


def test_b5_concurrent_ticks_during_docs_and_message_steps_make_one_of_each(world):
    w = world
    c = Chain(w)
    c.to_approval()
    c.decide("approve")
    second = Flows(w.store, w.caps, clock=c.clock)
    errors = []

    def both():
        c.clock.advance(31)
        gate = threading.Barrier(2)

        def work(engine):
            try:
                gate.wait()
                for _ in range(3):
                    engine.tick()
            except BaseException as exc:  # noqa: BLE001 - reported below
                errors.append(exc)
        threads = [threading.Thread(target=work, args=(e,)) for e in (c.engine, second)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(60)
        assert not errors, errors

    made, job = c.to_end(tick=both)
    both()
    assert c.run()["status"] == "partial"
    assert_one_document_and_message(c, w, made, job)
    for operation in (f"{DOCS_APP_ID}.project.from_dataset", f"{DOCS_APP_ID}.project.get", f"{DOCS_APP_ID}.job.create",
                      f"{DOCS_APP_ID}.job.artifacts", f"{MESSAGES_APP_ID}.conversation.get", f"{MESSAGES_APP_ID}.message.send"):
        c.op(operation)
    assert count(w.store, "app_events WHERE type IN ('docs.project.materialized','docs.job.submitted','messages.message.recorded')") == 3
    assert dispatches(c) == len(FLOW_CALLS)
