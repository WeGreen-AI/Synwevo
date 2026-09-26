"""Copying a personal Docs project into an organization workspace with the target's approval:
who may ask, who may decide, what the approver sees, what the copy contains, that the original
is untouched, and that a changed source or a withdrawn request never gets copied. Runs the real
composed server (accounts + Docs) with real sessions, the real document engine and a real job in
the target workspace whose output is read back."""
from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient

PW = "harbor lights guide the evening ferry home"
CSRF = "X-Synwevo-CSRF"
WS = "X-Synwevo-Workspace"


def _xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "客户"
    ws.append(["姓名", "编号"]); ws.append(["张三", "0007"]); ws.append(["李四", "0008"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _docx() -> bytes:
    import docx
    d = docx.Document(); d.add_paragraph("尊敬的示例客户名称：您好。")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


@pytest.fixture()
def world(tmp_path, monkeypatch):
    for k, v in {"UI_AUTH": "session", "COOKIE_SECURE": "0", "DATA_DIR": str(tmp_path / "d"), "MAIL_MODE": "file", "MAIL_DIR": str(tmp_path / "mail"),
                 "MAIL_FROM": "synwevo@example.test", "PUBLIC_URL": "http://testserver", "ROLE": "all", "REGISTRATION": "open"}.items():
        monkeypatch.setenv("SYNWEVO_" + k, v)
    from synwevo_server.app import create_app
    api = create_app(str(tmp_path / "d"), role="all", start_jobs=True)
    ident = api.state.identity
    users = {}
    for key, name in (("alice", "Alice Owner"), ("bob", "Bob Boss"), ("carol", "Carol Admin"), ("dave", "Dave Editor"), ("eve", "Eve Outside")):
        u = ident.create_user(f"{key}@example.test", PW, status="active", display_name=name, with_personal_workspace=True)
        ident.store.execute("UPDATE users SET email_verified_at=1 WHERE id=?", (u["id"],))
        users[key] = u
    from synwevo_core.organizations import Organizations
    created = Organizations(ident.store, ident).create(users["bob"]["id"], "Copy Org", "Copy Space", "fixture-idem-copy-org-0001")
    oid, wid = created["organization"]["id"], created["workspace"]["id"]
    for key, org_role, ws_role in (("alice", "member", "editor"), ("carol", "admin", "viewer"), ("dave", "member", "editor")):
        ident.store.execute("INSERT INTO organization_memberships (organization_id,user_id,role,created_at,updated_at,created_by) VALUES (?,?,?,1,1,?)",
                            (oid, users[key]["id"], org_role, users["bob"]["id"]))
        ident.store.execute("INSERT INTO memberships (workspace_id,user_id,role,can_export,created_at,updated_at,created_by) VALUES (?,?,?,1,1,1,?)",
                            (wid, users[key]["id"], ws_role, users["bob"]["id"]))
    personal = {k: ident.ensure_personal_workspace(u["id"])["id"] for k, u in users.items()}
    clients = {}
    for key in users:
        c = TestClient(api)
        r = c.post("/api/auth/login", json={"email": f"{key}@example.test", "password": PW})
        assert r.status_code == 200, r.text
        c.headers.update({CSRF: c.cookies.get("dw_csrf")})
        clients[key] = c
    try:
        yield {"api": api, "ident": ident, "users": users, "oid": oid, "wid": wid, "personal": personal, "c": clients}
    finally:
        for c in clients.values():
            c.close()
        if api.state.jobs:
            api.state.jobs.stop()


def acting(c: TestClient, wid: str) -> TestClient:
    c.headers.update({WS: wid})
    return c


def wait_parsed(c, pid, timeout=60):
    t0 = time.time()
    while time.time() - t0 < timeout:
        srcs = c.get(f"/api/projects/{pid}/sources").json()
        if all(s["status"] not in ("parsing", "scanning") for s in srcs):
            return srcs
        time.sleep(0.1)
    raise AssertionError("parse timeout")


def build_project(c, name="个人合同项目"):
    pid = c.post("/api/projects", json={"name": name}).json()["id"]
    r = c.post(f"/api/projects/{pid}/sources", files=[("files", ("客户.xlsx", _xlsx()))]); assert r.status_code < 400, r.text
    wait_parsed(c, pid)
    t = c.post("/api/templates/upload", files={"file": ("合同模板.docx", _docx())}, data={"project_id": pid}).json()
    view = c.get(f"/api/templates/{t['id']}/view").json()
    block = next(b for part in view["doc"]["parts"] for b in _walk(part["blocks"]) if "示例客户名称" in b["text"])
    s = block["text"].index("示例客户名称")
    slot = c.post(f"/api/templates/{t['id']}/slots", json={"kind": "text", "part": block["part"], "path": block["path"], "start": s, "end": s + 6, "name": "客户名称"}).json()["created_slot"]
    assert c.patch(f"/api/projects/{pid}", json={"record_rule": {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}}}).status_code == 200
    r = c.put(f"/api/projects/{pid}/mappings", json={"mappings": [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": slot["binding"]}}]})
    assert r.status_code == 200, r.text
    return pid, t["id"]


def _walk(blocks):
    for b in blocks:
        if b["t"] == "p":
            yield b
        else:
            for row in b["rows"]:
                for cell in row["cells"]:
                    yield from _walk(cell["blocks"])


def run_job(c, pid, timeout=120):
    j = c.post(f"/api/projects/{pid}/jobs", json={"formats": ["docx"]}).json()
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/jobs/{j['id']}").json()
        if j["status"] not in ("queued", "running", "cancelling"):
            return j
        time.sleep(0.2)
    raise AssertionError("job timeout")


def docx_text(data: bytes) -> str:
    import docx
    return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)


def settle(w, c, pid, timeout=30):
    """Run the copy runner once and wait for the request to leave approved/running. The runner
    thread is live in this fixture (a self-approval wakes it), so the tick here may find nothing."""
    w["api"].state.copy_runner.run_once()
    t0 = time.time()
    while time.time() - t0 < timeout:
        r = c.get(f"/api/projects/{pid}/copy-requests").json()["items"][0]
        if r["status"] not in ("approved", "running"):
            return r
        time.sleep(0.1)
    raise AssertionError(f"copy did not settle: {r}")


def reauth(c):
    assert c.post("/api/auth/reauth", json={"password": PW, "scope": "admin"}).status_code == 200


def test_copy_needs_the_target_to_approve_and_copies_exactly_the_approved_manifest(world):
    w = world; oid, wid = w["oid"], w["wid"]
    alice = acting(w["c"]["alice"], w["personal"]["alice"])
    pid, tid = build_project(alice)
    before = alice.get(f"/api/projects/{pid}").json()
    # who can be a target: only organization workspaces alice can write to; she cannot decide herself
    targets = alice.get("/api/copy-targets").json()
    assert targets["source_is_personal"] is True and [x["id"] for x in targets["items"]] == [wid] and targets["items"][0]["can_self_approve"] is False
    manifest = alice.get(f"/api/projects/{pid}/copy-manifest?sources=1").json()["manifest"]
    assert manifest["template"]["slots"] == 1 and manifest["mappings"]["count"] == 1 and manifest["sources"]["count"] == 1 and manifest["sources"]["items"][0]["sha256"]
    # a viewer of the target cannot ask; an outsider sees no targets
    eve = acting(w["c"]["eve"], w["personal"]["eve"])
    assert eve.get("/api/copy-targets").json()["items"] == []
    r = alice.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": wid, "name": "合同项目（组织副本）", "purpose": "给团队复用", "scope": {"sources": True}})
    assert r.status_code == 201, r.text
    req = r.json()
    assert req["status"] == "pending" and req["manifest"]["sources"]["included"] is True and "source_workspace_id" not in req
    assert alice.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": wid, "name": "again", "scope": {}}).json()["error"] == "copy_already_requested"
    # the deciders were told; the approval routes: dave (no permission) 403, eve (not a member) 404,
    # carol (org admin, holds the permission, but only a viewer of the target) may look but not decide
    n = w["api"].state.notifier
    assert [x["kind"] for x in n.list(w["users"]["bob"]["id"])["items"]] == ["copy.requested"]
    assert n.list(w["users"]["carol"]["id"])["items"] == [], "carol cannot decide, so she is not asked"
    A = f"/api/auth/organizations/{oid}/approvals"
    assert w["c"]["dave"].get(A).status_code == 403
    assert w["c"]["eve"].get(A).status_code == 404
    carol = w["c"]["carol"]; reauth(carol)
    listed = carol.get(A).json()["items"]
    assert len(listed) == 1 and listed[0]["can_decide"] is False and listed[0]["requester"]["name"] == "Alice Owner"
    assert carol.post(f"{A}/{req['id']}/decide", json={"action": "approve", "reason": "looks fine"}).status_code == 403
    # the approver sees the manifest, not the personal workspace
    bob = w["c"]["bob"]
    detail = bob.get(f"{A}/{req['id']}").json()
    assert detail["manifest"]["sources"]["items"][0]["filename"] == "客户.xlsx" and detail["can_decide"] is True
    assert acting(bob, wid).get(f"/api/projects/{pid}").status_code == 404
    # deciding is a sensitive write with a reason
    assert bob.post(f"{A}/{req['id']}/decide", json={"action": "approve", "reason": "team needs it"}).json()["error"] == "reauth_required"
    reauth(bob)
    assert bob.post(f"{A}/{req['id']}/decide", json={"action": "approve", "reason": "no"}).json()["error"] == "reason_required"
    out = bob.post(f"{A}/{req['id']}/decide", json={"action": "approve", "reason": "team needs it", "version": req["version"]}).json()
    assert out["status"] == "approved" and out["self_approved"] is False and out["decider"]["name"] == "Bob Boss"
    assert "copy.approved" in [x["kind"] for x in n.list(w["users"]["alice"]["id"])["items"]]
    assert bob.post(f"{A}/{req['id']}/decide", json={"action": "reject", "reason": "changed my mind"}).json()["error"] == "not_pending"
    # the Docs runner executes it: one new project in the target, original untouched
    done = settle(w, alice, pid)
    assert done["status"] == "done" and done["result"]["sources"] == 1 and done["result"]["mappings"] == 1, done
    new_pid = done["result"]["project_id"]
    after = alice.get(f"/api/projects/{pid}").json()
    assert after["rev"] == before["rev"] and after["mapping_rev"] == before["mapping_rev"] and after["template_id"] == tid
    assert alice.get(f"/api/projects/{new_pid}").status_code == 404, "the copy is not in the personal workspace"
    copy = bob.get(f"/api/projects/{new_pid}").json()
    assert copy["name"] == "合同项目（组织副本）" and copy["template_id"] != tid and copy["origin"].startswith("copy:") and copy["mapping_rev"] == 1
    # the copy says where it came from — names and the approval, never the source workspace
    cf = copy["copied_from"]
    assert cf["request_id"] == req["id"] and cf["source_project_name"] == "个人合同项目" and cf["requester"]["name"] == "Alice Owner"
    assert cf["decider"]["name"] == "Bob Boss" and cf["self_approved"] is False and cf["organization_id"] == oid and cf["sources_included"] is True
    assert "source_workspace_id" not in cf and alice.get(f"/api/projects/{pid}").json()["copied_from"] is None
    srcs = bob.get(f"/api/projects/{new_pid}/sources").json()
    assert [s["status"] for s in srcs] == ["ready"] and srcs[0]["id"] != manifest["sources"]["items"][0]["id"]
    blob = w["ident"].store.one("SELECT blob FROM source_files WHERE id=?", (srcs[0]["id"],))["blob"]
    assert blob == manifest["sources"]["items"][0]["sha256"], "same content, shared blob, new identity"
    tpl = bob.get(f"/api/templates/{copy['template_id']}").json()
    assert tpl["head"] == 1 and len(bob.get(f"/api/templates/{copy['template_id']}/view").json()["slots"]) == 1
    # the copy generates from its own rows: a real job in the target workspace, output read back
    job = run_job(bob, new_pid)
    assert job["status"] == "completed", job
    items = bob.get(f"/api/jobs/{job['id']}/items", params={"limit": 100}).json()
    texts = [docx_text(bob.get(f"/api/artifacts/{a['id']}").content) for it in items for a in it["artifacts"]]
    assert any("张三" in t for t in texts) and any("李四" in t for t in texts), texts
    assert "copy.done" in [x["kind"] for x in n.list(w["users"]["alice"]["id"])["items"]]
    # nothing of the workspace's credentials or the original's history moved
    assert w["ident"].store.one("SELECT COUNT(*) AS n FROM jobs WHERE project_id=?", (pid,))["n"] == 0
    assert w["ident"].store.one("SELECT COUNT(*) AS n FROM webhook_subscriptions WHERE tenant=?", (wid,))["n"] == 0


def test_a_changed_source_withdrawal_and_self_approval(world):
    w = world; oid, wid = w["oid"], w["wid"]
    alice = acting(w["c"]["alice"], w["personal"]["alice"])
    pid, _ = build_project(alice, "会变的项目")
    A = f"/api/auth/organizations/{oid}/approvals"
    bob = w["c"]["bob"]; reauth(bob)
    # approved, then the source changes before the copy runs: stale, never copied
    req = alice.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": wid, "name": "副本", "scope": {}}).json()
    assert bob.post(f"{A}/{req['id']}/decide", json={"action": "approve", "reason": "ok for now"}).json()["status"] == "approved"
    # a rename does not change what would be copied; the record rule does
    assert alice.patch(f"/api/projects/{pid}", json={"name": "会变的项目（改名）"}).status_code == 200
    assert alice.get(f"/api/projects/{pid}/copy-manifest").json()["digest"] == req["digest"]
    assert alice.patch(f"/api/projects/{pid}", json={"record_rule": {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0, "data_end": 1}}}).status_code == 200
    assert alice.get(f"/api/projects/{pid}/copy-manifest").json()["digest"] != req["digest"]
    stale = settle(w, alice, pid)
    assert stale["status"] == "stale" and stale["error"] == "source_changed" and stale["result"] is None
    assert w["ident"].store.one("SELECT COUNT(*) AS n FROM projects WHERE origin LIKE 'copy:%' AND tenant=?", (wid,))["n"] == 0
    assert "copy.failed" in [x["kind"] for x in w["api"].state.notifier.list(w["users"]["alice"]["id"])["items"]]
    # withdrawn before a decision: the decider cannot decide it any more
    req2 = alice.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": wid, "name": "副本 2", "scope": {}}).json()
    assert alice.post(f"/api/copy-requests/{req2['id']}/withdraw", json={}).json()["status"] == "withdrawn"
    assert bob.post(f"{A}/{req2['id']}/decide", json={"action": "approve", "reason": "too late"}).json()["error"] == "not_pending"
    # rejection reaches the requester with the reason
    req3 = alice.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": wid, "name": "副本 3", "scope": {}}).json()
    out = bob.post(f"{A}/{req3['id']}/decide", json={"action": "reject", "reason": "not this quarter"}).json()
    assert out["status"] == "rejected" and out["decision_reason"] == "not this quarter"
    assert any(x["kind"] == "copy.rejected" and "not this quarter" in x["body"] for x in w["api"].state.notifier.list(w["users"]["alice"]["id"])["items"])
    # a requester who also holds the deciding authority approves in the same step, recorded as such
    bobp = acting(w["c"]["bob"], w["personal"]["bob"])
    bpid, _ = build_project(bobp, "老板的个人项目")
    assert bobp.get("/api/copy-targets").json()["items"][0]["can_self_approve"] is True
    r = bobp.post(f"/api/projects/{bpid}/copy-requests", json={"target_workspace_id": wid, "name": "老板副本", "scope": {"sources": True, "corrections": True}, "decide": "approve", "reason": "my own workspace"})
    assert r.status_code == 201 and r.json()["status"] == "approved" and r.json()["self_approved"] is True, r.text
    mine = settle(w, bobp, bpid)
    assert mine["status"] == "done" and mine["result"]["sources"] == 1
    # an organization workspace project cannot be the source of a request
    r = acting(w["c"]["bob"], wid).post(f"/api/projects/{mine['result']['project_id']}/copy-requests", json={"target_workspace_id": wid, "name": "x", "scope": {}})
    assert r.status_code == 409 and r.json()["error"] == "source_not_personal"


def test_a_copy_carries_neither_the_source_s_live_room_nor_its_history(world):
    """9.13 D / 9.16 「个人复制不携带私人协同房间/未批准历史」: the copy is the approved snapshot
    (the template's head version), never the live draft; the source project's activity feed, the
    template's live draft and its change log stay with the source."""
    w = world; oid, wid = w["oid"], w["wid"]
    alice = acting(w["c"]["alice"], w["personal"]["alice"])
    pid, tid = build_project(alice)
    store = w["ident"].store
    assert store.one("SELECT COUNT(*) AS n FROM project_activity WHERE project_id=?", (pid,))["n"] >= 1, "the source has history"
    # a live draft nobody saved as a version, as the shutdown hook or the last person out would leave it
    store.execute("INSERT INTO collab_docs (room, tenant, template_id, state, head, dirty, updated_at) VALUES (?,?,?,?,?,?,?)",
                  (f"template:{tid}", w["personal"]["alice"], tid, b"\x00\x00", 1, 1, 1.0))
    r = alice.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": wid, "name": "只带已批准的快照", "purpose": "复用配置", "scope": {}})
    assert r.status_code == 201, r.text
    req = r.json()
    bob = w["c"]["bob"]; reauth(bob)
    out = bob.post(f"/api/auth/organizations/{oid}/approvals/{req['id']}/decide", json={"action": "approve", "reason": "team needs it", "version": req["version"]}).json()
    assert out["status"] == "approved", out
    done = settle(w, alice, pid)
    assert done["status"] == "done", done
    new_pid, new_tid = done["result"]["project_id"], done["result"]["template_id"]
    assert new_tid and new_tid != tid
    assert store.one("SELECT COUNT(*) AS n FROM project_activity WHERE project_id=?", (new_pid,))["n"] == 0, "no history travels with the copy"
    assert store.one("SELECT 1 AS x FROM collab_docs WHERE room=?", (f"template:{new_tid}",)) is None, "no live draft travels with the copy"
    assert store.one("SELECT COUNT(*) AS n FROM collab_changes WHERE room=?", (f"template:{new_tid}",))["n"] == 0
    t = acting(bob, wid).get(f"/api/templates/{new_tid}").json()
    assert t["head"] == 1 and len(t["versions"]) == 1, "the copy starts from the approved head version alone"
    assert store.one("SELECT dirty FROM collab_docs WHERE room=?", (f"template:{tid}",))["dirty"] == 1, "the source keeps its draft"
