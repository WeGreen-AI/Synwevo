"""Member import through the document engine: the composition root injects Synwevo Docs' readers
into the account service, so XLSX, DOCX, PDF and scanned images become rows there — while the
account service itself never imports the Docs package. Runs the real engine (openpyxl, lxml,
PyMuPDF/pdfium, Pillow, Tesseract when installed)."""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

PW = "harbor lights guide the evening ferry home"
CSRF = "X-Synwevo-CSRF"


@pytest.fixture()
def platform(tmp_path, monkeypatch):
    for k, v in {"UI_AUTH": "session", "COOKIE_SECURE": "0", "DATA_DIR": str(tmp_path / "d"), "MAIL_MODE": "file",
                 "MAIL_DIR": str(tmp_path / "mail"), "MAIL_FROM": "synwevo@example.test", "PUBLIC_URL": "http://testserver",
                 "ROLE": "all", "REGISTRATION": "open"}.items():
        monkeypatch.setenv("SYNWEVO_" + k, v)
    from synwevo_server.app import create_app
    api = create_app(str(tmp_path / "d"), role="all", start_jobs=False)
    ident = api.state.identity
    owner = ident.create_user("owner@example.test", PW, status="active", display_name="Ola Owner")
    ident.store.execute("UPDATE users SET email_verified_at=1 WHERE id=?", (owner["id"],))
    from synwevo_core.organizations import Organizations
    created = Organizations(ident.store, ident).create(owner["id"], "Engine Org", "Engine Team", "fixture-idem-engine-org-0001")
    with TestClient(api) as c:
        c.api, c.ident, c.oid, c.wid = api, ident, created["organization"]["id"], created["workspace"]["id"]
        r = c.post("/api/auth/login", json={"email": "owner@example.test", "password": PW})
        assert r.status_code == 200, r.text
        c.h = {CSRF: c.cookies.get("dw_csrf")}
        assert c.post("/api/auth/reauth", json={"password": PW, "scope": "admin"}, headers=c.h).status_code == 200
        yield c


def _upload(c, name, data, ctype="application/octet-stream", ocr=False):
    r = c.post(f"/api/auth/organizations/{c.oid}/imports", files={"file": (name, data, ctype)}, data={"ocr": "1" if ocr else ""}, headers=c.h)
    assert r.status_code == 201, r.text
    return r.json()


def _plan(c, imp, mapping):
    r = c.put(f"/api/auth/organizations/{c.oid}/imports/{imp['id']}/mapping", headers=c.h, json={"workspace_id": c.wid, "mapping": mapping})
    assert r.status_code == 200, r.text
    return r.json()


def test_the_engine_is_injected_and_reads_a_workbook(platform):
    c = platform
    cap = c.get(f"/api/auth/organizations/{c.oid}/imports/capabilities").json()
    assert cap["ingest"]["engine"] == "docs" and {"xlsx", "docx", "pdf", "png"} <= set(cap["ingest"]["formats"])
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Members"
    ws.append(["Name", "E-mail", "Workspace role", "Custom role"])
    ws.append(["Nina New", "nina.new@example.test", "editor", ""])
    ws.append(["Leading Zero", "007@example.test", "viewer", ""])            # numeric-looking local part stays text
    ws.append([None, None, None, None])
    hidden = wb.create_sheet("Old"); hidden.sheet_state = "hidden"; hidden.append(["ignored@example.test"])
    buf = io.BytesIO(); wb.save(buf)
    imp = _upload(c, "members.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ex = imp["extraction"]
    assert ex["engine"] == "docs" and ex["status"] == "ready" and [t["name"] for t in ex["tables"]] == ["Members", "Old"]
    assert ex["tables"][1]["source"].get("hidden") is True
    assert imp["mapping"] == {"table": 0, "header_row": True, "columns": {"email": 1, "name": 0, "workspace_role": 2, "org_roles": 3}, "mode": "table"}
    plan = _plan(c, imp, {**imp["mapping"], "defaults": {"workspace_role": "viewer"}})
    assert plan["summary"]["counts"]["invite"] == 2 and plan["summary"]["counts"]["blocked"] == 0, plan["summary"]
    rows = c.get(f"/api/auth/organizations/{c.oid}/imports/{imp['id']}/rows").json()["items"]
    assert rows[1]["fields"]["email"] == "007@example.test" and rows[1]["source"] == {"sheet": "Members", "table": "Members", "row": 3}
    assert rows[0]["fields"]["workspace_role"] == "editor"


def test_docx_tables_pdf_lines_and_json_records(platform):
    c = platform
    import docx
    d = docx.Document()
    d.add_paragraph("Project members")
    tb = d.add_table(rows=3, cols=2)
    for i, (n, e) in enumerate((("姓名", "邮箱"), ("陈一", "chen.yi@example.test"), ("Dana", "dana@example.test"))):
        tb.cell(i, 0).text, tb.cell(i, 1).text = n, e
    d.add_paragraph("External: ext@example.test (Ext)")
    buf = io.BytesIO(); d.save(buf)
    imp = _upload(c, "members.docx", buf.getvalue())
    ex = imp["extraction"]
    assert ex["status"] == "ready" and ex["tables"][0]["rows"] == 3 and ex["lines"] == 2
    assert imp["mapping"]["columns"] == {"email": 1, "name": 0} and imp["mapping"]["header_row"] is True
    plan = _plan(c, imp, imp["mapping"])
    assert plan["summary"]["counts"]["invite"] == 2
    lines = _plan(c, imp, {"mode": "lines"})
    assert lines["summary"]["counts"]["invite"] == 1          # the paragraph outside the table
    # PDF with a text layer: read line by line
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Team roster", fontsize=14)
    page.insert_text((72, 130), "Eve Example - eve@example.test", fontsize=12)
    page.insert_text((72, 150), "Frank F frank.f@example.test", fontsize=12)
    imp = _upload(c, "roster.pdf", doc.tobytes(), "application/pdf")
    ex = imp["extraction"]
    assert ex["status"] == "ready" and ex["pages"] == 1 and ex["lines"] == 3 and ex["tables"] == []
    plan = _plan(c, imp, {"mode": "lines", "defaults": {"workspace_role": "editor"}})
    rows = c.get(f"/api/auth/organizations/{c.oid}/imports/{imp['id']}/rows").json()["items"]
    assert [(r["fields"]["email"], r["fields"]["name"], r["source"]["page"]) for r in rows] == [("eve@example.test", "Eve Example -", 1), ("frank.f@example.test", "Frank F", 1)] or \
           [r["fields"]["email"] for r in rows] == ["eve@example.test", "frank.f@example.test"]
    # JSON records through the engine path take the shared JSON reader
    imp = _upload(c, "members.json", b'[{"email": "gia@example.test", "name": "Gia"}, {"email": "hal@example.test", "name": "Hal", "role": "editor"}]', "application/json")
    assert imp["extraction"]["tables"][0]["rows"] == 3 and imp["mapping"]["columns"]["email"] == 0


def test_a_scanned_image_needs_ocr_and_ocr_is_on_request(platform):
    c = platform
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new("RGB", (1400, 300), "white")
    draw = ImageDraw.Draw(im)
    try:
        font = ImageFont.load_default(size=44)
    except TypeError:
        font = ImageFont.load_default()
    draw.text((40, 60), "Ivy Import  ivy.import@example.test", fill="black", font=font)
    draw.text((40, 160), "Jon Jones  jon.jones@example.test", fill="black", font=font)
    buf = io.BytesIO(); im.save(buf, "PNG")
    imp = _upload(c, "roster.png", buf.getvalue(), "image/png")
    assert imp["extraction"]["status"] == "needs_ocr" and imp["extraction"]["lines"] == 0
    cap = c.get(f"/api/auth/organizations/{c.oid}/imports/capabilities").json()
    if not cap["ingest"]["ocr"]:
        pytest.skip("Tesseract is not installed here; OCR is reported as unavailable, not faked")
    r = c.post(f"/api/auth/organizations/{c.oid}/imports/{imp['id']}/ocr", headers=c.h)
    assert r.status_code == 200, r.text
    ex = r.json()["extraction"]
    assert ex["ocr"]["applied"] is True and ex["lines"] >= 2, ex
    plan = _plan(c, r.json(), {"mode": "lines"})
    emails = sorted(x["fields"]["email"] for x in c.get(f"/api/auth/organizations/{c.oid}/imports/{imp['id']}/rows").json()["items"])
    assert "ivy.import@example.test" in emails and "jon.jones@example.test" in emails, emails
    rows = c.get(f"/api/auth/organizations/{c.oid}/imports/{imp['id']}/rows").json()["items"]
    assert all(x["source"].get("ocr") is True and "conf" in x["source"] for x in rows)
