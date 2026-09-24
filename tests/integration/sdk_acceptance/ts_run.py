"""Run synthetic Node/n8n acceptance on a fresh _testenv instance. No existing server accepted."""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

sys.dont_write_bytecode = True

ROOT = Path(__file__).resolve().parents[3]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def run(argv, evidence, name, env=None, timeout=240):
    result = subprocess.run([str(x) for x in argv], env=env, capture_output=True, text=True, timeout=timeout)
    (evidence / f'{name}.stdout').write_text(result.stdout)
    (evidence / f'{name}.stderr').write_text(result.stderr)
    with (evidence / 'commands.jsonl').open('a') as f:
        f.write(json.dumps({'argv': [str(x) for x in argv], 'exit_code': result.returncode}) + '\n')
    if result.returncode:
        raise RuntimeError(f'{name} failed ({result.returncode}); see {evidence}/{name}.stderr')
    return result


def verify(result, expected, kit, out):
    by_name = {x['file_base'] + '.docx': x for x in expected['items']}
    assert len(result['artifacts']) == expected['samples']
    assert len({x['filename'] for x in result['artifacts']}) == expected['samples']
    records = []
    preview = result.get('preview', {}).get('download')
    if preview:
        preview_bytes = Path(preview['path']).read_bytes()
        assert hashlib.sha256(preview_bytes).hexdigest() == preview['sha256'] == preview['header_sha256']
        assert preview['project_signature'] and preview['template_version'] and preview['record_key']
        assert not kit.check_docx(preview_bytes, expected['items'][0], expected['forbidden_text'])
    for artifact in result['artifacts']:
        path = Path(artifact['path'])
        assert path.resolve().is_relative_to(out.resolve())
        data = path.read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        assert digest == artifact['sha256'] == artifact['manifest_sha256'] == artifact['header_sha256']
        problems = kit.check_docx(data, by_name[artifact['filename']], expected['forbidden_text'])
        assert not problems, (artifact['filename'], problems)
        records.append({'file': artifact['filename'], 'sha256': digest, 'content': 'passed'})
    if (out / 'archive.zip').exists():
        with zipfile.ZipFile(out / 'archive.zip') as archive:
            for artifact in result['artifacts']:
                assert hashlib.sha256(archive.read(artifact['filename'])).hexdigest() == artifact['sha256']
        with (out / 'report.csv').open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == expected['samples']
        assert {row['输出文件'] for row in rows} == set(by_name)
        assert all(row['状态'] == '成功' for row in rows)
        (out / 'archive-verification.json').write_text(json.dumps({
            'archive_entries_match_artifact_sha256': len(result['artifacts']),
            'report_success_rows': len(rows)}) + '\n')
    return records


def verify_builtin(result, expected, kit, out):
    """Read native contents, including the saved model's revision marker, after real HTTP rendering."""
    import docx
    import openpyxl

    fmt = 'docx' if result['kind'] == 'builtin_doc' else 'xlsx'
    by_name = {item['file_base'] + '.' + fmt: item for item in expected['items']}
    artifacts = result['artifacts']
    assert len(artifacts) == expected['samples']
    assert {artifact['filename'] for artifact in artifacts} == set(by_name)
    preview = result['preview']['download']
    assert preview['project_signature'] == result['preflight']['signature']
    assert preview['template_version'] == '2' and preview['record_key']
    checks = []
    outputs = [(preview, expected['items'][0]),
               *((artifact, by_name[artifact['filename']]) for artifact in artifacts)]
    for download, item in outputs:
        path = Path(download['path'])
        assert path.resolve().is_relative_to(out.resolve())
        data = path.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        assert sha == download['sha256'] == download['header_sha256']
        if 'manifest_sha256' in download:
            assert sha == download['manifest_sha256']
        if fmt == 'docx':
            assert not kit.check_docx(data, item, expected['forbidden_text'])
            values = [paragraph.text for paragraph in docx.Document(io.BytesIO(data)).paragraphs]
        else:
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            try:
                sheet = workbook['结果']
                # Expected cells come from the unchanged shared source workbook's column order.
                wanted = dict(zip(kit.HEADERS, item['must_contain']))
                for row, name in enumerate(result['field_names'], start=2):
                    assert sheet[f'A{row}'].value == name
                    assert str(sheet[f'B{row}'].value) == wanted[name]
                values = [str(cell.value) for row in sheet for cell in row if cell.value is not None]
                assert all(text not in '\n'.join(values) for text in expected['forbidden_text'])
            finally:
                workbook.close()
        assert result['marker'] in values and result['initial_marker'] not in '\n'.join(values)
        checks.append({'file': path.name, 'sha256': sha, 'content': 'passed', 'kind': result['kind']})
    return {'preview': checks[0], 'artifacts': checks[1:], 'operations': result['checks']}



def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--evidence', type=Path, required=True)
    parser.add_argument('--node', default='node')
    parser.add_argument('--n8n', type=Path)
    parser.add_argument('--builtin-only', action='store_true', help='Run only the real service-auth builtin-template smoke')
    args = parser.parse_args()
    if args.builtin_only and args.n8n:
        parser.error('--builtin-only cannot be combined with --n8n')
    evidence = args.evidence.resolve()
    evidence.mkdir(parents=True, exist_ok=False)
    kit = load('synthetic_shared_kit', ROOT / 'apps/docs/sdk/acceptance/make_kit.py')
    testenv = load('synthetic_testenv', ROOT / 'apps/docs/scripts/_testenv.py')
    identity_fixture = load('synthetic_account_fixture', Path(__file__).with_name('ts_fixture.py'))
    wait = testenv._wait

    def wait_with_evidence(base, proc, log):
        try:
            wait(base, proc, log)
        except BaseException:
            (evidence / 'server-startup.log').write_bytes(log.read_bytes())
            probe = subprocess.run(['curl', '--noproxy', '*', '-v', '--max-time', '5', base + '/healthz'],
                                   capture_output=True, text=True)
            (evidence / 'health-probe.log').write_text(probe.stdout + probe.stderr)
            raise
    testenv._wait = wait_with_evidence
    expected = kit.make(evidence / 'kit')
    account_base = f'http://127.0.0.1:{testenv._free_port()}'
    original_child_env = testenv._child_env
    def child_env(tmp, port):
        env = original_child_env(tmp, port)
        env.update(SYNWEVO_ISSUER=account_base, SYNWEVO_COOKIE_SECURE='0',
                   SYNWEVO_DOCS_WEBHOOK_ALLOW_HTTP='1', SYNWEVO_DOCS_WEBHOOK_ALLOW_PRIVATE='127.0.0.1',
                   NO_PROXY='127.0.0.1,localhost,::1', no_proxy='127.0.0.1,localhost,::1',
                   NODE_USE_ENV_PROXY='0', PYTHONDONTWRITEBYTECODE='1')
        return env
    testenv._child_env = child_env
    with testenv.isolated() as instance, identity_fixture.account_fixture(testenv, instance, evidence, account_base) as oauth:
        assert testenv.health(instance.base)['instance'] == 'test'
        child = testenv._child_env(instance.data_dir.parent, int(instance.base.rsplit(':', 1)[1]))
        child['PYTHONDONTWRITEBYTECODE'] = '1'
        child.update(NO_PROXY='127.0.0.1,localhost,::1', no_proxy='127.0.0.1,localhost,::1', NODE_USE_ENV_PROXY='0')
        secret_file = evidence / '.client-secret'
        secret_file.touch(mode=0o600)
        secret_file.write_text(oauth['clientSecret'] + '\n')
        child.update(SYNWEVO_BASE_URL=instance.base, SYNWEVO_CLIENT_SECRET_FILE=str(secret_file),
                     SYNWEVO_CLIENT_ID=oauth['clientId'], SYNWEVO_TOKEN_URL=oauth['tokenUrl'],
                     SYNWEVO_SCOPE=oauth['scope'], SYNWEVO_RESOURCE=oauth['resource'])
        (evidence / 'authorization.json').write_text(json.dumps({
            'kind': 'service_account', 'grant_type': 'client_credentials', 'scope': oauth['scope'],
            'workspace': oauth['workspace'], 'token_endpoint': oauth['tokenUrl'],
            'credential_setup': 'controlled synthetic identity fixture; no business records seeded'}) + '\n')
        try:
            builtin_out = evidence / 'builtin'
            run([args.node, ROOT / 'apps/docs/sdk/acceptance/ts_builtin.mjs', evidence / 'kit', builtin_out],
                evidence, 'ts-builtin', child)
            builtin = json.loads((builtin_out / 'result.json').read_text())
            builtin_checks = {result['kind']: verify_builtin(result, expected, kit, builtin_out) for result in builtin}
            (evidence / 'builtin-verification.json').write_text(json.dumps(builtin_checks, ensure_ascii=False, indent=2) + '\n')
            if args.builtin_only:
                print(json.dumps({'status': 'passed', 'selection': 'builtin-only',
                                  'documents': {k: len(v['artifacts']) for k, v in builtin_checks.items()}}))
                return
            out = evidence / 'typescript'
            run([args.node, ROOT / 'apps/docs/sdk/acceptance/ts_chain.mjs', evidence / 'kit', out], evidence, 'ts-chain', child)
            result = json.loads((out / 'result.json').read_text())
            checks = {'typescript': verify(result, expected, kit, out)}
            cli_out = evidence / 'cli'
            cli = run([args.node, ROOT / 'apps/docs/sdk/typescript/synwevo-docs.mjs', 'generate',
                       '--source', evidence / 'kit/customers.xlsx', '--template', evidence / 'kit/template.docx',
                       '--fields', evidence / 'kit/fields.json', '--definition', evidence / 'kit/project.json',
                       '--out', cli_out, '--key-prefix', 'synthetic-cli'], evidence, 'ts-cli', child)
            checks['cli'] = verify(json.loads(cli.stdout), expected, kit, cli_out)
            if args.n8n:
                n8n = load('synthetic_n8n', Path(__file__).with_name('n8n_run.py'))
                result = n8n.execute(args.node, args.n8n.resolve(), evidence, instance.base, oauth)
                checks['n8n'] = verify(result, expected, kit, evidence / 'n8n-output')
                checks['receiver'] = verify(json.loads((evidence / 'receiver-result.json').read_text()), expected, kit, evidence / 'receiver-output')
            (evidence / 'verification.json').write_text(json.dumps(checks, ensure_ascii=False, indent=2) + '\n')
            print(json.dumps({'status': 'passed', 'documents': {k: len(v) for k, v in checks.items()},
                              'builtin': {k: len(v['artifacts']) for k, v in builtin_checks.items()}}))
        finally:
            secret_file.unlink(missing_ok=True)


if __name__ == '__main__':
    main()
