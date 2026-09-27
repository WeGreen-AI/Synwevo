"""Synwevo Docs ↔ Synwevo Message, end to end on the composed server (`all`, CI-75): a real Docs
project, a real comment through the Docs API, a real generation run through the runner, a real copy
request decided through the core service the account service calls — each writing the outbox in
its own transaction — and the bridge worker (live in this process, as it is in the messages
process) turning them into cards in the project's conversation, rendered per reader.

Also the negative half: an event before any conversation exists is skipped and never creates
one; cards create no inbox items; a member who left the workspace sees a neutral card and is then
removed from the conversation."""
from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient

from synwevo_core.apps import DOCS_APP_ID
from synwevo_core.messaging_bridge import NEUTRAL_TITLE, SKIPPED_NO_CONVERSATION

PW = "harbor lights guide the evening ferry home"
CSRF, WS = "X-Synwevo-CSRF", "X-Synwevo-Workspace"
M = "/api/messaging"
ORIGIN = "http://testserver"


def _xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "客户"
    ws.append(["姓名", "编号"]); ws.append(["张三", "0007"]); ws.append(["李四", "0008"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _docx() -> bytes:
    import docx
    d = docx.Document(); d.add_paragraph("尊敬的示例客户名称：您好。")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def _walk(blocks):
    for b in blocks:
        if b["t"] == "p":
            yield b
        else:
            for row in b["rows"]:
                for cell in row["cells"]:
                    yield from _walk(cell["blocks"])


def _wait(fn, timeout=60, what="condition"):
    t0 = time.time()
    while time.time() - t0 < timeout:
        v = fn()
        if v:
            return v
        time.sleep(0.15)
    raise AssertionError(f"timeout waiting for {what}")


@pytest.fixture()
def world(tmp_path, monkeypatch):
    for k, v in {"UI_AUTH": "session", "COOKIE_SECURE": "0", "DATA_DIR": str(tmp_path / "d"), "MAIL_MODE": "file", "MAIL_DIR": str(tmp_path / "mail"),
                 "MAIL_FROM": "synwevo@example.test", "PUBLIC_URL": ORIGIN, "ISSUER": ORIGIN, "ROLE": "all", "APP_URL": ORIGIN}.items():
        monkeypatch.setenv("SYNWEVO_" + k, v)
    from synwevo_server.app import create_app
    api = create_app(str(tmp_path / "d"), role="all", start_jobs=True)
    ident = api.state.identity
    users = {}
    for key, name in (("ann", "Ann Editor"), ("bob", "Bob Viewer"), ("dan", "Dan Admin"), ("eve", "Eve Outside")):
        u = ident.create_user(f"{key}@example.test", PW, status="active", display_name=name, with_personal_workspace=True)
        ident.store.execute("UPDATE users SET email_verified_at=1 WHERE id=?", (u["id"],))
        users[key] = u
    from synwevo_core.organizations import Organizations
    created = Organizations(ident.store, ident).create(users["dan"]["id"], "Bridge Org", "Bridge Team", "fixture-idem-bridge-org-0001")
    oid, wid = created["organization"]["id"], created["workspace"]["id"]
    for key, role in (("ann", "editor"), ("bob", "viewer")):
        ident.store.execute("INSERT INTO organization_memberships (organization_id,user_id,role,created_at,updated_at,created_by) VALUES (?,?,'member',1,1,?)",
                            (oid, users[key]["id"], users["dan"]["id"]))
        ident.store.execute("INSERT INTO memberships (workspace_id,user_id,role,can_export,created_at,updated_at,created_by) VALUES (?,?,?,1,1,1,?)",
                            (wid, users[key]["id"], role, users["dan"]["id"]))
    clients = {}
    for key in users:
        c = TestClient(api)
        assert c.post("/api/auth/login", json={"email": f"{key}@example.test", "password": PW}).status_code == 200
        c.headers.update({CSRF: c.cookies.get("dw_csrf"), WS: wid})
        clients[key] = c
    try:
        yield {"api": api, "c": clients, "u": users, "wid": wid, "oid": oid, "ident": ident}
    finally:
        for c in clients.values():
            c.close()
        if api.state.jobs:
            api.state.jobs.stop()
        api.state.bridge_worker.stop()


def build(c) -> tuple[str, str]:
    pid = c.post("/api/projects", json={"name": "桥接项目"}).json()["id"]
    assert c.post(f"/api/projects/{pid}/sources", files=[("files", ("客户.xlsx", _xlsx()))]).status_code < 400
    _wait(lambda: all(s["status"] == "ready" for s in c.get(f"/api/projects/{pid}/sources").json()), what="parse")
    t = c.post("/api/templates/upload", files={"file": ("模板.docx", _docx())}, data={"project_id": pid}).json()
    view = c.get(f"/api/templates/{t['id']}/view").json()
    block = next(b for part in view["doc"]["parts"] for b in _walk(part["blocks"]) if "示例客户名称" in b["text"])
    s = block["text"].index("示例客户名称")
    slot = c.post(f"/api/templates/{t['id']}/slots", json={"kind": "text", "part": block["part"], "path": block["path"], "start": s, "end": s + 6, "name": "客户名称"}).json()["created_slot"]
    assert c.patch(f"/api/projects/{pid}", json={"record_rule": {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}}}).status_code == 200
    assert c.put(f"/api/projects/{pid}/mappings", json={"mappings": [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": slot["binding"]}}]}).status_code == 200
    return pid, t["id"]


def cards(c, cid):
    return [m for m in c.get(f"{M}/conversations/{cid}/messages").json()["items"] if m.get("card")]


def test_docs_events_become_cards_in_the_projects_conversation(world):
    C, U, store = world["c"], world["u"], world["ident"].store
    ann, bob, dan, eve = (C[k] for k in ("ann", "bob", "dan", "eve"))
    n = world["api"].state.notifier
    pid, _ = build(ann)
    # ---- a comment before anyone opened the discussion: the row is written with the comment,
    #      then skipped by the worker — an event never creates a conversation
    c0 = ann.post(f"/api/projects/{pid}/comments", json={"body": "还没有讨论群的时候"}).json()
    row = store.one("SELECT * FROM messaging_outbox WHERE event_id=?", (f"comment:{c0['id']}",))
    assert row and row["object_kind"] == "project" and row["object_id"] == pid and row["workspace_id"] == world["wid"]
    row = _wait(lambda: (r := store.one("SELECT * FROM messaging_outbox WHERE event_id=?", (f"comment:{c0['id']}",))) and r["delivered_at"] and r, what="skip")
    assert row["last_error"] == SKIPPED_NO_CONVERSATION
    assert store.one("SELECT 1 AS x FROM object_conversations WHERE object_id=?", (pid,)) is None
    # ---- 「在消息中讨论」: the Docs shell sends the person to the messages origin's
    #      /objects/{app}/project/{pid}; the shell there calls this and lands on /c/{cid}
    r = ann.post(f"{M}/objects/{DOCS_APP_ID}/project/{pid}/conversation")
    assert r.status_code == 200, r.text
    out = r.json()
    cid = out["conversation"]["id"]
    assert out["created"] is True and out["conversation"]["name"] == "桥接项目" and out["conversation"]["workspace_id"] == world["wid"]
    assert {m["user_id"] for m in out["conversation"]["members"]} == {U["ann"]["id"], U["bob"]["id"], U["dan"]["id"]}
    assert out["object"]["href"] == f"{ORIGIN}/p/{pid}"
    # an outsider (no workspace header: with one naming a workspace she is not in, the platform
    # gate itself answers 403 before any route) gets the same nothing as for a missing object
    eve.headers.pop(WS)
    r = eve.post(f"{M}/objects/{DOCS_APP_ID}/project/{pid}/conversation")
    assert r.status_code == 404 and "桥接" not in r.text, r.text
    # the job page's entry resolves through the job to the same conversation
    # ---- a reply through the Docs API → the worker (live in this process) delivers a card
    c1 = bob.post(f"/api/projects/{pid}/comments", json={"body": "取「姓名」列", "parent_id": c0["id"]}).json()
    got = _wait(lambda: cards(ann, cid), what="comment card")
    assert len(got) == 1
    m = got[0]
    assert m["sender_kind"] == "app" and m["sender"] == "Synwevo Docs" and m["body"] == "Synwevo Docs · 新评论"
    assert m["card"] == {"app": DOCS_APP_ID, "app_name": "Synwevo Docs", "title": "桥接项目", "subtitle": "Bob Viewer: 取「姓名」列", "state": "open",
                         "object": {"kind": "comment", "id": c1["id"]}, "actions": [{"label": "查看评论", "href": f"{ORIGIN}/p/{pid}?comment={c1['id']}"}]}
    # the comment produced its own docs.* items (ann, the earlier participant) and no conversation.* item for anyone
    assert [x["kind"] for x in n.list(U["ann"]["id"])["items"]] == ["docs.comment"]
    for key in ("ann", "bob", "dan"):
        assert not [x for x in n.list(U[key]["id"])["items"] if x["kind"].startswith("conversation.")], key
    # ---- a real run through the runner: one card per run, state and counts from the job row
    jid = ann.post(f"/api/projects/{pid}/jobs", json={"formats": ["docx"]}).json()["id"]
    job = _wait(lambda: (j := ann.get(f"/api/jobs/{jid}").json()) and j["status"] not in ("queued", "running", "cancelling") and j, timeout=120, what="job")
    assert job["status"] == "completed", job
    got = _wait(lambda: [m for m in cards(ann, cid) if m["card"].get("object", {}).get("kind") == "job"], what="job card")
    assert len(got) == 1
    assert got[0]["card"] == {"app": DOCS_APP_ID, "app_name": "Synwevo Docs", "title": "桥接项目", "subtitle": "2 份成功，0 份失败", "state": "completed",
                              "object": {"kind": "job", "id": jid}, "actions": [{"label": "查看任务", "href": f"{ORIGIN}/p/{pid}/jobs/{jid}"}]}
    assert store.one("SELECT attempts, last_error FROM messaging_outbox WHERE event_id=?", (f"job:{jid}:1",)) == {"attempts": 1, "last_error": None}
    assert ann.post(f"{M}/objects/{DOCS_APP_ID}/job/{jid}/conversation").json()["conversation"]["id"] == cid
    assert [x["kind"] for x in n.list(U["ann"]["id"])["items"] if x["kind"].startswith("docs.")] == ["docs.job_finished", "docs.comment"]
    assert not [x for x in n.list(U["ann"]["id"])["items"] if x["kind"].startswith("conversation.")]
    # ---- bob leaves the workspace (still in the organization): the same cards, neutral, without
    #      links; then the sweep removes him from the conversation
    world["ident"].remove_member(world["wid"], U["bob"]["id"])
    bob.headers.pop(WS)
    seen = [m["card"] for m in cards(bob, cid)]
    assert len(seen) == 2 and all(c_["title"] == NEUTRAL_TITLE["zh-Hans"] and c_["degraded"] is True and "actions" not in c_ and "subtitle" not in c_ for c_ in seen)
    assert [c_["state"] for c_ in seen] == ["open", "completed"] and [c_["object"] for c_ in seen] == [{"kind": "comment"}, {"kind": "job"}]
    assert world["api"].state.bridge_worker.sweep() == 1
    assert bob.get(f"{M}/conversations/{cid}").status_code == 404
    assert bob.get(f"{M}/conversations/{cid}/messages").status_code == 404
    # dan still reads the full cards
    assert [m["card"]["title"] for m in cards(dan, cid)] == ["桥接项目", "桥接项目"]


def test_a_copy_requests_decision_and_outcome_reach_its_own_conversation(world):
    """The approval flow's two writes — the decision (through the core service the account service
    calls, with the same re-authenticated route in front of it) and the Docs runner's outcome — each
    queue a card for the request's conversation, which the requester opened."""
    C, U, store = world["c"], world["u"], world["ident"].store
    ann, dan = C["ann"], C["dan"]
    personal = world["ident"].ensure_personal_workspace(U["ann"]["id"])["id"]
    ann.headers.update({WS: personal})
    pid, _ = build(ann)
    req = ann.post(f"/api/projects/{pid}/copy-requests", json={"target_workspace_id": world["wid"], "name": "桥接项目（组织副本）", "scope": {"sources": True}}).json()
    assert req["status"] == "pending", req
    out = ann.post(f"{M}/objects/{DOCS_APP_ID}/copy_request/{req['id']}/conversation").json()
    cid = out["conversation"]["id"]
    assert out["created"] is True and {m["user_id"] for m in out["conversation"]["members"]} == {U["ann"]["id"], U["dan"]["id"]}
    assert out["object"]["href"] == f"{ORIGIN}/organizations/{world['oid']}?approval={req['id']}"
    assert C["bob"].get(f"{M}/objects/{DOCS_APP_ID}/copy_request/{req['id']}/conversation").status_code == 404, "a viewer of the target cannot decide, so cannot see it"
    # the decision through the account service's route (a recent admin re-authentication is required there)
    assert dan.post("/api/auth/reauth", json={"password": PW, "scope": "admin"}).status_code == 200
    r = dan.post(f"/api/auth/organizations/{world['oid']}/approvals/{req['id']}/decide", json={"action": "approve", "reason": "team needs it", "version": req["version"]})
    assert r.status_code == 200 and r.json()["status"] == "approved", r.text
    got = _wait(lambda: cards(ann, cid), what="decision card")
    assert got[-1]["card"]["state"] == "approved" and got[-1]["card"]["subtitle"] == "Dan Admin: team needs it" and got[-1]["body"] == "Synwevo Docs · 复制申请已批准"
    assert got[-1]["card"]["actions"] == [{"label": "查看审批", "href": f"{ORIGIN}/organizations/{world['oid']}?approval={req['id']}"}]
    # the Docs runner executes the copy; its outcome is the next card, with the copy's own link first
    world["api"].state.copy_runner.run_once()
    done = _wait(lambda: (x := ann.get(f"/api/projects/{pid}/copy-requests").json()["items"][0]) and x["status"] not in ("approved", "running") and x, what="copy")
    assert done["status"] == "done", done
    got = _wait(lambda: [m for m in cards(ann, cid) if m["card"]["state"] == "done"], what="outcome card")
    assert [a["href"] for a in got[0]["card"]["actions"]] == [f"{ORIGIN}/p/{done['result']['project_id']}", f"{ORIGIN}/organizations/{world['oid']}?approval={req['id']}"]
    assert [r["event_id"] for r in store.all("SELECT event_id FROM messaging_outbox WHERE object_id=? ORDER BY created_at", (req["id"],))] == [f"copy:{req['id']}:approved", f"copy:{req['id']}:done"]
    # the requester's inbox has the approval's own items and nothing from the conversation
    kinds = [x["kind"] for x in world["api"].state.notifier.list(U["ann"]["id"])["items"]]
    assert "copy.approved" in kinds and "copy.done" in kinds and not any(k.startswith("conversation.") for k in kinds)
