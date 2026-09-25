"""Docs tells the person concerned, through the platform inbox, when their generation job ends
and when a file they uploaded fails the security check: real composed server, real session,
real engine, a scanner test double for the refusal. Machine uploads have nobody to tell."""
from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient

from synwevo_core.context import as_tenant
from synwevo_core.malware import ScanResult

PW = "harbor lights guide the evening ferry home"
CSRF, WS = "X-Synwevo-CSRF", "X-Synwevo-Workspace"


class _Infected:
    """A scanner that finds a signature in everything: the refusal path, not ClamAV, is under test."""
    engine = "test double"

    def scan(self, data: bytes, sha256: str | None = None) -> ScanResult:
        return ScanResult(status="infected", signature="Win.Test.EICAR_HDB-1", engine=self.engine, sha256=sha256)


def _xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "客户"
    ws.append(["姓名"]); ws.append(["张三"]); ws.append(["李四"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _docx() -> bytes:
    import docx
    d = docx.Document(); d.add_paragraph("尊敬的示例客户名称：您好。")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


@pytest.fixture()
def world(tmp_path, monkeypatch):
    for k, v in {"UI_AUTH": "session", "COOKIE_SECURE": "0", "DATA_DIR": str(tmp_path / "d"), "MAIL_MODE": "file", "MAIL_DIR": str(tmp_path / "mail"),
                 "MAIL_FROM": "synwevo@example.test", "PUBLIC_URL": "http://testserver", "ROLE": "all", "REGISTRATION": "open",
                 "APP_URL": "http://docs.testserver"}.items():
        monkeypatch.setenv("SYNWEVO_" + k, v)
    from synwevo_server.app import create_app
    api = create_app(str(tmp_path / "d"), role="all", start_jobs=True)
    ident = api.state.identity
    u = ident.create_user("alice@example.test", PW, status="active", display_name="Alice Owner", with_personal_workspace=True, locale="en")
    ident.store.execute("UPDATE users SET email_verified_at=1 WHERE id=?", (u["id"],))
    personal = ident.ensure_personal_workspace(u["id"])["id"]
    c = TestClient(api)
    assert c.post("/api/auth/login", json={"email": "alice@example.test", "password": PW}).status_code == 200
    c.headers.update({CSRF: c.cookies.get("dw_csrf"), WS: personal})
    try:
        yield {"api": api, "c": c, "uid": u["id"], "personal": personal}
    finally:
        c.close()
        if api.state.jobs:
            api.state.jobs.stop()


def _walk(blocks):
    for b in blocks:
        if b["t"] == "p":
            yield b
        else:
            for row in b["rows"]:
                for cell in row["cells"]:
                    yield from _walk(cell["blocks"])


def _wait(fn, timeout=60, what="condition"):
    t0 = time.time()
    while time.time() - t0 < timeout:
        v = fn()
        if v:
            return v
        time.sleep(0.15)
    raise AssertionError(f"timeout waiting for {what}")


def test_the_requester_is_told_when_their_job_ends(world):
    c, n = world["c"], world["api"].state.notifier
    pid = c.post("/api/projects", json={"name": "通知项目"}).json()["id"]
    assert c.post(f"/api/projects/{pid}/sources", files=[("files", ("客户.xlsx", _xlsx()))]).status_code < 400
    _wait(lambda: all(s["status"] == "ready" for s in c.get(f"/api/projects/{pid}/sources").json()), what="parse")
    t = c.post("/api/templates/upload", files={"file": ("模板.docx", _docx())}, data={"project_id": pid}).json()
    view = c.get(f"/api/templates/{t['id']}/view").json()
    block = next(b for part in view["doc"]["parts"] for b in _walk(part["blocks"]) if "示例客户名称" in b["text"])
    s = block["text"].index("示例客户名称")
    slot = c.post(f"/api/templates/{t['id']}/slots", json={"kind": "text", "part": block["part"], "path": block["path"], "start": s, "end": s + 6, "name": "客户名称"}).json()["created_slot"]
    assert c.patch(f"/api/projects/{pid}", json={"record_rule": {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}}}).status_code == 200
    assert c.put(f"/api/projects/{pid}/mappings", json={"mappings": [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": slot["binding"]}}]}).status_code == 200
    jid = c.post(f"/api/projects/{pid}/jobs", json={"formats": ["docx"]}).json()["id"]
    job = _wait(lambda: (j := c.get(f"/api/jobs/{jid}").json()) and j["status"] not in ("queued", "running", "cancelling") and j, timeout=120, what="job")
    assert job["status"] == "completed", job
    items = _wait(lambda: [x for x in n.list(world["uid"], locale="en")["items"] if x["kind"] == "docs.job_finished"], what="job_finished notification")
    assert len(items) == 1, "one notification per run, not per item"
    x = items[0]
    assert x["href"] == f"http://docs.testserver/p/{pid}/jobs/{jid}" and x["object_id"] == jid
    assert "通知项目" in x["title"] and "finished" in x["title"], x["title"]           # the reader's language (en)
    assert "2 succeeded, 0 failed" in x["body"], x["body"]


def test_the_uploader_is_told_when_their_file_is_refused_and_a_machine_upload_tells_nobody(world):
    api, c, n = world["api"], world["c"], world["api"].state.notifier
    app = api.state.app
    pid = c.post("/api/projects", json={"name": "安全检查项目"}).json()["id"]
    app.scanner = _Infected()
    try:
        r = c.post(f"/api/projects/{pid}/sources", files=[("files", ("名单.xlsx", _xlsx()))])
        assert r.status_code < 400, r.text
        sid = r.json()[0]["id"]
        src = _wait(lambda: (x := next(s for s in c.get(f"/api/projects/{pid}/sources").json() if s["id"] == sid)) and x["status"] == "rejected" and x, what="rejection")
        assert "安全检查未通过" in (src.get("error") or "")
        items = _wait(lambda: [x for x in n.list(world["uid"], locale="en")["items"] if x["kind"] == "docs.file_rejected"], what="file_rejected notification")
        assert len(items) == 1 and items[0]["href"] == f"http://docs.testserver/p/{pid}" and items[0]["object_id"] == sid
        assert "名单.xlsx" in items[0]["title"] and "安全检查项目" in items[0]["body"]
        # the same upload again is refused at once (known bad content) and does not notify twice
        again = c.post(f"/api/projects/{pid}/sources", files=[("files", ("名单.xlsx", _xlsx()))])
        assert again.status_code == 400 and again.json()["error"] == "malware_detected", again.text
        assert len([x for x in n.list(world["uid"], locale="en")["items"] if x["kind"] == "docs.file_rejected"]) == 1
        # an upload without a signed-in person behind it (SDK / API key) has nobody to tell
        with as_tenant(world["personal"]):
            row = app.add_source(pid, "机器上传.xlsx", _xlsx() + b"\n")
        _wait(lambda: next(s for s in c.get(f"/api/projects/{pid}/sources").json() if s["id"] == row["id"])["status"] == "rejected", what="machine rejection")
        assert len([x for x in n.list(world["uid"], locale="en")["items"] if x["kind"] == "docs.file_rejected"]) == 1
    finally:
        app.scanner = None
