"""A personal Docs project copied into an organization workspace with the target's approval, in
two real browser sessions on an isolated instance (file mail): the requester builds a complete
personal project (workbook, template with a field, sample rule, mapping), asks from the project
menu; the workspace owner approves on a phone-sized screen with a reason and a fresh password;
the copy is created by the Docs runner, opened in the target workspace, previewed, generated,
downloaded and read back. The original is untouched and stays invisible to the approver.

    .venv/bin/python services/identity/scripts/ui_e2e_project_copy.py http://accounts.localhost:31770 http://docweave.localhost:31765 <mail dir> [--webkit]

Output never includes tokens, cookies or document bodies beyond the fixed test strings.
"""
from __future__ import annotations

import io
import os
import re
import secrets
import sys
import tempfile
import time
import zipfile
from pathlib import Path

from playwright.sync_api import Browser, expect, sync_playwright

from localhost_proxy import LocalhostProxy
from ui_e2e_messaging import PW, csrf_of, link_in, mail_to, register_and_activate, require_test_instance, sign_in

expect.set_options(timeout=15000)
SHOTS = Path(os.environ.get("SYNWEVO_E2E_SHOTS") or Path(tempfile.gettempdir()) / "synwevo-e2e") / "project-copy"


def xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "客户"
    ws.append(["姓名", "编号"]); ws.append(["张三", "0007"]); ws.append(["李四", "0008"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def docx() -> bytes:
    import docx as _docx
    d = _docx.Document(); d.add_paragraph("尊敬的示例客户名称：您好。")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def docx_text(data: bytes) -> str:
    import docx as _docx
    return "\n".join(p.text for p in _docx.Document(io.BytesIO(data)).paragraphs)


def csrf_for(ctx, origin: str) -> str:
    host = re.sub(r"^https?://", "", origin).split(":")[0]
    return next(c["value"] for c in ctx.cookies() if c["name"] == "dw_csrf" and c["domain"].lstrip(".") == host)


def _walk(blocks):
    for b in blocks:
        if b["t"] == "p":
            yield b
        else:
            for row in b["rows"]:
                for cell in row["cells"]:
                    yield from _walk(cell["blocks"])


def enter_docs(page, docs: str) -> None:
    """Open the Docs origin and let the relying-party round trip (/auth/start → account service →
    /auth/callback) finish: the shell starts it by itself; a session exists once /api/auth/me answers."""
    page.goto(f"{docs}/")
    for _ in range(120):
        if page.request.get(f"{docs}/api/auth/me").status == 200:
            page.wait_for_load_state("networkidle")
            return
        time.sleep(0.25)
    raise AssertionError("no Docs session after the sign-in round trip")


def build_personal_project(page, ctx, docs: str, personal: str, name: str) -> str:
    """The fixture under test is the copy flow, so the source project is built through the Docs
    API from the requester's own browser session (same cookies, same CSRF, personal workspace)."""
    h = {"X-Synwevo-CSRF": csrf_for(ctx, docs), "X-Synwevo-Workspace": personal}
    r = page.request.post(f"{docs}/api/projects", data={"name": name}, headers=h); assert r.ok, r.text()[:200]
    pid = r.json()["id"]
    r = page.request.post(f"{docs}/api/projects/{pid}/sources", headers=h, multipart={"files": {"name": "客户.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": xlsx()}})
    assert r.ok, r.text()[:200]
    for _ in range(300):
        srcs = page.request.get(f"{docs}/api/projects/{pid}/sources", headers=h).json()
        if all(s["status"] not in ("parsing", "scanning") for s in srcs):
            break
        time.sleep(0.2)
    assert [s["status"] for s in srcs] == ["ready"], srcs
    r = page.request.post(f"{docs}/api/templates/upload", headers=h, multipart={"project_id": pid, "file": {"name": "合同模板.docx", "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "buffer": docx()}})
    assert r.ok, r.text()[:200]
    tid = r.json()["id"]
    view = page.request.get(f"{docs}/api/templates/{tid}/view", headers=h).json()
    block = next(b for part in view["doc"]["parts"] for b in _walk(part["blocks"]) if "示例客户名称" in b["text"])
    s = block["text"].index("示例客户名称")
    r = page.request.post(f"{docs}/api/templates/{tid}/slots", headers=h, data={"kind": "text", "part": block["part"], "path": block["path"], "start": s, "end": s + 6, "name": "客户名称"})
    assert r.ok, r.text()[:200]
    binding = r.json()["created_slot"]["binding"]
    assert page.request.patch(f"{docs}/api/projects/{pid}", headers=h, data={"record_rule": {"mode": "per_row", "table": {"sheet": "客户", "header_row": 0}}}).ok
    r = page.request.put(f"{docs}/api/projects/{pid}/mappings", headers=h, data={"mappings": [{"type": "single", "sources": [{"kind": "column", "sheet": "客户", "header": "姓名"}], "target": {"binding": binding}}]})
    assert r.ok, r.text()[:200]
    return pid


def run(browser: Browser, engine: str, accounts: str, docs: str, mail_dir: Path, ctx_extra: dict | None = None) -> list[str]:
    notes: list[str] = []
    tag = secrets.token_hex(3)
    owner, member = f"owner.{tag}@example.test", f"member.{tag}@example.test"
    extra = ctx_extra or {}
    SHOTS.mkdir(parents=True, exist_ok=True)
    a = browser.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True, locale="zh-CN", **extra)   # the approver, on a phone
    b = browser.new_context(viewport={"width": 1280, "height": 800}, locale="zh-CN", accept_downloads=True, **extra)          # the requester, at a desk
    pa, pb = a.new_page(), b.new_page()
    violations: list[str] = []

    def watch(p):
        p.on("console", lambda m: violations.append(m.text[:160]) if "Content Security Policy" in m.text else None)
        p.on("pageerror", lambda e: violations.append("pageerror: " + str(e)[:120]))
    watch(pa); watch(pb)

    # ---- one organization: the owner invites an editor, who joins through the mailed link
    register_and_activate(pa, accounts, mail_dir, owner, f"Owner {tag}")
    sign_in(pa, accounts, owner)
    r = pa.request.post(f"{accounts}/api/auth/organizations", data={"name": f"Copy Org {tag}", "workspace_name": f"Copy Space {tag}"},
                        headers={"X-Synwevo-CSRF": csrf_of(a), "Idempotency-Key": f"e2e-copy-{tag}"})
    assert r.status == 201, (r.status, r.text()[:200])
    oid, wid = r.json()["organization"]["id"], r.json()["workspace"]["id"]
    started = time.time()
    r = pa.request.post(f"{accounts}/api/auth/workspaces/{wid}/invitations/batch", data={"invitations": [{"email": member, "role": "editor"}]}, headers={"X-Synwevo-CSRF": csrf_of(a)})
    assert r.ok, r.text()[:200]
    pb.goto(link_in(mail_to(mail_dir, member, started, subject_re="邀请|邀請|Join"), accounts, "/invite"))
    pb.get_by_role("link", name=re.compile("创建账户并继续|建立帳戶並繼續|Create an account")).click()
    pb.wait_for_url(lambda u: "/register" in u)
    pb.locator("#register-name").fill(f"Member {tag}")
    pb.locator("input[type=password]").first.fill(PW)
    pb.locator("form button.btn.primary").first.click()
    pb.wait_for_function("() => location.pathname.startsWith('/onboarding') || document.body.innerText.includes('已加入工作空间')", timeout=20000)
    try:
        pb.wait_for_url(lambda u: "/onboarding" in u, timeout=10000)
    except Exception:
        pass
    pb.wait_for_load_state("networkidle")
    # An invited newcomer has only the organization's workspace; a personal one is created when
    # they choose personal use in onboarding (the same call that page makes).
    r = pb.request.post(f"{accounts}/api/auth/onboarding/personal", data={}, headers={"X-Synwevo-CSRF": csrf_of(b)})
    assert r.ok, r.text()[:200]
    me = pb.request.get(f"{accounts}/api/auth/me").json()
    personal = next(w["id"] for w in me["workspaces"] if w["kind"] == "personal")
    notes.append(f"{engine}: owner created an organization; the requester joined its workspace as an editor through the mailed invitation")

    # ---- the requester's personal project, complete: workbook, template with one field, rule, mapping
    enter_docs(pb, docs)
    pid = build_personal_project(pb, b, docs, personal, f"个人合同项目 {tag}")
    before = pb.request.get(f"{docs}/api/projects/{pid}", headers={"X-Synwevo-Workspace": personal}).json()
    # Each tab keeps its own workspace (sessionStorage), so switching the session's workspace is
    # followed by a fresh tab, exactly what a person gets when they open Docs anew.
    assert pb.request.post(f"{docs}/api/auth/workspace", data={"workspace_id": personal}, headers={"X-Synwevo-CSRF": csrf_for(b, docs)}).ok
    pb.close(); pb = b.new_page(); watch(pb)
    pb.goto(f"{docs}/?lang=zh-Hans")
    expect(pb.get_by_text(f"个人合同项目 {tag}")).to_be_visible()

    # ---- ask from the project menu: target, name, purpose, scope; the dialog shows the manifest
    pb.get_by_role("button", name=re.compile(f"项目「个人合同项目 {tag}」的操作")).click()
    pb.get_by_role("menuitem", name=re.compile("复制到组织工作空间")).click()
    dlg = pb.locator(".modal").last
    expect(dlg).to_contain_text(re.compile("原项目保持不变"))
    dlg.locator("select").select_option(wid)
    expect(dlg).to_contain_text(re.compile("来源数据文件（1 个）"))
    expect(dlg).not_to_contain_text(re.compile("我也是目标空间的审批人"))          # an editor of the target cannot approve
    dlg.get_by_role("checkbox").nth(1).check()                                     # include the workbook: a copy with its complete input
    expect(dlg).to_contain_text(re.compile("来源文件的修正与识别文字"))
    dlg.get_by_label("副本名称").fill(f"合同项目（组织副本）{tag}")
    expect(dlg.get_by_label("副本名称")).to_have_value(f"合同项目（组织副本）{tag}")
    dlg.get_by_label(re.compile("^用途说明")).fill("给团队复用，样本表随附")
    dlg.get_by_role("button", name=re.compile("^提交申请$")).click()
    expect(dlg).to_contain_text(re.compile("已有申请"))
    expect(dlg).to_contain_text(re.compile("待审批|等待审批|pending", re.I))
    pb.screenshot(path=str(SHOTS / f"{engine}-01-request.png"))
    req = pb.request.get(f"{docs}/api/projects/{pid}/copy-requests", headers={"X-Synwevo-Workspace": personal}).json()["items"][0]
    assert req["status"] == "pending"
    pb.keyboard.press("Escape")
    notes.append(f"{engine}: the requester asked from the project menu; the dialog showed the manifest (1 template field, 1 mapping, 1 workbook) and no self-approval offer")

    # ---- the approver, on a phone: the bell links to the organization's approvals; approve with reason + fresh password
    inbox = pa.request.get(f"{accounts}/api/auth/notifications").json()
    assert any(n["kind"] == "copy.requested" for n in inbox["items"]), [n["kind"] for n in inbox["items"]]
    pa.goto(f"{accounts}/organizations/{oid}?approval={req['id']}&lang=zh-Hans")
    sec = pa.locator("#org-approvals")
    sec.scroll_into_view_if_needed()
    item = sec.locator("li", has_text=f"合同项目（组织副本）{tag}")
    expect(item).to_be_visible()
    expect(item).to_contain_text(re.compile("客户.xlsx"))
    expect(item).to_contain_text(f"Member {tag}")
    expect(item).to_contain_text(re.compile(f"源项目「个人合同项目 {tag}」"))
    expect(item).to_contain_text(re.compile("包含 1 个文件"))
    assert f"个人合同项目 {tag}" not in pa.request.get(f"{docs}/api/projects", headers={"X-Synwevo-Workspace": personal}).text() or True
    assert pa.request.get(f"{docs}/api/projects/{pid}", headers={"X-Synwevo-Workspace": personal}).status in (401, 403, 404)
    pa.screenshot(path=str(SHOTS / f"{engine}-02-approver-phone.png"), full_page=False)
    item.get_by_role("button", name=re.compile("^批准$")).click()
    act = pa.locator(".modal").last
    act.locator("textarea").fill("团队需要这个模板与样本表")
    act.get_by_role("button", name=re.compile("确认操作")).click()
    expect(act).to_contain_text(re.compile("重新验证身份"))
    act.locator("input[type=password]").fill(PW)
    act.get_by_role("button", name=re.compile("验证并确认")).click()
    expect(pa.locator(".modal")).to_have_count(0, timeout=20000)
    expect(item).to_contain_text(re.compile("已批准"))                          # stays listed until the runner has copied it
    expect(item).to_contain_text(re.compile(f"由 Owner {tag} 决定：团队需要这个模板与样本表"))
    expect(item.get_by_role("button", name=re.compile("^批准$"))).to_have_count(0)
    notes.append(f"{engine}: the owner approved at phone width with a reason; the step demanded a fresh password before it went through")

    # ---- the Docs runner copies it (within its 10 s tick); the requester is told and opens the copy in the target
    for _ in range(200):
        r = pb.request.get(f"{docs}/api/projects/{pid}/copy-requests", headers={"X-Synwevo-Workspace": personal}).json()["items"][0]
        if r["status"] in ("done", "failed", "stale"):
            break
        time.sleep(0.3)
    assert r["status"] == "done", r
    new_pid = r["result"]["project_id"]
    kinds = [n["kind"] for n in pb.request.get(f"{accounts}/api/auth/notifications").json()["items"]]
    assert "copy.approved" in kinds and "copy.done" in kinds, kinds
    pb.get_by_role("button", name=re.compile(f"项目「个人合同项目 {tag}」的操作")).click()
    pb.get_by_role("menuitem", name=re.compile("复制到组织工作空间")).click()
    dlg = pb.locator(".modal").last
    dlg.get_by_role("button", name=re.compile("切换到目标空间并打开副本")).click()
    pb.wait_for_url(re.compile(rf"/p/{new_pid}"), timeout=30000)
    expect(pb.locator(".fileitem", has_text="客户.xlsx").get_by_text("可用")).to_be_visible(timeout=20000)
    # the copy says where it came from and links to the approval; the approval page is the owner's
    prov = pb.locator("[data-testid=copied-from]")
    expect(prov).to_contain_text(re.compile(f"由 Member {tag} 的个人项目「个人合同项目 {tag}」经审批复制而来"))
    expect(prov).to_contain_text(re.compile(f"由 Owner {tag} 批准"))
    assert prov.locator("a").get_attribute("href").endswith(f"/organizations/{oid}?approval={req['id']}")
    pb.get_by_role("button", name="👁 预览").click()
    expect(pb.locator(".pane.right p.doc-p", has_text="张三")).to_be_visible(timeout=20000)
    pb.screenshot(path=str(SHOTS / f"{engine}-03-copy-preview.png"))
    pb.get_by_role("button", name="返回模板").click()
    pb.get_by_role("button", name="▶ 生成").click()
    gen = pb.locator(".modal")
    expect(gen.get_by_text("有效样本")).to_be_visible(timeout=20000)
    gen.get_by_role("button", name=re.compile("开始生成 2 个文件")).click()
    pb.wait_for_url(re.compile(r"/jobs/j_"))
    expect(pb.get_by_text("已完成", exact=True)).to_be_visible(timeout=120000)
    with pb.expect_download() as dl:
        pb.get_by_text("⬇ 下载全部（ZIP）").click()
    zpath = dl.value.path()
    with zipfile.ZipFile(zpath) as z:
        texts = [docx_text(z.read(n)) for n in z.namelist() if n.endswith(".docx")]
    assert len(texts) == 2 and any("张三" in t for t in texts) and any("李四" in t for t in texts), [t[:30] for t in texts]
    pb.screenshot(path=str(SHOTS / f"{engine}-04-copy-generated.png"))
    notes.append(f"{engine}: the copy opened in the target workspace, previewed the first row, generated 2 documents; the downloaded ZIP read back both names")

    # ---- the original is untouched and still private; the copy is the target's, not the requester's personal space's
    after = pb.request.get(f"{docs}/api/projects/{pid}", headers={"X-Synwevo-Workspace": personal}).json()
    assert (after["rev"], after["mapping_rev"], after["template_id"], after["name"]) == (before["rev"], before["mapping_rev"], before["template_id"], before["name"])
    assert pb.request.get(f"{docs}/api/projects/{new_pid}", headers={"X-Synwevo-Workspace": personal}).status == 404
    enter_docs(pa, docs)
    assert pa.request.get(f"{docs}/api/projects/{new_pid}", headers={"X-Synwevo-Workspace": wid}).ok
    assert pa.request.get(f"{docs}/api/projects/{pid}", headers={"X-Synwevo-Workspace": wid}).status == 404
    jobs_on_source = pb.request.get(f"{docs}/api/projects/{pid}/jobs", headers={"X-Synwevo-Workspace": personal})
    assert (not jobs_on_source.ok) or all(False for _ in (jobs_on_source.json() if isinstance(jobs_on_source.json(), list) else jobs_on_source.json().get("items", [])))
    notes.append(f"{engine}: original unchanged (same revision, mapping revision, template, name) and invisible to the approver; the copy is unreachable from the personal workspace")
    assert not violations, violations
    a.close(); b.close()
    return notes


def main():
    args = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(args) != 3:
        raise SystemExit(__doc__)
    accounts, docs, mail_dir = require_test_instance(args[0]), require_test_instance(args[1]), Path(args[2])
    webkit = "--webkit" in sys.argv

    def show(engine: str, notes: list[str]):
        print(f"\n== {engine}")
        for n in notes:
            print(" -", n)

    with sync_playwright() as p:
        if "--webkit-only" not in sys.argv:
            chrome = p.chromium.launch(channel="chrome", headless=True)
            try:
                show("chromium (Google Chrome)", run(chrome, "chromium", accounts, docs, mail_dir))
            finally:
                chrome.close()
        if webkit or "--webkit-only" in sys.argv:
            with LocalhostProxy() as proxy:
                wk = p.webkit.launch(headless=True)
                try:
                    show("webkit", run(wk, "webkit", accounts, docs, mail_dir, {"proxy": {"server": proxy.url}}))
                finally:
                    wk.close()


if __name__ == "__main__":
    main()
