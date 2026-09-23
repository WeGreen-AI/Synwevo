"""Each generated workbook must carry ITS OWN computed values.

The engine writes computed results into the cached-value slots so other software shows the
right numbers without recalculating.  That makes cache provenance a correctness question:
if a second sample inherited the first sample's cache, every downstream reader would see
confidently wrong figures with no error anywhere.
"""
from __future__ import annotations

import io

import openpyxl
import pytest


def _template_with_formulas() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["A1"], ws["B1"], ws["C1"] = "客户", "数量", "单价"
    ws["A2"], ws["B2"], ws["C2"] = "示例", 0, 0
    ws["D1"], ws["D2"] = "金额", "=B2*C2"
    ws["E1"], ws["E2"] = "含税", "=ROUND(D2*1.13,2)"
    ws["F1"], ws["F2"] = "等级", '=IF(D2>=1000,"大客户","小额")'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _source(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "来源"
    ws.append(["客户", "数量", "单价"])
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS = [("甲公司", 12, 99.5), ("乙公司", 2, 30.0), ("丙公司", 100, 7.25)]
EXPECT = [(1194.0, 1349.22, "大客户"), (60.0, 67.8, "小额"), (725.0, 819.25, "小额")]


def _run(api, tpl_bytes):
    pid = api.project("公式样本隔离")
    api.upload_bytes(pid, "来源.xlsx", _source(ROWS))
    t = api.ok(api.c.post("/api/templates/upload", files={"file": ("模板.xlsx", tpl_bytes)}, data={"project_id": pid}))
    slots = {}
    for name, cell in (("客户", "A2"), ("数量", "B2"), ("单价", "C2")):
        slots[name] = api.slot(t["id"], kind="cell", sheet="订单", cell=cell, name=name)
    api.rule(pid, {"mode": "per_row", "table": {"sheet": "来源", "header_row": 0, "data_start": 1}})
    api.maps(pid, [{"type": "single", "sources": [{"kind": "column", "sheet": "来源", "header": n}],
                    "target": {"binding": slots[n]["binding"]}} for n in ("客户", "数量", "单价")])
    job = api.job(pid, formats=["xlsx"])
    assert job["status"] == "completed", api.items(job["id"])
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(api.c.get(f"/api/jobs/{job['id']}/zip").content))
    return [z.read(n) for n in sorted(n for n in z.namelist() if n.endswith(".xlsx"))]


def test_each_sample_carries_its_own_computed_values(api):
    outs = _run(api, _template_with_formulas())
    assert len(outs) == 3
    seen = []
    for data, (qty, price), (amount, taxed, grade) in zip(outs, [(r[1], r[2]) for r in ROWS], EXPECT):
        cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)["订单"]
        formulas = openpyxl.load_workbook(io.BytesIO(data))["订单"]
        # the formulas themselves must survive — the file stays a live spreadsheet
        assert formulas["D2"].value == "=B2*C2"
        assert str(formulas["F2"].value).startswith("=IF(")
        assert cached["B2"].value == qty and cached["C2"].value == price
        assert cached["D2"].value == pytest.approx(amount), (qty, price, cached["D2"].value)
        assert cached["E2"].value == pytest.approx(taxed)
        assert cached["F2"].value == grade
        seen.append(cached["D2"].value)
    assert len(set(seen)) == 3, f"三份样本的缓存值必须各不相同，得到 {seen}"


def test_a_later_sample_never_inherits_an_earlier_sample_cache(api):
    """The failure this guards against: sample 2 silently showing sample 1's numbers."""
    outs = _run(api, _template_with_formulas())
    values = [openpyxl.load_workbook(io.BytesIO(d), data_only=True)["订单"]["D2"].value for d in outs]
    assert values == [pytest.approx(1194.0), pytest.approx(60.0), pytest.approx(725.0)], values
    # and nothing equals the FIRST sample's value except the first
    assert values[1] != pytest.approx(values[0]) and values[2] != pytest.approx(values[0])


def test_stale_cached_values_in_the_template_are_replaced_not_trusted(api):
    """A template carrying an old cached result must not leak it into the output."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "订单"
    ws["A1"], ws["B1"], ws["C1"] = "客户", "数量", "单价"
    ws["A2"], ws["B2"], ws["C2"] = "示例", 0, 0
    ws["D1"], ws["D2"] = "金额", "=B2*C2"
    buf = io.BytesIO()
    wb.save(buf)

    # inject a bogus cached value straight into the sheet XML
    import re
    import zipfile
    src = buf.getvalue()
    zin = zipfile.ZipFile(io.BytesIO(src))
    out = io.BytesIO()
    POISON = "999999"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename.endswith("sheet1.xml"):
                txt = data.decode("utf-8")
                txt = re.sub(r'(<c r="D2"[^>]*>)(<f>[^<]*</f>)', r'\1\2<v>' + POISON + '</v>', txt)
                data = txt.encode("utf-8")
            zout.writestr(it, data)
    poisoned = out.getvalue()
    assert POISON.encode() in zipfile.ZipFile(io.BytesIO(poisoned)).read("xl/worksheets/sheet1.xml")

    outs = _run(api, poisoned)
    for i, data in enumerate(outs):
        cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)["订单"]
        assert cached["D2"].value != 999999, f"第 {i+1} 份样本仍带着模板里的旧缓存值"
        assert cached["D2"].value == pytest.approx(EXPECT[i][0])
