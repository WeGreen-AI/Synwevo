"""Generation jobs: snapshot inputs, run items in a background thread, verify
outputs, keep state in SQLite so a page refresh or server restart can resume."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import threading
import time
import traceback
import zipfile

from .engine import generate as G
from .engine import records as R
from .engine.plan import build_plan, static_check
from .engine.resolve import SourceDoc
from .services import App, ServiceError, not_found, safe_output_name
from .store import new_id, now

TOKEN_RE = re.compile(r"\{([^{}]{1,40})\}")


# ------------------------------------------------------------------ naming
def naming_tokens(slots: list[dict], columns: list[str] | None = None) -> list[str]:
    names = ["序号", "样本标识", "来源文件名", "工作表", "块号"]
    for s in slots:
        if s["kind"] not in ("repeat", "region") and s["name"] not in names:
            names.append(s["name"])
    for c in columns or []:
        if c and c not in names:
            names.append(c)
    return names


def render_name(pattern: str, record: dict, plan, slots: list[dict], index: int, width: int,
                column_value=None) -> tuple[str, list[str]]:
    """Build an output file name.  Tokens: built-ins, template field names, then source column
    headers (value of the current sample's row; first row of a group)."""
    warns = []
    by_name = {}
    for s in slots:
        fv = plan.values.get(s["id"]) if plan else None
        if fv and fv.get("kind") == "scalar":
            by_name.setdefault(s["name"], fv["value"].text)

    def sub(m):
        key = m.group(1).strip()
        if key == "序号":
            return str(index + 1).zfill(width)
        if key == "样本标识":
            return record.get("ident") or record.get("label", "")
        if key == "来源文件名":
            return (record.get("file_name") or "").rsplit(".", 1)[0]
        if key == "工作表":
            return record.get("sheet", "")
        if key == "块号":
            return str(record["block"] + 1) if record.get("block") is not None else ""
        if key in by_name:
            v = by_name[key]
            if not v:
                warns.append(f"文件名字段「{key}」为空")
            return v
        if column_value is not None:
            v = column_value(record, key)
            if v is not None:
                return v
        warns.append(f"文件名中的「{{{key}}}」不是可用字段")
        return ""
    name = TOKEN_RE.sub(sub, pattern or "{序号}")
    return safe_output_name(name.replace("\n", " ")), warns


def dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        k = n.lower()
        if k not in seen:
            seen[k] = 1
            out.append(n)
            continue
        i = seen[k] + 1
        while f"{n} ({i})".lower() in seen:
            i += 1
        seen[k] = i
        cand = f"{n} ({i})"
        seen[cand.lower()] = 1
        out.append(cand)
    return out


# ------------------------------------------------------------------ project context
class ProjectContext:
    """Everything needed to build plans for a project at a given moment."""

    def __init__(self, app: App, pid: str):
        self.app = app
        self.project = app.get_project(pid)
        self.rule = self.project.get("record_rule") or {}
        tid = self.project.get("template_id")
        self.template = app.get_template(tid) if tid else None
        self.version = app.get_version(tid, self.template["head"]) if self.template else None
        self.slots = (self.version or {}).get("slots") or []
        self.mappings = app.get_mappings(pid)["mappings"]
        self.source_rows = app.project_sources(pid)
        self.all_source_rows = app.project_sources(pid, ready_only=False)
        self.srcs: dict[str, SourceDoc] = {r["id"]: app._doc_from_row(r) for r in self.source_rows}
        self._records = None

    def signature(self) -> str:
        parts = [self.project["rev"], self.project["mapping_rev"], (self.template or {}).get("id"), (self.version or {}).get("version_no"),
                 json.dumps(self.rule, sort_keys=True), [(r["id"], r["version"], r["blob"]) for r in self.source_rows]]
        return hashlib.sha1(json.dumps(parts, default=str).encode()).hexdigest()[:16]

    @property
    def records(self) -> dict:
        if self._records is None:
            if not self.rule.get("mode"):
                raise ServiceError(400, "no_rule", "请先设置样本模式")
            try:
                self._records = R.build_records([self.srcs[r["id"]] for r in self.source_rows], self.rule)
            except ValueError as e:
                raise ServiceError(400, "bad_rule", str(e)) from e
            names = {r["id"]: r["filename"] for r in self.source_rows}
            for rec in self._records["records"]:
                rec["file_name"] = names.get(rec["file_ids"][0] if rec.get("file_ids") else rec["file_id"], "")
        return self._records

    def plan(self, record: dict):
        return build_plan(record, self.rule, self.srcs, self.mappings, self.slots)

    def column_headers(self) -> list[str]:
        return column_headers(self.rule, self.srcs, self.source_rows)

    def column_value(self, record: dict, header: str):
        return record_column_value(self.rule, self.srcs, record, header)

    def working(self) -> bytes:
        return self.app.store.get_blob(self.version["working_blob"])


def column_headers(rule: dict, srcs: dict, rows: list[dict]) -> list[str]:
    if rule.get("mode") not in ("per_row", "per_group") or not rule.get("table"):
        return []
    from .engine.resolve import ResolveError, table_view
    for r in rows:
        src = srcs.get(r["id"])
        if src is None or src.workbook_json is None:
            continue
        try:
            return [h for h in table_view(src, rule["table"]).headers if h]
        except ResolveError:
            continue
    return []


_TV_CACHE: dict = {}


def record_column_value(rule: dict, srcs: dict, record: dict, header: str):
    if rule.get("mode") not in ("per_row", "per_group"):
        return None
    from .engine import tabular
    from .engine.resolve import ResolveError, table_view
    fid = record.get("file_id")
    src = srcs.get(fid)
    if src is None or src.workbook_json is None:
        return None
    key = (src.file_id, repr(rule.get("table")))
    hit = _TV_CACHE.get(key)
    tv = hit[1] if hit and hit[0] is src else None      # identity check: never reuse a view of another file version
    if tv is None:
        try:
            tv = table_view(src, rule["table"])
        except ResolveError:
            return None
        if len(_TV_CACHE) > 64:
            _TV_CACHE.clear()
        _TV_CACHE[key] = (src, tv)
    idx = tv.header_index().get(tabular.normalize_header(header))
    if not idx or len(idx) > 1:
        return None
    row = record.get("row")
    if row is None and record.get("rows"):
        row = record["rows"][0][2]
    if row is None:
        return None
    return tv.sheet.cell(row, idx[0]).text


# ------------------------------------------------------------------ verification
def verify_output(fmt: str, data: bytes, slots: list[dict], plan) -> dict:
    """Reopen the produced file and check that filled values are present."""
    res = {"reopened": False, "checks": 0, "missing": [], "engine": None}
    try:
        if fmt == "docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            from .engine.docx_render import render_docx
            rend = render_docx(data)
            text = "\n".join(_blocks_text(p["blocks"]) for p in rend["parts"])
            res.update(reopened=True, engine="python-docx " + getattr(docx, "__version__", ""), paragraphs=len(d.paragraphs))
        elif fmt == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data))
            text = "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)
            res.update(reopened=True, engine="openpyxl", sheets=len(wb.worksheets))
        elif fmt == "pdf":
            from .engine import pdf_engine as pe
            text = pe.pdf_text(data)
            res.update(reopened=True, engine={"pymupdf": "PyMuPDF", "pdfium": "PDFium"}[pe.ENGINE], pages=pe.page_count(data))
        elif fmt == "pptx":
            import pptx
            prs = pptx.Presentation(io.BytesIO(data))
            from .engine.pptx_engine import render_pptx
            rend = render_pptx(data)
            text = "\n".join(p["text"] for s in rend["slides"] for sh in s["shapes"] for p in sh.get("paras", [])) + "\n" + \
                "\n".join(p["text"] for s in rend["slides"] for sh in s["shapes"] for row in sh.get("rows", []) for c in row for p in c["paras"])
            res.update(reopened=True, engine="python-pptx", slides=len(prs.slides))
        elif fmt == "csv":
            text = data.decode("utf-8-sig")
            res.update(reopened=True, engine="csv")
        elif fmt in ("png", "jpg"):
            from PIL import Image
            im = Image.open(io.BytesIO(data))
            im.verify()
            im = Image.open(io.BytesIO(data))
            res.update(reopened=True, engine="Pillow", size=list(im.size),
                       note="图片中的文字无法直接读取：已在栅格化前的页面中逐值核对（未写入则该样本失败）")
            return res
        else:
            return res
    except Exception as e:  # noqa: BLE001
        res["error"] = f"输出文件无法重新打开：{e}"
        return res
    norm = _squash(text)
    for s in slots:
        fv = plan.values.get(s["id"]) if plan else None
        if not fv:
            continue
        vals = []
        if fv["kind"] == "scalar":
            vals = [fv["value"].text]
        elif fv["kind"] == "rows":
            vals = [v.text for row in fv["rows"][:50] for v in row.values()]
        elif fv["kind"] == "matrix":
            vals = [c.text for row in fv["rows"] for c in row]
        for v in vals:
            if not v or (fmt == "xlsx" and _numeric(v)):
                continue
            res["checks"] += 1
            if _squash(v) not in norm:
                res["missing"].append({"slot": s["name"], "value": v[:80]})
    return res


def _numeric(s: str) -> bool:
    return bool(re.fullmatch(r"-?[\d,]+(\.\d+)?", s.strip()))


def _squash(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def _blocks_text(blocks) -> str:
    out = []
    for b in blocks:
        if b["t"] == "p":
            out.append(b["text"])
        else:
            for row in b["rows"]:
                for c in row["cells"]:
                    out.append(_blocks_text(c["blocks"]))
    return "\n".join(out)


# ------------------------------------------------------------------ job service
class JobService:
    def __init__(self, app: App):
        self.app = app
        self.store = app.store
        self._wake = threading.Event()
        self._stop = False
        self._thread = threading.Thread(target=self._loop, name="job-runner", daemon=True)
        self.recover()
        self._thread.start()

    def recover(self) -> None:
        """Jobs interrupted by a restart: mark running items back to pending and resume."""
        for j in self.store.all("SELECT id FROM jobs WHERE status IN ('running', 'queued')"):
            self.store.execute("UPDATE job_items SET status='pending' WHERE job_id=? AND status='running'", (j["id"],))
            self.store.update("jobs", {"id": j["id"]}, {"status": "queued"})

    def stop(self):
        self._stop = True
        self._wake.set()

    # -------------------------------------------------------------- summary before generation
    def summary(self, pid: str, formats: list[str] | None = None, naming: str | None = None) -> dict:
        ctx = ProjectContext(self.app, pid)
        if not ctx.template:
            raise ServiceError(400, "no_template", "请先选择或创建模板")
        recs = ctx.records
        valid, invalid, warnings = [], [], 0
        issues_by_code: dict[str, int] = {}
        sample_issues = []
        for rec in recs["records"]:
            pl = ctx.plan(rec)
            warnings += sum(1 for i in pl.issues if i["level"] == "warning")
            if pl.blocking:
                invalid.append(rec)
                for i in pl.blocking:
                    issues_by_code[i["code"]] = issues_by_code.get(i["code"], 0) + 1
                if len(sample_issues) < 50:
                    sample_issues.append({"record": rec["label"], "index": rec["index"], "issues": pl.blocking[:3]})
            else:
                valid.append(rec)
        exports = self.app._exports_for(ctx.template["format"])
        fmts = formats or [G.native_format(ctx.template["format"])]
        fmt_info = []
        for f in fmts:
            ex = next((e for e in exports if e["format"] == f), None)
            if ex is None:
                fmt_info.append({"format": f, "available": False, "reason": f"不支持导出 {f.upper()}"})
            else:
                fmt_info.append(ex)
        static = static_check(ctx.mappings, ctx.slots, ctx.rule)
        pattern = naming or (ctx.project.get("output") or {}).get("naming") or "{序号}"
        names = []
        width = max(3, len(str(len(recs["records"]))))
        for rec in valid[:5]:
            nm, _ = render_name(pattern, rec, ctx.plan(rec), ctx.slots, rec["index"], width, ctx.column_value)
            names.append(nm)
        n_ok_formats = len([f for f in fmt_info if f.get("available")])
        return {"input_files": len(ctx.all_source_rows), "ready_files": len(ctx.source_rows),
                "file_errors": recs["file_errors"], "records": len(recs["records"]), "valid": len(valid), "invalid": len(invalid),
                "excluded": recs["excluded"], "warnings": warnings, "record_warnings": recs["warnings"],
                "issues_by_code": issues_by_code, "invalid_samples": sample_issues, "formats": fmt_info,
                "expected_outputs": len(valid) * n_ok_formats, "expected_if_block": len(recs["records"]) * n_ok_formats if not invalid else 0,
                "naming_examples": names, "static_issues": static, "tokens": naming_tokens(ctx.slots, ctx.column_headers()),
                "signature": ctx.signature(), "template_version": ctx.version["version_no"]}

    # -------------------------------------------------------------- create
    def create(self, pid: str, options: dict) -> dict:
        req = options.get("request_id")
        if req:
            existing = self.store.one("SELECT id FROM jobs WHERE request_id=?", (req,))
            if existing:
                return self.get(existing["id"])
        ctx = ProjectContext(self.app, pid)
        if not ctx.template:
            raise ServiceError(400, "no_template", "请先选择或创建模板")
        formats = options.get("formats") or [G.native_format(ctx.template["format"])]
        exports = {e["format"]: e for e in self.app._exports_for(ctx.template["format"])}
        for f in formats:
            if f not in exports or not exports[f].get("available"):
                reason = (exports.get(f) or {}).get("reason") or f"不支持导出 {f.upper()}"
                raise ServiceError(400, "format_unavailable", f"无法导出 {f.upper()}：{reason}")
        policy = options.get("error_policy", "block")
        recs = ctx.records
        if not recs["records"]:
            raise ServiceError(400, "no_records", "没有识别到任何样本，无法生成")
        from .engine.limits import LIMITS
        if len(recs["records"]) > LIMITS.max_records_per_job:
            raise ServiceError(400, "too_many", f"样本数 {len(recs['records'])} 超过单次任务上限 {LIMITS.max_records_per_job}")
        plans = [(rec, ctx.plan(rec)) for rec in recs["records"]]
        invalid = [(rec, pl) for rec, pl in plans if pl.blocking]
        if invalid and policy == "block":
            raise ServiceError(409, "has_errors", f"{len(invalid)} 个样本存在错误，已按「有错误则不开始」停止。请修正或选择「跳过错误样本」",
                               {"invalid": [{"record": r["label"], "issues": p.blocking[:3]} for r, p in invalid[:50]]})
        pattern = options.get("naming") or (ctx.project.get("output") or {}).get("naming") or "{序号}"
        width = max(3, len(str(len(plans))))
        raw_names = [render_name(pattern, rec, pl, ctx.slots, rec["index"], width, ctx.column_value)[0] for rec, pl in plans]
        names = dedupe(raw_names)
        jid = new_id("j")
        snapshot = {
            "project_id": pid, "signature": ctx.signature(), "rule": ctx.rule, "mappings": ctx.mappings,
            "template": {"id": ctx.template["id"], "name": ctx.template["name"], "format": ctx.template["format"],
                         "version_no": ctx.version["version_no"], "working_blob": ctx.version["working_blob"], "slots": ctx.slots,
                         "fill_options": self.app.fill_options(ctx.template)},
            "sources": [{k: r.get(k) for k in ("id", "filename", "version", "blob", "working_blob", "parsed_blob", "format", "family",
                                              "ocr", "corrections")} for r in ctx.source_rows],
        }
        t = now()
        self.store.insert("jobs", {"id": jid, "project_id": pid, "status": "queued", "created_at": t, "snapshot": snapshot,
                                   "options": {"formats": formats, "error_policy": policy, "naming": pattern,
                                               "keep_fields": bool(options.get("keep_fields")),
                                               "layout": options.get("layout") if options.get("layout") in ("off", "warn", "fit")
                                               else ((ctx.project.get("output") or {}).get("layout") or "off")},
                                   "counts": {}, "request_id": req})
        for (rec, pl), name in zip(plans, names):
            status = "skipped" if pl.blocking else "pending"
            self.store.insert("job_items", {"id": new_id("i"), "job_id": jid, "seq": rec["index"], "record_key": rec["key"],
                                            "record": rec, "status": status, "issues": pl.blocking if pl.blocking else [],
                                            "file_base": name, "label": rec["label"]})
        self._update_counts(jid)
        self._wake.set()
        return self.get(jid)

    # -------------------------------------------------------------- queries
    def get(self, jid: str) -> dict:
        j = self.store.one("SELECT id, project_id, status, created_at, started_at, finished_at, options, counts, error, cancel_requested, request_id FROM jobs WHERE id=?", (jid,))
        if not j:
            not_found("任务")
        snap = self.store.one("SELECT snapshot FROM jobs WHERE id=?", (jid,))["snapshot"]
        j["template"] = {k: snap["template"][k] for k in ("id", "name", "format", "version_no")}
        j["source_count"] = len(snap["sources"])
        return j

    def items(self, jid: str, status: str | None = None, offset: int = 0, limit: int = 500) -> list[dict]:
        sql = "SELECT id, seq, record_key, label, status, issues, attempts, file_base, started_at, finished_at FROM job_items WHERE job_id=?"
        args: tuple = (jid,)
        if status:
            sql += " AND status=?"
            args += (status,)
        sql += " ORDER BY seq LIMIT ? OFFSET ?"
        args += (limit, offset)
        rows = self.store.all(sql, args)
        arts = self.store.all("SELECT id, item_id, format, filename, size, verification FROM artifacts WHERE job_id=?", (jid,))
        by_item: dict[str, list] = {}
        for a in arts:
            by_item.setdefault(a["item_id"], []).append(a)
        for r in rows:
            r["artifacts"] = by_item.get(r["id"], [])
        return rows

    def list_for_project(self, pid: str) -> list[dict]:
        return self.store.all("SELECT id, status, created_at, started_at, finished_at, counts, options FROM jobs WHERE project_id=? ORDER BY created_at DESC LIMIT 50", (pid,))

    def cancel(self, jid: str) -> dict:
        j = self.get(jid)
        if j["status"] in ("completed", "cancelled", "failed", "completed_with_errors"):
            return j
        self.store.update("jobs", {"id": jid}, {"cancel_requested": 1})
        if j["status"] == "queued":
            self._finish_cancel(jid)
        return self.get(jid)

    def _finish_cancel(self, jid: str) -> None:
        self.store.execute("UPDATE job_items SET status='cancelled' WHERE job_id=? AND status IN ('pending','running')", (jid,))
        self.store.update("jobs", {"id": jid}, {"status": "cancelled", "finished_at": now()})
        self._update_counts(jid)

    def retry(self, jid: str, include_skipped: bool = False) -> dict:
        self.get(jid)
        statuses = ("failed", "cancelled") + (("skipped",) if include_skipped else ())
        n = self.store.execute(f"UPDATE job_items SET status='pending' WHERE job_id=? AND status IN ({','.join('?' * len(statuses))})",
                               (jid,) + statuses)
        if not n:
            raise ServiceError(400, "nothing_to_retry", "没有失败或已取消的项目可重试")
        self.store.update("jobs", {"id": jid}, {"status": "queued", "cancel_requested": 0, "finished_at": None})
        self._update_counts(jid)
        self._wake.set()
        return self.get(jid)

    def _update_counts(self, jid: str) -> dict:
        rows = self.store.all("SELECT status, COUNT(*) AS n FROM job_items WHERE job_id=? GROUP BY status", (jid,))
        c = {r["status"]: r["n"] for r in rows}
        c["total"] = sum(c.values())
        c["files"] = self.store.one("SELECT COUNT(*) AS n FROM artifacts WHERE job_id=?", (jid,))["n"]
        self.store.update("jobs", {"id": jid}, {"counts": c})
        return c

    # -------------------------------------------------------------- runner
    def _loop(self) -> None:
        while not self._stop:
            job = self.store.one("SELECT id FROM jobs WHERE status='queued' ORDER BY created_at LIMIT 1")
            if not job:
                self._wake.wait(1.0)
                self._wake.clear()
                continue
            try:
                self._run(job["id"])
            except Exception as e:  # noqa: BLE001
                self.store.update("jobs", {"id": job["id"]}, {"status": "failed", "error": f"{e}", "finished_at": now()})

    def _run(self, jid: str) -> None:
        j = self.store.one("SELECT * FROM jobs WHERE id=?", (jid,))
        self.store.update("jobs", {"id": jid}, {"status": "running", "started_at": j.get("started_at") or now()})
        snap, opts = j["snapshot"], j["options"]
        srcs = {}
        for s in snap["sources"]:
            srcs[s["id"]] = self.app._doc_from_row({**s, "family": s["family"]})
        tpl = snap["template"]
        work = self.store.get_blob(tpl["working_blob"])
        slots = tpl["slots"]
        native = G.native_format(tpl["format"])
        layout_mode = opts.get("layout") or "off"
        tpl_pages = G.template_pages(tpl["format"], work) if layout_mode != "off" else None
        last_count = time.time()
        while True:
            item = self.store.one("SELECT * FROM job_items WHERE job_id=? AND status='pending' ORDER BY seq LIMIT 1", (jid,))
            if item is None:
                break
            if self.store.one("SELECT cancel_requested FROM jobs WHERE id=?", (jid,))["cancel_requested"]:
                self._finish_cancel(jid)
                return
            self.store.update("job_items", {"id": item["id"]}, {"status": "running", "started_at": now(), "attempts": item["attempts"] + 1})
            try:
                plan = build_plan(item["record"], snap["rule"], srcs, snap["mappings"], slots)
                if plan.blocking:
                    self.store.update("job_items", {"id": item["id"]}, {"status": "failed", "issues": plan.blocking, "finished_at": now()})
                    continue
                out, issues, lay = G.fill_checked(tpl["format"], work, slots, plan, {**(tpl.get("fill_options") or {}), "keep_fields": opts.get("keep_fields")},
                                                  layout_mode, tpl_pages)
                blocking = [i for i in issues if i.get("level") == "error"]
                if blocking or not out:
                    self.store.update("job_items", {"id": item["id"]}, {"status": "failed", "issues": blocking or issues, "finished_at": now()})
                    continue
                produced = []
                failed_fmt = None
                for f in opts["formats"]:
                    try:
                        data, conv_issues = G.convert_output(out, native, f)
                        issues.extend(conv_issues)
                    except Exception as e:  # noqa: BLE001
                        failed_fmt = (f, str(e))
                        break
                    ver = verify_output(f, data, slots if f == native else [], plan)
                    if lay and f == native:
                        ver["layout"] = lay
                    if not ver.get("reopened"):
                        failed_fmt = (f, ver.get("error") or "输出校验失败")
                        break
                    if ver["missing"]:
                        failed_fmt = (f, "输出文件中找不到已填写的值：" + "、".join(f"{m['slot']}={m['value']}" for m in ver["missing"][:3]))
                        break
                    produced.append((f, data, ver))
                if failed_fmt:
                    self.store.update("job_items", {"id": item["id"]}, {"status": "failed", "finished_at": now(),
                                                                        "issues": issues + [{"level": "error", "code": "export_failed", "message": f"{failed_fmt[0].upper()}：{failed_fmt[1]}"}]})
                    continue
                attempt = item["attempts"] + 1
                for f, data, ver in produced:
                    blob = self.store.put_blob(data)
                    self.store.insert("artifacts", {"id": new_id("a"), "job_id": jid, "item_id": item["id"], "format": f,
                                                    "filename": f"{item['file_base']}.{f}", "blob": blob, "size": len(data),
                                                    "trace": {"record": {k: item["record"].get(k) for k in ("key", "label", "file_id", "file_name", "where", "rows")},
                                                              "slots": plan.trace, "template_version": tpl["version_no"]},
                                                    "verification": ver, "created_at": now(), "attempt": attempt})
                self.store.update("job_items", {"id": item["id"]}, {"status": "success", "issues": [i for i in issues if i.get("level") != "info"][:20],
                                                                    "finished_at": now()})
            except Exception as e:  # noqa: BLE001
                self.store.update("job_items", {"id": item["id"]}, {"status": "failed", "finished_at": now(),
                                                                    "issues": [{"level": "error", "code": "exception", "message": f"{type(e).__name__}: {e}",
                                                                                "trace": traceback.format_exc()[-800:]}]})
            if time.time() - last_count > 0.3:
                self._update_counts(jid)
                last_count = time.time()
        c = self._update_counts(jid)
        status = "completed" if not c.get("failed") and not c.get("skipped") else "completed_with_errors"
        self.store.update("jobs", {"id": jid}, {"status": status, "finished_at": now()})

    # -------------------------------------------------------------- downloads
    def artifact(self, aid: str) -> tuple[dict, bytes]:
        a = self.store.one("SELECT * FROM artifacts WHERE id=?", (aid,))
        if not a:
            not_found("输出文件")
        return a, self.store.get_blob(a["blob"])

    def zip(self, jid: str, fmt: str | None = None) -> bytes:
        self.get(jid)
        arts = self.store.all("SELECT a.*, i.seq FROM artifacts a JOIN job_items i ON a.item_id=i.id WHERE a.job_id=? ORDER BY i.seq", (jid,))
        buf = io.BytesIO()
        used = set()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for a in arts:
                if fmt and a["format"] != fmt:
                    continue
                name = a["filename"]
                folder = f"{a['format']}/" if len({x['format'] for x in arts}) > 1 else ""
                full = folder + name
                k = 2
                while full.lower() in used:
                    base, _, ext = name.rpartition(".")
                    full = f"{folder}{base} ({k}).{ext}"
                    k += 1
                used.add(full.lower())
                z.writestr(full, self.store.get_blob(a["blob"]))
            z.writestr("生成清单.csv", self.report_csv(jid))
        return buf.getvalue()

    def report_csv(self, jid: str) -> bytes:
        items = self.items(jid, limit=1_000_000)
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(["序号", "样本", "状态", "输出文件", "来源文件", "来源位置", "问题说明"])
        status_cn = {"success": "成功", "failed": "失败", "skipped": "已跳过（生成前校验错误）", "cancelled": "已取消", "pending": "等待中", "running": "进行中"}
        for it in items:
            files = "；".join(a["filename"] for a in it["artifacts"])
            issues = "；".join(i.get("message", "") for i in (it["issues"] or []) if i.get("level") in ("error", "warning"))
            rec = self.store.one("SELECT record FROM job_items WHERE id=?", (it["id"],))["record"]
            w.writerow([it["seq"] + 1, it["label"], status_cn.get(it["status"], it["status"]), files, rec.get("file_name", ""),
                        rec.get("where", ""), issues])
        return ("﻿" + out.getvalue()).encode("utf-8")
