"""Render worksheets to paged JSON grids (values + basic styles + merges).

Used for XLSX templates, XLSX outputs (preview) and tabular sources.  Large
sheets are delivered in row windows; every row remains reachable.
"""
from __future__ import annotations

import io

from . import tabular
from .numfmt import format_value

THEME_COLORS = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]
INDEXED = ["000000", "FFFFFF", "FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF", "000000", "FFFFFF", "FF0000",
           "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF", "800000", "008000", "000080", "808000", "800080", "008080",
           "C0C0C0", "808080", "9999FF", "993366", "FFFFCC", "CCFFFF", "660066", "FF8080", "0066CC", "CCCCFF", "000080",
           "FF00FF", "FFFF00", "00FFFF", "800080", "800000", "008080", "0000FF", "00CCFF", "CCFFFF", "CCFFCC", "FFFF99",
           "99CCFF", "FF99CC", "CC99FF", "FFCC99", "3366FF", "33CCCC", "99CC00", "FFCC00", "FF9900", "FF6600", "666699",
           "969696", "003366", "339966", "003300", "333300", "993300", "993366", "333399", "333333"]


def _color(c) -> str | None:
    if c is None:
        return None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str):
            rgb = c.rgb[-6:]
            if c.rgb.startswith("00") and len(c.rgb) == 8 and rgb == "000000":
                return "#000000"
            return "#" + rgb
        if c.type == "theme" and c.theme is not None and 0 <= c.theme < len(THEME_COLORS):
            base = THEME_COLORS[c.theme]
            tint = c.tint or 0
            return "#" + _tint(base, tint)
        if c.type == "indexed" and c.indexed is not None and c.indexed < len(INDEXED):
            return "#" + INDEXED[c.indexed]
    except Exception:
        return None
    return None


def _tint(hexrgb: str, tint: float) -> str:
    r, g, b = int(hexrgb[0:2], 16), int(hexrgb[2:4], 16), int(hexrgb[4:6], 16)

    def t(x):
        return int(x + (255 - x) * tint) if tint > 0 else int(x * (1 + tint))
    return f"{max(0, min(255, t(r))):02X}{max(0, min(255, t(g))):02X}{max(0, min(255, t(b))):02X}"


class XlsxRenderer:
    """openpyxl-based renderer (read only; never saves)."""

    def __init__(self, data: bytes):
        import openpyxl
        self.wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, rich_text=False)
        # cached results (written by DocWeave's formula engine / Excel) shown instead of formula text
        self.vals = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
        self._styles: dict = {}
        self._style_list: list = []

    def sheets(self) -> list[dict]:
        out = []
        for ws in self.wb.worksheets:
            out.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column, "hidden": ws.sheet_state != "visible"})
        return out

    def _style_id(self, cell) -> int:
        key = tuple(cell._style)
        sid = self._styles.get(key)
        if sid is not None:
            return sid
        st = {}
        f = cell.font
        if f is not None:
            if f.b:
                st["b"] = 1
            if f.i:
                st["i"] = 1
            if f.u:
                st["u"] = 1
            if f.strike:
                st["s"] = 1
            if f.sz:
                st["sz"] = float(f.sz)
            col = _color(f.color)
            if col and col.upper() != "#000000":
                st["color"] = col
            if f.name:
                st["font"] = f.name
        fill = cell.fill
        if fill is not None and fill.fill_type == "solid":
            col = _color(fill.fgColor)
            if col:
                st["bg"] = col
        b = cell.border
        if b is not None:
            sides = {}
            for side in ("left", "right", "top", "bottom"):
                sd = getattr(b, side)
                if sd is not None and sd.style:
                    sides[side[0]] = sd.style
            if sides:
                st["bd"] = sides
        a = cell.alignment
        if a is not None:
            if a.horizontal:
                st["ha"] = a.horizontal
            if a.vertical:
                st["va"] = a.vertical
            if a.wrap_text:
                st["wrap"] = 1
        sid = len(self._style_list)
        self._style_list.append(st)
        self._styles[key] = sid
        return sid

    def window(self, sheet: str | None, offset: int = 0, limit: int = 200, max_cols: int = 1024) -> dict:
        ws = self.wb[sheet] if sheet else self.wb.worksheets[0]
        max_row, max_col = ws.max_row, min(ws.max_column, max_cols)
        rows = []
        end = min(max_row, offset + limit)
        for r in range(offset + 1, end + 1):
            row = []
            for c in range(1, max_col + 1):
                cell = ws.cell(r, c)
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    cached = self.vals[ws.title].cell(r, c).value
                    if cached is None:
                        item = {"v": v, "f": 1, "fx": v}
                    else:
                        text, _ = format_value(cached, cell.number_format)
                        item = {"v": text, "fx": v}
                else:
                    text, _ = format_value(v, cell.number_format)
                    item = {"v": text}
                    if cell.number_format and cell.number_format != "General" and v is not None:
                        item["nf"] = cell.number_format
                if cell.has_style:
                    item["s"] = self._style_id(cell)
                row.append(item)
            rows.append(row)
        widths = []
        for c in range(1, max_col + 1):
            letter = tabular.col_letter(c - 1)
            dim = ws.column_dimensions.get(letter)
            w = dim.width if dim is not None and dim.width else None
            hidden = bool(dim is not None and dim.hidden)
            widths.append({"w": round((w or 8.43) * 7 + 5), "hidden": hidden})
        heights = {}
        for r in range(offset + 1, end + 1):
            dim = ws.row_dimensions.get(r)
            if dim is not None and (dim.height or dim.hidden):
                heights[r - 1] = {"h": round((dim.height or 15) * 4 / 3), "hidden": bool(dim.hidden)}
        merges = []
        for mr in ws.merged_cells.ranges:
            merges.append([mr.min_row - 1, mr.min_col - 1, mr.max_row - 1, mr.max_col - 1])
        return {"sheet": ws.title, "offset": offset, "total_rows": max_row, "total_cols": ws.max_column, "cols_shown": max_col,
                "rows": rows, "widths": widths, "heights": heights, "merges": merges, "styles": self._style_list,
                "sheets": self.sheets()}


def workbook_window(wb_json: dict, sheet: str | None, offset: int = 0, limit: int = 200, max_cols: int = 1024) -> dict:
    """Window over a parsed source workbook (values only)."""
    wb = tabular.Workbook.from_json(wb_json)
    sh = wb.sheet(name=sheet) if sheet else (wb.sheets[0] if wb.sheets else None)
    if sh is None:
        raise ValueError(f"没有工作表「{sheet}」")
    ncols = min(sh.n_cols, max_cols)
    rows = []
    for r in range(offset, min(sh.n_rows, offset + limit)):
        row = []
        for c in range(ncols):
            v = sh.cell(r, c)
            item = {"v": v.text}
            if v.t not in ("str", "empty"):
                item["t"] = v.t
            row.append(item)
        rows.append(row)
    return {"sheet": sh.name, "offset": offset, "total_rows": sh.n_rows, "total_cols": sh.n_cols, "cols_shown": ncols,
            "rows": rows, "widths": [], "heights": {}, "merges": sh.merged, "styles": [],
            "sheets": [{"name": s.name, "rows": s.n_rows, "cols": s.n_cols, "hidden": s.hidden} for s in wb.sheets],
            "notes": sh.notes}
